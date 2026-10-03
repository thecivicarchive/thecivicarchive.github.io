"""
ballot/wy_place_votes.py - how each Wyoming place voted in past partisan general elections, so a page can show the
record of a county or a district without anyone labelling a candidate. The Wyoming twin of ballot/mn_place_votes.py;
the output has the same shape (the Democratic count is "dem" here, as in the Dakotas' and Montana's files, where
Minnesota's is "dfl").

    python ballot/wy_place_votes.py               reads (or downloads) the files, writes ballot/lean/wy_place_votes.json
    python ballot/wy_place_votes.py --refresh     asks for everything again even when the cached copies are there
    python ballot/wy_place_votes.py --out x.json  writes somewhere else; --cache DIR keeps the downloads somewhere else
    python ballot/wy_place_votes.py --selftest    the arithmetic and the readers on made-up tables; downloads nothing

What it is, and is not
----------------------
For President and U.S. Senator in 2024 and in 2020, and Governor in 2022 (Wyoming elects its Governor in the years
between presidential elections, and elected no senator in 2022): the votes for the Democratic ticket, the Republican
ticket, every other candidate and all write-ins together, and the total, for every one of the 23 counties and the nine
judicial districts, and for 2022 and 2024 for every state House and Senate district whose precincts voted in that
district alone (see below). It is how the people of a place voted then, on the lines in force at that election. It is
not a prediction, it says nothing about any candidate on a later ballot or about any voter, and it turns no
nonpartisan office into a partisan one. No database is opened for writing; ballot_local_2026.sqlite is opened
read-only at the end, only to count how many of our places are covered.

Where the numbers come from
---------------------------
All from the Wyoming Secretary of State's results pages (sos.wyo.gov/Elections/ElectionResults.aspx, one page for each
general election), which answer a plain script. Each election has a zip of the official workbooks:

  - "County General PbP" (precinct by precinct; one sheet for each county in 2022 and 2024, one workbook for each
    county in 2020): a row for each precinct, a column for each candidate of each contest, then Write-Ins, Overvotes
    and Undervotes; a dash where a precinct does not vote in a contest; a Total row. THE COUNTY FIGURES ARE THESE ROWS
    ADDED UP. The columns kept are the precinct, the contest and the candidate heading with its party letter; the
    over and under votes are left out of every sum (one county typed a full stop in such a cell).
  - "Results Summaries" (2022 and 2024) and "General Statewide Candidates Summary" (2020): the statewide contests county
    by county with a Total row, and (2022 and 2024) every state House and Senate contest county by county. They are
    the control for the counties and for the precinct rows of the legislative contests.
  - A state House or Senate district is the precincts that voted in its contest. A precinct is a county's numbered
    voting area, and its number does not name a district; it votes in the House contest of the district it lies in.
    Where a precinct lies in two or three House districts, its rows carry a vote in each of those contests, and its
    votes for President, Senator or Governor are given whole: they cannot be divided between the districts without
    an estimate, so a district with such a precinct is not given for that year, and the file names the precinct. A
    Senate district is up every four years (the even ones in 2024, the odd ones in 2022 and 2026), so a precinct's
    Senate district is read from the year its district was on the ballot: the same precinct of the same county in the
    other general election's workbook, or in the 2026 primary's (a check only, where the kit has it). A precinct found
    in two Senate districts, or whose other half the record does not show, leaves those districts out for that year.
  - A judicial district is whole counties, listed in section 5-3-101, Wyoming Statutes; a circuit court has the same
    boundaries (section 5-9-102). The section is read from the Legislature's own copy of title 5 (wyoleg.gov) and must
    equal the copy typed below.

Not given: cities and towns (a Wyoming precinct has a number, and the results do not say which city or town it lies
in); county commissioner districts, school, college, hospital, conservation and other districts (the results do not
name them for a precinct); House and Senate districts for 2020 (the plan of 2012, 60 House and 30 Senate districts;
today's 62 and 31 were first used in 2022).

The control
-----------
Nothing is written unless all of this holds: in each workbook every precinct is named once in its county and every
precinct votes in each statewide contest; for every contest the precincts add up, column by column, to each county's
Total row and to that county's row in the statewide summary, and the counties to the summary's Total row; for
President and U.S. Senator the totals equal the Clerk of the U.S. House of Representatives' "Statistics of the ...
Election" (clerk.house.gov, the Wyoming page, read from the PDF at each run; its under and over votes are left out, as
here); for 2022 and 2024 every House and Senate contest's precinct rows add up to its county rows in the summary,
every one of the 62 House districts has a contest each year and every one of the 31 Senate districts in one of the two
years; every kind of place that covers the state adds up to the statewide sum, and the districts given, with the
precincts of the districts left out, do too. The statewide figures this loader was checked against on 2026-10-02 are
typed below (CHECKED); the loader says so if a workbook reads differently later.
"""

import argparse
import collections
import datetime as dt
import io
import json
import os
import pathlib
import re
import sqlite3
import sys
import zipfile

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from states import net  # noqa: E402
from ballot import mn_place_votes as M  # noqa: E402  the shared small things: fingerprints, days, the file writer

Stop = M.Stop
METHOD = "1.0"                 # change how anything is added up or matched, and this goes up
CACHE = os.path.join(HERE, "states_cache", "wy_local")
OUT = os.path.join(HERE, "ballot", "lean", "wy_place_votes.json")
DB = os.path.join(HERE, "ballot_local_2026.sqlite")
PRIMARY_2026 = os.path.join(HERE, "ballot_cache", "wy", "wy_2026_primary_results.zip")   # the federal loader's copy; a check only
PRIMARY_BOOK = "2026 Primary County PbP Results - OFFICIAL.xlsx"
STATE_FIPS = "56"
FEW = 20
SIDES = ("dem", "rep", "other", "total")
KEEP_DAYS = 3650               # certified results of past elections do not change; --refresh asks again
HOUSE, SENATE = 62, 31
DISTRICT_YEARS = (2022, 2024)  # the elections held on today's House and Senate districts

AGENCY = "Wyoming Secretary of State, Elections Division"
RESULTS_PAGE = "https://sos.wyo.gov/Elections/ElectionResults.aspx"
DOCS = "https://sos.wyo.gov/Elections/Docs"
ELECTIONS = {
    2024: {"date": "2024-11-05", "page": f"{DOCS}/2024/2024GeneralResults.aspx",
           "zip": {"url": f"{DOCS}/2024/Results/General/2024_Wyoming_General_Results.zip", "file": "wy_2024_general_results.zip",
                   "title": "2024 General Election Results, the official workbooks (zip)"},
           "books": "2024 County General PbP - OFFICIAL.xlsx", "summary": "2024 General Results Summaries - OFFICIAL.xlsx",
           "candidates": "Statewide Candidates", "house": "Statewide House", "senate": "Statewide Senate Even"},
    2022: {"date": "2022-11-08", "page": f"{DOCS}/2022/2022GeneralResults.aspx",
           "zip": {"url": f"{DOCS}/2022/Results/General/2022_Wyoming_General_Results.zip", "file": "wy_2022_general_results.zip",
                   "title": "2022 General Election Results, the official workbooks (zip)"},
           "books": "2022_General_County_PbP.xlsx", "summary": "2022_General_Summaries_Results.xlsx",
           "candidates": "Statewide Candidates", "house": "Statewide House", "senate": "Statewide Senate"},
    2020: {"date": "2020-11-03", "page": f"{DOCS}/2020/2020GeneralResults.aspx",
           "zip": {"url": f"{DOCS}/2020/Results/General/2020_Wyoming_General_Results.zip", "file": "wy_2020_general_results.zip",
                   "title": "2020 General Election Results, the official workbooks (zip)"},
           "books": re.compile(r"^2020_(.+)_County_General_PbP\.xlsx$"), "summary": "2020_General_Statewide_Candidates_Summary.xlsx",
           "candidates": None, "house": None, "senate": None},
}
STATUTE = {"id": "wy-stat-5-3-101", "cite": "Section 5-3-101, Wyoming Statutes (Judicial districts enumerated), in title 5 as the Legislature publishes it",
           "url": "https://wyoleg.gov/statutes/compress/title05.pdf", "file": "wy_title05.pdf"}

# The 23 counties in the order of their FIPS codes (56001, 56003, ...).
COUNTIES = ("Albany|Big Horn|Campbell|Carbon|Converse|Crook|Fremont|Goshen|Hot Springs|Johnson|Laramie|Lincoln|Natrona|Niobrara|Park|"
            "Platte|Sheridan|Sublette|Sweetwater|Teton|Uinta|Washakie|Weston").split("|")


def bare(name):
    """A county's name for comparing only: capitals, letters alone ('Big Horn', 'Big_Horn', 'BIG HORN COUNTY')."""
    return re.sub(r"COUNTY$", "", re.sub(r"[^A-Z]", "", (name or "").upper()))


FIPS = {bare(name): f"{STATE_FIPS}{2 * i + 1:03d}" for i, name in enumerate(COUNTIES)}
COUNTY_NAME = {FIPS[bare(name)]: f"{name} County" for name in COUNTIES}

# Section 5-3-101 as this loader was checked against it on 2026-10-02 (read again at each run; it must be equal).
JUDICIAL = {1: ["Laramie"], 2: ["Albany", "Carbon"], 3: ["Sweetwater", "Lincoln", "Uinta"], 4: ["Johnson", "Sheridan"],
            5: ["Big Horn", "Hot Springs", "Park", "Washakie"], 6: ["Campbell", "Crook", "Weston"], 7: ["Natrona"],
            8: ["Converse", "Platte", "Goshen", "Niobrara"], 9: ["Fremont", "Sublette", "Teton"]}
ROMAN = {"i": 1, "ii": 2, "iii": 3, "iv": 4, "v": 5, "vi": 6, "vii": 7, "viii": 8, "ix": 9}
ORDINAL_WORDS = {"first": 1, "second": 2, "third": 3, "fourth": 4, "fifth": 5, "sixth": 6, "seventh": 7, "eighth": 8, "ninth": 9}

# The contests. "contest" is the workbooks' heading; "dem" and "rep" are the tickets, whose first person's surname must
# be in the heading the workbooks print for the candidate of that party, or the loader stops.
CONTESTS = [
    {"id": "2024-president", "year": 2024, "office": "President and Vice President of the United States", "contest": "United States President",
     "dem": "Harris and Walz", "rep": "Trump and Vance", "clerk": "FOR PRESIDENTIAL ELECTORS"},
    {"id": "2024-us-senate", "year": 2024, "office": "United States Senator", "contest": "United States Senator",
     "dem": "Scott D. Morrow", "rep": "John Barrasso", "clerk": "FOR UNITED STATES SENATOR"},
    {"id": "2022-governor", "year": 2022, "office": "Governor", "contest": "Governor",
     "dem": "Theresa A. Livingston", "rep": "Mark Gordon", "clerk": None},
    {"id": "2020-president", "year": 2020, "office": "President and Vice President of the United States", "contest": "United States President",
     "dem": "Biden and Harris", "rep": "Trump and Pence", "clerk": "FOR PRESIDENTIAL ELECTORS"},
    {"id": "2020-us-senate", "year": 2020, "office": "United States Senator", "contest": "United States Senator",
     "dem": "Merav Ben David", "rep": "Cynthia M. Lummis", "clerk": "FOR UNITED STATES SENATOR"},
]
PARTY_WORDS = {"D": "Democratic", "R": "Republican", "L": "Libertarian", "C": "Constitution", "I": "independent"}
NOT_CANDIDATES = ("Overvotes", "Undervotes")
WRITE_INS = "Write-Ins"

# The statewide totals this loader was checked against on 2026-10-02 (the precincts added up, the summaries' Total rows
# and, for President and U.S. Senator, the Clerk of the House all read these).
CHECKED = {
    "2024-president": {"dem": 69527, "rep": 192633, "other": 6888, "total": 269048},
    "2024-us-senate": {"dem": 63727, "rep": 198418, "other": 2017, "total": 264162},
    "2022-governor": {"dem": 30686, "rep": 143696, "other": 19618, "total": 194000},
    "2020-president": {"dem": 73491, "rep": 193559, "other": 9715, "total": 276765},
    "2020-us-senate": {"dem": 72766, "rep": 198100, "other": 1071, "total": 271937},
}
CLERK = {year: {"id": f"clerk-statistics-{year}", "kind": "official federal compilation", "agency": "Office of the Clerk, U.S. House of Representatives",
                "title": f"Statistics of the Presidential and Congressional Election from Official Sources for the Election of {day}",
                "url": f"https://clerk.house.gov/member_info/electionInfo/{year}/statistics{year}.pdf", "file": f"clerk_statistics{year}.pdf"}
         for year, day in ((2024, "November 5, 2024"), (2020, "November 3, 2020"))}

# The kinds of place: (kind, first election year given, what it is, what its key is)
KINDS = [
    ("county", 2020, "Counties", "five-digit county FIPS code, as sl_places kind county and the jurisdiction_id of our county races"),
    ("house", 2022, "State House districts of the plan first used in 2022", "district number, as the district of our House races"),
    ("senate", 2022, "State Senate districts of the plan first used in 2022", "district number, as the district of our Senate races"),
    ("judicial", 2020, "Judicial districts, as whole counties (a circuit court's district is the same)",
     "WY-JD and the district number (WY-JD4), as the jurisdiction_id of our district and circuit court races"),
]
KIND_NAMES = [k[0] for k in KINDS]
REDRAWN = ("The 2020 election was held on the House and Senate districts of 2012 (60 and 30 of them); the districts drawn after the "
           "2020 census were first used in 2022, so 2020 is not given.")
WHAT = ("How each Wyoming place voted in five contests of the 2020, 2022 and 2024 general elections: the votes for the Democratic ticket, "
        "the Republican ticket, every other candidate and all write-ins together and the total, added up from the Wyoming Secretary of "
        "State's official precinct results.")
NOTE = ("What this is: how the people of a place voted in that election, in the Secretary of State's official precinct results, added up "
        "here by county and by judicial district and, for 2022 and 2024, by state House and Senate district, on the lines in force at that "
        "election. A district is given only where every precinct of it voted in that district alone: where a precinct lies in two "
        "districts the results give its votes whole, and they are never divided by estimate. What this is not: it is not a prediction of "
        "any election; it says nothing about any candidate on a later ballot or about any voter; a nonpartisan office stays nonpartisan; "
        "and a place is not its lines for ever: where a district was drawn again, the figures are for the lines of that year. Wyoming "
        "elected no senator in 2022, and its Governor only in 2022. The tickets are named, as the Secretary's own results name them, only "
        "to say which election this was. In a place with very few voters the split would come close to saying how particular people "
        "voted; those contests are listed in the place's too_few, and a page should leave the split out.")
HOW_TO_READ = ("Each place's votes are four counts for each contest: dem (the Democratic ticket), rep (the Republican ticket), other (every "
               "other candidate and all write-ins, together) and total (the votes for candidates and write-ins; over and under votes are "
               "left out). Minnesota's file calls the first count dfl. A contest a place does not have was not counted on its lines: House "
               "and Senate districts are given for 2022 and 2024 only, and only where no precinct of the district lies in another district "
               "too.")
TOO_FEW = (f"A place's too_few lists the contests in which it had fewer than {FEW} votes, or in which every vote went the same way. There the "
           "split comes close to saying how particular people voted, so a page should leave it out. The counts are the official record all the "
           "same, and are kept here so that the sums can be checked.")
NOT_GIVEN = ("Cities and towns: a Wyoming precinct has a number, and the results do not say which city or town it lies in. County "
             "commissioner districts, school, college, hospital, conservation and other districts: the results do not name them for a "
             "precinct.")


# ---------------------------------------------------------------- the files

def fetch(spec, cache, refresh, say, magic):
    path = os.path.join(cache, spec["file"])
    try:
        net.download(spec["url"], path, 0 if refresh else KEEP_DAYS, tries=3, say=say)
    except Exception as e:  # noqa: BLE001
        raise Stop(f"    {spec.get('title') or spec['file']}: could not be fetched ({e}) and no copy is on disk. Wait a few minutes and run this again.")
    with open(path, "rb") as fh:
        if not fh.read(8).startswith(magic):
            raise Stop(f"    {spec['file']}: what came back is not the file (it does not begin {magic!r}); stopping")
    return path


def words(x):
    return " ".join(str(x).split()) if x is not None else ""


def label(cell):
    """A precinct as its row names it. Some counties typed '1-2' and the spreadsheet stored it as a date (2 January):
    it is written back as month-day, the way it was typed."""
    if isinstance(cell, (dt.datetime, dt.date)):
        return f"{cell.month}-{cell.day}"
    return words(cell)


def count(cell, what):
    """A vote count, or None where the precinct or county does not vote in the contest (a dash, '*' or blank; one
    county once typed a full stop in an over-vote cell, which is no count either, and those columns are never added)."""
    if cell is None or (isinstance(cell, str) and cell.strip() in ("", "-", "*", ".")):
        return None
    if isinstance(cell, bool):
        raise Stop(f"    {what}: a cell is {cell!r}, not a count of votes; stopping")
    if isinstance(cell, (int, float)) and cell == int(cell) and cell >= 0:
        return int(cell)
    if isinstance(cell, str) and re.fullmatch(r"\d[\d,]*", cell.strip()):
        return int(cell.strip().replace(",", ""))
    raise Stop(f"    {what}: a cell is {cell!r}, not a count of votes; stopping")


def present(c, _what):
    """For a sheet read only for which contests a precinct voted in (a primary's, whose sheets also carry percentages): 1
    for any cell that is not a dash or blank."""
    return None if c is None or (isinstance(c, str) and c.strip() in ("", "-", "*", ".")) else 1


def table(rows, what, parties=("Republican", "Democratic", "Libertarian", "Constitution", "Nonpartisan"), unique=True, cell=count):
    """One results sheet as {"rows": {label: {(contest, heading): count or None}}, "total": {...}, "keys": [...]}. The
    heading row is the first with a Write-Ins cell; the contest row is the row above it (above a row of party names, in
    a primary's sheets); a contest runs from its heading to the next. Rows run to the Total row; a "Precincts Continue
    on Next Page" line is passed over."""
    rows = [tuple(r) for r in rows]
    h = next((i for i, r in enumerate(rows) if any(words(c) == WRITE_INS for c in r)), None)
    if h is None or h == 0:
        raise Stop(f"    {what}: no heading row with Write-Ins; the sheet is not laid out as this loader reads it; stopping")
    c = h - 1
    while c > 0 and [words(x) for x in rows[c] if words(x)] and all(words(x) in parties for x in rows[c] if words(x)):
        c -= 1
    con, head = rows[c], rows[h]
    keys, cur = [], None
    for j in range(1, len(head)):
        if j < len(con) and words(con[j]):
            cur = re.sub(r",\s*Continued$", "", words(con[j]), flags=re.I)
        if not words(head[j]):
            continue
        if cur is None:
            raise Stop(f"    {what}: a column heading stands under no contest; stopping")
        keys.append((j, (cur, words(head[j]))))
    if len({k for _j, k in keys}) != len(keys):
        if unique:
            raise Stop(f"    {what}: a contest has the same heading twice; stopping")
        keys = [(j, (k[0], f"{k[1]} #{j}")) for j, k in keys]     # a primary's parties side by side: Write-Ins under each
    out, total = {}, None
    for r in rows[h + 1:]:
        if not r or r[0] is None or not label(r[0]):
            continue
        name = label(r[0])
        if name.startswith("Precincts Continue"):
            continue
        vals = {k: cell(r[j] if j < len(r) else None, f"{what}, {name}") for j, k in keys}
        if name == "Total":
            total = vals
            break
        if name in out:
            raise Stop(f"    {what}: {name} is named twice; stopping")
        out[name] = vals
    if total is None:
        raise Stop(f"    {what}: no Total row; stopping")
    return {"rows": out, "total": total, "keys": [k for _j, k in keys]}


def read_sheets(data, what, names=None):
    """{sheet name: rows} of a workbook given as bytes."""
    import openpyxl
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")                     # "Cannot parse header or footer": the print layout, not the data
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return {s: list(wb[s].iter_rows(values_only=True)) for s in (names or wb.sheetnames)}
    except KeyError as e:
        raise Stop(f"    {what}: no sheet {e}; stopping")
    finally:
        wb.close()


def read_year(path, year):
    """One general election's zip: ({county fips: precinct table}, {"candidates": table, "house": table, "senate": table})."""
    e = ELECTIONS[year]
    z = zipfile.ZipFile(path)
    books = {}
    if isinstance(e["books"], str):
        for sheet, rows in read_sheets(z.read(e["books"]), f"{year} {e['books']}").items():
            code = FIPS.get(bare(sheet))
            if code is None or code in books:
                raise Stop(f"    {year} precinct workbook: a sheet is named {sheet!r}, not one county once; stopping")
            books[code] = table(rows, f"{year} precinct workbook, {sheet}")
    else:
        for name in z.namelist():
            m = e["books"].match(name)
            if m:
                code = FIPS.get(bare(m.group(1)))
                if code is None or code in books:
                    raise Stop(f"    {year}: a workbook is named {name!r}, not one county once; stopping")
                rows = next(iter(read_sheets(z.read(name), f"{year} {name}").values()))
                books[code] = table(rows, f"{year} precinct workbook, {COUNTY_NAME[code]}")
    if len(books) != len(COUNTIES):
        raise Stop(f"    {year}: precinct results for {len(books)} counties, not {len(COUNTIES)}; stopping")
    sheets = read_sheets(z.read(e["summary"]), f"{year} {e['summary']}")
    summary = {}
    for key in ("candidates", "house", "senate"):
        if key == "candidates" or e[key]:
            name = e[key] or next(iter(sheets))
            if name not in sheets:
                raise Stop(f"    {year} summary workbook: no sheet {name!r}; stopping")
            summary[key] = table(sheets[name], f"{year} summary, {name}")
            summary[key]["sheet"] = name
    return books, summary


def read_primary(path):
    """{(county fips, precinct): {Senate districts}} from the 2026 primary's precinct workbook: which Senate contest each
    precinct voted in, and nothing else (no votes, no names)."""
    z = zipfile.ZipFile(path)
    out = {}
    for sheet, rows in read_sheets(z.read(PRIMARY_BOOK), "2026 primary precinct workbook").items():
        code = FIPS.get(bare(sheet))
        if code is None:
            raise Stop(f"    2026 primary precinct workbook: a sheet is named {sheet!r}; stopping")
        for p, vals in table(rows, f"2026 primary precinct workbook, {sheet}", unique=False, cell=present)["rows"].items():
            out[(code, p)] = (districts_of(vals, "House"), districts_of(vals, "Senate"))
    return out


def primary_whole(primary, senates=SENATE):
    """Whether the primary had a contest in every odd-numbered Senate district, so that a precinct with none is in none of them."""
    return bool(primary) and set().union(*(v[1] for v in primary.values())) >= set(range(1, senates + 1, 2))


def read_statute(path):
    """What section 5-3-101 says: {district: [county names]}, from title 5's text."""
    from ballot import pdftext
    text = " ".join(t for _p, _y, t in pdftext.lines(path))
    return statute_districts(text)


def statute_districts(text):
    text = " ".join(text.split())
    a = text.find("5-3-101. Judicial districts enumerated")
    b = text.find("5-3-102.", a)
    if a < 0 or b < 0:
        raise ValueError("title 5 has no section 5-3-101 as this loader reads it")
    part = text[a:b]
    out = {}
    for m in re.finditer(r"\(([ivx]+)\)\s*(.*?)\s+(?:is|are)\s+the\s+(\w+)\s+judicial\s+district", part):
        n = ORDINAL_WORDS.get(m.group(3).lower())
        if n is None or ROMAN.get(m.group(1)) != n:
            raise ValueError(f"section 5-3-101: paragraph ({m.group(1)}) does not read as a numbered district")
        rest, names = m.group(2), []
        for name in sorted(COUNTIES, key=len, reverse=True):
            if re.search(rf"\b{re.escape(name)}\b", rest):
                names.append((rest.index(name), name))
                rest = re.sub(rf"\b{re.escape(name)}\b", " " * len(name), rest, count=1)
        if re.sub(r"\b(?:the|counties|county|of|and|The)\b|[,\s]", "", rest, flags=re.I):
            raise ValueError(f"section 5-3-101, district {n}: words that are not counties")
        out[n] = [x for _i, x in sorted(names)]
    return out


def statute(cache, refresh, say):
    """Section 5-3-101, read from the Legislature's title 5 and required to equal the typed copy; the typed copy stands
    in, and the file says so, when the title cannot be read."""
    s = STATUTE
    path = os.path.join(cache, s["file"])
    rec = {"id": s["id"], "kind": "statute, as the Legislature publishes it", "agency": "Wyoming Legislature, Wyoming Statutes", "title": s["cite"],
           "url": s["url"]}
    try:
        net.download(s["url"], path, 0 if refresh else 30, tries=3, say=say)
        rec.update(fetched=M._day(path), sha256=M._sha_file(path))
    except Exception as e:  # noqa: BLE001
        say(f"      {s['cite']}: could not be read ({e}); the copy typed into this loader on 2026-10-02 is used")
        rec["unread"] = f"could not be read on {M._now()}; the copy typed into the loader on 2026-10-02 is used"
        return rec
    try:
        got = read_statute(path)
    except ValueError as e:
        raise Stop(f"    {e}; stopping")
    if {k: sorted(v) for k, v in got.items()} != {k: sorted(v) for k, v in JUDICIAL.items()}:
        raise Stop("    section 5-3-101, Wyoming Statutes, no longer reads as the copy this loader was checked against; nothing is written "
                   "until a person has read the change")
    return rec


# ---------------------------------------------------------------- adding up

def _add(into, v):
    for i in range(len(v)):
        into[i] += v[i]


def districts_of(vals, chamber):
    """The House or Senate districts whose contest a precinct row has a count in."""
    out = set()
    for (contest, _head), v in vals.items():
        m = re.fullmatch(rf"{chamber} District (\d+)", contest)
        if m and v is not None:
            out.add(int(m.group(1)))
    return out


def party_of(head):
    m = re.search(r"\(([A-Z]{1,3})\)\s*$", head)
    return m.group(1) if m else None


def surname(ticket):
    return ticket.split(" and ")[0].split()[-1].upper()


def sides_of(c, heads, what):
    """Which headings of a contest are the two tickets, which are the rest (candidates and write-ins), and which are
    left out (over and under votes)."""
    side = {}
    for s, letter in (("dem", "D"), ("rep", "R")):
        fits = [h for h in heads if party_of(h) == letter]
        if len(fits) != 1 or surname(c[s]) not in fits[0].upper():
            raise Stop(f"    {what}, {c['contest']}: the {PARTY_WORDS[letter]} candidate is not the one ({c[s]}) this loader was checked against; stopping")
        side[s] = fits[0]
    others = [h for h in heads if h not in side.values() and h not in NOT_CANDIDATES]
    if WRITE_INS not in others or any(h != WRITE_INS and not party_of(h) for h in others):
        raise Stop(f"    {what}, {c['contest']}: a heading is neither a candidate with a party letter nor Write-Ins; stopping")
    return side, others


def four(vals, contest, side, others):
    dem, rep = vals[(contest, side["dem"])], vals[(contest, side["rep"])]
    other = sum(vals[(contest, h)] for h in others)
    return [dem, rep, other, dem + rep + other]


def few(rec):
    hold = [cid for cid, v in rec["votes"].items() if v["total"] < FEW or max(v["dem"], v["rep"], v["other"]) == v["total"]]
    if hold:
        rec["too_few"] = hold
    return rec


def plain(key):
    """A (contest, heading) key with the heading's party letter set aside."""
    return (key[0], re.sub(r"\s*\([A-Z]{1,3}\)\s*$", "", key[1]))


def compare(mine, theirs, keys):
    """The keys whose counts differ ('-' and 0 are not the same: a county that has no part in a contest has a dash)."""
    return [k for k in keys if mine.get(k) != theirs.get(k)]


def build(years, clerk, contests=CONTESTS, judicial=None, primary=None, houses=HOUSE, senates=SENATE, expect=True, sources=None):
    """Everything added up and checked. years: {year: (books, summary)}; clerk: {contest id: {dem, rep, other, total,
    where}}; primary: {(county, precinct): {Senate districts}} or None. Returns (the file, what failed)."""
    judicial = JUDICIAL if judicial is None else judicial
    failed, ctl, state = [], {}, {}
    pack = lambda v: dict(zip(SIDES, v))                                                               # noqa: E731
    votes = {k: collections.defaultdict(dict) for k in KIND_NAMES}
    ids = [c["id"] for c in contests]
    first = {k[0]: k[1] for k in KINDS}
    district_years = [y for y in DISTRICT_YEARS if y in years]
    given = {k: [c["id"] for c in contests if c["year"] >= first[k] and (k not in ("house", "senate") or c["year"] in district_years)]
             for k in KIND_NAMES}

    # ---- the statewide contests, county by county
    for c in contests:
        cid, year = c["id"], c["year"]
        books, summary = years[year]
        heads = None
        for code, t in books.items():
            mine = [h for (k, h) in t["keys"] if k == c["contest"]]
            if heads is None:
                heads = mine
            elif sorted(mine) != sorted(heads):
                failed.append(f"{cid}: {COUNTY_NAME[code]}'s sheet heads the contest's columns {mine}, another county's {heads}")
        side, others = sides_of(c, heads or [], f"{year} precinct workbook")
        keys = [(c["contest"], h) for h in heads if h not in NOT_CANDIDATES]
        sums = collections.defaultdict(lambda: [0] * 4)
        county_bad, n_pre = [], 0
        for code, t in books.items():
            col = {k: 0 for k in keys}
            for p, vals in t["rows"].items():
                if any(vals.get(k) is None for k in keys):
                    failed.append(f"{cid}: precinct {p} of {COUNTY_NAME[code]} has no count in the contest")
                    continue
                n_pre += 1
                for k in keys:
                    col[k] += vals[k]
                _add(sums[code], four(vals, c["contest"], side, others))
            srow = summary["candidates"]["rows"].get(next((n for n in summary["candidates"]["rows"] if FIPS.get(bare(n)) == code), None), {})
            if compare(col, t["total"], keys) or compare(col, srow, keys):
                county_bad.append(code)
        state[cid] = [sum(v[i] for v in sums.values()) for i in range(4)]
        stotal = summary["candidates"]["total"]
        scounty = {FIPS.get(bare(n)) for n in summary["candidates"]["rows"]}
        total_ok = not compare({k: sum(sum(t["rows"][p][k] or 0 for p in t["rows"]) for t in books.values()) for k in keys}, stotal, keys)
        rec = ctl[cid] = {"sum_of_precincts": pack(state[cid]), "precincts": n_pre,
                          "official": pack(four(stotal, c["contest"], side, others)) if all(stotal.get(k) is not None for k in keys) else None,
                          "counties": {"compared": len(books), "equal": len(books) - len(county_bad)},
                          "equal": not county_bad and total_ok and scounty == set(books),
                          "source": f"wy-sos-results-{year}", "where": f"the {summary['candidates']['sheet']} sheet of {ELECTIONS[year]['summary']}, "
                          f"the columns headed {c['contest']}; and each county's Total row in the precinct workbook",
                          "read": f"from the zip fetched {(sources or {}).get(year, M._now())}"}
        if county_bad:
            rec["counties"]["differ"] = [COUNTY_NAME[x] for x in county_bad]
            failed.append(f"{cid}: the precincts do not add up to the Total row or the summary in {', '.join(COUNTY_NAME[x] for x in county_bad)}")
        if not total_ok or scounty != set(books):
            failed.append(f"{cid}: the summary's counties or its Total row do not equal the precincts added up")
        typed = CHECKED.get(cid) if expect else None
        if typed and typed != pack(state[cid]):
            failed.append(f"{cid}: the precincts add up to {pack(state[cid])}; this loader was checked against {typed}")
        if clerk.get(cid):
            k = clerk[cid]
            same = all(k[s] == state[cid][i] for i, s in enumerate(SIDES))
            rec["clerk"] = {"official": {s: k[s] for s in SIDES}, "equal": same, "where": k.get("where")}
            if not same:
                failed.append(f"{cid}: the precincts add up to {pack(state[cid])} and the Clerk of the House prints { {s: k[s] for s in SIDES} }")
        for code, v in sums.items():
            votes["county"][code][cid] = v
        for d, names in judicial.items():
            v = [0] * 4
            for name in names:
                _add(v, sums.get(FIPS[bare(name)], [0] * 4))
            votes["judicial"][f"WY-JD{d}"][cid] = v
        c["_side"], c["_others"] = side, others

    # ---- the legislative contests against the summary: the precinct rows that say which district a precinct is in
    leg_ctl = {}
    p_hd, p_sd = {}, {}
    for year in district_years:
        books, summary = years[year]
        p_hd[year], p_sd[year] = {}, {}
        seen = {"House": set(), "Senate": set()}
        bad = []
        for code, t in books.items():
            for p, vals in t["rows"].items():
                p_hd[year][(code, p)] = districts_of(vals, "House")
                p_sd[year][(code, p)] = districts_of(vals, "Senate")
            for chamber, sheet in (("House", "house"), ("Senate", "senate")):
                mine = sorted({k for k in t["keys"] if re.fullmatch(rf"{chamber} District \d+", k[0]) and k[1] not in NOT_CANDIDATES})
                seen[chamber] |= {int(k[0].split()[-1]) for k in mine}
                srow = summary[sheet]["rows"].get(next((n for n in summary[sheet]["rows"] if FIPS.get(bare(n)) == code), None), {})
                col = {k: sum(v[k] for v in t["rows"].values() if v.get(k) is not None) if any(v.get(k) is not None for v in t["rows"].values()) else None
                       for k in mine}
                # (a county's sheet can print a legislative candidate without the party letter the summary prints: compared by name)
                diff = compare({plain(k): v for k, v in col.items()}, {plain(k): v for k, v in srow.items()}, [plain(k) for k in mine])
                if diff:
                    bad.append(f"{COUNTY_NAME[code]}, {diff[0][0]}")
        no_hd = [k for k, v in p_hd[year].items() if not v]
        leg_ctl[str(year)] = {"house_contests": len(seen["House"]), "senate_contests": len(seen["Senate"]),
                              "precinct_rows_equal_the_summary": not bad}
        if bad:
            failed.append(f"{year}: the precinct rows of the legislative contests do not add up to the summary in {'; '.join(bad[:6])}")
        if sorted(seen["House"]) != list(range(1, houses + 1)):
            failed.append(f"{year}: House contests for {len(seen['House'])} districts, not {houses}")
        if no_hd:
            failed.append(f"{year}: {len(no_hd)} precincts voted in no House contest (first: {no_hd[0][1]}, {COUNTY_NAME[no_hd[0][0]]})")
    if district_years and expect:
        both = set().union(*(set().union(*p_sd[y].values()) for y in district_years))
        if sorted(both) != list(range(1, senates + 1)):
            failed.append(f"Senate contests for {len(both)} districts in {district_years}, not {senates}")

    # ---- the districts: a precinct's House districts from its own rows; its Senate district from the year each half was up
    split_house, split_senate, unknown_senate = {}, {}, {}
    hd_counties = collections.defaultdict(set)
    sd_counties = collections.defaultdict(set)
    sd_house = collections.defaultdict(set)
    parity = {}
    whole_primary = primary_whole(primary, senates)
    for year in district_years:
        ups = set().union(*p_sd[year].values())
        parity[year] = {d % 2 for d in ups}
    for year in district_years:
        books, _summary = years[year]
        others = [y for y in district_years if y != year]
        hd_of, sd_of = p_hd[year], {}
        for key, hs in hd_of.items():
            sds = set(p_sd[year][key])
            halves = set(parity[year])
            # the same precinct in another year's results: the same name in the same county, voting in the same House
            # districts (the plan did not change between them); a name that votes elsewhere is not taken as the same place
            for y in others:
                if key in p_sd[y] and p_hd[y][key] == hs:
                    sds |= p_sd[y][key]
                    halves |= parity[y]
            if primary and key in primary and primary[key][0] == hs and (primary[key][1] or whole_primary):
                sds |= primary[key][1]
                halves |= {1}
            for h in hs:
                hd_counties[h].add(key[0])
            if len(hs) > 1:
                split_house.setdefault(year, {})[key] = sorted(hs)
            if len(sds) > 1:
                split_senate.setdefault(year, {})[key] = sorted(sds)
            elif halves != {0, 1} or not sds:
                unknown_senate.setdefault(year, {})[key] = sorted(hs)
                sds = set()
            sd_of[key] = sds
            for s in sds:
                sd_counties[s].add(key[0])
                sd_house[s] |= hs
        # which Senate districts an unknown precinct could be in: those its House districts reach in any precinct
        reach = collections.defaultdict(set)
        for key, sds in sd_of.items():
            for h in hd_of[key]:
                reach[h] |= sds
        out_h = set().union(*(set(v) for v in split_house.get(year, {}).values()))
        out_s = set().union(*(set(v) for v in split_senate.get(year, {}).values()))
        for key, hs in unknown_senate.get(year, {}).items():
            could = set().union(*(reach[h] for h in hs)) | set(p_sd[year][key])
            # (where the record reaches no Senate district from it at all, any district could be missing it: none is given that year)
            out_s |= could or set(range(1, senates + 1))
        mine = [c for c in contests if c["year"] == year]
        for c in mine:
            cid = c["id"]
            hsum, ssum, rest = collections.defaultdict(lambda: [0] * 4), collections.defaultdict(lambda: [0] * 4), [0] * 4
            for code, t in books.items():
                for p, vals in t["rows"].items():
                    v = four(vals, c["contest"], c["_side"], c["_others"])
                    hs, sds = hd_of[(code, p)], sd_of[(code, p)]
                    if len(hs) == 1 and next(iter(hs)) not in out_h:
                        _add(hsum[next(iter(hs))], v)
                    else:
                        _add(rest, v)
                    if len(sds) == 1 and next(iter(sds)) not in out_s:
                        _add(ssum[next(iter(sds))], v)
            for h in range(1, houses + 1):
                if h not in out_h and h in hsum:
                    votes["house"][str(h)][cid] = hsum[h]
            for s in range(1, senates + 1):
                if s not in out_s and s in ssum:
                    votes["senate"][str(s)][cid] = ssum[s]
            back = [sum(v[i] for v in hsum.values()) + rest[i] for i in range(4)]
            if back != state[cid]:
                failed.append(f"{cid}: the House districts given and the precincts left out add up to {back}, not {state[cid]}")
        leg_ctl[str(year)].update(
            house_districts_given=houses - len(out_h), house_districts_left_out=sorted(out_h),
            senate_districts_given=senates - len(out_s & set(range(1, senates + 1))), senate_districts_left_out=sorted(out_s),
            precincts_in_more_than_one_house_district=[{"county": COUNTY_NAME[k[0]], "precinct": k[1], "house_districts": v}
                                                       for k, v in sorted(split_house.get(year, {}).items())],
            precincts_in_more_than_one_senate_district=[{"county": COUNTY_NAME[k[0]], "precinct": k[1], "senate_districts": v}
                                                        for k, v in sorted(split_senate.get(year, {}).items())],
            precincts_whose_other_senate_half_is_not_shown=[{"county": COUNTY_NAME[k[0]], "precinct": k[1], "house_districts": v}
                                                            for k, v in sorted(unknown_senate.get(year, {}).items())])

    if expect:
        missing = sorted(set(FIPS.values()) - set(votes["county"]))
        if missing or len(votes["judicial"]) != 9:
            failed.append(f"{len(votes['county'])} counties and {len(votes['judicial'])} judicial districts, not 23 and 9")
        if sorted(bare(x) for names in judicial.values() for x in names) != sorted(FIPS):
            failed.append("judicial districts: the statute's lists are not every county once")

    # ---- the places
    places = {k: {} for k in KIND_NAMES}
    packed = lambda d: {cid: pack(d[cid]) for cid in ids if cid in d}                                  # noqa: E731
    for code in sorted(votes["county"]):
        places["county"][code] = few({"name": COUNTY_NAME[code], "votes": packed(votes["county"][code])})
    for h in sorted(votes["house"], key=int):
        places["house"][h] = few({"name": f"House District {h}", "counties": sorted(hd_counties[int(h)]), "votes": packed(votes["house"][h])})
    for s in sorted(votes["senate"], key=int):
        places["senate"][s] = few({"name": f"Senate District {s}", "house_districts": sorted(sd_house[int(s)]), "counties": sorted(sd_counties[int(s)]),
                                   "votes": packed(votes["senate"][s])})
    for d in sorted(judicial):
        places["judicial"][f"WY-JD{d}"] = few({"name": f"{['', 'First', 'Second', 'Third', 'Fourth', 'Fifth', 'Sixth', 'Seventh', 'Eighth', 'Ninth'][d]} Judicial District",
                                               "counties": sorted(FIPS[bare(x)] for x in judicial[d]), "votes": packed(votes["judicial"][f"WY-JD{d}"])})
    kinds_ctl = {}
    for kind in ("county", "judicial"):
        bad =[cid for cid in given[kind] if [sum(p["votes"][cid][s] for p in places[kind].values() if cid in p["votes"]) for s in SIDES] != state[cid]]
        kinds_ctl[kind] = "equal" if not bad else f"differs for {', '.join(bad)}"
        failed += [f"{kind}: its places do not add up to the statewide sum for {cid}" for cid in bad]
    for kind in ("house", "senate"):
        kinds_ctl[kind] = ("not every district is given; the House districts given and the precincts of those left out add up to the statewide sum"
                           if kind == "house" else "not every district is given; each district given is whole precincts that voted in it alone")

    records = []
    for c in contests:
        rec = {"id": c["id"], "date": ELECTIONS[c["year"]]["date"], "office": c["office"], "table": f"wy-sos-results-{c['year']}",
               "kinds": [k for k in KIND_NAMES if c["id"] in given[k] and any(c["id"] in p["votes"] for p in places[k].values())]}
        for s, word in (("dem", "Democratic"), ("rep", "Republican")):
            rec[s] = {"party": word, "ticket": c[s], "column": f"the column headed {c['_side'][s]}"}
        rec["other"] = {"what": "every other candidate and all write-ins, together", "columns": [h for h in c["_others"]]}
        rec["total"] = {"what": "the votes cast for candidates and write-ins; the over votes and under votes the workbooks also give are left out"}
        rec["statewide"] = pack(state[c["id"]])
        rec["official_source"] = f"clerk-statistics-{c['year']}" if clerk.get(c["id"]) else f"wy-sos-results-{c['year']}"
        records.append(rec)
    kinds = {}
    for kind, year, what, key in KINDS:
        kinds[kind] = {"what": what, "key": key, "places": len(places[kind]), "contests": given[kind],
                       "covers_the_state": kind in ("county", "judicial"),
                       "places_with_a_contest_too_few": sum(1 for p in places[kind].values() if p.get("too_few"))}
        if year > min(c["year"] for c in contests):
            kinds[kind]["why_not_2020"] = REDRAWN + (" A district missing 2022 or 2024 had a precinct that year that lay in another "
                                                     "district too, or whose district the record does not show, and the results give a "
                                                     "precinct's votes whole." if kind in ("house", "senate") else "")
    for kind, word in (("house", "House"), ("senate", "Senate")):
        lc = {y: leg_ctl[str(y)][f"{kind}_districts_left_out"] for y in district_years}
        out_all = sorted(set().union(*map(set, lc.values()))) if lc else []
        if kind == "house":
            why = ("a precinct of the district lies in another House district too, and the results give that precinct's votes whole")
        else:
            why = ("a precinct of the district lies in another Senate district too, or the record does not show the Senate district of a precinct "
                   "it may share, and the results give a precinct's votes whole")
        kinds[kind]["note"] = (f"A {word} district is given only where every precinct of it voted in that district alone. "
                               + (f"Left out ({'; '.join(f'{y}: {len(v)}' for y, v in lc.items())}): districts "
                                  f"{', '.join(str(x) for x in out_all)}, because {why}; nothing is estimated."
                                  if out_all else "Every district is given."))
        kinds[kind]["left_out"] = {str(y): v for y, v in lc.items()}
    kinds["senate"]["note"] += (" A precinct's Senate district is read from the election its district was on the ballot (the even districts "
                                "in 2024, the odd ones in 2022 and again in the 2026 primary): the same precinct of the same county, "
                                "voting in the same House districts, in the other election's results.")
    kinds["judicial"]["note"] = ("The counties of each district are those section 5-3-101, Wyoming Statutes, lists; a circuit court's district "
                                 "is the same (section 5-9-102). Every year's figures are those counties' votes.")
    control = {"result": "equal" if not failed else "differs",
               "statement": ("For every contest the precinct rows add up, county by county and column by column, to each county's Total row and "
                             "to the statewide summary, whose counties add up to its Total row; for President and U.S. Senator the totals equal "
                             "the Clerk of the House's statistics; for 2022 and 2024 the precinct rows of every House and Senate contest add up "
                             "to that contest's county rows in the summary; counties and judicial districts add up to the statewide sum, and "
                             "the House districts given, with the precincts of those left out, do too."
                             if not failed else "The sums do not all agree; see the differences."),
               "precincts": {str(y): sum(len(t["rows"]) for t in years[y][0].values()) for y in sorted(years)},
               "contests": ctl, "kinds": kinds_ctl, "districts": leg_ctl,
               "notes": ["Labels some counties typed as 1-2 were stored by the spreadsheet as dates; they are written back as month-day, as typed.",
                         "The 2024 summary workbook also holds a sheet named Statewide Senate Odd that carries the 2022 primary's heading and "
                         "figures; it is not read. The 2024 Senate contests are checked against its Statewide Senate Even sheet."]}
    if primary is not None:
        control["notes"].append("The 2026 primary's precinct workbook was read for which odd-numbered Senate contest each precinct voted in, "
                                "as a further record of the Senate district of the precincts of 2022 and 2024 that bear the same name; no "
                                "votes or names were taken from it.")
    doc = {"what": WHAT, "note": NOTE, "state": "WY", "generated": M._now(), "method_version": METHOD, "how_to_read": HOW_TO_READ,
           "vote_keys": list(SIDES), "too_few": {"fewer_than": FEW, "why": TOO_FEW}, "contests": records, "kinds": kinds, "control": control,
           "coverage": {"not_given": NOT_GIVEN, "no_2022_senate": "Wyoming elected no United States senator in 2022.",
                        "governor": "Wyoming elected its Governor in 2022, not in 2020 or 2024."},
           "sources": [], "places": places}
    for c in contests:
        c.pop("_side", None)
        c.pop("_others", None)
    return doc, failed


# ---------------------------------------------------------------- the run

def read_clerk(path, state="WYOMING"):
    """The state's page of the Clerk of the House's statistics: {section: [(label, votes)]} for the presidential
    electors and United States Senator; the page number line and the representatives' part end it."""
    from ballot import pdftext
    lines = pdftext.lines(path)
    start = next((i for i, (_p, _y, t) in enumerate(lines) if t.strip() == state), None)
    if start is None:
        raise ValueError(f"no {state} heading")
    page = lines[start][0]
    sections, section = {}, None
    for p, _y, t in lines[start + 1:]:
        t = t.strip()
        if p != page or t.startswith("FOR UNITED STATES REPRESENTATIVE"):
            break
        if t.startswith("FOR "):
            section = t
            sections[section] = []
            continue
        m = re.match(r"^(.*?)\s*\.{3,}\s*([\d,]+)$", t)
        if section and m:
            sections[section].append((m.group(1).strip(), int(m.group(2).replace(",", ""))))
        elif section and t:
            raise ValueError(f"a line under {section} is not a name, dots and a number")
    top = next((t.strip() for p, _y, t in lines if p == page and t.strip().isdigit()), None)
    return {"where": f"Wyoming, page {top}" if top else f"Wyoming, page {page} of the PDF", "sections": sections}


def clerk_figures(doc, section):
    lines = [(lab, v) for lab, v in doc["sections"].get(section) or [] if lab not in ("Under Votes", "Over Votes")]
    party = (lambda s: s) if section == "FOR PRESIDENTIAL ELECTORS" else (lambda s: s.rsplit(",", 1)[-1].strip())
    dem = [v for lab, v in lines if party(lab).startswith("Democrat")]
    rep = [v for lab, v in lines if party(lab) == "Republican"]
    if len(dem) != 1 or len(rep) != 1:
        raise ValueError(f"{section}: not one Democratic and one Republican line")
    total = sum(v for _l, v in lines)
    return {"dem": dem[0], "rep": rep[0], "other": total - dem[0] - rep[0], "total": total}


def clerk_totals(contests, cache, refresh, say):
    out, recs = {}, []
    for year in sorted({c["year"] for c in contests if c.get("clerk")}, reverse=True):
        src = CLERK[year]
        path = os.path.join(cache, src["file"])
        rec = {k: src[k] for k in ("id", "kind", "agency", "title", "url")}
        try:
            net.download(src["url"], path, 0 if refresh else KEEP_DAYS, tries=3, say=say)
            doc = read_clerk(path)
            rec.update(fetched=M._day(path), sha256=M._sha_file(path), where=doc["where"],
                       read="The Wyoming lines for presidential electors and United States Senator; its under and over votes are left out.")
            for c in contests:
                if c["year"] == year and c.get("clerk"):
                    out[c["id"]] = dict(clerk_figures(doc, c["clerk"]), where=doc["where"])
        except Exception as e:  # noqa: BLE001  the summaries are the control; the Clerk's page is a second one
            say(f"      {src['title']}: could not be read ({e}); the Secretary's summaries alone are the control for {year}")
            rec["unread"] = f"could not be read on {M._now()}: {e}"
        recs.append(rec)
    return out, recs


def our_places(db):
    """{kind: ids} of our Wyoming places and races, read only; None when the database is not there."""
    if not db or not os.path.exists(db):
        return None
    con = sqlite3.connect(pathlib.Path(os.path.abspath(db)).as_uri() + "?mode=ro", uri=True)
    try:
        out = collections.defaultdict(set)
        for kind, pid in con.execute("SELECT kind, id FROM sl_places WHERE source_id LIKE 'wy-%'"):
            out[{"county": "county", "mcd": "mcd"}.get(kind, kind)].add(pid)
        for office, level, jid, district in con.execute("SELECT office_kind, level, jurisdiction_id, district FROM sl_races WHERE state = 'WY'"):
            if office == "state_senate":
                out["senate"].add(str(district))
            elif office == "state_house":
                out["house"].add(str(district))
            elif level == "court" and str(jid).startswith("WY-JD"):
                out["judicial"].add(str(jid))
            elif level == "county":
                out["county"].add(str(jid))
        return dict(out)
    except sqlite3.Error:
        return None
    finally:
        con.close()


def load(say=print, out=OUT, cache=CACHE, db=DB, refresh=False):
    say("    Wyoming place votes: the Secretary of State's official precinct results and summaries (sos.wyo.gov/Elections/ElectionResults.aspx)")
    net.patient_lookups()
    os.makedirs(cache, exist_ok=True)
    years, srcs, fetched = {}, [], {}
    for year in sorted(ELECTIONS, reverse=True):
        e = ELECTIONS[year]
        path = fetch(e["zip"], cache, refresh, say, b"PK")
        years[year] = read_year(path, year)
        fetched[year] = M._day(path)
        n = sum(len(t["rows"]) for t in years[year][0].values())
        say(f"      {year}: {n:,} precincts in {len(years[year][0])} counties (copy of {fetched[year]})")
        inside = (e["books"] if isinstance(e["books"], str) else "the 23 files named 2020_<county>_County_General_PbP.xlsx") + " and " + e["summary"]
        srcs.append({"id": f"wy-sos-results-{year}", "kind": "official results by precinct, with the statewide summaries", "agency": AGENCY,
                     "title": e["zip"]["title"], "url": e["zip"]["url"], "listed_on": e["page"], "fetched": fetched[year], "sha256": M._sha_file(path),
                     "precincts": n, "files_read": inside,
                     "read": "Of each row: the precinct and the counts under each contest and candidate heading; of the headings, only the "
                             "candidates' names and party letters in the contests named here, to check which ticket a column is, and the "
                             "House and Senate contest headings, to see which district each precinct voted in."})
    primary = None
    if os.path.exists(PRIMARY_2026):
        try:
            primary = read_primary(PRIMARY_2026)
            srcs.append({"id": "wy-sos-precinct-2026-primary", "kind": "official results by precinct (a check only)", "agency": AGENCY,
                         "title": "2026 Primary Election, County Precinct by Precinct Results, Official", "url": f"{DOCS}/2026/Results/Primary/2026_Wyoming_Primary_Results.zip",
                         "listed_on": f"{DOCS}/2026/2026PrimaryResults.aspx", "sha256": M._sha_file(PRIMARY_2026),
                         "read": "Only which Senate contest each precinct voted in; no votes and no names are taken from it."})
        except (Stop, KeyError, zipfile.BadZipFile) as e:
            say(f"      the 2026 primary workbook could not be read ({e}); the two general elections alone show the Senate districts")
            primary = None
    srcs.append(statute(cache, refresh, say))
    say("    the second control: Clerk of the U.S. House (President and U.S. Senator)")
    clerk, clerk_recs = clerk_totals(CONTESTS, cache, refresh, say)
    doc, failed = build(years, clerk, contests=[dict(c) for c in CONTESTS], primary=primary, sources=fetched)
    doc["sources"] = srcs + clerk_recs

    p = doc["places"]
    say(f"    added up: {len(p['county'])} counties; {len(p['judicial'])} judicial districts; for 2022 and 2024, {len(p['house'])} House and "
        f"{len(p['senate'])} Senate districts")
    for y, d in doc["control"]["districts"].items():
        say(f"      {y}: House districts given {d.get('house_districts_given')}, left out {d.get('house_districts_left_out')}; Senate given "
            f"{d.get('senate_districts_given')}, left out {d.get('senate_districts_left_out')}")
    for c in doc["contests"]:
        k, s = doc["control"]["contests"].get(c["id"], {}), c["statewide"]
        say(f"    control, {c['id']}: Democratic {s['dem']:,}, Republican {s['rep']:,}, other {s['other']:,}, total {s['total']:,}: "
            + ("equal to" if k.get("equal") else "DIFFERS from") + f" the Total rows and the summary ({(k.get('counties') or {}).get('equal')} of "
            f"{(k.get('counties') or {}).get('compared')} counties)"
            + (f"; Clerk of the House {'equal' if k['clerk']['equal'] else 'DIFFERS'}" if k.get("clerk") else ""))
    if failed:
        for line in failed:
            say(f"    CHECK: {line}")
        raise Stop(f"    Wyoming place votes: the control did not hold; {os.path.basename(out)} was not written. Nothing was changed to make it fit.")
    ours = our_places(db)
    if ours:
        covered = {}
        for kind in KIND_NAMES:
            idset = ours.get(kind, set())
            without = sorted((i for i in idset if not p[kind].get(i, {}).get("votes")), key=lambda x: (len(x), x))
            covered[kind] = {"places": len(idset), "with_votes": len(idset) - len(without), "without": without}
            say(f"      our {kind} places with votes: {len(idset) - len(without):,} of {len(idset):,}" + (f" (none for {', '.join(without[:16])})" if without else ""))
        covered["not_given_kinds"] = {k: len(v) for k, v in sorted(ours.items()) if k not in KIND_NAMES}
        doc["coverage"]["our_places"] = dict(covered, read=f"sl_places and sl_races in {os.path.basename(db)}, opened read-only on {M._now()}")
    M.write(out, doc)
    say(f"    Wyoming place votes: {len(doc['contests'])} contests, {sum(len(v) for v in p.values()):,} places, control {doc['control']['result']}; "
        f"{os.path.relpath(out, HERE)} ({os.path.getsize(out) / 1e6:.2f} MB)")
    return doc


# ---------------------------------------------------------------- the arithmetic and the readers on made-up tables

def selftest(say=print):
    checks = []

    def check(what, got, want):
        checks.append(got == want)
        say(f"      {'ok  ' if got == want else 'FAIL'} {what}" + ("" if got == want else f": got {got!r}, expected {want!r}"))

    def sheet(rows, title="X County Official Precinct-by-Precinct Summary"):
        con = [None, "Governor", None, None, None, None, None, "House District 1", None, None, None, None, "House District 2", None, None, None, None,
               "Senate District 1", None, None, None, None]
        head = ["Precinct", "Ann\nAlder (D)", "Bo\nBirch (R)", "Cy Cedar (L)", "Write-Ins", "Overvotes", "Undervotes",
                "H One (R)", "Write-Ins", "Overvotes", "Undervotes", None, "H Two (R)", "Write-Ins", "Overvotes", "Undervotes", None,
                "S One (R)", "Write-Ins", "Overvotes", "Undervotes", None]
        out = [[None, title], [None], con, head]
        tot = [0] * (len(head) - 1)
        for i, r in enumerate(rows):
            out.append(list(r))
            if i == 1:
                out.append(["Precincts Continue\non Next Page"])
            for j, v in enumerate(r[1:]):
                tot[j] += v if isinstance(v, int) else 0
        out.append(["Total"] + [t if any(isinstance(r[j + 1], int) for r in rows) else "-" for j, t in enumerate(tot)])
        return out
    d = "-"
    a_rows = [["1-1", 30, 50, 2, 1, 0, 3, 70, 1, 0, 2, None, d, d, d, d, None, 75, 0, 0, 0, None],
              [dt.datetime(2022, 1, 2), 10, 5, 0, 0, 0, 1, 14, 0, 0, 0, None, d, d, d, d, None, d, d, d, d, None],
              ["2-1", 40, 20, 1, 0, 0, 0, 20, 0, 0, 0, None, 40, 1, 0, 0, None, 60, 0, 0, 0, None]]
    b_rows = [["01-01", 0, 12, 0, 0, 0, 0, d, d, d, d, None, 12, 0, 0, 0, None, d, d, d, d, None]]
    ta, tb = table(sheet(a_rows), "made-up A"), table(sheet(b_rows), "made-up B")
    check("a sheet is read by its headings, a date label written back as typed, the continue line passed over",
          (sorted(ta["rows"]), ta["rows"]["1-1"][("Governor", "Ann Alder (D)")], ta["total"][("Governor", "Bo Birch (R)")]), (["1-1", "1-2", "2-1"], 30, 75))
    check("a dash is no count", ta["rows"]["1-2"][("Senate District 1", "S One (R)")], None)

    def summ(rows):
        t = table(sheet(rows), "made-up summary")
        t["sheet"] = "Statewide Candidates"
        return t
    def col_sums(rows):
        return [sum(r[j + 1] for r in rows if isinstance(r[j + 1], int)) if any(isinstance(r[j + 1], int) for r in rows) else "-"
                for j in range(len(rows[0]) - 1)]
    summary = summ([["Albany"] + col_sums(a_rows), ["Big Horn"] + col_sums(b_rows)])
    summary_h = dict(summary, sheet="Statewide House")
    summary_s = dict(summary, sheet="Statewide Senate")
    books = {"56001": ta, "56003": tb}
    contests = [{"id": "2022-governor", "year": 2022, "office": "Governor", "contest": "Governor", "dem": "Ann Alder", "rep": "Bo Birch", "clerk": None}]
    judicial = {1: ["Albany"], 2: ["Big Horn"]}
    years = {2022: (books, {"candidates": summary, "house": summary_h, "senate": summary_s})}
    doc, failed = build(years, {}, contests=[dict(c) for c in contests], judicial=judicial, houses=2, senates=1, expect=False)
    p = doc["places"]
    check("the control holds on the made-up tables", (doc["control"]["result"], failed), ("equal", []))
    check("a county adds up its precincts, write-ins with the others", p["county"]["56001"]["votes"]["2022-governor"],
          {"dem": 80, "rep": 75, "other": 4, "total": 159})
    check("a judicial district is whole counties", p["judicial"]["WY-JD2"]["votes"]["2022-governor"], {"dem": 0, "rep": 12, "other": 0, "total": 12})
    check("a precinct in two House districts leaves both out", sorted(p["house"]), [])
    check("and is named", [x["precinct"] for x in doc["control"]["districts"]["2022"]["precincts_in_more_than_one_house_district"]], ["2-1"])
    check("a place where every vote went one way is marked too_few", p["county"]["56003"].get("too_few"), ["2022-governor"])
    # without the split precinct, district 1 is given, and its Senate district only where the other half is shown
    a2 = [r for r in a_rows if r[0] != "2-1"]
    b2 = [["01-01", 0, 12, 0, 0, 0, 0, d, d, d, d, None, 12, 0, 0, 0, None, d, d, d, d, None]]
    ta2, tb2 = table(sheet(a2), "made-up A2"), table(sheet(b2), "made-up B2")
    s2 = summ([["Albany"] + col_sums(a2), ["Big Horn"] + col_sums(b2)])
    years2 = {2022: ({"56001": ta2, "56003": tb2}, {"candidates": s2, "house": dict(s2, sheet="h"), "senate": dict(s2, sheet="s")})}
    doc2, failed2 = build(years2, {}, contests=[dict(c) for c in contests], judicial=judicial, houses=2, senates=1, expect=False)
    check("a House district of whole precincts is given", (failed2, doc2["places"]["house"]["1"]["votes"]["2022-governor"]),
          ([], {"dem": 40, "rep": 55, "other": 3, "total": 98}))
    check("a Senate district whose precincts' other half is not shown is left out", sorted(doc2["places"]["senate"]), [])
    wrong = dict(ta2, total={**ta2["total"], ("Governor", "Ann Alder (D)"): 41})
    years3 = {2022: ({"56001": wrong, "56003": tb2}, years2[2022][1])}
    check("a county whose Total row differs is caught", build(years3, {}, contests=[dict(c) for c in contests], judicial=judicial, houses=2, senates=1,
                                                               expect=False)[0]["control"]["result"], "differs")
    try:
        table(sheet(a_rows + [a_rows[0]]), "made-up twice")
        caught = False
    except Stop:
        caught = True
    check("a precinct named twice stops the loader", caught, True)
    text = ("5-3-101. Judicial districts enumerated; terms of court. (a) The state of Wyoming is divided into judicial districts as follows: "
            "(i) The county of Laramie is the first judicial district; (ii) The counties of Albany and Carbon are the second judicial district; "
            "(A) Repealed by Laws 2019. (iii) The counties of Sweetwater, Lincoln and Uinta are the third judicial district; "
            "(vii) Natrona county is the seventh judicial district; 5-3-102. Next")
    check("the statute is read", statute_districts(text), {1: ["Laramie"], 2: ["Albany", "Carbon"], 3: ["Sweetwater", "Lincoln", "Uinta"], 7: ["Natrona"]})
    check("county codes follow the Census order", (FIPS["ALBANY"], FIPS[bare("Big_Horn")], FIPS["WESTON"], FIPS[bare("Hot Springs County")]),
          ("56001", "56003", "56045", "56017"))
    say(f"    self-test: {sum(checks)} of {len(checks)} checks passed")
    return all(checks)


def main(argv=None):
    ap = argparse.ArgumentParser(description="How each Wyoming place voted in past partisan general elections -> ballot/lean/wy_place_votes.json")
    ap.add_argument("--out", default=OUT, help="the file to write (default: ballot/lean/wy_place_votes.json)")
    ap.add_argument("--cache", default=CACHE, help="where downloads are kept (default: states_cache/wy_local)")
    ap.add_argument("--db", default=DB, help="the ballot database to count our places in, opened read-only (default: ballot_local_2026.sqlite)")
    ap.add_argument("--refresh", action="store_true", help="ask for every file again, even when the cached copies are there")
    ap.add_argument("--selftest", action="store_true", help="check the arithmetic and the readers on made-up tables; downloads and writes nothing")
    a = ap.parse_args(argv)
    if a.selftest:
        raise SystemExit(0 if selftest() else 1)
    load(out=a.out, cache=a.cache, db=a.db, refresh=a.refresh)


if __name__ == "__main__":
    main()
