"""
ballot/state_local_ca.py - California's state races on the November 3, 2026 ballot, into ballot_local_2026.sqlite:

  - the eight statewide offices elected this year for four-year terms: Governor, Lieutenant Governor (a contest of its
    own, not a ticket), Secretary of State, Controller, Treasurer, Attorney General, Insurance Commissioner, and
    Superintendent of Public Instruction (nonpartisan);
  - the four seats of the State Board of Equalization, elected by district;
  - the 20 even-numbered State Senate districts (senators serve four years, 20 seats beginning every two years) and all
    80 State Assembly districts (two-year terms);
  - the retention questions for justices of the Supreme Court and the six Courts of Appeal (yes or no, no opponent);
with the June 2, 2026 top-two primary and its official votes. The federal rows (United States Representative) are
left to ballot/lists/ca.py and the federal database, which is never opened here.

Sources, all the Secretary of State's own:
  - the November ballot: the Certified List of Candidates for the November 3, 2026 General Election (the PDF the
    federal loader caches as ballot_cache/ca_cert_list_2026_general.pdf). Its certificate says it holds the name, office
    sought, ballot designation and party preference of each person nominated at the June 2 primary, and the names of
    the justices entitled to receive votes; it prints no address or contact detail. Each office is headed in 14-point
    type; each candidate takes a 9-point name row (the name, * for an incumbent, and the party preference in a column to
    the right) and an indented row beneath with the ballot designation. Only the name and party preference cells are
    turned into text; the ballot designation rows are never joined. The court pages are 12-point prose: each appellate
    district's counties, then for each justice "For <office>" and "Shall <title> <NAME> be elected to the office for
    the term provided by law?"; those words are read.
  - the June 2 primary: the Statement of Vote's "CSV Files - Voter Nominated" workbook (one row per county, contest and
    candidate: name, party preference, incumbent and write-in flags, votes; the federal loader's cached copy), with the
    county rows added up; and, to reconcile every figure, the Statement of Vote's own workbook for each office
    (19-gov, 46-lt-gov, 55-sos, 58-sco, 61-treasurer, 64-ag, 67-ic, 73-boe, 90-state-senator, 95-state-assembly,
    113-spi), linked from the Statement of Vote page and cached in ballot_cache/ca/. Every candidate's county figures and
    State or District Totals there must equal the county rows of the CSV workbook. These files carry names, party
    preferences and votes only.
Who holds each seat comes from the Open States roster in state_ca.sqlite (legislators is_current = 1 by chamber and
district; the officials table for the Governor, Lieutenant Governor, Attorney General and Secretary of State, the only
statewide offices it carries): ids, names and party only. County codes come from the Census Bureau's 2024 county file
(states_cache/census/), matched by name; the counties a district reaches are the ones its primary was counted in
(derived, and said so), and an appellate district's are the ones the certified list prints under its heading.

California's primary is top-two: every candidate, of every party preference, is on one primary ballot, and the two with
the most votes advance whatever their parties (a tie for second advances everyone tied). The Superintendent of Public
Instruction is nonpartisan and works the same way unless someone wins more than half the June vote. So the primary
field is election "primary" (as on the federal side), shown only where two or more candidates were on the June ballot,
and the November ballot is the top two, checked against the certified list race by race (family name and party
preference); where they differ the certified list is what the ballot prints, and its names are used. The party shown
is the candidate's own stated preference as the list prints it ("Democratic", "No Party Preference"); it is not a
party's nomination. party_code: D and R, L Libertarian, G Green, I No Party Preference, O for American Independent,
Peace and Freedom and anything else; the nonpartisan office is "Nonpartisan office", code N.

Ballot order. Neither official file gives the order of names on the ballot, so no ballot_order is stored.

Privacy: only office, district, candidate name, party preference, incumbent and write-in marks and votes are read.
Ballot designations (occupations) are never turned into text; no address, city, ZIP code, phone, website, e-mail or
treasurer is read, printed, logged, cached or stored (none of these files carries one), and the roster's contact columns
are never selected. Before anything is written every stored name, party and note is checked for anything that looks
like a contact detail, and the load stops (without showing it) if one does.

    python -m ballot.state_local_ca <path to a test database>
"""

import collections
import datetime as dt
import hashlib
import io
import json
import os
import re
import sqlite3
import sys
import zipfile

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

import openpyxl  # noqa: E402

from ballot.common import CACHE, fold, name_parts  # noqa: E402
from ballot.lists import ca as C  # noqa: E402
from ballot.match import fits  # noqa: E402
from ballot.pdftext import PDF, join, rows as pdf_rows  # noqa: E402
from states import net  # noqa: E402

STATE, FIPS, NAME = "CA", "06", "California"
GENERAL, PRIMARY = "2026-11-03", C.PRIMARY
ROSTER = os.path.join(HERE, "state_ca.sqlite")
COUNTY_ZIP = os.path.join(HERE, "states_cache", "census", "cb_2024_us_county_500k.zip")
COUNTY_URL = "https://www2.census.gov/geo/tiger/GENZ2024/shp/cb_2024_us_county_500k.zip"
CSV_FILE = "ca_sov_2026_primary_voter_nominated.xlsx"      # the federal loader's cached copy, same address
CERT_FILE = "ca_cert_list_2026_general.pdf"                 # likewise
BOOK_BASE = "https://elections.cdn.sos.ca.gov/sov/2026-primary/sov/"
BOOK_FILE = "ca_2026_primary_sov_{}"
AGENCY = "California Secretary of State"
COUNTIES = 58
SENATE_UP = [str(d) for d in range(2, 41, 2)]
ASSEMBLY = [str(d) for d in range(1, 81)]
SRC_CSV, SRC_CERT = "ca-sos-2026-sl-primary-csv", "ca-sos-2026-sl-general-certified-list"
SRC_ROSTER, SRC_COUNTY = "ca-openstates-roster", "ca-census-2024-county-codes"
NONPARTISAN = "Nonpartisan office"

# contest name (as both official files write it) -> (key, office_kind, office, roster office, workbook, partisan)
STATEWIDE = {
    "Governor": ("GOV", "governor", "Governor", "governor", "19-gov.xlsx", 1),
    "Lieutenant Governor": ("LTG", "lieutenant_governor", "Lieutenant Governor", "lt_governor", "46-lt-gov.xlsx", 1),
    "Secretary of State": ("SOS", "secretary_of_state", "Secretary of State", "secretary of state", "55-sos.xlsx", 1),
    "Controller": ("CTRL", "state_controller", "Controller", None, "58-sco.xlsx", 1),
    "Treasurer": ("TREAS", "state_treasurer", "Treasurer", None, "61-treasurer.xlsx", 1),
    "Attorney General": ("AG", "attorney_general", "Attorney General", "attorney general", "64-ag.xlsx", 1),
    "Insurance Commissioner": ("INS", "insurance_commissioner", "Insurance Commissioner", None, "67-ic.xlsx", 1),
    "Superintendent of Public Instruction": ("SPI", "superintendent_of_public_instruction", "Superintendent of Public Instruction", None,
                                             "113-spi.xlsx", 0),
}
DISTRICT_BOOKS = {"73-boe.xlsx": ("Board of Equalization District", "Board of Equalization Member District {}"),
                  "90-state-senator.xlsx": ("State Senate District", "State Senate District {}"),
                  "95-state-assembly.xlsx": ("Assembly District", "State Assembly Member District {}")}
CONTEST_RE = [
    (re.compile(r"State Senate District (\d+)"), "SS"),
    (re.compile(r"State Assembly Member District (\d+)"), "SH"),
    (re.compile(r"Board of Equalization Member District (\d+)"), "BOE"),
]
# the Statement of Vote workbooks' party codes -> the CSV workbook's words (checked, not assumed: see load)
BOOK_PARTY = {"DEM": "Democratic", "REP": "Republican", "NPP": "No Party Preference", "LIB": "Libertarian", "GRN": "Green",
              "PF": "Peace and Freedom", "AI": "American Independent", "NP": "Non-Partisan"}
ORDINAL_WORDS = {"First": 1, "Second": 2, "Third": 3, "Fourth": 4, "Fifth": 5, "Sixth": 6}
DIVISION_WORDS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5, "Six": 6, "Seven": 7, "Eight": 8}

SCHEMA = """
CREATE TABLE IF NOT EXISTS sl_races (race_id TEXT PRIMARY KEY, state TEXT NOT NULL, level TEXT NOT NULL, office_kind TEXT NOT NULL,
  office TEXT NOT NULL, jurisdiction TEXT, jurisdiction_id TEXT, county_ids TEXT, district TEXT, seat TEXT, special INTEGER NOT NULL DEFAULT 0,
  partisan INTEGER NOT NULL, holder_id TEXT, holder_name TEXT, holder_party TEXT, election_date TEXT NOT NULL, note TEXT);
CREATE TABLE IF NOT EXISTS sl_candidates (race_id TEXT NOT NULL, election TEXT NOT NULL, election_date TEXT, name TEXT NOT NULL, party TEXT,
  party_code TEXT, ballot_order INTEGER, incumbent INTEGER NOT NULL DEFAULT 0, write_in INTEGER NOT NULL DEFAULT 0, votes INTEGER, pct REAL,
  outcome TEXT, state_member_id TEXT, source_id TEXT, note TEXT, PRIMARY KEY (race_id, election, name));
CREATE TABLE IF NOT EXISTS sl_sources (source_id TEXT PRIMARY KEY, state TEXT, kind TEXT, agency TEXT, title TEXT, url TEXT, published TEXT,
  fetched TEXT, sha256 TEXT, rows INTEGER, note TEXT);
CREATE TABLE IF NOT EXISTS sl_places (kind TEXT NOT NULL, id TEXT NOT NULL, name TEXT NOT NULL, county_ids TEXT, source_id TEXT,
  PRIMARY KEY (kind, id));
"""

CONTACT = re.compile(r"@|https?:|www\.|\.(?:com|org|net|gov|us)\b|\d{3}|P\.?\s?O\.?\s+Box|\bSuite\b", re.I)
CONTACT_NOTE = re.compile(r"@|https?:|www\.|\b\d{3}[-.)\s]\s?\d{3}[-.\s]\d{4}\b|\bP\.?\s?O\.?\s+Box\b|\bSuite\b", re.I)

TOP_TWO = ("California's primary is top-two: every candidate, of every party preference, was on one June 2 ballot, and the two "
           "with the most votes advanced to November, whatever their parties. The party shown is the candidate's own stated "
           "preference as the ballot prints it, not a party's nomination.")
SPI_NOTE = ("A nonpartisan office: no party is printed on the ballot. No candidate won more than half the June 2 vote, so the two "
            "with the most votes are on the November ballot.")
FOUR_YEARS = "Elected statewide for a four-year term."
LTG_NOTE = "California elects its Lieutenant Governor in a contest of its own, not on a ticket with the Governor."
BOE_NOTE = ("The State Board of Equalization has four members elected by district, each for a four-year term, all four "
            "this year.")
SENATE_NOTE = ("The State Constitution gives senators four-year terms, 20 of the 40 beginning every two years (Article IV, "
               "section 2); this year's official lists carry the 20 even-numbered districts.")
ASSEMBLY_NOTE = ("Members of the Assembly serve two-year terms (State Constitution, Article IV, section 2), so all 80 seats are "
                 "on the ballot every even year.")
NO_ROSTER = ("The Open States roster carries only the Governor, Lieutenant Governor, Attorney General and Secretary of State, "
             "so no holder is shown for this office.")
NO_HOLDER = "The Open States roster shows no one holding this seat today, so no holder is shown."
SINGLE = "Only one candidate was on the June 2 primary ballot, so the November ballot names one."
ORDER_NOTE = "Neither official file gives the order of names on the ballot, so none is shown."
COURT_NOTE = ("A retention question: voters answer yes or no, and no one runs against the justice. Nonpartisan. The Open States "
              "roster does not carry judges, so no holder is shown.")


# ------------------------------------------------------------------------------------------------ small helpers

def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def mtime(path):
    return dt.datetime.fromtimestamp(os.path.getmtime(path)).strftime("%Y-%m-%d")


def ws(text):
    return re.sub(r"\s+", " ", str(text or "")).strip()


def county_key(text):
    return re.sub(r"[^a-z]", "", re.sub(r"\s+County$", "", ws(text), flags=re.I).lower())


def num(v):
    """A vote cell (the workbooks write some as numbers and some as text) -> int, or None if it is not a count."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else None
    t = str(v).strip().replace(",", "")
    return int(t) if t.isdigit() else None


def census_counties(path):
    """{county key: (GEOID, name)} for California from the Census Bureau's county file (names and codes only)."""
    import shapefile                                                     # pyshp, in the kit's environment
    z = zipfile.ZipFile(path)
    dbf = next(n for n in z.namelist() if n.endswith(".dbf"))
    rdr = shapefile.Reader(dbf=io.BytesIO(z.read(dbf)))
    return {county_key(r["NAME"]): (r["GEOID"], r["NAME"]) for r in (x.as_dict() for x in rdr.iterRecords()) if r["STATEFP"] == FIPS}


def shown_party(p):
    return NONPARTISAN if fold(p) in ("non partisan", "nonpartisan") else ws(p)


def code(p):
    t = fold(p)
    if t in ("non partisan", "nonpartisan", "nonpartisan office"):
        return "N"
    return {"democratic": "D", "republican": "R", "libertarian": "L", "green": "G", "no party preference": "I"}.get(t, "O")


# ------------------------------------------------------------------------------------------------ the races

def contest_race(contest):
    """(race_id, level, office_kind, office, jurisdiction, jurisdiction_id, district, partisan) for a contest name, or None."""
    contest = ws(contest)
    if contest in STATEWIDE:
        key, kind, office, _r, _b, partisan = STATEWIDE[contest]
        return f"2026-{STATE}-{key}", "statewide", kind, office, NAME, STATE, None, partisan
    for rx, key in CONTEST_RE:
        m = rx.fullmatch(contest)
        if m:
            d = str(int(m.group(1)))
            if key == "SS":
                return f"2026-{STATE}-SS{d}", "legislature", "state_senate", "State Senator", f"Senate District {d}", d, d, 1
            if key == "SH":
                return f"2026-{STATE}-SH{d}", "legislature", "state_house", "Member of the State Assembly", f"Assembly District {d}", d, d, 1
            return (f"2026-{STATE}-BOE{d}", "statewide", "board_of_equalization", "Member of the State Board of Equalization",
                    f"Board of Equalization District {d}", f"BOE{d}", d, 1)
    if re.search(r"senat|assembly|equalization|governor|controller|treasurer|attorney|secretary|insurance|superintendent", contest, re.I):
        raise SystemExit(f"California (state races): a state contest that is not read ({contest!r})")
    return None


# ------------------------------------------------------------------------------------------------ the June 2 primary

def read_csv_export(path):
    """{contest: {"cands": {candidate id: [name, party, incumbent flag, write-in flag, votes]}, "county": {(county, id): votes}}}
    for the state contests, and {county id: county name}. The columns are found by their headings."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    it = wb.worksheets[0].iter_rows(values_only=True)
    head = [ws(h) for h in next(it)]
    need = ("Election Date", "County Id", "County Name", "Contest Name", "Candidate ID", "Candidate Name", "Incumbent Flag",
            "Write-in Flag", "Party Name", "Vote Total")
    if not all(k in head for k in need):
        raise SystemExit(f"California (state races): the Statement of Vote CSV workbook's columns changed ({head})")
    ix = {k: head.index(k) for k in need}
    out, names, dates, blank = {}, {}, collections.Counter(), 0
    for r in it:
        if all(v is None for v in r):
            blank += 1
            continue
        contest = ws(r[ix["Contest Name"]])
        if contest.startswith("United States ") or contest_race(contest) is None:
            continue
        dates[ws(r[ix["Election Date"]])] += 1
        f = out.setdefault(contest, {"cands": {}, "county": collections.Counter()})
        cid = ws(r[ix["Candidate ID"]])
        c = f["cands"].setdefault(cid, [ws(r[ix["Candidate Name"]]), ws(r[ix["Party Name"]]), ws(r[ix["Incumbent Flag"]]) == "Y",
                                        ws(r[ix["Write-in Flag"]]) == "Y", 0])
        v = num(r[ix["Vote Total"]])
        if v is None:
            raise SystemExit(f"California (state races): a vote total in the CSV workbook is not a number ({contest})")
        c[4] += v
        county = ws(r[ix["County Name"]])
        names[ws(r[ix["County Id"]])] = county
        f["county"][(county, cid)] += v
    wb.close()
    if set(dates) != {"06/02/2026"}:
        raise SystemExit(f"California (state races): the CSV workbook's state rows are not all of the June 2, 2026 primary ({dict(dates)})")
    return out, names, blank


def read_book(path, fname):
    """{contest: {candidate name: {"code": party code, "total": votes, "county": {county: votes}}}} from one of the Statement
    of Vote's office workbooks (first sheet), with the problems found inside it (county rows that do not add up)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    wb.close()
    out, problems = {}, []
    if fname in DISTRICT_BOOKS:
        words, fmt = DISTRICT_BOOKS[fname]
        contest = None
    else:
        contest = next(c for c, v in STATEWIDE.items() if v[4] == fname)
        out[contest] = {}
    block, i = None, 0
    while i < len(rows):
        r = rows[i]
        c0 = ws(r[0]) if r[0] is not None else ""
        if fname in DISTRICT_BOOKS and re.fullmatch(rf"(\d+)(?:st|nd|rd|th) {words}", c0):
            contest = fmt.format(int(re.match(r"\d+", c0).group(0)))
            if contest in out:
                raise SystemExit(f"California (state races): {fname} has two blocks for {contest}")
            out[contest] = {}
            block = None
        elif not c0 and any(v is not None for v in r[1:]) and all(num(v) is None for v in r[1:] if v is not None):
            # a names row, and the party row under it
            if contest is None or i + 1 >= len(rows):
                raise SystemExit(f"California (state races): {fname} has a names row outside any contest")
            names = {j: ws(v) for j, v in enumerate(r) if j and v is not None}
            parties = {j: ws(rows[i + 1][j]) for j in names}
            block = {j: out[contest].setdefault(names[j], {"code": parties[j], "total": None, "county": {}}) for j in names}
            if len(block) != len(names):
                raise SystemExit(f"California (state races): {fname} names a candidate twice in {contest}")
            i += 2
            continue
        elif c0 and block is not None and not c0.startswith("Percent"):
            vals = {j: num(r[j]) for j in block}
            if re.fullmatch(r"(State|District) Totals?", c0):
                for j, c in block.items():
                    c["total"] = vals[j]
                    if vals[j] is None or sum(c["county"].values()) != vals[j]:
                        problems.append(f"{fname} {contest}: {c0} for a candidate is {vals[j]} and the county rows add to "
                                        f"{sum(c['county'].values())}")
                block = None
            else:
                for j, c in block.items():
                    if vals[j] is None:
                        raise SystemExit(f"California (state races): {fname} {contest}: a county row with a cell that is not a count")
                    c["county"][c0] = vals[j]
        i += 1
    for contest, cands in out.items():
        for n, c in cands.items():
            if c["total"] is None:
                problems.append(f"{fname} {contest}: no totals row for a candidate ({n})")
    return out, problems


def book_name(n):
    return fold(n)


# ------------------------------------------------------------------------------------------------ the November ballot

def certified(path):
    """(title, published, dated, {contest: [(name, party preference, incumbent mark)]}, [court question], problems).
    Only the name and party preference cells of a state office's name rows, the headings, the page furniture's title and
    date, and the court pages' prose are turned into text; ballot designation rows are never joined."""
    pdf = PDF(open(path, "rb").read())
    pages = pdf.pages()
    if len(pages) < 2:
        raise SystemExit("California (state races): the Certified List of Candidates has no pages after its certificate")
    cover = " ".join(join(rs) for _y, rs in pdf_rows(pdf, *pages[0]))
    m = re.search(r"Dated at Sacramento, California, this (\d{1,2})\s*(?:st|nd|rd|th)? day of ([A-Z][a-z]+), (\d{4})", cover)
    dated = dt.date(int(m.group(3)), C.MONTHS[m.group(2)], int(m.group(1))).isoformat() if m and m.group(2) in C.MONTHS else ""
    if "names of the justices" not in cover:
        raise SystemExit("California (state races): the certificate no longer says the list names the justices")
    top = [join(rs) for _y, rs in pdf_rows(pdf, *pages[1]) if min(r[0] for r in rs) < 50][:4]
    if len(top) < 4 or not re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}", top[2]) or not top[3].startswith("Page "):
        raise SystemExit("California (state races): the Certified List of Candidates' page heading was not found where it was")
    mo, day, yr = (int(x) for x in top[2].split("/"))
    published = dt.date(yr, mo, day).isoformat()
    title = f"{top[1]}, {top[0]}"
    problems = [] if dated in ("", published) else [f"the certificate is dated {dated} and the pages {published}"]
    body, courts, office = [], [], None
    for page, res in pages[1:]:
        for _y, rs in pdf_rows(pdf, page, res):
            if min(r[0] for r in rs) < 50:      # title, date, page number and the footer: page furniture
                continue
            size = max(r[2] for r in rs)
            if size >= 13:
                office = join(rs)
                if re.match(r"(Supreme Court|Court of Appeal)\b", office):
                    courts.append({"heading": office, "counties": [], "items": []})
            elif 8.5 <= size <= 9.5:
                if office and not office.startswith("United States ") and contest_race(office) is not None:
                    body.append((office, C.cells(rs)))
                elif office and office.startswith("United States "):
                    continue          # the federal rows: the federal loader's
                else:
                    problems.append(f"9-point rows under a heading this loader does not read ({office!r}; not read)")
            elif 11.5 <= size <= 12.5 and courts and courts[-1]["heading"] == office:
                s = courts[-1]
                cl = C.cells(rs)
                texts = [join(x) for _x, x in cl]
                if len(cl) == 2 and texts[1].startswith("For ") and len(texts[0]) <= 2:
                    s["items"].append([texts[1], ""])
                elif not s["items"]:
                    s["counties"].append(" ".join(texts))
                else:
                    s["items"][-1][1] += " " + " ".join(texts)
            elif office and not office.startswith("United States "):
                problems.append(f"a row of {size:g}-point type under {office} (not read)")
    pairs = collections.Counter((round(row[0][0]), round(row[-1][0])) for _o, row in body if len(row) == 2)
    if not pairs:
        raise SystemExit("California (state races): no state office's row has a name and a party preference")
    name_x, party_x = pairs.most_common(1)[0][0]
    out, odd = collections.defaultdict(list), collections.Counter()
    for office, row in body:
        if abs(row[0][0] - name_x) <= 1.5 and len(row) == 2 and abs(row[1][0] - party_x) <= 1.5:
            name = join(row[0][1])
            out[ws(office)].append((name.rstrip("* ").strip(), join(row[1][1]), name.endswith("*")))
        elif row[0][0] - name_x > 4 and all(abs(x - party_x) > 1.5 for x, _rs in row):
            continue      # the ballot designation beneath a name: never turned into text
        else:
            odd[office] += 1
    problems += [f"{n} row(s) under {o} were neither a name nor a ballot designation (not read)" for o, n in sorted(odd.items())]
    return title, published, dated, dict(out), courts, problems


def court_races(courts, counties, problems):
    """The retention questions -> [(race dict, candidate name, incumbent, note)]."""
    out = []
    for s in courts:
        h = s["heading"].replace("�", "-").replace("–", "-")
        if re.fullmatch(r"Supreme Court - For all 58 Counties", h):
            dnum = None
        else:
            m = re.fullmatch(r"Court of Appeal\s*-\s*(\w+) Appellate District", h)
            if not m or m.group(1) not in ORDINAL_WORDS:
                raise SystemExit(f"California (state races): a court heading that is not read ({h!r})")
            dnum = ORDINAL_WORDS[m.group(1)]
        cids = None
        if dnum:
            words = [w.strip() for w in " ".join(s["counties"]).split(",") if w.strip()]
            geo = []
            for w in words:
                g = counties.get(county_key(w))
                if g:
                    geo.append(g[0])
                else:
                    problems.append(f"{h}: a county the Census file does not name ({w})")
            cids = sorted(geo) or None
        for head, q in s["items"]:
            q = ws(q)
            if dnum is None:
                fm = re.fullmatch(r"For (Associate Justice|Chief Justice) of the Supreme Court", head)
                if not fm:
                    raise SystemExit(f"California (state races): a Supreme Court office that is not read ({head!r})")
                office_title, division = fm.group(1), None
            else:
                fm = re.fullmatch(r"For ((?:Administrative )?Presiding Justice|Associate Justice), Court of Appeal, (\w+) District"
                                  r"(?:, Division (\w+))?", head)
                if not fm or ORDINAL_WORDS.get(fm.group(2)) != dnum or (fm.group(3) and fm.group(3) not in DIVISION_WORDS):
                    raise SystemExit(f"California (state races): a Court of Appeal office that is not read ({head!r})")
                office_title, division = fm.group(1), fm.group(3)
            qm = re.fullmatch(r"Shall (?P<pre>.*?)\s*(?P<name>(?:[A-Z][A-Z.'\-]*\s)*[A-Z][A-Z.'\-]*) be elected to the office for the "
                              r"(?P<term>term provided by law|unexpired term provided by law)\?", q)
            if not qm or not name_parts(qm.group("name"))[1]:
                raise SystemExit(f"California (state races): a retention question that is not read ({head})")
            pre, name = ws(qm.group("pre")), ws(qm.group("name"))
            family = re.sub(r"[^A-Z]", "", name_parts(name)[1].upper())
            wanted = office_title + (" of the Supreme Court" if dnum is None else "")
            inc = int(pre == wanted)
            if dnum is None:
                rid = f"2026-{STATE}-SC-{family}"
                race = dict(level="court", office_kind="supreme_court", office=f"{office_title} of the Supreme Court", jurisdiction=NAME,
                            jurisdiction_id=STATE, county_ids=None, district=None, seat=None)
            else:
                dv = DIVISION_WORDS[division] if division else None
                rid = f"2026-{STATE}-COA{dnum}-" + (f"{dv}-" if dv else "") + family
                dname = f"{h.split('-', 1)[1].strip()}"
                race = dict(level="court", office_kind="court_of_appeals", office=f"{office_title}, Court of Appeal",
                            jurisdiction=dname + (f", Division {division}" if division else ""), jurisdiction_id=f"COA{dnum}",
                            county_ids=json.dumps(cids) if cids else None, district=str(dnum), seat=f"Division {division}" if division else None)
            race.update(race_id=rid, _district_name=None if dnum is None else h.split("-", 1)[1].strip(), _county_list=cids)
            said = f"Shall {pre + ' ' if pre else ''}{name} be elected to the office for the {qm.group('term')}?"
            note = [f"The ballot asks: \"{said}\""]
            if not inc:
                note.append(f"The question names this person as \"{pre}\", not by the title of this office; the certified list does not "
                            "name the seat's present holder." if pre else
                            "The question gives no title before the name; the certified list does not name the seat's present holder.")
            if qm.group("term").startswith("unexpired"):
                note.append("The question is for the rest of a term (\"the unexpired term\").")
            note.append("The certified list prints the name in capitals; it is shown as printed.")
            out.append((race, name, inc, " ".join(note)))
    return out


# ------------------------------------------------------------------------------------------------ the roster

def roster(path):
    """Sitting legislators and the statewide officials the roster carries: ids, names, party, chamber, district and
    office only (the roster's contact columns are never selected)."""
    con = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    legs = [dict(zip(("id", "first", "last", "full", "other", "party", "chamber", "district"), r)) for r in con.execute(
        "SELECT bioguide_id, first_name, last_name, official_full, other_names, party_name, chamber, district "
        "FROM legislators WHERE is_current = 1")]
    offs = [dict(zip(("id", "first", "last", "full", "office", "label", "party"), r)) for r in con.execute(
        "SELECT bioguide_id, first_name, last_name, official_full, office, office_label, party_name FROM officials")]
    con.close()
    for p in legs:
        p["district"] = str(int(p["district"])) if str(p["district"] or "").isdigit() else str(p["district"] or "")
    return legs, offs


def person_fits(name, p):
    """A listed name fits a roster person: same family name and a given name that fits (the roster's first name or any
    other form of the name it keeps)."""
    readings = [name_parts(name), name_parts(re.sub(r'"[^"]*"|\([^)]*\)', " ", name))]
    forms = [(fold(p["first"] or "").split(), " ".join(fold(p["last"] or "").split()))]
    if p.get("full"):
        forms.append(name_parts(p["full"]))
    forms += [name_parts(o.strip()) for o in (p.get("other") or "").split(";") if o.strip() and "," not in o and "." not in o]
    return any(f[1] and r[1] and fits(r, f) for r in readings for f in forms)


def party_letter(roster_party):
    t = (roster_party or "").lower()
    return "D" if t.startswith("democrat") else "R" if t.startswith("republican") else "O"


# ------------------------------------------------------------------------------------------------ loading

def load(db_path, say=print, cache=CACHE, roster_db=ROSTER, county_zip=COUNTY_ZIP):
    net.patient_lookups()
    folder = os.path.join(cache, "ca")
    os.makedirs(folder, exist_ok=True)
    report = []

    # ---- the files
    csv_path = os.path.join(cache, CSV_FILE)
    net.download(C.URL, csv_path, max_age_days=60, say=say)
    cert_path = os.path.join(cache, CERT_FILE)
    net.download(C.CERT, cert_path, max_age_days=30, say=say)
    book_paths = {}
    for fname in [v[4] for v in STATEWIDE.values()] + list(DISTRICT_BOOKS):
        p = os.path.join(folder, BOOK_FILE.format(fname))
        net.download(BOOK_BASE + fname, p, max_age_days=365, say=say)
        if open(p, "rb").read(2) != b"PK":
            raise SystemExit(f"California (state races): {os.path.basename(p)} is not a workbook (a web page?); delete it and run again")
        book_paths[fname] = p
    if not os.path.exists(county_zip):
        net.download(COUNTY_URL, county_zip, max_age_days=3650, say=say)
    counties = census_counties(county_zip)
    if len(counties) != COUNTIES:
        report.append(f"the Census county file names {len(counties)} California counties, not {COUNTIES}")

    primary, county_names, _blank = read_csv_export(csv_path)
    books, book_counts = {}, {}
    for fname, p in book_paths.items():
        got, probs = read_book(p, fname)
        report += probs
        book_counts[fname] = sum(len(v) for v in got.values())
        for contest, cands in got.items():
            if contest in books:
                raise SystemExit(f"California (state races): {contest} is in two office workbooks")
            books[contest] = cands
    title, published, dated, cert, courts, cert_probs = certified(cert_path)
    report += cert_probs
    legs, offs = roster(roster_db)
    geo_of = {}
    for cid, n in county_names.items():
        g = counties.get(county_key(n))
        if not g:
            report.append(f"a county in the Statement of Vote the Census file does not name ({n})")
        else:
            geo_of[n] = g[0]
    if len(county_names) != COUNTIES:
        report.append(f"the Statement of Vote's state contests name {len(county_names)} counties, not {COUNTIES}")

    # ---- reconcile: every candidate's county rows in the CSV workbook against the office workbooks
    unreconciled, party_pairs = [], collections.Counter()
    if set(primary) != set(books):
        report.append(f"contests in the CSV workbook and the office workbooks differ (CSV only: {sorted(set(primary) - set(books))}; "
                      f"workbooks only: {sorted(set(books) - set(primary))})")
    for contest in sorted(set(primary) & set(books), key=str):
        f, b = primary[contest], books[contest]
        by_name = {book_name(c[0]): (cid, c) for cid, c in f["cands"].items()}
        if len(by_name) != len(f["cands"]):
            report.append(f"{contest}: two candidates share a name in the CSV workbook")
        if set(by_name) != {book_name(n) for n in b}:
            unreconciled.append(f"{contest}: candidates differ (CSV only: {sorted(set(by_name) - {book_name(n) for n in b})}; "
                                f"workbook only: {sorted({book_name(n) for n in b} - set(by_name))})")
        for n, bc in b.items():
            hit = by_name.get(book_name(n))
            if not hit:
                continue
            cid, c = hit
            if bc["total"] != c[4]:
                unreconciled.append(f"{contest} {c[0]}: the CSV county rows add to {c[4]:,}, the workbook's totals row says {bc['total']}")
            csv_county = {k[0]: v for k, v in f["county"].items() if k[1] == cid}
            if csv_county != bc["county"]:
                diff = sorted(k for k in set(csv_county) | set(bc["county"]) if csv_county.get(k) != bc["county"].get(k))
                unreconciled.append(f"{contest} {c[0]}: county figures differ in {', '.join(diff[:5])}")
            wcode, wi = re.fullmatch(r"(.+?)(\s*\(W/I\))?", bc["code"]).groups()
            party_pairs[(c[1], wcode)] += 1
            if bool(wi) != c[3]:
                unreconciled.append(f"{contest} {c[0]}: the write-in mark differs between the CSV workbook and the office workbook")
    for (words, wcode), n in sorted(party_pairs.items()):
        if BOOK_PARTY.get(wcode) != words:
            report.append(f"party preference: the CSV workbook writes {words!r} where the office workbook writes {wcode!r} ({n} candidates)")

    # ---- the races
    races, fields = {}, {}
    contests = sorted(set(primary) | set(cert), key=lambda c: contest_race(c)[0])
    offs_by = {o["office"]: o for o in offs}
    for contest in contests:
        rid, level, kind, office, juris, jid, district, partisan = contest_race(contest)
        notes, h = [], None
        if level == "statewide" and kind != "board_of_equalization":
            ro = STATEWIDE[contest][3]
            h = offs_by.get(ro) if ro else None
            notes.append(FOUR_YEARS)
            if kind == "lieutenant_governor":
                notes.append(LTG_NOTE)
            if ro is None:
                notes.append(NO_ROSTER)
            elif h is None:
                report.append(f"{rid}: the roster's officials table has no {ro}")
                notes.append(NO_HOLDER)
        elif kind == "board_of_equalization":
            notes += [BOE_NOTE, NO_ROSTER.replace("this office", "the Board of Equalization")]
        else:
            chamber = "Senate" if kind == "state_senate" else "House"
            hs = [p for p in legs if p["chamber"] == chamber and p["district"] == district]
            h = hs[0] if len(hs) == 1 else None
            if len(hs) > 1:
                report.append(f"{rid}: {len(hs)} sitting members in the roster for this seat; none is taken as the holder")
            notes.append(SENATE_NOTE if kind == "state_senate" else ASSEMBLY_NOTE)
            if h is None:
                report.append(f"{rid}: no sitting member in the roster for this seat")
                notes.append(NO_HOLDER)
        notes.append(TOP_TWO if partisan else SPI_NOTE)
        cids = None
        if level != "statewide" or kind == "board_of_equalization":
            cids = sorted({geo_of[k[0]] for k in primary.get(contest, {"county": {}})["county"] if k[0] in geo_of}) or None
            if cids:
                notes.append(f"The counties listed for this district are the {len(cids)} in which its June 2 primary was counted "
                             "(the Statement of Vote), a derived list.")
        races[rid] = dict(race_id=rid, state=STATE, level=level, office_kind=kind, office=office, jurisdiction=juris, jurisdiction_id=jid,
                          county_ids=json.dumps(cids) if cids else None, district=district, seat=None, special=0, partisan=partisan,
                          holder_id=h["id"] if h else None, holder_name=h["full"] if h else None, holder_party=h["party"] if h else None,
                          election_date=GENERAL, note=notes, _holder=h, _contest=contest, _roster=(kind in ("state_senate", "state_house")
                                                                                                     or (level == "statewide" and kind != "board_of_equalization"
                                                                                                         and STATEWIDE[contest][3])))

    # ---- which seats: every Senate seat on the lists even-numbered, all 20 of them; all 80 Assembly seats
    sen = sorted((r["district"] for r in races.values() if r["office_kind"] == "state_senate"), key=int)
    if sen != SENATE_UP:
        report.append(f"the Senate seats on the lists are {sen}, not the 20 even-numbered districts")
    asm = sorted((r["district"] for r in races.values() if r["office_kind"] == "state_house"), key=int)
    if asm != ASSEMBLY:
        report.append(f"Assembly seats missing from the lists: {sorted(set(ASSEMBLY) - set(asm), key=int)}")
    for contest in STATEWIDE:
        if contest not in cert:
            report.append(f"{contest}: not on the certified list")
    for d in range(1, 5):
        if f"Board of Equalization Member District {d}" not in cert:
            report.append(f"Board of Equalization District {d}: not on the certified list")

    def where(p):
        if "chamber" in p:
            return f"the California {'Senate' if p['chamber'] == 'Senate' else 'Assembly'}, District {p['district']}"
        return f"California's {p['label']}"

    def identify(race, name, pcode):
        """(incumbent, state_member_id, note): the seat's holder when the name fits; else a sitting legislator or official
        of a compatible party anywhere when the name fits exactly one."""
        h = race["_holder"]
        if h is not None and person_fits(name, h):
            return 1, h["id"], None
        pool = [p for p in legs + offs if person_fits(name, p) and (pcode not in ("D", "R") or party_letter(p["party"]) in (pcode, "O"))]
        if len(pool) == 1 and not (h is not None and pool[0]["id"] == h["id"]):
            p = pool[0]
            return 0, p["id"], (f"Serves today in {where(p)}." if "chamber" in p else f"Serves today as {where(p)}.")
        return 0, None, None

    # ---- the candidates: the June 2 field, and the November ballot checked against the certified list
    cands, singles, agree, differ, single_rids = [], 0, 0, [], []
    for rid in sorted(races):
        race = races[rid]
        contest = race["_contest"]
        partisan = race["partisan"]
        f = primary.get(contest)
        top = []
        if f is None:
            report.append(f"{rid}: on the certified list but not in the June 2 Statement of Vote")
        else:
            ranked = sorted(f["cands"].values(), key=lambda c: (-c[4], c[0]))
            total = sum(c[4] for c in ranked)
            second = ranked[1][4] if len(ranked) > 1 else ranked[0][4]
            if not partisan and ranked and total and ranked[0][4] * 2 > total:
                report.append(f"{rid}: {ranked[0][0]} won more than half the June vote; this nonpartisan office should not be on the "
                              "November ballot")
            top = [c for c in ranked if c[4] >= second]
            if len(ranked) >= 2:
                fields[rid] = len(ranked)
                for c in ranked:
                    pcode = code(c[1]) if partisan else "N"
                    if partisan and pcode == "N":
                        report.append(f"{rid}: {c[0]} is nonpartisan in a partisan contest")
                    if not partisan and fold(c[1]) not in ("non partisan", "nonpartisan"):
                        report.append(f"{rid}: {c[0]} has a party preference in a nonpartisan contest")
                    inc, mid, n2 = identify(race, c[0], pcode)
                    if race["_roster"] and bool(inc) != c[2]:
                        report.append(f"{rid} primary: {c[0]}: the Statement of Vote's incumbent flag is {'Y' if c[2] else 'N'}, the roster "
                                      f"match says {'the holder' if inc else 'not the holder'}")
                    note = [n2] if n2 else []
                    if c[3]:
                        note.append("A write-in candidate in the June 2 primary; the Statement of Vote counts this candidate's votes by name.")
                    if not race["_roster"] and c[2]:
                        inc = 1
                        note.append("The Statement of Vote marks this candidate as the incumbent.")
                    cands.append([rid, "primary", PRIMARY, c[0], shown_party(c[1]) if partisan else NONPARTISAN, pcode, None, inc, int(c[3]),
                                  c[4], round(100 * c[4] / total, 2) if total else None, "advanced" if c[4] >= second else "lost", mid,
                                  SRC_CSV, " ".join(note) or None])
            else:
                singles += 1
                single_rids.append(rid)
                race["note"].insert(-1, SINGLE)
        listed = cert.get(contest, [])
        if not listed:
            report.append(f"{rid}: no candidates on the certified list")
        key = lambda n, p: (name_parts(n)[1], fold(shown_party(p)))
        if sorted(key(c[0], c[1]) for c in top) == sorted(key(n, p) for n, p, _i in listed):
            agree += 1
            same = True
        else:
            same = False
            differ.append(rid)
            report.append(f"{rid}: the Statement of Vote's top two were {'; '.join(f'{c[0]} ({c[1]})' for c in top) or 'nobody'}; the "
                          f"certified list prints {'; '.join(f'{n} ({p})' for n, p, _i in listed) or 'nobody'} (the certified list's names are used)")
        by_key = {key(c[0], c[1]): c for c in top}
        for n, p, mark in listed:
            pcode = code(p) if partisan else "N"
            if partisan and pcode == "N":
                report.append(f"{rid}: {n} is nonpartisan on the certified list in a partisan contest")
            if not partisan and pcode != "N":
                report.append(f"{rid}: {n} has a party preference on the certified list in a nonpartisan contest")
            inc, mid, n2 = identify(race, n, pcode)
            note = [n2] if n2 else []
            if race["_roster"] and bool(inc) != mark:
                report.append(f"{rid}: {n}: the certified list's incumbent mark is {'*' if mark else 'absent'}, the roster match says "
                              f"{'the holder' if inc else 'not the holder'}")
            if not race["_roster"] and mark:
                inc = 1
                note.append("The Secretary of State's certified list marks this candidate as the incumbent.")
            sov = by_key.get(key(n, p))
            if not same:
                note.append("Named as the Secretary of State's certified list prints it, which differs from the June 2 Statement of "
                            "Vote's top two for this race.")
            elif sov and fold(sov[0]) != fold(n):
                note.append(f"The June 2 Statement of Vote writes the name {sov[0]}.")
            if same and sov and sov[3]:
                note.append("Ran as a write-in candidate in the June 2 primary and finished in the top two; the November ballot prints "
                            "the name.")
            cands.append([rid, "general", GENERAL, n, shown_party(p) if partisan else NONPARTISAN, pcode, None, inc, 0, None, None, None, mid,
                          SRC_CERT, " ".join(note) or None])
        race["note"].append(ORDER_NOTE)

    # ---- the courts
    court = court_races(courts, counties, report)
    for race, name, inc, note in court:
        rid = race["race_id"]
        if rid in races:
            raise SystemExit(f"California (state races): two retention questions share the id {rid}")
        races[rid] = dict(race, state=STATE, special=0, partisan=0, holder_id=None, holder_name=None, holder_party=None,
                          election_date=GENERAL, note=[COURT_NOTE], _holder=None, _contest=None, _roster=False)
        cands.append([rid, "general", GENERAL, name, NONPARTISAN, "N", None, inc, 0, None, None, None, None, SRC_CERT, note])

    # ---- the last look at every stored text: nothing that looks like a contact detail reaches the database
    for c in cands:
        if CONTACT.search(c[3]) or CONTACT.search(c[4] or "") or (c[14] and CONTACT_NOTE.search(c[14])):
            raise SystemExit(f"California (state races): a stored cell for {c[0]} failed the contact-detail check (not shown)")
    for r in races.values():
        if CONTACT_NOTE.search(" ".join(r["note"])) or CONTACT.search(r["holder_name"] or ""):
            raise SystemExit(f"California (state races): the note or holder for {r['race_id']} failed the contact-detail check (not shown)")
    keys = [(c[0], c[1], c[3]) for c in cands]
    if len(keys) != len(set(keys)):
        raise SystemExit("California (state races): two candidate rows share race, election and name")
    listed_rows = sum(len(v) for v in cert.values()) + len(court)
    general = [c for c in cands if c[1] == "general"]
    if listed_rows != len(general):
        raise SystemExit(f"California (state races): {listed_rows} names read from the certified list, {len(general)} stored")

    # ---- places: counties, the districts on the ballot, the appellate districts
    place_rows = [("county", g, f"{n} County", json.dumps([g]), SRC_COUNTY) for g, n in sorted(counties.values())]
    seen_places = set()
    for rid, r in sorted(races.items()):
        kind = {"state_senate": "senate", "state_house": "house", "board_of_equalization": "boe", "court_of_appeals": "appellate"}.get(r["office_kind"])
        if not kind:
            continue
        pid = f"{STATE}-{r['district']}" if kind in ("senate", "house") else f"{STATE}-{r['jurisdiction_id']}"
        if (kind, pid) in seen_places:
            continue
        seen_places.add((kind, pid))
        pname = r["jurisdiction"] if kind != "appellate" else r["_district_name"]
        place_rows.append((kind, pid, pname, r["county_ids"], SRC_CERT if kind == "appellate" else SRC_CSV))

    # ---- write California's rows only, in one transaction
    cols = ["race_id", "state", "level", "office_kind", "office", "jurisdiction", "jurisdiction_id", "county_ids", "district", "seat",
            "special", "partisan", "holder_id", "holder_name", "holder_party", "election_date", "note"]
    race_rows = [tuple(" ".join(races[r]["note"]) or None if k == "note" else races[r][k] for k in cols) for r in sorted(races)]
    prim = [c for c in cands if c[1] == "primary"]
    when = dt.date.fromisoformat(published).strftime("%B %d, %Y").replace(" 0", " ")
    sources = [
        (SRC_CERT, STATE, "official candidate list", AGENCY, title, C.CERT, published, mtime(cert_path), sha(cert_path), listed_rows,
         f"Found on {C.CERT_PAGE} as \"Certified List of Candidates\"; the Secretary's certificate is dated at Sacramento, "
         f"{dt.date.fromisoformat(dated).strftime('%B %d, %Y').replace(' 0', ' ') if dated else '(date not read)'}. For the state "
         "offices only the name and party preference cells of each name row are read (an asterisk marks an incumbent); ballot "
         "designations are never turned into text. The court pages' retention questions are read in full (each justice's name and "
         f"title, and each appellate district's counties): {len(court)} questions. Checked race by race against the June 2 Statement "
         f"of Vote's top two by family name and party preference: {agree} of {agree + len(differ)} races agree"
         + (f"; where they differ ({', '.join(differ)}) the certified list's names are used." if differ else ".")
         + " The list gives no ballot order."),
        (SRC_CSV, STATE, "official results", AGENCY, "Statement of Vote, June 2, 2026 Primary Election: CSV Files - Voter Nominated",
         C.URL, "", mtime(csv_path), sha(csv_path), len(prim),
         f"Found on {C.PAGE}. The state contests' county rows (name, party preference, incumbent and write-in flags, votes) added "
         f"up per candidate. A field is shown only where two or more candidates were on the June ballot ({len(fields)} fields; "
         f"{singles} races had one candidate). "
         + ("Every candidate's county figures and totals agree with the Statement of Vote's office workbooks."
            if not unreconciled else "Did not reconcile: " + "; ".join(unreconciled) + ".")),
    ]
    for fname, p in book_paths.items():
        sources.append((f"ca-sos-2026-sl-sov-{fname.rsplit('.', 1)[0]}", STATE, "official results", AGENCY,
                        f"Statement of Vote, June 2, 2026 Primary Election: {fname}", BOOK_BASE + fname, "", mtime(p), sha(p), book_counts[fname],
                        f"Linked from {C.PAGE}. Read only to reconcile the CSV workbook: each candidate's county rows and State or "
                        "District Totals (names, party codes and votes only)."))
    sources += [
        (SRC_ROSTER, STATE, "member roster", "Open States people project (CC0), as loaded into state_ca.sqlite",
         "Sitting California legislators and statewide officials", "https://github.com/openstates/people", "", mtime(roster_db), sha(roster_db),
         len(legs) + len(offs), "Who holds each seat today (legislators by chamber and district; the Governor, Lieutenant Governor, "
         "Attorney General and Secretary of State, the only statewide offices the roster carries), and which candidate is a sitting "
         "member (same seat, the name fits, one fit only; a member of another seat only when the name fits exactly one sitting "
         "legislator or official of a compatible party). Ids, names, party, chamber and district only."),
        (SRC_COUNTY, STATE, "county codes", "U.S. Census Bureau", "Cartographic boundary file, counties, 2024 (1:500,000)", COUNTY_URL, "",
         mtime(county_zip), sha(county_zip), len(counties), "California's 58 counties: names and GEOIDs only, matched by name."),
    ]

    con = sqlite3.connect(db_path)
    con.executescript(SCHEMA)
    with con:
        con.execute("DELETE FROM sl_candidates WHERE race_id IN (SELECT race_id FROM sl_races WHERE state = ?)", (STATE,))
        con.execute("DELETE FROM sl_candidates WHERE race_id LIKE ?", (f"2026-{STATE}-%",))
        con.execute("DELETE FROM sl_races WHERE state = ?", (STATE,))
        con.execute("DELETE FROM sl_sources WHERE state = ?", (STATE,))
        con.execute("DELETE FROM sl_places WHERE source_id LIKE 'ca-%' OR (kind = 'county' AND id GLOB '06[0-9][0-9][0-9]') OR id GLOB 'CA-*'")
        con.executemany(f"INSERT INTO sl_races VALUES ({','.join('?' * len(cols))})", race_rows)
        con.executemany("INSERT INTO sl_candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", cands)
        con.executemany("INSERT INTO sl_sources VALUES (?,?,?,?,?,?,?,?,?,?,?)", sources)
        con.executemany("INSERT INTO sl_places VALUES (?,?,?,?,?)", place_rows)
    con.close()

    by = lambda kinds, rows_: sum(1 for c in rows_ if races[c[0]]["office_kind"] in kinds)
    nk = lambda kinds: sum(1 for r in races.values() if r["office_kind"] in kinds)
    sw = tuple(v[1] for v in STATEWIDE.values())
    say(f"    California: {len(races)} races ({nk(sw)} statewide, {nk(('board_of_equalization',))} Board of Equalization, "
        f"{nk(('state_senate',))} Senate, {nk(('state_house',))} Assembly, {nk(('supreme_court', 'court_of_appeals'))} retention questions); "
        f"{len(general)} on the November list (statewide {by(sw, general)}, Board of Equalization {by(('board_of_equalization',), general)}, "
        f"Senate {by(('state_senate',), general)}, Assembly {by(('state_house',), general)}, justices "
        f"{by(('supreme_court', 'court_of_appeals'), general)}); {len(fields)} top-two primary fields, {len(prim)} primary rows (statewide "
        f"{by(sw, prim)}, Board of Equalization {by(('board_of_equalization',), prim)}, Senate {by(('state_senate',), prim)}, Assembly "
        f"{by(('state_house',), prim)}); the certified list ({when}) agrees with the top two in {agree} of {agree + len(differ)} races; "
        + ("every primary figure reconciled with the office workbooks" if not unreconciled else f"{len(unreconciled)} figures did not reconcile"))
    for c in cands:
        if c[12] and not c[7]:
            say(f"      matched elsewhere: {c[0]} {c[1]}: {c[3]} -> {c[12]}")
    for line in unreconciled:
        say(f"      check: did not reconcile: {line}")
    for line in report:
        say(f"      check: {line}")
    return len(general)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: python -m ballot.state_local_ca <database>")
    load(sys.argv[1])
