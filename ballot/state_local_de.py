"""
ballot/state_local_de.py - Delaware's state and county races on the November 3, 2026 ballot: all 41 seats of the House of
Representatives, the 11 State Senate seats up this year, and the three statewide offices Delaware elects in the midterm
year (Attorney General, State Treasurer, Auditor of Accounts), with the September 15 party primaries and their official
votes; and the county offices on the same ballot (see "The county level" below). Written into ballot_local_2026.sqlite
(never ballot_2026.sqlite), Delaware's rows only.

Which seats. Representatives serve two years, so every Representative District is on the ballot. Senators serve four
years, staggered: the Department of Elections' official 2024 general election results (reports/GE2024.csv) name the ten
Senate districts elected in 2024 (2, 3, 4, 6, 10, 11, 16, 17, 18, 21), and the other eleven elect a senator in 2026. The
November list must name exactly those eleven; a 2024 district on it would be a special election and stops the loader
until it is read. The Governor, Lieutenant Governor and Insurance Commissioner were elected in 2024 and are not on this
ballot; if one appears on the list the loader stops.

Sources, all the Department of Elections' own (elections.delaware.gov), the same lists the federal loader
(ballot/lists/de.py) reads for Congress:
  - "Filed Candidates by Office, General Election 11/3/2026" (candidates/candidatelist/genl_fcddt_2026.xlsx): one row per
    filing, every office on the ballot. Only County, Office, Withdrawal Date, BallotName, Party and DisplayedStatus are
    read, by heading. The workbook also carries the filer's name parts, filing date, residential and mailing addresses,
    website, two e-mail addresses and two telephone numbers: none of those cells is ever turned into text, printed or
    kept. The workbook is held in memory only; the kept cells of the state-office rows go to
    ballot_cache/de/sl_de_2026_general_list.json with the file's SHA-256, and the other rows are only counted there
    (federal, county; the county part below keeps the county rows' kept cells in a file of its own). For the state part
    the page's "Last Updated" line is read for the date and nothing else.
  - "Write-In Candidates by Office, General Election 11/3/2026" (genl_wcddt_2026.xlsx), read the same way (it has no
    Party column). Delaware counts write-in votes only for declared write-in candidates, and these are they.
  - "Filed Candidates by Office, Primary Election 9/15/2026" (prim_fcddt_2026.xlsx), read the same way: who was on each
    party's primary ballot (contested primaries only) and how their names are printed.
  - The official primary results behind results/enr/PR2026.html: Election_StatewideResults_ID_PR2026.json (each
    candidate's machine, absentee, early-voting and total votes; every row must say OFFICIAL RESULTS with every election
    district reported) and, as a control, Election_ByCountyWithWilmington_ID_PR2026.json. Both are the files the federal
    loader caches in ballot_cache/de/ (names and votes only), shared, not fetched twice.
  - The official 2024 general election results (reports/GE2024.csv): the office column only, for which Senate districts
    were elected in 2024. Kept as the list of districts and the file's SHA-256.
  - The official 2022 general election report (reports/GE2022.html), the last election with every Senate and
    Representative District on the ballot under today's lines (drawn in 2021): its "By Election District" section says
    which election districts make up each district's contest, and the three county sheriff contests say which county each
    election district lies in, so each district's counties come from the record, not from a map. Only contest titles and
    election district labels are read; the kept result is ballot_cache/de/sl_de_2022_district_counties.json with the
    page's SHA-256. Control: the primary's county file must show votes only in counties the district touches.
  - Who holds each seat today, from state_de.sqlite (the Open States roster the state pages use): ids, names, parties,
    chambers, districts and offices only. The roster carries the Attorney General but no Treasurer or Auditor.
  - The Census Bureau's cartographic county file (cb_2024_us_county_500k, in states_cache/census/) for the three county
    GEOIDs: names and GEOIDs only.

What is stored. 2026-DE-SS<n> (State Senator), 2026-DE-SH<n> (State Representative), 2026-DE-AG, 2026-DE-TREAS,
2026-DE-AUD. The November candidates are the list's rows marked Qualified (or Provisional, noted: the State Election
Commissioner is still reviewing the filing). A row marked Withdrawn, or carrying a withdrawal date, is left off and
counted; any other status stops the loader. The list gives no ballot positions and no general-election sample ballots are
posted yet, so the list's own order within each contest is kept as the ballot order, as on the federal side, after checking
it is the pattern Delaware's lists follow (Democratic, then Republican, then other parties); a contest that departs from it
is stored with no ballot order and reported. Parties are kept as printed ("Ind Pty of DE" written out as Independent Party
of Delaware). Declared write-ins are stored with write_in 1, no ballot position and the party "Write-in", as on the federal
side. A party primary is a field when two or more of its candidates were on the ballot (Delaware's primary list names only
those); the results' candidates must be exactly the list's, each candidate's machine, absentee and early votes must add up
to the total, the county figures must add up to the state figure, and the leader is the nominee, who must be the party's
candidate on the November list or the row says so. No write-in votes are reported, so a field's total is the sum of its
candidates' votes.

A candidate is the incumbent (incumbent 1, state_member_id) only when the name fits exactly one sitting member of the same
chamber and district; the Attorney General the same way. A sitting legislator running for another office is given the
roster id (not incumbent) when the name fits exactly one legislator of the same party.

The county level
----------------
One Department of Elections runs every state and county election in Delaware, so the same two lists carry the county
contests: in 2026 the sheriff and recorder of deeds of each of the three counties, the register of wills of New Castle
and Sussex, New Castle County Council districts 1 to 6, the Kent County Levy Court's at-large seat and districts 2, 4 and
6, and Sussex County Council districts 4 and 5 (20 contests, all partisan). No city, town, school board or special
district office is on the November ballot: school boards are elected on the second Tuesday in May (14 Del. C. 1072(c)),
cities and towns on dates set by their charters, and Wilmington's city offices with the President (next in 2028).

  - The county rows of the same two Excel files (general and declared write-ins), read the same way: the six kept
    columns by heading, the workbook in memory only, and only the county rows' kept cells written to
    ballot_cache/de/local/ (with the sheet's row number, so a problem can be named by row and never by its text). A kept
    cell that reads like contact details is blanked and counted, never printed. Within a run each workbook and page is
    asked for once, whichever part reads it first.
  - Each list's own page (genl_fcddt_2026.html, genl_wcddt_2026.html), as a second reading: the page's table is read for
    the cells under Office, County, Party and Status only, and its county rows must come to the same counts as the
    workbook's. The Candidate cell (name with contact details) and the Filed cell are never turned into text, and the
    page is not kept.
  - The Department's "Schedule of Elections: Elected Office Table 2026-2034" (a PDF: office titles and a mark under each
    year an office is elected): the control on which county contests there should be. SCHEDULE_2026 below is that table
    as read on 2026-09-30; each run reads the table again and says so if it no longer agrees. A scheduled office with no
    row on the list is kept as a race with no candidate and a note; an office on the list that the schedule does not
    have for 2026, or a title this loader does not know, is not guessed at: it goes to sl_gaps.
  - The Department's 2026 Municipal Elections Calendar (a PDF): only the election dates are read, to confirm none is
    November 3; town names and websites are not kept. And the Department's general election page, read only for
    whether the sample ballots (which would give ballot order) are posted yet.

What is stored for the county level. sl_races: level county, jurisdiction the county ("Kent County", id its 5-digit
code), office_kind county_council (New Castle, Sussex), county_commissioner (Kent's Levy Court commissioners),
county_recorder, register_of_wills, sheriff; district the council or Levy Court district number, seat "At Large" for
Kent's at-large commissioner; race_id 2026-DE-<county code>-<office kind>[-D<n> or -AL]. sl_candidates: the rows marked
Qualified or Provisional, the party as the list prints it, no ballot order (the list states none), incumbent 0, no
member id; declared write-ins with write_in 1 and the party "Write-in"; withdrawn rows left off and counted. sl_gaps
and sl_notes (ballot.check_local.EXTRA_SCHEMA; local_calendar and local_coverage) are rewritten for Delaware on every
run. The county party primaries of September 15 are not loaded.

    python -m ballot.state_local_de <path to a test database>
"""

import csv
import datetime as dt
import hashlib
import html as H
import io
import json
import os
import re
import sqlite3
import sys
import time
import zipfile
from urllib.error import HTTPError

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from ballot import pdftext  # noqa: E402
from ballot.check_local import EXTRA_SCHEMA, contact_like  # noqa: E402
from ballot.common import CACHE, fold, name_parts, party_code  # noqa: E402
from ballot.lists import de as DE  # noqa: E402
from ballot.match import fits  # noqa: E402
from states import net  # noqa: E402

STATE, FIPS, NAME = "DE", "10", "Delaware"
GENERAL, PRIMARY = "2026-11-03", DE.PRIMARY
ROSTER = os.path.join(HERE, "state_de.sqlite")
COUNTY_ZIP = os.path.join(HERE, "states_cache", "census", "cb_2024_us_county_500k.zip")
COUNTY_URL = "https://www2.census.gov/geo/tiger/GENZ2024/shp/cb_2024_us_county_500k.zip"
REPORT_2022 = DE.BASE + "reports/GE2022.html"
CSV_2024 = DE.BASE + "reports/GE2024.csv"

SRC = {"general": "de-doe-2026-sl-general-list", "write_in": "de-doe-2026-sl-write-ins", "primary": "de-doe-2026-sl-primary-list",
       "results": "de-doe-2026-sl-primary-results", "county": "de-doe-2026-sl-primary-results-county",
       "ge2024": "de-doe-2024-general-results", "ge2022": "de-doe-2022-general-report", "roster": "de-openstates-roster",
       "census": "de-census-cb-2024-county"}
KEPT = {k: f"sl_de_2026_{k}_list.json" for k in DE.LISTS}
SEATS_FILE, SENATE_FILE = "sl_de_2022_district_counties.json", "sl_de_2024_senate_districts.json"

SCHEMA = """
CREATE TABLE IF NOT EXISTS sl_races (race_id TEXT PRIMARY KEY, state TEXT NOT NULL, level TEXT NOT NULL, office_kind TEXT NOT NULL,
  office TEXT NOT NULL, jurisdiction TEXT, jurisdiction_id TEXT, county_ids TEXT, district TEXT, seat TEXT, special INTEGER NOT NULL DEFAULT 0,
  partisan INTEGER NOT NULL, holder_id TEXT, holder_name TEXT, holder_party TEXT, election_date TEXT NOT NULL, note TEXT);
CREATE TABLE IF NOT EXISTS sl_candidates (race_id TEXT NOT NULL, election TEXT NOT NULL, election_date TEXT, name TEXT NOT NULL, party TEXT,
  party_code TEXT, ballot_order INTEGER, incumbent INTEGER NOT NULL DEFAULT 0, write_in INTEGER NOT NULL DEFAULT 0, votes INTEGER, pct REAL,
  outcome TEXT, state_member_id TEXT, source_id TEXT, note TEXT, PRIMARY KEY (race_id, election, name));
CREATE TABLE IF NOT EXISTS sl_sources (source_id TEXT PRIMARY KEY, state TEXT, kind TEXT, agency TEXT, title TEXT, url TEXT, published TEXT,
  fetched TEXT, sha256 TEXT, rows INTEGER, note TEXT);
CREATE TABLE IF NOT EXISTS sl_places (kind TEXT NOT NULL, id TEXT NOT NULL, name TEXT NOT NULL, county_ids TEXT, source_id TEXT,
  PRIMARY KEY (kind, id));
"""

KEEP = DE.KEEP                                      # County, Office, Withdrawal Date, BallotName, Party, DisplayedStatus
STATEWIDE = {   # office on the list: (race key, office_kind, roster office or None)
    "Attorney General": ("AG", "attorney_general", "attorney general"),
    "State Treasurer": ("TREAS", "state_treasurer", None),
    "Auditor of Accounts": ("AUD", "state_auditor", None),
}
OTHER_YEARS = ("Governor", "Lieutenant Governor", "Insurance Commissioner")
LEG = re.compile(r"State (Senator|Representative) District (\d+)")
LOCAL = re.compile(r"^(?:New Castle County|Kent County|Sussex County|City of |Town of )")
SENATE_SEATS, HOUSE_SEATS = 21, 41
COUNTIES = ("New Castle", "Kent", "Sussex")
SHERIFF = {"NCC Sheriff": "New Castle", "KC Sheriff": "Kent", "SC Sheriff": "Sussex"}
WRITTEN_OUT = DE.WRITTEN_OUT
CODES = DE.CODES
ON = DE.ON
PROVISIONAL, WRITE_IN = DE.PROVISIONAL, DE.WRITE_IN
RANK = {"Democratic": 0, "Republican": 1}

ORDER_NOTE = ("The Department of Elections' list gives no ballot positions; the order shown is the list's own (Democratic, then "
              "Republican, then other parties), as on the federal side.")
HOUSE_NOTE = "Representatives serve two-year terms; all 41 Representative Districts elect a representative in 2026."
NO_HOLDER = "The Open States roster this site uses does not carry the {office}, so today's holder is not shown."

# ---- the county level
LSRC = {"general": "de-doe-2026-local-general-list", "write_in": "de-doe-2026-local-write-ins",
        "general_page": "de-doe-2026-local-general-list-page", "write_in_page": "de-doe-2026-local-write-ins-page",
        "schedule": "de-doe-schedule-of-elections", "municipal": "de-doe-2026-municipal-calendar",
        "election_page": "de-doe-2026-general-election-page"}
LOCAL_KEPT = {"general": "sl_de_2026_general_county.json", "write_in": "sl_de_2026_write_in_county.json"}
SCHEDULE_FILE, MUNI_FILE, ELECTION_PAGE_FILE = "sl_de_schedule_of_elections.json", "sl_de_2026_municipal_calendar.json", "sl_de_2026_general_election_page.json"
SCHEDULE_URL = DE.BASE + "candidates/pdfs/Schedule_of_Elections_Table.pdf"
MUNI_URL = DE.BASE + "public/calendar/pdfs/2026ElectionCalendarMunicipalities.pdf"
ELECTION_PAGE = DE.BASE + "elections/general/general.html"
LIST_TITLES = {"general": "Filed Candidates by Office, General Election 11/3/2026", "write_in": "Write-In Candidates by Office, General Election 11/3/2026"}
COUNTY_TITLE = re.compile(r"(New Castle|Kent|Sussex) County (.+)")
# The county offices on the 2026 ballot: the Department's "Schedule of Elections: Elected Office Table 2026-2034" (updated
# 6/12/2025) as read on 2026-09-30, written the way the candidate lists title them after the county's name. Every run
# reads the table again and reports any difference.
SCHEDULE_2026 = {
    "New Castle": ["Recorder of Deeds", "Register of Wills", "Sheriff"] + [f"Council District {n}" for n in range(1, 7)],
    "Kent": ["Recorder of Deeds", "Sheriff", "Levy Court At-Large", "Levy Court District 2", "Levy Court District 4", "Levy Court District 6"],
    "Sussex": ["Recorder of Deeds", "Register of Wills", "Sheriff", "Council District 4", "Council District 5"],
}
ROW_OFFICES = {   # a county office elected by the whole county: (office_kind, the title in plain words)
    "Recorder of Deeds": ("county_recorder", "Recorder of Deeds"),
    "Register of Wills": ("register_of_wills", "Register of Wills"),
    "Sheriff": ("sheriff", "Sheriff"),
    "Clerk of the Peace": ("clerk_of_the_peace", "Clerk of the Peace"),                     # elected with the President, next in 2028
    "County Executive": ("county_executive", "County Executive"),
    "President of County Council": ("county_council_president", "President of County Council"),
}
LEVY_NOTE = ("The Levy Court is Kent County's governing body: seven commissioners, one chosen by the voters of each of six districts "
             "and one by the whole county.")
NO_ROW_NOTE = ("The Department of Elections' schedule has this office on the 2026 ballot, but its list of candidates names no one "
               "for it.")
ONLY_WRITE_IN_NOTE = ("No name is printed on the ballot for this office: the Department of Elections' list names only declared "
                      "write-in candidates.")
NO_PARTY = "No party listed"
MARK = "●"                                      # the round mark the schedule prints under a year
MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")
DATE_LINE = re.compile(r"\b(?:Mon|Tues|Wednes|Thurs|Fri|Satur|Sun)day,?\s+(" + "|".join(MONTHS) + r")\s+(\d{1,2}),?\s+(20\d\d)\b")
# the page builder's own, wider test for a street address (it also counts Court and Place), so no text written here is dropped there
BUILDER_STREET = re.compile(r"\b\d{1,6}\s+(?:[NSEW]\.?\s+)?[A-Za-z0-9.' -]{1,40}?\s(?:St|Street|Ave|Avenue|Rd|Road|Blvd|Boulevard|Dr|Drive|Ln|Lane|Way|Ct|Court|"
                            r"Cir|Circle|Pkwy|Parkway|Hwy|Highway|Trl|Trail|Pl|Place|Ter|Terrace)\b\.?", re.I)


class NotRead(Exception):
    """A control source that could not be read as this loader knows it. The message names the source and the check,
    never a line of it."""


# ------------------------------------------------------------------------------------------------ fetching and caching

_MEMO = {}                                          # the candidate lists' workbooks and pages, asked for once a run


def fetch(url, accept="*/*", say=print, memo=False):
    """One request, asked at most twice more on a refusal or a server error (never past a challenge page). With memo, the
    answer is kept in memory for the rest of the run, so the state part and the county part share one download."""
    if memo and url in _MEMO:
        return _MEMO[url]
    for attempt in range(3):
        try:
            raw = net.get(url, accept=accept)
        except HTTPError as e:
            if e.code not in (403, 429, 500, 502, 503, 504) or attempt == 2:
                raise
            say(f"      {url.rsplit('/', 1)[1]}: HTTP {e.code}; asking again in {10 * (attempt + 1)} s")
            time.sleep(10 * (attempt + 1))
            continue
        head = raw[:4000].lower()
        if b"captcha" in head or b"challenge-platform" in head or b"incapsula" in head:
            raise SystemExit(f"Delaware (state races): {url} answered with a bot check; it was not worked around. A person in a "
                             "browser would have to fetch it.")
        if memo:
            _MEMO[url] = raw
        return raw


def fresh(path, days):
    return os.path.exists(path) and time.time() - os.path.getmtime(path) < days * 86400


def kept(path, max_age_days, make, say=print):
    """A small JSON of what is read from a source, refreshed when older than max_age_days; on a failed refresh the older
    copy is used."""
    if fresh(path, max_age_days):
        return json.load(open(path, encoding="utf-8"))
    try:
        data = make()
    except (HTTPError, OSError) as e:
        if os.path.exists(path):
            say(f"      could not refresh {os.path.basename(path)} ({e}); using the copy read earlier")
            return json.load(open(path, encoding="utf-8"))
        raise
    os.makedirs(os.path.dirname(path), exist_ok=True)
    json.dump(data, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return data


# ------------------------------------------------------------------------------------------------ what an office is

def office_of(office):
    """'federal', 'local', or a dict describing a state office on Delaware's lists and results; anything else stops."""
    office = DE.squash(office)
    if office in DE.FEDERAL or DE.LOOKS_FEDERAL.search(office):
        return "federal"
    if LOCAL.search(office):
        return "local"
    if office in STATEWIDE:
        key, kind, _r = STATEWIDE[office]
        return {"race_id": f"2026-{STATE}-{key}", "level": "statewide", "office_kind": kind, "office": office,
                "jurisdiction": NAME, "jurisdiction_id": FIPS, "district": None, "chamber": None}
    m = LEG.fullmatch(office)
    if m:
        d, senate = str(int(m.group(2))), m.group(1) == "Senator"
        if not 1 <= int(d) <= (SENATE_SEATS if senate else HOUSE_SEATS):
            raise SystemExit(f"Delaware (state races): a district that does not exist ({office!r})")
        return {"race_id": f"2026-{STATE}-{'SS' if senate else 'SH'}{d}", "level": "legislature",
                "office_kind": "state_senate" if senate else "state_house",
                "office": "State Senator" if senate else "State Representative",
                "jurisdiction": f"{'Senate' if senate else 'Representative'} District {d}", "jurisdiction_id": f"{STATE}-{d}",
                "district": d, "chamber": "Senate" if senate else "House"}
    if any(office.startswith(o) for o in OTHER_YEARS):
        raise SystemExit(f"Delaware (state races): {office!r} is on a 2026 list, but Delaware elects it in presidential years; read it")
    raise SystemExit(f"Delaware (state races): an office this loader does not know ({office!r})")


# ------------------------------------------------------------------------------------------------ the candidate lists

def candidate_list(key, say=print):
    """The kept cells of the state-office rows of one candidate list's Excel file, the other rows only counted, and the
    workbook's SHA-256. The workbook is read in memory; the cells outside KEEP are never turned into text."""
    import openpyxl
    page, book, heading = DE.LISTS[key]
    url = DE.LISTS_PAGE + book
    raw = fetch(url, say=say, memo=True)
    if raw[:2] != b"PK":
        raise SystemExit(f"Delaware (state races): {url} is not an Excel workbook")
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True)
    if len(wb.worksheets) != 1:
        raise SystemExit(f"Delaware (state races): {book} has {len(wb.worksheets)} sheets; one was expected")
    rows = wb.worksheets[0].iter_rows(values_only=True)
    heads = [DE.squash(h) for h in next(rows)]
    need = [k for k in KEEP if not (k == "Party" and key == "write_in")]
    if any(k not in heads for k in need):
        raise SystemExit(f"Delaware (state races): {book}'s columns changed ({[k for k in need if k not in heads]} missing)")
    idx = {k: heads.index(k) for k in need}
    out, counts, total = [], {"federal": 0, "local": 0, "state": 0}, 0
    for r in rows:
        if not r or not any(v not in (None, "") for v in r):
            continue
        total += 1
        office = DE.cell(r[idx["Office"]]) if idx["Office"] < len(r) else ""
        kind = office_of(office)
        if isinstance(kind, str):
            counts[kind] += 1
            continue
        counts["state"] += 1
        out.append({k: DE.cell(r[i]) if i < len(r) else "" for k, i in idx.items()})
    wb.close()
    time.sleep(1.5)
    try:
        published = last_updated(key, say)
    except (HTTPError, OSError):
        published = ""
    time.sleep(1.5)
    return {"url": url, "page": DE.LISTS_PAGE + page, "published": published, "fetched": dt.date.today().isoformat(),
            "sha256": hashlib.sha256(raw).hexdigest(), "rows_in_file": total, "counts": counts, "rows": out}


def list_page(key, say=print):
    """One candidate list's own page, asked for once a run: (the bytes as fetched, the text), its heading checked."""
    page, _book, heading = DE.LISTS[key]
    raw = fetch(DE.LISTS_PAGE + page, accept="text/html", say=say, memo=True)
    text = raw.decode("utf-8-sig", "replace")
    if heading not in text:
        raise SystemExit(f"Delaware: {DE.LISTS_PAGE + page} is no longer headed {heading!r}")
    return raw, text


def last_updated(key, say=print):
    """The "Last Updated" date printed on a candidate list's page (YYYY-MM-DD), or ""."""
    m = re.search(r"Last Updated:\s*(\d{4}-\d\d-\d\d)", list_page(key, say)[1])
    return m.group(1) if m else ""


# ------------------------------------------------------------------------------------------------ which seats, which counties

def senate_2024(say=print):
    """The Senate districts on the 2024 general election ballot, from the office column of the official results CSV."""
    raw = fetch(CSV_2024, say=say)
    time.sleep(1.5)
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig", "replace"))))
    heads = [h.strip().lower() for h in rows[0]]
    if "office" not in heads or "electionname" not in heads:
        raise SystemExit("Delaware (state races): the 2024 general results CSV's columns changed")
    io_, ie = heads.index("office"), heads.index("electionname")
    names = {DE.squash(r[ie]) for r in rows[1:] if len(r) > ie}
    if not names or not all("2024" in n and "General" in n for n in names):
        raise SystemExit(f"Delaware (state races): {CSV_2024} is not the 2024 general election ({sorted(names)[:3]})")
    ds = sorted({int(m.group(1)) for r in rows[1:] if len(r) > io_ for m in [re.fullmatch(r"State Senator District (\d+)", DE.squash(r[io_]))] if m})
    if not ds:
        raise SystemExit("Delaware (state races): the 2024 general results name no Senate district")
    return {"url": CSV_2024, "fetched": dt.date.today().isoformat(), "sha256": hashlib.sha256(raw).hexdigest(),
            "rows": len(rows) - 1, "districts": ds}


def heading_text(s):
    return DE.squash(H.unescape(re.sub(r"<[^>]+>", " ", s)))


def district_counties(say=print):
    """{"SS": {district: [county names]}, "SH": {...}} from the official 2022 general election report's "By Election
    District" section: each district contest's election districts, and which sheriff contest (one per county) each
    election district voted in. Only contest titles and election district labels are read."""
    raw = fetch(REPORT_2022, accept="text/html", say=say)
    time.sleep(1.5)
    page = raw.decode("utf-8", "replace")
    if "2022 General Election Report" not in page:
        raise SystemExit(f"Delaware (state races): {REPORT_2022} is no longer the 2022 general election report")
    tabs = [(m.start(), heading_text(m.group(1))) for m in re.finditer(r'<h2 class="resp-accordion[^"]*"[^>]*>(.*?)</h2>', page, re.S)]
    starts = [i for i, (_p, t) in enumerate(tabs) if t.startswith("By Election District")]
    if len(starts) != 1:
        raise SystemExit("Delaware (state races): the 2022 report has no single \"By Election District\" section")
    a = tabs[starts[0]][0]
    b = tabs[starts[0] + 1][0] if starts[0] + 1 < len(tabs) else len(page)
    seg, contest, eds = page[a:b], None, {}
    for m in re.finditer(r'<h3 class="contest-title[^"]*"[^>]*>(.*?)</h3>|<h4 class="electiondistrict-title[^"]*"[^>]*>(.*?)</h4>', seg, re.S):
        if m.group(1) is not None:
            contest = heading_text(m.group(1))
            eds.setdefault(contest, set())
            continue
        label = heading_text(m.group(2))
        e = re.fullmatch(r"Election District (\d\d)-(\d\d)", label)
        if not e or contest is None:
            raise SystemExit(f"Delaware (state races): an election district label in the 2022 report that is not read ({label!r})")
        eds[contest].add(f"{e.group(1)}-{e.group(2)}")
    county_of = {}
    for contest, county in SHERIFF.items():
        if not eds.get(contest):
            raise SystemExit(f"Delaware (state races): the 2022 report has no {contest} contest")
        for ed in eds[contest]:
            if ed in county_of:
                raise SystemExit(f"Delaware (state races): election district {ed} is in two county contests of the 2022 report")
            county_of[ed] = county
    statewide = eds.get("Attorney General", set())
    if set(county_of) != statewide:
        raise SystemExit(f"Delaware (state races): the three county contests do not cover the state's election districts "
                         f"({len(set(county_of) ^ statewide)} differ)")
    out = {"SS": {}, "SH": {}}
    for d in range(1, SENATE_SEATS + 1):
        got = eds.get(f"State Senator District {d}")
        if not got:
            raise SystemExit(f"Delaware (state races): the 2022 report has no State Senator District {d} contest")
        out["SS"][str(d)] = sorted({county_of[e] for e in got}, key=COUNTIES.index)
    for d in range(1, HOUSE_SEATS + 1):
        got = eds.get(f"State Representative District {d}")
        if not got:
            raise SystemExit(f"Delaware (state races): the 2022 report has no State Representative District {d} contest")
        if any(int(e[:2]) != d for e in got):
            raise SystemExit(f"Delaware (state races): Representative District {d}'s contest lists another district's election districts")
        out["SH"][str(d)] = sorted({county_of[e] for e in got}, key=COUNTIES.index)
    # every election district is in exactly one Senate district and one Representative district
    for kind, word in (("State Senator", "Senate"), ("State Representative", "Representative")):
        seen = {}
        for c, s in eds.items():
            if c.startswith(kind + " District "):
                for e in s:
                    seen[e] = seen.get(e, 0) + 1
        if set(seen) != statewide or any(v != 1 for v in seen.values()):
            raise SystemExit(f"Delaware (state races): the 2022 {word} contests do not divide the state's election districts exactly")
    say(f"      2022 report: {len(statewide)} election districts read into {SENATE_SEATS} Senate and {HOUSE_SEATS} Representative districts")
    return {"url": REPORT_2022, "fetched": dt.date.today().isoformat(), "sha256": hashlib.sha256(raw).hexdigest(),
            "election_districts": len(statewide), "SS": out["SS"], "SH": out["SH"]}


def census_counties():
    """{county name: (GEOID, NAMELSAD)} for Delaware from the Census Bureau's cartographic county file."""
    import shapefile                                   # pyshp
    if not os.path.exists(COUNTY_ZIP):
        net.download(COUNTY_URL, COUNTY_ZIP, max_age_days=3650)
    z = zipfile.ZipFile(COUNTY_ZIP)
    base = next(n[:-4] for n in z.namelist() if n.endswith(".dbf"))
    rdr = shapefile.Reader(dbf=io.BytesIO(z.read(base + ".dbf")))
    fields = [f[0] for f in rdr.fields[1:]]
    out = {}
    for rec in rdr.iterRecords():
        rec = dict(zip(fields, rec))
        if str(rec.get("STATEFP")) == FIPS:
            out[str(rec["NAME"])] = (str(rec["GEOID"]), str(rec["NAMELSAD"]))
    if sorted(out) != sorted(COUNTIES):
        raise SystemExit(f"Delaware (state races): the county file gives {sorted(out)}, not Delaware's three counties")
    return out


# ------------------------------------------------------------------------------------------------ the primary results

def primary_results(statewide, counties, report):
    """{(race_id, code): [(name as reported, votes, percentage)]} for every state party primary, with the report time and
    the counties each contest had votes in; every check made."""
    out, times, local = {}, set(), 0
    for r in statewide:
        if r.get("Election Id") != DE.ELECTION_ID or r.get("Election Date") != PRIMARY:
            raise SystemExit(f"Delaware (state races): the results file is for {r.get('Election Name')} {r.get('Election Date')}, not the 2026 primary")
        if r.get("Results Type") != "OFFICIAL RESULTS":
            raise SystemExit(f"Delaware (state races): the primary results say {r.get('Results Type')!r}; unofficial figures are never stored")
        if r.get("Precincts Reported") != r.get("Total Precincts"):
            raise SystemExit(f"Delaware (state races): {r.get('Contest Title')} has {r.get('Precincts Reported')} of {r.get('Total Precincts')} election districts reported")
        times.add(r.get("ReportTime"))
        kind = office_of(r.get("Contest Title"))
        if kind == "local":
            local += 1
        if isinstance(kind, str):
            continue
        code, name = DE.code_of(r.get("Party Name")), DE.squash(r.get("Candidate Name"))
        total = DE.votes_of(r["Total Votes"])
        parts = sum(DE.votes_of(r[k]) for k in ("Machine Votes", "Absentee Votes", "Early Voting Votes"))
        if parts != total:
            raise SystemExit(f"Delaware (state races): {name}'s machine, absentee and early votes ({parts}) do not add up to the total ({total})")
        if re.search(r"write[- ]?in", name, re.I):
            raise SystemExit(f"Delaware (state races): write-in votes in the {kind['race_id']} {code} results; read how they are reported")
        field = out.setdefault((kind["race_id"], code), [])
        if any(nkey(n) == nkey(name) for n, _v, _p in field):
            raise SystemExit(f"Delaware (state races): {name} appears twice in the {kind['race_id']} {code} results")
        field.append((name, total, r.get("Percentage")))
    if len(times) != 1:
        raise SystemExit(f"Delaware (state races): the statewide results carry more than one report time ({sorted(map(str, times))})")
    # control: the county file adds up to the same totals; and the counties with votes
    seen, where = {}, {}
    for r in counties:
        if r.get("Election Id") != DE.ELECTION_ID or r.get("Results Type") != "OFFICIAL RESULTS":
            raise SystemExit("Delaware (state races): the county results file is not the official 2026 primary results")
        kind = office_of(r.get("Contest Title"))
        if isinstance(kind, str):
            continue
        name = DE.squash(r.get("Candidate Name"))
        nc, kent, sussex, state = (DE.votes_of(r[k]) for k in ("New Castle", "Kent", "Sussex", "State"))
        if nc + kent + sussex != state or DE.votes_of(r["Wilmington"]) + DE.votes_of(r["Rest of New Castle"]) != nc:
            raise SystemExit(f"Delaware (state races): {name}'s county votes in {kind['race_id']} do not add up to the state figure")
        seen[(kind["race_id"], DE.code_of(r.get("Party Name")), nkey(name))] = state
        w = where.setdefault(kind["race_id"], set())
        w.update(c for c, v in zip(COUNTIES, (nc, kent, sussex)) if v)
    mine = {(race, code, nkey(n)): v for (race, code), field in out.items() for n, v, _p in field}
    if seen != mine:
        raise SystemExit(f"Delaware (state races): the county results differ from the statewide results "
                         f"({sorted(set(seen.items()) ^ set(mine.items()))[:4]})")
    return out, times.pop(), where, local


# ------------------------------------------------------------------------------------------------ names and the roster

def nkey(text):
    """Letters only, one space between words: a name as the lists and the results both write it."""
    return " ".join(fold(text).split())


def roster(path):
    """Sitting legislators and the statewide officials the roster carries: ids, names, parties, chambers, districts and
    offices only (the roster's contact columns are never selected)."""
    con = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    legs = [dict(zip(("id", "first", "last", "full", "other", "party", "chamber", "district"), r)) for r in con.execute(
        "SELECT bioguide_id, first_name, last_name, official_full, other_names, party_name, chamber, district "
        "FROM legislators WHERE is_current = 1")]
    offs = [dict(zip(("id", "first", "last", "full", "office", "party"), r)) for r in con.execute(
        "SELECT bioguide_id, first_name, last_name, official_full, office, party_name FROM officials")]
    con.close()
    for p in offs:
        p["other"] = ""
    return legs, offs


def readings(name):
    """The (given names, family name) readings of a name as printed: with the quoted nickname set aside, and the nickname
    alone as a given name."""
    g, f = name_parts(re.sub(r'"[^"]*"', " ", name))
    out = [(g, f)]
    out += [(fold(n).split(), f) for n in re.findall(r'"([^"]*)"', name) if fold(n)]
    return out


def person_fits(reads, p):
    forms = [(fold(p["first"] or "").split(), " ".join(fold(p["last"] or "").split()))]
    forms += [name_parts(o.strip()) for o in (p.get("other") or "").split(";") if o.strip() and "," not in o and "." not in o]
    return any(f[1] and r[1] and fits(r, f) for r in reads for f in forms)


def chamber_words(p):
    return f"Delaware {'Senate, Senate' if p['chamber'] == 'Senate' else 'House of Representatives, Representative'} District {p['district']}"


# ------------------------------------------------------------------------------------------------ the county level

def reads_like_contact(text, strict=True):
    """The trial check's own test for contact details and, in strict mode, the page builder's wider street test."""
    t = str(text if text is not None else "")
    return bool(t) and (contact_like(t, strict) or bool(strict and BUILDER_STREET.search(t)))


def rows_digest(rows):
    """A fingerprint of rows' kept cells, to tell whether two downloads of a list carry the same rows."""
    return hashlib.sha256(json.dumps([[r.get(k, "") for k in KEEP] for r in rows], ensure_ascii=False).encode("utf-8")).hexdigest()


def long_date(iso):
    y, m, d = (int(x) for x in iso.split("-"))
    return f"{MONTHS[m - 1]} {d}, {y}"


def plural(n, word, many=None):
    return f"{n} {word if n == 1 else (many or word + 's')}"


def and_words(items):
    items = list(items)
    return (", ".join(items[:-1]) + (" and " if len(items) > 1 else "") + items[-1]) if items else ""


def county_office(county, words):
    """What the words after a county's name in an office title mean: (office_kind, the office in plain words, district,
    seat), or None for words this loader does not know. The candidate lists write "Levy Court District 2", the
    schedule "Levy Court Commissioner District 2"."""
    words = DE.squash(words)
    if words in ROW_OFFICES:
        return ROW_OFFICES[words] + (None, None)
    m = re.fullmatch(r"Council District (\d+)", words)
    if m and county in ("New Castle", "Sussex"):
        return "county_council", "County Council Member", str(int(m.group(1))), None
    m = re.fullmatch(r"Levy Court(?: Commissioner)? District (\d+)", words)
    if m and county == "Kent":
        return "county_commissioner", "Levy Court Commissioner", str(int(m.group(1))), None
    if county == "Kent" and re.fullmatch(r"Levy Court(?: Commissioner)? At[- ]Large", words):
        return "county_commissioner", "Levy Court Commissioner", None, "At Large"
    return None


def county_contest(title, ctab):
    """A county office title as the Department's lists write it ("Kent County Levy Court District 2") -> the contest's
    fields, or None for a title this loader does not know."""
    m = COUNTY_TITLE.fullmatch(DE.squash(title))
    got = county_office(m.group(1), m.group(2)) if m else None
    if not got:
        return None
    kind, office, district, seat = got
    geoid, full = ctab[m.group(1)]
    rid = f"2026-{STATE}-{geoid}-{kind.replace('_', '-')}" + (f"-D{district}" if district else "-AL" if seat else "")
    return {"race_id": rid, "county": m.group(1), "geoid": geoid, "jurisdiction": full, "office_kind": kind, "office": office,
            "district": district, "seat": seat, "title": DE.squash(title)}


def page_reading(key, say=print):
    """A candidate list's own page, read as a second count of the list: the "Last Updated" date, how many rows the page's
    table has, and its county rows counted by the cells under Office, County, Party and Status. Only those cells are
    turned into text; the Candidate cell (the name with contact details) and the Filed cell never are, and the page is
    not kept."""
    raw, text = list_page(key, say)
    m = re.search(r"Last Updated:\s*(\d{4}-\d\d-\d\d)", text)
    tables = [t for t in re.finditer(r"<table\b([^>]*)>(.*?)</table>", text, re.S | re.I) if "no-screen" not in t.group(1)]
    if len(tables) != 1:
        raise NotRead(f"the {key} list's page shows {len(tables)} tables on screen; one was expected")
    body = tables[0].group(2)
    heads = [heading_text(h) for h in re.findall(r"<th\b[^>]*>(.*?)</th>", body, re.S | re.I)]
    want = [h for h in ("Office", "County", "Party", "Status") if h in heads]
    if not {"Office", "County", "Status"} <= set(want) or len(set(heads)) != len(heads):
        raise NotRead(f"the {key} list's page no longer has the headings Office, County and Status, once each")
    idx = {h: heads.index(h) for h in want}
    total, county = 0, {}
    for n, tr in enumerate(re.findall(r"<tr\b[^>]*>(.*?)</tr>", body, re.S | re.I), 1):
        cells = re.split(r"<td\b[^>]*>", tr, flags=re.I)[1:]
        if not cells:
            continue                                       # the heading row
        if len(cells) != len(heads):
            raise NotRead(f"row {n} of the {key} list's page has {len(cells)} cells under {len(heads)} headings")
        total += 1
        got = {h: heading_text(cells[i]) for h, i in idx.items()}
        if LOCAL.search(got["Office"]):
            k = "|".join(got.get(h, "") for h in ("Office", "County", "Party", "Status"))
            county[k] = county.get(k, 0) + 1
    return {"published": m.group(1) if m else "", "sha256": hashlib.sha256(raw).hexdigest(), "rows": total, "county": county}


def local_list(key, say=print):
    """The county-office rows of one candidate list's Excel file: the kept cells only, each with the sheet's row number.
    The other rows are counted, and the state rows' kept cells are fingerprinted, so that this download can be compared
    with the one the state races were read from. The workbook is read in memory and the cells outside KEEP are never
    turned into text; a kept cell that reads like contact details is blanked and counted. Then the list's own page, as
    a second count."""
    import openpyxl
    page, book, _heading = DE.LISTS[key]
    url = DE.LISTS_PAGE + book
    asked = [u not in _MEMO for u in (url, DE.LISTS_PAGE + page)]      # a pause only after a request really made
    raw = fetch(url, say=say, memo=True)
    time.sleep(1.5 if asked[0] else 0)
    if raw[:2] != b"PK":
        raise SystemExit(f"Delaware (county races): {url} is not an Excel workbook")
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True)
    if len(wb.worksheets) != 1:
        raise SystemExit(f"Delaware (county races): {book} has {len(wb.worksheets)} sheets; one was expected")
    rows = wb.worksheets[0].iter_rows(values_only=True)
    heads = [DE.squash(h) for h in next(rows)]
    need = [k for k in KEEP if not (k == "Party" and key == "write_in")]
    if any(k not in heads for k in need):
        raise SystemExit(f"Delaware (county races): {book}'s columns changed ({[k for k in need if k not in heads]} missing)")
    idx = {k: heads.index(k) for k in need}
    out, state_rows, counts, total, blanked = [], [], {"federal": 0, "local": 0, "state": 0}, 0, 0
    for n, r in enumerate(rows, 2):                        # the sheet's own row number: row 1 holds the headings
        if not r or not any(v not in (None, "") for v in r):
            continue
        total += 1
        rec = {k: DE.cell(r[i]) if i < len(r) else "" for k, i in idx.items()}
        kind = office_of(rec["Office"])
        if isinstance(kind, dict):
            counts["state"] += 1
            state_rows.append(rec)
            continue
        counts[kind] += 1
        if kind != "local":
            continue
        for k, v in rec.items():
            if v and (contact_like(v, True) or (k == "BallotName" and ("@" in v or re.search(r"\d{3,}", v)))):
                rec[k] = ""
                blanked += 1
        rec["row"] = n
        out.append(rec)
    wb.close()
    try:
        seen = page_reading(key, say)
    except (HTTPError, OSError, NotRead) as e:
        seen = {"error": str(e)}
    time.sleep(1.5 if asked[1] else 0)
    return {"url": url, "page": DE.LISTS_PAGE + page, "published": seen.get("published", ""), "fetched": dt.date.today().isoformat(),
            "sha256": hashlib.sha256(raw).hexdigest(), "rows_in_file": total, "counts": counts, "state_digest": rows_digest(state_rows),
            "blanked": blanked, "rows": out, "page_read": seen}


def pdf_rows(raw, what):
    """Every printed row of a PDF as pdftext's runs, top to bottom, page after page."""
    if raw[:5] != b"%PDF-":
        raise NotRead(f"{what} is not a PDF any more")
    try:
        pdf = pdftext.PDF(raw)
        return [runs for page, res in pdf.pages() for _y, runs in pdftext.rows(pdf, page, res)]
    except Exception as e:      # noqa: BLE001  a file pdftext cannot take apart is a control not read, not a stop
        raise NotRead(f"{what} could not be taken apart ({type(e).__name__})") from None


def schedule_table(rows):
    """[(section, office title, [years])] and the "Updated" date from the Schedule of Elections table's printed rows. A
    section's heading row ends with the table's years; an office's row carries a round mark under each year the office
    is elected, and a mark belongs to the year whose printed place is nearest it."""
    out, section, centres, updated = [], None, None, ""
    for runs in rows:
        runs = sorted(runs, key=lambda r: r[0])
        text = DE.squash(pdftext.join(runs))
        m = re.fullmatch(r"Updated (\d{1,2})/(\d{1,2})/(\d{4})", text)
        if m:
            updated = f"{m.group(3)}-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
            continue
        m = re.fullmatch(r"(.+?) ((?:20\d\d ){1,9}20\d\d)", text)
        if m:
            years = m.group(2).split()
            boxes = [(x0 + (x1 - x0) * i / len(t), x0 + (x1 - x0) * (i + 1) / len(t), ch)
                     for x0, _y, _size, t, x1 in runs for i, ch in enumerate(t) if not ch.isspace()]
            digits = boxes[-4 * len(years):]
            if "".join(b[2] for b in digits) != "".join(years):
                raise NotRead("the Schedule of Elections: a heading's years are not its last printed characters")
            section = m.group(1)
            centres = [(int(years[i]), (digits[4 * i][0] + digits[4 * i + 3][1]) / 2) for i in range(len(years))]
            continue
        marks = [r for r in runs if r[3].strip() == MARK]
        if not marks:
            continue
        if centres is None:
            raise NotRead("the Schedule of Elections: an office row stands before any heading with years")
        edge = min(c for _year, c in centres) - 12
        title = DE.squash(pdftext.join([r for r in runs if r[4] <= edge and r[3].strip() != MARK]))
        years = []
        for x0, _y, _size, _t, x1 in marks:
            year, c = min(centres, key=lambda yc: abs(yc[1] - (x0 + x1) / 2))
            if abs(c - (x0 + x1) / 2) > 12:
                raise NotRead("the Schedule of Elections: a mark does not stand under a year")
            years.append(year)
        if not title or len(set(years)) != len(years):
            raise NotRead("the Schedule of Elections: an office row with no title, or with two marks under one year")
        out.append((section, title, sorted(years)))
    return out, updated


def schedule(say=print):
    """The Department's Schedule of Elections table, cut down to the county offices (title, office kind, district or
    seat, the years marked) and the City of Wilmington's (title, years). Nothing else of the PDF is kept."""
    raw = fetch(SCHEDULE_URL, say=say)
    time.sleep(1.5)
    table, updated = schedule_table(pdf_rows(raw, "the Schedule of Elections"))
    offices, city = [], []
    for section, title, years in table:
        m = re.fullmatch(r"(New Castle|Kent|Sussex) County Offices", section or "")
        if m:
            county = m.group(1)
            got = county_office(county, re.sub(rf"^{county} County ", "", title))
            if not got or reads_like_contact(title):
                raise NotRead(f"the Schedule of Elections lists a {county} County office this loader does not know")
            if any(b - a != 4 for a, b in zip(years, years[1:])):
                raise NotRead("the Schedule of Elections: a county office's marked years are not four years apart")
            offices.append({"county": county, "title": title, "kind": got[0], "district": got[2], "seat": got[3], "years": years})
        elif section == "City of Wilmington":
            if reads_like_contact(title):
                raise NotRead("the Schedule of Elections: a City of Wilmington row that does not read like an office")
            city.append({"title": title, "years": years})
    if {o["county"] for o in offices} != set(COUNTIES):
        raise NotRead("the Schedule of Elections no longer has a section for each of the three counties")
    return {"url": SCHEDULE_URL, "fetched": dt.date.today().isoformat(), "sha256": hashlib.sha256(raw).hexdigest(), "updated": updated,
            "rows": len(table), "offices": offices, "wilmington": city}


def municipal_calendar(say=print):
    """The Department's 2026 Municipal Elections Calendar: the election dates only, and how many entries carry each. A
    line's other words (the town, its website, its county) are never kept."""
    raw = fetch(MUNI_URL, say=say)
    time.sleep(1.5)
    dates, updated, titled = {}, "", False
    for runs in pdf_rows(raw, "the municipal elections calendar"):
        text = DE.squash(pdftext.join(runs))
        titled = titled or "MUNICIPAL ELECTIONS CALENDAR" in text.upper()
        m = re.fullmatch(r"Updated (\d{1,2})/(\d{1,2})/(\d{4})", text)
        if m:
            updated = f"{m.group(3)}-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
        for m in DATE_LINE.finditer(text):
            iso = f"{m.group(3)}-{MONTHS.index(m.group(1)) + 1:02d}-{int(m.group(2)):02d}"
            dates[iso] = dates.get(iso, 0) + 1
    if not titled or not dates or any(not d.startswith("2026-") for d in dates):
        raise NotRead("the municipal elections calendar is no longer the 2026 calendar this loader knows")
    return {"url": MUNI_URL, "fetched": dt.date.today().isoformat(), "sha256": hashlib.sha256(raw).hexdigest(), "updated": updated,
            "dates": dict(sorted(dates.items()))}


def election_page(say=print):
    """The Department's general election page, read only for whether it still says that the general election's sample
    ballots are not yet available. Nothing else of the page is kept."""
    raw = fetch(ELECTION_PAGE, accept="text/html", say=say)
    time.sleep(1.5)
    text = heading_text(raw.decode("utf-8", "replace"))
    if "General Election" not in text or "November 3, 2026" not in text:
        raise NotRead("the general election page is no longer about November 3, 2026")
    return {"url": ELECTION_PAGE, "fetched": dt.date.today().isoformat(), "sha256": hashlib.sha256(raw).hexdigest(),
            "not_yet": bool(re.search(r"Sample Ballots are not yet available", text, re.I))}


def optional(path, days, make, what, report):
    """A control source's kept JSON: a fresh copy, else one read now, else the older copy, else None. The run's report
    says when it is not a fresh reading."""
    if fresh(path, days):
        return json.load(open(path, encoding="utf-8"))
    try:
        data = make()
    except (HTTPError, OSError, NotRead, SystemExit) as e:
        if os.path.exists(path):
            report.append(f"{what} could not be read again ({e}); the copy read earlier is used")
            return json.load(open(path, encoding="utf-8"))
        report.append(f"{what} could not be read ({e})")
        return None
    os.makedirs(os.path.dirname(path), exist_ok=True)
    json.dump(data, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return data


def local_level(folder, state_lists, ctab, say=print):
    """Delaware's county contests on the November lists, as rows ready to write (races, candidates, sources, gaps, notes)
    with the counts behind them and the lines to report. Nothing here touches the database."""
    report, gaps = [], []
    os.makedirs(folder, exist_ok=True)

    # ---- the contests there should be: the Department's schedule, as read on 2026-09-30 and read again now
    expected = {}
    for county in COUNTIES:
        for words in SCHEDULE_2026[county]:
            c = county_contest(f"{county} County {words}", ctab)
            expected[c["race_id"]] = c
    if len(expected) != sum(len(v) for v in SCHEDULE_2026.values()):
        raise SystemExit("Delaware (county races): two scheduled offices share a race id")
    sched = optional(os.path.join(folder, SCHEDULE_FILE), 30, lambda: schedule(say), "the Department's Schedule of Elections table", report)
    agrees, city_2026 = None, []
    if sched:
        theirs = {(o["county"], o["kind"], o["district"] or "", o["seat"] or "") for o in sched["offices"] if 2026 in o["years"]}
        mine = {(c["county"], c["office_kind"], c["district"] or "", c["seat"] or "") for c in expected.values()}
        agrees = theirs == mine
        if not agrees:
            report.append(f"the Department's Schedule of Elections no longer marks for 2026 the county offices this loader was written from "
                          f"({len(theirs - mine)} it has and this loader does not, {len(mine - theirs)} the other way): read it and bring "
                          "SCHEDULE_2026 and the calendar note up to date")
        city_2026 = [o["title"] for o in sched["wilmington"] if 2026 in o["years"]]
    muni = optional(os.path.join(folder, MUNI_FILE), 7, lambda: municipal_calendar(say), "the Department's municipal elections calendar", report)
    nov3 = muni["dates"].get(GENERAL, 0) if muni else None
    epage = optional(os.path.join(folder, ELECTION_PAGE_FILE), 2, lambda: election_page(say), "the Department's general election page", report)

    # ---- the county rows of the two lists (a failed first reading stops the loader, so earlier rows stay as they were)
    lists = {k: kept(os.path.join(folder, LOCAL_KEPT[k]), 2, lambda k=k: local_list(k, say), say) for k in ("general", "write_in")}
    words = {"general": "general", "write_in": "write-in"}
    contests = {rid: dict(c, printed=[], write_ins=0, withdrawn=0, unread=0, rows=0) for rid, c in expected.items()}
    counts = {k: 0 for k in ("rows", "placed", "withdrawn", "status", "off_schedule", "unknown", "nameless", "repeat")}
    off_schedule, unknown, cands, seen, provisional, withdrew = {}, {}, [], set(), 0, {"general": 0, "write_in": 0}
    for k in ("general", "write_in"):
        for r in lists[k]["rows"]:
            counts["rows"] += 1
            c = county_contest(r["Office"], ctab)
            if c is None:
                counts["unknown"] += 1
                unknown[r["Office"]] = unknown.get(r["Office"], 0) + 1
                continue
            if r["County"] != c["county"]:
                report.append(f"row {r['row']} of the {words[k]} list: the office is {c['jurisdiction']}'s, but the row's County cell says "
                              "otherwise; it is filed by the office's own title")
            if c["race_id"] not in contests:
                counts["off_schedule"] += 1
                off_schedule.setdefault(c["race_id"], dict(c, rows=0))["rows"] += 1
                continue
            have = contests[c["race_id"]]
            have["rows"] += 1
            status = r["DisplayedStatus"]
            if status == "Withdrawn" or r["Withdrawal Date"]:
                counts["withdrawn"] += 1
                have["withdrawn"] += 1
                withdrew[k] += 1
                continue
            if status not in ON:
                counts["status"] += 1
                have["unread"] += 1
                report.append(f"row {r['row']} of the {words[k]} list ({c['race_id']}): a status this loader does not know; the candidate is not stored")
                continue
            name = r["BallotName"]
            if not name:
                counts["nameless"] += 1
                report.append(f"row {r['row']} of the {words[k]} list ({c['race_id']}): no ballot name, or one that read like contact details and "
                              "was blanked; the candidate is not stored")
                continue
            if (c["race_id"], nkey(name)) in seen:
                counts["repeat"] += 1
                report.append(f"row {r['row']} of the {words[k]} list repeats a name already stored in {c['race_id']}; it is stored once")
                continue
            seen.add((c["race_id"], nkey(name)))
            notes = []
            if k == "general":
                party = WRITTEN_OUT.get(r.get("Party") or "", r.get("Party") or "")
                if not party:
                    party = NO_PARTY
                    notes.append("The Department of Elections' list gives no party for this candidate.")
                    report.append(f"row {r['row']} of the general list ({c['race_id']}): no party given")
                code, write_in, src = party_code(party), 0, LSRC["general"]
                have["printed"].append(party)
            else:
                party, code, write_in, src = "Write-in", "W", 1, LSRC["write_in"]
                notes.append(WRITE_IN)
                have["write_ins"] += 1
            if status == "Provisional":
                notes.append(PROVISIONAL)
                provisional += 1
            cands.append([c["race_id"], "general", GENERAL, name, party, code, None, 0, write_in, None, None, None, None, src, " ".join(notes) or None])
            counts["placed"] += 1
    if (counts["rows"] != sum(v for k, v in counts.items() if k != "rows") or counts["placed"] != len(cands)
            or len(cands) != len({(c[0], c[3]) for c in cands}) or counts["rows"] != len(lists["general"]["rows"]) + len(lists["write_in"]["rows"])):
        raise SystemExit("Delaware (county races): the county rows read and the rows accounted for differ; stopping")
    for rid, c in sorted(contests.items()):
        twice = sorted({p for p in c["printed"] if p in CODES and c["printed"].count(p) > 1})
        if twice:
            report.append(f"{rid}: two candidates of one party are printed ({', '.join(twice)}); both are stored")

    # ---- the second reading (each list's own page) and the download the state races came from
    second = {}
    for k in ("general", "write_in"):
        L, seen_page = lists[k], lists[k].get("page_read") or {}
        mine = {}
        for r in L["rows"]:
            key = "|".join((r["Office"], r["County"], r.get("Party", ""), r["DisplayedStatus"]))
            mine[key] = mine.get(key, 0) + 1
        if "county" not in seen_page:
            second[k] = None
            report.append(f"the {words[k]} list's page could not be read as a second count ({seen_page.get('error', 'not read')})")
        else:
            second[k] = seen_page["rows"] == L["rows_in_file"] and seen_page["county"] == mine and not L["blanked"]
            if not second[k]:
                report.append(f"the {words[k]} list's page and its Excel file differ: the page has {seen_page['rows']} rows, "
                              f"{sum(seen_page['county'].values())} of them county rows; the file {L['rows_in_file']} and {len(L['rows'])}"
                              + (f"; {L['blanked']} kept cells of the file were blanked" if L["blanked"] else ""))
        S = state_lists[k]
        if S["sha256"] != L["sha256"] and (rows_digest(S["rows"]) != L["state_digest"] or S["counts"] != L["counts"]):
            report.append(f"the {words[k]} list changed between the download the state races were read from ({S.get('fetched') or 'earlier'}) and the "
                          f"one the county races were read from ({L['fetched']}); the state rows are read again when their kept copy is two days old")

    # ---- races: every scheduled contest, with a note where the list has no printed name for it
    races, n_empty = [], 0
    for rid, c in sorted(contests.items()):
        notes = [LEVY_NOTE] if c["office_kind"] == "county_commissioner" else []
        if not c["rows"]:
            notes.append(NO_ROW_NOTE)
        elif not c["printed"] and c["write_ins"]:
            notes.append(ONLY_WRITE_IN_NOTE)
        elif not c["printed"]:
            left = ([plural(c["withdrawn"], "filing") + " withdrawn"] if c["withdrawn"] else []) + \
                   ([plural(c["unread"], "filing") + " with a status that is not read"] if c["unread"] else [])
            notes.append("The Department of Elections' list names no candidate printed on the ballot for this office"
                         + (f" ({' and '.join(left)})." if left else "."))
        n_empty += 0 if c["printed"] or c["write_ins"] else 1
        races.append((rid, STATE, "county", c["office_kind"], c["office"], c["jurisdiction"], c["geoid"], json.dumps([c["geoid"]]),
                      c["district"], c["seat"], 0, 1, None, None, None, GENERAL, " ".join(notes) or None))

    # ---- what could not be loaded
    G, W = lists["general"], lists["write_in"]
    for rid, o in sorted(off_schedule.items()):
        gaps.append((STATE, "race", rid, o["jurisdiction"], o["title"],
                     f"The Department of Elections' list of candidates carries this office ({plural(o['rows'], 'filing')}), but the Department's "
                     "schedule of elections does not have it on the 2026 ballot, so it may be an election to fill a vacancy. It is not shown "
                     "until the Department's notice of that election has been read.", G["page"]))
        report.append(f"{rid}: on the list but not on the schedule for 2026 ({plural(o['rows'], 'row')}); not stored, named in sl_gaps")
    for title, n in sorted(unknown.items()):
        gaps.append((STATE, "state", STATE, NAME, title or "an office with no title",
                     f"The Department of Elections' list of candidates carries this office ({plural(n, 'filing')}) under a title this loader "
                     "does not know how to file, so it is not shown until the title has been read.", G["page"]))
        report.append(f"an office title this loader does not know ({plural(n, 'row')}); not stored, named in sl_gaps")
    for rid, c in sorted(contests.items()):
        if c["unread"]:
            gaps.append((STATE, "race", rid, c["jurisdiction"], f"{plural(c['unread'], 'candidate')} for {c['title']}",
                         "The Department of Elections' list gives these filings a status other than Qualified, Provisional or Withdrawn. What "
                         "that status means for the ballot has not been read, so they are not shown.", G["page"]))
    if city_2026:
        gaps.append((STATE, "state", STATE, NAME, "City of Wilmington offices",
                     f"The Department's schedule of elections now marks {plural(len(city_2026), 'City of Wilmington office')} for 2026, which this "
                     "loader does not read from the list of candidates yet.", SCHEDULE_URL))
        report.append("the Schedule of Elections marks City of Wilmington offices for 2026; named in sl_gaps")
    if nov3:
        gaps.append((STATE, "state", STATE, NAME, "city and town elections on November 3",
                     f"The Department of Elections' municipal elections calendar lists {plural(nov3, 'city or town election')} on November 3, 2026. "
                     "Cities and towns run their own elections and take their own filings, so those contests are not on the Department's list "
                     "of candidates and are not shown here.", MUNI_URL))
        report.append(f"the municipal elections calendar lists {plural(nov3, 'election')} on November 3; named in sl_gaps")
    if epage and epage["not_yet"]:
        order_why = ("The Department of Elections' list of candidates gives no ballot positions, and the Department's general election page said "
                     "its sample ballots were not yet available when it was read; until they are posted, candidates are shown by surname.")
    else:
        order_why = ("The Department of Elections' list of candidates gives no ballot positions; the order is printed on the Department's sample "
                     "ballots, which are not read here, so candidates are shown by surname.")
        if epage:
            report.append("the Department's general election page no longer says its sample ballots are not yet available: ballot order may now "
                          "be readable from them")
    gaps.append((STATE, "state", STATE, NAME, "ballot order of county candidates", order_why, ELECTION_PAGE))

    # ---- the two notes
    towns = ("none of them on November 3 this year" if nov3 == 0 else "at other times of the year" if nov3 is None
             else f"{nov3} of them on November 3 this year, which the Department's list of candidates does not carry")
    calendar = ("On November 3, 2026 each of Delaware's three counties elects a sheriff and a recorder of deeds, New Castle and Sussex counties "
                "elect a register of wills, and part of each county's governing body is chosen: New Castle County Council districts 1 to 6, the "
                "Kent County Levy Court's at-large seat and districts 2, 4 and 6, and Sussex County Council districts 4 and 5; all are partisan "
                "offices. The other county offices and council seats, and the City of Wilmington's mayor, treasurer and council, are elected "
                "with the President, next in 2028. School board members are elected every year on the second Tuesday in May (May 12 in 2026), "
                f"and the other cities and towns hold their own elections on dates set by their charters, {towns}.")
    n_printed, n_write = sum(1 for c in cands if not c[8]), sum(1 for c in cands if c[8])
    by_county = {name: sum(1 for c in contests.values() if c["county"] == name) for name in COUNTIES}
    left_off = ([plural(counts["withdrawn"], "candidate") + " who withdrew"] if counts["withdrawn"] else []) + \
               ([plural(counts["status"] + counts["nameless"] + counts["repeat"], "row") + " that could not be read"]
                if counts["status"] + counts["nameless"] + counts["repeat"] else []) + \
               ([plural(counts["off_schedule"] + counts["unknown"], "row") + " for offices named among the gaps"]
                if counts["off_schedule"] + counts["unknown"] else [])
    coverage = ("Loaded from the Department of Elections' lists of filed candidates and of declared write-in candidates for the November 3, 2026 "
                f"general election, as read on {long_date(G['fetched'])}: all {len(races)} county contests on the Department's schedule for 2026 "
                f"({by_county['New Castle']} in New Castle County, {by_county['Kent']} in Kent County and {by_county['Sussex']} in Sussex County), "
                f"with {plural(n_printed, 'candidate')} printed on the ballot and {plural(n_write, 'declared write-in candidate')}"
                + (f"; {and_words(left_off)} {'is' if len(left_off) == 1 and left_off[0].startswith('1 ') else 'are'} left off" if left_off else "")
                + (f"; {plural(n_empty, 'contest has', 'contests have')} no candidate on the lists" if n_empty else "")
                + ". The lists give no ballot positions, so no ballot order is shown. Not loaded: the September 15 party primaries for county "
                "offices, ballot questions (a list of candidates carries none), and anything about a candidate beyond the name as filed, the "
                "office, the party and the status. "
                + ("The contests shown are county offices only; what the lists or the Department's calendars carry beyond them is named among the gaps."
                   if nov3 or city_2026 or unknown or off_schedule
                   else "No city, town, school board or special district office is on Delaware's November 3 ballot: the Department's schedule, "
                        "calendars and lists have none for that day." if nov3 == 0
                   else "The Department's lists of candidates carry no city, town, school board or special district office."))
    notes = [(STATE, "local_calendar", calendar,
              "Delaware Department of Elections, Schedule of Elections: Elected Office Table 2026-2034, and 2026 Municipal Elections Calendar; "
              "14 Del. C. section 1072(c) (school board elections)", SCHEDULE_URL),
             (STATE, "local_coverage", coverage,
              "Delaware Department of Elections, Filed Candidates by Office and Write-In Candidates by Office, General Election 11/3/2026", G["page"])]

    # ---- sources: every file and page read for the county level
    def second_words(k):
        if second[k] is None:
            return "The list's own page could not be read as a second count on this run."
        return ("Read a second way, from the list's own page, the county rows come to the same counts by office, county" +
                (", party" if k == "general" else "") + " and status, and the page has as many rows as the file." if second[k]
                else "Read a second way, from the list's own page, the county rows do NOT come to the same counts; the run's report says how.")
    aside = counts["status"] + counts["nameless"] + counts["repeat"] + counts["off_schedule"] + counts["unknown"]
    g_on, w_on = sum(1 for c in cands if c[13] == LSRC["general"]), sum(1 for c in cands if c[13] == LSRC["write_in"])
    no_rows = sum(1 for c in contests.values() if not c["rows"])
    sources = [
        (LSRC["general"], STATE, "official candidate list", "Delaware Department of Elections", LIST_TITLES["general"] + ": county offices (Excel)",
         G["url"], G["published"], G["fetched"], G["sha256"], len(G["rows"]),
         f"The county rows of the list's own Excel file ({G['rows_in_file']} filings: {G['counts']['state']} state, {G['counts']['federal']} federal, "
         f"{G['counts']['local']} county). Six columns are read, by their headings: county, office, withdrawal date, ballot name, party and status. "
         "The name parts, filing date, residential and mailing addresses, website, e-mail addresses and telephone numbers in the file are never "
         "read, and only the six kept cells of the county rows are kept on disk; the fingerprint is of the workbook as it was fetched. "
         f"Of the {len(G['rows'])} county rows, {g_on} are candidates printed on the ballot, each stored in one contest, and "
         f"{withdrew['general']} withdrew and {'is' if withdrew['general'] == 1 else 'are'} left off"
         + (f"; {aside} more rows on the two lists are set aside, as the gaps say" if aside else "")
         + f". Provisional: {provisional}. The list gives no ballot positions. " + second_words("general")),
        (LSRC["write_in"], STATE, "official candidate list", "Delaware Department of Elections", LIST_TITLES["write_in"] + ": county offices (Excel)",
         W["url"], W["published"], W["fetched"], W["sha256"], len(W["rows"]),
         f"Declared write-in candidates, whose write-in votes alone Delaware counts: the county rows of the list's own Excel file "
         f"({W['rows_in_file']} filings for every office, {W['counts']['local']} of them for county offices, {w_on} stored, "
         f"{withdrew['write_in']} withdrawn and left off). The same columns are read, less party (the list gives none); the contact columns "
         "are never read. The fingerprint is of the workbook as it was fetched. " + second_words("write_in")),
    ]
    for k in ("general", "write_in"):
        p = lists[k].get("page_read") or {}
        if "county" in p:
            sources.append((LSRC[k + "_page"], STATE, "official candidate list (page)", "Delaware Department of Elections",
                            LIST_TITLES[k] + " (the list's own page)", lists[k]["page"], p["published"], lists[k]["fetched"], p["sha256"],
                            sum(p["county"].values()),
                            f"The same list as the Department's page shows it ({p['rows']} rows), read for the date it was last updated and as a "
                            "second count of the county rows. Only the cells under the headings Office, County" + (", Party" if k == "general" else "")
                            + " and Status are read; the Candidate cell, which carries the name with contact details, and the Filed cell are never "
                            "read, and the page is not kept. The fingerprint is of the page as it was fetched."))
    if sched:
        on = sum(1 for o in sched["offices"] if 2026 in o["years"])
        sources.append((LSRC["schedule"], STATE, "official election schedule", "Delaware Department of Elections",
                        "Schedule of Elections: Elected Office Table 2026-2034 (PDF)", sched["url"], sched["updated"], sched["fetched"], sched["sha256"],
                        len(sched["offices"]),
                        f"Which office is elected in which year. Read: the titles and year marks of the {len(sched['offices'])} county offices and "
                        f"the {len(sched['wilmington'])} City of Wilmington rows; {on} county offices are marked for 2026"
                        + ((", the same contests the candidate lists carry." if not no_rows else
                            f"; {plural(no_rows, 'of them has', 'of them have')} no filing on the candidate lists.")
                           if agrees and not off_schedule and not unknown
                           else ", which is NOT the set this loader was written from or the lists carry; the run's report says how.")
                        + " Nothing else of the file is kept."))
    if muni:
        ds = sorted(muni["dates"])
        sources.append((LSRC["municipal"], STATE, "official election calendar", "Delaware Department of Elections",
                        "2026 State of Delaware Municipal Elections Calendar (PDF)", muni["url"], muni["updated"], muni["fetched"], muni["sha256"],
                        sum(muni["dates"].values()),
                        f"Only the election dates are read: {sum(muni['dates'].values())} dated entries, from {long_date(ds[0])} to "
                        f"{long_date(ds[-1])}, {f'{nov3} of them' if nov3 else 'none of them'} on November 3, 2026. The towns' names and "
                        "websites on the calendar are not kept."))
    if epage:
        sources.append((LSRC["election_page"], STATE, "official election page", "Delaware Department of Elections",
                        "General Election, November 3, 2026 (the Department's page)", epage["url"], "", epage["fetched"], epage["sha256"], 1,
                        "Read only for whether the general election's sample ballots, which would give the ballot order, are posted: on the day it "
                        "was read the page " + ("said they were not yet available." if epage["not_yet"] else "no longer said they were unavailable.")
                        + " Nothing else of the page is kept."))

    # ---- the last look before anything is written: nothing that reads like contact details, by the trial check's test and the page builder's
    for table, strict, items in (("sl_races", True, [(r[0], (r[4], r[5], r[8], r[9], r[16])) for r in races]),
                                 ("sl_candidates", True, [(c[0], (c[3], c[4], c[14])) for c in cands]),
                                 ("sl_gaps", True, [(g[2], (g[3], g[4], g[5])) for g in gaps]),
                                 ("sl_notes", True, [(n[1], (n[2], n[3])) for n in notes]),
                                 ("sl_sources", False, [(s[0], (s[3], s[4], s[10])) for s in sources])):
        for ident, texts in items:
            if any(t and reads_like_contact(t, strict) for t in texts):
                raise SystemExit(f"Delaware (county races): a text for {table} ({ident}) reads like contact details; stopping (the text is not printed)")

    by_kind = {}
    for r in races:
        by_kind[r[3]] = by_kind.get(r[3], 0) + 1
    return {"races": races, "cands": cands, "sources": sources, "gaps": gaps, "notes": notes, "report": report, "counts": counts,
            "by_kind": by_kind, "by_county": by_county, "printed": n_printed, "write_ins": n_write, "empty": n_empty, "second": second,
            "fetched": G["fetched"], "provisional": provisional, "schedule": agrees}


# ------------------------------------------------------------------------------------------------------------ loading

def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def mtime(path):
    return dt.datetime.fromtimestamp(os.path.getmtime(path)).strftime("%Y-%m-%d")


def load(db_path, say=print, cache=CACHE, roster_db=ROSTER, local_cache=None):
    """Delaware's rows into the database at db_path: the state races, then the county ones. local_cache is the folder for
    the county level's kept files (ballot_cache/de/local/ unless told otherwise)."""
    net.patient_lookups()
    folder = os.path.join(cache, "de")
    os.makedirs(folder, exist_ok=True)
    report = []

    paths = {k: os.path.join(folder, KEPT[k]) for k in DE.LISTS}
    lists = {k: kept(paths[k], 30 if k == "primary" else 2, lambda k=k: candidate_list(k, say), say) for k in DE.LISTS}
    s24path, dcpath = os.path.join(folder, SENATE_FILE), os.path.join(folder, SEATS_FILE)
    s24 = kept(s24path, 3650, lambda: senate_2024(say), say)
    dc = kept(dcpath, 3650, lambda: district_counties(say), say)
    stpath = os.path.join(folder, "Election_StatewideResults_ID_PR2026.json")
    copath = os.path.join(folder, "Election_ByCountyWithWilmington_ID_PR2026.json")
    asked = [not fresh(p, 30) for p in (stpath, copath)]      # one request at a time, a moment apart
    statewide = DE.results_file(DE.STATEWIDE, stpath, say)
    time.sleep(1.5 if asked[0] else 0)
    counties_file = DE.results_file(DE.COUNTIES, copath, say)
    time.sleep(1.5 if asked[1] else 0)
    legs, offs = roster(roster_db)
    ctab = census_counties()

    # ---- which seats are on the ballot
    senate_up = sorted(set(range(1, SENATE_SEATS + 1)) - set(s24["districts"]))
    if len(s24["districts"]) + len(senate_up) != SENATE_SEATS:
        raise SystemExit(f"Delaware (state races): the 2024 results name Senate districts outside 1-21 ({s24['districts']})")
    listed_senate = sorted({int(office_of(r["Office"])["district"]) for k in ("general", "write_in") for r in lists[k]["rows"]
                            if office_of(r["Office"])["office_kind"] == "state_senate"})
    special = sorted(set(listed_senate) & set(s24["districts"]))
    if special:
        raise SystemExit(f"Delaware (state races): Senate districts elected in 2024 are on the 2026 list ({special}): a special "
                         "election, to be read before it is loaded")

    def status_of(r, where):
        status = r["DisplayedStatus"]
        if status == "Withdrawn" or r["Withdrawal Date"]:
            return "off"
        if status not in ON:
            raise SystemExit(f"Delaware (state races): a status on the {where} list that is not read ({status!r}, {r['Office']})")
        return status

    def party_of(r, where):
        if not r.get("Party"):
            raise SystemExit(f"Delaware (state races): a candidate for {r['Office']} with no party on the {where} list")
        return WRITTEN_OUT.get(r["Party"], r["Party"])

    # ---- the races: every seat up this year, and every state office on the November list
    races = {}

    def holders(race):
        if race["level"] == "legislature":
            return [p for p in legs if p["chamber"] == race["chamber"] and str(p["district"]) == race["district"]]
        rk = STATEWIDE[race["office"]][2]
        return [p for p in offs if (p["office"] or "").lower() == rk] if rk else []

    def add(info):
        rid = info["race_id"]
        if rid in races:
            return races[rid]
        race = dict(info, state=STATE, county_ids=None, seat=None, special=0, partisan=1, holder_id=None, holder_name=None,
                    holder_party=None, election_date=GENERAL, note=None)
        if race["level"] == "legislature":
            names = dc["SS" if race["office_kind"] == "state_senate" else "SH"][race["district"]]
            race["county_ids"] = json.dumps(sorted(ctab[n][0] for n in names))
            race["_counties"] = names
        hs = holders(race)
        if len(hs) > 1:
            raise SystemExit(f"Delaware (state races): the roster has {len(hs)} sitting members for {rid}")
        race["_holders"] = hs
        notes = []
        if hs:
            race.update(holder_id=hs[0]["id"], holder_name=hs[0]["full"], holder_party=hs[0]["party"])
        elif race["level"] == "legislature":
            report.append(f"{rid}: no sitting member in the roster for this seat")
            notes.append("The roster shows no sitting member for this seat.")
        else:
            notes.append(NO_HOLDER.format(office=race["office"]))
        if race["office_kind"] == "state_senate":
            notes.append(f"Senators serve four-year terms, staggered: {len(senate_up)} of the {SENATE_SEATS} Senate districts elect a "
                         f"senator in 2026, and the other {len(s24['districts'])} ({', '.join(map(str, s24['districts']))}) were elected in "
                         "2024 (the Department of Elections' official 2024 general election results).")
        elif race["office_kind"] == "state_house":
            notes.append(HOUSE_NOTE)
        race["_notes"] = notes
        races[rid] = race
        return race

    for d in senate_up:
        add(office_of(f"State Senator District {d}"))
    for d in range(1, HOUSE_SEATS + 1):
        add(office_of(f"State Representative District {d}"))
    for k in ("general", "write_in"):
        for r in lists[k]["rows"]:
            add(office_of(r["Office"]))
    for office in STATEWIDE:
        if office_of(office)["race_id"] not in races:
            report.append(f"{office}: not on the November list")

    # the list files each legislative contest under one county: it must be one the district touches
    for k in ("general", "write_in", "primary"):
        for r in lists[k]["rows"]:
            info = office_of(r["Office"])
            if info["level"] == "statewide":
                if r["County"] != "Statewide":
                    raise SystemExit(f"Delaware (state races): a {r['Office']} row on the {k} list is filed under {r['County']!r}")
                continue
            names = dc["SS" if info["office_kind"] == "state_senate" else "SH"][info["district"]]
            if r["County"] not in names:
                report.append(f"{info['race_id']}: the {k} list files it under {r['County']}, a county the 2022 report does not place it in ({names})")

    def identify(race, reads, party):
        """(incumbent, state_member_id, note) for one name: the seat's holder when the name fits, else a sitting legislator
        of the same party elsewhere when the name fits exactly one."""
        hs = race["_holders"]
        got = [p for p in hs if person_fits(reads, p)]
        if len(got) == 1:
            return 1, got[0]["id"], None
        pool = [p for p in legs if (p["party"] or "") == party and person_fits(reads, p) and p not in hs]
        if len(pool) == 1:
            return 0, pool[0]["id"], f"Serves today in the {chamber_words(pool[0])}."
        return 0, None, None

    def one_holder_each(rows):
        """A sitting member's id goes to one candidate per election only."""
        seen = {}
        for i, c in enumerate(rows):
            if c[12]:
                seen.setdefault((c[1], c[12]), []).append(i)
        for (election, mid), idx in seen.items():
            if len(idx) < 2:
                continue
            inc = [i for i in idx if rows[i][7]]
            keep = inc[0] if len(inc) == 1 else None
            for i in idx:
                if i == keep:
                    continue
                report.append(f"{rows[i][0]} {election}: {rows[i][3]} fits the sitting member {mid}, who is matched elsewhere or twice; not tied")
                rows[i][7], rows[i][12] = 0, None
                rows[i][14] = re.sub(r"\s*Serves today in the [^.]*\.", "", rows[i][14] or "").strip() or None

    # ---- November: printed candidates in list order, then declared write-ins
    printed, off, provisional = {}, [], []
    for r in lists["general"]["rows"]:
        rid, name = office_of(r["Office"])["race_id"], r["BallotName"]
        if not name:
            raise SystemExit(f"Delaware (state races): a {r['Office']} row on the general list has no ballot name")
        status = status_of(r, "general")
        if status == "off":
            off.append(f"{name} ({r['Office']}, {r.get('Party') or 'no party'}, withdrew {r['Withdrawal Date'] or 'on a date not given'})")
            continue
        printed.setdefault(rid, []).append((name, party_of(r, "general"), status))
    cands, nominee, seen, order_odd = [], {}, set(), []
    for rid, rows in printed.items():
        race = races[rid]
        parties = [p for _n, p, _s in rows]
        blocks = [p for i, p in enumerate(parties) if i == 0 or parties[i - 1] != p]
        ok = len(blocks) == len(set(blocks)) and [RANK.get(p, 2) for p in parties] == sorted(RANK.get(p, 2) for p in parties)
        if not ok:
            order_odd.append(rid)
        for pos, (name, party, status) in enumerate(rows, 1):
            if (rid, nkey(name)) in seen:
                raise SystemExit(f"Delaware (state races): {name} is on the {rid} list twice")
            seen.add((rid, nkey(name)))
            if party in CODES:
                if (rid, CODES[party]) in nominee:
                    raise SystemExit(f"Delaware (state races): two {party} candidates printed for {rid}")
                nominee[(rid, CODES[party])] = name
            notes = []
            if status == "Provisional":
                notes.append(PROVISIONAL)
                provisional.append(name)
            inc, mid, n2 = identify(race, readings(name), party)
            if n2:
                notes.append(n2)
            cands.append([rid, "general", GENERAL, name, party, party_code(party), pos if ok else None, inc, 0, None, None, None, mid,
                          SRC["general"], " ".join(notes) or None])
    if order_odd:
        report.append(f"contests whose list order is not Democratic, Republican, others (stored with no ballot order): {order_odd}")

    write_ins, write_off = [], []
    for r in lists["write_in"]["rows"]:
        race = races[office_of(r["Office"])["race_id"]]
        name = r["BallotName"]
        status = status_of(r, "write-in")
        if status == "off":
            write_off.append(f"{name} ({r['Office']})")
            continue
        if (race["race_id"], nkey(name)) in seen:
            raise SystemExit(f"Delaware (state races): {name} is both printed and a declared write-in for {r['Office']}")
        seen.add((race["race_id"], nkey(name)))
        write_ins.append(f"{name} ({r['Office']}{', Provisional' if status == 'Provisional' else ''})")
        inc, mid, n2 = identify(race, readings(name), "")
        notes = [WRITE_IN] + ([PROVISIONAL] if status == "Provisional" else []) + ([n2] if n2 else [])
        cands.append([race["race_id"], "general", GENERAL, name, "Write-in", "W", None, inc, 1, None, None, None, mid,
                      SRC["write_in"], " ".join(notes)])
    one_holder_each(cands)

    for rid, race in races.items():
        if not any(c[0] == rid and not c[8] for c in cands):
            report.append(f"{rid}: no printed candidate on the November list")
            race["_notes"].append("The Department of Elections' list names no candidate printed on the ballot for this seat.")
        if race["partisan"] and (race["level"] == "legislature" or race["level"] == "statewide"):
            race["_notes"].append(ORDER_NOTE)

    # ---- the primaries: who was on each party's ballot, and the official votes
    filed, withdrew = {}, []
    for r in lists["primary"]["rows"]:
        rid, name = office_of(r["Office"])["race_id"], r["BallotName"]
        if status_of(r, "primary") == "off":
            withdrew.append(f"{name} ({r['Office']}, {r.get('Party')})")
            continue
        party = party_of(r, "primary")
        if party not in CODES:
            raise SystemExit(f"Delaware (state races): a primary candidate for {r['Office']} of a party that is not read ({party!r})")
        filed.setdefault((rid, CODES[party]), {})[nkey(name)] = (name, party)
    votes, report_time, voted_in, local_rows = primary_results(statewide, counties_file, report)
    if set(votes) != set(filed):
        raise SystemExit(f"Delaware (state races): the party primaries in the results and on the primary list differ ({sorted(set(votes) ^ set(filed))})")
    for rid, cs in voted_in.items():
        race = races.get(rid)
        if race and race["level"] == "legislature" and not cs <= set(race["_counties"]):
            report.append(f"{rid}: primary votes in {sorted(cs - set(race['_counties']))}, a county the 2022 report does not place it in")

    primary_rows, fields, not_on, later = [], {}, [], []
    for (rid, code), field in sorted(votes.items()):
        if rid not in races:
            raise SystemExit(f"Delaware (state races): primary results for {rid}, which is not on the November list or up this year")
        race, listed = races[rid], filed[(rid, code)]
        if {nkey(n) for n, _v, _p in field} != set(listed):
            raise SystemExit(f"Delaware (state races): the results' candidates for {rid} {code} are not the primary list's "
                             f"({sorted({nkey(n) for n, _v, _p in field} ^ set(listed))})")
        top = max(v for _n, v, _p in field)
        if sum(1 for _n, v, _p in field if v == top) > 1:
            raise SystemExit(f"Delaware (state races): the {rid} {code} primary is tied at the top; read how it was settled")
        winner = next(listed[nkey(n)][0] for n, v, _p in field if v == top)
        nom = nominee.get((rid, code))
        if nom is not None and nkey(nom) != nkey(winner) and not person_fits(readings(nom), {"first": " ".join(readings(winner)[0][0]),
                                                                                          "last": readings(winner)[0][1]}):
            later.append(f"{nom} ({rid}, {code}; the primary was won by {winner})")
        if len(field) < 2:
            continue
        fields[race["office_kind"]] = fields.get(race["office_kind"], 0) + 1
        total = sum(v for _n, v, _p in field)
        for name, v, pct in sorted(field, key=lambda t: (-t[1], nkey(t[0]))):
            shown, party = listed[nkey(name)]
            if isinstance(pct, (int, float)) and total and abs(100 * v / total - pct) > 0.011:
                raise SystemExit(f"Delaware (state races): {name}'s share in the {rid} {code} results ({pct}) is not the votes' share ({100 * v / total:.2f})")
            won, notes = v == top, []
            if won and (nom is None or (nkey(nom) != nkey(shown) and not person_fits(readings(nom), {"first": " ".join(readings(shown)[0][0]),
                                                                                                      "last": readings(shown)[0][1]}))):
                notes.append("Won the primary but is not on the November list.")
                not_on.append(f"{shown} ({rid}, {code})")
            inc, mid, n2 = identify(race, readings(shown), party)
            if n2:
                notes.append(n2)
            primary_rows.append([rid, f"primary-{code}", PRIMARY, shown, party, party_code(party), None, inc, 0, v,
                                 round(100 * v / total, 1) if total else None, "advanced" if won else "lost", mid, SRC["results"],
                                 " ".join(notes) or None])
    one_holder_each(primary_rows)
    cands += primary_rows
    keys = set()
    for c in cands:
        k = (c[0], c[1], c[3])
        if k in keys:
            raise SystemExit(f"Delaware (state races): {c[3]} is listed twice in {c[0]} {c[1]}")
        keys.add(k)

    # ---- check the counts against the lists
    general = [c for c in cands if c[1] == "general"]
    n_printed = sum(1 for r in lists["general"]["rows"] if status_of(r, "general") != "off")
    n_write = sum(1 for r in lists["write_in"]["rows"] if status_of(r, "write-in") != "off")
    if sum(1 for c in general if not c[8]) != n_printed or sum(1 for c in general if c[8]) != n_write:
        raise SystemExit("Delaware (state races): the stored November candidates do not match the lists' counts")
    sen = sorted(int(r["district"]) for r in races.values() if r["office_kind"] == "state_senate")
    if sen != senate_up:
        report.append(f"Senate districts: stored {sen}, up this year {senate_up}")
    if sorted(int(r["district"]) for r in races.values() if r["office_kind"] == "state_house") != list(range(1, HOUSE_SEATS + 1)):
        report.append("the House districts stored are not 1 to 41")

    for race in races.values():
        race["note"] = " ".join(race["_notes"]) or None

    # ---- places: the three counties, and each district on the ballot
    place_rows = [("county", g, full, json.dumps([g]), SRC["census"]) for g, full in sorted(ctab.values())]
    for race in sorted(races.values(), key=lambda r: (r["office_kind"], int(r["district"] or 0))):
        if race["level"] == "legislature":
            place_rows.append(("senate" if race["office_kind"] == "state_senate" else "house", race["jurisdiction_id"],
                               race["jurisdiction"], race["county_ids"], SRC["ge2022"]))

    # ---- the county level: the county rows of the same lists (its races use the three county places above)
    local = local_level(local_cache or os.path.join(folder, "local"), lists, ctab, say)

    # ---- write: Delaware's rows only, in one transaction
    cols = ["race_id", "state", "level", "office_kind", "office", "jurisdiction", "jurisdiction_id", "county_ids", "district", "seat",
            "special", "partisan", "holder_id", "holder_name", "holder_party", "election_date", "note"]
    race_rows = [tuple(r[c] for c in cols) for r in sorted(races.values(), key=lambda r: r["race_id"])]
    if {r[0] for r in race_rows} & {r[0] for r in local["races"]}:
        raise SystemExit("Delaware: a county race shares its id with a state race")
    g, w, p = lists["general"], lists["write_in"], lists["primary"]
    published = (report_time or "")[:10]
    src = [
        (SRC["general"], STATE, "official candidate list", "Delaware Department of Elections",
         "Filed Candidates by Office, General Election 11/3/2026: state offices (Excel)", g["url"], g["published"], g.get("fetched") or mtime(paths["general"]),
         g["sha256"], len(g["rows"]),
         f"From the list's own Excel file ({g['rows_in_file']} filings: {g['counts']['state']} state, {g['counts']['federal']} federal, "
         f"{g['counts']['local']} county). Six columns read by name (county, office, withdrawal date, ballot name, "
         "party, status); name parts, filing date, addresses, website, e-mail and telephones never read. No ballot positions are "
         f"given; the list's order is kept. Withdrawn, left off: {len(off)} ({'; '.join(off) or 'none'}). Provisional: "
         f"{', '.join(provisional) or 'none'}. The county rows of the same file are read by the county part of this loader."),
        (SRC["write_in"], STATE, "official candidate list", "Delaware Department of Elections",
         "Write-In Candidates by Office, General Election 11/3/2026: state offices (Excel)", w["url"], w["published"], w.get("fetched") or mtime(paths["write_in"]),
         w["sha256"], len(w["rows"]),
         f"Declared write-in candidates, whose write-in votes alone Delaware counts ({w['rows_in_file']} for every office). "
         f"The same columns read, less party (the list gives none). State offices: {'; '.join(write_ins) or 'none'}. "
         f"Withdrawn, left off: {', '.join(write_off) or 'none'}."),
        (SRC["primary"], STATE, "official candidate list", "Delaware Department of Elections",
         "Filed Candidates by Office, Primary Election 9/15/2026: state offices (Excel)", p["url"], p["published"], p.get("fetched") or mtime(paths["primary"]),
         p["sha256"], len(p["rows"]),
         "Who was on each party's primary ballot and how the names are printed; the results' candidates must be exactly these. "
         f"Contested primaries only. The same kept columns only. Withdrawn before the primary: {'; '.join(withdrew) or 'none'}."),
        (SRC["results"], STATE, "official results", "Delaware Department of Elections",
         "2026 Primary Election (September 15, 2026), Official Results: statewide results (JSON)", DE.STATEWIDE, published, mtime(stpath),
         sha(stpath), len(primary_rows),
         f"The file behind the Department's results page; every row says OFFICIAL RESULTS, report time {report_time}, "
         "all election districts reported. Each candidate's machine, absentee and early-voting votes add up to the total. No write-in "
         f"votes are reported, so a field's total is its candidates' votes. {local_rows} county-office rows not read."),
        (SRC["county"], STATE, "official results", "Delaware Department of Elections",
         "2026 Primary Election (September 15, 2026), Official Results: by county, with Wilmington (JSON)", DE.COUNTIES, published, mtime(copath),
         sha(copath), len(primary_rows),
         "Control: for every state candidate, New Castle, Kent and Sussex add up to the state figure, Wilmington and the rest of New "
         "Castle add up to New Castle, and the state figure equals the statewide file's; a district's votes lie only in its counties."),
        (SRC["ge2024"], STATE, "official results", "Delaware Department of Elections",
         "2024 General Election, Official Results (CSV)", s24["url"], "2024-11-05", s24["fetched"], s24["sha256"], len(s24["districts"]),
         f"Office column only: the Senate districts elected in 2024 ({', '.join(map(str, s24['districts']))}), which are not on the 2026 "
         f"ballot; the other {len(senate_up)} are."),
        (SRC["ge2022"], STATE, "official results", "Delaware Department of Elections",
         "2022 General Election Report, By Election District", dc["url"], "2022-11-08", dc["fetched"], dc["sha256"], dc["election_districts"],
         "Contest titles and election district labels only: each Senate and Representative District's election districts (every one "
         "of the state's election districts in exactly one of each), and the county of each from the three county sheriff contests. "
         "Gives each district's counties (county_ids). The 2022 election was the first under the lines drawn in 2021."),
        (SRC["roster"], STATE, "roster", "Open States people project (CC0), as loaded into state_de.sqlite",
         "Sitting legislators and statewide officials", "https://github.com/openstates/people", "", mtime(roster_db), sha(roster_db),
         len(legs) + len(offs), "Who holds each seat today; the roster carries the Attorney General but not the State Treasurer or the "
         "Auditor of Accounts. Ids, names, parties, chambers, districts and offices only."),
        (SRC["census"], STATE, "county codes", "U.S. Census Bureau", "Cartographic boundary file, counties, 2024 (1:500,000)",
         COUNTY_URL, "2024", mtime(COUNTY_ZIP), sha(COUNTY_ZIP), len(ctab), "Delaware's three counties: names and GEOIDs only."),
    ]
    # Delaware's rows only: its races and candidates by state and race id, its sources, gaps and notes by state, its places
    # by their de- source ids (the counties' codes begin 10, so the id alone cannot pick them out)
    con = sqlite3.connect(db_path)
    try:
        con.executescript(SCHEMA + EXTRA_SCHEMA)
        with con:
            con.execute("DELETE FROM sl_candidates WHERE race_id IN (SELECT race_id FROM sl_races WHERE state = ?)", (STATE,))
            con.execute("DELETE FROM sl_candidates WHERE race_id LIKE ?", (f"2026-{STATE}-%",))
            con.execute("DELETE FROM sl_races WHERE state = ?", (STATE,))
            con.execute("DELETE FROM sl_places WHERE source_id LIKE ?", (STATE.lower() + "-%",))
            con.execute("DELETE FROM sl_sources WHERE state = ?", (STATE,))
            con.execute("DELETE FROM sl_gaps WHERE state = ?", (STATE,))
            con.execute("DELETE FROM sl_notes WHERE state = ?", (STATE,))
            con.executemany(f"INSERT INTO sl_races ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})", race_rows + local["races"])
            con.executemany("INSERT INTO sl_candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", [tuple(c) for c in cands + local["cands"]])
            con.executemany("INSERT INTO sl_places VALUES (?,?,?,?,?)", place_rows)
            con.executemany("INSERT INTO sl_sources VALUES (?,?,?,?,?,?,?,?,?,?,?)", src + local["sources"])
            con.executemany("INSERT INTO sl_gaps VALUES (?,?,?,?,?,?,?)", local["gaps"])
            con.executemany("INSERT INTO sl_notes VALUES (?,?,?,?,?)", local["notes"])
    finally:
        con.close()

    by = lambda kind: sum(1 for c in general if races[c[0]]["office_kind"] == kind)
    nrace = lambda kind: sum(1 for r in races.values() if r["office_kind"] == kind)
    say(f"    Delaware (state races): {len(races)} races ({nrace('state_senate')} Senate, {nrace('state_house')} House, "
        f"{len(races) - nrace('state_senate') - nrace('state_house')} statewide); {len(general)} November candidates (Senate "
        f"{by('state_senate')}, House {by('state_house')}, statewide {len(general) - by('state_senate') - by('state_house')}; "
        f"{len(write_ins)} declared write-in, {len(off)} withdrawn left off); {sum(fields.values())} primary fields "
        f"({', '.join(f'{k} {v}' for k, v in sorted(fields.items()))}), {len(primary_rows)} primary rows, official votes, county sums checked")
    for c in cands:
        if c[12]:
            say(f"      matched: {c[0]} {c[1]}: {c[3]} -> {c[12]}{' (holds this seat)' if c[7] else ''}")
    if not_on:
        report.append(f"primary winners not on the November list: {', '.join(not_on)}")
    if later:
        report.append(f"November nominees who did not win their party's primary: {', '.join(later)}")
    for line in report:
        say(f"      check: {line}")
    lc = local["counts"]
    say(f"    Delaware (county races): {len(local['races'])} races ({', '.join(f'{k} {v}' for k, v in sorted(local['by_county'].items()))}; "
        f"{', '.join(f'{k} {v}' for k, v in sorted(local['by_kind'].items()))}); {local['printed']} candidates printed on the ballot and "
        f"{local['write_ins']} declared write-in, no ballot order (the lists state none); {len(local['gaps'])} in sl_gaps; lists of {local['fetched']}")
    aside = {"with a status not read": lc["status"], "with no usable name": lc["nameless"], "repeated": lc["repeat"],
             "for offices not on the 2026 schedule": lc["off_schedule"], "for offices not known": lc["unknown"]}
    pages = {True: "agree", False: "DIFFER", None: "not read"}
    say(f"      count: {lc['rows']} county rows on the two lists = {lc['placed']} candidates stored, each in one race, + {lc['withdrawn']} withdrawn"
        + "".join(f" + {v} {k}" for k, v in aside.items() if v)
        + f"; the lists' own pages, as a second count: general {pages[local['second']['general']]}, write-in {pages[local['second']['write_in']]}"
        + f"; the Department's schedule for 2026: {pages[local['schedule']]}"
        + (f"; {local['empty']} scheduled contests with no candidate" if local["empty"] else ""))
    for line in local["report"]:
        say(f"      check: {line}")
    return len(general)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: python -m ballot.state_local_de <database>")
    load(sys.argv[1])
