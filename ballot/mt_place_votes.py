"""
ballot/mt_place_votes.py - how each Montana place voted in past partisan general elections, so a page can show the
record of a county or a district without anyone labelling a candidate. The Montana twin of ballot/mn_place_votes.py;
the output has the same shape (the Democratic count is "dem" here, as in the Dakotas' files, where Minnesota's is
"dfl").

    python ballot/mt_place_votes.py               reads (or downloads) the files, writes ballot/lean/mt_place_votes.json
    python ballot/mt_place_votes.py --refresh     asks for everything again even when the cached copies are there
    python ballot/mt_place_votes.py --out x.json  writes somewhere else; --cache DIR keeps the downloads somewhere else
    python ballot/mt_place_votes.py --selftest    the arithmetic and the readers on made-up tables; downloads nothing

What it is, and is not
----------------------
For President, U.S. Senator and Governor in 2024 and in 2020 (Montana elects its Governor in presidential years, and
elected no senator and no statewide partisan officer in 2022, so 2022 has no contest here): the votes for the
Democratic ticket, the Republican ticket, every other candidate on the ballot together, and the total, for every one of
the 56 counties and 22 judicial districts, and for 2024 for every one of the 100 state House districts, 50 state
Senate districts and 5 Public Service Commission districts. It is how the people of a place voted then, on the lines in
force at that election. It is not a prediction, it says nothing about any candidate on a later ballot or about any
voter, and it turns no nonpartisan office into a partisan one. No database is opened for writing;
ballot_local_2026.sqlite is opened read-only at the end, only to count how many of our places are covered.

Where the numbers come from
---------------------------
All from the Montana Secretary of State's results page (sosmt.gov/elections/results/), which answers a plain script:

  - "Precinct by Precinct" (a workbook for each general election): one row for each candidate in each precinct, with
    the county, the precinct, the race, the party and the votes. Read by its column headings; the columns kept are
    the county, the precinct, the race, its area, the party, the votes and the name on the ballot. THE COUNTY FIGURES
    ARE THESE ROWS ADDED UP. The workbook carries no write-in votes (the Board canvasses declared write-ins on a
    separate abstract: a few dozen votes for President in 2024), so "other" here is the other candidates printed on
    the ballot, and the total is the votes for printed candidates.
  - "Statewide" (the State Canvass, a PDF): the Board of State Canvassers' table of each statewide contest, a row a
    county and a Total row. It is the control for the counties: every county, every candidate.
  - "Statewide by HD" (the State Canvass by House district, a PDF, 2024): each statewide contest by House district
    and, inside a district, by county. THE HOUSE DISTRICT FIGURES ARE THIS DOCUMENT'S TOTAL ROWS. They are not added
    up from precincts because four precincts of 727 (in Blaine, Hill and Roosevelt Counties) are split between two
    House districts, and the workbook gives a split precinct's votes for President whole. Everywhere else the
    workbook's precincts are added up by the House district each one voted in and must equal this document, county
    part by county part.
  - A Senate district is two House districts (district n is House districts 2n-1 and 2n in the plan first used in
    2024); the loader checks that against the precincts of every State Senator contest in the 2024 workbook and,
    where the kit has it, the 2026 primary workbook, which between them hold all 50. A Public Service Commission
    district is twenty House districts, listed in section 69-1-104, MCA; a judicial district is whole counties,
    listed in section 3-5-101, MCA. Both sections are read from the Legislature's own site (archive.legmt.gov) and
    must equal the copies typed below.

Not given: cities and towns (a Montana precinct has a number and a local name, and the results do not say which city
or town it lies in); county commissioner districts, wards, conservation and school districts (the results do not name
them for a precinct); House, Senate and Public Service Commission districts for 2020 (the House lines drawn after the
2010 census; today's were first used in 2024).

The control
-----------
Nothing is written unless all of this holds: in each workbook every precinct has every candidate of a contest once;
for every contest the State Canvass table has all 56 counties, each column adds up to its Total row, and every
county's row equals the workbook's precincts added up, candidate by candidate; for President and U.S. Senator the
totals equal the Clerk of the U.S. House of Representatives' "Statistics of the ... Election" (clerk.house.gov, the
Montana page, read from the PDF at each run); for 2024 the State Canvass by House district has all 100 districts, each
district's county rows add up to its Total row, the districts add up to the statewide totals, and the workbook's
precincts equal it wherever no precinct is split; every kind of place that covers the state adds up to the statewide
sum. The statewide figures this loader was checked against on 2026-10-02 are typed below (CHECKED); the loader says
so if a document reads differently later.
"""

import argparse
import collections
import html
import json
import os
import pathlib
import re
import sqlite3
import sys

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from states import net  # noqa: E402
from ballot import mn_place_votes as M  # noqa: E402  the shared small things: fingerprints, days, the file writer
from ballot import nd_place_votes as N  # noqa: E402  the reader of the Clerk of the House's page

Stop = M.Stop
METHOD = "1.0"                 # change how anything is added up or matched, and this goes up
CACHE = os.path.join(HERE, "states_cache", "mt_local")
OUT = os.path.join(HERE, "ballot", "lean", "mt_place_votes.json")
DB = os.path.join(HERE, "ballot_local_2026.sqlite")
PRIMARY_2026 = os.path.join(HERE, "ballot_cache", "mt", "mt_2026_primary_precinct.xlsx")   # the federal loader's copy; a check only
STATE_FIPS = "30"
FEW = 20
SIDES = ("dem", "rep", "other", "total")
KEEP_DAYS = 3650               # certified results of past elections do not change; --refresh asks again
HOUSE, SENATE = 100, 50

AGENCY = "Montana Secretary of State"
RESULTS_PAGE = "https://sosmt.gov/elections/results/"
ELECTIONS = {
    2024: {"date": "2024-11-05",
           "book": {"url": "https://sosmt.gov/docs/32/results/66992/2024_general_precinct_by_precinct", "file": "mt_2024_general_precinct.xlsx",
                    "title": "2024 General Election, Precinct by Precinct"},
           "canvass": {"url": "https://sosmt.gov/docs/31/post-election/66775/2024-general-election-report-state-canvass",
                       "file": "mt_2024_general_state_canvass.pdf", "title": "2024 Statewide General Election Canvass"},
           "by_hd": {"url": "https://sosmt.gov/docs/31/post-election/66772/2024-general-election-report-state-canvass-by-hd",
                     "file": "mt_2024_general_state_canvass_by_hd.pdf", "title": "2024 Statewide General Election Canvass, by House district"}},
    2020: {"date": "2020-11-03",
           "book": {"url": "https://sosmt.gov/wp-content/uploads/2020_General_Precinct-by-Precinct.xlsx", "file": "mt_2020_general_precinct.xlsx",
                    "title": "2020 General Election, Precinct by Precinct Results"},
           "canvass": {"url": "https://sosmt.gov/wp-content/uploads/State_Canvass_Report.pdf", "file": "mt_2020_general_state_canvass.pdf",
                       "title": "2020 Statewide General Election Canvass"}},
}
DISTRICT_YEAR = 2024           # the first election on today's House, Senate and commission districts
STATUTES = {
    "judicial": {"id": "mca-3-5-101", "cite": "Section 3-5-101, Montana Code Annotated (Judicial districts defined)", "file": "mca_3-5-101.html",
                 "url": "https://archive.legmt.gov/bills/mca/title_0030/chapter_0050/part_0010/section_0010/0030-0050-0010-0010.html"},
    "psc": {"id": "mca-69-1-104", "cite": "Section 69-1-104, Montana Code Annotated (Public service commission districts)", "file": "mca_69-1-104.html",
            "url": "https://archive.legmt.gov/bills/mca/title_0690/chapter_0010/part_0010/section_0040/0690-0010-0010-0040.html"},
}

# The 56 counties in the order of their FIPS codes (30001, 30003, ...): alphabetical, McCone before Madison.
COUNTIES = ("Beaverhead|Big Horn|Blaine|Broadwater|Carbon|Carter|Cascade|Chouteau|Custer|Daniels|Dawson|Deer Lodge|Fallon|Fergus|Flathead|"
            "Gallatin|Garfield|Glacier|Golden Valley|Granite|Hill|Jefferson|Judith Basin|Lake|Lewis and Clark|Liberty|Lincoln|McCone|Madison|"
            "Meagher|Mineral|Missoula|Musselshell|Park|Petroleum|Phillips|Pondera|Powder River|Powell|Prairie|Ravalli|Richland|Roosevelt|Rosebud|"
            "Sanders|Sheridan|Silver Bow|Stillwater|Sweet Grass|Teton|Toole|Treasure|Valley|Wheatland|Wibaux|Yellowstone").split("|")


def bare(name):
    """A county's name for comparing only: capitals, & as AND, letters alone ('Lewis & Clark', 'LEWIS AND CLARK')."""
    return re.sub(r"[^A-Z]", "", (name or "").upper().replace("&", " AND "))


FIPS = {bare(name): f"{STATE_FIPS}{2 * i + 1:03d}" for i, name in enumerate(COUNTIES)}
COUNTY_NAME = {FIPS[bare(name)]: f"{name} County" for name in COUNTIES}

# The two sections as this loader was checked against them on 2026-10-02 (read again at each run; they must be equal).
JUDICIAL = {1: ["Lewis and Clark", "Broadwater"], 2: ["Silver Bow"], 3: ["Deer Lodge", "Granite", "Powell"], 4: ["Missoula", "Mineral"],
            5: ["Beaverhead", "Jefferson", "Madison"], 6: ["Park", "Sweet Grass"], 7: ["Dawson", "McCone", "Richland", "Prairie", "Wibaux"],
            8: ["Cascade"], 9: ["Teton", "Pondera", "Toole", "Glacier"], 10: ["Fergus", "Judith Basin", "Petroleum"], 11: ["Flathead"],
            12: ["Liberty", "Hill", "Chouteau"], 13: ["Yellowstone"], 14: ["Meagher", "Wheatland", "Golden Valley", "Musselshell"],
            15: ["Roosevelt", "Daniels", "Sheridan"], 16: ["Custer", "Carter", "Fallon", "Powder River", "Garfield", "Treasure", "Rosebud"],
            17: ["Phillips", "Blaine", "Valley"], 18: ["Gallatin"], 19: ["Lincoln"], 20: ["Lake", "Sanders"], 21: ["Ravalli"],
            22: ["Stillwater", "Carbon", "Big Horn"]}
PSC = {1: [19, 20, 21, 22, 23, 26, 27, 28, 29, 30, 31, 32, 33, 34, 35, 36, 38, 43, 44, 45],
       2: [39, 40, 41, 42, 46, 47, 48, 49, 50, 51, 52, 53, 54, 55, 56, 57, 58, 59, 61, 62],
       3: [37, 60, 63, 64, 65, 66, 67, 68, 69, 70, 71, 72, 73, 74, 75, 77, 78, 79, 85, 86],
       4: [1, 2, 6, 8, 9, 10, 12, 13, 14, 87, 88, 89, 90, 93, 94, 95, 96, 97, 98, 100],
       5: [3, 4, 5, 7, 11, 15, 16, 17, 18, 24, 25, 76, 80, 81, 82, 83, 84, 91, 92, 99]}
ORDINALS = {"first": 1, "second": 2, "third": 3, "fourth": 4, "fifth": 5}

# The contests. "race" is the workbook's race name and the canvass's heading; "dem" and "rep" are the tickets, whose
# first surnames must be in the name the workbook prints for the candidate of that party, or the loader stops.
CONTESTS = [
    {"id": "2024-president", "year": 2024, "office": "President and Vice President of the United States", "race": "PRESIDENT & VICE PRESIDENT",
     "dem": "Harris and Walz", "rep": "Trump and Vance", "clerk": "FOR PRESIDENTIAL ELECTORS"},
    {"id": "2024-us-senate", "year": 2024, "office": "United States Senator", "race": "UNITED STATES SENATOR",
     "dem": "Jon Tester", "rep": "Tim Sheehy", "clerk": "FOR UNITED STATES SENATOR"},
    {"id": "2024-governor", "year": 2024, "office": "Governor and Lieutenant Governor", "race": "GOVERNOR & LT. GOVERNOR",
     "dem": "Busse and Graybill", "rep": "Gianforte and Juras", "clerk": None},
    {"id": "2020-president", "year": 2020, "office": "President and Vice President of the United States", "race": "PRESIDENT",
     "dem": "Biden and Harris", "rep": "Trump and Pence", "clerk": "FOR PRESIDENTIAL ELECTORS"},
    {"id": "2020-us-senate", "year": 2020, "office": "United States Senator", "race": "UNITED STATES SENATOR",
     "dem": "Steve Bullock", "rep": "Steve Daines", "clerk": "FOR UNITED STATES SENATOR"},
    {"id": "2020-governor", "year": 2020, "office": "Governor and Lieutenant Governor", "race": "GOVERNOR & LT. GOVERNOR",
     "dem": "Cooney and Schreiner", "rep": "Gianforte and Juras", "clerk": None},
]
PARTY_WORDS = {"DEM": "Democratic", "REP": "Republican", "LIB": "Libertarian", "GRN": "Green", "WTP": "We the People", "IND": "independent"}

# The statewide totals this loader was checked against on 2026-10-02: the workbooks, the State Canvass and (President
# and U.S. Senator) the Clerk of the House all read these.
CHECKED = {
    "2024-president": {"dem": 231906, "rep": 352079, "other": 18978, "total": 602963},
    "2024-us-senate": {"dem": 276305, "rep": 319682, "other": 11275, "total": 607262},
    "2024-governor": {"dem": 232644, "rep": 354569, "other": 15191, "total": 602404},
    "2020-president": {"dem": 244786, "rep": 343602, "other": 15252, "total": 603640},
    "2020-us-senate": {"dem": 272463, "rep": 333174, "other": 0, "total": 605637},
    "2020-governor": {"dem": 250860, "rep": 328548, "other": 24179, "total": 603587},
}
CLERK = {year: {"id": f"clerk-statistics-{year}", "kind": "official federal compilation", "agency": "Office of the Clerk, U.S. House of Representatives",
                "title": f"Statistics of the Presidential and Congressional Election from Official Sources for the Election of {day}",
                "url": f"https://clerk.house.gov/member_info/electionInfo/{year}/statistics{year}.pdf", "file": f"clerk_statistics{year}.pdf"}
         for year, day in ((2024, "November 5, 2024"), (2020, "November 3, 2020"))}

# The kinds of place: (kind, first election year given, what it is, what its key is)
KINDS = [
    ("county", 2020, "Counties", "five-digit county FIPS code, as sl_places kind county and the jurisdiction_id of our county races"),
    ("house", DISTRICT_YEAR, "State House districts of the plan first used in 2024", "district number, as the district of our House races"),
    ("senate", DISTRICT_YEAR, "State Senate districts of the plan first used in 2024", "district number, as the district of our Senate races"),
    ("judicial", 2020, "Judicial districts, as whole counties", "JD and the district number (JD4), as the jurisdiction_id of our district court races"),
    ("psc", DISTRICT_YEAR, "Public Service Commission districts", "PSC and the district number (PSC1), as the jurisdiction_id of our commission races"),
]
KIND_NAMES = [k[0] for k in KINDS]
REDRAWN = ("The 2020 election was held on the House districts drawn after the 2010 census; the districts drawn after the 2020 census were "
           "first used in 2024, so 2020 is not given.")
WHAT = ("How each Montana place voted in six contests of the 2020 and 2024 general elections: the votes for the Democratic ticket, the "
        "Republican ticket, every other candidate on the ballot together and the total, added up from the Montana Secretary of State's "
        "official precinct results and the State Canvass.")
NOTE = ("What this is: how the people of a place voted in that election, in the Secretary of State's official precinct results and the "
        "State Canvass, added up here by county and by judicial district and, for 2024, given by state House, state Senate and Public "
        "Service Commission district, on the lines in force at that election. What this is not: it is not a prediction of any election; it "
        "says nothing about any candidate on a later ballot or about any voter; a nonpartisan office stays nonpartisan; and a place is not "
        "its lines for ever: where a district was drawn again, the figures are for the lines of that year. Montana elected no senator and no "
        "statewide partisan officer in 2022, so there is no 2022 contest. Declared write-in votes are canvassed apart and are not in these "
        "figures. The tickets are named, as the Secretary's own results name them, only to say which election this was. In a place with very "
        "few voters the split would come close to saying how particular people voted; those contests are listed in the place's too_few, and "
        "a page should leave the split out.")
HOW_TO_READ = ("Each place's votes are four counts for each contest: dem (the Democratic ticket), rep (the Republican ticket), other (every "
               "other candidate printed on the ballot) and total. Minnesota's file calls the first count dfl. A contest a place does not have "
               "was not counted on its lines: House, Senate and Public Service Commission districts are given for 2024 only.")
TOO_FEW = (f"A place's too_few lists the contests in which it had fewer than {FEW} votes, or in which every vote went the same way. There the "
           "split comes close to saying how particular people voted, so a page should leave it out. The counts are the official record all the "
           "same, and are kept here so that the sums can be checked.")
NOT_GIVEN = ("Cities and towns: a Montana precinct has a number and a local name, and the results do not say which city or town it lies in. "
             "County commissioner districts, wards, conservation districts and school districts: the results do not name them for a precinct.")


# ---------------------------------------------------------------- the files

def fetch(spec, cache, refresh, say, magic):
    path = os.path.join(cache, spec["file"])
    try:
        net.download(spec["url"], path, 0 if refresh else KEEP_DAYS, tries=3, say=say)
    except Exception as e:  # noqa: BLE001
        raise Stop(f"    {spec.get('title') or spec['file']}: could not be fetched ({e}) and no copy is on disk. Wait a few minutes and run this again.")
    with open(path, "rb") as fh:
        if not fh.read(8).startswith(magic):
            raise Stop(f"    {spec['file']}: what came back is not the file (it does not begin {magic!r}); stopping")
    return path


BOOK_COLUMNS = {"county": ("CountyName", "County"), "precinct": ("PrecinctName", "Precinct"), "race": ("Race Name", "RaceName", "Race"),
                "area": ("Area", "District"), "party": ("Party", "PartyCode"), "votes": ("Votes",),
                "name": ("Name On Ballot", "NameOnBallot", "Candidate Last Name")}


def book_rows(sheet_rows, what, need=("county", "precinct", "race", "party", "votes", "name")):
    """A "Precinct by Precinct" sheet as rows of {county, precinct, race, area, party, votes, name}: the heading row is
    the first with a Votes cell, and only the columns named in BOOK_COLUMNS are taken, by their headings."""
    it = iter(sheet_rows)
    at = None
    for row in it:
        heads = [str(c).strip() if c is not None else "" for c in row]
        if "Votes" in heads:
            at = {key: next((heads.index(h) for h in names if h in heads), None) for key, names in BOOK_COLUMNS.items()}
            break
    if at is None or any(at[k] is None for k in need):
        raise Stop(f"    {what}: the workbook's columns are not the ones this loader reads; stopping")
    out = []
    for row in it:
        if not any(c is not None for c in row):
            continue
        r = {key: (row[i] if i is not None and i < len(row) else None) for key, i in at.items()}
        v = r["votes"]
        if isinstance(v, bool) or not isinstance(v, (int, float)) or v != int(v) or v < 0:
            raise Stop(f"    {what}: a Votes cell is {v!r}, not a count of votes; stopping")
        out.append({"county": str(r["county"] or "").strip(), "precinct": str(r["precinct"] or "").strip(), "race": " ".join(str(r["race"] or "").split()),
                    "area": str(r["area"] or "").strip(), "party": str(r["party"] or "").strip().upper(), "votes": int(v),
                    "name": " ".join(str(r["name"] or "").split())})
    return out


def read_book(path, what, need=("county", "precinct", "race", "party", "votes", "name")):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return book_rows(wb[wb.sheetnames[0]].iter_rows(values_only=True), what, need)
    finally:
        wb.close()


def county_of(name, more=()):
    """The FIPS code of a county as a results row writes it, with any pieces of the name wrapped onto following lines."""
    for n in range(len(more) + 1):
        code = FIPS.get(bare(" ".join([name, *more[:n]])))
        if code:
            return code
    return None


def _numbers(text):
    m = re.match(r"^(.*?)((?:\s+\d+)+)$", text.strip())
    return (m.group(1).strip(), [int(x) for x in m.group(2).split()]) if m else (text.strip(), None)


def canvass_tables(lines, titles):
    """The State Canvass's county tables: {title: {"rows": {fips: [votes]}, "total": [votes], "page": n}} for each
    title wanted. `lines` are (page, y, text) as ballot/pdftext.lines gives them. A table starts at a line that is
    the title and runs to its Total row; a county's name can wrap onto the next line (LEWIS AND / CLARK)."""
    out = {}
    texts = [(p, t.strip()) for p, _y, t in lines]
    for title in titles:
        starts = [i for i, (_p, t) in enumerate(texts) if t == title]
        if not starts:
            raise Stop(f"    the State Canvass has no table headed {title}; stopping")
        i = starts[0]
        rows, total, pending = {}, None, None
        for j in range(i + 1, len(texts)):
            name, nums = _numbers(texts[j][1])
            if nums is None:
                if pending is not None:
                    pending[1].append(texts[j][1])
                continue
            if pending is not None:
                code = county_of(pending[0], pending[1])
                if code is None or code in rows:
                    raise Stop(f"    the State Canvass, {title}: a row is headed {pending[0]!r}, not one county once; stopping")
                rows[code], pending = pending[2], None
            if name == "Total":
                total = nums
                break
            if texts[j][0] != texts[i][0]:
                raise Stop(f"    the State Canvass, {title}: the table runs past its page before a Total row; stopping")
            pending = (name, [], nums)
        if total is None or any(len(v) != len(total) for v in rows.values()):
            raise Stop(f"    the State Canvass, {title}: no Total row, or rows of different lengths; stopping")
        out[title] = {"rows": rows, "total": total, "page": texts[i][0]}
    return out


def by_hd_tables(lines, titles, houses=HOUSE):
    """The State Canvass by House district: {title: {district: {"rows": {fips: [votes]}, "total": [votes]}}}. A
    district opens with "HD n", lists its counties and closes with a Total row; names wrap (YELLOWSTO / NE)."""
    out = {}
    texts = [t.strip() for _p, _y, t in lines]
    for title in titles:
        starts = [i for i, t in enumerate(texts) if t == title]
        if not starts:
            raise Stop(f"    the State Canvass by House district has no table headed {title}; stopping")
        table, hd, pending, width = {}, None, None, None

        def settle():
            nonlocal pending
            if pending is not None:
                code = county_of(pending[0], pending[1])
                if code is None or hd is None or code in table[hd]["rows"]:
                    raise Stop(f"    the State Canvass by House district, {title}: a row is headed {pending[0]!r}, not one county once; stopping")
                table[hd]["rows"][code], pending = pending[2], None
        for t in texts[starts[0] + 1:]:
            name, nums = _numbers(t)
            if nums is None:
                if pending is not None:
                    pending[1].append(t)
                continue
            m = re.match(r"^HD\s+(\d+)\s+(.*)$", name)
            if not (m or name == "Total" or (hd is not None and re.fullmatch(r"[A-Za-z .&]+", name or ""))):
                settle()                                                    # a page number ("6 of 125"), not a row
                continue
            settle()
            if name == "Total":
                if hd is None:
                    raise Stop(f"    the State Canvass by House district, {title}: a Total row before any district; stopping")
                table[hd]["total"] = nums
                if hd == houses:
                    break
                continue
            if m:
                hd = int(m.group(1))
                if hd in table:
                    raise Stop(f"    the State Canvass by House district, {title}: district {hd} appears twice; stopping")
                table[hd] = {"rows": {}, "total": None}
                name, width = m.group(2), width or len(nums)
            if len(nums) != width:
                raise Stop(f"    the State Canvass by House district, {title}: a row of district {hd} has {len(nums)} counts, not {width}; stopping")
            pending = (name, [], nums)
        out[title] = table
    return out


def read_statutes(pages):
    """What the two sections say: ({district: [county names]}, {district: [house districts]}). `pages` are the pages'
    HTML as bytes, keyed judicial and psc."""
    text = {k: re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", v.decode("utf-8", "replace")))) for k, v in pages.items()}
    judicial = {}
    for m in re.finditer(r"\(\d+\)\s*(\d+)(?:st|nd|rd|th) district:\s*(.*?)\s*Count(?:y|ies)\s*[;.]", text["judicial"]):
        rest, names = m.group(2), []
        for name in sorted(COUNTIES, key=len, reverse=True):
            if re.search(rf"\b{re.escape(name)}\b", rest):
                names.append((rest.index(name), name))
                rest = re.sub(rf"\b{re.escape(name)}\b", " " * len(name), rest, count=1)
        if re.sub(r"\band\b|[,\s]", "", rest):
            raise ValueError(f"section 3-5-101, district {m.group(1)}: words that are not counties")
        judicial[int(m.group(1))] = [n for _i, n in sorted(names)]
    psc = {ORDINALS[m.group(1)]: [int(x) for x in re.findall(r"\d+", m.group(2))]
           for m in re.finditer(r"\([a-e]\)\s*(first|second|third|fourth|fifth) district:\s*([\d,\s]+)", text["psc"])}
    return judicial, psc


def statutes(cache, refresh, say):
    """The two sections, read from the Legislature's site and required to equal the typed copies; the typed copies
    stand in, and the file says so, when a page cannot be read."""
    recs, pages = [], {}
    for key, s in STATUTES.items():
        path = os.path.join(cache, s["file"])
        rec = {"id": s["id"], "kind": "statute, as the Legislature publishes it", "agency": "Montana Legislature, Montana Code Annotated",
               "title": s["cite"], "url": s["url"]}
        try:
            net.download(s["url"], path, 0 if refresh else 30, tries=3, say=say)
            with open(path, "rb") as fh:
                pages[key] = fh.read()
            rec.update(fetched=M._day(path), sha256=M._sha(pages[key]))
        except Exception as e:  # noqa: BLE001
            say(f"      {s['cite']}: could not be read ({e}); the copy typed into this loader on 2026-10-02 is used")
            rec["unread"] = f"could not be read on {M._now()}; the copy typed into the loader on 2026-10-02 is used"
        recs.append(rec)
    if len(pages) == 2:
        try:
            judicial, psc = read_statutes(pages)
        except ValueError as e:
            raise Stop(f"    {e}; stopping")
        if {k: sorted(v) for k, v in judicial.items()} != {k: sorted(v) for k, v in JUDICIAL.items()} or psc != PSC:
            raise Stop("    section 3-5-101 or 69-1-104, MCA, no longer reads as the copy this loader was checked against; nothing is written "
                       "until a person has read the change")
    return recs


# ---------------------------------------------------------------- adding up

def _add(into, v):
    for i in range(len(v)):
        into[i] += v[i]


def hd_number(area):
    m = re.fullmatch(r"HD\s*0*(\d+)", (area or "").strip(), re.I)
    return int(m.group(1)) if m else None


def tally_book(rows, contests, what):
    """What one workbook says of its contests: for each, the candidates (party, name) in a fixed order, their votes
    by precinct, and which of them are the two tickets. Every precinct must have every candidate once."""
    by_race = collections.defaultdict(list)
    wanted = {c["race"]: c for c in contests}
    for r in rows:
        if r["race"] in wanted:
            by_race[r["race"]].append(r)
    out = {}
    for race, c in wanted.items():
        mine = by_race.get(race) or []
        cands = sorted({(r["party"], r["name"]) for r in mine})
        if not mine:
            raise Stop(f"    {what}: no rows for {race}; stopping")
        pos = {k: i for i, k in enumerate(cands)}
        pre, seen = {}, set()
        for r in mine:
            code = county_of(r["county"])
            if code is None:
                raise Stop(f"    {what}: {r['county']!r} is not one of the 56 counties; stopping")
            key = (code, r["precinct"])
            if (key, r["party"], r["name"]) in seen:
                raise Stop(f"    {what}: precinct {r['precinct']} of {COUNTY_NAME[code]} has a candidate of {race} twice; stopping")
            seen.add((key, r["party"], r["name"]))
            pre.setdefault(key, [None] * len(cands))[pos[(r["party"], r["name"])]] = r["votes"]
        short = [k for k, v in pre.items() if None in v]
        if short:
            raise Stop(f"    {what}: {len(short)} precincts lack a candidate of {race} (first: {short[0][1]}, {COUNTY_NAME[short[0][0]]}); stopping")
        side = {}
        for s, code in (("dem", "DEM"), ("rep", "REP")):
            fits = [k for k in cands if k[0] == code]
            surname = re.split(r"\s+and\s+|\s+", c[s] if " and " in c[s] else c[s].split()[-1])[0].upper()
            if len(fits) != 1 or surname not in fits[0][1].upper():
                raise Stop(f"    {what}, {race}: the {PARTY_WORDS[code]} candidate is not the one ({c[s]}) this loader was checked against; stopping")
            side[s] = pos[fits[0]]
        out[c["id"]] = {"candidates": cands, "precincts": pre, "side": side}
    return out


def four(vec, side):
    total = sum(vec)
    return [vec[side["dem"]], vec[side["rep"]], total - vec[side["dem"]] - vec[side["rep"]], total]


def match_columns(sums, total, what):
    """Which column of an official table is which candidate of the workbook: the one whose Total equals the
    candidate's statewide sum. None when the totals are not the same numbers, or two are equal (never guessed)."""
    if sorted(sums) != sorted(total) or len(set(sums)) != len(sums):
        return None
    return [total.index(v) for v in sums]


def precinct_districts(rows, prefix, areas=True):
    """{(county, precinct): {district numbers}} from the rows of the contests whose race begins with `prefix`."""
    out = collections.defaultdict(set)
    for r in rows:
        if r["race"].upper().startswith(prefix):
            m = re.search(r"(\d+)\s*$", r["area"] if areas and r["area"] else r["race"])
            if m:
                out[(county_of(r["county"]), r["precinct"])].add(int(m.group(1)))
    return out


def few(rec):
    hold = [cid for cid, v in rec["votes"].items() if v["total"] < FEW or max(v["dem"], v["rep"], v["other"]) == v["total"]]
    if hold:
        rec["too_few"] = hold
    return rec


def build(books, canvass, by_hd, clerk, contests=CONTESTS, judicial=None, psc=None, houses=HOUSE, primary=None, sources=None, expect=True):
    """Everything added up and checked. books: {year: rows}; canvass: {contest id: table}; by_hd: {contest id: table}
    for the district year; clerk: {contest id: {dem, rep, other, total, where}}. Returns (the file, what failed)."""
    judicial = JUDICIAL if judicial is None else judicial
    psc = PSC if psc is None else psc
    failed, ctl, state = [], {}, {}
    pack = lambda v: dict(zip(SIDES, v))                                                               # noqa: E731
    votes = {k: collections.defaultdict(dict) for k in KIND_NAMES}
    hd_counties = collections.defaultdict(set)
    tallies = {}
    for year, rows in books.items():
        tallies.update(tally_book(rows, [c for c in contests if c["year"] == year], f"{year} precinct workbook"))
    ids = [c["id"] for c in contests]
    first = {k[0]: k[1] for k in KINDS}
    given = {k: [c["id"] for c in contests if c["year"] >= first[k]] for k in KIND_NAMES}

    # which House district each precinct of the district year voted in, and which precincts are split
    drows = books.get(DISTRICT_YEAR) or []
    p_hd = precinct_districts(drows, "STATE REPRESENTATIVE")
    split = {k: sorted(v) for k, v in p_hd.items() if len(v) > 1}
    group = {h: h for h in range(1, houses + 1)}                            # House districts tied together by a split precinct

    def root(h):
        while group[h] != h:
            h = group[h]
        return h
    for hs in split.values():
        for h in hs[1:]:
            group[root(h)] = root(hs[0])

    for c in contests:
        cid, t = c["id"], tallies[c["id"]]
        side, n = t["side"], len(t["candidates"])
        sums = [sum(v[i] for v in t["precincts"].values()) for i in range(n)]
        state[cid] = four(sums, side)
        county = collections.defaultdict(lambda: [0] * n)
        for (code, _p), v in t["precincts"].items():
            _add(county[code], v)
        rec = ctl[cid] = {"sum_of_precincts": pack(state[cid]), "precincts": len(t["precincts"])}
        # the State Canvass: every county, every candidate
        table = canvass[cid]
        cols = match_columns(sums, table["total"], cid)
        col_sums = [sum(v[i] for v in table["rows"].values()) for i in range(len(table["total"]))]
        if len(table["rows"]) != len(county) or col_sums != table["total"]:
            failed.append(f"{cid}: the State Canvass table has {len(table['rows'])} counties and its columns add up to {col_sums}, its Total row "
                          f"reads {table['total']}")
        if cols is None:
            failed.append(f"{cid}: the workbook's candidates add up to {sums} and the State Canvass Total row reads {table['total']}")
            rec.update(official=None, equal=False)
        else:
            off = four([table["total"][j] for j in cols], side)
            differ = sorted(code for code in set(county) | set(table["rows"])
                            if [table["rows"].get(code, [None] * n)[j] for j in cols] != county.get(code))
            rec.update(official=pack(off), equal=off == state[cid] and not differ,
                       counties={"compared": len(county), "equal": len(county) - len(differ)})
            if differ:
                rec["counties"]["differ"] = [COUNTY_NAME[x] for x in differ]
                failed.append(f"{cid}: the workbook's precincts do not add up to the State Canvass row in {', '.join(COUNTY_NAME[x] for x in differ)}")
        rec.update(source=f"mt-sos-state-canvass-{c['year']}", where=f"the table headed {c['race']}, page {table.get('page')}",
                   read=f"from the document on {(sources or {}).get(('canvass', c['year']), {}).get('fetched', M._now())}")
        typed = CHECKED.get(cid) if expect else None
        if typed and typed != pack(state[cid]):
            failed.append(f"{cid}: the precincts add up to {pack(state[cid])}; this loader was checked against {typed}")
        if clerk.get(cid):
            k = clerk[cid]
            same = all(k[s] == state[cid][i] for i, s in enumerate(SIDES))
            rec["clerk"] = {"official": {s: k[s] for s in SIDES}, "equal": same, "where": k.get("where")}
            if not same:
                failed.append(f"{cid}: the precincts add up to {pack(state[cid])} and the Clerk of the House prints { {s: k[s] for s in SIDES} }")
        for code, v in county.items():
            votes["county"][code][cid] = four(v, side)
        for d, names in judicial.items():
            v = [0] * n
            for name in names:
                _add(v, county.get(FIPS[bare(name)], [0] * n))
            votes["judicial"][f"JD{d}"][cid] = four(v, side)

        # the districts of the district year, from the State Canvass by House district
        if c["year"] >= DISTRICT_YEAR:
            hd = by_hd[cid]
            hd_tot = {h: d["total"] for h, d in hd.items()}
            if sorted(hd) != list(range(1, houses + 1)) or any(v is None for v in hd_tot.values()):
                failed.append(f"{cid}: the State Canvass by House district has {len(hd)} districts with a Total row, not {houses}")
                continue
            width = len(next(iter(hd_tot.values())))
            hsums = [sum(v[i] for v in hd_tot.values()) for i in range(width)]
            hcols = match_columns(sums, hsums, cid)
            rows_bad = [h for h, d in hd.items() if [sum(v[i] for v in d["rows"].values()) for i in range(width)] != d["total"]]
            rec["house_districts"] = {"districts": len(hd), "add_up_to_statewide": hcols is not None, "rows_add_up_to_totals": not rows_bad}
            if rows_bad:
                failed.append(f"{cid}: in the State Canvass by House district the county rows of districts {rows_bad} do not add up to their Total rows")
            if hcols is None:
                failed.append(f"{cid}: the State Canvass by House district adds up to {hsums} and the precincts to {sums}")
                continue
            # the workbook's precincts against it: every district and county part where no precinct is split
            mine = collections.defaultdict(lambda: [0] * n)
            for key, v in t["precincts"].items():
                hs = p_hd.get(key)
                if not hs:
                    failed.append(f"{cid}: precinct {key[1]} of {COUNTY_NAME[key[0]]} voted in no State Representative contest")
                    continue
                _add(mine[(root(min(hs)), key[0])], v)
            theirs = collections.defaultdict(lambda: [0] * n)
            for h, d in hd.items():
                for code, v in d["rows"].items():
                    _add(theirs[(root(h), code)], [v[j] for j in hcols])
                    hd_counties[h].add(code)
            bad = sorted(k for k in set(mine) | set(theirs) if mine.get(k, [0] * n) != theirs.get(k, [0] * n))
            unsplit = sum(1 for h in hd if not any(h in hs for hs in split.values()))
            rec["house_districts"].update(compared_with_precincts=f"{unsplit} districts with no split precinct, county part by county part, and "
                                          f"the {len({root(h) for hs in split.values() for h in hs})} groups of districts that share a split "
                                          "precinct, each group together", equal=not bad)
            if bad:
                failed.append(f"{cid}: the precincts do not add up to the State Canvass by House district in "
                              + "; ".join(f"district {h} ({COUNTY_NAME[code]})" for h, code in bad[:8]))
            for h, v in hd_tot.items():
                votes["house"][str(h)][cid] = four([v[j] for j in hcols], side)
            for s in range(1, houses // 2 + 1):
                v = [0] * n
                for h in (2 * s - 1, 2 * s):
                    _add(v, [hd_tot[h][j] for j in hcols])
                votes["senate"][str(s)][cid] = four(v, side)
            for d, hs in psc.items():
                v = [0] * n
                for h in hs:
                    _add(v, [hd_tot[h][j] for j in hcols])
                votes["psc"][f"PSC{d}"][cid] = four(v, side)

    # which districts are which: the Senate pairs and the commission's lists, against the precincts that voted in them
    whole = {k: next(iter(v)) for k, v in p_hd.items() if len(v) == 1}
    seen_senate, senate_bad = set(), []
    for label, rws in (("2024 general", drows), ("2026 primary", primary or [])):
        hmap = whole if rws is drows else {k: next(iter(v)) for k, v in precinct_districts(rws, "STATE REPRESENTATIVE").items() if len(v) == 1}
        for key, ss in precinct_districts(rws, "STATE SENATOR").items():
            if key in hmap and len(ss) == 1:
                s = next(iter(ss))
                seen_senate.add(s)
                if hmap[key] not in (2 * s - 1, 2 * s):
                    senate_bad.append(f"{label}: precinct {key[1]} voted for Senate district {s} and House district {hmap[key]}")
    psc_bad, psc_seen = [], set()
    for key, ds in precinct_districts(drows, "PUBLIC SERVICE COMMISSIONER", areas=True).items():
        if key in whole and len(ds) == 1:
            d = next(iter(ds))
            psc_seen.add(d)
            if whole[key] not in psc.get(d, []):
                psc_bad.append(f"precinct {key[1]} voted for commission district {d} and House district {whole[key]}")
    jud_bad, jud_seen = [], set()
    for r in drows:
        m = re.fullmatch(r"JUDICIAL DISTRICT\s+(\d+)", r["area"].upper())
        if m:
            jud_seen.add(int(m.group(1)))
            names = {bare(x) for x in judicial.get(int(m.group(1)), [])}
            if bare(r["county"]) not in names and f"{r['county']} in district {m.group(1)}" not in jud_bad:
                jud_bad.append(f"{r['county']} in district {m.group(1)}")
    failed += [f"Senate districts: {x}" for x in senate_bad[:6]] + [f"commission districts: {x}" for x in psc_bad[:6]] \
        + [f"judicial districts: the workbook has {x}, the statute does not" for x in jud_bad[:6]]
    if drows and sorted(h for hs in psc.values() for h in hs) != list(range(1, houses + 1)):
        failed.append("commission districts: the statute's lists are not every House district once")
    if sorted(bare(x) for names in judicial.values() for x in names) != sorted(FIPS) and expect:
        failed.append("judicial districts: the statute's lists are not every county once")

    # the places
    places = {k: {} for k in KIND_NAMES}
    packed = lambda d: {cid: pack(d[cid]) for cid in ids if cid in d}                                  # noqa: E731
    for code in sorted(votes["county"]):
        places["county"][code] = few({"name": COUNTY_NAME[code], "votes": packed(votes["county"][code])})
    for h in sorted(votes["house"], key=int):
        places["house"][h] = few({"name": f"House District {h}", "counties": sorted(hd_counties[int(h)]), "votes": packed(votes["house"][h])})
    for s in sorted(votes["senate"], key=int):
        hs = [2 * int(s) - 1, 2 * int(s)]
        places["senate"][s] = few({"name": f"Senate District {s}", "house_districts": hs, "counties": sorted(set().union(*(hd_counties[h] for h in hs))),
                                   "votes": packed(votes["senate"][s])})
    for d in sorted(judicial):
        places["judicial"][f"JD{d}"] = few({"name": f"{M.ordinal(d)} Judicial District", "counties": sorted(FIPS[bare(x)] for x in judicial[d]),
                                            "votes": packed(votes["judicial"][f"JD{d}"])})
    for d in sorted(psc):
        if votes["psc"].get(f"PSC{d}"):
            places["psc"][f"PSC{d}"] = few({"name": f"Public Service Commission District {d}", "house_districts": psc[d],
                                            "counties": sorted(set().union(*(hd_counties[h] for h in psc[d]))), "votes": packed(votes["psc"][f"PSC{d}"])})
    kinds_ctl = {}
    for kind in KIND_NAMES:
        bad = [cid for cid in given[kind] if places[kind]
               and [sum(p["votes"][cid][s] for p in places[kind].values() if cid in p["votes"]) for s in SIDES] != state[cid]]
        kinds_ctl[kind] = "equal" if not bad else f"differs for {', '.join(bad)}"
        failed += [f"{kind}: its places do not add up to the statewide sum for {cid}" for cid in bad]
    if expect:
        if len(places["county"]) != 56 or len(places["judicial"]) != 22:
            failed.append(f"{len(places['county'])} counties and {len(places['judicial'])} judicial districts, not 56 and 22")
        if len(places["house"]) != HOUSE or len(places["senate"]) != SENATE or len(places["psc"]) != 5:
            failed.append(f"{len(places['house'])} House, {len(places['senate'])} Senate and {len(places['psc'])} commission districts, not 100, 50 and 5")

    records = []
    for c in contests:
        t = tallies[c["id"]]
        others = [PARTY_WORDS.get(p, p) for i, (p, _n) in enumerate(t["candidates"]) if i not in t["side"].values()]
        rec = {"id": c["id"], "date": ELECTIONS[c["year"]]["date"] if c["year"] in ELECTIONS else c.get("date"), "office": c["office"],
               "table": f"mt-sos-precinct-{c['year']}", "kinds": [k for k in KIND_NAMES if c["id"] in given[k] and places[k]]}
        for s, word in (("dem", "Democratic"), ("rep", "Republican")):
            rec[s] = {"party": word, "ticket": c[s], "column": f"the rows whose party is {'DEM' if s == 'dem' else 'REP'}"}
        rec["other"] = {"what": "every other candidate printed on the ballot, together; the workbook carries no write-in votes", "columns": others}
        rec["total"] = {"what": "the votes cast for the candidates printed on the ballot; the workbook gives no write-ins, under votes or over votes"}
        rec["statewide"] = pack(state[c["id"]])
        rec["official_source"] = f"mt-sos-state-canvass-{c['year']}"
        records.append(rec)
    kinds = {}
    for kind, year, what, key in KINDS:
        kinds[kind] = {"what": what, "key": key, "places": len(places[kind]), "contests": given[kind], "covers_the_state": True,
                       "places_with_a_contest_too_few": sum(1 for p in places[kind].values() if p.get("too_few"))}
        if year > min(c["year"] for c in contests):
            kinds[kind]["why_not_2020"] = REDRAWN
    kinds["house"]["note"] = ("The figures are the Total rows of the State Canvass by House district. " + (
        f"{len(split)} precincts are split between two House districts ({'; '.join(f'{k[1]}, {COUNTY_NAME[k[0]]}' for k in sorted(split))}); "
        "everywhere else the precincts add up to the same figures." if split else "No precinct is split between House districts."))
    kinds["senate"]["note"] = ("Senate district n is House districts 2n-1 and 2n. The precincts of the State Senator contests in the results "
                               f"show this for {len(seen_senate)} of the {houses // 2} districts.")
    kinds["judicial"]["note"] = "The counties of each district are those section 3-5-101, MCA, lists; every year's figures are those counties' votes."
    kinds["psc"]["note"] = ("The House districts of each district are those section 69-1-104, MCA, lists (the districts in force from 2024). The "
                            f"precincts of the 2024 commission contests agree for districts {', '.join(str(d) for d in sorted(psc_seen)) or 'none'}.")
    control = {"result": "equal" if not failed else "differs",
               "statement": ("For every contest the precinct rows add up, county by county and candidate by candidate, to the State Canvass, whose "
                             "columns add up to its own Total rows; for President and U.S. Senator the totals equal the Clerk of the House's "
                             "statistics; for 2024 the State Canvass by House district has every district, adds up to the statewide totals and "
                             "equals the precincts wherever no precinct is split; and every kind of place adds up to the statewide sum."
                             if not failed else "The sums do not all agree; see the differences."),
               "precincts": {str(y): len({(county_of(r["county"]), r["precinct"]) for r in rows}) for y, rows in sorted(books.items())},
               "contests": ctl, "kinds": kinds_ctl,
               "districts": {"split_precincts": [{"county": COUNTY_NAME[k[0]], "precinct": k[1], "house_districts": v} for k, v in sorted(split.items())],
                             "senate_districts_seen_in_results": sorted(seen_senate), "commission_districts_seen_in_results": sorted(psc_seen),
                             "judicial_districts_seen_in_results": sorted(jud_seen)},
               "notes": ["Declared write-in votes are canvassed on a separate abstract (State Canvass, Write-in) and are in neither the "
                         "workbook nor these figures; the Clerk of the House's totals leave them out too."]}
    doc = {"what": WHAT, "note": NOTE, "state": "MT", "generated": M._now(), "method_version": METHOD, "how_to_read": HOW_TO_READ,
           "vote_keys": list(SIDES), "too_few": {"fewer_than": FEW, "why": TOO_FEW}, "contests": records, "kinds": kinds, "control": control,
           "coverage": {"not_given": NOT_GIVEN, "no_2022": "Montana elected no senator and no statewide partisan officer in 2022."},
           "sources": [], "places": places}
    return doc, failed


# ---------------------------------------------------------------- the run

def clerk_totals(contests, cache, refresh, say):
    out, recs = {}, []
    for year in sorted({c["year"] for c in contests if c.get("clerk")}, reverse=True):
        src = CLERK[year]
        path = os.path.join(cache, src["file"])
        rec = {k: src[k] for k in ("id", "kind", "agency", "title", "url")}
        try:
            net.download(src["url"], path, 0 if refresh else KEEP_DAYS, tries=3, say=say)
            doc = N.read_clerk(path, "MONTANA")
            where = doc["where"].replace("North Dakota", "Montana")
            rec.update(fetched=M._day(path), sha256=M._sha_file(path), where=where)
            for c in contests:
                if c["year"] == year and c.get("clerk"):
                    lines = doc["sections"].get(c["clerk"]) or []
                    party = (lambda s: s) if c["clerk"] == "FOR PRESIDENTIAL ELECTORS" else (lambda s: s.rsplit(",", 1)[-1].strip())
                    dem = [v for label, v in lines if party(label).startswith("Democrat")]
                    rep = [v for label, v in lines if party(label) == "Republican"]
                    if len(dem) != 1 or len(rep) != 1:
                        raise ValueError(f"{c['clerk']}: not one Democratic and one Republican line")
                    total = sum(v for _l, v in lines)
                    out[c["id"]] = {"dem": dem[0], "rep": rep[0], "other": total - dem[0] - rep[0], "total": total, "where": where}
        except Exception as e:  # noqa: BLE001  the State Canvass is the control; the Clerk's page is a second one
            say(f"      {src['title']}: could not be read ({e}); the State Canvass alone is the control for {year}")
            rec["unread"] = f"could not be read on {M._now()}: {e}"
        recs.append(rec)
    return out, recs


def our_places(db):
    """{kind: ids} of our Montana places and races, read only; None when the database is not there."""
    if not db or not os.path.exists(db):
        return None
    con = sqlite3.connect(pathlib.Path(os.path.abspath(db)).as_uri() + "?mode=ro", uri=True)
    try:
        out = collections.defaultdict(set)
        for pid, in con.execute("SELECT id FROM sl_places WHERE source_id LIKE 'mt-%' AND kind = 'county'"):
            out["county"].add(pid)
        for office, level, jid, district in con.execute("SELECT office_kind, level, jurisdiction_id, district FROM sl_races WHERE state = 'MT'"):
            if office == "state_senate":
                out["senate"].add(str(district))
            elif office == "state_house":
                out["house"].add(str(district))
            elif office == "district_court":
                out["judicial"].add(str(jid))
            elif office == "public_service_commissioner":
                out["psc"].add(str(jid))
            elif level == "county":
                out["county"].add(str(jid))
        return dict(out)
    except sqlite3.Error:
        return None
    finally:
        con.close()


def load(say=print, out=OUT, cache=CACHE, db=DB, refresh=False):
    from ballot import pdftext
    say("    Montana place votes: the Secretary of State's precinct workbooks and the State Canvass (sosmt.gov/elections/results/)")
    net.patient_lookups()
    os.makedirs(cache, exist_ok=True)
    books, canvass, by_hd, facts, srcs = {}, {}, {}, {}, []
    for year in sorted(ELECTIONS, reverse=True):
        e = ELECTIONS[year]
        mine = [c for c in CONTESTS if c["year"] == year]
        path = fetch(e["book"], cache, refresh, say, b"PK")
        books[year] = read_book(path, f"{year} precinct workbook")
        precincts = len({(r["county"], r["precinct"]) for r in books[year]})
        say(f"      {year}: {precincts:,} precincts, {len(books[year]):,} rows (copy of {M._day(path)})")
        srcs.append({"id": f"mt-sos-precinct-{year}", "kind": "official results by precinct", "agency": AGENCY, "title": e["book"]["title"],
                     "url": e["book"]["url"], "listed_on": RESULTS_PAGE, "fetched": M._day(path), "sha256": M._sha_file(path), "precincts": precincts,
                     "read": "Of each row: the county, the precinct, the race and its area, the party, the votes and the name on the ballot; "
                             "of the names, only those of the candidates in the contests named here are looked at, to check which ticket a row is."})
        cpath = fetch(e["canvass"], cache, refresh, say, b"%PDF")
        tables = canvass_tables(pdftext.lines(cpath), sorted({c["race"] for c in mine}))
        for c in mine:
            canvass[c["id"]] = tables[c["race"]]
        facts[("canvass", year)] = {"fetched": M._day(cpath)}
        srcs.append({"id": f"mt-sos-state-canvass-{year}", "kind": "official canvass, by county", "agency": AGENCY + " (Board of State Canvassers)",
                     "title": e["canvass"]["title"], "url": e["canvass"]["url"], "listed_on": RESULTS_PAGE, "fetched": M._day(cpath),
                     "sha256": M._sha_file(cpath)})
        if "by_hd" in e:
            hpath = fetch(e["by_hd"], cache, refresh, say, b"%PDF")
            htables = by_hd_tables(pdftext.lines(hpath), sorted({c["race"] for c in mine}))
            for c in mine:
                by_hd[c["id"]] = htables[c["race"]]
            srcs.append({"id": f"mt-sos-state-canvass-by-hd-{year}", "kind": "official canvass, by House district", "agency": AGENCY + " (Board of State Canvassers)",
                         "title": e["by_hd"]["title"], "url": e["by_hd"]["url"], "listed_on": RESULTS_PAGE, "fetched": M._day(hpath),
                         "sha256": M._sha_file(hpath)})
    primary = None
    if os.path.exists(PRIMARY_2026):
        try:
            primary = [{k: r[k] for k in ("county", "precinct", "race", "area")}
                       for r in read_book(PRIMARY_2026, "2026 primary precinct workbook", need=("county", "precinct", "race", "area", "votes"))]
            srcs.append({"id": "mt-sos-precinct-2026-primary", "kind": "official results by precinct (a check only)", "agency": AGENCY,
                         "title": "2026 Primary Election Precinct by Precinct Report", "url": RESULTS_PAGE, "sha256": M._sha_file(PRIMARY_2026),
                         "read": "Only which Senate and House district each precinct voted in; no votes and no names are taken from it."})
        except Stop:
            primary = None
    srcs += statutes(cache, refresh, say)
    say("    the second control: Clerk of the U.S. House (President and U.S. Senator)")
    clerk, clerk_recs = clerk_totals(CONTESTS, cache, refresh, say)
    doc, failed = build(books, canvass, by_hd, clerk, primary=primary, sources=facts)
    doc["sources"] = srcs + clerk_recs

    p = doc["places"]
    say(f"    added up: {len(p['county'])} counties; {len(p['judicial'])} judicial districts; for 2024, {len(p['house'])} House, {len(p['senate'])} "
        f"Senate and {len(p['psc'])} Public Service Commission districts")
    for c in doc["contests"]:
        k, s = doc["control"]["contests"].get(c["id"], {}), c["statewide"]
        say(f"    control, {c['id']}: Democratic {s['dem']:,}, Republican {s['rep']:,}, other {s['other']:,}, total {s['total']:,}: "
            + ("equal to" if k.get("equal") else "DIFFERS from") + f" the State Canvass ({(k.get('counties') or {}).get('equal')} of "
            f"{(k.get('counties') or {}).get('compared')} counties)"
            + (f"; Clerk of the House {'equal' if k['clerk']['equal'] else 'DIFFERS'}" if k.get("clerk") else "")
            + (f"; by House district {'equal' if k['house_districts'].get('equal') else 'DIFFERS'}" if k.get("house_districts") else ""))
    if failed:
        for line in failed:
            say(f"    CHECK: {line}")
        raise Stop(f"    Montana place votes: the control did not hold; {os.path.basename(out)} was not written. Nothing was changed to make it fit.")
    ours = our_places(db)
    if ours:
        covered = {}
        for kind in KIND_NAMES:
            idset = ours.get(kind, set())
            without = sorted(i for i in idset if not p[kind].get(i, {}).get("votes"))
            covered[kind] = {"places": len(idset), "with_votes": len(idset) - len(without), "without": without}
            say(f"      our {kind} places with votes: {len(idset) - len(without):,} of {len(idset):,}" + (f" (none for {', '.join(without[:12])})" if without else ""))
        doc["coverage"]["our_places"] = dict(covered, read=f"sl_places and sl_races in {os.path.basename(db)}, opened read-only on {M._now()}")
    M.write(out, doc)
    say(f"    Montana place votes: {len(doc['contests'])} contests, {sum(len(v) for v in p.values()):,} places, control {doc['control']['result']}; "
        f"{os.path.relpath(out, HERE)} ({os.path.getsize(out) / 1e6:.2f} MB)")
    return doc


# ---------------------------------------------------------------- the arithmetic and the readers on made-up tables

def selftest(say=print):
    checks = []

    def check(what, got, want):
        checks.append(got == want)
        say(f"      {'ok  ' if got == want else 'FAIL'} {what}" + ("" if got == want else f": got {got!r}, expected {want!r}"))

    head = ("County ID", "CountyName", "PrecinctName", "Race Name", "Area", "Last Name", "Party", "Votes", "Name On Ballot")
    pre = [("BEAVERHEAD", "P1", 1, (30, 50, 2)), ("BEAVERHEAD", "P2", 2, (10, 5, 0)), ("LEWIS AND CLARK", "P1", 3, (40, 20, 1)),
           ("LEWIS AND CLARK", "P2", (3, 4), (7, 9, 1)), ("BIG HORN", "P1", 4, (0, 12, 0))]
    sheet = [head]
    for county, p, hds, (d, r, o) in pre:
        for party, name, v in (("DEM", "ANN ALDER", d), ("REP", "BO BIRCH", r), ("LIB", "CY CEDAR", o)):
            sheet.append(("01", county, p, "GOVERNOR", "Montana", name.split()[-1], party, v, name))
        for h in (hds if isinstance(hds, tuple) else (hds,)):
            sheet.append(("01", county, p, f"STATE REPRESENTATIVE DISTRICT {h}", f"HD {h:03d}", "X", "NON", 1, "X"))
            sheet.append(("01", county, p, f"STATE SENATOR DISTRICT {(h + 1) // 2}", f"SD {(h + 1) // 2:02d}", "X", "NON", 1, "X"))
    rows = book_rows(sheet, "made-up workbook")
    check("the workbook is read by its headings", (len(rows), rows[0]["party"], rows[0]["votes"]), (len(sheet) - 1, "DEM", 30))
    canvass_lines = [(1, 0, t) for t in ("2024 STATEWIDE GENERAL ELECTION CANVASS", "GOVERNOR", "BO BIRCH ANN ALDER CY CEDAR", "BEAVERHEAD 55 40 2",
                                         "BIG HORN 12 0 0", "LEWIS AND 29 47 2", "CLARK", "Total 96 87 4", "Page 1 of 1")]
    table = canvass_tables(canvass_lines, ["GOVERNOR"])["GOVERNOR"]
    check("the canvass table is read, a wrapped county name too", (table["rows"]["30049"], table["total"]), ([29, 47, 2], [96, 87, 4]))
    hd_lines = [(1, 0, t) for t in ("2024 STATEWIDE GENERAL ELECTION CANVASS", "GOVERNOR", "ANN ALDER CY CEDAR BO BIRCH", "HD 1 BEAVERHEAD 30 2 50",
                                    "Total 30 2 50", "HD 2 BEAVERHEAD 10 0 5", "Total 10 0 5", "1 of 2", "2024 STATEWIDE GENERAL ELECTION CANVASS",
                                    "HD 3 LEWIS AND 43 1 24", "CLARK", "Total 43 1 24", "HD 4 BIG HORN 0 0 12", "LEWIS AND 4 1 5", "CLARK", "Total 4 1 17")]
    hd = by_hd_tables(hd_lines, ["GOVERNOR"], houses=4)["GOVERNOR"]
    check("the canvass by House district is read across a page", (sorted(hd), hd[4]["rows"], hd[3]["total"]),
          ([1, 2, 3, 4], {"30003": [0, 0, 12], "30049": [4, 1, 5]}, [43, 1, 24]))
    contests = [{"id": "2024-governor", "year": 2024, "office": "Governor", "race": "GOVERNOR", "dem": "Ann Alder", "rep": "Bo Birch", "clerk": None}]
    judicial, psc = {1: ["Beaverhead", "Big Horn"], 2: ["Lewis and Clark"]}, {1: [1, 2, 3, 4]}
    doc, failed = build({2024: rows}, {"2024-governor": table}, {"2024-governor": hd}, {}, contests, judicial, psc, houses=4, expect=False)
    p = doc["places"]
    check("the control holds on the made-up tables", (doc["control"]["result"], failed), ("equal", []))
    check("a county adds up its precincts", p["county"]["30001"]["votes"]["2024-governor"], {"dem": 40, "rep": 55, "other": 2, "total": 97})
    check("a House district takes the canvass's figures, a split precinct's share included", p["house"]["4"]["votes"]["2024-governor"],
          {"dem": 4, "rep": 17, "other": 1, "total": 22})
    check("and reaches both its counties", p["house"]["4"]["counties"], ["30003", "30049"])
    check("a Senate district is two House districts", p["senate"]["2"]["votes"]["2024-governor"], {"dem": 47, "rep": 41, "other": 2, "total": 90})
    check("a judicial district is whole counties", p["judicial"]["JD1"]["votes"]["2024-governor"], {"dem": 40, "rep": 67, "other": 2, "total": 109})
    check("a commission district is the statute's House districts", p["psc"]["PSC1"]["votes"]["2024-governor"], {"dem": 87, "rep": 96, "other": 4, "total": 187})
    check("the split precinct is named", [x["precinct"] for x in doc["control"]["districts"]["split_precincts"]], ["P2"])
    check("a place where every vote went one way is marked too_few", p["county"]["30003"].get("too_few"), ["2024-governor"])
    wrong = dict(table, rows=dict(table["rows"], **{"30001": [54, 41, 2]}))
    check("a county that differs from the canvass is caught", build({2024: rows}, {"2024-governor": wrong}, {"2024-governor": hd}, {}, contests,
                                                                   judicial, psc, houses=4, expect=False)[0]["control"]["result"], "differs")
    moved = {h: dict(d) for h, d in hd.items()}
    moved[1], moved[2] = dict(hd[1], rows={"30001": [29, 2, 50]}, total=[29, 2, 50]), dict(hd[2], rows={"30001": [11, 0, 5]}, total=[11, 0, 5])
    check("a vote filed under the wrong House district is caught", build({2024: rows}, {"2024-governor": table}, {"2024-governor": moved}, {}, contests,
                                                                        judicial, psc, houses=4, expect=False)[0]["control"]["result"], "differs")
    try:
        tally_book(book_rows(sheet + [sheet[1]], "made-up workbook"), contests, "made-up workbook")
        caught = False
    except Stop:
        caught = True
    check("a candidate twice in a precinct stops the loader", caught, True)
    pages = {"judicial": b"<p>(1) 1st district: Lewis and Clark and Broadwater Counties; (2) 2nd district: Silver Bow County; (3) 3rd district: "
                         b"Deer Lodge, Granite, and Powell Counties.</p>",
             "psc": b"<p>(a) first district: 19, 20, 21; (b) second district: 39, 40; and (e) fifth district: 3, 99. (2) During</p>"}
    j, q = read_statutes(pages)
    check("the statutes are read", (j, q), ({1: ["Lewis and Clark", "Broadwater"], 2: ["Silver Bow"], 3: ["Deer Lodge", "Granite", "Powell"]},
                                           {1: [19, 20, 21], 2: [39, 40], 5: [3, 99]}))
    check("county codes follow the Census order, McCone before Madison", (FIPS["MCCONE"], FIPS["MADISON"], FIPS["YELLOWSTONE"], FIPS[bare("Lewis & Clark")]),
          ("30055", "30057", "30111", "30049"))
    say(f"    self-test: {sum(checks)} of {len(checks)} checks passed")
    return all(checks)


def main(argv=None):
    ap = argparse.ArgumentParser(description="How each Montana place voted in past partisan general elections -> ballot/lean/mt_place_votes.json")
    ap.add_argument("--out", default=OUT, help="the file to write (default: ballot/lean/mt_place_votes.json)")
    ap.add_argument("--cache", default=CACHE, help="where downloads are kept (default: states_cache/mt_local)")
    ap.add_argument("--db", default=DB, help="the ballot database to count our places in, opened read-only (default: ballot_local_2026.sqlite)")
    ap.add_argument("--refresh", action="store_true", help="ask for every file again, even when the cached copies are there")
    ap.add_argument("--selftest", action="store_true", help="check the arithmetic and the readers on made-up tables; downloads and writes nothing")
    a = ap.parse_args(argv)
    if a.selftest:
        raise SystemExit(0 if selftest() else 1)
    load(out=a.out, cache=a.cache, db=a.db, refresh=a.refresh)


if __name__ == "__main__":
    main()
