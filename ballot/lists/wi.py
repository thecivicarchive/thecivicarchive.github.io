"""
Wisconsin: the Elections Commission's "Candidates on Ballot by Election" for the 2026 General Election (the PDF on
elections.wi.gov/elections, printed 8/25/2026) and its County by County Report for the August 11, 2026 Partisan Primary
(a workbook on the 2026 Partisan Primary Election Results page). The Commission's site answers scripts with 403
Forbidden and serves browsers, so both files are carried out of the Browser pane into ballot_cache/wi/ (the New York
method in CLAUDE.md: an in-page fetch, handed back as base64 and rebuilt on disk). This loader downloads nothing.

The list gives, under each office ("Office : REPRESENTATIVE IN CONGRESS DISTRICT 3 Incumbent: ..."), one line per
candidate in ballot order: the order number, a committee number for some, the name and the party (a long party name
wraps onto the next line: "Wisconsin" / "Green"), then "Total Number of ... Candidates :N", which is checked. The
primary workbook has one sheet per office and party, named in its Document map; the "Office Totals:" row gives each
candidate's votes, and SCATTERING (write-ins) counts toward the total but is not a candidate. A field is a party
primary with two candidates or more; the one on the November list for that party advanced.
"""

import os
import re

import openpyxl

from ballot.common import fold, house_id, party_code, record_source, senate_id
from ballot.pdftext import lines

LIST = "wi_candidates_on_ballot_2026_general.pdf"
PRIMARY_BOOK = "wi_county_by_county_2026_partisan_primary.xlsx"
LIST_URL = "https://elections.wi.gov/media/40951/download"
PRIMARY_URL = ("https://elections.wi.gov/sites/default/files/documents/"
               "County%20by%20County%20Report_Partisan%20Primary%202026_All%20State%20Contests.xlsx")
PRIMARY = "2026-08-11"
OFFICE = re.compile(r"^Office : (?P<office>.+?)(?: Incumbent: .*)?$")
TOTAL = re.compile(r"^Total Number of (?P<office>.+) Candidates :(?P<n>\d+)$")
ROW = re.compile(r"^(?P<order>\d+) (?:\d{7} )?(?P<rest>.+)$")
PARTIES = ("Wisconsin Green", "Republican", "Democratic", "Libertarian", "Constitution", "Independent", "Wisconsin", "Green")


def race_of(office):
    m = re.fullmatch(r"REPRESENTATIVE IN CONGRESS DISTRICT (\d+)", office.strip())
    if m:
        return house_id("WI", int(m.group(1)))
    return senate_id("WI", 1) if office.strip() == "UNITED STATES SENATOR" else None


def general_list(path):
    out, counts, race, office = [], {}, None, None
    L = lines(path)
    i = 0
    while i < len(L):
        text = L[i][2]
        i += 1
        m = OFFICE.match(text)
        if m:
            office = m.group("office").strip()
            race = race_of(office)
            continue
        m = TOTAL.match(text)
        if m:
            if race:
                got = sum(1 for r in out if r[0] == race)
                if got != int(m.group("n")):
                    raise SystemExit(f"Wisconsin: the list counts {m.group('n')} candidates for {office}; {got} were read")
            race = None
            continue
        if not race:
            continue
        m = ROW.match(text)
        if not m:
            continue
        rest = m.group("rest")
        party = next((p for p in PARTIES if rest.endswith(" " + p)), None)
        if not party:
            raise SystemExit(f"Wisconsin: no party read on the line for {office}: {text!r}")
        name = rest[: -len(party)].strip()
        if party == "Wisconsin" and i < len(L) and L[i][2] == "Green":      # the party's name wrapped onto the next line
            party, i = "Wisconsin Green", i + 1
        out.append((race, int(m.group("order")), name, party))
    return out


def primary_fields(path):
    """{(race, party): [(name, votes, total)]} for every congressional party primary, from the Office Totals rows."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = [r[1] for r in wb["Document map"].iter_rows(values_only=True) if len(r) > 1 and r[1]]
    out = {}
    for k, title in enumerate(names):
        m = re.fullmatch(r"(REPRESENTATIVE IN CONGRESS DISTRICT \d+|UNITED STATES SENATOR) - (\w+)", str(title).strip())
        if not m:
            continue
        ws = wb.worksheets[k + 1]
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        if not any(str(c or "").strip() == str(title).strip() for r in rows[:5] for c in r):
            raise SystemExit(f"Wisconsin: sheet {ws.title} does not carry the office its Document map names ({title})")
        head = next(j for j, r in enumerate(rows) if r and str(r[0] or "").strip() == "County")
        cands = rows[head + 1]
        totals = next(r for r in rows if r and str(r[0] or "").strip().startswith("Office Totals"))
        cast_col = next(j for j, c in enumerate(rows[head]) if str(c or "").strip() == "Total Votes Cast")
        total = int(totals[cast_col] or 0)
        field = [(str(c).strip(), int(totals[j] or 0)) for j, c in enumerate(cands) if c and str(c).strip() and str(c).strip() != "SCATTERING"]
        out[(race_of(m.group(1)), m.group(2))] = [(n, v, total) for n, v in field]
    return out


def load(con, cache, say=print):
    folder = os.path.join(cache, "wi")
    path, book = os.path.join(folder, LIST), os.path.join(folder, PRIMARY_BOOK)
    if not os.path.exists(path):
        say(f"    Wisconsin: {LIST} is not in {folder}. The Commission's site refuses scripts; carry the file out of the Browser pane "
            "from elections.wi.gov/elections (\"Candidates on Ballot By Election_November 3 2026 General Election.pdf\").")
        return 0
    listed = general_list(path)
    rows = [(race, "general", "2026-11-03", name, party, party_code(party), order, 0, 0, None, None, None, None, None, "wi-wec-2026-candidates-on-ballot", None)
            for race, order, name, party in listed]
    fields = primary_fields(book) if os.path.exists(book) else {}
    nominee = {(race, party): fold(name) for race, _o, name, party in listed}
    for (race, party), field in fields.items():
        if len(field) < 2:
            continue
        code = {"Democratic": "DEM", "Republican": "REP"}.get(party, party[:3].upper())
        for name, votes, total in field:
            rows.append((race, f"primary-{code}", PRIMARY, name, party, party_code(party), None, 0, 0, votes,
                         round(100 * votes / total, 1) if total else None, "advanced" if nominee.get((race, party)) == fold(name) else "lost",
                         None, None, "wi-wec-2026-primary-county", None))
    with con:
        con.execute("DELETE FROM candidates WHERE race_id LIKE '2026-WI-%'")
        con.executemany("INSERT OR REPLACE INTO candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", rows)
        record_source(con, "wi-wec-2026-candidates-on-ballot", path=path, level="federal", state="WI", kind="official candidate list",
                      agency="Wisconsin Elections Commission", title="Candidates on Ballot by Election: 2026 General Election, 11/3/2026",
                      url=LIST_URL, rows=len(listed),
                      note="The Commission's site refuses scripts; the file was carried out of the Browser pane. Ballot order as printed; each office's total checked.")
        if fields:
            record_source(con, "wi-wec-2026-primary-county", path=book, level="federal", state="WI", kind="official results",
                          agency="Wisconsin Elections Commission", title="County by County Report: 2026 Partisan Primary, all state contests (August 11, 2026)",
                          url=PRIMARY_URL, rows=sum(len(v) for v in fields.values()),
                          note="Votes from each office's Office Totals row; write-ins (SCATTERING) count in the total but are not listed.")
    n = len(listed)
    say(f"    Wisconsin: {len({r[0] for r in listed})} House districts, {n} candidates on the November ballot; "
        f"{sum(1 for v in fields.values() if len(v) > 1)} party primaries with a field, votes from the official county-by-county report")
    return n
