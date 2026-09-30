"""
Arizona: the Secretary of State's candidate listing for the November 3, 2026 general election, its 2026 General
Write-In and Withdrawn Candidate List, and its Official Canvass of the July 21, 2026 primary (the statewide canvass of
August 6, 2026). Arizona has nine House seats, on the 2024 lines, and no Senate race in 2026.

Every host of the Secretary of State answered scripts on 2026-09-30 with a Cloudflare challenge ("Just a moment...",
Cf-Mitigated: challenge): azsos.gov, apps.azsos.gov (where the canvass and the candidate listings are filed),
results.arizona.vote and my.arizona.vote. The challenge is never worked around, so this loader downloads nothing. It
reads the files John saves from his own browser into ballot_cache/az/ under these names:

  az_2026_general_candidates.pdf   the "Candidate Listing" for the 2026 General Election, linked from the 2026
                                   Election Info page (azsos.gov/elections/election-information/2026-election-info).
                                   For 2024 the same listing was "2024 General Election: Federal, Statewide, and
                                   Legislative Candidates", apps.azsos.gov/election/2024/ge_cand/2024_General_Candidates_Web.pdf.
                                   A workbook or CSV of the listing may be saved instead (.xlsx or .csv, same name).
  az_2026_general_write_in_withdrawn.pdf   "2026 General Write-In and Withdrawn Candidate List" (azsos.gov/media/718);
                                   optional: its write-in candidates are added (write_in 1, no ballot position) and its
                                   withdrawn candidates are taken off the November ballot if the listing still has them.
  az_2026_primary_canvass.pdf      "2026 Primary Election", the Official Canvass,
                                   apps.azsos.gov/election/2026/canvass/20260806_Primary_Canvass.pdf

Saved files are checked by their first bytes (%PDF, or PK for a workbook): a saved challenge page is refused.

None of these files had been seen when this loader was written (every copy was behind the challenge), so its layout
rules are general and its checks strict: whatever it cannot read whole stops it with a message, and nothing is stored.
It was tried only on made-up files of each layout below; check the first run on the real files row by row against
the PDFs, and use --probe (at the end of this note) to adjust the rules.

The listings are read as tables. On each page the row of column headings is found by its words (Candidate Name or
Name, or Last Name and First Name; Party or Office; and, where present, District, Status and Ballot Order) and each
heading's position gives its column's band; only the bands of those headings are turned into text. Any other column
(a city, a mailing address, a telephone, an e-mail, a website) stays in the file and is never read. An office printed
as a heading across the page ("U.S. Representative in Congress - District No. 1"), or once for the rows under it in
an Office column, applies to those rows; so does a party printed alone on a line. A line with nothing in the name
column is used only when it is such a heading, and never when it carries an address, a ZIP code, a telephone or an
e-mail. A listing with no row of headings is read line by line: under a congressional heading a line is taken only
when it is a name and a party (printed apart, or marked "(DEM)" or "- Democratic"), or a name under a line naming
the party alone; any other line there stops the loader, which prints no line of the file. Workbooks and CSV files are
read by the same headings. No ballot order is printed unless a Ballot Order column says so; the listing's own order
is kept. A status saying withdrawn, removed or disqualified, or a Withdrawn section, leaves the candidate off; a
Write-In section or status makes a write-in candidate (write_in 1, no ballot position).

The canvass is a results book, with no personal details. Each U.S. Representative contest ("U.S. Representative in
Congress - District No. 1", with the party in the title, on the line under it, or after each candidate's name) is a
table: either a row per county and a column per candidate, closed (or opened) by a Total row, where each column's
counties must add up to its total; or a row per candidate and a column per county and Total, where each row's counties
must add up to its total and the rows to the Total row. A title with no figures under it (a certificate's sentence, a
table of contents) is passed over; a title with figures that make neither table stops the loader. Columns of
registered voters, ballots cast, turnout, over- and undervotes are left out; unnamed write-in votes count toward a
party primary's total but are not a candidate. A contest printed in two halves is joined, and the same table printed
twice is counted once. A party primary becomes a field when two or more names were printed on that party's ballot (a
name marked as a write-in candidate counts toward the total but not toward that threshold); votes are the canvass's,
pct is the share of that party primary's votes including write-ins, and the candidate with the most votes (Arizona
nominates by plurality) advanced; a tie stops the loader. A nominee missing from the November listing keeps
"advanced" and is noted, and any district where the canvass's leader and the listing's party candidate differ is
printed.

Parties are kept as printed, written out where abbreviated: DEM Democratic, REP Republican, LBT/LIB Libertarian,
GRN/GRE Green, AZI/AIP Arizona Independent, NL No Labels, IND Independent, PND Party Not Designated. A bare word
after a name is never taken for a party (Green is a surname too); a party is read after a name only in brackets,
after a dash or comma, as a short code, or from its own column. Primary elections are stored as primary-DEM,
primary-REP, primary-LIB, primary-GRE, primary-AZI (Arizona Independent Party) and primary-NL. Names written "Last,
First" are turned round; names printed in capitals are shown in ordinary capitals (with the capitals the
congress-legislators roster gives a sitting member's name when its letters are the same), with a note.

python -m ballot.lists.az --probe canvass|general|writein prints what a reader needs to fix the layout rules: for
the canvass, its lines on the pages with a House contest (results only); for the listings, the heading rows and the
allowlisted cells of the congressional rows, never a whole row.
"""

import csv
import io
import os
import re
import sqlite3
import statistics
import sys

import openpyxl

from ballot.common import HERE, fold, house_id, name_parts, party_code, record_source
from ballot.lists.tx import proper
from ballot.pdftext import PDF, join, rows as pdf_rows

STATE = "AZ"
PRIMARY = "2026-07-21"
GENERAL_DATE = "2026-11-03"
SEATS = 9
GENERAL_STEM = "az_2026_general_candidates"
WRITEIN_FILE = "az_2026_general_write_in_withdrawn.pdf"
CANVASS_FILE = "az_2026_primary_canvass.pdf"
INFO_PAGE = "https://azsos.gov/elections/election-information/2026-election-info"
WRITEIN_URL = "https://azsos.gov/media/718"
CANVASS_URL = "https://apps.azsos.gov/election/2026/canvass/20260806_Primary_Canvass.pdf"
AGENCY = "Arizona Secretary of State, Elections Division"
BLOCKED = ("azsos.gov, apps.azsos.gov, results.arizona.vote and my.arizona.vote answer scripts with a Cloudflare "
           "challenge; the file was saved from a browser.")

COUNTIES = ("apache", "cochise", "coconino", "gila", "graham", "greenlee", "la paz", "maricopa", "mohave", "navajo",
            "pima", "pinal", "santa cruz", "yavapai", "yuma")
# (as printed, upper case) -> (written out, election code); tried in this order, so Arizona Independent comes before Independent
PARTIES = (
    (r"ARIZONA INDEPENDENT(?: PARTY)?|AZ INDEPENDENT(?: PARTY)?|AZI|AIP", "Arizona Independent", "AZI"),
    (r"DEMOCRATIC(?: PARTY)?|DEMOCRAT|DEM", "Democratic", "DEM"),
    (r"REPUBLICAN(?: PARTY)?|REP|GOP", "Republican", "REP"),
    (r"LIBERTARIAN(?: PARTY)?|LBT|LIB", "Libertarian", "LIB"),
    (r"GREEN(?: PARTY)?|GRN|GRE", "Green", "GRE"),
    (r"NO LABELS(?: PARTY)?|NL|NLP", "No Labels", "NL"),
    (r"INDEPENDENT|IND", "Independent", None),
    (r"PARTY NOT DESIGNATED|PND|NONPARTISAN|NON-PARTISAN", "Party Not Designated", None),
)
PARTY_ANY = "|".join(p for p, _l, _c in PARTIES)
CODES = r"DEM|REP|LBT|LIB|GRN|GRE|AZI|AIP|NL|NLP|IND|PND|GOP"
PARTY_PAREN = re.compile(rf"\(\s*({PARTY_ANY})\s*\)", re.I)
# a party after a name: in brackets, after a dash or comma, or a short code; a bare word such as Green is a surname
PARTY_TAIL = re.compile(rf"(?:\s*[-,–]\s*({PARTY_ANY})|\s+({CODES}))\s*$", re.I)


def split_party(text):
    """(the text without its party mark, the party as printed or None): 'SHAH, AMISH (DEM)' -> ('SHAH, AMISH', 'DEM')."""
    m = PARTY_PAREN.search(text or "") or PARTY_TAIL.search(text or "")
    if not m or m.start() == 0:
        return text, None
    return (text[:m.start()] + text[m.end():]).strip(" -,–"), next(g for g in m.groups() if g)

HOUSE = re.compile(r"(?i)(?:\b(?:U\.?\s*S\.?|United\s+States)\s+)?\bRep(?:resentative|\.)?\s+in\s+Congress\b\D{0,40}?\b(\d{1,2})\b"
                   r"|\bU\.?\s*S\.?\s+(?:House|Rep(?:resentative|\.)?)\b\D{0,40}?\b(\d{1,2})\b"
                   r"|\bCongressional\s+Dist(?:rict|\.)?\s*(?:No\.?\s*)?(\d{1,2})\b|\bCD\s*-?\s*(\d{1,2})\b")
REP_WORDS = re.compile(r"(?i)\bRep(?:resentative|\.)?\s+in\s+Congress\b|\bU\.?\s*S\.?\s+(?:House|Representative)\b")
OTHER = re.compile(r"(?i)^(?:U\.?\s*S\.?\s+Senat|United States Senat|Governor|Secretary of State|Attorney General|State Treasurer|"
                   r"Treasurer|Superintendent of Public|State Mine Inspector|Mine Inspector|Corporation Commission|State Senat|"
                   r"State Rep|Legislative District|Justice|Judge|Supreme Court|Court of Appeals|Superior Court|Precinct Committee)")
CONTACTISH = re.compile(r"@|www\.|https?:|\b\d{5}(?:-\d{4})?\b|\(\d{3}\)|\b\d{3}[-.]\d{3}[-.]\d{4}\b")
NOT_HEAD = re.compile(r"(?i)^\(?(?:vote for|number of|precincts|official|unofficial|canvass|page \d|printed|\d{1,2}/\d{1,2}/\d{2,4})")
SECTION_WI = re.compile(r"(?i)\bwrite[- ]?ins?\b")
SECTION_WD = re.compile(r"(?i)\bwithdr[a-z]*\b")
SECTION_LINE = re.compile(r"(?i)^\W*(?:official |certified |qualified )?(?:write[- ]?ins?|withdrawn|withdrawals?)"
                          r"(?: candidates?)?(?: list)?\W*$")
GONE = re.compile(r"(?i)withdr|remov|disqual|deceas|denied|reject|stricken|struck|off ballot")
NUM = re.compile(r"^\d{1,3}(?:,\d{3})+$|^\d+$")
TOTAL_ROW = re.compile(r"(?i)^(?:grand |statewide |state |district |contest )?totals?:?$")
STATS = re.compile(r"(?i)\bregist|\bballots\b|\bturnout\b|\btimes (?:cast|counted)\b|\bprecincts?\b|\bunder ?votes?\b|"
                   r"\bover ?votes?\b|\bblanks?\b|\btotal votes\b|\bvotes cast\b|\bcast votes\b|\bvoters\b")
WRITE_COL = re.compile(r"(?i)^(?:unassigned |unqualified |unofficial |unresolved |other |total |misc\.? )?write[- ]?ins?(?: votes?)?$")
WRITE_MARK = re.compile(r"(?i)\s*[\(\[]?\b(?:official |certified )?write[- ]in(?: candidate)?\b[\)\]]?\s*")
WIN_MARK = re.compile(r"^\s*[*✓✔]\s*|\s*[*✓✔]\s*$")
HEAD_NOISE = re.compile(r"(?i)\b(?:votes?|percent|pct|count|candidates?)\b|%")
LABEL_HEAD = re.compile(r"(?i)^(?:county|counties|jurisdiction|name|candidate)$")
SUFFIX = re.compile(r"(?i)^(?:JR|SR|II|III|IV|V)\.?$")
NAME_OK = re.compile(r"^[A-Za-zÀ-ɏ .,'\"()’“”-]+$")
FOOTER = re.compile(r"(?i)^(?:page \d+(?: of \d+)?|printed|run date|report date|\d{1,2}/\d{1,2}/\d{2,4})\b")
CAPS_LIST = "Arizona's list prints names in capitals; they are shown here in ordinary capitals."
CAPS_CANVASS = "Arizona's canvass prints names in capitals; they are shown here in ordinary capitals."
WRITE_IN_NOTE = "Write-in candidate: the name is not printed on the ballot."
NOT_ON_LIST = "Won the primary but is not on the November list."

HEADS = {
    "name": ("candidate name", "candidate", "name", "ballot name", "name on ballot", "candidate ballot name",
             "name as it appears on ballot", "name as it will appear on ballot", "candidate name on ballot"),
    "last": ("last name", "last"),
    "first": ("first name", "first"),
    "middle": ("middle name", "middle", "middle initial"),
    "suffix": ("suffix",),
    "party": ("party", "party affiliation", "political party", "party name", "party preference"),
    "office": ("office", "contest", "office sought", "race", "office name", "contest name", "office title"),
    "district": ("district", "dist", "district no", "district number", "district name"),
    "status": ("status", "candidate status", "type", "filing type", "write-in/withdrawn", "withdrawn/write-in", "category"),
    "order": ("ballot order", "ballot position", "order", "position"),
}


# ---------------------------------------------------------------------------------------------------------- words

def party_of(text):
    """(written out, election code or None) for a party as printed, or None when the text is not a party."""
    t = re.sub(r"\s+", " ", (text or "").strip().strip("()").strip()).upper()
    if not t:
        return None
    for pattern, label, code in PARTIES:
        if re.fullmatch(pattern, t):
            printed = (text or "").strip().strip("()").strip()
            if len(t) <= 4 or printed == printed.upper():          # an abbreviation, or a name in capitals: written out
                return label, code
            return re.sub(r"\s+", " ", printed), code
    return None


def election_code(label, code):
    return code or re.sub(r"[^A-Z]", "", (label or "").upper())[:3] or "OTH"


def race_of(text):
    m = HOUSE.search(text or "")
    if not m:
        return None
    d = int(next(g for g in m.groups() if g))
    if not 1 <= d <= SEATS:
        raise SystemExit(f"Arizona: a U.S. Representative heading names district {d}; Arizona has {SEATS}")
    return house_id(STATE, d)


def roster_spellings():
    """Upper-case forms of the sitting Arizona members' names, with the capitals the congress-legislators roster gives them."""
    path = os.path.join(HERE, "congress_119.sqlite")
    if not os.path.exists(path):
        return {}
    rec = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    out = {}
    try:
        for full, first, last in rec.execute("SELECT official_full, first_name, last_name FROM legislators WHERE is_current = 1 AND state = 'AZ'"):
            for form in (full, f"{first} {last}"):
                if form:                                   # the letters as printed, with the roster's capitals: never another name
                    out[re.sub(r"\s+", " ", form.upper())] = form
    finally:
        rec.close()
    return out


ROSTER = None


def shown(raw):
    """(name as shown, whether it was printed in capitals). 'SHAH, AMISH' -> 'Amish Shah'; 'Crane, Eli Jr.' -> 'Eli Crane Jr.'"""
    global ROSTER
    name = re.sub(r"\s+", " ", str(raw or "")).strip()
    name = WIN_MARK.sub("", name).strip()
    name = re.sub(r"(?i)\s*\((?:i|inc\.?|incumbent)\)\s*$", "", name).strip()
    given, family = None, None
    last, sep, rest = name.partition(",")
    if sep and rest.strip() and not SUFFIX.match(rest.strip()):
        words = last.split()
        rest_words = rest.replace(",", " ").split()
        suffixes = [w for w in words + rest_words if SUFFIX.match(w)]
        given = " ".join(w for w in rest_words if not SUFFIX.match(w))
        family = " ".join([w for w in words if not SUFFIX.match(w)] + suffixes)
        name = f"{given} {family}"
    caps = any(c.isalpha() for c in name) and name == name.upper()
    if not caps:
        return name, False
    if ROSTER is None:
        ROSTER = roster_spellings()
    if name in ROSTER:
        return ROSTER[name], True
    # a family name printed on its own ("DEL VECCHIO, BRIAN") keeps a capital on its first word: Brian Del Vecchio
    fixed = f"{proper(given)} {proper(family)}" if family else proper(name)
    return re.sub(r"(['\"(“])([a-z])", lambda m: m.group(1) + m.group(2).upper(), fixed), True


def same_person(a, b):
    ga, fa = name_parts(a)
    gb, fb = name_parts(b)
    return bool(fa) and fa == fb and (not ga or not gb or ga[0][:1] == gb[0][:1])


def saved(path, kind):
    """The file's bytes, refused if it is not what it should be (a saved Cloudflare page, a Page Not Found)."""
    data = open(path, "rb").read()
    want = {"pdf": b"%PDF", "xlsx": b"PK"}.get(kind)
    if want and not data.lstrip()[:4].startswith(want):
        raise SystemExit(f"Arizona: {os.path.basename(path)} is not a {kind.upper()} file (it starts {data[:15]!r}); "
                         "save the file itself from the browser, not the page around it")
    return data


# ---------------------------------------------------------------------------------------------- the listings

def head_field(text):
    t = re.sub(r"\s+", " ", re.sub(r"[^a-z/ -]+", "", (text or "").lower())).strip(" -:")
    for field, words in HEADS.items():
        if t in words:
            return field
    return None


def cells(runs, gap=6.0):
    """A printed row's runs grouped into cells [(x0, x1, runs)]: a gap wider than `gap` points starts the next cell."""
    out = []
    for r in sorted(runs, key=lambda r: r[0]):
        if out and r[0] - out[-1][1] <= max(gap, 0.6 * r[2]):
            out[-1][2].append(r)
            out[-1][1] = max(out[-1][1], r[4])
        else:
            out.append([r[0], r[4], [r]])
    return [(a, b, rs) for a, b, rs in out]


def header_bands(runs):
    """The column bands [(lo, hi, field or None)] when this row is a row of column headings, else None. A row of
    headings names the candidate (Candidate Name, or Last Name and First Name) and the party or the office."""
    cs = cells(runs)
    fields = [head_field(join(rs)) for _a, _b, rs in cs]
    named = "name" in fields or ("last" in fields and "first" in fields)
    if not named or not ("party" in fields or "office" in fields):
        return None
    bands = []
    for i, (a, b, _rs) in enumerate(cs):
        lo = -1e9 if i == 0 else (cs[i - 1][1] + a) / 2
        hi = 1e9 if i == len(cs) - 1 else (b + cs[i + 1][0]) / 2
        bands.append((lo, hi, fields[i]))
    return bands


def compose(rec):
    """A name from separate Last / First / Middle / Suffix columns, written 'Last, First Middle Suffix'."""
    if "last" in rec or "first" in rec:
        given = " ".join(y for y in (rec.get("first", ""), rec.get("middle", "")) if y)
        rec["name"] = rec.get("name") or ", ".join(x for x in (rec.get("last", ""), given) if x)
        if rec.get("suffix") and rec["name"]:
            rec["name"] += " " + rec["suffix"]
    return rec


def no_name_row(t, state, report):
    """A printed row with nothing in the name column: an office heading, a party alone, a write-in or withdrawn
    section, or anything else (a title, a footer, a second line of an entry), which is let go unread."""
    if CONTACTISH.search(t):
        return
    if HOUSE.search(t) or OTHER.match(t):
        state.update(heading=t, party=None, office=None)
        report["headings"].append(t if HOUSE.search(t) else t[:40])
        return
    p = party_line(t)
    if p and len(t) <= 40:
        state["party"] = t
        return
    if len(t) <= 80 and SECTION_WI.search(t) and not SECTION_WD.search(t):
        state["section"] = "write-in"
    elif len(t) <= 80 and SECTION_WD.search(t) and not SECTION_WI.search(t):
        state["section"] = "withdrawn"


def records_from_bands(pdf, report):
    """The rows of a listing read by column bands: [{field: text, "heading": ..., "section": ...}], or None when no
    page carries a row of headings. Only allowlisted bands are turned into text."""
    out, bands, found = [], None, False
    state = {"heading": "", "party": None, "section": None, "office": None}
    for pno, (page, res) in enumerate(pdf.pages(), start=1):
        for _y, runs in pdf_rows(pdf, page, res):
            b = header_bands(runs)
            if b:
                bands, found = b, True
                report["heading_rows"].append([join(rs) for _a, _b, rs in cells(runs)])
                continue
            if not bands:
                continue
            got = {}
            for r in runs:
                field = next((f for lo, hi, f in bands if lo <= r[0] < hi), None)
                if field:
                    got.setdefault(field, []).append(r)
            rec = compose({f: join(rs) for f, rs in got.items()})
            nm = rec.get("name", "")
            if not nm or HOUSE.search(nm) or REP_WORDS.search(nm) or OTHER.match(nm) or SECTION_LINE.match(nm) \
                    or (party_line(nm) and len(got) == 1):
                no_name_row(join(runs), state, report)      # a heading printed from the left margin across the columns
                continue
            if rec.get("office"):
                state["office"] = (rec["office"], rec.get("district", ""))
            elif any(f == "office" for _lo, _hi, f in bands) and state["office"]:
                rec["office"], d = state["office"]            # an office printed once for the rows under it
                rec.setdefault("district", d)
            if not rec.get("party") and state["party"]:
                rec["party"] = state["party"]
            rec.update({"heading": state["heading"], "section": state["section"], "page": pno})
            out.append(rec)
    return out if found else None


def records_from_lines(pdf, report):
    """A listing with no row of headings, line by line: under a congressional heading, a line is taken only when it
    is a name and a party (either way round), or a name under a line naming the party alone. Any other line there
    stops the loader, which prints none of them."""
    out, odd = [], 0
    state = {"heading": "", "party": None, "section": None, "office": None}
    party_first = re.compile(rf"^\(?({PARTY_ANY})\)?\s+(.+)$", re.I)
    for pno, (page, res) in enumerate(pdf.pages(), start=1):
        for _y, runs in pdf_rows(pdf, page, res):
            t = join(runs)
            if FOOTER.match(t) or NOT_HEAD.match(t):
                continue
            if HOUSE.search(t) or OTHER.match(t) or (party_line(t) and len(t) <= 40) or SECTION_LINE.match(t) or \
                    (len(t) <= 80 and SECTION_WI.search(t) and SECTION_WD.search(t)):
                no_name_row(t, state, report)
                continue
            if not race_of(state["heading"]):
                continue
            name = party = None
            cs = [join(rs) for _a, _b, rs in cells(runs)]
            if len(cs) == 2 and party_of(cs[1]):                # a name and a party printed apart
                name, party = cs
            elif len(cs) == 2 and party_of(cs[0]):
                party, name = cs
            elif len(cs) == 1 and split_party(t)[1]:            # 'Name (DEM)', 'Name - Democratic'
                name, party = split_party(t)
            elif len(cs) == 1 and party_first.match(t) and len(party_first.match(t).group(1)) <= 4:
                party, name = party_first.match(t).group(1), party_first.match(t).group(2)
            elif len(cs) == 1 and (state["party"] or state["section"] == "write-in"):
                name, party = t, state["party"]
            if name and NAME_OK.match(name) and len(name.split()) <= 7 and not CONTACTISH.search(t):
                out.append({"name": name, "party": party or "", "heading": state["heading"], "section": state["section"], "page": pno})
            else:
                odd += 1
    if odd:
        raise SystemExit(f"Arizona: the listing has no row of column headings, and {odd} line(s) under U.S. Representative "
                         "headings are not a name and a party alone; run python -m ballot.lists.az --probe general")
    return out


def records_from_book(path):
    """A listing saved as a workbook or CSV: the heading row found by its words, the allowlisted columns only."""
    if path.lower().endswith(".csv"):
        raw = open(path, "rb").read().decode("utf-8-sig", "replace")
        table = list(csv.reader(io.StringIO(raw)))
    else:
        wb = openpyxl.load_workbook(io.BytesIO(saved(path, "xlsx")), read_only=True, data_only=True)
        table = [["" if c is None else str(c) for c in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
    for i, r in enumerate(table[:30]):
        fields = [head_field(c) for c in r]
        if "party" in fields and ("name" in fields or ("last" in fields and "first" in fields)):
            break
    else:
        raise SystemExit(f"Arizona: {os.path.basename(path)} has no heading row with a name and a party column")
    idx = {}
    for j, f in enumerate(fields):
        if f and f not in idx:
            idx[f] = j
    out, last_office = [], None
    for r in table[i + 1:]:
        rec = compose({f: re.sub(r"\s+", " ", r[j]).strip() for f, j in idx.items() if j < len(r) and r[j]})
        if rec.get("office"):
            last_office = (rec["office"], rec.get("district", ""))
        elif rec.get("name") and "office" in idx and last_office:
            rec["office"], d = last_office                 # an office written once for the rows under it
            rec.setdefault("district", d)
        if rec.get("name"):
            rec.update({"heading": "", "section": None, "page": None})
            out.append(rec)
    return out, [table[i][j] for j in idx.values()]


def listing(path, report=None):
    """[(race, name as printed, party as printed, status, order, section)] for the congressional rows of a listing."""
    report = report if report is not None else {"heading_rows": [], "headings": []}
    if path.lower().endswith((".xlsx", ".csv")):
        recs, heads = records_from_book(path)
        report["heading_rows"].append(heads)
    else:
        pdf = PDF(saved(path, "pdf"))
        recs = records_from_bands(pdf, report)
        if recs is None:
            recs = records_from_lines(pdf, report)
    out, no_party = [], 0
    for rec in recs:
        if rec.get("office"):
            district = rec.get("district", "")
            if re.fullmatch(r"(?:No\.?\s*)?\d{1,2}", district):
                district = "District " + district
            race = race_of(f"{rec['office']} {district}") if REP_WORDS.search(rec["office"]) or HOUSE.search(rec["office"]) else None
        else:
            race = race_of(rec["heading"])
        if not race:
            continue
        party = rec.get("party", "")
        if not party:
            rec["name"], party = split_party(rec["name"])
            party = party or ""
        if not party and not SECTION_WI.search(rec.get("status", "") + " " + (rec.get("section") or "")):
            no_party += 1
            continue
        out.append((race, rec["name"], party, rec.get("status", ""), rec.get("order", ""), rec.get("section")))
    if no_party:
        raise SystemExit(f"Arizona: {no_party} row(s) under U.S. Representative on {os.path.basename(path)} have a name and no party; "
                         "run python -m ballot.lists.az --probe general")
    return out


# ------------------------------------------------------------------------------------------------ the canvass

def words(runs):
    """Runs split into words [(x0, x1, text, size)], a run's width shared out by its characters."""
    out = []
    for x0, _y, size, t, x1 in runs:
        t = t.replace("\xa0", " ")
        if len(t.split()) <= 1:
            if t.strip():
                out.append((x0, x1, t.strip(), size))
            continue
        n = len(t)
        for m in re.finditer(r"\S+", t):
            out.append((x0 + (x1 - x0) * m.start() / n, x0 + (x1 - x0) * m.end() / n, m.group(), size))
    return sorted(out, key=lambda w: w[0])


def label_and_numbers(runs):
    """(the words before the first number, [(x0, x1, number)]) for a printed row."""
    ws = words(runs)
    k = next((i for i, w in enumerate(ws) if NUM.match(w[2])), None)
    if k is None:
        return " ".join(w[2] for w in ws), []
    return " ".join(w[2] for w in ws[:k]), [(w[0], w[1], int(w[2].replace(",", ""))) for w in ws[k:] if NUM.match(w[2])]


def county_of(label):
    t = fold(label).replace(" county", "").strip()
    return t if t in COUNTIES else None


OFFICE_WORDS = re.compile(r"(?i)\b(?:U\.?\s*S\.?|United\s+States)\s+Rep(?:resentative|\.)?(?:\s+in\s+Congress)?\b"
                          r"|\bRep(?:resentative|\.)?\s+in\s+Congress\b"
                          r"|\bU\.?\s*S\.?\s+House\b|\bCongressional\b|\bDist(?:rict|\.)?\s*(?:No\.?\s*)?\d{1,2}\b")


def contest_party(texts):
    """The party a contest's title (or a line alone just above it) names, as (label, code), or None. The office's own
    words are set aside first, so the "Rep." of "U.S. Rep. in Congress" is never read as Republican."""
    for t in texts:
        rest = OFFICE_WORDS.sub(" ", t)
        for m in reversed(list(re.finditer(rf"\b({PARTY_ANY})\b", rest, re.I))):
            p = party_of(m.group(1))
            if p:
                return p
    return None


def party_line(t):
    """(label, code) when a printed line is a party alone ('DEMOCRATIC', 'Republican Party', 'LBT Primary'), else None."""
    return party_of(re.sub(r"(?i)\s*\b(?:party|primary|nomination|candidates?)\b.*$", "", t or "").strip())


def header_text(rs):
    """A column's heading runs as one text: printed lines top to bottom (text set on its side, left to right)."""
    flat = [r for r in rs if r[4] - r[0] >= 0.5]
    side = [r for r in rs if r[4] - r[0] < 0.5]
    lines_ = {}
    for r in flat:
        lines_.setdefault(round(r[1] / 2), []).append(r)
    parts = [join(v) for _k, v in sorted(lines_.items(), key=lambda kv: -kv[0])]
    parts += [r[3].strip() for r in sorted(side, key=lambda r: (r[0], r[1]))]
    return re.sub(r"\s+", " ", HEAD_NOISE.sub(" ", " ".join(parts))).strip(" -")


def table_a(block, total_at, first):
    """Counties as rows, candidates as columns, closed (or opened) by a Total row: ([(heading, total)], county rows)."""
    _label, cols = label_and_numbers(block[total_at][1])
    rights = [b for _a, b, _v in cols]
    gaps = [rights[i + 1] - rights[i] for i in range(len(rights) - 1)]
    width = statistics.median(gaps) if gaps else 60.0
    if total_at > first:
        county_rows = block[first:total_at]
    else:                                                   # the Total row first, the counties under it
        county_rows = []
        for y, runs in block[total_at + 1:]:
            lab, ns = label_and_numbers(runs)
            if not ns or not county_of(lab):
                break
            county_rows.append((y, runs))
    sums, counties = [0] * len(cols), 0
    for _y, runs in county_rows:
        lab, ns = label_and_numbers(runs)
        if not ns:
            continue
        if not county_of(lab):
            raise SystemExit(f"Arizona: in the canvass, a row labelled {lab!r} sits among the county rows of a U.S. Representative table")
        counties += 1
        seen = set()
        for _a, b, v in ns:
            k = min(range(len(cols)), key=lambda i: abs(rights[i] - b))
            if abs(rights[k] - b) > max(6.0, 0.45 * width) or k in seen:
                raise SystemExit(f"Arizona: in the canvass, {lab}'s figures do not line up with the Total row's columns")
            seen.add(k)
            sums[k] += v
        if len(seen) != len(cols):
            raise SystemExit(f"Arizona: in the canvass, {lab}'s row has {len(seen)} figures; the Total row has {len(cols)}")
    if counties and sums != [v for _a, _b, v in cols]:
        raise SystemExit(f"Arizona: in the canvass, the counties do not add up to the Total row ({sums} against {[v for *_x, v in cols]})")
    heads = [[] for _ in cols]
    for _y, runs in block[:min(first, total_at)]:
        if NOT_HEAD.match(join(runs)):
            continue
        for r in runs:
            text = r[3].strip()
            if not text or not HEAD_NOISE.sub("", text).strip(" %()") or LABEL_HEAD.match(text):
                continue
            c = r[0] if r[4] - r[0] < 0.5 else (r[0] + r[4]) / 2          # text set on its side has no width across
            if c < rights[0] - 1.3 * width:
                continue                                                   # the column of county names
            k = next((i for i, rt in enumerate(rights) if rt + 3 >= c), None)
            if k is not None:
                heads[k].append(r)
    out = []
    for k, rs in enumerate(heads):
        text = header_text(rs)
        if not text:
            raise SystemExit(f"Arizona: in the canvass, column {k + 1} of a U.S. Representative table has no heading")
        out.append((text, cols[k][2]))
    return out, counties


def table_b(block, hdr):
    """Candidates as rows, counties (and Total) as columns: ([(row label, total)], county columns)."""
    heads = words(block[hdr][1])
    cols, i = [], 0
    while i < len(heads):                                   # county names of two words (La Paz, Santa Cruz) are joined
        a, b, t, _s = heads[i]
        if i + 1 < len(heads) and county_of(t + " " + heads[i + 1][2]):
            cols.append(((a + heads[i + 1][1]) / 2, county_of(t + " " + heads[i + 1][2])))
            i += 2
            continue
        if county_of(t):
            cols.append(((a + b) / 2, county_of(t)))
        elif TOTAL_ROW.match(t):
            cols.append(((a + b) / 2, "total"))
        i += 1
    centres = [c for c, _n in cols]
    out, pending, n_counties, closing = [], [], sum(1 for _c, n in cols if n != "total"), None
    for _y, runs in block[hdr + 1:]:
        lab, ns = label_and_numbers(runs)
        if not ns:
            if lab and not FOOTER.match(lab) and not NOT_HEAD.match(lab):
                pending.append(lab)
            continue
        lab = " ".join(pending + [lab]).strip()
        pending = []
        vals = {}
        for a, b, v in ns:
            k = min(range(len(cols)), key=lambda j: abs(centres[j] - (a + b) / 2))
            if cols[k][1] in vals:
                raise SystemExit(f"Arizona: in the canvass, the row {lab!r} has two figures under {cols[k][1]}")
            vals[cols[k][1]] = v
        county_sum = sum(v for n, v in vals.items() if n != "total")
        total = vals.get("total", county_sum)
        if "total" in vals and n_counties and county_sum != total:
            raise SystemExit(f"Arizona: in the canvass, the row {lab!r} does not add up ({county_sum} against {total})")
        if TOTAL_ROW.match(lab):
            closing = total
            break
        out.append((lab, total))
    if closing is not None and closing != sum(v for lab, v in out if not STATS.search(lab)):
        raise SystemExit(f"Arizona: in the canvass, a U.S. Representative table's rows do not add up to its Total row")
    return out, n_counties


def is_title(texts, i):
    """How many printed lines (1 to 3) make up a U.S. Representative contest's title starting at line i, else 0."""
    if HOUSE.search(texts[i]):
        return 1
    if REP_WORDS.search(texts[i]):
        for k in (2, 3):
            if HOUSE.search(" ".join(texts[i:i + k])):
                return k
    return 0


def canvass_tables(path, report=None):
    """[{race, party, page, title, columns: [(heading, votes)], counties, shape}] for every U.S. Representative table."""
    pdf = PDF(saved(path, "pdf"))
    out, section = [], None                                 # section: a party printed alone above the tables it heads
    current = None                                          # the House contest whose tables are being read
    for pno, (page, res) in enumerate(pdf.pages(), start=1):
        rws = pdf_rows(pdf, page, res)
        texts = [join(rs) for _y, rs in rws]
        starts, i = [], 0
        while i < len(texts):
            used = is_title(texts, i)
            if used:
                current = " ".join(texts[i:i + used])
                title = current
                if i + used < len(texts) and party_line(texts[i + used]) and len(texts[i + used]) <= 40:
                    title += " " + texts[i + used]
                    used += 1
                starts.append((i, used, title, section))
                i += used
                continue
            if OTHER.match(texts[i]):
                starts.append((i, 0, None, section))        # another office's contest: it closes the one before
                current = None
            elif party_line(texts[i]) and len(texts[i]) <= 40 and not re.search(r"\d", texts[i]):
                section = party_line(texts[i])
                if current:                                 # a party alone under a House contest opens that party's table
                    starts.append((i, 1, f"{current} {texts[i]}", section))
            i += 1
        for n, (i, used, title, section_then) in enumerate(starts):
            if title is None:
                continue
            end = starts[n + 1][0] if n + 1 < len(starts) else len(rws)
            block = rws[i + used:end]
            race = race_of(title)
            party = contest_party([title]) or section_then
            read = [label_and_numbers(rs) for _y, rs in block]
            first = next((k for k, (lab, ns) in enumerate(read) if ns and (county_of(lab) or TOTAL_ROW.match(lab))), None)
            total_at = next((k for k, (lab, ns) in enumerate(read) if ns and TOTAL_ROW.match(lab)), None)
            hdr = next((k for k, (_y, rs) in enumerate(block)       # a row of county names (two, or one and Total): a candidate
                        if not read[k][1] and                         # named Graham or Yuma is not a heading of counties
                        sum(1 for w in words(rs) if county_of(w[2])) + sum(1 for w in words(rs) if TOTAL_ROW.match(w[2])) >= 2), None)
            if total_at is not None and first is not None and (hdr is None or hdr > first):
                cols, counties = table_a(block, total_at, first)
                shape = "counties as rows"
            elif hdr is not None:
                cols, counties = table_b(block, hdr)
                shape = "candidates as rows"
            elif not any(ns for _lab, ns in read):
                if report is not None:                      # a certificate's sentence or a table of contents: no figures under it
                    report.append(f"page {pno}: {title} (no table under it)")
                continue
            else:
                raise SystemExit(f"Arizona: page {pno} of the canvass has a U.S. Representative contest whose table is not read "
                                 f"({title!r}); run python -m ballot.lists.az --probe canvass")
            out.append({"race": race, "party": party, "page": pno, "title": title, "columns": cols, "counties": counties, "shape": shape})
            if report is not None:
                report.append(f"page {pno}: {title} -> {race} {party[0] if party else '(party after each name)'}: "
                              f"{'; '.join(f'{t} {v}' for t, v in cols)} ({counties} counties, {shape})")
    return out


def primary_contests(tables):
    """{(race, code): {"label", "code", "cands": {printed name: [votes, write_in]}, "writein": votes}} from the tables.
    A contest printed in two halves is joined; a table whose named candidates are all already read with the same
    figures is the same table printed again, and is counted once."""
    out = {}
    for tb in tables:
        named, unnamed = [], []
        for text, votes in tb["columns"]:
            if STATS.search(text):
                continue
            core, mark = split_party(WIN_MARK.sub("", text).strip())
            party = (party_of(mark) if mark else None) or tb["party"]
            if not party:
                raise SystemExit(f"Arizona: page {tb['page']} of the canvass has a column for {tb['race']} with no party "
                                 "in the contest's title or after the name")
            if WRITE_COL.match(core):
                unnamed.append((party, votes))
            else:
                named.append((party, WRITE_MARK.sub(" ", core).strip(" -,*"), votes, bool(WRITE_MARK.search(core))))
        def slot(party):
            return out.setdefault((tb["race"], party[1] or party[0]), {"label": party[0], "code": party[1], "cands": {}, "writein": 0})
        if named and all(slot(p)["cands"].get(n, [None])[0] == v for p, n, v, _w in named):
            continue
        for party, name, votes, write_in in named:
            c = slot(party)
            if name in c["cands"] and c["cands"][name][0] != votes:
                raise SystemExit(f"Arizona: the canvass gives {name} two different totals for {tb['race']} ({c['cands'][name][0]} and {votes})")
            c["cands"][name] = [votes, write_in]
        for party, votes in unnamed:
            slot(party)["writein"] += votes
    return {k: v for k, v in out.items() if v["cands"] or v["writein"]}


# ------------------------------------------------------------------------------------------------------ load

def find_general(folder):
    for ext in (".pdf", ".xlsx", ".csv"):
        p = os.path.join(folder, GENERAL_STEM + ext)
        if os.path.exists(p):
            return p
    return None


def published(path):
    """The date a listing says it was updated or printed ('Updated 9/15/2026', 'As of September 15, 2026'), from
    the lines of its first page that say so, else ''. The table's own columns are never searched."""
    try:
        pdf = PDF(open(path, "rb").read())
        page, res = pdf.pages()[0]
        said = [join(rs) for _y, rs in pdf_rows(pdf, page, res)]
    except Exception:  # noqa: BLE001  a date is a courtesy; never a reason to stop
        return ""
    months = ("january", "february", "march", "april", "may", "june", "july", "august", "september", "october", "november", "december")
    for t in said:
        m = re.search(r"(?i)\b(?:updated|as of|printed|published|revised|report date|run date|date prepared)\b:?\s*(.*)$", t)
        if not m:
            continue
        rest = m.group(1)
        d = re.match(r"(\d{1,2})/(\d{1,2})/(20\d\d)\b", rest)
        if d:
            return f"{d.group(3)}-{int(d.group(1)):02d}-{int(d.group(2)):02d}"
        d = re.match(r"(?:[A-Z][a-z]+day,?\s+)?([A-Z][a-z]+)\.? (\d{1,2}),? (20\d\d)\b", rest)
        if d and d.group(1).lower() in months:
            return f"{d.group(3)}-{months.index(d.group(1).lower()) + 1:02d}-{int(d.group(2)):02d}"
    return ""


def what_to_save(folder):
    return (f"Save, from a browser, into {folder}: {GENERAL_STEM}.pdf (the 'Candidate Listing' for the 2026 General "
            f"Election on {INFO_PAGE}), {WRITEIN_FILE} ({WRITEIN_URL}, optional) and {CANVASS_FILE} ({CANVASS_URL}).")


def load(con, cache, say=print):
    folder = os.path.join(cache, "az")
    os.makedirs(folder, exist_ok=True)
    gpath, wpath, cpath = find_general(folder), os.path.join(folder, WRITEIN_FILE), os.path.join(folder, CANVASS_FILE)
    wpath = wpath if os.path.exists(wpath) else None
    cpath = cpath if os.path.exists(cpath) else None
    if not gpath and not cpath:
        say(f"    Arizona: nothing stored. The Secretary of State's hosts answer scripts with a Cloudflare challenge. {what_to_save(folder)}")
        return 0
    races = [r for (r,) in con.execute("SELECT race_id FROM races WHERE state = ? ORDER BY race_id", (STATE,))]

    # the November ballot
    general, gone, write_ins, order = [], [], [], {}
    if gpath:
        for race, raw, praw, status, given, section in listing(gpath):
            if GONE.search(status) or section == "withdrawn":
                gone.append((race, raw))
                continue
            wi = bool(SECTION_WI.search(status) or section == "write-in" or WRITE_MARK.search(raw))
            name, caps = shown(WRITE_MARK.sub(" ", raw).strip())
            p = party_of(praw) if praw else None
            label = p[0] if p else (praw.strip() or ("Write-in" if wi else ""))
            note = " ".join(x for x in (WRITE_IN_NOTE if wi else "", CAPS_LIST if caps else "") if x) or None
            if wi:
                write_ins.append(name)
                general.append([race, name, label, None, 1, note, "az-sos-2026-general-list"])
                continue
            order[race] = order.get(race, 0) + 1
            pos = int(given) if str(given).strip().isdigit() else order[race]
            general.append([race, name, label, pos, 0, note, "az-sos-2026-general-list"])
        missing = [r for r in races if not any(g[0] == r and not g[4] for g in general)]
        if missing:
            raise SystemExit(f"Arizona: no candidate read for {', '.join(missing)} on {os.path.basename(gpath)}; "
                             "run python -m ballot.lists.az --probe general")
    wi_rows, wi_gone = 0, []
    if wpath:
        for race, raw, praw, status, _given, section in listing(wpath):
            name, caps = shown(WRITE_MARK.sub(" ", raw).strip())
            if GONE.search(status) or section == "withdrawn":
                wi_gone.append(name)
                before = len(general)
                general = [g for g in general if not (g[0] == race and same_person(g[1], name))]
                if len(general) < before:
                    gone.append((race, raw))
                continue
            if SECTION_WI.search(status) or section == "write-in":
                if any(g[0] == race and same_person(g[1], name) for g in general):
                    continue
                p = party_of(praw) if praw else None
                wi_rows += 1
                write_ins.append(name)
                general.append([race, name, p[0] if p else (praw.strip() or "Write-in"), None, 1,
                                " ".join(x for x in (WRITE_IN_NOTE, CAPS_LIST if caps else "") if x), "az-sos-2026-general-write-in"])

    # the primary fields
    rows, fields, contests, unmatched, differ = [], 0, {}, [], []
    if cpath:
        contests = primary_contests(canvass_tables(cpath))
        if not contests:
            raise SystemExit("Arizona: no U.S. Representative contest was read from the canvass; run python -m ballot.lists.az --probe canvass")
        found = sorted({r for r, _k in contests})
        if races and found != races:
            raise SystemExit(f"Arizona: the canvass's U.S. Representative contests are {found}; the races are {races}")
        for (race, _key), c in sorted(contests.items()):
            label, code = c["label"], c["code"]
            ranked = sorted(c["cands"].items(), key=lambda kv: -kv[1][0])
            nominee = [g for g in general if g[0] == race and not g[4] and party_of(g[2]) and party_of(g[2])[1] == code and code]
            if ranked and gpath and code and nominee and not any(same_person(g[1], shown(ranked[0][0])[0]) for g in nominee):
                differ.append(f"{race} {label}: the canvass's leader {shown(ranked[0][0])[0]}, the November list's {nominee[0][1]}")
            printed = [n for n, (_v, wi) in c["cands"].items() if not wi]
            if len(printed) < 2:
                continue
            fields += 1
            total = sum(v for v, _wi in c["cands"].values()) + c["writein"]
            if ranked[0][1][0] == ranked[1][1][0]:
                raise SystemExit(f"Arizona: the {race} {label} primary is a tie in the canvass; read the canvass's own note")
            winner = ranked[0][0]
            for name, (votes, wi) in ranked:
                shown_name, caps = shown(name)
                outcome = "advanced" if name == winner else "lost"
                note = [WRITE_IN_NOTE if wi else "", CAPS_CANVASS if caps else ""]
                if outcome == "advanced" and gpath and not any(same_person(g[1], shown_name) for g in nominee):
                    note.append(NOT_ON_LIST)
                    unmatched.append(f"{race} {label}: {shown_name}")
                rows.append((race, f"primary-{election_code(label, code)}", PRIMARY, shown_name, label, party_code(label), None, 0,
                             int(wi), votes, round(100 * votes / total, 1) if total else None, outcome, None, None,
                             "az-sos-2026-primary-canvass", " ".join(x for x in note if x) or None))

    seen, doubled = set(), 0
    kept = []
    for g in general:
        if (g[0], g[1]) in seen:
            doubled += 1
            continue
        seen.add((g[0], g[1]))
        kept.append(g)
    general = kept
    out = [(g[0], "general", GENERAL_DATE, g[1], g[2], "W" if g[4] and not party_of(g[2]) else party_code(g[2]), g[3], 0, g[4],
            None, None, None, None, None, g[6], g[5]) for g in general]
    with con:
        con.execute("DELETE FROM candidates WHERE race_id LIKE '2026-AZ-%'")
        con.executemany("INSERT OR REPLACE INTO candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", out + rows)
        if gpath:
            record_source(con, "az-sos-2026-general-list", path=gpath, level="federal", state=STATE, kind="official candidate list",
                          agency=AGENCY, title="Candidate Listing, 2026 General Election (November 3, 2026): U.S. Representative in Congress",
                          url=INFO_PAGE, published=published(gpath) if gpath.endswith(".pdf") else "",
                          rows=sum(1 for g in general if g[6] == "az-sos-2026-general-list"),
                          note=f"{BLOCKED} Only the name, party, office, district, status and ballot order columns are read. "
                               f"Withdrawn or removed, left off: {len(gone)}. Write-in candidates listed: {len(write_ins) - wi_rows}. "
                               "Ballot order as the listing gives it, else the listing's own order.")
        if wpath:
            record_source(con, "az-sos-2026-general-write-in", path=wpath, level="federal", state=STATE, kind="official candidate list",
                          agency=AGENCY, title="2026 General Write-In and Withdrawn Candidate List: U.S. Representative in Congress",
                          url=WRITEIN_URL, published=published(wpath), rows=wi_rows + len(wi_gone),
                          note=f"{BLOCKED} Write-in candidates added: {wi_rows}. Withdrawn: {len(wi_gone)}.")
        if cpath:
            record_source(con, "az-sos-2026-primary-canvass", path=cpath, level="federal", state=STATE, kind="official results",
                          agency=AGENCY, title="Official Canvass, 2026 Primary Election (July 21, 2026): U.S. Representative in Congress",
                          url=CANVASS_URL, published="2026-08-06", rows=len(rows),
                          note=f"{BLOCKED} Each table's counties checked against its totals. Unnamed write-in votes count toward "
                               "each party primary's total. The candidate with the most votes advanced (Arizona nominates by plurality)"
                               + (f"; not on the November list: {'; '.join(unmatched)}." if unmatched else "."))
    n = len(out)
    if gpath:
        say(f"    Arizona: {len({g[0] for g in general})} House districts, {n} candidates on the November ballot "
            f"({len(write_ins)} write-in, {len(gone)} withdrawn or removed left off); "
            + (f"{fields} party primaries with a field, votes from the Official Canvass" if cpath else
               f"primary fields not loaded (save {CANVASS_FILE} from {CANVASS_URL})"))
    else:
        say(f"    Arizona: the November listing is not saved, so no November candidates are stored; {fields} party primaries "
            f"with a field, votes from the Official Canvass. {what_to_save(folder)}")
    if doubled:
        say(f"      the listing names {doubled} candidate(s) twice for the same race; each is stored once")
    if differ:
        say("      the canvass's leader and the November list's party candidate differ: " + "; ".join(differ))
    return n


# ------------------------------------------------------------------------------------------------------ probe

def probe(kind, cache):
    """What a reader needs to fix the layout rules, and nothing a privacy rule keeps back."""
    folder = os.path.join(cache, "az")
    if kind == "canvass":
        path = os.path.join(folder, CANVASS_FILE)
        pdf = PDF(saved(path, "pdf"))
        for pno, (page, res) in enumerate(pdf.pages(), start=1):
            texts = [join(rs) for _y, rs in pdf_rows(pdf, page, res)]
            if any(HOUSE.search(t) or REP_WORDS.search(t) for t in texts):
                print(f"--- page {pno}")
                for t in texts:
                    print("   ", t)
        report = []
        try:
            canvass_tables(path, report)
        finally:
            print("\n".join(report))
        return
    path = find_general(folder) if kind == "general" else os.path.join(folder, WRITEIN_FILE)
    report = {"heading_rows": [], "headings": []}
    got = listing(path, report)
    print(f"heading rows found: {len(report['heading_rows'])}")
    for h in report["heading_rows"][:3]:
        print("   ", [c for c in h if head_field(c)], f"and {sum(1 for c in h if not head_field(c))} other heading(s)")
    print(f"office headings: {len(report['headings'])}; congressional: {sum(1 for h in report['headings'] if HOUSE.search(h))}")
    for race, name, party, status, order, section in got:
        print("   ", race, "|", name, "|", party, "|", status, "|", order, "|", section or "")


if __name__ == "__main__":
    sys.path.insert(0, HERE)
    from ballot.common import CACHE
    if len(sys.argv) == 3 and sys.argv[1] == "--probe":
        probe(sys.argv[2], CACHE)
    else:
        print("python -m ballot.lists.az --probe canvass|general|writein")
