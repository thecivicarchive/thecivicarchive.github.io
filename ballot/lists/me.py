"""
Maine: the Secretary of State's own files (Bureau of Corporations, Elections and Commissions, Division of Elections).
Maine has two House districts and the class 2 Senate seat (Susan M. Collins). The primary was on June 9, 2026.

  November ballot   "2026 General Election Candidates List" (an Excel workbook, "2026 General Candidate List -
                    FINAL.xlsx", posted 7/27/2026 at the replacement filing deadline), linked from
                    maine.gov/sos/elections-voting/upcoming-elections. One sheet, "November 3, 2026", every office;
                    columns are taken by name, only Office, Dist, Party, Last Name, First Name, Middle Name and
                    Suffix. Offices are the Division's codes ("2026 Office Listing of Candidates Guide", the page's
                    "2026 Office Abbreviation Key"): US = US SENATOR, CG = REPRESENTATIVE TO CONGRESS. The list gives
                    no ballot order; its own order (alphabetical by surname within each office) is kept. Parties are
                    the qualified parties' letters, written out as the page names the parties (D Democratic, R
                    Republican, G Green Independent, L Libertarian); a non-party candidate's designation is printed in
                    full and kept as printed.
  withdrawals       "2026 Candidate Withdrawals and Replacement Candidate Nominations" (a PDF, updated as they come;
                    "after the June 9, 2026 Primary"), linked from the same page. The page says a primary winner who
                    withdrew has been taken off the general list and the replacement added. Read by fixed column bands
                    (Ofc., Dist., Pty, Withdrawn Candidate, Date, Replacement Candidate), measured on the file, with
                    every heading checked where the layout says it sits (a moved heading stops the loader; headings
                    are centred over their columns, so a band begins left of its heading); only those bands become
                    text, and the Residence Municipality and Date Received columns are never read. A federal
                    candidate who withdrew and is still on the general list would be left off and counted.
  write-ins         "Declared Write-In Candidates" for the November 3, 2026 General Election (a PDF, dated 8/25/26),
                    linked from the same page; read the same way (Office, District, Party, Name; the Residence and
                    Municipalities Affected columns never read). Its federal rows are added as declared write-in
                    candidates: write_in 1, no ballot position, the designation as printed.
  primary list      "2026 Primary Candidate List" (Excel, 3/16/2026) and "2026 Primary Candidate Withdrawals and
                    Replacement Candidate Nominations" (PDF, as of May 22, 2026), linked from
                    maine.gov/sos/elections-voting/election-results-data: who was printed on each party's ballot.
  primary votes     the Secretary of State's final tabulations of the June 9 primary on the same page: one workbook
                    per contest, one row per municipality (or ward and precinct) and a state total. Workbooks are read
                    for the party primaries with a field only: "U.S. Senate - Democratic", "Representative to Congress,
                    District 1 - Republican" and, for the one federal contest that went to the central ranked-choice
                    count, District 2's Democratic "RCV Central Count first choice votes" with the "Jun26 DEM CG2 RCV
                    Summary Report" (a PDF: the rounds, each candidate's votes in each round, exhausted ballots, the
                    threshold and the winner). The cast vote records beside them are not read.

Ranked choice. The Secretary's Candidate's Guide to Ballot Access (2026, "Ranked-Choice Voting (RCV)") says a primary
for U.S. Senate or U.S. Congress is decided by ranked choice if three or more candidates qualify for the ballot, or two
and there is a declared write-in candidate, and otherwise by plurality; municipalities count first choices only, and
when no candidate has more than half of them the ballots go to the Secretary's central count. So:
  - votes and pct are each candidate's first-choice votes as the tabulation gives them (for the contest that went to
    the central count, its "first choice votes" workbook); pct is of the candidates plus declared write-ins (blank
    ballots left out); a declared write-in candidate with votes is kept in the field (write_in 1) and counts in the
    total;
  - in a ranked-choice contest where the leader had more than half of the first choices, no further round was counted
    (the Secretary lists such contests with the "Non-Ranked Choice Offices"), the leader advanced, and the row note
    says so; in one that went to the central count, the winner of the RCV Summary Report's final round advanced, and
    the row note gives the rounds and the final-round result in plain words. The report's first round must account
    for every ballot in the first-choice workbook (candidates plus exhausted = ballots cast), and the difference
    between its first round and the first-choice column is explained when it is the ballots that skipped a first
    choice (blank first choices less first-round exhausted ballots);
  - a contest with two candidates and no declared write-in was decided by plurality, and the note says that.
Every workbook's municipal rows must add up to its state total and to each county subtotal it prints ("CUM Totals"),
and each row's candidates plus blanks must equal its ballots cast. The printed candidates of each workbook must be
exactly the primary list's for that office and party, less any withdrawn before the primary. The winner must be the
party's candidate on the November list, or have withdrawn after the primary with the November list's candidate named as
the replacement; that row then says so. Neither the results page nor the files use the words "official" or
"certified": the tabulations are the Secretary's own, the non-ranked ones named "FINAL" (all last saved June 29, 2026,
twenty days after the primary), and the RCV report is the central count's result.

By the same rule, the November races for U.S. Senate and District 2 (two candidates on the ballot and declared
write-ins) will be counted by ranked choice, and District 1 (two, no declared write-in) by plurality. The rows do not
say so; that belongs in a state note.

Names: the lists print names in ordinary capitals, in separate columns; they are shown first name first, and a lone
middle initial is given its full stop, as the general list prints it ("Susan M." where the primary list has "M"). The
tabulations print names in capitals ("LAST, FIRST M"); they are matched to the primary list by family name and first
given name, and a declared write-in, on no list, is shown in ordinary capitals with a note. The files carry each
candidate's town of residence (no street address, telephone or e-mail); that column is never read.
"""

import collections
import datetime as dt
import html as H
import os
import re
import time
from urllib.parse import urljoin

import openpyxl

from ballot.common import fold, house_id, party_code, record_source, senate_id
from ballot.lists.tx import proper
from ballot.pdftext import PDF, join, lines, rows as pdf_rows
from states import net

STATE = "ME"
AGENCY = "Maine Secretary of State, Bureau of Corporations, Elections and Commissions"
SITE = "https://www.maine.gov"
FILES = SITE + "/sos/sites/maine.gov.sos/files/inline-files/"
UPCOMING = SITE + "/sos/elections-voting/upcoming-elections"
GENERAL_DATE = "2026-11-03"
PRIMARY = "2026-06-09"

# the November files, found on the Upcoming Elections page by their link text (a new version gets a new address)
LINKED = {
    "general": (r"2026 General Election Candidates? Lis", "me_2026_general_candidate_list.xlsx", "PK",
                FILES + "2026%20General%20Candidate%20List%20-%20FINAL.xlsx"),
    "withdrawals": (r"2026 Candidate Withdrawals and Replacement Candidate Nominations", "me_2026_post_primary_withdrawals.pdf", "%PDF",
                    FILES + "2026%20Post%20Primary%20Withdrawals%20%26%20Replacement%20Candidates%20-%20Updated%209.22.26_0.pdf"),
    "write_ins": (r"Declared Write-In Candidates", "me_2026_general_declared_write_ins.pdf", "%PDF",
                  FILES + "Nov26%20Declared%20Write%20in%20List.pdf"),
}
# the primary's files: final, so fixed
PRIMARY_LIST = ("me_2026_primary_candidate_list.xlsx", FILES + "2026%20Primary%20Candidate%20List%20posting%20FINAL%203.16.26.xlsx")
PRIMARY_WITHDRAWALS = ("me_2026_pre_primary_withdrawals.pdf", FILES + "2026%20PrePrimary%20Withdrawal%20List%205%2022%202026.pdf")
TALLIES = {       # (race, party letter): (cache file, address, title as the results page gives it)
    ("2026-ME-S2", "D"): ("me_2026_primary_us_senate_dem.xlsx", FILES + "US%20Senate%20DEM%20-%20FINAL.xlsx",
                          "June 9, 2026 Primary Election: U.S. Senate - Democratic"),
    ("2026-ME-H01", "R"): ("me_2026_primary_cg1_rep.xlsx", FILES + "Rep%20to%20Congress%20Dis%201%20REP%20-%20FINAL.xlsx",
                           "June 9, 2026 Primary Election: Representative to Congress, District 1 - Republican"),
    ("2026-ME-H02", "D"): ("me_2026_primary_cg2_dem_first_choice.xlsx",
                           FILES + "Representative%20to%20Congress%20District%202%20-%20Democratic.xlsx",
                           "June 9, 2026 Primary Election: Representative to Congress, District 2 - Democratic, "
                           "RCV Central Count first choice votes"),
}
RCV_REPORTS = {
    ("2026-ME-H02", "D"): ("me_2026_primary_cg2_dem_rcv_summary.pdf", FILES + "CG2%20Democratic%20RCV%20Summary%20Report.pdf",
                           "Jun26 DEM CG2 RCV Summary Report (Representative to Congress, District 2 - Democratic)"),
}

PARTIES = {"D": "Democratic", "R": "Republican", "G": "Green Independent", "L": "Libertarian"}
CODES = {"D": "DEM", "R": "REP", "G": "GRN", "L": "LIB"}
NON_PARTY = {"unenrolled", "un enrolled", "undeclared", "unaffiliated", "independent", "no party"}
LIST_KEEP = ("Office", "Dist", "Party", "Last Name", "First Name", "Middle Name", "Suffix")
FEDERAL = ("US", "CG")
WRITE_IN = "Write-in candidate: the name is not printed on the ballot."
CAPS_RESULTS = "Maine's results print names in capitals; they are shown here in ordinary capitals."

# the PDFs: where each heading sits (checked, to 6 points), where each kept column's cells begin and end, and the one
# cell that marks the last heading row; anything outside the kept bands is never turned into text
LAYOUTS = {
    "withdrawals": {"heads": {"Ofc.": 86, "Dist.": 123, "Pty": 160, "Withdrawn": 189, "Date": 330, "Replacement": 395, "Residence": 536},
                    "bands": {"office": (0, 110), "district": (110, 150), "party": (150, 176), "name": (176, 305),
                              "date": (305, 376), "replacement": (376, 520)},
                    "mark": ("office", "Ofc.")},
    "primary_withdrawals": {"heads": {"Ofc.": 27, "Pty": 100, "Withdrawn": 145, "Replacement": 382, "Residence": 543},
                            "bands": {"office": (0, 50), "district": (50, 95), "party": (95, 120), "name": (120, 280),
                                      "date": (280, 350), "replacement": (350, 520)},
                            "mark": ("district", "Cty")},
    "write_ins": {"heads": {"Office": 21, "District": 46, "County": 75, "Party": 132, "Name": 224, "Residence": 307},
                  "bands": {"office": (0, 44), "district": (44, 70), "party": (97, 178), "name": (178, 270)},
                  "mark": ("office", "Office")},
}
MONTHS = ("january", "february", "march", "april", "may", "june", "july", "august", "september", "october", "november", "december")


# ------------------------------------------------------------------------------------------------------ small helpers

def clean(value):
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def person(first, middle, last, suffix=""):
    """First name first; a lone initial gets its full stop, as the general list prints it."""
    words = []
    for part in (first, middle):
        part = clean(part)
        words.append(part + "." if re.fullmatch(r"[A-Z]", part) else part)
    words += [clean(last), clean(suffix)]
    return " ".join(w for w in words if w)


def key_of(family, given):
    """Family name and first given name, letters only: how a tabulation's capitals are matched to a list."""
    g = fold(given).split()
    return fold(family).replace(" ", ""), g[0] if g else ""


def race_of(office, district, where):
    office, district = clean(office).upper(), clean(district)
    if office == "US":
        if district:
            raise SystemExit(f"Maine: a U.S. Senator row on {where} names a district ({district!r})")
        return senate_id(STATE, 2)
    if office == "CG":
        if district not in ("1", "2"):
            raise SystemExit(f"Maine: a Representative to Congress row on {where} names a district that is not read ({district!r})")
        return house_id(STATE, int(district))
    return None


def party_label(code, where):
    code = clean(code)
    if code in PARTIES:
        return PARTIES[code]
    if len(code) <= 2:
        raise SystemExit(f"Maine: a party code on {where} that is not read ({code!r})")
    return code                                              # a non-party candidate's designation, printed in full


def colour(label):
    """party_code, knowing Maine's words for a candidate of no party ("Unenrolled", "Un-enrolled", "Undeclared")."""
    return "I" if fold(label) in NON_PARTY else party_code(label)


def date_of(text):
    """'As of September 22, 2026', '8/25/26', '7/10/2026' -> ISO date, else ''."""
    m = re.search(r"\b(\d{1,2})/(\d{1,2})/(\d{2}|\d{4})\b", text)
    if m:
        year = int(m.group(3)) + (2000 if len(m.group(3)) == 2 else 0)
        return f"{year}-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
    m = re.search(r"\b([A-Z][a-z]+) (\d{1,2}), (20\d\d)\b", text)
    if m and m.group(1).lower() in MONTHS:
        return f"{m.group(3)}-{MONTHS.index(m.group(1).lower()) + 1:02d}-{int(m.group(2)):02d}"
    return ""


def published_of(title, fallback=""):
    """The date a PDF's title lines say it was made: 'As of September 22, 2026', else a date written in figures
    ('Declared Write-in Candidates 8/25/26'); never the election date the title names."""
    for t in title:
        m = re.search(r"(?i)\bas of\b(.*)$", t)
        if m and date_of(m.group(1)):
            return date_of(m.group(1))
    for t in title:
        m = re.search(r"\b\d{1,2}/\d{1,2}/(\d{2}|\d{4})\b", t)
        if m:
            return date_of(m.group(0))
    return fallback


def plain_date(iso):
    d = dt.date.fromisoformat(iso)
    return f"{MONTHS[d.month - 1].title()} {d.day}, {d.year}"


def fetch(url, path, first_bytes, max_age_days, say):
    net.download(url, path, max_age_days=max_age_days, say=say)
    if not open(path, "rb").read(8).startswith(first_bytes.encode()):
        os.replace(path, path + ".bad")
        raise SystemExit(f"Maine: {url} did not answer with the file expected (it does not begin {first_bytes}); "
                         f"what came is kept as {os.path.basename(path)}.bad")


def linked(say):
    """{kind: (address, date printed beside the link)} from the Upcoming Elections page; the known addresses if the
    page cannot be read or a link is not found there."""
    found = {k: (v[3], "") for k, v in LINKED.items()}
    try:
        page = net.get(UPCOMING, accept="text/html").decode("utf-8", "replace")
    except Exception as e:  # noqa: BLE001  the files themselves are what matter; the page only finds new versions
        say(f"      Maine: the Upcoming Elections page could not be read ({e}); using the known addresses")
        return found
    for kind, (pattern, *_rest) in LINKED.items():
        for m in re.finditer(r'<a\s[^>]*href="([^"]+)"[^>]*>(.*?)</a>', page, re.S):
            label = clean(H.unescape(re.sub(r"<[^>]+>", " ", m.group(2))))
            if re.search(pattern, label, re.I):
                after = clean(H.unescape(re.sub(r"<[^>]+>", " ", page[m.end():m.end() + 400])))
                found[kind] = (urljoin(UPCOMING, H.unescape(m.group(1))), date_of(after))
                break
        else:
            say(f"      Maine: no link for {kind} on the Upcoming Elections page; using the known address")
    return found


# ------------------------------------------------------------------------------------------------------ the lists

def read_list(path, sheet, where):
    """[(race, party code, name as printed)] for the federal rows of a candidate list workbook, in the list's order."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    if clean(ws.title) != sheet:
        raise SystemExit(f"Maine: {where} is headed {ws.title!r}, not {sheet!r}")
    it = ws.iter_rows(values_only=True)
    heads = [clean(c) for c in next(it)]
    if not all(k in heads for k in LIST_KEEP):
        raise SystemExit(f"Maine: the columns of {where} changed ({[h for h in heads if h]})")
    idx = {k: heads.index(k) for k in LIST_KEEP}
    out, total = [], 0
    for r in it:
        r = list(r) + [None] * (len(heads) - len(r))
        cell = {k: r[i] for k, i in idx.items()}
        if not clean(cell["Office"]):
            if clean(cell["Last Name"]):
                raise SystemExit(f"Maine: {where} has a candidate row with no office")
            continue
        total += 1
        race = race_of(cell["Office"], cell["Dist"], where)
        if race:
            name = person(cell["First Name"], cell["Middle Name"], cell["Last Name"], cell["Suffix"])
            out.append({"race": race, "party": clean(cell["Party"]), "name": name,
                        "key": key_of(cell["Last Name"], cell["First Name"])})
    wb.close()
    return out, total


def words(runs):
    """Runs that touch, joined into words: [(x, text)]."""
    out = []
    for x0, _y, size, t, x1 in sorted(runs, key=lambda r: r[0]):
        if out and x0 - out[-1][2] < 0.18 * size and not t.startswith(" "):
            out[-1][1] += t
            out[-1][2] = max(out[-1][2], x1)
        else:
            out.append([x0, t, x1])
    return [(x, t.strip()) for x, t, _e in out if t.strip()]


def read_pdf(path, kind):
    """(title lines, [{band: text}]) for a one-table PDF. On the first page the headings are checked where the layout
    says they sit, and the title lines above them are read whole; below them, and on every later page, only the kept
    bands are turned into text, row by row (a later page that repeats the headings has them dropped, up to the row
    that marks their end). A row with nothing in its office and date bands continues the row above."""
    lay = LAYOUTS[kind]
    band, mark = lay["mark"]
    pdf = PDF(open(path, "rb").read())
    title, out, seen = [], [], set()

    def cells_of(runs):
        return {b: join([r for r in runs if lo <= r[0] < hi]) for b, (lo, hi) in lay["bands"].items()}

    for pno, (page, res) in enumerate(pdf.pages(), start=1):
        rs = pdf_rows(pdf, page, res)
        if pno == 1:
            in_heads, start = False, None
            for i, (_y, runs) in enumerate(rs):
                ws = words(runs)
                hits = {h for h, x in lay["heads"].items() for wx, wt in ws if wt.startswith(h) and abs(wx - x) <= 6}
                if hits:
                    in_heads = True
                    seen |= hits
                elif not in_heads:
                    title.append(join(runs))                  # the title and date above the headings
                if in_heads and cells_of(runs)[band] == mark:
                    start = i + 1
                    break
            if start is None:
                raise SystemExit(f"Maine: {os.path.basename(path)} has no table heading where one is expected")
        else:
            top = [cells_of(runs) for _y, runs in rs[:8]]
            start = next((i + 1 for i, c in enumerate(top) if c[band] == mark), 0)
        for _y, runs in rs[start:]:
            cells = cells_of(runs)
            if not any(cells.values()):
                continue
            if out and not cells["office"] and not cells.get("date"):
                for b, text in cells.items():
                    if text:
                        out[-1][b] = clean(f"{out[-1][b]} {text}")
                continue
            out.append(cells)
    missing = set(lay["heads"]) - seen
    if missing:
        raise SystemExit(f"Maine: the headings of {os.path.basename(path)} moved or changed ({sorted(missing)} not where expected)")
    return title, out


def federal_rows(path, kind, where):
    """The federal rows of a withdrawal or write-in PDF, each checked: [(race, cells)]."""
    title, table = read_pdf(path, kind)
    out = []
    for cells in table:
        race = race_of(cells["office"], cells["district"], where) if cells["office"] in FEDERAL else None
        if not race:
            continue
        if not re.search(r"[A-Za-z]{2}", cells["name"]) or not cells["party"]:
            raise SystemExit(f"Maine: a federal row of {where} has no name or no party where they are expected")
        if "date" in cells and not re.fullmatch(r"\d{1,2}/\d{1,2}/(\d{2}|\d{4})", cells["date"]):
            raise SystemExit(f"Maine: a federal row of {where} has no date where one is expected ({cells['date']!r})")
        out.append((race, cells))
    return title, out, len(table)


# ------------------------------------------------------------------------------------------------------ the votes

NAME_COL = re.compile(r"[A-Z][A-Z'.\- ]*, [A-Z][A-Z'.\- ]*")


def tally(path, where):
    """{"cands": {column: printed name}, "declared": {columns}, "votes": {column: votes}, "blank", "ballots", "places"}
    from a tabulation workbook; every municipal row checked, and their sum checked against the state total."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    grid = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    wb.close()
    def word(c, w):
        return isinstance(c, str) and c.strip().upper() == w
    hi = next((i for i, r in enumerate(grid[:8]) if any(word(c, "BLANK") for c in r)), None)
    tbc = {j for r in grid[:8] for j, c in enumerate(r) if word(c, "TBC")}
    if hi is None or len(tbc) != 1:
        raise SystemExit(f"Maine: {where} has no BLANK and TBC columns where they are expected")
    head, tbc = grid[hi], tbc.pop()
    blank = next(j for j, c in enumerate(head) if word(c, "BLANK"))
    cands = {j: clean(head[j]) for j in range(blank) if isinstance(head[j], str) and NAME_COL.fullmatch(clean(head[j]).upper())}
    first = min(cands) if cands else None
    if not cands or sorted(cands) != list(range(first, blank)) or tbc < blank:
        raise SystemExit(f"Maine: the candidate columns of {where} are not where they are expected")
    declared = {j for j in cands for r in grid[hi + 1:hi + 4] if j < len(r) and isinstance(r[j], str)
                and fold(r[j]) == "declared write in"}
    cols = list(cands) + [blank, tbc]
    cty = {j for r in grid[:8] for j, c in enumerate(r) if word(c, "CTY") and j < first}
    cty = cty.pop() if len(cty) == 1 else None
    sums, total, places, done = collections.Counter(), None, 0, False
    by_county, subtotals = collections.defaultdict(collections.Counter), 0
    for n, r in enumerate(grid[hi + 1:], start=hi + 2):
        r = r + [None] * (tbc + 1 - len(r))
        nums = [r[j] for j in cols]
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in nums) or (n <= hi + 5 and not any(isinstance(c, (int, float)) for c in nums)):
            continue                                        # a blank row, or the furniture under the names (party, a declared write-in's mark)
        if done and all(isinstance(c, (int, float)) and 0 <= c <= 1 for c in nums) and any(isinstance(c, float) for c in nums):
            continue                                        # the shares printed under the total
        if not all(isinstance(c, int) and not isinstance(c, bool) for c in nums):
            raise SystemExit(f"Maine: row {n} of {where} is not read (its vote cells are not whole numbers)")
        if done:
            raise SystemExit(f"Maine: {where} has figures after its state total (row {n})")
        labels = [clean(c) for c in r[:first] if clean(c)]
        if sum(r[j] for j in cands) + r[blank] != r[tbc]:
            raise SystemExit(f"Maine: row {n} of {where}: candidates plus blanks are not the ballots cast")
        if not labels or re.fullmatch(r"(?i)state totals?", labels[-1]):
            total, done = {j: r[j] for j in cols}, True
            continue
        sub = re.fullmatch(r"(?i)([A-Z]{3}) totals?", labels[-1])
        if sub:                                             # a county's subtotal: checked against its municipal rows, not added again
            if cty is None or any(by_county[sub.group(1).upper()][j] != r[j] for j in cols):
                raise SystemExit(f"Maine: the {sub.group(1)} subtotal of {where} (row {n}) is not the sum of that county's rows")
            subtotals += 1
            continue
        places += 1
        for j in cols:
            sums[j] += r[j]
            if cty is not None and clean(r[cty]):
                by_county[clean(r[cty]).upper()][j] += r[j]
    if total is None:
        raise SystemExit(f"Maine: {where} has no state total")
    if any(sums[j] != total[j] for j in cols):
        raise SystemExit(f"Maine: the municipal rows of {where} do not add up to its state total")
    return {"cands": cands, "declared": declared, "votes": {j: total[j] for j in cands}, "blank": total[blank],
            "ballots": total[tbc], "places": places, "subtotals": subtotals}


def rcv_report(path, where):
    """{"winner", "threshold", "rounds", "cands": {printed name: [votes by round]}, "exhausted": [by round]} from an RCV
    Summary Report, checked: each round's candidates plus exhausted ballots are the same number of ballots, a
    candidate once dropped stays at nought, and the winner has the most votes in the final round, above the threshold."""
    got, cands, exhausted = {}, {}, None
    for _p, _y, t in lines(path):
        m = re.fullmatch(r"(Contest|Office|Winner\(s\)|Threshold|Rounds|Eliminated|Elected) ?(.*)", t)
        if m:
            got[m.group(1)] = m.group(2).strip()
            continue
        m = re.fullmatch(r"Exhausted Ballots((?: \d+)+)", t)
        if m:
            exhausted = [int(x) for x in m.group(1).split()]
            continue
        m = re.fullmatch(r"([A-Z][A-Za-z'.\- ]+, [A-Z][A-Za-z'.\- ]+?)((?: \d+)+)", t)
        if m:
            cands[m.group(1).strip()] = [int(x) for x in m.group(2).split()]
    rounds = len(re.findall(r"Round \d+", got.get("Rounds", "")))
    if not rounds or exhausted is None or len(exhausted) != rounds or len(cands) < 3 or any(len(v) != rounds for v in cands.values()):
        raise SystemExit(f"Maine: {where} is not laid out as an RCV Summary Report is expected to be")
    ballots = {sum(v[k] for v in cands.values()) + exhausted[k] for k in range(rounds)}
    if len(ballots) != 1:
        raise SystemExit(f"Maine: the rounds of {where} do not each account for the same ballots")
    for name, v in cands.items():
        if any(v[k] == 0 and v[k + 1] for k in range(rounds - 1)):
            raise SystemExit(f"Maine: {where} gives {name} votes after dropping them")
    final = sorted(cands, key=lambda c: -cands[c][-1])
    threshold = int(got.get("Threshold", "0") or 0)
    continuing = sum(v[-1] for v in cands.values())
    if fold(got.get("Winner(s)", "")) != fold(final[0]) or fold(got.get("Elected", "")) != fold(final[0]):
        raise SystemExit(f"Maine: the winner named in {where} does not have the most votes in its final round")
    if threshold != continuing // 2 + 1 or cands[final[0]][-1] < threshold:
        raise SystemExit(f"Maine: the threshold in {where} is not half the final round's votes plus one, or the winner is under it")
    dropped = [[c for c, v in cands.items() if v[k] and not v[k + 1]] for k in range(rounds - 1)]
    if fold(got.get("Eliminated", "")) != fold(" ".join(c for step in dropped for c in step)):
        raise SystemExit(f"Maine: the candidates {where} says were eliminated are not the ones its figures drop")
    return {"winner": final[0], "threshold": threshold, "rounds": rounds, "cands": cands, "exhausted": exhausted,
            "dropped": dropped, "contest": got.get("Office", ""), "ballots": ballots.pop()}


def source_of(race, party):
    """me-sos-2026-primary-s2-dem, me-sos-2026-primary-h01-rep."""
    return f"me-sos-2026-primary-{race.split('-')[-1].lower()}-{CODES[party].lower()}"


def turned(printed):
    """'Dunlap, Matthew G.' -> 'Matthew G. Dunlap'."""
    last, _, first = printed.partition(",")
    return clean(f"{first} {last}")


# ------------------------------------------------------------------------------------------------------ load

def load(con, cache, say=print):
    net.patient_lookups()
    folder = os.path.join(cache, "me")
    os.makedirs(folder, exist_ok=True)
    races = [r for (r,) in con.execute("SELECT race_id FROM races WHERE state = ? ORDER BY race_id", (STATE,))]

    links = linked(say)
    paths = {}
    for kind, (_pattern, name, magic, _default) in LINKED.items():
        paths[kind] = os.path.join(folder, name)
        fetch(links[kind][0], paths[kind], magic, 2, say)
        time.sleep(1)
    plist, pwd = os.path.join(folder, PRIMARY_LIST[0]), os.path.join(folder, PRIMARY_WITHDRAWALS[0])
    fetch(PRIMARY_LIST[1], plist, "PK", 30, say)
    fetch(PRIMARY_WITHDRAWALS[1], pwd, "%PDF", 30, say)
    for (name, url, _t) in list(TALLIES.values()) + list(RCV_REPORTS.values()):
        fetch(url, os.path.join(folder, name), "%PDF" if name.endswith(".pdf") else "PK", 30, say)

    # ---- the November ballot
    general, general_rows = read_list(paths["general"], "November 3, 2026", "the 2026 General Election Candidates List")
    wd_title, withdrawals, wd_rows = federal_rows(paths["withdrawals"], "withdrawals", "the list of withdrawals after the primary")
    wi_title, write_ins, wi_rows = federal_rows(paths["write_ins"], "write_ins", "the list of declared write-in candidates")
    on_list = {(g["race"], fold(g["name"])) for g in general}
    left_off, replaced = [], {}                       # replaced: (race, party letter) -> {withdrawn, date, replacement}
    for race, c in withdrawals:
        party, gone, repl = c["party"], c["name"], c.get("replacement", "")
        party_label(party, "the list of withdrawals after the primary")
        if re.search(r"(?i)\b(not|no|none|permitted|nominee|declined|vacancy)\b", repl):
            repl = ""                                          # the column's own remark ("Not permitted - ..."), not a name
        match = [g for g in general if g["race"] == race and g["party"] == party and fold(g["name"]) == fold(repl)]
        if repl and not match:
            raise SystemExit(f"Maine: the replacement named for {gone} ({race}) is not the general list's {PARTIES.get(party, party)} "
                             f"candidate ({repl!r})")
        replaced[(race, party)] = {"withdrawn": gone, "date": date_of(c["date"]), "replacement": match[0]["name"] if match else None}
        if (race, fold(gone)) in on_list:
            left_off.append((race, fold(gone)))
    rows, order = [], collections.Counter()
    for g in general:
        if (g["race"], fold(g["name"])) in left_off:
            continue
        label = party_label(g["party"], "the 2026 General Election Candidates List")
        order[g["race"]] += 1
        rows.append((g["race"], "general", GENERAL_DATE, g["name"], label, colour(label), order[g["race"]], 0, 0,
                     None, None, None, None, None, "me-sos-2026-general-list", None))
    for race, c in write_ins:
        rows.append((race, "general", GENERAL_DATE, c["name"], c["party"], colour(c["party"]), None, 0, 1,
                     None, None, None, None, None, "me-sos-2026-general-write-ins", WRITE_IN))
    missing = sorted(set(races) - {r[0] for r in rows})
    if missing:
        raise SystemExit(f"Maine: no November candidate was read for {missing}")
    nominee = {(g["race"], g["party"]): g for g in general if g["party"] in PARTIES and (g["race"], fold(g["name"])) not in left_off}

    # ---- the primary: who was printed on each party's ballot
    plisted, _ptotal = read_list(plist, "June 9, 2026", "the 2026 Primary Candidate List")
    pw_title, pwithdrawn, pw_rows = federal_rows(pwd, "primary_withdrawals", "the list of withdrawals before the primary")
    gone_before = {(race, c["party"], fold(c["name"])) for race, c in pwithdrawn}
    ballot = collections.defaultdict(list)
    for p in plisted:
        party_label(p["party"], "the 2026 Primary Candidate List")
        if (p["race"], p["party"], fold(p["name"])) not in gone_before:
            ballot[(p["race"], p["party"])].append(p)
    fields = {k for k, v in ballot.items() if len(v) >= 2}
    if fields != set(TALLIES):
        raise SystemExit(f"Maine: the party primaries with a field are {sorted(fields)}; the loader reads tabulations for {sorted(TALLIES)}")

    tallies_used, rcv_used, primary_rows = {}, {}, 0
    for (race, party), (fname, _url, title) in sorted(TALLIES.items()):
        t = tally(os.path.join(folder, fname), title)
        printed = {j: n for j, n in t["cands"].items() if j not in t["declared"]}
        listed = {p["key"]: p for p in ballot[(race, party)]}
        read_keys = {}
        for j, n in printed.items():
            last, _, given = n.partition(",")
            read_keys[j] = key_of(last, given)
        if sorted(read_keys.values()) != sorted(listed):
            raise SystemExit(f"Maine: the candidates of {title} are not the primary list's for that office and party")
        total = sum(t["votes"].values())
        label, code = PARTIES[party], CODES[party]
        rcv = len(printed) >= 3 or (len(printed) == 2 and t["declared"])
        ranked = sorted(t["cands"], key=lambda j: -t["votes"][j])
        if t["votes"][ranked[0]] == t["votes"][ranked[1]]:
            raise SystemExit(f"Maine: {title} is a tie on first choices; read the Secretary's own result")
        shown = {j: listed[read_keys[j]]["name"] if j in read_keys else proper(turned(n)) for j, n in t["cands"].items()}
        report = RCV_REPORTS.get((race, party))
        majority = 2 * t["votes"][ranked[0]] > total
        if not rcv:
            if report:
                raise SystemExit(f"Maine: {title} has an RCV report but was not a ranked-choice contest")
            winner = ranked[0]
            note = ("Decided by plurality: with two candidates on the ballot and no declared write-in, Maine does not count a "
                    "primary by ranked choice.")
        elif majority:
            if report:
                raise SystemExit(f"Maine: {title} has an RCV report although its leader had a majority of first choices")
            winner = ranked[0]
            note = (f"Counted by ranked choice. {shown[winner]} had more than half of the first choices "
                    f"({t['votes'][winner]:,} of {total:,}), so no further round was counted.")
        else:
            if not report:
                raise SystemExit(f"Maine: no candidate in {title} had a majority of first choices, and no RCV report is configured")
            rpath = os.path.join(folder, report[0])
            rep = rcv_report(rpath, report[2])
            by_key = {}
            for printed_name in rep["cands"]:
                last, _, given = printed_name.partition(",")
                by_key[key_of(last, given)] = printed_name
            col_of = {read_keys[j]: j for j in read_keys}
            if set(by_key) != set(col_of) or rep["ballots"] != t["ballots"]:
                raise SystemExit(f"Maine: {report[2]} does not count the same candidates and ballots as {title}")
            rep_col = {by_key[k]: col_of[k] for k in by_key}
            winner = rep_col[rep["winner"]]
            steps = []
            for k, out in enumerate(rep["dropped"], start=1):
                if out:
                    steps.append(f"{' and '.join(shown[rep_col[c]] for c in out)} after round {k}")
            last2 = sorted(rep["cands"], key=lambda c: -rep["cands"][c][-1])
            final = [c for c in last2 if rep["cands"][c][-1]]
            if len(final) != 2:
                raise SystemExit(f"Maine: the final round of {report[2]} does not have two candidates")
            first_round = {rep_col[c]: v[0] for c, v in rep["cands"].items()}
            more = sum(first_round.values()) - sum(t["votes"][j] for j in printed)
            skipped = t["blank"] - rep["exhausted"][0]
            note = (f"Counted by ranked choice. The votes shown are first choices; no candidate had more than half of them, so "
                    f"the Secretary of State's central count dropped the last-placed candidate in each round: {', then '.join(steps)}. "
                    f"In the final round {shown[rep_col[final[0]]]} had {rep['cands'][final[0]][-1]:,} votes and "
                    f"{shown[rep_col[final[1]]]} {rep['cands'][final[1]][-1]:,} ({rep['exhausted'][-1]:,} ballots counted for "
                    f"neither), so {shown[winner]} won the nomination.")
            if more and more == skipped and all(first_round[j] >= t["votes"][j] for j in printed):
                note += (f" The count's first round credits {more:,} more votes than the first choices shown: ballots that "
                         f"skipped the first choice were counted for their next one.")
            elif more:
                note += f" The count's first round credits {more:,} more votes than the first choices shown."
            rcv_used[(race, party)] = (rpath, rep)
        nom = nominee.get((race, party))
        repl = replaced.get((race, party))
        after = None
        if nom and fold(nom["name"]) == fold(shown[winner]):
            pass
        elif repl and fold(repl["withdrawn"]) == fold(shown[winner]) and (repl["replacement"] or not nom):
            after = (f"Won the nomination, then withdrew on {plain_date(repl['date'])}"
                     + (f"; {repl['replacement']} was nominated as the replacement candidate." if repl["replacement"] else
                        "; no replacement is on the November list."))
            repl["winner"] = True
        else:
            raise SystemExit(f"Maine: the winner of {title} ({shown[winner]}) is not the party's candidate on the November list")
        for j in ranked:
            wi = j in t["declared"]
            parts = [WRITE_IN if wi else "", CAPS_RESULTS if j not in read_keys else "", note, after if j == winner and after else ""]
            rows.append((race, f"primary-{code}", PRIMARY, shown[j], label, colour(label), None, 0, int(wi), t["votes"][j],
                         round(100 * t["votes"][j] / total, 1) if total else None, "advanced" if j == winner else "lost",
                         None, None, source_of(race, party),
                         " ".join(p for p in parts if p) or None))
            primary_rows += 1
        t["title"], t["path"] = title, os.path.join(folder, fname)
        tallies_used[(race, party)] = t

    # a replacement candidate on the November list says whom they replaced
    for i, r in enumerate(rows):
        repl = replaced.get((r[0], next((k for k, v in PARTIES.items() if v == r[4]), None))) if r[1] == "general" else None
        if repl and repl["replacement"] == r[3]:
            won = ", who won the June 9 primary," if repl.get("winner") else ""
            rows[i] = r[:15] + (f"Replacement candidate: nominated after {repl['withdrawn']}{won} withdrew on {plain_date(repl['date'])}.",)
    general_out = [r for r in rows if r[1] == "general"]
    with con:
        con.execute("DELETE FROM candidates WHERE race_id LIKE '2026-ME-%'")
        con.execute("DELETE FROM list_gaps WHERE state = 'ME'")
        con.executemany("INSERT OR REPLACE INTO candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", rows)
        record_source(con, "me-sos-2026-general-list", path=paths["general"], level="federal", state=STATE,
                      kind="official candidate list", agency=AGENCY,
                      title="2026 General Election Candidates List (November 3, 2026): U.S. Senator and Representative to Congress",
                      url=links["general"][0], published=links["general"][1], rows=len(general),
                      note=f"Every office's {general_rows} rows read; only Office, Dist, Party and the name columns taken, the town of "
                           f"residence never read. No ballot order is given; the list's order (alphabetical by surname) is kept. "
                           f"Withdrawn and still listed, left off: {len(left_off)}.")
        gone_words = "; ".join(f"{r['withdrawn']} ({'U.S. Senator' if race.endswith('S2') else 'Representative to Congress'}, "
                               f"{PARTIES.get(p, p)}), withdrew {plain_date(r['date'])}"
                               + (f", replaced by {r['replacement']}" if r["replacement"] else "")
                               for (race, p), r in sorted(replaced.items()))
        record_source(con, "me-sos-2026-general-withdrawals", path=paths["withdrawals"], level="federal", state=STATE,
                      kind="official candidate list", agency=AGENCY,
                      title="Candidate Withdrawals and Replacement Candidate Nominations after the June 9, 2026 Primary",
                      url=links["withdrawals"][0], published=published_of(wd_title, links["withdrawals"][1]),
                      rows=len(withdrawals),
                      note=f"{wd_rows} rows, every office; federal: {gone_words or 'none'}. The general list had already taken off "
                           f"the candidates who withdrew; {len(left_off)} still listed and left off here. Only the office, district, "
                           f"party, name, date and replacement columns are read.")
        record_source(con, "me-sos-2026-general-write-ins", path=paths["write_ins"], level="federal", state=STATE,
                      kind="official candidate list", agency=AGENCY,
                      title="Declared Write-in Candidates for the November 3, 2026 General Election",
                      url=links["write_ins"][0], published=published_of(wi_title, links["write_ins"][1]),
                      rows=len(write_ins),
                      note=f"{wi_rows} rows, every office; {len(write_ins)} federal, added as declared write-in candidates. Only the "
                           f"office, district, party and name columns are read.")
        record_source(con, "me-sos-2026-primary-list", path=plist, level="federal", state=STATE, kind="official candidate list",
                      agency=AGENCY, title="2026 Primary Candidate List (June 9, 2026)", url=PRIMARY_LIST[1], published="2026-03-16",
                      rows=len(plisted),
                      note="Used to check the tabulations: who was printed on each party's primary ballot, and the names in ordinary "
                           "capitals. Only Office, Dist, Party and the name columns taken.")
        record_source(con, "me-sos-2026-primary-withdrawals", path=pwd, level="federal", state=STATE, kind="official candidate list",
                      agency=AGENCY, title="Candidate Withdrawals and Replacement Candidate Nominations for the June 9, 2026 State Primary Election",
                      url=PRIMARY_WITHDRAWALS[1], published=published_of(pw_title), rows=len(pwithdrawn),
                      note=f"{pw_rows} rows, every office; federal withdrawals before the primary: "
                           f"{'; '.join(c['name'] for _r, c in pwithdrawn) or 'none'}.")
        for (race, party), t in sorted(tallies_used.items()):
            declared = [t["cands"][j] for j in t["declared"]]
            what = ("The Secretary of State's ranked-choice central count, first choices" if (race, party) in RCV_REPORTS
                    else "The Secretary of State's final tabulation")
            record_source(con, source_of(race, party), path=t["path"], level="federal", state=STATE,
                          kind="official results", agency=AGENCY, title=t["title"], url=TALLIES[(race, party)][1],
                          rows=len(t["cands"]),
                          note=f"{what}: {t['places']} rows (municipalities or their wards and precincts, and the UOCAVA ballots) "
                               f"add up to the state total"
                               + (f" and to each of {t['subtotals']} county subtotals" if t["subtotals"] else "")
                               + f", and each row's candidates plus blanks to its ballots cast ({t['ballots']:,} ballots, "
                               f"{t['blank']:,} blank). Shares are of the candidates"
                               + (f" and the declared write-in ({', '.join(proper(turned(d)) for d in declared)})" if declared else "")
                               + "; blank ballots left out.")
        for (race, party), (rpath, rep) in sorted(rcv_used.items()):
            record_source(con, source_of(race, party) + "-rcv", path=rpath, level="federal",
                          state=STATE, kind="official results", agency=AGENCY, title=RCV_REPORTS[(race, party)][2],
                          url=RCV_REPORTS[(race, party)][1], rows=len(rep["cands"]),
                          note=f"{rep['rounds']} rounds; threshold {rep['threshold']:,}; winner {turned(rep['winner'])}. Each round's "
                               f"candidates plus exhausted ballots equal the {rep['ballots']:,} ballots cast in the first-choice "
                               "tabulation. Gives the outcome and the row note; the votes stored are the first choices.")
    rcv_count = len(rcv_used)
    say(f"    Maine: 2 House districts and the Senate race, {len(general_out)} candidates on the November ballot "
        f"({len(write_ins)} declared write-ins, {len(replaced)} primary winner withdrawn and replaced, {len(left_off)} left off); "
        f"{len(TALLIES)} party primaries with a field ({primary_rows} rows), votes from the Secretary of State's final tabulations, "
        f"{rcv_count} decided in the ranked-choice central count")
    return len(general_out)
