"""
Colorado: the Secretary of State's own candidate lists and its certified results of the June 30, 2026 primary.
Colorado has eight House seats (the 2024 lines) and the class 2 Senate seat.

  November ballot   the "2026 General Election Official Candidate List" (coloradosos.gov/pubs/elections/vote/
                    generalCandidates.html), "official and reflects what was certified to the counties on September
                    4". The web page's table is read, because a withdrawn candidate is shown there struck through
                    (text-decoration: line-through) and the page is the one kept up to date; its "Excel version
                    (XLSX)" is fetched too and must list the same federal candidates (it keeps struck names without
                    marking them, and on 2026-09-30 it wrote one name shorter than the page: "Samir Witta" for the
                    page's "Samir Ezzeldin Witta"; the page's spelling is kept and any such difference is recorded).
                    Columns are taken by name: Candidate name, Office, District, Party, Write in?. Nothing else is
                    printed in either file.
  ballot order      the "2026 General Election Ballot Order for Major & Minor Party Candidates" workbook, linked from
                    vote/generalBallotOrder.html: the lot drawing of July 28, 2026, one position per party per office.
                    Colorado law (1-5-404, C.R.S.) orders a race in three tiers, major parties, then minor parties,
                    then unaffiliated candidates, each tier by lot; the drawing's numbers already put the major
                    parties first. A party that drew a position but has no candidate on the certified list is not
                    printed, so the positions stored are the places on the printed ballot (gaps closed up); a race's
                    one unaffiliated candidate follows the parties. (Should a race ever have two unaffiliated
                    candidates, their drawing is not in this file and they are given no position.) Certified write-in
                    candidates (Write in? = Y) are on the list but not printed on the ballot: write_in 1, no position.
  primary fields    the Secretary's "2026 State Primary Election Statewide Abstract of Votes Cast", certified by the
                    Secretary on July 24, 2026 (Results/Archives.html). It is a scan of the signed and sealed pages
                    (JBIG2 images, no text layer), so no script here can read it. The votes are therefore
                    taken from the Secretary's election night reporting site, results.enr.clarityelections.com/CO
                    (election 126592), whose county-by-county detail file (reports/detailxml.zip, last written
                    7/23/2026) carries every contest; that site is headed "Unofficial Results", so its figures are
                    stored only because each federal party primary's figures, candidate by candidate, equal the
                    certified abstract's Total row, read by eye on 2026-09-30 (the scan's images decoded to pictures
                    with a one-off JBIG2 decoder outside the kit; pages 2-5 and 22-25) and written below as
                    ABSTRACT (with the scan's SHA-256). Each candidate's county figures must also add up to the
                    candidate's total. If the Secretary replaces the scan, or the site's figures ever differ, the
                    fields are stored without votes (who advanced read from the November list) and the loader says so.
  write-ins         the "2026 State Primary Election Certified Write-In Results" workbook (Office, Candidate, Party,
                    County, Votes, with a Total Votes line): Colorado counts write-in votes only for certified
                    write-in candidates, and the abstract lists them beside the printed names. A certified write-in
                    counts in a field's total and is shown with write_in 1; the county lines must add up to its
                    Total Votes line and to the abstract.
  primary list      the "2026 Official Primary Election Candidate List" (vote/primaryCandidates.html, certified to the
                    counties May 1), read the same way: every printed name on the results must be a name on the list
                    that is not struck through, and the other way round.

A field is a party primary with two or more names printed on that party's ballot for the race; a certified write-in
alone beside one printed name does not make a field (in 2026: District 4's Democratic primary, Eileen Laubacher
72,483 against the write-in Jenna Preston's 566). pct is the share of that party primary's votes (candidates and
certified write-ins). The candidate of that party on the November list advanced and must be the one with the most
votes. Mel Tewahade withdrew from the District 6 Republican primary on June 24; the party named Jason Clark in his
place, and by 1-4-1004(4)(b), C.R.S. the votes cast for Tewahade are counted for Clark, who then stood alone, so that
primary is no field. The minor parties nominated no one for Congress by primary (the Libertarian and Unity primaries
were for state offices).

Parties are kept as the lists print them: Democratic Party, Republican Party, Libertarian Party, Unity Party,
American Constitution Party, Approval Voting Party, Forward Party, and Unaffiliated. Names are printed first name
first in ordinary capitals and are kept as printed, nicknames in quotation marks included. Only federal rows are
kept, as JSON in ballot_cache/co/ with the SHA-256 of each file they came from; the abstract scan is kept whole. None
of these files carries an address, telephone number or e-mail.
"""

import datetime as dt
import hashlib
import html as H
import io
import json
import os
import re
import time
import zipfile
import xml.etree.ElementTree as ET
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin

import openpyxl

from ballot.common import house_id, party_code, record_source, senate_id
from states import net

SOS = "https://www.coloradosos.gov/pubs/elections/"
PAGES = {"general": SOS + "vote/generalCandidates.html", "primary": SOS + "vote/primaryCandidates.html",
         "order": SOS + "vote/generalBallotOrder.html", "archive": SOS + "Results/Archives.html"}
TITLES = {"general": "2026 General Election Official Candidate List", "primary": "2026 Official Primary Election Candidate List",
          "order": "2026 General Election Ballot Order for Major & Minor Party Candidates"}
ABSTRACT_LABEL = "2026 State Primary Election Statewide Abstract of Votes Cast"
WRITE_IN_LABEL = "2026 State Primary Election Certified Write-In Results"
ENR = "https://results.enr.clarityelections.com/CO/126592/"
PRIMARY = "2026-06-30"
HEADS = ("candidate name", "office", "district", "party", "write in?")
CODES = {"Democratic Party": "DEM", "Republican Party": "REP", "Libertarian Party": "LIB", "Unity Party": "UNI",
         "American Constitution Party": "ACN", "Approval Voting Party": "APV", "Forward Party": "FWD", "Center Party": "CCP"}
WRITE_IN = "Write-in candidate: the name is not printed on the ballot."

# The certified abstract (a scan; see the note above), read by eye on 2026-09-30: each federal party primary's Total
# row, candidate by candidate, certified write-ins included. Only a control: the stored votes come from the site's data.
ABSTRACT_SHA = "807e10067804ca143853d15d13c8a0807dde64aed918a645ffd9428f035ff7aa"
ABSTRACT = {
    ("2026-CO-S2", "DEM"): {"Julie Gonzales": 422488, "John Hickenlooper": 472870},
    ("2026-CO-H01", "DEM"): {"Melat Kiros": 83855, "Diana DeGette": 62715, "Wanda James": 11075},
    ("2026-CO-H02", "DEM"): {"Joe Neguse": 136146},
    ("2026-CO-H03", "DEM"): {"Alex Kelloff": 39394, "Dwayne L. Romero": 47848},
    ("2026-CO-H04", "DEM"): {"Eileen Laubacher": 72483, "Jenna Preston": 566},
    ("2026-CO-H05", "DEM"): {"Jessica Killin": 45511, "Joe Reagan": 28864},
    ("2026-CO-H06", "DEM"): {"Jason Crow": 94207},
    ("2026-CO-H07", "DEM"): {"Brittany Pettersen": 113707},
    ("2026-CO-H08", "DEM"): {"Manny Rutinel": 50435, "Evan Munsing": 3717, "Shannon Bird": 26122},
    ("2026-CO-S2", "REP"): {"Mark Baisley": 445352},
    ("2026-CO-H01", "REP"): {"Christy Peterson": 15010},
    ("2026-CO-H02", "REP"): {"Kelley Anne Dennison": 21584, "Christina Blunt (Ducommun)": 15666},
    ("2026-CO-H03", "REP"): {"Ron Hanks": 31715, "Jeff Hurd": 62536},
    ("2026-CO-H04", "REP"): {"Lauren Boebert": 92651},
    ("2026-CO-H05", "REP"): {"Jeff Crank": 76547},
    ("2026-CO-H06", "REP"): {"Jason Clark": 38437},
    ("2026-CO-H07", "REP"): {"Tim Bennett": 58806},
    ("2026-CO-H08", "REP"): {"Gabe Evans": 48012},
}


def text(cell):
    return re.sub(r"\s+", " ", H.unescape(re.sub(r"<[^>]+>", " ", cell)).replace("\xa0", " ")).strip()


def links(page, base):
    """{label: absolute address} for every link on a page."""
    out = {}
    for m in re.finditer(r'<a[^>]+href="([^"]+)"[^>]*>(.*?)</a>', page, re.S | re.I):
        label = text(m.group(2))
        if label and label not in out:
            out[label] = urljoin(base, H.unescape(m.group(1)))
    return out


def updated(page):
    """The date in a page's "Last updated ..." line, as YYYY-MM-DD."""
    m = re.search(r"Last updated(?: on)? ([A-Z][a-z]+ \d{1,2}, \d{4})", text(page))
    return dt.datetime.strptime(m.group(1), "%B %d, %Y").date().isoformat() if m else ""


def race_of(office, district):
    office, district = (office or "").strip(), str(district or "").strip()
    if office == "US Senate":
        return senate_id("CO", 2)
    if office in ("US House of Representatives", "US House"):
        if not district.isdigit():
            raise SystemExit(f"Colorado: a U.S. House row names a district that is not read ({district!r})")
        return house_id("CO", int(district))
    return None


def fetch(url, magic, say, accept="*/*"):
    """A file's bytes, checked by its first bytes (a missing file can come back as a web page)."""
    data = net.get(url, accept=accept)
    time.sleep(1.5)
    if not data.startswith(magic):
        raise SystemExit(f"Colorado: {url} did not give the expected file (it starts {data[:12]!r})")
    return data


def table_rows(page):
    """The federal rows of a candidate list page's table: the named columns, and whether the row is struck through."""
    tables = re.findall(r"<table.*?</table>", page, re.S | re.I)
    if len(tables) != 1:
        raise SystemExit(f"Colorado: the candidate list page has {len(tables)} tables, not one")
    heads = [text(h).lower() for h in re.findall(r"<th[^>]*>(.*?)</th>", tables[0], re.S | re.I)]
    if not all(h in heads for h in HEADS):
        raise SystemExit(f"Colorado: the candidate list's columns changed ({heads})")
    idx = {h: heads.index(h) for h in HEADS}
    out = []
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", tables[0], re.S | re.I):
        cells = re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S | re.I)
        if not cells:
            continue
        if len(cells) != len(heads):
            raise SystemExit("Colorado: a row of the candidate list does not line up with its headings")
        rec = {h: text(cells[i]) for h, i in idx.items()}
        race = race_of(rec["office"], rec["district"])
        if race:
            out.append({"race": race, "name": rec["candidate name"], "party": rec["party"], "write_in": rec["write in?"].upper() == "Y",
                        "struck": "line-through" in tr.lower()})
    return out


def workbook_rows(data):
    """The federal rows of a candidate list workbook: (race, name, party, write in) by the named columns."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    it = wb.worksheets[0].iter_rows(values_only=True)
    heads = [str(c or "").strip().lower() for c in next(it)]
    if not all(h in heads for h in HEADS):
        raise SystemExit(f"Colorado: the candidate list workbook's columns changed ({heads})")
    idx = {h: heads.index(h) for h in HEADS}
    out = []
    for r in it:
        rec = {h: (str(r[i]).strip() if i < len(r) and r[i] is not None else "") for h, i in idx.items()}
        race = race_of(rec["office"], rec["district"])
        if race:
            out.append((race, rec["candidate name"], rec["party"], rec["write in?"].upper() == "Y"))
    wb.close()
    return out


def read_list(kind, folder, say, max_age_days):
    """A candidate list's federal rows (page and workbook compared), as JSON in the cache."""
    path = os.path.join(folder, f"co_{kind}_list_2026_federal.json")
    if os.path.exists(path) and time.time() - os.path.getmtime(path) < max_age_days * 86400:
        return json.load(open(path, encoding="utf-8"))
    url = PAGES[kind]
    page = net.get(url, accept="text/html").decode("utf-8", "replace")
    time.sleep(1.5)
    if TITLES[kind] not in text(page):
        raise SystemExit(f"Colorado: {url} is no longer the {TITLES[kind]}")
    rows = table_rows(page)
    book_url = links(page, url).get("Excel version (XLSX)")
    differ = []
    book_sha = ""
    if book_url:
        data = fetch(book_url, b"PK", say)
        book_sha = hashlib.sha256(data).hexdigest()
        book = workbook_rows(data)
        mine = [(r["race"], r["name"], r["party"], r["write_in"]) for r in rows]
        for r in mine:
            if r not in book:
                differ.append(f"page {r[1]!r} ({r[0]}) not in the workbook as written")
        for r in book:
            if r not in mine:
                differ.append(f"workbook {r[1]!r} ({r[0]}) not on the page as written")
        loose = lambda rs: sorted((r[0], r[2], r[3]) for r in rs)       # the same seats, parties and write-in marks
        if loose(mine) != loose(book):
            raise SystemExit(f"Colorado: the {TITLES[kind]} page and its workbook list different candidates: {'; '.join(differ)}")
    else:
        differ.append("the page no longer links its Excel version")
    kept = {"url": url, "workbook": book_url, "page_sha256": hashlib.sha256(page.encode("utf-8")).hexdigest(),
            "workbook_sha256": book_sha, "updated": updated(page), "read": dt.date.today().isoformat(), "differ": differ, "rows": rows}
    os.makedirs(folder, exist_ok=True)
    json.dump(kept, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    say(f"      {TITLES[kind]}: {len(rows)} federal rows ({sum(r['struck'] for r in rows)} struck through)")
    return kept


def read_order(folder, say):
    """{race: {party: drawn position}} from the ballot order workbook, as JSON in the cache."""
    path = os.path.join(folder, "co_2026_general_ballot_order_federal.json")
    if os.path.exists(path) and time.time() - os.path.getmtime(path) < 30 * 86400:
        return json.load(open(path, encoding="utf-8"))
    url = PAGES["order"]
    page = net.get(url, accept="text/html").decode("utf-8", "replace")
    time.sleep(1.5)
    if TITLES["order"] not in text(page):
        raise SystemExit(f"Colorado: {url} is no longer the {TITLES['order']}")
    book_url = links(page, url).get("Excel version (XLSX)")
    if not book_url:
        raise SystemExit(f"Colorado: {url} no longer links its Excel version")
    data = fetch(book_url, b"PK", say)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    it = wb.worksheets[0].iter_rows(values_only=True)
    heads = [str(c or "").strip() for c in next(it)]
    if heads[:2] != ["Office", "District"]:
        raise SystemExit(f"Colorado: the ballot order workbook's columns changed ({heads})")
    parties = [(i, h) for i, h in enumerate(heads) if i >= 2 and h]
    order = {}
    for r in it:
        race = race_of(str(r[0] or "").strip(), r[1])
        if race:
            order[race] = {h: int(r[i]) for i, h in parties if i < len(r) and r[i] not in (None, "")}
    wb.close()
    kept = {"url": url, "workbook": book_url, "workbook_sha256": hashlib.sha256(data).hexdigest(), "updated": updated(page),
            "read": dt.date.today().isoformat(), "order": order}
    json.dump(kept, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return kept


def contest_of(title):
    """(race, party label) from an ENR contest title, or None for any other office."""
    m = re.fullmatch(r"United States Senator - (.+)", title)
    if m:
        return senate_id("CO", 2), m.group(1)
    m = re.fullmatch(r"Representative to the 120th United States Congress - District (\d+) - (.+)", title)
    if m:
        return house_id("CO", int(m.group(1))), m.group(2)
    return None


def read_enr(folder, say):
    """The federal contests of the results site's county detail file, as JSON in the cache; None if it cannot be had."""
    path = os.path.join(folder, "co_2026_primary_enr_federal.json")
    if os.path.exists(path) and time.time() - os.path.getmtime(path) < 30 * 86400:
        return json.load(open(path, encoding="utf-8"))
    try:
        ver = net.get(ENR + "current_ver.txt", accept="text/plain").decode("ascii").strip()
        time.sleep(1.5)
        if not ver.isdigit():
            raise SystemExit(f"Colorado: the results site's version is not read ({ver[:40]!r})")
        url = ENR + ver + "/reports/detailxml.zip"
        data = fetch(url, b"PK", say)
    except (HTTPError, URLError, OSError) as e:
        say(f"      the results site could not be read ({e})")
        return json.load(open(path, encoding="utf-8")) if os.path.exists(path) else None
    root = ET.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("detail.xml"))
    if root.findtext("ElectionName") != "2026 Primary" or root.findtext("ElectionDate") != "6/30/2026":
        raise SystemExit(f"Colorado: the results site holds {root.findtext('ElectionName')!r} of {root.findtext('ElectionDate')}, not the 2026 primary")
    contests = []
    for c in root.iter("Contest"):
        found = contest_of(c.get("text", ""))
        if not found:
            continue
        choices = []
        for ch in c.findall("Choice"):
            vt = ch.findall("VoteType")
            if len(vt) != 1:
                raise SystemExit(f"Colorado: {c.get('text')} splits its votes by type; the loader reads one Total Votes line")
            choices.append({"name": ch.get("text"), "total": int(ch.get("totalVotes")),
                            "counties": {k.get("name"): int(k.get("votes")) for k in vt[0].findall("County")}})
        contests.append({"title": c.get("text"), "race": found[0], "party": found[1], "counties": int(c.get("countiesParticipating")),
                         "reported": int(c.get("countiesReported")), "choices": choices})
    kept = {"url": url, "version": ver, "sha256": hashlib.sha256(data).hexdigest(), "written": root.findtext("Timestamp"),
            "read": dt.date.today().isoformat(), "contests": contests}
    json.dump(kept, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    say(f"      results site, version {ver} ({kept['written']}): {len(contests)} federal party primaries")
    return kept


def read_writeins(folder, found, say):
    """{(race, name): {"party", "total", "counties"}} for the certified write-in candidates for Congress, as JSON."""
    path = os.path.join(folder, "co_2026_primary_writeins_federal.json")
    if os.path.exists(path) and time.time() - os.path.getmtime(path) < 30 * 86400:
        return json.load(open(path, encoding="utf-8"))
    url = found.get(f"{WRITE_IN_LABEL} (XLSX)")
    if not url:
        raise SystemExit(f"Colorado: the results archive no longer links \"{WRITE_IN_LABEL}\"")
    data = fetch(url, b"PK", say)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    it = wb.worksheets[0].iter_rows(values_only=True)
    heads = [str(c or "").strip() for c in next(it)]
    want = ("Office", "Candidate", "Party", "County", "Votes")
    if not all(h in heads for h in want):
        raise SystemExit(f"Colorado: the write-in results workbook's columns changed ({heads})")
    idx = {h: heads.index(h) for h in want}
    out = {}
    for r in it:
        office = str(r[idx["Office"]] or "").strip()
        m = re.fullmatch(r"US House of Representatives - District (\d+)", office)
        race = house_id("CO", int(m.group(1))) if m else (senate_id("CO", 2) if office == "US Senate" else None)
        if not race:
            continue
        name, county = str(r[idx["Candidate"]]).strip(), str(r[idx["County"]] or "").strip()
        rec = out.setdefault(f"{race}|{name}", {"race": race, "name": name, "party": str(r[idx["Party"]]).strip(), "total": None, "counties": {}})
        if county == "Total Votes":
            rec["total"] = int(r[idx["Votes"]])
        elif county != "Votes %":
            rec["counties"][county] = int(r[idx["Votes"]] or 0)
    wb.close()
    for rec in out.values():
        if rec["total"] is None or sum(rec["counties"].values()) != rec["total"]:
            raise SystemExit(f"Colorado: {rec['name']}'s county write-in votes do not add up to the Total Votes line")
    kept = {"url": url, "sha256": hashlib.sha256(data).hexdigest(), "read": dt.date.today().isoformat(), "candidates": out}
    json.dump(kept, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return kept


def load(con, cache, say=print):
    net.patient_lookups()
    folder = os.path.join(cache, "co")
    os.makedirs(folder, exist_ok=True)
    gen = read_list("general", folder, say, 2)
    pri = read_list("primary", folder, say, 30)
    order = read_order(folder, say)
    archive = net.get(PAGES["archive"], accept="text/html").decode("utf-8", "replace")
    time.sleep(1.5)
    found = links(archive, PAGES["archive"])
    abstract_url = found.get(f"{ABSTRACT_LABEL} (PDF)")
    if not abstract_url:
        raise SystemExit(f"Colorado: the results archive no longer links \"{ABSTRACT_LABEL}\"")
    abstract = os.path.join(folder, "co_2026_primary_abstract.pdf")
    net.download(abstract_url, abstract, max_age_days=30)
    if not open(abstract, "rb").read(5).startswith(b"%PDF"):
        raise SystemExit(f"Colorado: {abstract_url} did not give a PDF")
    abstract_sha = hashlib.sha256(open(abstract, "rb").read()).hexdigest()
    wins = read_writeins(folder, found, say)
    enr = read_enr(folder, say)

    # the November ballot
    rows, off, write_ins, unplaced, empty = [], [], [], [], []
    nominee = {}
    by_race = {}
    for r in gen["rows"]:
        if r["struck"]:
            off.append(f"{r['name']} ({r['race']})")
            continue
        by_race.setdefault(r["race"], []).append(r)
    for race in sorted(by_race):
        drawn = order["order"].get(race)
        if drawn is None:
            raise SystemExit(f"Colorado: the ballot order drawing has no row for {race}")
        printed = [r for r in by_race[race] if not r["write_in"]]
        parties = [r["party"] for r in printed if r["party"] != "Unaffiliated"]
        if len(parties) != len(set(parties)):
            raise SystemExit(f"Colorado: {race} has two candidates of one party on the November list")
        missing = [p for p in parties if p not in drawn]
        if missing:
            raise SystemExit(f"Colorado: {race} has candidates of parties with no drawn position ({missing})")
        empty += [f"{p} ({race}, drawn {n})" for p, n in drawn.items() if p not in parties]
        placed = sorted((r for r in printed if r["party"] != "Unaffiliated"), key=lambda r: drawn[r["party"]])
        independents = [r for r in printed if r["party"] == "Unaffiliated"]
        if len(independents) == 1:
            placed.append(independents[0])
        elif independents:
            unplaced += [f"{r['name']} ({race})" for r in independents]
        position = {id(r): n for n, r in enumerate(placed, start=1)}
        for r in printed:
            rows.append((race, "general", "2026-11-03", r["name"], r["party"], party_code(r["party"]), position.get(id(r)), 0, 0,
                         None, None, None, None, None, "co-sos-2026-general-list", None))
            if r["party"] in CODES:
                nominee[(race, CODES[r["party"]])] = r["name"]
        for r in by_race[race]:
            if r["write_in"]:
                write_ins.append(r["name"])
                rows.append((race, "general", "2026-11-03", r["name"], r["party"], party_code(r["party"]), None, 0, 1,
                             None, None, None, None, None, "co-sos-2026-general-list", WRITE_IN))
    if unplaced:
        say(f"      two or more unaffiliated candidates in one race; their drawing is not in the file, so no position: {'; '.join(unplaced)}")

    # the primary list: who was on each party's ballot, who withdrew, who were certified write-ins
    listed, withdrew, certified = {}, [], {}
    for r in pri["rows"]:
        code = CODES.get(r["party"])
        if not code:
            raise SystemExit(f"Colorado: a party on the primary list that is not read ({r['party']!r})")
        if r["struck"]:
            withdrew.append(f"{r['name']} ({r['party']}, {r['race']})")
        elif r["write_in"]:
            certified.setdefault((r["race"], code), []).append(r["name"])
        else:
            listed.setdefault((r["race"], code), set()).add(r["name"])

    # the primary fields
    problems, fields, alone_with_write_in = [], 0, []
    usable = enr is not None and abstract_sha == ABSTRACT_SHA
    if enr is None:
        problems.append("the results site could not be read")
    elif abstract_sha != ABSTRACT_SHA:
        problems.append("the Secretary's abstract is not the scan whose totals were read (its SHA-256 changed); read it again and update ABSTRACT")
    contests = {}
    for c in (enr or {}).get("contests", []):
        code = CODES.get(c["party"])
        if not code:
            raise SystemExit(f"Colorado: a party on the results site that is not read ({c['party']!r})")
        contests[(c["race"], code)] = c
    keys = set(listed) | set(contests)
    for key in sorted(keys):
        race, code = key
        c = contests.get(key)
        names = listed.get(key, set())
        cert = certified.get(key, [])
        votes = {}
        if c:
            if c["reported"] != c["counties"]:
                problems.append(f"{c['title']}: {c['reported']} of {c['counties']} counties reported")
            if {ch["name"] for ch in c["choices"]} != names:
                problems.append(f"{race} {code}: the results name {sorted(ch['name'] for ch in c['choices'])}, the primary list {sorted(names)}")
            for ch in c["choices"]:
                if sum(ch["counties"].values()) != ch["total"]:
                    problems.append(f"{race} {code}: {ch['name']}'s county votes do not add up to the total")
                votes[ch["name"]] = ch["total"]
        elif usable:
            problems.append(f"{race} {code}: on the primary list but not on the results site")
        for name in cert:
            w = wins["candidates"].get(f"{race}|{name}")
            if w is None:
                problems.append(f"{race} {code}: the certified write-in {name} has no line in the write-in results")
            else:
                votes[name] = w["total"]
        if usable and ABSTRACT.get(key) != votes:
            problems.append(f"{race} {code}: the results site's votes {votes} are not the certified abstract's {ABSTRACT.get(key)}")
        if len(names) < 2:
            if cert:
                alone_with_write_in.append(f"{race} {code}: {', '.join(sorted(names))} and the certified write-in {', '.join(cert)}")
            continue
        fields += 1
    if problems:
        usable = False
        say("      primary votes not stored: " + "; ".join(problems))
    for key in sorted(listed):
        race, code = key
        names = sorted(listed[key])
        if len(names) < 2:
            continue
        party = next(r["party"] for r in pri["rows"] if r["race"] == race and CODES[r["party"]] == code)
        cert = certified.get(key, [])
        won = nominee.get(key)
        if won not in names:
            raise SystemExit(f"Colorado: the {race} {code} nominee on the November list ({won}) was not on that party's primary ballot")
        votes = {}
        if usable:
            votes = {ch["name"]: ch["total"] for ch in contests[key]["choices"]}
            votes.update({n: wins["candidates"][f"{race}|{n}"]["total"] for n in cert})
            top = max(votes, key=votes.get)
            if top != won:
                raise SystemExit(f"Colorado: {race} {code}: {top} had the most votes but {won} is on the November list")
        total = sum(votes.values())
        for name in sorted(names + cert, key=lambda n: (-votes.get(n, 0), n)):
            v = votes.get(name) if usable else None
            wi = name in cert
            rows.append((race, f"primary-{code}", PRIMARY, name, party, party_code(party), None, 0, int(wi), v,
                         round(100 * v / total, 1) if v is not None and total else None, "advanced" if name == won else "lost",
                         None, None, "co-sos-2026-primary-abstract" if usable else "co-sos-2026-primary-list", WRITE_IN if wi else None))

    general = [r for r in rows if r[1] == "general"]
    primary_rows = len(rows) - len(general)
    with con:
        con.execute("DELETE FROM candidates WHERE race_id LIKE '2026-CO-%'")
        con.executemany("INSERT OR REPLACE INTO candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", rows)
        record_source(con, "co-sos-2026-general-list", path=os.path.join(folder, "co_general_list_2026_federal.json"), level="federal",
                      state="CO", kind="official candidate list", agency="Colorado Secretary of State",
                      title="2026 General Election Official Candidate List (certified to the counties September 4, 2026): US Senate and US House of Representatives",
                      url=gen["url"], published=gen["updated"], rows=len(gen["rows"]),
                      note=f"The web page's table read (a struck-through name has withdrawn), columns Candidate name, Office, District, Party, "
                           f"Write in?; its workbook (SHA-256 {gen['workbook_sha256'][:16]}...) lists the same candidates"
                           + (f", written differently: {'; '.join(gen['differ'])}; the page's spelling is kept" if gen["differ"] else "")
                           + f". Withdrawn, left off: {len(off)}{' (' + '; '.join(off) + ')' if off else ''}. Certified write-ins, not printed on "
                           f"the ballot: {len(write_ins)}. Parties as printed.")
        record_source(con, "co-sos-2026-ballot-order", path=os.path.join(folder, "co_2026_general_ballot_order_federal.json"), level="federal",
                      state="CO", kind="official ballot order", agency="Colorado Secretary of State",
                      title="2026 General Election Ballot Order for Major & Minor Party Candidates (lot drawing of July 28, 2026)",
                      url=order["workbook"], published=order["updated"], rows=len(order["order"]),
                      note="Major parties, then minor parties, then unaffiliated candidates (1-5-404, C.R.S.), each tier by lot; the positions "
                           "stored are the places on the printed ballot, a party with no candidate left out. Drawn positions with no "
                           f"candidate on the certified list: {'; '.join(empty) or 'none'}.")
        record_source(con, "co-sos-2026-primary-list", path=os.path.join(folder, "co_primary_list_2026_federal.json"), level="federal",
                      state="CO", kind="official candidate list", agency="Colorado Secretary of State",
                      title="2026 Official Primary Election Candidate List (certified to the counties May 1, 2026): US Senate and US House of Representatives",
                      url=pri["url"], published=pri["updated"], rows=len(pri["rows"]),
                      note=f"Used to check the results: every printed name on each party's ballot. Withdrawn (struck through): "
                           f"{'; '.join(withdrew) or 'none'}; the Republican Party named Jason Clark in Mel Tewahade's place in District 6, "
                           f"and the votes cast for Tewahade count for Clark (1-4-1004(4)(b), C.R.S.). Certified write-ins: "
                           f"{'; '.join(f'{n} ({k[0]} {k[1]})' for k, ns in certified.items() for n in ns) or 'none'}.")
        record_source(con, "co-sos-2026-primary-abstract", path=abstract, level="federal", state="CO", kind="official results",
                      agency="Colorado Secretary of State", title="2026 State Primary Election Statewide Abstract of Votes Cast (June 30, 2026), certified July 24, 2026",
                      url=abstract_url, published="2026-07-24", rows=primary_rows if usable else 0,
                      note=("A scan with no text layer. Its federal Total rows were read by eye on 2026-09-30 and every one equals the "
                            "results site's figures, candidate by candidate, so the votes stored are the certified ones. "
                            if usable else "Votes not stored: " + "; ".join(problems) + ". Who advanced is read from the November list. ")
                           + "pct is of the party primary's votes, certified write-ins included (Colorado counts no other write-ins). "
                           + (f"One printed name and a certified write-in, not shown as a field: {'; '.join(alone_with_write_in)}." if alone_with_write_in else ""))
        if enr:
            record_source(con, "co-sos-2026-primary-enr", path=os.path.join(folder, "co_2026_primary_enr_federal.json"), level="federal",
                          state="CO", kind="results data, checked against the certified abstract", agency="Colorado Secretary of State (election night reporting, Clarity)",
                          title=f"2026 Primary: county detail (detail.xml), version {enr['version']}, written {enr['written']}",
                          url=enr["url"], rows=sum(len(c["choices"]) for c in enr["contests"]),
                          note=f"From the site at {ENR}, headed \"Unofficial Results\"; zip SHA-256 {enr['sha256'][:16]}.... The federal contests kept, "
                               "every county's votes per candidate; each candidate's counties add up to the candidate's total.")
        record_source(con, "co-sos-2026-primary-writeins", path=os.path.join(folder, "co_2026_primary_writeins_federal.json"), level="federal",
                      state="CO", kind="official results", agency="Colorado Secretary of State",
                      title="2026 State Primary Election Certified Write-In Results", url=wins["url"], rows=len(wins["candidates"]),
                      note="Certified write-in candidates for Congress, county by county; the counties add up to each Total Votes line.")
    say(f"    Colorado: 8 House districts and the Senate race, {len(general)} candidates on the November ballot ({len(write_ins)} certified "
        f"write-ins, {len(off)} withdrawn left off); {fields} party primaries with a field, "
        + ("votes from the Secretary's results site, equal to the certified abstract" if usable else "votes not stored (see above)"))
    return len(general)
