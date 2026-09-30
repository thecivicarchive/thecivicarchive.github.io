"""
Wyoming: the Secretary of State's Elections Division, two files linked from its 2026 Election page
(sos.wyo.gov/Elections/2026ElectionInformation.aspx).

The November ballot is the "2026 General Election Candidate Roster" (a PDF of every office, printed with the date and
time at the foot of each page, September 9, 2026 on the copy read first). Its columns are Office Sought, Party
Affiliation, Candidate Name, Mailing Address, Date Filed and Campaign Telephone, with City, State & ZIP, Date Withdrawn
and Email on each entry's second line. The column edges come from the page's own headings; only the office, the party
and the name are turned into text, and on an entry's second line only a date under Date Withdrawn is looked for.
Addresses, telephones and e-mail stay in the file and are never read. A candidate with a withdrawal date is left off.
The roster gives no ballot order, so the order is the roster's own (by party). Parties are printed in full
(Republican, Democratic, Libertarian, Constitution) and kept as printed; names are printed in ordinary capitals.

The August 18 primary is the Division's official results, "2026 Primary Results Summaries - OFFICIAL.xlsx" inside the
zip file of data files on its 2026 Official Primary Election Results page (the same figures as the Statewide
Candidates Summary PDF). Its "Statewide Candidates" sheet gives, across the columns, the office (a block continued is
labelled ", Continued"), the party, each candidate as "First\\nLast", then Write-Ins, Overvotes and Undervotes, with a
Total row. A party primary becomes a field when two or more names were printed on that party's ballot; write-ins count
toward the total but are not listed, and over- and undervotes are left out of it. A candidate who withdrew after the
ballots were printed is headed "* Withdrawn Candidate", named in the sheet's footnote, and kept in the field with the
votes the summary counts for them, and a note. The one on the November roster for that party advanced. Each field is
checked against the "Statewide Total Ballots Cast" sheet: candidates, write-ins, overvotes and undervotes add up to the
party's ballots. The minor parties (Libertarian, Constitution) nominate outside the primary and have no field.
Both hosts answer scripts.
"""

import io
import os
import re
import zipfile

import openpyxl

from ballot.common import fold, house_id, party_code, record_source, senate_id
from ballot.pdftext import PDF, join, rows as pdf_rows
from states import net

LIST_URL = "https://sos.wyo.gov/Elections/Docs/2026/2026_WY_General_Election_Candidates.pdf"
RESULTS_URL = "https://sos.wyo.gov/Elections/Docs/2026/Results/Primary/2026_Wyoming_Primary_Results.zip"
BOOK = "2026 Primary Results Summaries - OFFICIAL.xlsx"
PRIMARY = "2026-08-18"
OFFICES = {"UNITED STATES SENATOR": senate_id("WY", 2), "UNITED STATES REPRESENTATIVE": house_id("WY", 0)}
SHEET_OFFICES = {"United States Senator": senate_id("WY", 2), "United States Representative": house_id("WY", 0)}
PRINTED = re.compile(r"^[A-Z][a-z]+day, ([A-Z][a-z]+ \d\d?, \d{4}) - \d\d?:\d\d ?[AP]M$")
DATE = re.compile(r"\d\d/\d\d/\d{4}")
MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")
WITHDREW = re.compile(r"^\*\s*(?P<name>.+?) withdrew (?:his|her|their) candidacy after official ballots had been printed")
WITHDREW_NOTE = "Withdrew after the primary ballots had been printed; the official summary still counts the votes cast for this candidate."


def cells(runs):
    """A printed row's pieces grouped into cells: [(x, pieces)]; a gap wider than six points starts the next cell."""
    out = []
    for r in sorted(runs, key=lambda r: r[0]):
        if out and r[0] - out[-1][2] <= 6:
            out[-1][1].append(r)
            out[-1][2] = max(out[-1][2], r[4])
        else:
            out.append([r[0], [r], r[4]])
    return [(x, rs) for x, rs, _end in out]


def general_rows(path):
    """([(race, party, name, withdrawn date)], printed date) for the candidates for Congress on the roster."""
    pdf = PDF(open(path, "rb").read())
    out, office, printed, titled = [], None, "", False
    for page, res in pdf.pages():
        edges, body = {}, False
        for _y, runs in pdf_rows(pdf, page, res):
            row = [(x, rs) for x, rs in cells(runs) if x < 600]      # the Division's own address block sits at the right
            if not body:      # the title and the two heading lines; the column edges come from the headings
                heads = {join(rs): x for x, rs in row}
                titled = titled or "2026 General Election Candidate Roster" in heads
                if {"Party Affiliation", "Candidate Name", "Mailing Address"} <= set(heads):
                    edges.update(party=heads["Party Affiliation"] - 12, name=heads["Candidate Name"] - 12, addr=heads["Mailing Address"] - 12)
                elif "Date Withdrawn" in heads and edges:
                    edges.update(wd=heads["Date Withdrawn"] - 12, email=heads.get("Email", 10 ** 6) - 12)
                    body = True
                continue
            left = join([r for x, rs in row if x < edges["party"] for r in rs])
            party = join([r for x, rs in row if edges["party"] <= x < edges["name"] for r in rs])
            name = join([r for x, rs in row if edges["name"] <= x < edges["addr"] for r in rs])
            if party or name:
                if not (party and name):
                    raise SystemExit(f"Wyoming: a roster entry under {office} has a party or a name but not both; read the roster again")
                out.append([office, party, re.sub(r"\s+", " ", name), ""])
                continue
            if left:
                m = PRINTED.match(left)
                if m:
                    printed = re.sub(r" 0(\d),", r" \1,", m.group(1))      # September 09, 2026 -> September 9, 2026
                else:
                    office = left
                continue
            # an entry's second line: only a date under Date Withdrawn is looked for (city, ZIP and e-mail are not read)
            dates = [join(rs) for x, rs in row if edges["wd"] <= x < edges["email"]]
            dates = [d for d in dates if DATE.fullmatch(d)]
            if dates and out:
                out[-1][3] = dates[0]
    if not titled:
        raise SystemExit("Wyoming: the file is not the 2026 General Election Candidate Roster")
    fed = [(OFFICES[o], p, n, w) for o, p, n, w in out if o in OFFICES]
    missing = [o for o, race in OFFICES.items() if not any(r[0] == race for r in fed)]
    if missing:
        raise SystemExit(f"Wyoming: no candidates read under {', '.join(missing)} on the roster")
    return fed, printed


def iso(printed):
    m = re.fullmatch(r"([A-Z][a-z]+) (\d\d?), (\d{4})", printed or "")
    return f"{m.group(3)}-{MONTHS.index(m.group(1)) + 1:02d}-{int(m.group(2)):02d}" if m and m.group(1) in MONTHS else ""


def primary_fields(path):
    """{(race, party): {"names": [(name, votes, note)], "write_ins", "over", "under", "ballots"}} from the summaries workbook."""
    with zipfile.ZipFile(path) as z:
        if BOOK not in z.namelist():
            raise SystemExit(f"Wyoming: {BOOK} is not in the primary results zip ({z.namelist()})")
        wb = openpyxl.load_workbook(io.BytesIO(z.read(BOOK)), read_only=True, data_only=True)
    text = lambda c: re.sub(r"\s+", " ", str(c)).strip() if c is not None else ""
    ballots = {}
    cast = [list(r) for r in wb["Statewide Total Ballots Cast"].iter_rows(values_only=True)]
    head = next(r for r in cast if "Republican" in [text(c) for c in r])
    total = next(r for r in cast if r and text(r[0]) == "Total")
    for j, c in enumerate(head):
        if text(c):
            ballots[text(c)] = int(total[j] or 0)
    g = [list(r) for r in wb["Statewide Candidates"].iter_rows(values_only=True)]
    oi = next(i for i, r in enumerate(g) if any(text(c) in SHEET_OFFICES for c in r))
    offices, parties, names = g[oi], g[oi + 1], g[oi + 2]
    ti = next(i for i, r in enumerate(g) if r and text(r[0]) == "Total")
    totals = g[ti]
    notes = [(j, text(c)) for r in g[ti + 1:] for j, c in enumerate(r) if text(c).startswith("*")]

    def left_of(row, j, floor=0):
        k = next((k for k in range(j, floor - 1, -1) if text(row[k])), None)
        return (k, text(row[k])) if k is not None else (None, "")

    out = {}
    for j in range(1, len(names)):
        label = text(names[j])
        if not label:
            continue
        ocol, office = left_of(offices, j)
        race = SHEET_OFFICES.get(re.sub(r", Continued$", "", office))
        if not race:
            continue
        _pcol, party = left_of(parties, j, ocol)
        if not party:
            raise SystemExit(f"Wyoming: no party heads column {j + 1} of the primary summary ({office})")
        f = out.setdefault((race, party), {"names": [], "write_ins": 0, "over": 0, "under": 0, "ballots": ballots.get(party)})
        votes = int(totals[j] or 0)
        if label in ("Write-Ins", "Overvotes", "Undervotes"):
            f[{"Write-Ins": "write_ins", "Overvotes": "over", "Undervotes": "under"}[label]] = votes
        elif label.startswith("* Withdrawn"):
            said = [WITHDREW.match(n) for k, n in notes if SHEET_OFFICES.get(re.sub(r", Continued$", "", left_of(offices, k)[1])) == race]
            said = [m.group("name") for m in said if m]
            if len(said) != 1:
                raise SystemExit(f"Wyoming: the primary summary marks a withdrawn candidate for {office} without naming one in its footnote")
            f["names"].append((said[0], votes, WITHDREW_NOTE))
        else:
            f["names"].append((label, votes, None))
    return out


def load(con, cache, say=print):
    net.patient_lookups()
    folder = os.path.join(cache, "wy")
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, "wy_2026_general_candidate_roster.pdf")
    zpath = os.path.join(folder, "wy_2026_primary_results.zip")
    net.download(LIST_URL, path, max_age_days=2)
    net.download(RESULTS_URL, zpath, max_age_days=30)
    listed, printed = general_rows(path)
    on = [r for r in listed if not r[3]]
    gone = [r for r in listed if r[3]]
    rows, order = [], {}
    for race, party, name, _wd in on:
        order[race] = order.get(race, 0) + 1
        rows.append((race, "general", "2026-11-03", name, party, party_code(party), order[race], 0, 0, None, None, None, None, None,
                     "wy-sos-2026-general-roster", None))
    nominee = {(race, party): fold(name) for race, party, name, _wd in on}
    fields, unreconciled, upset = primary_fields(zpath), [], []
    for (race, party), f in sorted(fields.items()):
        if len(f["names"]) < 2:
            continue
        counted = sum(v for _n, v, _note in f["names"]) + f["write_ins"]
        if f["ballots"] is not None and counted + f["over"] + f["under"] != f["ballots"]:
            unreconciled.append(f"{race} {party}: {counted + f['over'] + f['under']:,} votes, over- and undervotes against {f['ballots']:,} ballots")
        code = {"Republican": "REP", "Democratic": "DEM"}.get(party, party[:3].upper())
        top = max(f["names"], key=lambda n: n[1])[0]
        if nominee.get((race, party)) != fold(top):
            upset.append(f"{race} {party}")
        for name, votes, note in f["names"]:
            rows.append((race, f"primary-{code}", PRIMARY, name, party, party_code(party), None, 0, 0, votes,
                         round(100 * votes / counted, 1) if counted else None,
                         "advanced" if nominee.get((race, party)) == fold(name) else "lost", None, None, "wy-sos-2026-primary-summary", note))
    with con:
        con.execute("DELETE FROM candidates WHERE race_id LIKE '2026-WY-%'")
        con.executemany("INSERT OR REPLACE INTO candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", rows)
        record_source(con, "wy-sos-2026-general-roster", path=path, level="federal", state="WY", kind="official candidate list",
                      agency="Wyoming Secretary of State, Elections Division", title="2026 General Election Candidate Roster"
                      + (f" (printed {printed})" if printed else ""), url=LIST_URL, published=iso(printed), rows=len(listed),
                      note=f"Every office; the candidates for Congress read (office, party and name only; addresses, telephones and e-mail are never "
                           f"read). The roster gives no ballot order; its own order, by party, is kept. Withdrawn, left off: {len(gone)}.")
        record_source(con, "wy-sos-2026-primary-summary", path=zpath, level="federal", state="WY", kind="official results",
                      agency="Wyoming Secretary of State, Elections Division",
                      title="2026 Official Primary Election Results (August 18, 2026): Statewide Candidates Official Summary, "
                            "from the zip file of data files (\"2026 Primary Results Summaries - OFFICIAL.xlsx\")",
                      url=RESULTS_URL, rows=sum(len(f["names"]) for f in fields.values() if len(f["names"]) > 1),
                      note="Votes from each office's Total row; write-ins count in the total but are not listed; over- and undervotes are left out. "
                           "A candidate who withdrew after the ballots were printed is kept with the votes the summary counts. "
                           + ("Every field adds up to its party's ballots cast." if not unreconciled else "Did not add up: " + "; ".join(unreconciled) + "."))
    if unreconciled:
        say("    Wyoming: primary fields that do not add up to the party's ballots cast: " + "; ".join(unreconciled))
    if upset:
        say("    Wyoming: the November nominee is not the primary's top vote-getter in " + ", ".join(upset) + "; read the files again")
    n = len(rows) - sum(1 for r in rows if r[1] != "general")
    say(f"    Wyoming: the Senate race and the at-large House seat, {n} candidates on the November ballot ({len(gone)} withdrawn left off); "
        f"{sum(1 for f in fields.values() if len(f['names']) > 1)} party primaries with a field, votes from the official statewide summary")
    return n
