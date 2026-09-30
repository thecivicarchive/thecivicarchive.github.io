"""
Ohio: the candidate lists of Ohio's county boards of elections for the November 3, 2026 general election. In Ohio a
candidate for the House files with, and is certified by, the board of the most populous county in the district; every
county board then prints the race on its own ballot and publishes its own list. The Secretary of State's own site
(ohiosos.gov, publicfiles.ohiosos.gov) and the county boards hosted on boe.ohio.gov and lookup.boe.ohio.gov (Hamilton,
Lucas, Summit, Montgomery, Licking and most small counties) answer scripts with a Cloudflare challenge, so they are not
read. Eight boards that publish on their own sites are read instead, each race from the first of them in this order that
carries it, and every other list that carries the race is checked against it:

  Franklin   "2026 General Election Candidates", revised September 22, 2026 (a spreadsheet printed turned on its side,
             one band across the page per column): the Senate, districts 3 and 15 (Franklin certifies both)
  Cuyahoga   "November 3, 2026 General Election Candidate List", 09/17/2026 (English and Spanish): districts 7 and 11
  Lake       "November 3, 2026 General Election Candidate Filings", updated 9/8/2026: district 14
  Lorain     "Candidate List", 09/14/2026: district 5 (also on Wood's list, checked)
  Stark      "Candidates List General Election Tuesday, November 3, 2026", printed 8/31/2026: districts 6 and 13
  Wood       "Candidate List for November 3, 2026", certified 08/26/2026: district 9
  Butler     "Candidate & Petition Activity, November 3, 2026, General Election", printed 9/28/2026: districts 8 and 10
  Union      the 46-day "Election Notice for Use with the Federal Write-in Absentee Ballot" (R.C. 3511.16), revised
             9/15/2026: district 4

  Hamilton  the 46-day "Election Notice" (R.C. 3511.16), 9/15/2026: district 1 (and checks the Senate and district 8).
             votehamiltoncountyohio.gov answers scripts with a Cloudflare challenge and serves browsers, so the notice
             was carried out of the Browser pane into ballot_cache/oh/ (the New York method in CLAUDE.md); it is read
             from there and never downloaded. The notice prints names, offices, parties and precincts, no addresses.

The Senate race is on all nine and every list agrees on its printed names. Districts 2 and 12 are left out: no board
that prints them publishes a list a script or the Browser pane could reach on 2026-09-30 (see MISSING). Printed names are fixed 70 days out (August 25;
R.C. 3513.30(E) keeps a later withdrawal on the ballot), and the write-in deadline was August 24, so lists from late
August on are final for the ballot; a board that does not certify a race may still leave out its write-in candidates.

Only the name, party, office and status columns are read, each by its position on the page; the lists also print
addresses, telephones, e-mail and websites, which are never turned into text. A candidate a list marks withdrawn or
removed is left off; certified write-in candidates are kept only where the list names them (write_in 1; Butler's list
prints "Write-in Candidate" without names, and those lines are skipped). Ohio rotates the order of names
from precinct to precinct (R.C. 3505.03) and the lists order them differently, so no ballot order is stored and the page
shows them by surname.

Parties are shown as the list prints them. Written out: D, R, L (Lake, Butler) and Dem, Rep, Lib (Union) as Democratic,
Republican, Libertarian; Cuyahoga's "Democratic | Democrata" as its English half. Ohio's terms for candidates of no
recognised party, "other-party" and "nonparty", are printed differently from list to list (Greg Levy is "Other-party" in
Franklin, "Other Party" in Lorain, "Other-party candidate" in Wood, "Nonpartisan" in Cuyahoga); a candidate the list gives
no party (Franklin's "N/A", or nothing, as on Union's notice) is shown as "No party", and a write-in with no party as
"Write-in". Union's notice prints names in capitals; they are shown in ordinary capitals and the row says so.

The May 5, 2026 primaries come from the Secretary of State's official canvass: one "Summary Level Official Results for
2026 Primary Election" workbook per party (Democratic, Republican, Libertarian), listed in
publicfiles.ohiosos.gov/election-results/files-index.json. That host answers scripts with a Cloudflare challenge and
serves the Secretary's own Data Portal in a browser, so the three workbooks were carried out of the Browser pane into
ballot_cache/oh/ and are read from there. Their "U.S. Congress" sheet has one column per candidate: the office in the
first row (over the office's first column only), "Name (R)" or "Name (WI)" for a write-in in the second, and the
statewide Total in the third. A field is a party primary with two printed candidates or more; write-in votes count
toward its total, and a named write-in is shown as one. The winner is the top vote-getter, checked against the
November lists.
"""

import datetime as dt
import os
import re

from ballot.common import fold, house_id, party_code, record_source, senate_id
from ballot.lists.tx import proper
from ballot.pdftext import PDF, join, lines, page_runs, rows as pdf_rows
from states import net

SENATE = senate_id("OH", 3)
WRITE_IN_NOTE = "Write-in candidate: the name is not printed on the ballot."
CAPS_NOTE = "Union County's notice prints names in capitals; they are shown here in ordinary capitals."
LISTS = {      # key: (board, title, address); read in this order, and a race is taken from the first list that has it
    "franklin": ("Franklin County Board of Elections", "2026 General Election Candidates",
                 "https://vote.franklincountyohio.gov/getmedia/e0a4ca45-6893-45c1-8190-87696f8804b4/2026-General-Candidate-List-5"),
    "cuyahoga": ("Cuyahoga County Board of Elections", "November 3, 2026 General Election Candidate List",
                 "https://boe.cuyahogacounty.gov/docs/default-source/boe/candidates-page/candidate-list.pdf"),
    "lake": ("Lake County Board of Elections", "November 3, 2026 General Election Candidate Filings",
             "https://wpassets.lakecountyohio.gov/wp-content/uploads/sites/12/2026/09/15104255/2026-Candidate-Filings.pdf"),
    "lorain": ("Lorain County Board of Elections", "Candidate List, November 3, 2026 General Election",
               "https://www.voteloraincountyohio.gov/_files/ugd/2568d0_ef2bd7f0adc244659d5357bfa1879d07.pdf"),
    "stark": ("Stark County Board of Elections", "Candidates List, General Election Tuesday, November 3, 2026",
              "https://www.starkcountyohio.gov/Board%20of%20Elections/Candidates%20List%20General.pdf"),
    "wood": ("Wood County Board of Elections", "Candidate List for November 3, 2026 (certified 08/26/2026)",
             "https://boe.woodcountyohio.gov/DocumentCenter/View/1006/Candidate-List"),
    "butler": ("Butler County Board of Elections", "Candidate & Petition Activity, November 3, 2026, General Election",
               "https://cms2.revize.com/revize/repository/butlercountyboe/2026/November/PetitionActivityReport.pdf"),
    "union": ("Union County Board of Elections", "Election Notice for Use with the Federal Write-in Absentee Ballot (FWAB), "
              "46 Day Notice, November 3, 2026", "https://www.unioncountyohio.gov/media/Agencies/Board%20of%20Elections/"
              "45%20Day%20notice%20with%20candidates.pdf"),
    "hamilton": ("Hamilton County Board of Elections", "Election Notice (46 Day Notice), November 3, 2026 General Election",
                 "https://votehamiltoncountyohio.gov/wp-content/uploads/2026/09/G26-46-Day-Notice-9-15-26.pdf"),
}
CARRIED = {"hamilton": "oh_hamilton_2026_46day_notice.pdf"}      # carried out of the Browser pane (a Cloudflare challenge for scripts)
PRIMARY = "2026-05-05"
PRIMARY_BOOKS = {      # party code: (file carried out of the Browser pane, its address in the Secretary's file index)
    code: (f"oh_primary_2026_summary_{word}.xlsx",
           "https://publicfiles.ohiosos.gov/election-results/past-elections/2026/Primary+Special%20Election%20-%20May%205,%202026/group1/"
           f"summary-level-official-results-2026-primary---{word}.xlsx")
    for code, word in (("DEM", "democratic"), ("REP", "republican"), ("LIB", "libertarian"))}
PARTY_OF = {"DEM": "Democratic", "REP": "Republican", "LIB": "Libertarian"}
MISSING = {      # why the districts no reachable board prints are left out, and what a browser would carry out
    2: "no board in the district that a script can reach publishes a candidate list (the boards on boe.ohio.gov are behind a Cloudflare challenge)",
    12: "Licking County's board prints it, and its certified candidate list (boe.ohio.gov/licking, ELECTION_Candidates.pdf) is behind a Cloudflare challenge for scripts, and boe.ohio.gov was down (Cloudflare 522) when tried in the Browser pane on 2026-09-30",
}
WORDS = {"D": "Democratic", "R": "Republican", "L": "Libertarian", "DEM": "Democratic", "REP": "Republican", "LIB": "Libertarian",
         "N/A": "", "O": "Other-party"}
PRINTED = {      # where each list prints its date, read from its first page
    "franklin": r"Revised ([A-Z][a-z]+ \d{1,2}, \d{4})", "cuyahoga": r"^(\d\d/\d\d/\d{4}) \d", "lake": r"Updated: (\d{1,2}/\d{1,2}/\d{4})",
    "lorain": r"^(\d\d/\d\d/\d{4})$", "stark": r"^(\d{1,2}/\d{1,2}/\d{4}) \d", "wood": r"(?:Mon|Tues|Wednes|Thurs|Fri|Satur|Sun)day, ([A-Z][a-z]+ \d{1,2}, \d{4})",
    "butler": r"Page 1 of \d+ (\d{1,2}/\d{1,2}/\d{4})", "union": r"Revised: (\d{1,2}/\d{1,2}/\d{4})", "hamilton": r"(\d{1,2}/\d{1,2}/\d{4})"}


def code(party):
    """party_code, with Ohio's "Non-Party" and "nonparty" read as no party."""
    return "I" if re.fullmatch(r"non-?party(?: candidate)?", party.strip().lower()) else party_code(party)


def printed_on(key, path):
    for _page, _y, text in lines(path, 1, 1):
        m = re.search(PRINTED[key], text)
        if m:
            d = m.group(1)
            fmt = "%B %d, %Y" if d[0].isalpha() else "%m/%d/%Y"
            return dt.datetime.strptime(d, fmt).strftime("%Y-%m-%d")
    return ""


def race_of(text):
    """The race a heading names, or None: U.S. Senator (this year's special election) or a numbered House district."""
    t = re.sub(r"\s+", " ", text.upper())
    if re.search(r"\bU\. ?S\. SENAT|UNITED STATES SENAT", t):
        return SENATE
    m = re.search(r"(?:REPRESENTATIVE TO CONGRESS|UNITED STATES CONGRESS|U\. ?S\. HOUSE|CONGRESSIONAL DISTRICT)\D{0,12}?(\d{1,2})(?!\d)", t)
    if m:
        n = int(m.group(1))
        return house_id("OH", n) if 1 <= n <= 15 else None
    return None


def row(race, name, party, write_in=False, note=None):
    party = WORDS.get(party.strip().upper(), party.strip()) if party else ""
    if not party:
        party = "Write-in" if write_in else "No party"
    return {"race": race, "name": re.sub(r"\s+", " ", name).strip(), "party": party, "write_in": bool(write_in), "note": note}


def printed_rows(path):
    """(page, y, runs left to right) for every printed row of a PDF, top to bottom."""
    pdf = PDF(open(path, "rb").read())
    for n, (page, res) in enumerate(pdf.pages(), start=1):
        for y, rs in pdf_rows(pdf, page, res):
            yield n, y, sorted(rs, key=lambda r: r[0])


def col(rs, lo, hi):
    return join([r for r in rs if lo <= r[0] < hi])


def franklin(path):
    """Printed turned on its side: each column of the spreadsheet is a band across the page, each candidate a position
    along it. Only the Name on Ballot, Party, Office and Write-In bands are read, matched by position."""
    pdf = PDF(open(path, "rb").read())
    keys = ("Name on Ballot", "Party", "Office", "Write-In")
    band, out, gone = {}, [], []
    for n, (page, res) in enumerate(pdf.pages(), start=1):
        runs = page_runs(pdf, page, res)
        if not band:
            band = {r[3].strip(): r[1] for r in runs if r[0] < 70 and r[3].strip() in keys}
            if len(band) != len(keys):
                raise SystemExit(f"Franklin: the list's columns changed (found {sorted(band)})")
        cells = {k: [(r[0], r[3].strip()) for r in runs if abs(r[1] - band[k]) < 3 and r[0] >= 60 and r[3].strip() != k] for k in keys}
        for x, office in cells["Office"]:
            race = race_of(office)
            if not race:
                continue
            name = [t for x2, t in cells["Name on Ballot"] if abs(x2 - x) < 3]
            party = [t for x2, t in cells["Party"] if abs(x2 - x) < 3]
            if len(name) != 1 or len(party) > 1:
                raise SystemExit(f"Franklin: a row for {office} does not line up across the list's columns (page {n})")
            wi = any(abs(x2 - x) < 3 and t == "Yes" for x2, t in cells["Write-In"])
            out.append(row(race, name[0], party[0].strip('" ') if party else "", wi))
    return out, gone


def cuyahoga(path):
    """Candidate Name, Party, Filed and Valid columns; the office is named in the heading of each block."""
    out, gone, race = [], [], None
    for _page, _y, rs in printed_rows(path):
        left = col(rs, 0, 340)
        if left.startswith("For "):
            race = SENATE if left.startswith("For United States Senator") else ("H?" if left.startswith("For Representative to Congress") else None)
            continue
        m = re.match(r"\((\d+)(?:st|nd|rd|th) District\)", left)
        if race == "H?" and m:
            race = house_id("OH", int(m.group(1)))
            continue
        filed = col(rs, 390, 470)
        if not race or race == "H?" or not re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}", filed):
            continue
        name, party, valid = col(rs, 30, 230), col(rs, 230, 390).split("|")[0].strip(), col(rs, 470, 700)
        r = row(race, name, party, "Write-In" in valid)
        (out if valid.startswith("Yes") else gone).append(r)
    return out, gone


def lake(path):
    """Office, Party and Name columns (the address columns to their right are never read). Each block opens with a
    PARTY / NAME heading row; the office is named in the first column beside it."""
    table = list(printed_rows(path))
    heads = [(p, y) for p, y, rs in table if col(rs, 185, 250).replace(" ", "") == "PARTY"]
    out, gone, race = [], [], None
    for page, y, rs in table:
        if (page, y) in heads:
            label = " ".join(col(r2, 0, 185) for p2, y2, r2 in table if p2 == page and abs(y2 - y) <= 16)
            race = race_of(label)
            continue
        if re.fullmatch(r"[A-Z]+OFFICES", join(rs).replace(" ", "")):
            race = None
            continue
        name = col(rs, 250, 470)
        if race and name:
            wi = "(Write-in)" in name.replace(" ", "") or "(Write-" in name
            out.append(row(race, re.sub(r"\s*\(Write-?\s*in\)", "", name, flags=re.I), col(rs, 185, 250), wi))
    return out, gone


LORAIN_PARTY = {"Democratic", "Republican", "Libertarian", "Green", "Other Party", "Nonparty", "Non-Party", "Independent", "Write-In"}


def lorain(path):
    """Candidate and Party columns (the contact column is never read). A candidate's name starts at the left margin;
    an office heading is set in a little from it and runs across both columns."""
    out, gone, race = [], [], None
    for _page, _y, rs in printed_rows(path):
        name, party = col(rs, 15, 185), col(rs, 185, 300)
        if not name or name == "Candidate":
            continue
        if party in LORAIN_PARTY:
            if race:
                wi = party == "Write-In"
                out.append(row(race, name, "" if wi else party, wi))
            continue
        if race and rs[0][0] < 25:
            raise SystemExit(f"Lorain: a candidate for {race} with a party the loader does not know ({party!r})")
        race = race_of(join([r for r in rs if r[0] < 300]))
    return out, gone


STARK_ON = {"Nominated", "Filed", "Write In"}
STARK_OFF = {"Withdrawn", "Withdrew", "Rejected", "Invalid", "Disqualified", "Removed", "Not Certified"}


def stark(path):
    """Status and Name columns under a party heading (the address and telephone columns are never read). Office
    headings (which carry nothing else) are read whole."""
    out, gone, race, party = [], [], None, ""
    for _page, _y, rs in printed_rows(path):
        left, name = col(rs, 0, 150), col(rs, 150, 280)
        if re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4} \d{1,2}:\d\d:\d\d [AP]M|Page \d+.*", left):
            continue
        if left in STARK_ON or left in STARK_OFF:
            if race:
                wi = left == "Write In"
                r = row(race, name, "" if wi else party, wi)
                (out if left in STARK_ON else gone).append(r)
            continue
        if left:
            race, party = race_of(join(rs)), ""
            continue
        if name:
            party = name
    return out, gone


WOOD_PARTY = re.compile(r"^(Democratic|Republican|Libertarian|Green|Other-party candidate|Nonparty candidate|Write-In Candidate)$")


def wood(path):
    """A party (or Write-In Candidate) in the first column, the name after it; addresses sit on the next row, which
    has no party, and are never read."""
    out, gone, race = [], [], None
    for _page, _y, rs in printed_rows(path):
        left = col(rs, 0, 190)
        if WOOD_PARTY.match(left):
            if race:
                wi = left == "Write-In Candidate"
                out.append(row(race, col(rs, 190, 480), "" if wi else left, wi))
            continue
        if left and not re.match(r"(Mon|Tues|Wednes|Thurs|Fri|Satur|Sun)day, ", left):
            race = race_of(join(rs))
    return out, gone


def butler(path):
    """One line per candidate: (party) name, then where the candidate files or the petition dates, and Withdrawn."""
    out, gone, race = [], [], None
    for _page, _y, rs in printed_rows(path):
        text = join(rs)
        if not text.startswith("(") and (" - (Vote for not" in text or " - For " in text or " - Representative" in text):
            race = race_of(text)
            continue
        m = re.match(r"^\((\w)\) (.+?)(?: Files with .*| \(P\) .*)?$", text)
        if not (m and race):
            continue
        name = m.group(2).strip()
        if name == "Write-in Candidate":      # a write-in the report does not name
            continue
        r = row(race, name, m.group(1))
        (gone if "Withdrawn" in text else out).append(r)
    seen, kept = set(), []
    for r in out:      # the report prints some blocks twice
        if (r["race"], r["name"]) not in seen:
            seen.add((r["race"], r["name"]))
            kept.append(r)
    return kept, gone


def union(path):
    """Name of Candidate (Party) and Office columns of the notice; each office's block opens on the row that names its
    precincts. Write-in candidates follow a WRITE-IN CANDIDATES line, each after a bullet."""
    blocks, out = [], []
    for _page, _y, rs in printed_rows(path):
        right = join([r for r in rs if r[0] >= 245])
        if re.search(r"Precincts$", right) and "Name of Candidate" not in join(rs):
            blocks.append({"office": [], "names": []})
        if not blocks:
            continue
        blocks[-1]["office"].append(re.sub(r"\s*All Precincts$", "", right))
        name = col(rs, 30, 245)
        if name:
            blocks[-1]["names"].append(name)
    for b in blocks:
        race = race_of(" ".join(b["office"]))
        if not race:
            continue
        wi = False
        for name in b["names"]:
            if name.upper().startswith("WRITE-IN CANDIDATES"):
                wi = True
                continue
            name = re.sub(r"^[\u2022\x95\ufffd·•]\s*", "", name)
            m = re.match(r"^(.+?)\s*\((Dem|Rep|Lib|Grn)\)$", name)
            out.append(row(race, proper(m.group(1) if m else name), m.group(2) if m else "", wi, CAPS_NOTE))
    return out, []


HAM_OFFICE = re.compile(r"^(U\.\s?S\.|State |Judge|Justice|Member|County|Governor|Attorney|Auditor|Secretary|Treasurer|Ohio )")


def hamilton(path):
    """The notice's four columns: name (x 38-175), office (175-295), party (295-385) and precincts, which are not read.
    An office opens a block where its title starts; the rows under it, until the next title, are its candidates. A long
    name or party wraps onto the next row ("Anthony Holliman (Write-" / "In)", "Other-party" / "candidate")."""
    blocks = []
    for page, _y, rs in printed_rows(path):
        if page == 1:
            continue
        whole = join(rs)
        if "Ohio General Assembly" in whole:
            break
        name, office, party = col(rs, 38, 175), col(rs, 175, 295), col(rs, 295, 385)
        if "Name of Candidate" in whole:
            continue
        if office and HAM_OFFICE.match(office):
            blocks.append({"office": [], "cands": []})
        if not blocks:
            continue
        b = blocks[-1]
        if office:
            b["office"].append(office)
        if name:
            last = b["cands"][-1][0] if b["cands"] else ""
            if last.endswith("(Write-") or (b["cands"] and re.fullmatch(r"\(?(Write-)?\s?In\)", name)):
                b["cands"][-1][0] = re.sub(r"\(Write-\s*\(?(?:Write-)?\s?In\)", "(Write-In)", last + " " + name)
            else:
                b["cands"].append([name, party])
        elif party and b["cands"]:
            b["cands"][-1][1] = (b["cands"][-1][1] + " " + party).strip()
    out = []
    for b in blocks:
        text = " ".join(b["office"]).upper()
        m = re.search(r"REPRESENTATIVE\D*?(\d{1,2})\s*(?:ST|ND|RD|TH)?\s*DISTRICT", text)
        race = house_id("OH", int(m.group(1))) if m else (SENATE if re.search(r"U\.\s?S\. SENAT", text) else None)
        if not race:
            continue
        for name, party in b["cands"]:
            wi = bool(re.search(r"\(Write-?\s?In\)", name, re.I))
            name = re.sub(r"\s*\(Write-?\s?In\)\s*", " ", name, flags=re.I).strip()
            party = re.sub(r"\s+candidate$", "", party.strip(), flags=re.I)
            out.append(row(race, name, party, wi))
    return out, []


def primary_fields(cache):
    """{(race, party code): [(name, votes, write_in)]} from the Secretary of State's canvass workbooks, and the files read."""
    import openpyxl
    fields, files = {}, []
    for pcode, (name, _url) in PRIMARY_BOOKS.items():
        path = os.path.join(cache, "oh", name)
        if not os.path.exists(path):
            continue
        if open(path, "rb").read(2) != b"PK":
            raise SystemExit(f"Ohio: {name} is not a workbook")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if "U.S. Congress" not in wb.sheetnames:
            raise SystemExit(f"Ohio: {name} has no U.S. Congress sheet ({wb.sheetnames})")
        head, names, total = [list(r) for r in wb["U.S. Congress"].iter_rows(min_row=1, max_row=3, values_only=True)]
        if str(total[0]).strip() != "Total" or "Official Canvass" not in str(head[0]):
            raise SystemExit(f"Ohio: {name}'s U.S. Congress sheet does not open with the official canvass and a Total row")
        files.append(path)
        office = None
        for i in range(6, len(names)):
            if head[i]:
                office = re.sub(r"\s+", " ", str(head[i])).strip()
            if not names[i] or office is None:
                continue
            m = re.search(r"District (\d{1,2})", office)
            race = house_id("OH", int(m.group(1))) if m and "Congress" in office else (SENATE if office.startswith("U.S. Senator") else None)
            if not race:
                continue
            label = re.sub(r"\s+", " ", str(names[i])).strip()
            wi = "(WI)" in label
            wm = re.match(r"^(.*?)\s*\((R|D|L)\)$", label)
            if not wm and not wi:
                raise SystemExit(f"Ohio: a candidate heading in {name} does not end with a party mark: {label}")
            person = re.sub(r"\s*\(WI\)\s*\*?|\s*\((?:R|D|L)\)$", " ", label).strip(" *")
            person = re.sub(r"\s*-\s+|\s+-\s*", "-", re.sub(r"\s+", " ", person))
            fields.setdefault((race, pcode), []).append((person, int(total[i] or 0), wi))
    return fields, files


READERS = {"franklin": franklin, "cuyahoga": cuyahoga, "lake": lake, "lorain": lorain, "stark": stark, "wood": wood,
           "butler": butler, "union": union, "hamilton": hamilton}


def load(con, cache, say=print):
    net.patient_lookups()
    folder = os.path.join(cache, "oh")
    got, paths = {}, {}
    for key, (_board, _title, url) in LISTS.items():
        paths[key] = os.path.join(folder, CARRIED.get(key, f"oh_{key}_2026_general.pdf"))
        if key in CARRIED and not os.path.exists(paths[key]):
            say(f"    Ohio: the {key} notice has not been carried out of the Browser pane into {folder}")
            continue
        try:
            if key not in CARRIED:
                net.download(url, paths[key], max_age_days=7, say=say)
        except Exception as e:  # noqa: BLE001  a board's site can be down; its races come from another list or wait
            say(f"    Ohio: {key} list not fetched ({e})")
            continue
        got[key] = READERS[key](paths[key])
        if SENATE not in {r["race"] for r in got[key][0]}:
            say(f"      Ohio: no Senate race read from the {key} list; its layout may have changed")
    chosen, checks = {}, []
    for key, (on, _gone) in got.items():
        for race in sorted({r["race"] for r in on}):
            here = [r for r in on if r["race"] == race]
            if race not in chosen:
                chosen[race] = (key, here)
                continue
            first_key, first = chosen[race]
            def printed(cands):      # family name, and the party where it is one of the big four
                return {(fold(r["name"]).split()[-1], code(r["party"]) if code(r["party"]) in "DRLG" else "-") for r in cands if not r["write_in"]}
            def written(cands):
                return {fold(r["name"]).split()[-1] for r in cands if r["write_in"]}
            if printed(first) != printed(here):
                checks.append(f"{race}: {first_key} prints {sorted(printed(first))}, {key} prints {sorted(printed(here))}")
            elif written(here) - written(first):
                checks.append(f"{race}: {key} names write-ins {sorted(written(here) - written(first))} that {first_key} does not")
    rows, used = [], {}
    for race, (key, cands) in sorted(chosen.items()):
        used.setdefault(key, []).append(race)
        for c in cands:
            note = "; ".join(n for n in (c["note"], WRITE_IN_NOTE if c["write_in"] else None) if n) or None
            rows.append((race, "general", "2026-11-03", c["name"], c["party"], code(c["party"]), None, 0, int(c["write_in"]),
                         None, None, None, None, None, f"oh-{key}-2026-general", note))
    nominee = {(r[0], r[4]): fold(r[3]) for r in rows}
    shown = {(r[0], r[4]): r[3] for r in rows}
    fields, pfiles = primary_fields(cache)
    nfields, disagree = 0, []
    for (race, pcode), cands in sorted(fields.items()):
        if sum(1 for _n, _v, wi in cands if not wi) < 2:
            continue
        nfields += 1
        total = sum(v for _n, v, _wi in cands)
        top = max(cands, key=lambda c: c[1])
        party = PARTY_OF[pcode]
        listed = nominee.get((race, party))
        if listed and fold(top[0]).split()[-1] != listed.split()[-1]:
            disagree.append(f"{race} {pcode}: the canvass's top vote-getter is {top[0]}; the November list names {shown[(race, party)]} (a replacement nominee)")
        replaced = bool(listed) and fold(top[0]).split()[-1] != listed.split()[-1]
        for person, votes, wi in cands:
            won = person == top[0]
            note = "; ".join(x for x in (
                WRITE_IN_NOTE if wi else None,
                "Won the nomination; not on the November lists read." if won and race in chosen and not listed else None,
                f"Won the primary; the November list names {shown[(race, party)]} as the party's candidate instead." if won and replaced else None) if x) or None
            rows.append((race, f"primary-{pcode}", PRIMARY, person, "Write-in" if wi else party, "W" if wi else party_code(party), None, 0, int(wi),
                         votes, round(100 * votes / total, 1) if total else None, "advanced" if won else "lost", None, None,
                         f"oh-sos-2026-primary-{pcode.lower()}", note))
    races = {r[0] for r in con.execute("SELECT race_id FROM races WHERE state = 'OH'")}
    missing = sorted(races - set(chosen))
    with con:
        con.execute("DELETE FROM list_gaps WHERE state = 'OH'")
        con.executemany("INSERT INTO list_gaps VALUES (?, 'OH', ?)",
                        [(race, "the county board that certifies it does not post its list anywhere we could read it yet") for race in missing])
        con.execute("DELETE FROM candidates WHERE race_id LIKE '2026-OH-%'")
        con.executemany("INSERT OR REPLACE INTO candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", rows)
        for key, (on, gone) in got.items():
            board, title, url = LISTS[key]
            mine = used.get(key, [])
            also = [r for r in sorted({r["race"] for r in on}) if r not in mine]
            note = (f"Races taken from this list: {', '.join(mine) or 'none'}" + (f"; also carries {', '.join(also)}, checked against the list used"
                    if also else "") + ". Name, party, office and status columns read by position; addresses, telephones, e-mail and websites "
                    "never read." + (f" Marked withdrawn or removed, left off: {len(gone)}." if gone else "")
                    + (" Write-in candidates the report prints without names are not read." if key == "butler" else ""))
            if key in CARRIED:
                note += " Carried out of the Browser pane: the board's site answers scripts with a Cloudflare challenge."
            record_source(con, f"oh-{key}-2026-general", path=paths[key], level="federal", state="OH", kind="official candidate list",
                          agency=board, title=title, url=url, published=printed_on(key, paths[key]), rows=len(on), note=note)
        for pcode, (name, url) in PRIMARY_BOOKS.items():
            path = os.path.join(folder, name)
            if path in pfiles:
                record_source(con, f"oh-sos-2026-primary-{pcode.lower()}", path=path, level="federal", state="OH", kind="official results",
                              agency="Ohio Secretary of State", title=f"Summary Level Official Results for 2026 Primary Election - {PARTY_OF[pcode]} "
                              "(Official Canvass, May 5, 2026; U.S. Congress sheet)", url=url,
                              rows=sum(len(v) for (r, pc), v in fields.items() if pc == pcode),
                              note="Statewide Total row of each candidate. Carried out of the Browser pane: the Secretary of State's file host "
                                   "answers scripts with a Cloudflare challenge. Write-in votes count toward a field's total.")
    n = sum(1 for r in rows if r[1] == "general")
    say(f"    Ohio: {len([r for r in chosen if '-H' in r])} House districts and {'the' if SENATE in chosen else 'no'} Senate race from "
        f"{len(used)} county boards' lists, {n} candidates on the November ballot ({sum(1 for r in rows if r[1] == 'general' and r[8])} "
        f"certified write-ins); {nfields} party primaries with a field, votes from the official canvass")
    for d in disagree:
        say(f"      check: {d}")
    for c in checks:
        say(f"      check: {c}")
    for race in missing:
        d = int(race[-2:]) if "-H" in race else None
        say(f"      {race} not loaded: {MISSING.get(d, 'no reachable board prints it')}")
    return n
