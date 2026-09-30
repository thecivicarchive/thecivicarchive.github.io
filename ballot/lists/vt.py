"""
Vermont: the Secretary of State's Elections Division (sos.vermont.gov/elections). Vermont has one House seat, elected
at large, and no U.S. Senate race in 2026 (its senators' seats are classes 1 and 3). The primary was held on
August 11, 2026 ("08/11/2026 - AUGUST PRIMARY" on the Secretary's own files).

  November ballot   "2026 General Election Candidate Listing/Financial Disclosure" (the Candidates page,
                    candidates/2026_general_election_qualified_candidates.xlsx): one sheet, "Candidate Listing", one row
                    per qualified candidate for every office on the ballot, and a "Selection Criteria" sheet whose one
                    line says when the list was last updated (9/23/2026 when this was written).
  primary list      "2026 Primary Election Candidate Listing/Financial Disclosure" (2026_statewide_primary_qualified_
                    candidates.xlsx): one sheet per party (DEMOCRATIC, PROGRESSIVE, REPUBLICAN) with the same columns,
                    a "REGISTERED WRITE-IN" sheet (office, districts, first, middle and last name) and a "Last Updated"
                    sheet. Used to check that the canvass names exactly the candidates on each party's primary ballot,
                    and to count the registered write-ins.
  primary results   "2026 August Primary Official Canvass - Town by Town" (the Elections Results & Data page, a PDF of
                    411 pages printed 08/17/2026). It opens with the Official Report of the Canvassing Committee (17
                    V.S.A. 2368-2371): one column per candidate, write-ins, overvotes, blank votes and total votes counted,
                    party by party, with the winners starred; then, office by office and party by party, a page of
                    statewide totals by county and a page (or two) per county with every town and the county total. The
                    pages are printed turned on their side; each column is read by its position (pdftext's page_runs):
                    a town's figures are its whole-number cells, top to bottom, against the row labels, top to bottom,
                    and a heading belongs to the one column of figures it sits over (a long town name, ESSEX JUNCTION
                    CITY, wraps onto a second line beside its column).
  control           "2026 August Primary Winner Listing" (2026-august-primary-winner-listing.xlsx): the winner mark,
                    votes and percentage of every candidate. Its votes and winners must agree with the canvass.

Checks on the canvass, for every party's primary for Representative to Congress: every town's candidates, write-ins,
overvotes and blank votes add up to its total; the towns of each county add up to the county total; the county totals
are the statewide page's county columns and add up to its statewide totals; and those are the canvassing committee's
figures. A field's total is its candidates' votes plus the write-in votes (the canvass reports write-ins as one figure
per town, without names, so no write-in candidate is shown); overvotes and blank votes are left out of it, so the
percentages here are not the canvass's, which are of total votes counted. A field is a party primary with two
candidates or more on its ballot; the Democratic primary (Becca Balint alone) and the Progressive primary (no candidate
printed, write-ins only) are not fields. The committee's starred winner must be the top vote-getter, the Winner
Listing's winner, and the party's candidate on the November list, or the row says plainly that the winner is not on it.

The candidate workbooks also carry every candidate's town of residence, mailing address, city, state, ZIP, day and
evening phones, e-mail, website and financial disclosure; the winner listing carries addresses and a phone. Workbooks
are read in memory and never saved; columns are taken by name, only Contest, District Name, Name On Ballot, Party and
Vote for Count (and Office and the three name columns of the write-in sheets; Winner, Name on Ballot, Party, Office
Name, District, Votes and Percent(%) of the winner listing); nothing else is read, and the cache (ballot_cache/vt/)
keeps only those columns of the federal rows, as JSON. The canvass holds results only and is kept as published.

The November ballot is every "REPRESENTATIVE TO CONGRESS" row of the general list. The list gives no ballot positions;
Vermont's ballot lists each office's candidates "in alphabetical order by surname" (17 V.S.A. 2472(b)(2)), the list's
own order, which the loader checks and stores as the ballot order. The list has no status column: a candidate who
withdrew is no longer listed, so none can be counted. A write-in candidate for November may register until 5 p.m. on
the Thursday before the election (the Candidates page); the list names none (no write-in sheet), and if one appears
its Congress rows are stored with write_in 1 and no ballot position. Names and parties are printed in capitals and
shown in ordinary capitals (a sitting member as the congress-legislators roster spells the same words), and the row
says so; the parties are printed in full (DEMOCRATIC, REPUBLICAN, INDEPENDENT) and shown as words.
"""

import io
import json
import os
import re
import sqlite3
import time

import openpyxl

from ballot.common import HERE, fold, house_id, name_parts, party_code, record_source
from ballot.lists.tx import proper
from ballot.pdftext import PDF, page_runs
from states import net

CANDIDATES_PAGE = "https://sos.vermont.gov/elections/election-info-resources/candidates/"
RESULTS_PAGE = "https://sos.vermont.gov/elections/election-info-resources/elections-results-data/"
FILES = "https://outside.vermont.gov/dept/sos/Elections_Division/election_info_resources/"
GENERAL_XLSX = FILES + "candidates/2026_general_election_qualified_candidates.xlsx"
PRIMARY_XLSX = FILES + "candidates/2026_statewide_primary_qualified_candidates.xlsx"
CANVASS_PDF = FILES + "elections_results_data/2026-august-primary-official-canvass-town-by-town.pdf"
WINNERS_XLSX = FILES + "elections_results_data/2026-august-primary-winner-listing.xlsx"
RACE = house_id("VT", 0)
PRIMARY = "2026-08-11"
OFFICE = "REPRESENTATIVE TO CONGRESS"
HEADING = "FOR " + OFFICE
KEEP = ("Contest", "District Name", "Name On Ballot", "Party", "Vote for Count")
WRITE_KEEP = ("OFFICE", "FIRST NAME", "MIDDLE NAME", "LAST NAME")
WIN_KEEP = ("Winner", "Name on Ballot", "Party", "Office Name", "District", "Votes", "Percent(%)")
PARTY_SHEETS = ("DEMOCRATIC", "PROGRESSIVE", "REPUBLICAN")
CODE = {"DEMOCRATIC": "DEM", "REPUBLICAN": "REP", "PROGRESSIVE": "PRO"}          # election code: primary-DEM, primary-REP
SHORT = {"DEM": "Democratic", "REP": "Republican", "PROG": "Progressive"}          # the lists' short forms in joint labels
COUNTIES = 14
OTHER = ("Write-In", "Overvotes", "Blank votes")
LABEL = {"WRITE-IN": "Write-In", "OVERVOTES": "Overvotes", "BLANK VOTES": "Blank votes", "TOTAL VOTES COUNTED": "Total",
         "Town Total": "Total"}
NUMBER = re.compile(r"\d{1,3}(?:,\d{3})+|\d+")
CAPS = "Vermont's list prints names in capitals; they are shown here in ordinary capitals."
WRITE_IN = "Write-in candidate: the name is not printed on the ballot."


def squash(value):
    return re.sub(r"\s+", " ", str(value if value is not None else "").replace("\xa0", " ")).strip()


def workbook(url):
    """A workbook of the Secretary's, read in memory (it carries contact columns, so it is never saved)."""
    raw = net.get(url)
    if raw[:2] != b"PK":
        raise SystemExit(f"Vermont: {url} is not a workbook (it begins {raw[:8]!r})")
    return openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)


def sheet_rows(ws, keep, first):
    """The rows under a sheet's header row (the first row whose first cell is `first`), the kept columns only, by name."""
    rows = ws.iter_rows(values_only=True)
    for r in rows:
        if r and squash(r[0]) == first:
            head = [squash(h) for h in r]
            break
    else:
        raise SystemExit(f"Vermont: no header row beginning {first!r} on the sheet {ws.title!r}")
    missing = [k for k in keep if k not in head]
    if missing:
        raise SystemExit(f"Vermont: the sheet {ws.title!r} no longer has the columns {missing}")
    idx = {k: head.index(k) for k in keep}
    out = []
    for r in rows:
        row = {k: squash(r[i]) if i < len(r) else "" for k, i in idx.items()}
        if any(row.values()):
            out.append(row)
    return out


def updated(wb, sheet):
    """The "Last Updated: M/D/YYYY" line of a workbook's criteria sheet, as YYYY-MM-DD; nothing else on it is kept."""
    if sheet not in wb.sheetnames:
        return ""
    for r in wb[sheet].iter_rows(values_only=True):
        for c in r:
            m = re.fullmatch(r"Last Updated:\s*(\d{1,2})/(\d{1,2})/(\d{4})", squash(c))
            if m:
                return f"{m.group(3)}-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
    return ""


def federal_check(rows, where):
    """Every row for Congress is for the one House seat; any other federal office stops the loader."""
    for r in rows:
        c = r.get("Contest", r.get("OFFICE", "")).upper()
        if c.replace("U.S. ", "") == OFFICE:
            continue
        if "CONGRESS" in c or c.startswith(("U.S.", "UNITED STATES")):
            raise SystemExit(f"Vermont: a federal office on the {where} that is not read ({c!r})")


def write_ins(wb, where):
    """Congress rows of a registered write-in sheet: office and the three name columns only."""
    out = []
    for name in wb.sheetnames:
        if "WRITE-IN" in name.upper():
            rows = sheet_rows(wb[name], WRITE_KEEP, "OFFICE")
            federal_check(rows, where)
            out += [r for r in rows if r["OFFICE"].upper().replace("U.S. ", "") == OFFICE]
    return out


def kept(path, max_age_days, fetch):
    if os.path.exists(path) and time.time() - os.path.getmtime(path) < max_age_days * 86400:
        return json.load(open(path, encoding="utf-8"))
    data = fetch()
    os.makedirs(os.path.dirname(path), exist_ok=True)
    json.dump(data, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return data


def general_list():
    wb = workbook(GENERAL_XLSX)
    sheet = "Candidate Listing" if "Candidate Listing" in wb.sheetnames else wb.sheetnames[0]
    rows, contest = sheet_rows(wb[sheet], KEEP, "Contest"), ""
    for r in rows:                                               # a contest named once for the rows below it
        contest = r["Contest"] = r["Contest"] or contest
    federal_check(rows, "general list")
    return {"url": GENERAL_XLSX, "sheet": sheet, "updated": updated(wb, "Selection Criteria"), "rows_read": len(rows),
            "congress": [r for r in rows if r["Contest"].upper() == OFFICE], "write_in": write_ins(wb, "general list")}


def primary_list():
    wb = workbook(PRIMARY_XLSX)
    parties, contest = {}, ""
    for sheet in wb.sheetnames:
        if sheet.upper() in PARTY_SHEETS:
            rows = sheet_rows(wb[sheet], KEEP, "Contest")
            for r in rows:                                       # a contest named once for the rows below it
                contest = r["Contest"] = r["Contest"] or contest
            federal_check(rows, f"primary list ({sheet})")
            parties[sheet.upper()] = [r for r in rows if r["Contest"].upper() == OFFICE]
        elif "WRITE-IN" not in sheet.upper() and sheet != "Last Updated":
            raise SystemExit(f"Vermont: the primary list has a sheet that is not read ({sheet!r})")
    if set(parties) != set(PARTY_SHEETS):
        raise SystemExit(f"Vermont: the primary list's party sheets changed ({sorted(parties)})")
    return {"url": PRIMARY_XLSX, "updated": updated(wb, "Last Updated"), "parties": parties,
            "write_in": write_ins(wb, "primary list")}


def winner_listing():
    wb = workbook(WINNERS_XLSX)
    rows = sheet_rows(wb["Winner Listing"], WIN_KEEP, "Winner")
    crit = {}
    if "Selection Criteria" in wb.sheetnames:
        for r in wb["Selection Criteria"].iter_rows(values_only=True):
            cells = [squash(c) for c in r if c is not None]
            if len(cells) == 2 and cells[0].endswith(":"):
                crit[cells[0].rstrip(":")] = cells[1]
    for r in rows:
        if "CONGRESS" in r["Office Name"].upper() and r["Office Name"].upper() != OFFICE:
            raise SystemExit(f"Vermont: a federal office in the winner listing that is not read ({r['Office Name']!r})")
    return {"url": WINNERS_XLSX, "criteria": crit, "congress": [r for r in rows if r["Office Name"].upper() == OFFICE]}


def number(text):
    t = text.strip()
    return int(t.replace(",", "")) if NUMBER.fullmatch(t) else None


def canvass(path):
    """(committee, stars, statewide, counties, towns, printed) for Representative to Congress, party by party:
    the canvassing committee's figures {party: {label: votes}}, its starred winners {(party, name)}, the statewide page
    {party: {column: {label: votes}}}, each county's County Total {party: {county: {label: votes}}}, every town's
    figures {party: {county: [{label: votes}]}} and the date the tabulation pages were printed."""
    raw = open(path, "rb").read()
    if raw[:5] != b"%PDF-":
        raise SystemExit("Vermont: the canvass is not a PDF")
    pdf = PDF(raw)
    committee, stars, statewide, county_total, towns = {}, set(), {}, {}, {}
    office, party, seen, printed, attested = None, None, False, "", False
    for n, (page, res) in enumerate(pdf.pages(), start=1):
        runs = page_runs(pdf, page, res)
        footer = [r for r in runs if re.fullmatch(r"Page \d+ of \d+", r[3].strip())]
        if not footer:                                          # the canvassing committee's report
            if seen:
                raise SystemExit(f"Vermont: canvass page {n} is laid out like the committee's report, after the tabulation")
            text = " ".join(r[3] for r in runs)
            attested = attested or ("17 V.S.A. 2368-2371" in text and "08/11/2026 - AUGUST PRIMARY" in text)
            events = [(r[0], "office", r[3].strip()) for r in runs if 115 <= r[1] <= 125 and r[3].strip().startswith("FOR ")]
            events += [(r[0], "column", r[3].strip()) for r in runs if 165 <= r[1] <= 175]
            for x, kind, label in sorted(events):
                if kind == "office":
                    office, party = label, None
                    continue
                if office != HEADING:
                    continue
                if label.endswith(" PARTY"):
                    party = label[:-6]
                    committee.setdefault(party, {})
                    continue
                vals = [number(r[3]) for r in runs if 415 <= r[1] <= 425 and abs(r[0] - x) <= 3 and number(r[3]) is not None]
                if party is None or len(vals) != 1:
                    raise SystemExit(f"Vermont: canvass page {n}: the committee's column {label!r} is not read ({vals})")
                committee[party][LABEL.get(label, label)] = vals[0]
                if any(r[3].strip() == "*" and 200 <= r[1] <= 300 and abs(r[0] - x) <= 3 for r in runs):
                    stars.add((party, label))
            continue
        fx = footer[0][0]
        date = [r[3].strip() for r in runs if abs(r[0] - fx) < 1 and re.fullmatch(r"\d\d/\d\d/\d{4}", r[3].strip())]
        xs = sorted({round(r[0], 1) for r in runs if abs(r[0] - fx) >= 1})
        titles = [r[3].strip() for r in sorted(runs, key=lambda r: r[1]) if round(r[0], 1) == xs[0]]
        if not titles or titles[0] != HEADING:
            if seen:
                break
            continue
        if len(titles) != 3:
            raise SystemExit(f"Vermont: canvass page {n} does not carry an office, an area and a party ({titles})")
        seen = True
        printed = printed or (date[0] if date else "")
        area, party = titles[1], titles[2]
        heads = [LABEL.get(r[3].strip(), r[3].strip()) for r in sorted(runs, key=lambda r: r[1]) if round(r[0], 1) == xs[1]]
        if len(heads) < 4 or heads[-1] != "Total" or set(heads[-4:-1]) != set(OTHER) or set(heads[:-4]) & set(OTHER):
            raise SystemExit(f"Vermont: canvass page {n}'s rows are not read ({heads})")
        body = [r for r in runs if abs(r[0] - fx) >= 1 and round(r[0], 1) not in xs[:2]]
        cols = sorted({round(r[0], 1) for r in body if r[1] >= 60})       # a column is where its figures are printed
        names = [r for r in body if r[1] < 60]                            # a long town name wraps onto a second line
        if any(sum(1 for x in cols if abs(r[0] - x) < 10) != 1 for r in names):
            raise SystemExit(f"Vermont: canvass page {n}: a column heading does not sit over exactly one column of figures")
        for x in cols:
            col = [r for r in body if round(r[0], 1) == x and r[1] >= 60]
            name = " ".join(r[3].strip() for r in sorted(names, key=lambda r: r[0]) if abs(r[0] - x) < 10)
            vals = [number(r[3]) for r in sorted(col, key=lambda r: r[1]) if "%" not in r[3]]
            if not name or None in vals or len(vals) != len(heads):
                raise SystemExit(f"Vermont: canvass page {n}, column {name!r}: {len(vals)} figures for {len(heads)} rows")
            fig = dict(zip(heads, vals))
            if sum(v for k, v in fig.items() if k != "Total") != fig["Total"]:
                raise SystemExit(f"Vermont: canvass page {n}, {name} ({party}): the figures do not add up to the total")
            if area == "STATEWIDE TOTALS":
                statewide.setdefault(party, {})[name] = fig
            elif area.endswith(" COUNTY"):
                county = area[:-7]
                if name == "County Total":
                    if county in county_total.get(party, {}):
                        raise SystemExit(f"Vermont: {county} County's total for {party} is printed twice")
                    county_total.setdefault(party, {})[county] = fig
                else:
                    towns.setdefault(party, {}).setdefault(county, []).append(fig)
            else:
                raise SystemExit(f"Vermont: canvass page {n} is for an area that is not read ({area!r})")
    if not attested:
        raise SystemExit("Vermont: the canvass does not open with the canvassing committee's report of the August 11 primary")
    return committee, stars, statewide, county_total, towns, printed


def primary_votes(path, listed):
    """{party: [(name, votes)], write-in votes, starred} for Representative to Congress; every sum checked."""
    committee, stars, statewide, county_total, towns, printed = canvass(path)
    if set(committee) != set(listed) or set(statewide) != set(listed):
        raise SystemExit(f"Vermont: the canvass's parties for Congress ({sorted(committee)}, {sorted(statewide)}) are not the primary list's")
    out, ntowns = {}, 0
    for party in sorted(listed):
        sw = statewide[party]
        counties = {c for c in sw if c != "Statewide Totals"}
        if len(counties) != COUNTIES or set(county_total.get(party, {})) != counties or set(towns.get(party, {})) != counties:
            raise SystemExit(f"Vermont: the canvass does not carry all {COUNTIES} counties for the {party} primary")
        labels = list(sw["Statewide Totals"])
        for county in counties:
            summed = {k: sum(t[k] for t in towns[party][county]) for k in labels}
            if summed != county_total[party][county] or county_total[party][county] != sw[county]:
                raise SystemExit(f"Vermont: {county} County's towns, county total and statewide column differ ({party})")
            ntowns += len(towns[party][county])
        if {k: sum(sw[c][k] for c in counties) for k in labels} != sw["Statewide Totals"]:
            raise SystemExit(f"Vermont: the counties do not add up to the statewide totals ({party})")
        if committee[party] != sw["Statewide Totals"]:
            raise SystemExit(f"Vermont: the canvassing committee's figures differ from the statewide totals ({party})")
        field = [(k, v) for k, v in sw["Statewide Totals"].items() if k not in OTHER and k != "Total"]
        if {fold(k) for k, _v in field} != {fold(r["Name On Ballot"]) for r in listed[party]}:
            raise SystemExit(f"Vermont: the canvass's {party} candidates are not the primary list's")
        out[party] = (field, sw["Statewide Totals"]["Write-In"], {k for p, k in stars if p == party})
    return out, printed, ntowns


def load(con, cache, say=print):
    net.patient_lookups()
    folder = os.path.join(cache, "vt")
    gen_path = os.path.join(folder, "vt_2026_general_federal.json")
    pri_path = os.path.join(folder, "vt_2026_primary_federal.json")
    win_path = os.path.join(folder, "vt_2026_primary_winner_listing_federal.json")
    pdf_path = os.path.join(folder, "vt_2026_primary_official_canvass_town_by_town.pdf")
    gen = kept(gen_path, 2, general_list)
    pri = kept(pri_path, 30, primary_list)
    net.download(CANVASS_PDF, pdf_path, max_age_days=30, say=say)
    win = kept(win_path, 30, winner_listing)

    rec = sqlite3.connect(f"file:{os.path.join(HERE, 'congress_119.sqlite')}?mode=ro", uri=True)
    fixed = {}
    for full, first, last in rec.execute("SELECT official_full, first_name, last_name FROM legislators WHERE is_current = 1 AND state = 'VT'"):
        for form in (full, f"{first} {last}"):
            if form:
                fixed[fold(form)] = form
    rec.close()

    def shown(caps):
        caps = squash(caps)
        if fold(caps) in fixed:                     # the roster's capitals, only where its words are the same words
            return fixed[fold(caps)]
        return re.sub(r"(['\"(])([a-z])", lambda m: m.group(1) + m.group(2).upper(), proper(caps))

    def party_of(label):
        """DEMOCRATIC -> Democratic; a joint nomination written short (DEM/PROG) -> Democratic/Progressive."""
        if not label:
            raise SystemExit("Vermont: a candidate for Congress with no party on the list")
        return "/".join(SHORT.get(p.strip(), p.strip().title()) for p in label.split("/"))

    # the November ballot
    congress = gen["congress"]
    if not congress:
        raise SystemExit("Vermont: no candidate for Representative to Congress on the general list")
    for r in congress:
        if r["District Name"] not in ("N/A", "") or r["Vote for Count"] != "ONE":
            raise SystemExit(f"Vermont: a Congress row names a district or seat count that is not read ({r['District Name']!r}, {r['Vote for Count']!r})")
    surnames = [name_parts(shown(r["Name On Ballot"]))[1] for r in congress]
    if surnames != sorted(surnames):
        raise SystemExit(f"Vermont: the general list is not in alphabetical order by surname ({surnames}); read the ballot order")
    rows, nominee = [], {}
    for i, r in enumerate(congress, start=1):
        party = party_of(r["Party"])
        rows.append((RACE, "general", "2026-11-03", shown(r["Name On Ballot"]), party, party_code(party), i, 0, 0,
                     None, None, None, None, None, "vt-sos-2026-general-list", CAPS))
        if r["Party"] in CODE:
            if (r["Party"]) in nominee:
                raise SystemExit(f"Vermont: two {r['Party']} candidates for Congress on the general list")
            nominee[r["Party"]] = fold(r["Name On Ballot"])
    for r in gen["write_in"]:
        caps = " ".join(x for x in (r["FIRST NAME"], r["MIDDLE NAME"], r["LAST NAME"]) if x)
        rows.append((RACE, "general", "2026-11-03", shown(caps), "Write-in", "W", None, 0, 1,
                     None, None, None, None, None, "vt-sos-2026-general-list", f"{WRITE_IN} {CAPS}"))

    # the August primary
    listed = pri["parties"]
    votes, printed, ntowns = primary_votes(pdf_path, listed)
    listing = {}
    for r in win["congress"]:
        if r["District"] not in ("", "None"):
            raise SystemExit(f"Vermont: a winner-listing row for Congress names a district ({r['District']!r})")
        listing[(r["Party"], fold(r["Name on Ballot"]))] = (int(r["Votes"].replace(",", "")), bool(r["Winner"]))
    if win["criteria"].get("Election Name") != "08/11/2026 - AUGUST PRIMARY":
        raise SystemExit(f"Vermont: the winner listing is for {win['criteria'].get('Election Name')!r}")
    fields = 0
    for party, (field, write_in_votes, starred) in sorted(votes.items()):
        for name, v in field:
            if listing.get((party, fold(name)), (None,))[0] != v:
                raise SystemExit(f"Vermont: the winner listing's votes for {name} ({party}) differ from the canvass")
        won = {fold(n) for n in starred}
        top = max(field, key=lambda nv: nv[1]) if field else None
        if field and (sum(1 for _n, v in field if v == top[1]) > 1 or won != {fold(top[0])}):
            raise SystemExit(f"Vermont: the {party} primary's starred winner is not its only top vote-getter")
        if {k[1] for k, (_v, w) in listing.items() if k[0] == party and w} != won:
            raise SystemExit(f"Vermont: the winner listing's {party} winner is not the canvass's")
        if len(field) < 2:
            continue
        fields += 1
        total = sum(v for _n, v in field) + write_in_votes
        printed_as = {fold(r["Name On Ballot"]): r["Name On Ballot"] for r in listed[party]}
        on_nov = nominee.get(party)
        for name, v in sorted(field, key=lambda nv: -nv[1]):
            advanced = fold(name) == fold(top[0])
            note = CAPS
            if advanced and on_nov != fold(name):
                note = f"Won the primary; not on the November list as the {party.title()} candidate. {CAPS}"
            party_label = party_of(party)
            rows.append((RACE, f"primary-{CODE[party]}", PRIMARY, shown(printed_as[fold(name)]), party_label, party_code(party_label),
                         None, 0, 0, v, round(100 * v / total, 1) if total else None, "advanced" if advanced else "lost",
                         None, None, "vt-sos-2026-primary-canvass", note))

    general = [r for r in rows if r[1] == "general"]
    write_in_n = sum(1 for r in general if r[8])
    primary_rows = len(rows) - len(general)
    m = re.fullmatch(r"(\d\d)/(\d\d)/(\d{4})", printed)
    printed_iso = f"{m.group(3)}-{m.group(1)}-{m.group(2)}" if m else ""
    pri_write = ", ".join(" ".join(x for x in (r["FIRST NAME"], r["MIDDLE NAME"], r["LAST NAME"]) if x).title() for r in pri["write_in"])
    with con:
        con.execute("DELETE FROM candidates WHERE race_id LIKE '2026-VT-%'")
        con.execute("DELETE FROM list_gaps WHERE state = 'VT'")
        con.executemany("INSERT OR REPLACE INTO candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", rows)
        record_source(con, "vt-sos-2026-general-list", path=gen_path, level="federal", state="VT", kind="official candidate list",
                      agency="Vermont Secretary of State, Elections Division",
                      title="2026 General Election Candidate Listing/Financial Disclosure (qualified candidates): Representative to Congress",
                      url=GENERAL_XLSX, published=gen["updated"], rows=len(general),
                      note=f"Linked from the Candidates page ({CANDIDATES_PAGE}). Workbook read in memory, never saved; columns taken by "
                           "name, only Contest, District Name, Name On Ballot, Party and Vote for Count; town of residence, addresses, "
                           "phones, e-mail, websites and financial disclosures never read. The list gives no ballot positions; its "
                           "order, alphabetical by surname as 17 V.S.A. 2472(b)(2) prints the ballot, is stored as the ballot order. "
                           "No status column, so a withdrawn candidate, no longer listed, cannot be counted. Registered write-in "
                           f"candidates named on the list: {write_in_n} (registration stays open until the Thursday before the election).")
        record_source(con, "vt-sos-2026-primary-list", path=pri_path, level="federal", state="VT", kind="official candidate list",
                      agency="Vermont Secretary of State, Elections Division",
                      title="2026 Primary Election Candidate Listing/Financial Disclosure (qualified candidates): Representative to Congress",
                      url=PRIMARY_XLSX, published=pri["updated"], rows=sum(len(v) for v in listed.values()) + len(pri["write_in"]),
                      note="Used for who was on each party's primary ballot and how their names are printed; the canvass's candidates "
                           "must be exactly these. The same kept columns only. Its REGISTERED WRITE-IN sheet names "
                           f"{len(pri['write_in'])} write-in candidates for Congress ({pri_write or 'none'}), without a party; the canvass "
                           "counts write-in votes without names, so they are not shown.")
        record_source(con, "vt-sos-2026-primary-canvass", path=pdf_path, level="federal", state="VT", kind="official results",
                      agency="Vermont Secretary of State, Elections Division",
                      title="2026 August Primary Official Canvass - Town by Town (Official Report of the Canvassing Committee, "
                            "August 11, 2026): For Representative to Congress",
                      url=CANVASS_PDF, published=printed_iso, rows=primary_rows,
                      note=f"Linked from the Elections Results & Data page ({RESULTS_PAGE}). Read by position: the committee's report "
                           f"and the statewide, county and town pages for Congress. {ntowns} town columns; every town adds up to its total, "
                           "the towns to each county's total, the counties to the statewide totals, and those equal the committee's "
                           "figures. A field's total is its candidates' votes plus write-in votes (overvotes and blank votes left out).")
        record_source(con, "vt-sos-2026-primary-winners", path=win_path, level="federal", state="VT", kind="official results",
                      agency="Vermont Secretary of State, Elections Division",
                      title="2026 August Primary Winner Listing: Representative to Congress",
                      url=WINNERS_XLSX, published="", rows=len(win["congress"]),
                      note="Control: every candidate's votes and the winner marks agree with the canvass. Workbook read in memory, "
                           "never saved; Winner, Name on Ballot, Party, Office Name, District, Votes and Percent(%) only; addresses "
                           "and phones never read.")
    say(f"    Vermont: the at-large House seat (no Senate race in 2026), {len(general)} candidates on the November ballot "
        f"({write_in_n} declared write-in); {fields} party primary with a field, votes from the Secretary of State's official "
        f"canvass, {ntowns} town columns summed to county, statewide and canvassing-committee totals and checked against the "
        "winner listing")
    return len(general)
