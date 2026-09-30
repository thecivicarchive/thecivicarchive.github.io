"""
Kansas: the Secretary of State's own lists and results. Kansas has four House seats and the Senate seat of class 2
(Roger Marshall) in 2026; the districts are the 2022 lines.

  November ballot   the Secretary's "Candidate List" page (sos.ks.gov/elections/elections_upcoming_candidate.aspx)
                    with "2026 General" chosen in its list of elections and Submit pressed. The site sits behind
                    Amazon CloudFront, which answers scripts on that page with 403 (three times on 2026-09-30, with
                    pauses; the site's plain pages and its results files answer), so it is not asked again. John
                    saves the page in his own browser (Ctrl+S, "Webpage, Complete" or "Webpage, Single File") into
                    ballot_cache/ks/, and this loader reads every saved page there (.html, .htm, .mhtml, .mht) and
                    takes the newest whose list of elections shows "2026 General" chosen (or, with no such list,
                    which names "2026 General" outside its table). Until one is there no November candidates are
                    stored: the primary winners alone would leave off the Libertarian and independent candidates,
                    and any nominee replaced since August.
                    The page's table (gvCandidateList, as laid out in 2020) has one heading row: Candidate, Office,
                    District, Position, Division, Party, Title, First Name, Middle, Last Name, Suffix, then home and
                    mailing addresses, phones, e-mail, web address, date filed and ballot cities. Columns are taken
                    by name, only Candidate (else First Name, Middle, Last Name and Suffix), Office, District and
                    Party, and a Status or Ballot Order column if the list carries one; no other cell is ever turned
                    into text. Only the federal rows are kept, as JSON in ballot_cache/ks/ with the saved page's
                    SHA-256. A page the grid's pager splits into several pages stops the loader (every row must be
                    on the saved page). A row whose status says withdrawn, removed or the like is left off and
                    counted. The list prints no ballot order, so each race keeps the list's own order. Parties are
                    printed in full and kept as printed; a name printed in capitals would be shown in ordinary
                    capitals and noted, and one written "Last, First" turned round.
  primary fields    the Secretary's "2026 Primary Election Official Vote Totals" (a PDF linked as "2026 Official
                    Primary Election Results" from sos.ks.gov/elections/election-results.html; it prints no date,
                    its file properties say September 1, 2026), read with ballot/pdftext.py: under each race
                    ("United States Senate", "United States House of Representatives 3") one line per candidate:
                    the party's initial and the name as on the ballot ("D-Sharice L. Davids"; D is Democratic and
                    R Republican, as the precinct workbooks write them out), the votes and the percent. A declared
                    write-in candidate who got votes is marked "(Write-in)" (none for Congress in 2026) and is kept
                    as one; the totals carry no count of other write-in votes, so a candidate's share is of the
                    listed candidates' votes, which is how the totals print their percent (each is checked).
                    A field is a party primary with two or more names printed on its ballot. Kansas nominates the
                    candidate with the most votes and has no runoff, so the leader advanced; once the November list
                    is loaded, the one on it for that party must be the same person, and a row says so if not.
                    Libertarians nominate by convention and have no primary.
  cross-check       the Secretary's precinct results workbooks for the same primary, "U.S. Senate 2026 Primary
                    Election" and "U.S. House of Representatives 2026 Primary Election" on the same page. The first
                    sheet has a row per precinct and candidate for every county but the four largest (County,
                    Precinct, Race, Candidate as "Last, First", Party, Votes, VTD; only Race, Candidate, Party and
                    Votes are read); Johnson, Sedgwick, Shawnee and Wyandotte each have a sheet of their own, the
                    candidates across the columns and a COUNTY TOTALS row. The county sheets' COUNTY TOTALS rows are
                    used, because some of Shawnee's precinct cells are printed "redacted per KSA 25-2701(a)(2)".
                    Every candidate's county sum must equal the statewide votes in the Official Vote Totals.
"""

import datetime as dt
import email
import email.policy
import glob
import hashlib
import html as H
import json
import os
import re
import time
from collections import defaultdict
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin

import openpyxl

from ballot.common import fold, house_id, name_parts, party_code, record_source, senate_id
from ballot.lists.tx import proper
from ballot.pdftext import lines
from states import net

LIST_PAGE = "https://sos.ks.gov/elections/elections_upcoming_candidate.aspx"
RESULTS_PAGE = "https://sos.ks.gov/elections/election-results.html"
PRIMARY = "2026-08-04"
ELECTION = "2026 General"
FILES = {      # what the results page calls each file, the address it had on 2026-09-30, the name kept here, its first bytes
    "totals": ("2026 Official Primary Election Results", "26elec/2026-Primary-Election-Official-Vote-Totals.pdf",
               "ks_2026_primary_official_vote_totals.pdf", b"%PDF"),
    "senate": ("U.S. Senate 2026 Primary Election", "26elec/2026-Primary-Election-United-States-Senate-Precinct-Level-Results.xlsx",
               "ks_2026_primary_us_senate_precincts.xlsx", b"PK"),
    "house": ("U.S. House of Representatives 2026 Primary Election",
              "26elec/2026-Primary-Election-United-States-House-of-Representatives-Precinct-Level-Results.xlsx",
              "ks_2026_primary_us_house_precincts.xlsx", b"PK")}
RACES = [senate_id("KS", 2)] + [house_id("KS", d) for d in range(1, 5)]
INITIAL = {"D": "Democratic", "R": "Republican"}
CODE = {"Democratic": "DEM", "Republican": "REP", "Libertarian": "LIB"}
LINE = re.compile(r"^(?P<p>[A-Z])-(?P<name>.+?) (?P<votes>\d[\d,]*) (?P<pct>\d+\.\d\d)%$")
FURNITURE = re.compile(r"^(Kansas Secretary of State|2026 Primary Election|Official Vote Totals|Page \d+ of \d+|Race Candidate Votes Percent)$")
NAME_COL, PARTS = "Candidate", ("First Name", "Middle", "Last Name", "Suffix")
KEEP, OPTIONAL = ("Office", "District", "Party"), ("Status", "Ballot Order")
GONE = ("withdr", "disqual", "remov", "deceas", "denied", "inactive", "void")
ON = ("", "active", "filed", "certified", "nominated", "qualified", "accepted")
SUFFIX = re.compile(r"(Jr|Sr|II|III|IV|V)\.?", re.I)
CAPS_NOTE = "Kansas's list prints names in capitals; they are shown here in ordinary capitals."
WRITE_IN_NOTE = "Write-in candidate: the name is not printed on the ballot."
NOT_ON_LIST = "Won the primary, but is not on the November list for this party."
JSON_NAME = "ks_candidate_list_2026_general_federal.json"


def key(name):
    """A name for matching: letters only, spaces made single."""
    return " ".join(fold(name).split())


def same(a, b):
    """The same person under two spellings: the same family name and the same first letter of the first given name
    (the list and the ballot may differ by a middle initial, a nickname or a suffix). One nominee per party makes it safe."""
    ga, fa = name_parts(a)
    gb, fb = name_parts(b)
    return bool(fa) and fa == fb and bool(ga) and bool(gb) and ga[0][0] == gb[0][0]


def text(fragment):
    return re.sub(r"\s+", " ", H.unescape(re.sub(r"<[^>]+>", " ", fragment or ""))).strip()


def race_of(office, district=""):
    """'United States Senate' -> 2026-KS-S2; 'United States House of Representatives 3' (or the office with a District
    of 3) -> 2026-KS-H03; any other office -> None."""
    o = re.sub(r"\s+", " ", office or "").strip()
    if re.fullmatch(r"(?:United States|U\.S\.) Senat(?:e|or)", o, re.I):
        return senate_id("KS", 2)
    m = re.fullmatch(r"(?:United States|U\.S\.) (?:House of Representatives|Representative)(?: (?:District )?(\d+))?", o, re.I)
    if not m:
        return None
    d = m.group(1) or re.sub(r"\D", "", str(district or ""))
    if not d or not 1 <= int(d) <= 4:
        raise SystemExit(f"Kansas: a row for the U.S. House names no district from 1 to 4 ({o!r}, district {district!r})")
    return house_id("KS", int(d))


def turned(raw):
    """'Young, Chad E' -> 'Chad E Young'; 'Smith, John Jr.' -> 'John Smith Jr.'; 'Thomas D. Hall, Jr.' stays."""
    name = re.sub(r"\s+", " ", str(raw or "")).strip()
    last, sep, rest = name.partition(",")
    if not sep or not rest.strip() or SUFFIX.fullmatch(rest.strip()):
        return name
    given = rest.split()
    tail = [w for w in given if SUFFIX.fullmatch(w)]
    return " ".join([w for w in given if not SUFFIX.fullmatch(w)] + last.split() + tail)


def shown(raw):
    """(name as shown, whether the list printed it in capitals)."""
    name = turned(raw)
    caps = any(c.isalpha() for c in name) and name == name.upper()
    return (proper(name) if caps else name), caps


def number(v):
    """(votes, redacted?) from a workbook cell."""
    if v is None or v == "":
        return 0, False
    if isinstance(v, (int, float)):
        return int(v), False
    s = str(v).strip().replace(",", "")
    return (int(s), False) if s.isdigit() else (0, True)


# ---------- the official results files ----------

def fetch_results(folder, say):
    """The three results files, downloaded when missing or older than a month. The results page is read first and each
    link followed by its label, so a corrected address is picked up; the address read on 2026-09-30 is the fallback."""
    paths = {k: os.path.join(folder, f[2]) for k, f in FILES.items()}
    urls = {k: urljoin(RESULTS_PAGE, f[1]) for k, f in FILES.items()}
    if all(os.path.exists(p) and time.time() - os.path.getmtime(p) < 30 * 86400 for p in paths.values()):
        return paths, urls
    try:
        page = net.get(RESULTS_PAGE, accept="text/html").decode("utf-8", "replace")
        found = {}
        for m in re.finditer(r'<a[^>]+href="([^"]+)"[^>]*>(.*?)</a>', page, re.S | re.I):
            found.setdefault(text(m.group(2)), urljoin(RESULTS_PAGE, H.unescape(m.group(1))))
        for k, (label, *_rest) in FILES.items():
            if label in found:
                urls[k] = found[label]
            else:
                say(f"      the Election Results page no longer links \"{label}\"; trying the address read on 2026-09-30")
    except (HTTPError, URLError, OSError) as e:
        say(f"      could not read the Election Results page ({e}); trying the addresses read on 2026-09-30")
    for k, p in paths.items():
        net.download(urls[k], p, max_age_days=30, say=say)
        with open(p, "rb") as fh:
            head = fh.read(8)
        if not head.startswith(FILES[k][3]):
            os.remove(p)      # the site answered with a web page, not the file: leave nothing that could be read as it
            raise SystemExit(f"Kansas: {urls[k]} did not give the file ({FILES[k][0]}); the site answered with something else")
    return paths, urls


def vote_totals(path):
    """{race: [(party, name, votes, printed percent, write_in)]} for the congressional races in the Official Vote Totals."""
    L = lines(path)
    first = " ".join(t for p, _y, t in L if p == 1)
    if "2026 Primary Election" not in first or "Official Vote Totals" not in first:
        raise SystemExit("Kansas: the file is not the 2026 Primary Election Official Vote Totals")
    out, race = {}, None
    for _p, _y, t in L:
        if FURNITURE.match(t):
            continue
        m = LINE.match(t)
        if not m:      # a race's heading; only the five federal races are read
            race = race_of(t)
            if race:
                if race in out:
                    raise SystemExit(f"Kansas: the Official Vote Totals list {t!r} twice")
                out[race] = []
            continue
        if race is None:
            continue
        if m.group("p") not in INITIAL:
            raise SystemExit(f"Kansas: a party initial the loader does not know under {race}: {m.group('p')!r}")
        name, wi = m.group("name"), False
        w = re.fullmatch(r"(.+?) \(Write-in\)", name)
        if w:
            name, wi = w.group(1), True
        out[race].append((INITIAL[m.group("p")], name, int(m.group("votes").replace(",", "")), float(m.group("pct")), wi))
    missing = [r for r in RACES if not out.get(r)]
    if missing:
        raise SystemExit(f"Kansas: the Official Vote Totals give no candidates for {', '.join(missing)}")
    return out


def file_date(path):
    """The creation date in a PDF's own properties (the Official Vote Totals print none on the page)."""
    m = re.search(rb"/CreationDate\s*\(D:(\d{4})(\d\d)(\d\d)", open(path, "rb").read())
    return f"{m.group(1).decode()}-{m.group(2).decode()}-{m.group(3).decode()}" if m else ""


def county_sums(path, known):
    """({(race, party, name key): votes}, rows read, [problems]) from a precinct workbook. known maps each name key in the
    Official Vote Totals for this workbook's office to (race, party)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sums, problems, nrows = defaultdict(int), [], 0
    it = wb.worksheets[0].iter_rows(values_only=True)
    head = [str(c).strip() if c is not None else "" for c in next(it)]
    need = ("Race", "Candidate", "Party", "Votes")
    if not all(k in head for k in need):
        raise SystemExit(f"Kansas: the first sheet of {os.path.basename(path)} has no {', '.join(k for k in need if k not in head)} column")
    ix = {k: head.index(k) for k in need}
    for r in it:
        race = race_of(str(r[ix["Race"]] or ""))
        if not race:
            continue
        votes, redacted = number(r[ix["Votes"]])
        if redacted:
            problems.append(f"{wb.worksheets[0].title}: a precinct's votes for {turned(r[ix['Candidate']])} are not a number")
        nrows += 1
        sums[(race, str(r[ix["Party"]] or "").strip(), key(turned(r[ix["Candidate"]])))] += votes
    for ws in wb.worksheets[1:]:
        g = [list(r) for r in ws.iter_rows(values_only=True)]
        if not g:
            continue
        names = max(g[:6], key=lambda row: sum(1 for c in row if c is not None and key(str(c)) in known))
        total = next((row for row in g if row and row[0] is not None and str(row[0]).strip().upper() == "COUNTY TOTALS"), None)
        if total is None:
            problems.append(f"the {ws.title} sheet has no COUNTY TOTALS row")
            continue
        for j, c in enumerate(names):
            if j == 0 or c is None or not str(c).strip():
                continue
            k = key(str(c))
            if k not in known:
                problems.append(f"the {ws.title} sheet has a column for {str(c).strip()!r}, who is not in the Official Vote Totals")
                continue
            votes, redacted = number(total[j] if j < len(total) else None)
            if redacted:
                problems.append(f"the {ws.title} sheet's COUNTY TOTALS for {str(c).strip()} is not a number")
            nrows += 1
            sums[known[k] + (k,)] += votes
    return sums, nrows, problems


# ---------- the November list, from a page John saved ----------

def saved_pages(folder):
    found = [p for ext in ("*.html", "*.htm", "*.mhtml", "*.mht") for p in glob.glob(os.path.join(folder, ext))]
    return sorted(found, key=os.path.getmtime, reverse=True)


def page_html(path):
    raw = open(path, "rb").read()
    if path.lower().endswith((".mhtml", ".mht")):
        msg = email.message_from_bytes(raw, policy=email.policy.default)
        for part in msg.walk():
            if part.get_content_type() == "text/html":
                return raw, part.get_content()
        raise SystemExit(f"Kansas: {os.path.basename(path)} holds no web page")
    try:
        return raw, raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw, raw.decode("cp1252", "replace")


def list_rows(page, name):
    """(election shown, [federal rows]) from a saved Candidate List page, or None when the page is not the 2026 General
    list. Columns are found from the heading cells; only the allowlisted cells of each row are turned into text."""
    chosen = None
    for sel in re.findall(r"<select\b[^>]*>(.*?)</select>", page, re.S | re.I):
        m = re.search(r"<option\b[^>]*\bselected\b[^>]*>(.*?)</option>", sel, re.S | re.I)
        if m and re.search(r"\b20\d\d\b", text(m.group(1))):
            chosen = text(m.group(1))
            break
    trs = list(re.finditer(r"<tr\b[^>]*>(.*?)</tr>", page, re.S | re.I))
    head_at, head = None, None
    for i, tr in enumerate(trs):
        cells = [text(c) for c in re.findall(r"<t[hd]\b[^>]*>(.*?)</t[hd]>", tr.group(1), re.S | re.I)]
        if all(k in cells for k in KEEP) and (NAME_COL in cells or all(k in cells for k in PARTS)):
            head_at, head = i, cells
            break
    if chosen is not None and chosen != ELECTION:
        return None
    if head is None:
        if chosen == ELECTION:      # the right page, but its table is not laid out as expected: name its headings only
            heads = [text(c) for c in re.findall(r"<th\b[^>]*>(.*?)</th>", page, re.S | re.I)][:40]
            raise SystemExit(f"Kansas: {name} shows \"{ELECTION}\" but no table with the columns {', '.join(KEEP)} and a name; "
                             f"its headings read {heads}")
        return None
    end = page.find("</table>", trs[head_at].end())
    end = len(page) if end < 0 else end
    if chosen is None:      # no list of elections on the page: the election must be named outside the table
        outside = text(page[:trs[head_at].start()] + page[end:])
        if ELECTION not in outside:
            return None
        chosen = ELECTION
    if re.search(r"Page\$(?:\d+|Next|Last)", H.unescape(page)):
        raise SystemExit(f"Kansas: {name} is one page of a list the site splits into several; save it with every row showing")
    ix = {k: head.index(k) for k in (NAME_COL,) + PARTS + KEEP + OPTIONAL if k in head}
    out = []
    for tr in trs[head_at + 1:]:
        if tr.start() > end:
            break
        cells = re.findall(r"<td\b[^>]*>(.*?)</td>", tr.group(1), re.S | re.I)
        if len(cells) != len(head):
            if len(cells) <= 1:      # an empty row or a footer spanning the table
                continue
            raise SystemExit(f"Kansas: a row of {name} has {len(cells)} cells under {len(head)} headings; read the page again")
        rec = {k: text(cells[i]) for k, i in ix.items()}      # the allowlisted cells only
        race = race_of(rec["Office"], rec["District"])
        if not race:
            continue
        full = rec.get(NAME_COL) or " ".join(rec.get(k, "") for k in PARTS if rec.get(k))
        out.append({"race": race, "name": full, "party": rec["Party"], "status": rec.get("Status", ""),
                    "order": rec.get("Ballot Order", "")})
    return chosen, out


def general_list(folder, say):
    """(meta, path of the file it came from) for the November list: from the newest saved page that is the 2026 General
    list (its federal rows kept as JSON), else from the JSON kept at an earlier reading, else (None, None)."""
    jpath = os.path.join(folder, JSON_NAME)
    skipped = []
    for path in saved_pages(folder):
        raw, page = page_html(path)
        got = list_rows(page, os.path.basename(path))
        if got is None:
            skipped.append(os.path.basename(path))
            continue
        shown_as, rows = got
        meta = {"url": LIST_PAGE, "origin": "saved in a browser: " + os.path.basename(path), "election": shown_as,
                "sha256": hashlib.sha256(raw).hexdigest(), "saved": dt.date.fromtimestamp(os.path.getmtime(path)).isoformat(),
                "read": dt.date.today().isoformat(), "rows": rows}
        json.dump(meta, open(jpath, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        if skipped:
            say(f"      saved pages that are not the 2026 General candidate list, not read: {', '.join(skipped)}")
        return meta, path
    if skipped:
        say(f"      saved pages that are not the 2026 General candidate list, not read: {', '.join(skipped)}")
    if os.path.exists(jpath):
        meta = json.load(open(jpath, encoding="utf-8"))
        page = os.path.join(folder, meta["origin"].split(": ", 1)[-1])
        return meta, (page if os.path.exists(page) else jpath)
    return None, None


# ---------- the loader ----------

def load(con, cache, say=print):
    net.patient_lookups()
    folder = os.path.join(cache, "ks")
    os.makedirs(folder, exist_ok=True)
    paths, urls = fetch_results(folder, say)
    totals = vote_totals(paths["totals"])

    # the November ballot
    meta, list_path = general_list(folder, say)
    rows, gone, nominees, names_seen = [], [], defaultdict(list), set()
    if meta:
        order = defaultdict(int)
        for r in meta["rows"]:
            st = r["status"].strip().lower()
            if any(g in st for g in GONE):
                gone.append(r)
                continue
            if st not in ON:
                raise SystemExit(f"Kansas: a candidate for {r['race']} has a status the loader does not know: {r['status']!r}")
            name, caps = shown(r["name"])
            party = r["party"].strip()
            write_in = int(bool(re.search(r"write[- ]?in", party + " " + r["name"], re.I)))
            if write_in:
                name = re.sub(r"\s*\(write[- ]?in\)\s*", " ", name, flags=re.I).strip()
            if (r["race"], name) in names_seen:
                raise SystemExit(f"Kansas: {name} is on the November list twice for {r['race']}")
            names_seen.add((r["race"], name))
            order[r["race"]] += 1
            ballot_order = int(r["order"]) if str(r["order"]).isdigit() else (None if write_in else order[r["race"]])
            note = "; ".join(n for n in (CAPS_NOTE if caps else None, WRITE_IN_NOTE if write_in else None) if n) or None
            if not write_in and party in CODE:
                nominees[(r["race"], party)].append(name)
            rows.append((r["race"], "general", "2026-11-03", name, party, party_code("write-in" if write_in and not party else party),
                         ballot_order, 0, write_in, None, None, None, None, None, "ks-sos-2026-general-list", note))
        doubled = [f"{race} {party}" for (race, party), ks in nominees.items() if len(ks) > 1]
        if doubled:
            raise SystemExit(f"Kansas: two November candidates under one party in {', '.join(doubled)}; read the list again")
        empty = [r for r in RACES if not any(row[0] == r for row in rows)]
        if empty:
            raise SystemExit(f"Kansas: the saved November list gives no candidates for {', '.join(empty)}")
    n = len(rows)

    # the August 4 primary
    fields, upset, printed_off = 0, [], []
    for race, entries in sorted(totals.items()):
        by_party = defaultdict(list)
        for e in entries:
            by_party[e[0]].append(e)
        for party, es in sorted(by_party.items()):
            for _p, name, votes, pct, _wi in es:      # the printed percent is of the party's listed candidates
                party_total = sum(x[2] for x in es)
                if party_total and abs(100 * votes / party_total - pct) > 0.006:
                    printed_off.append(f"{race} {name}")
            if sum(1 for e in es if not e[4]) < 2:
                continue
            fields += 1
            total = sum(e[2] for e in es)
            best = max(e[2] for e in es)
            leaders = [e[1] for e in es if e[2] == best]
            if len(leaders) > 1:
                raise SystemExit(f"Kansas: a tie for first in the {party} primary for {race}; the loader cannot say who advanced")
            off_list = meta is not None and not any(same(leaders[0], x) for x in nominees.get((race, party), []))
            if off_list:
                upset.append(f"{race} {party}")
            for _p, name, votes, _pct, wi in es:
                won = name == leaders[0]
                note = "; ".join(x for x in (WRITE_IN_NOTE if wi else None, NOT_ON_LIST if won and off_list else None) if x) or None
                rows.append((race, f"primary-{CODE[party]}", PRIMARY, name, party, party_code(party), None, 0, int(wi), votes,
                             round(100 * votes / total, 1) if total else None, "advanced" if won else "lost",
                             None, None, "ks-sos-2026-primary-vote-totals", note))

    # the cross-check: county sums against the statewide totals
    checks = {}
    for k, prefix in (("senate", "-S"), ("house", "-H")):
        known = {}
        for race, entries in totals.items():
            if prefix in race:
                for party, name, *_rest in entries:
                    if key(name) in known:
                        raise SystemExit(f"Kansas: two candidates for the same office are both named {name}; the cross-check cannot tell them apart")
                    known[key(name)] = (race, party)
        sums, nrows, problems = county_sums(paths[k], known)
        state = {(race, party, key(name)): votes for race, entries in totals.items() if prefix in race for party, name, votes, *_r in entries}
        called = {key(name): name for race, entries in totals.items() for _party, name, *_r in entries}
        for kk in sorted(set(sums) | set(state)):
            if sums.get(kk, 0) != state.get(kk, 0):
                problems.append(f"{kk[0]} {called.get(kk[2], kk[2])} ({kk[1]}): counties {sums.get(kk, 0):,}, statewide {state.get(kk, 0):,}")
        checks[k] = (nrows, problems)
    problems = [p for _n, ps in checks.values() for p in ps]

    with con:
        con.execute("DELETE FROM candidates WHERE race_id LIKE '2026-KS-%'")
        con.executemany("INSERT INTO candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", rows)
        if meta:
            record_source(con, "ks-sos-2026-general-list", path=list_path, level="federal", state="KS", kind="official candidate list",
                          agency="Kansas Secretary of State, Elections Division",
                          title=f"Candidate List, {meta['election']} (the Candidate List page with \"{ELECTION}\" chosen, saved in a browser "
                                f"on {meta['saved']})", url=LIST_PAGE, rows=len(meta["rows"]),
                          note="The site answers scripts on this page with 403 (CloudFront), so the page was saved in a browser. Office, "
                               "district, party and the candidate's name read (and status and ballot order if the list carries them); the "
                               "address, phone, e-mail and web columns are never read. The list prints no ballot order; its own order is kept. "
                               f"Withdrawn or removed, left off: {len(gone)}.")
        pub = file_date(paths["totals"])
        record_source(con, "ks-sos-2026-primary-vote-totals", path=paths["totals"], level="federal", state="KS", kind="official results",
                      agency="Kansas Secretary of State, Elections Division",
                      title="2026 Primary Election Official Vote Totals (August 4, 2026)", url=urls["totals"],
                      rows=sum(len(v) for v in totals.values()),
                      note="Statewide votes for every candidate for Congress on each party's primary ballot. The totals give no count of "
                           "write-in votes, so shares are of the listed candidates' votes, as the totals print them. The page prints no date"
                           + (f"; the file's own properties date it {pub}" if pub else "") + ". "
                           + ("Every candidate's county sums in the precinct workbooks equal these totals." if not problems
                              else "The precinct workbooks did not agree: " + "; ".join(problems) + "."))
        for k, title in (("senate", "U.S. Senate"), ("house", "U.S. House of Representatives")):
            nrows, ps = checks[k]
            record_source(con, f"ks-sos-2026-primary-precincts-{k}", path=paths[k], level="federal", state="KS", kind="official results",
                          agency="Kansas Secretary of State, Elections Division",
                          title=f"{title} 2026 Primary Election, precinct level results (workbook)", url=urls[k], rows=nrows,
                          note="Used only to check the Official Vote Totals: the first sheet's precinct rows plus the COUNTY TOTALS rows of "
                               "the Johnson, Sedgwick, Shawnee and Wyandotte sheets (some Shawnee precinct cells are printed \"redacted per "
                               "KSA 25-2701(a)(2)\"). " + ("They add up to the statewide totals, candidate by candidate." if not ps
                                                           else "Did not add up: " + "; ".join(ps) + "."))
    if problems:
        say("    Kansas: the precinct workbooks do not add up to the Official Vote Totals: " + "; ".join(problems))
    if printed_off:
        say("    Kansas: the printed percent does not match the votes for " + ", ".join(printed_off))
    if upset:
        say("    Kansas: the primary's leader is not on the November list for " + ", ".join(upset) + "; read the files again")
    ballot = (f"{n} candidates on the November ballot" + (f" ({len(gone)} withdrawn left off)" if gone else "") if meta else
              f"no November list yet (the Candidate List page answers scripts with 403: save it with \"{ELECTION}\" chosen into {folder})")
    say(f"    Kansas: the Senate race and 4 House districts, {ballot}; {fields} party primaries with a field, votes from the "
        f"Official Vote Totals, " + ("checked against the county sums" if not problems else "NOT reconciled with the county sums"))
    return n
