"""
ballot/nd_place_votes.py - how each North Dakota place voted in past partisan general elections, added up from the
Secretary of State's official precinct results, so a page can show the record of a county, a legislative district or
a judicial district without anyone labelling a candidate. The North Dakota twin of ballot/mn_place_votes.py; the
output has the same shape (the Democratic-NPL count is "dem" here, as in Wisconsin's and Iowa's files, where
Minnesota's is "dfl").

    python ballot/nd_place_votes.py               reads (or downloads) the files, writes ballot/lean/nd_place_votes.json
    python ballot/nd_place_votes.py --refresh     asks for everything again even when the cached copies are there
    python ballot/nd_place_votes.py --out x.json  writes somewhere else; --cache DIR keeps the downloads somewhere else
    python ballot/nd_place_votes.py --selftest    the arithmetic and the readers on made-up workbooks; downloads nothing

What it is, and is not
----------------------
For President, U.S. Senator and Governor in 2024, U.S. Senator in 2022, and President and Governor in 2020 (North
Dakota elects its Governor in presidential years, and elected no senator in 2020): the votes for the Democratic-NPL
ticket, the Republican ticket, everyone else together (write-ins included) and the total, for every county, every
judicial district (whole counties) and, for 2022 and 2024, every legislative district. It is how the people of a
place voted then. It is not a prediction, it says nothing about any candidate on a later ballot or about any voter,
and it turns no nonpartisan office into a partisan one (every county office in North Dakota is nonpartisan). No
database is opened for writing; ballot_local_2026.sqlite is opened read-only at the end, only to count how many of
our places are covered.

Where the numbers come from
---------------------------
The Secretary of State's results site, results.sos.nd.gov ("North Dakota Election Officials, County Auditors and
Secretary of State"), which the Secretary's Election Results page (www.sos.nd.gov/elections/election-results) links
for every general election since 2000. Each statewide contest there has an EXPORT button with two choices, and this
loader presses both, as a reader would, one polite request a second:
  - "Precinct": a workbook with one sheet per county and one row per precinct, a column per candidate, and a TOTALS
    row. In 2022 and 2024 a precinct is a six-digit number: the county's number (01 Adams to 53 Williams, in the
    order of the sheets), the legislative district, and the precinct within it (022401 is Barnes County, District 24,
    precinct 01). That is how a precinct is placed in its district, with no map. In 2020 precincts carry names, not
    numbers, so 2020 is added up by county only.
  - "County": a workbook with one row per county, the number of precincts, and each candidate's party under the name
    ("Republican", "Democratic-NPL"). The parties are read from there; the Democratic-NPL and the Republican column
    must each be exactly one, and must name the ticket typed in CONTESTS, or the loader stops.
The files are small (about 75 KB each), kept in states_cache/nd_local/results/. Of the candidates only the two
tickets are kept, to say which election this was; every other column is added into "other" under its party label.

  - Legislative districts are given for 2022 and 2024 only. The 2020 election was held on the districts drawn after
    the 2010 census, and its precincts are not numbered by district. Each year is on the lines in force at that
    election: the loader compares the counties each district reaches in the two years and names the districts where
    they differ (a change inside a county would not show). One district elects one senator and two representatives,
    so "senate" and "house" carry the same figures; House subdistricts (4A, 4B) are not given apart.
  - The district reading of the precinct number is checked against the Secretary's own "All Legislative" export: for
    every district that elected a senator that year, the number of precincts the export counts (county by county
    where it says) must equal the number of precinct numbers carrying that district.
  - Judicial districts are whole counties. Which counties is read from the 2022 general election's judicial
    contests (the "County" export of one contest in each of the eight districts lists exactly the counties that
    voted in it), and compared with the 2024 judicial contests where that export names counties.
  - Not given: cities and townships (a North Dakota precinct is numbered, often takes in several townships and
    cities, and none of their offices is on the November ballot), soil conservation districts (some share a county),
    school districts.

The control
-----------
Nothing is written unless all of this holds: in every county sheet each candidate's precinct rows add up to the
sheet's TOTALS row; every county's sum, and its number of precincts, equals that county's row in the "County"
workbook, for all 53 counties and every contest; the number of precincts equals the number the results page states;
counties, legislative districts and judicial districts each add up to the statewide sum; and the statewide sum equals
the official statewide totals:
  - President and U.S. Senator: the Clerk of the U.S. House of Representatives, "Statistics of the Presidential and
    Congressional Election" / "Statistics of the Congressional Election" (clerk.house.gov; compiled from official
    sources), the North Dakota page, read from the PDF at each run with ballot/pdftext.py.
  - Governor: the TOTALS row of the Secretary's own "County" workbook. That is a second report of the same results
    system, not an independent document, and the file says so; no canvass document a script may read was found.
The figures this loader was checked against on 2026-10-02 are typed below (CHECKED); they are used, and the file
says so, when a document cannot be read again.

Found on 2026-10-02: the workbooks of 2022 and 2020 say "Unofficial General Election Results" in their first line
although the pages they are exported from are headed "Official 2022 General Election Results" (last updated
11/29/2022) and likewise for 2020; the file records both. The site's bulk "Exports" page gives county rows only for
the newest election there (2024) and statewide rows for the older ones, and never precincts, which is why the
per-contest buttons are used. The newer results site (resultsnd.sos.nd.gov) holds 2026 only.
"""

import argparse
import collections
import csv
import datetime as dt
import html
import http.cookiejar
import io
import json
import os
import pathlib
import re
import sqlite3
import sys
import time
import urllib.parse
import urllib.request

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from states import net  # noqa: E402
from ballot import mn_place_votes as M  # noqa: E402  the shared small things: caching, fingerprints, the file writer

Stop = M.Stop
METHOD = "1.0"                 # change how anything is added up or matched, and this goes up
CACHE = os.path.join(HERE, "states_cache", "nd_local")
OUT = os.path.join(HERE, "ballot", "lean", "nd_place_votes.json")
DB = os.path.join(HERE, "ballot_local_2026.sqlite")
STATE_FIPS = "38"
FEW = 20
SIDES = ("dem", "rep", "other", "total")
KEEP_DAYS = 3650               # certified results of past elections do not change; --refresh asks again

AGENCY = "North Dakota Secretary of State (results reported by the county auditors and the Secretary of State)"
SITE = "https://results.sos.nd.gov/"
ARCHIVE = "https://www.sos.nd.gov/elections/election-results"
# The elections, as the Secretary's Election Results page links them (the newest on this site has no id in its link).
ELECTIONS = {
    2024: {"eid": "", "date": "2024-11-05", "site_date": "11/5/2024", "link": "General Election - November 5, 2024"},
    2022: {"eid": "vxUYQ0lrpP4.", "date": "2022-11-08", "site_date": "11/8/2022", "link": "General Election - November 8, 2022"},
    2020: {"eid": "StuhWbgeuSk.", "date": "2020-11-03", "site_date": "11/3/2020", "link": "General Election - November 3, 2020"},
}
DISTRICT_YEARS = (2022, 2024)  # the elections whose precincts are numbered county, district, precinct
JUDICIAL_YEAR = 2022           # the election whose judicial contests name the counties of all eight districts
DISTRICTS = 47

# North Dakota's 53 counties in the order of their FIPS codes (38001, 38003, ... 38105) and of the workbooks' sheets.
COUNTIES = ("Adams|Barnes|Benson|Billings|Bottineau|Bowman|Burke|Burleigh|Cass|Cavalier|Dickey|Divide|Dunn|Eddy|Emmons|Foster|Golden Valley|"
            "Grand Forks|Grant|Griggs|Hettinger|Kidder|LaMoure|Logan|McHenry|McIntosh|McKenzie|McLean|Mercer|Morton|Mountrail|Nelson|Oliver|"
            "Pembina|Pierce|Ramsey|Ransom|Renville|Richland|Rolette|Sargent|Sheridan|Sioux|Slope|Stark|Steele|Stutsman|Towner|Traill|Walsh|Ward|"
            "Wells|Williams").split("|")
FIPS = {name: f"{STATE_FIPS}{2 * i + 1:03d}" for i, name in enumerate(COUNTIES)}
JUDICIAL = {"East Central": "EC", "North Central": "NC", "Northeast": "NE", "Northeast Central": "NEC", "Northwest": "NW", "South Central": "SC",
            "Southeast": "SE", "Southwest": "SW"}
DEM_PARTY, REP_PARTY = "Democratic-NPL", "Republican"

# The contests. "heading" is the contest's heading on the results page; "dem" and "rep" are surnames the Secretary's
# own candidate column must carry under that party, or the loader stops.
CONTESTS = [
    {"id": "2024-president", "year": 2024, "office": "President and Vice President of the United States",
     "heading": "President & Vice-President of the United States", "dem": ("Harris", "Walz"), "rep": ("Trump", "Vance"),
     "official": "clerk-statistics-2024", "section": "FOR PRESIDENTIAL ELECTORS"},
    {"id": "2024-us-senate", "year": 2024, "office": "United States Senator", "heading": "United States Senator",
     "dem": ("Christiansen",), "rep": ("Cramer",), "official": "clerk-statistics-2024", "section": "FOR UNITED STATES SENATOR"},
    {"id": "2024-governor", "year": 2024, "office": "Governor and Lieutenant Governor", "heading": "Governor and Lt. Governor",
     "dem": ("Piepkorn", "Hart"), "rep": ("Armstrong", "Strinden"), "official": "nd-sos-county-results", "section": None},
    {"id": "2022-us-senate", "year": 2022, "office": "United States Senator", "heading": "United States Senator",
     "dem": ("Christiansen",), "rep": ("Hoeven",), "official": "clerk-statistics-2022", "section": "FOR UNITED STATES SENATOR"},
    {"id": "2020-president", "year": 2020, "office": "President and Vice President of the United States",
     "heading": "President & Vice-President of the United States", "dem": ("Biden", "Harris"), "rep": ("Trump", "Pence"),
     "official": "clerk-statistics-2020", "section": "FOR PRESIDENTIAL ELECTORS"},
    {"id": "2020-governor", "year": 2020, "office": "Governor and Lieutenant Governor", "heading": "Governor and Lt. Governor",
     "dem": ("Lenz", "Vig"), "rep": ("Burgum", "Sanford"), "official": "nd-sos-county-results", "section": None},
]

# The official statewide totals this loader was checked against on 2026-10-02. They stand in, and the file says so,
# only when a document cannot be read again.
CHECKED = {
    "2024-president": {"dem": 112327, "rep": 246505, "other": 9323, "total": 368155},
    "2024-us-senate": {"dem": 121602, "rep": 241569, "other": 1156, "total": 364327},
    "2022-us-senate": {"dem": 59995, "rep": 135474, "other": 44671, "total": 240140},
    "2020-president": {"dem": 114902, "rep": 235595, "other": 11322, "total": 361819},
}
OFFICIAL_DOCS = {
    "clerk-statistics-2024": {
        "kind": "official federal compilation", "agency": "Office of the Clerk, U.S. House of Representatives",
        "title": "Statistics of the Presidential and Congressional Election from Official Sources for the Election of November 5, 2024",
        "url": "https://clerk.house.gov/member_info/electionInfo/2024/statistics2024.pdf", "file": "clerk_statistics2024.pdf"},
    "clerk-statistics-2022": {
        "kind": "official federal compilation", "agency": "Office of the Clerk, U.S. House of Representatives",
        "title": "Statistics of the Congressional Election from Official Sources for the Election of November 8, 2022",
        "url": "https://clerk.house.gov/member_info/electionInfo/2022/statistics2022.pdf", "file": "clerk_statistics2022.pdf"},
    "clerk-statistics-2020": {
        "kind": "official federal compilation", "agency": "Office of the Clerk, U.S. House of Representatives",
        "title": "Statistics of the Presidential and Congressional Election from Official Sources for the Election of November 3, 2020",
        "url": "https://clerk.house.gov/member_info/electionInfo/2020/statistics2020.pdf", "file": "clerk_statistics2020.pdf"},
}
COUNTY_REPORT = {"id": "nd-sos-county-results", "kind": "the Secretary of State's own county report (a second report of the same results system)",
                 "agency": AGENCY, "title": "Statewide County Results: the \"County\" export of each contest on results.sos.nd.gov", "url": SITE,
                 "why": "For Governor no canvass document a script may read was found, so the statewide control is the TOTALS row of the "
                        "Secretary's county report. It is not an independent document. For every contest the same report is compared county "
                        "by county, all 53."}

# The kinds of place: (kind, first election year given, what it is, what its key is)
KINDS = [
    ("county", 2020, "Counties", "five-digit county FIPS code, as sl_places kind county and the jurisdiction_id of our county races"),
    ("senate", 2022, "Legislative districts (each elects one senator)", "district number, as the jurisdiction_id of our Senate races"),
    ("house", 2022, "Legislative districts (each elects two representatives)", "district number, as the jurisdiction_id of our House races"),
    ("judicial", 2020, "Judicial districts, as whole counties", "ND- and the district's letters (ND-EC), as sl_places kind judicial_district"),
]
KIND_NAMES = [k[0] for k in KINDS]
REDRAWN = ("The 2020 election was held on the districts drawn after the 2010 census, and its precincts are not numbered by district; the "
           "districts drawn after the 2020 census were first used in 2022, so 2020 is not given.")

WHAT = ("How each North Dakota place voted in six contests of the 2020, 2022 and 2024 general elections: the votes for the Democratic-NPL "
        "ticket, the Republican ticket, everyone else together (write-ins included) and the total, added up from the North Dakota Secretary "
        "of State's official precinct results.")
NOTE = ("What this is: how the people of a place voted in that election, in the precinct results the county auditors and the Secretary of "
        "State publish, added up here by county, by judicial district and, for 2022 and 2024, by legislative district, on the precinct and "
        "district lines in force at that election. What this is not: it is not a prediction of any election; it says nothing about any "
        "candidate on a later ballot or about any voter; a nonpartisan office stays nonpartisan (every county office in North Dakota is one); "
        "and a place is not its lines for ever: where a district was drawn again, the figures are for the lines of that year. The tickets are "
        "named, as the Secretary's own results name them, only to say which election this was. In a place with very few voters the split "
        "would come close to saying how particular people voted; those contests are listed in the place's too_few, and a page should leave "
        "the split out.")
HOW_TO_READ = ("Each place's votes are four counts for each contest: dem (the Democratic-NPL ticket), rep (the Republican ticket), other (every "
               "other candidate and all write-ins) and total. Minnesota's file calls the first count dfl. A contest a place does not have was "
               "not counted on its lines: legislative districts are given for 2022 and 2024 only.")
TOO_FEW = (f"A place's too_few lists the contests in which it had fewer than {FEW} votes, or in which every vote went the same way. There the "
           "split comes close to saying how particular people voted, so a page should leave it out. The counts are the official record all the "
           "same, and are kept here so that the sums can be checked.")
NOT_GIVEN = ("Cities and townships: a North Dakota precinct is numbered, often takes in several townships and cities, and the results do not "
             "say which. Soil conservation districts: some share a county, and the results do not say which precincts are in which. School "
             "districts and county commissioner districts: the precinct results do not name them. House subdistricts (4A, 4B): the precinct "
             "number carries the district only. Legislative districts for 2020: other lines, and precincts not numbered by district.")


def _clean(s):
    return re.sub(r"\s+", " ", str(s if s is not None else "")).strip()


# ---------------------------------------------------------------- the results site

class Site:
    """One visit to results.sos.nd.gov: the pages and their EXPORT buttons, one request a second, with the kit's honest
    User-Agent and the cookies the site hands out."""

    def __init__(self):
        net.patient_lookups()
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))

    def _open(self, url, data=None):
        headers = {"User-Agent": net.UA, "Accept": "*/*"}
        if data is not None:
            headers["Content-Type"] = "application/x-www-form-urlencoded"
        last = None
        for attempt in range(3):
            time.sleep(1.0 if attempt == 0 else 6.0)
            try:
                with self.opener.open(urllib.request.Request(url, data=data, headers=headers), timeout=180) as r:
                    return r.headers, r.read()
            except OSError as e:
                last = e
        raise OSError(f"{url.split('?')[0]}: {last}")

    def page(self, path, eid):
        url = SITE + path + (("&" if "?" in path else "?") + "eid=" + eid if eid else "")
        return url, self._open(url)[1].decode("utf-8", "replace")

    def press(self, url, page, target=None, fields=None):
        """Posts the page's own form back: its hidden fields, and either the button pressed (a __doPostBack target) or
        the fields a submit button sends."""
        form = {html.unescape(n): html.unescape(v) for n, v in re.findall(r'<input type="hidden" name="([^"]+)" id="[^"]*" value="([^"]*)"', page)}
        if target:
            form["__EVENTTARGET"], form["__EVENTARGUMENT"] = target, ""
        form.update(fields or {})
        headers, body = self._open(url, urllib.parse.urlencode(form).encode("utf-8"))
        return (headers.get("Content-Disposition") or ""), body


def page_facts(page):
    """What a results page says of its election: the heading, the date, when it was last updated, how many precincts."""
    find = lambda pat: _clean(html.unescape(re.sub(r"<[^>]+>", " ", (re.search(pat, page, re.S) or [None, ""])[1])))   # noqa: E731
    return {"heading": find(r'<div class="col header-box1">\s*<h1>(.*?)</h1>'), "date": find(r'id="hidElectionDate" value="([^"]*)"'),
            "updated": find(r'<div class="last-updated">\s*Results last updated:(.*?)</div>'),
            "precincts": find(r"Total Precincts<span>(\d+)</span>")}


def page_contests(page):
    """[(heading, the Precinct button, the County button)] in page order."""
    heads = [(m.start(), _clean(html.unescape(re.sub(r"<[^>]+>", " ", m.group(1)))))
             for m in re.finditer(r'<div class="display-results-box-a">\s*<h1>(.*?)</h1>', page, re.S)]
    buttons = [(m.start(), m.group(1), m.group(2)) for m in
               re.finditer(r'name="(ctl00\$MainContent\$rptRace\$ctl\d+\$ctl0\d)"[^>]*value="(Precinct|County)"', page)]
    out = []
    for i, (pos, name) in enumerate(heads):
        end = heads[i + 1][0] if i + 1 < len(heads) else len(page)
        mine = {kind: b for p, b, kind in buttons if pos < p < end}
        out.append((name, mine.get("Precinct"), mine.get("County")))
    return out


def _workbook(body, disposition, what):
    if not body.startswith(b"PK"):
        raise OSError(f"{what}: the site answered with something that is not a workbook ({disposition or 'no file'})")
    return body


def _write(path, data):
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    with open(path + ".part", "wb") as fh:
        fh.write(data)
    os.replace(path + ".part", path)


def book_path(cache, cid, kind):
    return os.path.join(cache, "results", f"{cid}_{kind}.xlsx")


def fetch_year(year, contests, cache, refresh, say, site):
    """The Precinct and County workbooks of one election's contests, and what the page says of the election. Kept on
    disk; nothing is asked for when everything is there. Returns the page facts."""
    mine = [c for c in contests if c["year"] == year]
    facts_path = os.path.join(cache, "results", f"{year}_page.json")
    paths = [book_path(cache, c["id"], k) for c in mine for k in ("precinct", "county")]
    have = all(os.path.exists(p) and os.path.getsize(p) > 0 for p in paths + [facts_path])
    if have and all(M._fresh(p, refresh, KEEP_DAYS) for p in paths + [facts_path]):
        return M._load(facts_path)
    e = ELECTIONS[year]
    try:
        url, page = site().page("ResultsSW.aspx?text=All&type=SW&map=CTY", e["eid"])
        facts = dict(page_facts(page), url=url, fetched=M._now())
        if facts["date"] != e["site_date"]:
            raise OSError(f"the page is the election of {facts['date'] or 'no date'}, not {e['site_date']}")
        listed = page_contests(page)
        for c in mine:
            fits = [x for x in listed if x[0] == c["heading"]]
            if len(fits) != 1 or not fits[0][1] or not fits[0][2]:
                raise OSError(f"{len(fits)} contests headed '{c['heading']}' with both export buttons, not one")
            for kind, target in (("precinct", fits[0][1]), ("county", fits[0][2])):
                disposition, body = site().press(url, page, target=target)
                _write(book_path(cache, c["id"], kind), _workbook(body, disposition, f"{c['id']} {kind} export"))
            say(f"      {c['id']}: the precinct and county workbooks fetched")
        M._save(facts_path, facts)
        return facts
    except (OSError, ValueError) as err:
        if have:
            say(f"      {year} results: could not be fetched again ({err}); using the copies of {M._day(paths[0])}")
            return M._load(facts_path)
        raise Stop(f"    {year} results: could not be fetched ({err}) and no complete copy is on disk. Wait a few minutes and run this again.")


def fetch_legislative(year, cache, refresh, say, site):
    """What the Secretary's "All Legislative" export says of each legislative contest: the contest's kind, the district,
    the number of precincts, and the county where the export names one. Candidates' names are not kept. None when it
    cannot be read and no extract is on disk."""
    path = os.path.join(cache, "results", f"{year}_legislative_precinct_counts.json")
    if M._fresh(path, refresh, KEEP_DAYS):
        return M._load(path)
    try:
        url, page = site().page("ResultsExport.aspx", ELECTIONS[year]["eid"])
        if page_facts(page)["date"] != ELECTIONS[year]["site_date"]:
            raise OSError("the export page is another election's")
        _disp, body = site().press(url, page, fields={"ctl00$MainContent$rblTypes": "1", "ctl00$MainContent$hidCountyID": "",
                                                      "ctl00$MainContent$btnAllLegislative": "All Legislative"})
        rows = legislative_rows(body.decode("utf-8-sig", "replace"))
        if not rows:
            raise OSError("the export has no legislative rows")
        doc = {"url": url, "export": "All Legislative, CSV", "fetched": M._now(), "sha256": M._sha(body), "rows": rows,
               "kept": "the contest's kind, the district, the precinct count and the county; no candidate's name"}
        M._save(path, doc)
        return doc
    except (OSError, ValueError) as err:
        if os.path.exists(path) and os.path.getsize(path) > 0:
            say(f"      {year} legislative export: could not be fetched again ({err}); using the extract of {M._day(path)}")
            return M._load(path)
        say(f"      {year} legislative export: could not be fetched ({err}); the district reading of the precinct numbers is not checked for {year}")
        return None


def legislative_rows(text):
    """[[chamber, district, precincts, county or None]] from the export, one for each contest, district and county."""
    out = set()
    for r in list(csv.reader(io.StringIO(text)))[1:]:
        if len(r) < 8:
            continue
        chamber = "senate" if r[0].startswith("State Senator") else "house" if r[0].startswith("State Representative") else None
        m, n = re.match(r"District 0*(\d+[A-Za-z]?)$", r[2].strip()), re.match(r"\d+/(\d+)$", r[7].strip())
        if chamber and m and n:
            out.add((chamber, m.group(1).upper(), int(n.group(1)), (r[8].strip() or None) if len(r) > 8 else None))
    return [list(x) for x in sorted(out, key=lambda x: (x[0], M.sort_key(x[1]), x[3] or ""))]


def fetch_judicial(cache, refresh, say, site):
    """{district name: [county names]}: the counties that voted in a judicial contest of each district in the 2022
    general election, from that contest's "County" export. Only the county names are read (never the candidates')."""
    path = os.path.join(cache, "results", f"{JUDICIAL_YEAR}_judicial_district_counties.json")
    if M._fresh(path, refresh, KEEP_DAYS):
        return M._load(path)
    try:
        url, page = site().page("ResultsSW.aspx?text=Race&type=JD&map=CTY&area=District", ELECTIONS[JUDICIAL_YEAR]["eid"])
        if page_facts(page)["date"] != ELECTIONS[JUDICIAL_YEAR]["site_date"]:
            raise OSError("the judicial page is another election's")
        districts, contests = {}, {}
        for heading, _precinct, county in page_contests(page):
            m = re.match(r"(Judge of the District Court No\. \d+)(?: Unexpired \d+-Year Term)? (.+)$", heading)
            if m and county and m.group(2) in JUDICIAL and m.group(2) not in districts:
                _disp, body = site().press(url, page, target=county)
                book = read_county_book(_workbook(body, _disp, f"judicial district {m.group(2)}"), names=False)
                districts[m.group(2)] = sorted(book["rows"])
                contests[m.group(2)] = m.group(1)
        doc = {"url": url, "election": ELECTIONS[JUDICIAL_YEAR]["link"], "fetched": M._now(), "districts": districts, "contests": contests,
               "kept": "for one contest of each district, the names of the counties in its County export; no candidate's name"}
        problem = judicial_problem(districts)
        if problem:
            raise OSError(problem)
        M._save(path, doc)
        say(f"      judicial districts: the counties of {len(districts)} districts read from the {JUDICIAL_YEAR} judicial contests")
        return doc
    except (OSError, ValueError) as err:
        if os.path.exists(path) and os.path.getsize(path) > 0:
            say(f"      judicial districts: could not be fetched again ({err}); using the extract of {M._day(path)}")
            return M._load(path)
        raise Stop(f"    judicial districts: could not be read ({err}) and no extract is on disk. Wait a few minutes and run this again.")


def judicial_problem(districts):
    seen = [c for cs in districts.values() for c in cs]
    if set(districts) != set(JUDICIAL):
        return f"the judicial contests name the districts {sorted(districts)}, not the eight"
    if sorted(seen) != sorted(COUNTIES):
        return "the judicial districts' counties are not the 53 counties, each once"
    return None


def fetch_judicial_check(year, cache, refresh, say, site):
    """{district: [counties]} as the "All Judicial" export of another election names them, where it names counties."""
    path = os.path.join(cache, "results", f"{year}_judicial_district_counties_check.json")
    if M._fresh(path, refresh, KEEP_DAYS):
        return M._load(path)
    try:
        url, page = site().page("ResultsExport.aspx", ELECTIONS[year]["eid"])
        _disp, body = site().press(url, page, fields={"ctl00$MainContent$rblTypes": "1", "ctl00$MainContent$hidCountyID": "",
                                                      "ctl00$MainContent$btnAllJudicial": "All Judicial"})
        districts = collections.defaultdict(set)
        for r in list(csv.reader(io.StringIO(body.decode("utf-8-sig", "replace"))))[1:]:
            if len(r) > 8 and r[2].strip() in JUDICIAL and r[8].strip():
                districts[r[2].strip()].add(r[8].strip())
        doc = {"url": url, "export": "All Judicial, CSV", "fetched": M._now(), "districts": {d: sorted(cs) for d, cs in sorted(districts.items())},
               "kept": "the district and the county of each row; no candidate's name"}
        M._save(path, doc)
        return doc
    except (OSError, ValueError) as err:
        if os.path.exists(path) and os.path.getsize(path) > 0:
            return M._load(path)
        say(f"      {year} judicial export: could not be fetched ({err}); the districts' counties are not compared with {year}")
        return None


# ---------------------------------------------------------------- the workbooks

def _sheets(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for name in wb.sheetnames:
            yield name, [list(r) for r in wb[name].iter_rows(values_only=True)]
    finally:
        wb.close()


def _votes(cells, what):
    out = []
    for v in cells:
        if isinstance(v, bool) or not isinstance(v, (int, float)) or v != int(v) or v < 0:
            raise Stop(f"    {what}: {v!r} is not a count of votes; stopping")
        out.append(int(v))
    return out


def _table(rows, label, what):
    """The header row whose second cell is `label`, and the rows under it down to TOTALS."""
    at = next((i for i, r in enumerate(rows) if len(r) > 1 and _clean(r[1]) == label), None)
    if at is None:
        raise Stop(f"    {what}: no row headed '{label}'; stopping")
    body, totals = [], None
    for r in rows[at + 1:]:
        if len(r) < 2 or r[1] is None or _clean(r[1]) == "":
            continue
        if _clean(r[1]) == "TOTALS":
            totals = r
            break
        body.append(r)
    if totals is None:
        raise Stop(f"    {what}: no TOTALS row; stopping")
    return rows[at], body, totals


def read_precinct_book(data, what):
    """{"says", "candidates": [names], "counties": {county: [(precinct, [votes])]}} from a "Precinct" workbook. Every
    sheet's precinct rows must add up to its TOTALS row."""
    out = {"says": None, "candidates": None, "counties": {}}
    for sheet, rows in _sheets(data):
        head, body, totals = _table(rows, "Precinct", f"{what}, sheet {sheet}")
        names = [_clean(x) for x in head[2:] if x is not None and _clean(x)]
        if out["candidates"] is None:
            out["says"], out["candidates"] = _clean(rows[0][0]), names
        elif names != out["candidates"]:
            raise Stop(f"    {what}, sheet {sheet}: its candidate columns are not those of the first sheet; stopping")
        n = len(names)
        precincts = [(_clean(r[1]), _votes(r[2:2 + n], f"{what}, sheet {sheet}, precinct {_clean(r[1])}")) for r in body]
        if len({p for p, _v in precincts}) != len(precincts):
            raise Stop(f"    {what}, sheet {sheet}: a precinct appears twice; stopping")
        sums = [sum(v[i] for _p, v in precincts) for i in range(n)]
        if sums != _votes(totals[2:2 + n], f"{what}, sheet {sheet}, TOTALS"):
            raise Stop(f"    {what}, sheet {sheet}: the precincts add up to {sums} and the sheet's TOTALS row is {totals[2:2 + n]}; stopping")
        out["counties"][sheet] = precincts
    return out


def read_county_book(data, what="county workbook", names=True):
    """{"says", "candidates": [(name, party)], "rows": {county: (precincts, [votes])}, "totals": (precincts, [votes])}
    from a "County" workbook. With names=False only the county names and precinct counts are read."""
    sheets = list(_sheets(data))
    if len(sheets) != 1:
        raise Stop(f"    {what}: {len(sheets)} sheets, not one; stopping")
    rows = sheets[0][1]
    head, body, totals = _table(rows, "County", what)
    out = {"says": _clean(rows[0][0]), "candidates": [], "rows": {}, "totals": None}
    if not names:
        out["rows"] = {_clean(r[1]): (int(r[2]), []) for r in body}
        return out
    for cell in head[3:]:
        if cell is None or not _clean(cell):
            continue
        parts = [p.strip() for p in str(cell).split("\n")]
        out["candidates"].append((_clean(parts[0]), _clean(" ".join(parts[1:]))))
    n = len(out["candidates"])
    for r in body:
        out["rows"][_clean(r[1])] = (int(r[2]), _votes(r[3:3 + n], f"{what}, {_clean(r[1])}"))
    out["totals"] = (int(totals[2]), _votes(totals[3:3 + n], f"{what}, TOTALS"))
    return out


def precinct_place(code, county, year):
    """A 2022 or 2024 precinct number -> its legislative district; stops if the number is not county, district, precinct."""
    if not re.fullmatch(r"\d{6}", code) or int(code[:2]) != COUNTIES.index(county) + 1 or not 1 <= int(code[2:4]) <= DISTRICTS:
        raise Stop(f"    {year}, {county} County: precinct '{code}' is not the county's number, a district from 1 to {DISTRICTS} and a precinct; stopping")
    return str(int(code[2:4]))


def contest_table(c, pbook, cbook):
    """One contest, precinct by precinct: {"columns": ..., "rows": [(county name, precinct, [dem, rep, other, total])]}.
    The parties come from the County workbook; the two tickets must be the ones typed in CONTESTS."""
    cands = cbook["candidates"]
    if [n for n, _p in cands] != pbook["candidates"]:
        raise Stop(f"    {c['id']}: the Precinct and the County workbook do not have the same candidate columns; stopping")
    sides = {}
    for side, party in (("dem", DEM_PARTY), ("rep", REP_PARTY)):
        at = [i for i, (_n, p) in enumerate(cands) if p == party]
        if len(at) != 1:
            raise Stop(f"    {c['id']}: {len(at)} columns under {party}, not one; stopping")
        name = cands[at[0]][0]
        if not all(re.search(rf"\b{re.escape(s)}\b", name) for s in c[side]):
            raise Stop(f"    {c['id']}: the {party} column is not the ticket this loader was checked against ({' and '.join(c[side])}); stopping")
        sides[side] = at[0]
    if sorted(pbook["counties"]) != sorted(COUNTIES):
        raise Stop(f"    {c['id']}: the Precinct workbook's sheets are not the 53 counties; stopping")
    rows = []
    for county in COUNTIES:
        for precinct, v in pbook["counties"][county]:
            total = sum(v)
            rows.append((county, precinct, [v[sides["dem"]], v[sides["rep"]], total - v[sides["dem"]] - v[sides["rep"]], total]))
    pack = lambda v: [v[sides["dem"]], v[sides["rep"]], sum(v) - v[sides["dem"]] - v[sides["rep"]], sum(v)]              # noqa: E731
    return {"rows": rows,
            "tickets": {side: cands[i][0] for side, i in sides.items()},
            "other_parties": [p or "write-in" if not n.lower().startswith("write-in") else "write-in"
                              for i, (n, p) in enumerate(cands) if i not in sides.values()],
            "says": {"precinct": pbook["says"], "county": cbook["says"]},
            "county_report": {county: {"precincts": n, "votes": pack(v)} for county, (n, v) in cbook["rows"].items()},
            "county_report_totals": {"precincts": cbook["totals"][0], "votes": pack(cbook["totals"][1])}}


# ---------------------------------------------------------------- the official statewide totals

def read_clerk(path, state="NORTH DAKOTA"):
    """The state's page of the Clerk of the House's election statistics: {section heading: [(label, votes)]} for the
    presidential electors (by party) and United States Senator (candidate, party), up to the representatives."""
    from ballot import pdftext
    lines = pdftext.lines(path)
    start = next((i for i, (_p, _y, t) in enumerate(lines) if t.strip() == state), None)
    if start is None:
        raise ValueError(f"no {state} heading")
    page = lines[start][0]
    top = next(t.strip() for p, _y, t in lines if p == page)
    sections, section = {}, None
    for _p, _y, t in lines[start + 1:]:
        t = t.strip()
        if t.startswith("FOR UNITED STATES REPRESENTATIVE") or (t.isupper() and not t.startswith("FOR ") and section):
            break
        if t.startswith("FOR "):
            section = t
            sections[section] = []
            continue
        m = re.match(r"^(.*?)\s*\.{3,}\s*([\d,]+)$", t)
        if section and m:
            sections[section].append((m.group(1).strip(), int(m.group(2).replace(",", ""))))
        elif section and t:
            raise ValueError(f"a line under {section} is not a name, dots and a number")
    return {"where": f"North Dakota, page {top}" if top.isdigit() else f"North Dakota, page {page} of the PDF", "sections": sections}


def clerk_official(doc, section):
    lines = doc["sections"].get(section) or []
    party = (lambda s: s) if section == "FOR PRESIDENTIAL ELECTORS" else (lambda s: s.rsplit(",", 1)[-1].strip())   # one line a party, or "Name, Party"
    dem = [v for label, v in lines if party(label).startswith("Democratic")]
    rep = [v for label, v in lines if party(label) == "Republican"]
    if len(dem) != 1 or len(rep) != 1:
        raise ValueError(f"{section}: {len(dem)} Democratic and {len(rep)} Republican lines, not one of each")
    total = sum(v for _label, v in lines)
    return {"dem": dem[0], "rep": rep[0], "other": total - dem[0] - rep[0], "total": total}


def official_totals(contests, tables, cache, refresh, say):
    """({contest id: {dem, rep, other, total, source, where, read}}, the documents' records)."""
    docs, out = {}, {}
    for sid, src in OFFICIAL_DOCS.items():
        if not any(c["official"] == sid for c in contests):
            continue
        path = os.path.join(cache, src["file"])
        rec = {"id": sid, "kind": src["kind"], "agency": src["agency"], "title": src["title"], "url": src["url"]}
        try:
            net.download(src["url"], path, 0 if refresh else KEEP_DAYS, tries=3, say=say)
            parsed = read_clerk(path)
            rec.update(fetched=M._day(path), sha256=M._sha_file(path), where=parsed["where"])
            docs[sid] = (rec, parsed)
        except Exception as e:  # noqa: BLE001  the control then runs on the typed figures, and the file says so
            say(f"      {src['title'][:60]}: could not be read again ({e}); the control uses the figures typed in on 2026-10-02")
            rec.update(unread=f"could not be read on {M._now()}: {e}")
            docs[sid] = (rec, None)
    for c in contests:
        if c["official"] == COUNTY_REPORT["id"]:
            out[c["id"]] = dict(zip(SIDES, tables[c["id"]]["county_report_totals"]["votes"]), source=COUNTY_REPORT["id"],
                                where="the TOTALS row of the contest's County workbook", read=f"from the workbook kept in {os.path.basename(cache)}")
            continue
        rec, parsed = docs[c["official"]]
        typed, got, read = CHECKED.get(c["id"]), None, None
        if parsed:
            try:
                got, read = clerk_official(parsed, c["section"]), f"from the document on {rec.get('fetched')}"
            except Exception as e:  # noqa: BLE001
                say(f"      {rec['title'][:60]}: {e}; the control uses the figures typed in on 2026-10-02")
        if got is None:
            if typed is None:
                raise Stop(f"    {c['id']}: no official total could be read and none is typed in this loader")
            got, read = dict(typed), "typed into the loader from the document on 2026-10-02; the document could not be read again today"
        elif typed is not None and got != typed:
            say(f"      {c['id']}: the document now reads {got}; this loader was checked against {typed}")
        out[c["id"]] = dict(got, source=c["official"], where=rec.get("where", ""), read=read)
    used = [rec for rec, _p in docs.values()]
    if any(c["official"] == COUNTY_REPORT["id"] for c in contests):
        used.append(dict(COUNTY_REPORT))
    return out, used


# ---------------------------------------------------------------- adding up

def tally(tables, judicial, contests):
    """Everything the workbooks say, added up: (places, statewide sums, precinct counts by year, districts' counties by year)."""
    votes = {k: collections.defaultdict(dict) for k in KIND_NAMES}
    state, precincts, reach = {}, {}, collections.defaultdict(lambda: collections.defaultdict(set))
    district_of = {county: d for d, cs in judicial.items() for county in cs}
    codes = collections.defaultdict(collections.Counter)                   # year -> (district, county) -> precincts
    for c in contests:
        year, rows = c["year"], tables[c["id"]]["rows"]
        keys = [(county, p) for county, p, _v in rows]
        if precincts.setdefault(year, keys) != keys:
            raise Stop(f"    {c['id']}: its precincts are not those of the other {year} contests; stopping")
        s = state.setdefault(c["id"], [0, 0, 0, 0])
        first = c is next(x for x in contests if x["year"] == year)
        for county, precinct, v in rows:
            where = {"county": FIPS[county], "judicial": f"ND-{JUDICIAL[district_of[county]]}"}
            if year in DISTRICT_YEARS:
                d = precinct_place(precinct, county, year)
                where["senate"] = where["house"] = d
                reach[year][d].add(FIPS[county])
                if first:
                    codes[year][(d, county)] += 1
            for i in range(4):
                s[i] += v[i]
            for kind, key in where.items():
                cur = votes[kind][key].setdefault(c["id"], [0, 0, 0, 0])
                for i in range(4):
                    cur[i] += v[i]
    order = [c["id"] for c in contests]
    pack = lambda d: {cid: dict(zip(SIDES, d[cid])) for cid in order if cid in d}                                      # noqa: E731

    def few(rec):
        hold = [cid for cid, v in rec["votes"].items() if v["total"] < FEW or max(v["dem"], v["rep"], v["other"]) == v["total"]]
        if hold:
            rec["too_few"] = hold
        return rec
    places = {k: {} for k in KIND_NAMES}
    for county in COUNTIES:
        if FIPS[county] in votes["county"]:
            places["county"][FIPS[county]] = few({"name": f"{county} County", "votes": pack(votes["county"][FIPS[county]])})
    years = sorted(reach)
    for kind in ("senate", "house"):
        for d in sorted(votes[kind], key=M.sort_key):
            rec = few({"name": f"Legislative District {d}", "votes": pack(votes[kind][d]),
                       "counties": {str(y): sorted(reach[y][d]) for y in years if d in reach[y]}})
            if len({tuple(x) for x in rec["counties"].values()}) > 1:
                rec["note"] = ("This district reaches different counties in the " + " and ".join(str(y) for y in years) + " results: its lines "
                               "were drawn again between those elections, and each year's figures are for the lines of that year.")
            places[kind][d] = rec
    for name in sorted(judicial):
        key = f"ND-{JUDICIAL[name]}"
        if key in votes["judicial"]:
            places["judicial"][key] = few({"name": f"{name} Judicial District", "votes": pack(votes["judicial"][key]),
                                           "counties": sorted(FIPS[c] for c in judicial[name])})
    return places, state, {y: len(k) for y, k in precincts.items()}, codes


def district_check(codes, legislative):
    """The district reading of the precinct numbers against the Secretary's legislative export: for each district that
    elected a senator that year, the precincts the export counts (by county where it says) against the precinct numbers
    carrying the district."""
    out, failed = {}, []
    for year in sorted(codes):
        doc = legislative.get(year)
        if not doc:
            out[str(year)] = {"checked": 0, "why": "the Secretary's legislative export could not be read at this run"}
            continue
        theirs = collections.defaultdict(dict)
        for chamber, d, n, county in doc["rows"]:
            if chamber == "senate" and d.isdigit():
                theirs[d][county or "*"] = n
        differ = []
        for d, by in sorted(theirs.items(), key=lambda x: int(x[0])):
            mine = {county: n for (dd, county), n in codes[year].items() if dd == d}
            same = (sum(mine.values()) == by["*"]) if "*" in by else (mine == by)
            if not same:
                differ.append(d)
                failed.append(f"{year}, District {d}: the legislative export counts {by} precincts and the precinct numbers give {mine}")
        out[str(year)] = {"checked": len(theirs), "of": DISTRICTS, "equal": len(theirs) - len(differ), "differ": differ,
                          "by_county": any("*" not in by for by in theirs.values()),
                          "what": "districts that elected a senator that year: the precincts the Secretary's All Legislative export counts in the "
                                  "district against the precinct numbers carrying the district"}
    return out, failed


def control(tables, places, state, precincts, official, contests, facts, codes, legislative, judicial, judicial_check):
    failed, per = [], {}
    for c in contests:
        t = tables[c["id"]]
        mine, off = dict(zip(SIDES, state[c["id"]])), official[c["id"]]
        theirs = {n: off[n] for n in SIDES}
        diff = {n: mine[n] - theirs[n] for n in SIDES if mine[n] != theirs[n]}
        differ = []
        for county in COUNTIES:
            rep = t["county_report"].get(county)
            got = places["county"].get(FIPS[county], {}).get("votes", {}).get(c["id"]) or dict.fromkeys(SIDES, 0)   # no precinct: no votes
            n = sum(1 for cn, _p, _v in t["rows"] if cn == county)
            if rep is None or dict(zip(SIDES, rep["votes"])) != got or rep["precincts"] != n:
                differ.append(county)
        if sorted(t["county_report"]) != sorted(COUNTIES):
            failed.append(f"{c['id']}: the County workbook's rows are not the 53 counties")
        if t["county_report_totals"]["votes"] != state[c["id"]] or t["county_report_totals"]["precincts"] != precincts[c["year"]]:
            failed.append(f"{c['id']}: the County workbook's TOTALS row is {t['county_report_totals']} and the precincts add up to {state[c['id']]} "
                          f"in {precincts[c['year']]} precincts")
        per[c["id"]] = {"sum_of_precincts": mine, "official": theirs, "equal": not diff, "source": off["source"], "where": off["where"],
                        "read": off["read"], "counties": {"compared": len(COUNTIES), "equal": len(COUNTIES) - len(differ), "differ": differ,
                                                          "with": "the contest's County workbook: each county's votes and its number of precincts"}}
        if off["source"] == COUNTY_REPORT["id"]:
            per[c["id"]]["independent"] = False
        if diff:
            per[c["id"]]["difference"] = diff
            failed.append(f"{c['id']}: the precincts add up to {mine} and the official totals ({off['source']}) are {theirs}; precincts minus official: {diff}")
        if differ:
            failed.append(f"{c['id']}: {len(differ)} counties do not equal the County workbook: {', '.join(differ[:8])}")
    for year, n in sorted(precincts.items()):
        if str(n) != str(facts[year].get("precincts")):
            failed.append(f"{year}: {n} precincts in the workbooks; the results page says {facts[year].get('precincts')}")
    kinds = {}
    for kind, first, _what, _key in KINDS:
        bad = [c["id"] for c in contests if c["year"] >= first and c["year"] in (DISTRICT_YEARS if kind in ("senate", "house") else ELECTIONS)
               and [sum(p["votes"].get(c["id"], dict.fromkeys(SIDES, 0))[x] for p in places[kind].values()) for x in SIDES] != state[c["id"]]]
        kinds[kind] = "equal" if not bad else f"differs for {', '.join(bad)}"
        failed += [f"{kind}: its places do not add up to the statewide sum for {cid}" for cid in bad]
    dcheck, dfailed = district_check(codes, legislative)
    failed += dfailed
    jcheck = {"read_from": f"the {JUDICIAL_YEAR} general election's judicial contests", "districts": len(judicial),
              "counties": sum(len(v) for v in judicial.values())}
    for year, doc in sorted(judicial_check.items()):
        if doc and doc.get("districts"):
            bad = sorted(d for d, cs in doc["districts"].items() if sorted(cs) != sorted(judicial.get(d, [])))
            jcheck[f"compared_with_{year}"] = {"districts_with_a_contest": len(doc["districts"]), "equal": len(doc["districts"]) - len(bad), "differ": bad}
            failed += [f"judicial district {d}: its counties in the {year} judicial contests are not those of {JUDICIAL_YEAR}" for d in bad]
    if len(places["county"]) != len(COUNTIES):
        failed.append(f"{len(places['county'])} counties, not {len(COUNTIES)}")
    for kind in ("senate", "house"):
        if any(c["year"] in DISTRICT_YEARS for c in contests) and sorted(places[kind], key=int) != [str(i) for i in range(1, DISTRICTS + 1)]:
            failed.append(f"{len(places[kind])} {kind} districts, not the {DISTRICTS} numbered 1 to {DISTRICTS}")
    ctl = {"result": "equal" if not failed else "differs",
           "statement": ("For every contest the precinct rows add up to the official statewide totals named here and, county by county, to the "
                         "Secretary's own county report for each of the 53 counties; in every county sheet the precincts add up to the "
                         "sheet's TOTALS row; and counties, legislative districts and judicial districts each add up to the statewide sum."
                         if not failed else "The sums do not all agree; see the differences."),
           "precincts": {str(y): n for y, n in sorted(precincts.items())}, "contests": per, "kinds": kinds,
           "district_numbers": dcheck, "judicial_districts": jcheck,
           "notes": ["For Governor the statewide control is the TOTALS row of the Secretary's own county report, a second report of the same "
                     "results system and not an independent document; President and U.S. Senator are checked against the Clerk of the U.S. "
                     "House's statistics.",
                     "The precinct and county workbooks of 2022 and 2020 call themselves 'Unofficial' in their first line; the pages they are "
                     "exported from are headed Official. Both are recorded under sources."]}
    return ctl, failed


# ---------------------------------------------------------------- the file

def build(tables, facts, judicial_doc, legislative, judicial_check, official, official_docs, contests=CONTESTS, files=None):
    judicial = judicial_doc["districts"]
    problem = judicial_problem(judicial)
    if problem:
        raise Stop(f"    {problem}; stopping")
    places, state, precincts, codes = tally(tables, judicial, contests)
    ctl, failed = control(tables, places, state, precincts, official, contests, facts, codes, legislative, judicial, judicial_check)
    given = {kind: [c["id"] for c in contests if c["year"] >= first and (kind not in ("senate", "house") or c["year"] in DISTRICT_YEARS)]
             for kind, first, _w, _k in KINDS}
    kinds = {}
    for kind, first, what, key in KINDS:
        kinds[kind] = {"what": what, "key": key, "places": len(places[kind]), "contests": given[kind], "covers_the_state": True,
                       "places_with_a_contest_too_few": sum(1 for p in places[kind].values() if p.get("too_few"))}
        if first > min(c["year"] for c in contests):
            kinds[kind]["why_not_2020"] = REDRAWN
    changed = sorted((d for d, p in places["senate"].items() if p.get("note")), key=int)
    lines = ("A precinct's district is the third and fourth digits of its number. Each year is on the lines in force at that election. "
             + (f"District{'s' if len(changed) != 1 else ''} {', '.join(changed)} reach{'es' if len(changed) == 1 else ''} different counties in "
                "the 2022 and the 2024 results, so the lines there were drawn again between the two elections; a change inside a county would "
                "not show in this comparison. " if changed else "")
             + "These files do not say which lines a later election uses.")
    kinds["senate"]["note"] = lines
    kinds["house"]["note"] = lines + " The figures are the whole district's, the same as under senate; House subdistricts (4A, 4B) are not given apart."
    kinds["judicial"]["note"] = (f"The counties of each district are those that voted in its judicial contests at the {JUDICIAL_YEAR} general "
                                 "election; every year's figures are those counties' votes.")
    records = []
    for c in contests:
        t = tables[c["id"]]
        records.append({"id": c["id"], "date": ELECTIONS[c["year"]]["date"], "office": c["office"], "table": f"nd-sos-precinct-results-{c['year']}",
                        "kinds": [k for k in KIND_NAMES if c["id"] in given[k]],
                        "dem": {"party": DEM_PARTY, "ticket": t["tickets"]["dem"], "column": "the column the County workbook puts under Democratic-NPL"},
                        "rep": {"party": REP_PARTY, "ticket": t["tickets"]["rep"], "column": "the column the County workbook puts under Republican"},
                        "other": {"what": "every other candidate and all write-ins, together", "columns": t["other_parties"]},
                        "total": {"what": "the votes cast for candidates and write-ins; the workbooks give no under votes or over votes"},
                        "statewide": dict(zip(SIDES, state[c["id"]])), "official_source": official[c["id"]]["source"]})
    sources = []
    for year in sorted({c["year"] for c in contests}, reverse=True):
        f = facts[year]
        mine = [c for c in contests if c["year"] == year]
        sources.append({"id": f"nd-sos-precinct-results-{year}", "kind": "official results by precinct", "agency": AGENCY,
                        "title": f"{f.get('heading')}: the Precinct and County exports of {', '.join(c['heading'] for c in mine)}",
                        "url": f.get("url"), "linked_from": ARCHIVE, "linked_as": ELECTIONS[year]["link"],
                        "the_page_says": {"heading": f.get("heading"), "last_updated": f.get("updated"), "total_precincts": f.get("precincts")},
                        "the_workbooks_say": sorted({tables[c["id"]]["says"][k] for c in mine for k in ("precinct", "county")}),
                        "fetched": f.get("fetched"), "precincts": precincts[year],
                        "tables": {c["id"]: {"rows_sha256": M.rows_fingerprint(tables[c["id"]]["rows"]),
                                             "files": (files or {}).get(c["id"])} for c in mine},
                        "read": "Of each precinct: its county (the sheet), its number or name, and each candidate's votes. Of the County "
                                "workbook: each candidate's party, each county's votes and number of precincts. Only the two tickets are named."})
    sources.append({"id": "nd-sos-judicial-district-counties", "kind": "official results (the counties of each judicial contest)", "agency": AGENCY,
                    "title": f"{judicial_doc.get('election')}: the County export of one judicial contest in each district",
                    "url": judicial_doc.get("url"), "fetched": judicial_doc.get("fetched"), "read": judicial_doc.get("kept")})
    for year, doc in sorted(legislative.items()):
        if doc:
            sources.append({"id": f"nd-sos-legislative-export-{year}", "kind": "official results (used only to check the district numbers)",
                            "agency": AGENCY, "title": f"{ELECTIONS[year]['link']}: Exports, All Legislative (CSV)", "url": doc.get("url"),
                            "fetched": doc.get("fetched"), "sha256": doc.get("sha256"), "read": doc.get("kept")})
    doc = {"what": WHAT, "note": NOTE, "state": "ND", "generated": M._now(), "method_version": METHOD, "how_to_read": HOW_TO_READ,
           "vote_keys": list(SIDES), "too_few": {"fewer_than": FEW, "why": TOO_FEW},
           "contests": records, "kinds": kinds, "control": ctl,
           "coverage": {"districts_whose_counties_changed": changed, "not_given": NOT_GIVEN},
           "sources": sources + official_docs, "places": places}
    return doc, failed


def our_places(db):
    """{kind: ids} of our North Dakota places and races, read only; None when the database is not there."""
    if not db or not os.path.exists(db):
        return None
    con = sqlite3.connect(pathlib.Path(os.path.abspath(db)).as_uri() + "?mode=ro", uri=True)
    try:
        out = collections.defaultdict(set)
        for kind, pid in con.execute("SELECT kind, id FROM sl_places WHERE source_id LIKE 'nd-%' AND kind IN ('county', 'judicial_district')"):
            out["county" if kind == "county" else "judicial"].add(pid)
        for office, jid in con.execute("SELECT office_kind, jurisdiction_id FROM sl_races WHERE state = 'ND' AND office_kind IN ('state_senate', 'state_house')"):
            out["senate" if office == "state_senate" else "house"].add(str(jid))
        for jid, in con.execute("SELECT DISTINCT jurisdiction_id FROM sl_races WHERE state = 'ND' AND level = 'county'"):
            out["county"].add(str(jid))
        return dict(out)
    except sqlite3.Error:
        return None
    finally:
        con.close()


def load(say=print, out=OUT, cache=CACHE, db=DB, refresh=False):
    say("    North Dakota place votes: the Secretary of State's official precinct results (results.sos.nd.gov)")
    net.patient_lookups()
    os.makedirs(os.path.join(cache, "results"), exist_ok=True)
    visit = []

    def site():                                                            # one visit, begun only if something must be fetched
        if not visit:
            visit.append(Site())
        return visit[0]
    years = sorted({c["year"] for c in CONTESTS}, reverse=True)
    facts = {year: fetch_year(year, CONTESTS, cache, refresh, say, site) for year in years}
    tables, files = {}, {}
    for c in CONTESTS:
        raw = {}
        for kind in ("precinct", "county"):
            with open(book_path(cache, c["id"], kind), "rb") as fh:
                raw[kind] = fh.read()
        tables[c["id"]] = contest_table(c, read_precinct_book(raw["precinct"], f"{c['id']} precinct workbook"),
                                        read_county_book(raw["county"], f"{c['id']} county workbook"))
        files[c["id"]] = {kind: {"file": os.path.basename(book_path(cache, c["id"], kind)), "sha256": M._sha(data),
                                 "fetched": M._day(book_path(cache, c["id"], kind))} for kind, data in raw.items()}
    for year in years:
        say(f"      {year}: {sum(1 for _r in tables[next(c['id'] for c in CONTESTS if c['year'] == year)]['rows']):,} precincts "
            f"({facts[year].get('heading')}; last updated {facts[year].get('updated')}; copy of {facts[year].get('fetched')})")
    judicial_doc = fetch_judicial(cache, refresh, say, site)
    legislative = {year: fetch_legislative(year, cache, refresh, say, site) for year in DISTRICT_YEARS}
    judicial_check = {2024: fetch_judicial_check(2024, cache, refresh, say, site)}
    say("    the official statewide totals: Clerk of the U.S. House (President and U.S. Senator); the Secretary's own county report (Governor)")
    official, official_docs = official_totals(CONTESTS, tables, cache, refresh, say)
    doc, failed = build(tables, facts, judicial_doc, legislative, judicial_check, official, official_docs, files=files)

    p = doc["places"]
    say(f"    added up: {len(p['county'])} counties; {len(p['senate'])} legislative districts (2022 and 2024); {len(p['judicial'])} judicial districts")
    if doc["coverage"]["districts_whose_counties_changed"]:
        say("      districts reaching different counties in 2022 and 2024: " + ", ".join(doc["coverage"]["districts_whose_counties_changed"]))
    for c in doc["contests"]:
        ctl, s = doc["control"]["contests"][c["id"]], c["statewide"]
        say(f"    control, {c['id']}: Democratic-NPL {s['dem']:,}, Republican {s['rep']:,}, other {s['other']:,}, total {s['total']:,}: "
            + ("equal to" if ctl["equal"] else "DIFFERS from") + f" {ctl['source']} ({ctl['where']}; {ctl['read']}); "
            f"{ctl['counties']['equal']} of {ctl['counties']['compared']} counties equal to the county report")
    for year, d in doc["control"]["district_numbers"].items():
        say(f"    district numbers, {year}: " + (f"{d['equal']} of {d['checked']} districts with a Senate contest have the precincts the "
                                                  f"legislative export counts" if d.get("checked") else d.get("why", "not checked")))
    for kind, result in doc["control"]["kinds"].items():
        if result != "equal":
            say(f"    control, {kind}: {result}")
    if failed:
        for line in failed:
            say(f"    CHECK: {line}")
        raise Stop(f"    North Dakota place votes: the control did not hold; {os.path.basename(out)} was not written. Nothing was changed to make it fit.")
    ours = our_places(db)
    if ours:
        covered = {}
        for kind in KIND_NAMES:
            ids = ours.get(kind, set())
            without = sorted(i for i in ids if not p[kind].get(i, {}).get("votes"))
            covered[kind] = {"places": len(ids), "with_votes": len(ids) - len(without), "without": without}
            say(f"      our {kind} places with votes: {len(ids) - len(without):,} of {len(ids):,}" + (f" (none for {', '.join(without[:12])})" if without else ""))
        doc["coverage"]["our_places"] = dict(covered, read=f"sl_places and sl_races in {os.path.basename(db)}, opened read-only on {M._now()}")
    M.write(out, doc)
    say(f"    North Dakota place votes: {len(doc['contests'])} contests, {sum(len(v) for v in p.values()):,} places, control {doc['control']['result']}; "
        f"{os.path.relpath(out, HERE)} ({os.path.getsize(out) / 1e6:.2f} MB)")
    return doc


# ---------------------------------------------------------------- the arithmetic and the readers on made-up workbooks

def _made_up_precinct_book(first_line, heading, candidates, counties, spoil=None):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for county in COUNTIES:
        ws = wb.create_sheet(county)
        for line in (first_line, "State of North Dakota", "Downloaded at some time", None, None, None):
            ws.append([line])
        ws.append([heading + " ", "Precinct"] + [c + " " for c in candidates])
        rows = counties.get(county, [])
        for precinct, v in rows:
            ws.append([None, precinct] + list(v))
        totals = [sum(v[i] for _p, v in rows) for i in range(len(candidates))]
        if spoil == county:
            totals[0] += 1
        ws.append([None, "TOTALS"] + totals)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _made_up_county_book(first_line, heading, candidates, counties):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for line in (first_line, "State of North Dakota", "Downloaded at some time", None, None, None):
        ws.append([line])
    ws.append([heading + " ", "County", "Number of Precincts"] + [f"{n} \n{p}" if p else n + " " for n, p in candidates])
    n = len(candidates)
    for county in COUNTIES:
        rows = counties.get(county, [])
        ws.append([None, county, len(rows)] + [sum(v[i] for _p, v in rows) for i in range(n)])
    ws.append([None, "TOTALS", sum(len(r) for r in counties.values())] + [sum(v[i] for r in counties.values() for _p, v in r) for i in range(n)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def selftest(say=print):
    checks = []

    def check(what, got, want):
        checks.append(got == want)
        say(f"      {'ok  ' if got == want else 'FAIL'} {what}" + ("" if got == want else f": got {got!r}, expected {want!r}"))
    cands24 = [("Ash and Birch", "Republican"), ("Cedar and Dogwood", "Democratic-NPL"), ("Elm and Fir", "Libertarian"), ("write-in", "")]
    cands20 = [("Gum and Hazel", "Democratic-NPL"), ("Ivy and Juniper", "Republican"), ("write-in", "")]
    c24 = {"Adams": [("013901", (30, 10, 2, 1))], "Barnes": [("022401", (20, 25, 0, 0)), ("022402", (5, 5, 1, 0))],
           "Benson": [("030901", (7, 0, 0, 0)), ("031501", (9, 10, 0, 1))], "Cass": [("091101", (100, 120, 5, 2))]}
    c20 = {"Adams": [("Adams County", (12, 30, 1))], "Barnes": [("Precinct 2401-00", (22, 21, 0))], "Cass": [("1101 A Church", (90, 80, 3))]}
    contests = [{"id": "2024-president", "year": 2024, "office": "President", "heading": "President", "dem": ("Cedar", "Dogwood"), "rep": ("Ash", "Birch"),
                 "official": "x", "section": None},
                {"id": "2020-president", "year": 2020, "office": "President", "heading": "President", "dem": ("Gum",), "rep": ("Ivy", "Juniper"),
                 "official": "x", "section": None}]
    names = lambda cs: [n for n, _p in cs]                                                                             # noqa: E731
    p24 = read_precinct_book(_made_up_precinct_book("2024 Official General Election Results", "President", names(cands24), c24), "made up 2024")
    k24 = read_county_book(_made_up_county_book("2024 Official General Election Results", "President", cands24, c24), "made up 2024")
    p20 = read_precinct_book(_made_up_precinct_book("2020 Unofficial General Election Results", "President", names(cands20), c20), "made up 2020")
    k20 = read_county_book(_made_up_county_book("2020 Unofficial General Election Results", "President", cands20, c20), "made up 2020")
    tables = {"2024-president": contest_table(contests[0], p24, k24), "2020-president": contest_table(contests[1], p20, k20)}
    check("the parties come from the County workbook, whatever the column order", tables["2024-president"]["tickets"],
          {"dem": "Cedar and Dogwood", "rep": "Ash and Birch"})
    check("other columns are kept by party only", tables["2024-president"]["other_parties"], ["Libertarian", "write-in"])
    judicial = {"districts": {d: [] for d in JUDICIAL}, "election": "made up", "url": "-", "fetched": "-", "kept": "-"}
    for i, county in enumerate(COUNTIES):
        judicial["districts"][sorted(JUDICIAL)[i % 8]].append(county)
    facts = {2024: {"precincts": "6", "heading": "made up"}, 2020: {"precincts": "3", "heading": "made up"}}
    official = {"2024-president": {"dem": 170, "rep": 171, "other": 12, "total": 353, "source": "x", "where": "-", "read": "-"},
                "2020-president": {"dem": 124, "rep": 131, "other": 4, "total": 259, "source": "x", "where": "-", "read": "-"}}
    legislative = {2024: {"rows": [["senate", "24", 2, "Barnes"], ["senate", "9", 1, "Benson"], ["house", "4A", 9, None]], "url": "-"}}

    def run(tables=tables, official=official, legislative=legislative, facts=facts):
        ctl_places, state, precincts, codes = tally(tables, judicial["districts"], contests)
        return ctl_places, control(tables, ctl_places, state, precincts, official, contests, facts, codes, legislative, judicial["districts"], {})
    places, (ctl, failed) = run()
    expected_only = [f for f in failed if "districts, not the" not in f and "counties, not" not in f]   # the made-up state is four counties, four districts
    check("the control holds on the made-up workbooks", (expected_only, ctl["kinds"]), ([], dict.fromkeys(KIND_NAMES, "equal")))
    check("a county adds up its precincts", places["county"]["38003"]["votes"]["2024-president"], {"dem": 30, "rep": 25, "other": 1, "total": 56})
    check("a district is read from the precinct number", places["senate"]["24"]["votes"]["2024-president"], {"dem": 30, "rep": 25, "other": 1, "total": 56})
    check("and reaches across counties", sorted(places["senate"]["9"]["counties"]["2024"]), ["38005"])
    check("house carries the same figures as senate", places["house"], places["senate"])
    check("districts are not given for 2020", "2020-president" in places["senate"]["24"]["votes"], False)
    check("counties are given for 2020, where precincts are names", places["county"]["38001"]["votes"]["2020-president"],
          {"dem": 12, "rep": 30, "other": 1, "total": 43})
    check("a judicial district is whole counties", sum(p["votes"]["2024-president"]["total"] for p in places["judicial"].values()), 353)
    check("a place where every vote went one way is marked too_few", places["senate"]["9"].get("too_few"), ["2024-president"])
    check("a place of twenty votes, not all one way, is not marked", places["senate"]["15"].get("too_few"), None)
    check("the district numbers are checked against the legislative export", ctl["district_numbers"]["2024"]["equal"], 2)
    wrong = dict(official, **{"2024-president": dict(official["2024-president"], dem=171, total=354)})
    check("a wrong official total is caught", run(official=wrong)[1][0]["result"], "differs")
    bad_leg = {2024: {"rows": [["senate", "24", 3, "Barnes"]], "url": "-"}}
    check("a district with another number of precincts in the legislative export is caught", run(legislative=bad_leg)[1][0]["district_numbers"]["2024"]["differ"], ["24"])
    check("another number of precincts than the page states is caught", run(facts={**facts, 2024: {"precincts": "7"}})[1][0]["result"], "differs")
    for what, make in (("a sheet whose precincts do not add up to its TOTALS row stops the loader",
                        lambda: read_precinct_book(_made_up_precinct_book("x", "President", names(cands24), c24, spoil="Barnes"), "made up")),
                       ("a precinct number of another county stops the loader", lambda: precinct_place("032401", "Barnes", 2024)),
                       ("two columns under one party stop the loader",
                        lambda: contest_table(contests[0], p24, dict(k24, candidates=[(n, "Republican") for n, _p in cands24]))),
                       ("a ticket other than the one typed in stops the loader",
                        lambda: contest_table(dict(contests[0], dem=("Walnut",)), p24, k24))):
        try:
            make()
            caught = False
        except Stop:
            caught = True
        check(what, caught, True)
    check("the legislative export keeps districts and counts, not names",
          legislative_rows('ContestName,PartyCode,AreaNum,CandidateName,VoteFor,CandidateVotes,CandidatePercentage,PrecinctsReporting\r\n'
                           '"State Senator","REP","District 24","A Name",1,10,50%,15/15,"Barnes",\r\n'
                           '"State Senator","DEM","District 24","B Name",1,10,50%,15/15,"Barnes",\r\n'
                           '"State Representative","REP","District 04a","C Name",1,10,50%,9/9,\r\n'),
          [["house", "4A", 9, None], ["senate", "24", 15, "Barnes"]])
    check("county codes run in the order of the sheets", (FIPS["Adams"], FIPS["LaMoure"], FIPS["McHenry"], FIPS["Williams"], len(FIPS)),
          ("38001", "38045", "38049", "38105", 53))
    say(f"    self-test: {sum(checks)} of {len(checks)} checks passed")
    return all(checks)


def main(argv=None):
    ap = argparse.ArgumentParser(description="How each North Dakota place voted in past partisan general elections -> ballot/lean/nd_place_votes.json")
    ap.add_argument("--out", default=OUT, help="the file to write (default: ballot/lean/nd_place_votes.json)")
    ap.add_argument("--cache", default=CACHE, help="where downloads are kept (default: states_cache/nd_local)")
    ap.add_argument("--db", default=DB, help="the ballot database to count our places in, opened read-only (default: ballot_local_2026.sqlite)")
    ap.add_argument("--refresh", action="store_true", help="ask for every workbook and document again, even when copies are on disk")
    ap.add_argument("--selftest", action="store_true", help="check the arithmetic and the readers on made-up workbooks; downloads and writes nothing")
    a = ap.parse_args(argv)
    if a.selftest:
        raise SystemExit(0 if selftest() else 1)
    load(out=a.out, cache=a.cache, db=a.db, refresh=a.refresh)


if __name__ == "__main__":
    main()
