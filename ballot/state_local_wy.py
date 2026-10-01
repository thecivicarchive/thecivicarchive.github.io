"""
ballot/state_local_wy.py - Wyoming's state races on the November 3, 2026 ballot: the State Senate seats up this year
(the sixteen odd-numbered districts, whose four-year terms end in January, and District 6, filled early), all 62 State
House seats (one member a district, two-year terms), and the five statewide offices Wyoming elects: Governor, Secretary
of State, State Auditor, State Treasurer and Superintendent of Public Instruction (Wyoming has no Lieutenant Governor,
and its Attorney General is appointed). Every one of these offices is partisan.

Sources, the Secretary of State's own (Elections Division, sos.wyo.gov) and nothing else:

  - The "2026 General Election Candidate Roster" (PDF), the same file ballot/lists/wy.py reads for Congress. Its columns
    are Office Sought, Party Affiliation, Candidate Name, Mailing Address, Date Filed and Campaign Telephone, with City,
    State & ZIP, Date Withdrawn and Email on each entry's second line. The column edges come from the page's own
    headings; only the office, the party and the name cells, and a date under Date Withdrawn, are ever joined into text
    or kept. Addresses, cities, ZIP codes, telephones and e-mail stay in the file and are never read, printed or stored.
    The loader reads the copy the federal loader keeps (ballot_cache/wy/); if there is none it holds a fresh copy in
    memory only, never on disk. The state rows it reads are kept, allowed columns only, as
    ballot_cache/wy/wy_2026_state_roster.json. The roster gives no ballot order (Wyoming's county clerks print the
    ballots), so no ballot order is stored. A candidate with a withdrawal date is left off.
  - The August 18 primary: "2026 Primary Results Summaries - OFFICIAL.xlsx" in the zip of data files on the 2026
    Official Primary Election Results page (the federal loader's cached copy). Its sheets "Statewide Candidates",
    "Statewide Senate Odd" and "Statewide House" give each office's candidates by party, with Write-Ins, Overvotes and
    Undervotes and a Total row. The workbook also holds a sheet "Statewide Senate Even" (districts 2 to 30, even) whose
    districts are not on the 2026 roster, are not in the county precinct-by-precinct results, and whose District 6 block
    differs from the one in "Statewide Senate Odd": it is set aside, and every summary block is kept only when the
    county precinct-by-precinct results add up to it, candidate by candidate. The sheets BLANKSS and BLANKSH say they
    are blank reference sheets and are not read.
  - The same zip's "2026 Primary County PbP Results - OFFICIAL.xlsx" and "2026 Primary County PbP Total Ballots Cast -
    OFFICIAL.xlsx", used as checks: every precinct's figures for a block are summed and must equal the summary's; the
    party ballots cast in the precincts where a race was on the ballot must equal its candidates' votes plus write-ins,
    overvotes and undervotes; and the counties where a district's race was on the ballot are its county_ids. Some
    counties' precinct numbers (01-02) were stored by the spreadsheet as dates; they are read back as month-day.

A party primary is a field when two or more names were printed on that party's ballot (a candidate who withdrew after
the ballots were printed, "* Withdrawn Candidate" in the sheets and named in a footnote, is kept with the votes the
summary counts and a note). Write-ins count toward the share but are not listed; over- and undervotes are left out.
Who advanced is the party's candidate on the November roster, checked against the most votes. The minor parties
(Libertarian, Constitution) nominate outside the primary; independents file by petition.

Today's holders come from state_wy.sqlite (the Open States roster): current legislators by chamber and district, and
the officials table for Governor and Secretary of State (the roster does not carry the Auditor, the Treasurer or the
Superintendent). Only names, parties, districts, ids and term starts are read from it. A candidate is marked as the
sitting member (incumbent 1, state_member_id) only when the name fits that seat's holder, one to one. A candidate who
sits today in the other chamber, in a district that shares a precinct with this one in the primary results, or (for a
statewide race) who holds any seat or office in the roster, gets state_member_id with incumbent 0 and a note, again
only when exactly one roster person fits.

The county and local level (John, 2026-09-30)
---------------------------------------------
Wyoming has no statewide list of county, city, school or special district candidates: each of the 23 county clerks
publishes its own. What is read, and what never is:

  - Judges standing for retention (level "court"): the Secretary of State's "General Election Candidates" data file
    (CSV, the roster's own list plus 21 judges). It has mailing address, city, telephone, e-mail and web address
    columns: the file is held in memory, cut down by its header to Election, Office Sought, Party Affiliation, Ballot
    Name and Date Withdrawn, and only that cut-down copy is kept (ballot_cache/wy/local/, as JSON). Which counties vote
    on a district or circuit judge comes from the statute that lists the judicial districts (W.S. 5-3-101, read from
    the Legislature's own Title 5 file).
  - County sample ballots: the November 3 ballots the county clerks post, precinct by precinct (text PDFs made by the
    counties' ballot-printing system; a ballot carries no contact details at all, and each file is read in memory and
    checked before it is kept). Every precinct's ballot is read: section headings, contest titles, the term and
    vote-for lines, party lines, names and the write-in lines; the wording of ballot questions is not read. The
    precinct ballots of a county are merged into one list of contests; a contest printed on two counties' ballots (a
    school or special district that crosses the county line) is one race, counted once. Names rotate from precinct to
    precinct (W.S. 22-6-122), so no ballot order is stored. Type sizes and column edges differ from county to county,
    so lines are told apart by their words and by where each file's own lines start. Two counties' files draw letters
    more than once on the same spot (Johnson's every letter two or three times; Converse's again wherever the SAMPLE
    mark lies across them): such a file is read letter by letter, each letter kept once, and its county is kept only
    when every state and federal name on its ballots reads exactly as the Secretary of State's file has it. A sample
    ballot that is a scanned picture has no text to read, and its county is named in sl_gaps. A file that is not a
    November 3, 2026 general election ballot (the primary's sample, an older year's) is set aside. The clerks who had
    posted no November sample ballot have their pages looked at again on --refresh for a link that says so.
  - Laramie County, which had posted no sample ballot: the county clerk's three general election rosters (XLSX). Each
    is held in memory and cut down by its header to Office, Party, Ballot Name, Withdrawn Date and Municipality; the
    mailing address, city, phone, e-mail and website columns are never read, and only the cut-down rows are kept.
  - Natrona County, the same way: the clerk's general election roster (PDF, a table). Its column edges come from each
    page's own row of headings; the cells of the first four columns (whose office, Office, Party, Ballot Name) are put
    together by where they start, a piece of text that starts at or beyond the Mailing Address column is never joined
    into anything, and only the four columns are kept.
  - Big Horn County, the same way: the clerk's general election roster (PDF: nine tables, for the county's offices,
    its towns, and its school and special districts, with cells centred under their headings). A cell is taken for a
    name only when it sits under the middle of the Candidate Name heading; whatever starts to the right of that and is
    not a name belongs to the mailing address, phone and e-mail columns and is never put together into text. This
    roster says how many are to be elected in each contest, and lists a contest even when no one filed for it.
  - Fremont County, in part: the clerk's list of general election filings for school, college and special district
    offices (PDF, a table read the same way: District Name, Office and Ballot Name only). The county's own offices
    and the cities' are not on that list and stay a gap until the clerk posts sample ballots.
  - Place names and codes: the Census Bureau's 2020 place codes file for cities and towns, and its 2024 county file.

A roster or a list of filings that does not say how many are to be elected has that said in each of its races' notes.
A candidate who withdrew is left off where the list marks it (Laramie's withdrawal date, a Fremont office that says
withdrawn); a ballot simply does not print one. Where a district's title does not say what its board's members are
called, the office is shown as "Board Member" (a conservation district's are supervisors on every list that names
them). A district on two counties' lists is one place: "X COUNTY FIRE PROTECTION DISTRICT #5" on a neighbour's ballot
and "FIRE # 5" on X County's own roster are filed under one id, and a contest with the same candidates on both lists
is one race.

Local ids: a county office is filed under the county's five-digit code; a city or town under WY-M-<Census place code>;
a school district under WY-S-<its county's three-digit code>-<its number>; a hospital district under WY-H-<key> and
any other district under WY-X-<key>, where a key is the county's three-digit code and the district's name, or the
name alone when the district is on more than one county's ballots. Everything that could not be loaded is a row in
sl_gaps, with the reason; sl_notes says which local offices are on this ballot and what the loaded lists cover.

Usage: python ballot/state_local_wy.py <database file> [--cache <folder>] [--local-cache <folder>] [--refresh]
       --refresh asks the county pages, the candidates file and the statute again; without it a re-run downloads
       nothing that is already in ballot_cache/wy/local/.
"""

import csv
import datetime as dt
import hashlib
import html
import io
import json
import os
import re
import sqlite3
import sys
import time
import unicodedata
import zipfile
from collections import Counter, OrderedDict, defaultdict
from html.parser import HTMLParser
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urljoin, urlparse

import openpyxl

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from ballot.check_local import EXTRA_SCHEMA, contact_like                   # noqa: E402
from ballot.common import fold, name_parts, party_code                     # noqa: E402
from ballot.lists.wy import BOOK, DATE, LIST_URL, PRINTED, RESULTS_URL, cells, iso   # noqa: E402
from ballot.match import fits, initials_clash                               # noqa: E402
from ballot.pdftext import PDF, join, page_runs, rows as pdf_rows           # noqa: E402
from states import net                                                      # noqa: E402

STATE, FIPS, NAME = "WY", "56", "Wyoming"
GENERAL, PRIMARY = "2026-11-03", "2026-08-18"
ROSTER_DB = os.path.join(HERE, "state_wy.sqlite")
CACHE = os.path.join(HERE, "ballot_cache", "wy")
COUNTY_ZIP = os.path.join(HERE, "states_cache", "census", "cb_2024_us_county_500k.zip")
COUNTY_URL = "https://www2.census.gov/geo/tiger/GENZ2024/shp/cb_2024_us_county_500k.zip"
PRECINCT_BOOK = "2026 Primary County PbP Results - OFFICIAL.xlsx"
BALLOTS_BOOK = "2026 Primary County PbP Total Ballots Cast - OFFICIAL.xlsx"

# the roster's office words -> (race key, office_kind, office shown, roster office, the results sheets' words)
STATEWIDE = {
    "GOVERNOR": ("GOV", "governor", "Governor", "governor", "Governor"),
    "SECRETARY OF STATE": ("SOS", "secretary_of_state", "Secretary of State", "secretary of state", "Secretary of State"),
    "STATE AUDITOR": ("AUD", "state_auditor", "State Auditor", None, "State Auditor"),
    "STATE TREASURER": ("TREAS", "state_treasurer", "State Treasurer", None, "State Treasurer"),
    "SUPERINTENDENT OF PUBLIC INSTRUCTION": ("SPI", "superintendent_of_public_instruction", "Superintendent of Public Instruction",
                                             None, "Superintendent of Public Instruction"),
}
FEDERAL = ("UNITED STATES SENATOR", "UNITED STATES REPRESENTATIVE")          # the federal loader's
LEG_ROSTER = re.compile(r"STATE (SENATOR|REPRESENTATIVE) 0*(\d+)")
LEG_RESULTS = re.compile(r"(Senate|House) District 0*(\d+)")
OFFICE_ANY = re.compile(r"^(United States (Senator|Representative)|Governor|Secretary of State|State Auditor|State Treasurer|"
                        r"Superintendent of Public Instruction|(Senate|House) District \d+)(, Continued)?$")
CONT = re.compile(r", Continued$")
TALLY = ("Write-Ins", "Overvotes", "Undervotes")
WITHDRAWN_LABEL = "* Withdrawn Candidate"
WITHDREW = re.compile(r"^\*\s*(?P<name>.+?) withdrew (?:his|her|their) candidacy after official ballots had been printed"
                      r"(?: by (?P<county>[A-Z][A-Za-z ]+?) County)?", re.I)
CHAMBER = {"Senate": ("SS", "state_senate", "State Senator"), "House": ("SH", "state_house", "State Representative")}
CHAMBER_WORDS = {"Senate": "the Wyoming Senate", "House": "the Wyoming House of Representatives"}
FIELD_CODE = {"Republican": "REP", "Democratic": "DEM"}

SRC_GENERAL = "wy-sos-2026-state-general-roster"
SRC_PRIMARY = "wy-sos-2026-state-primary-summary"
SRC_PRECINCTS = "wy-sos-2026-state-primary-precincts"
SRC_ROSTER = "wy-openstates-roster"
SRC_COUNTIES = "wy-census-cb-2024-county"

SCHEMA = """
CREATE TABLE IF NOT EXISTS sl_races (race_id TEXT PRIMARY KEY, state TEXT NOT NULL, level TEXT NOT NULL, office_kind TEXT NOT NULL, office TEXT NOT NULL, jurisdiction TEXT, jurisdiction_id TEXT, county_ids TEXT, district TEXT, seat TEXT, special INTEGER NOT NULL DEFAULT 0, partisan INTEGER NOT NULL, holder_id TEXT, holder_name TEXT, holder_party TEXT, election_date TEXT NOT NULL, note TEXT);
CREATE TABLE IF NOT EXISTS sl_candidates (race_id TEXT NOT NULL, election TEXT NOT NULL, election_date TEXT, name TEXT NOT NULL, party TEXT, party_code TEXT, ballot_order INTEGER, incumbent INTEGER NOT NULL DEFAULT 0, write_in INTEGER NOT NULL DEFAULT 0, votes INTEGER, pct REAL, outcome TEXT, state_member_id TEXT, source_id TEXT, note TEXT, PRIMARY KEY (race_id, election, name));
CREATE TABLE IF NOT EXISTS sl_sources (source_id TEXT PRIMARY KEY, state TEXT, kind TEXT, agency TEXT, title TEXT, url TEXT, published TEXT, fetched TEXT, sha256 TEXT, rows INTEGER, note TEXT);
CREATE TABLE IF NOT EXISTS sl_places (kind TEXT NOT NULL, id TEXT NOT NULL, name TEXT NOT NULL, county_ids TEXT, source_id TEXT, PRIMARY KEY (kind, id));
"""


def fail(msg):
    raise SystemExit(f"Wyoming (state races): {msg}")


def sha_of(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest() if path and os.path.exists(path) else ""


def mdate(path):
    return dt.date.fromtimestamp(os.path.getmtime(path)).isoformat() if path and os.path.exists(path) else ""


# ---------- the November roster: office, party, name and a withdrawal date only ----------

def roster_rows(data):
    """([office, party, name, withdrawn date] for every entry, printed date). Only those cells are joined into text."""
    pdf = PDF(data)
    out, office, printed, titled = [], None, "", False
    for page, res in pdf.pages():
        edges, body = {}, False
        for _y, runs in pdf_rows(pdf, page, res):
            row = [(x, rs) for x, rs in cells(runs) if x < 600]      # the Division's own address block sits at the right
            if not body:      # the title and the two heading lines; the column edges come from the headings
                heads = {join(rs): x for x, rs in row}
                titled = titled or "2026 General Election Candidate Roster" in heads
                if {"Party Affiliation", "Candidate Name", "Mailing Address"} <= set(heads):
                    edges.update(party=heads["Party Affiliation"] - 12, name=heads["Candidate Name"] - 12, addr=heads["Mailing Address"] - 12)
                elif "Date Withdrawn" in heads and edges:
                    edges.update(wd=heads["Date Withdrawn"] - 12, email=heads.get("Email", 10 ** 6) - 12)
                    body = True
                continue
            left = join([r for x, rs in row if x < edges["party"] for r in rs])
            party = join([r for x, rs in row if edges["party"] <= x < edges["name"] for r in rs])
            name = join([r for x, rs in row if edges["name"] <= x < edges["addr"] for r in rs])
            if party or name:
                if not (party and name):
                    fail(f"a roster entry under {office} has a party or a name but not both; read the roster again")
                if left:
                    fail(f"a roster entry under {office} has text in the office column; read the roster again")
                out.append([office, party, re.sub(r"\s+", " ", name), ""])
                continue
            if left:
                m = PRINTED.match(left)
                if m:
                    printed = re.sub(r" 0(\d),", r" \1,", m.group(1))
                else:
                    office = left
                continue
            # an entry's second line: only a date under Date Withdrawn is looked for (city, ZIP and e-mail are never joined)
            dates = [join(rs) for x, rs in row if edges["wd"] <= x < edges["email"]]
            dates = [d for d in dates if DATE.fullmatch(d)]
            if dates and out:
                out[-1][3] = dates[0]
    if not titled:
        fail("the file is not the 2026 General Election Candidate Roster")
    return out, printed


def read_roster(cache, say=print):
    """The state rows of the roster (allowed columns only), from the federal loader's copy or a fresh copy held in memory;
    the kept JSON if neither can be read. Returns (rows, printed, fetched, sha256, how, all_rows)."""
    pdf_path = os.path.join(cache, "wy_2026_general_candidate_roster.pdf")
    json_path = os.path.join(cache, "wy_2026_state_roster.json")
    data, how, fetched = None, "", ""
    if os.path.exists(pdf_path):
        data, how, fetched = open(pdf_path, "rb").read(), "the copy the federal loader keeps in ballot_cache/wy/", mdate(pdf_path)
    else:
        last = None
        for attempt in range(3):
            try:
                data = net.get(LIST_URL, accept="application/pdf")
                break
            except (HTTPError, URLError, OSError) as e:
                last = e
                if attempt < 2:
                    time.sleep(5 * (attempt + 1))
        if data is None:
            say(f"    Wyoming (state races): the roster could not be read ({last})")
        how, fetched = "read afresh and held in memory only (the PDF is not saved)", dt.date.today().isoformat()
    if data is not None and not data.startswith(b"%PDF"):
        say("    Wyoming (state races): the roster address answered without a PDF (a bot check or a changed page)")
        data = None
    if data is None:
        if not os.path.exists(json_path):
            fail("no roster to read, and no kept copy of its state rows; stopped")
        got = json.load(open(json_path, encoding="utf-8"))
        return got["rows"], got["printed"], got["fetched"], got["sha256"], f"the state rows kept on {got['fetched']}", got["all_rows"]
    rows, printed = roster_rows(data)
    state_rows = [r for r in rows if r[0] not in FEDERAL]
    sha = hashlib.sha256(data).hexdigest()
    os.makedirs(cache, exist_ok=True)
    with open(json_path, "w", encoding="utf-8") as fh:
        json.dump({"url": LIST_URL, "fetched": fetched, "printed": printed, "sha256": sha, "all_rows": len(rows),
                   "columns": ["office", "party", "name", "withdrawn"], "rows": state_rows}, fh, ensure_ascii=False, indent=0)
    return state_rows, printed, fetched, sha, how, len(rows)


# ---------- the primary workbooks (results only: no contact columns in them) ----------

def txt(c):
    return re.sub(r"\s+", " ", str(c)).strip() if c is not None else ""


def num(c):
    """A figure; None for '-' (the race was not on that ballot) or an empty cell."""
    if c is None:
        return None
    if isinstance(c, bool):
        fail("a figure cell holds a true/false value")
    if isinstance(c, (int, float)):
        return int(c)
    s = txt(c).replace(",", "")
    if s in ("", "-"):
        return None
    if s.isdigit():
        return int(s)
    fail("a figure cell holds something other than a number")


def pkey(v):
    """A precinct number as printed; the spreadsheet stored some (01-02) as dates, read back as month-day."""
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{v.month:02d}-{v.day:02d}"
    return txt(v)


def left_of(row, j, floor=0):
    k = next((k for k in range(min(j, len(row) - 1), floor - 1, -1) if txt(row[k])), None)
    return (k, txt(row[k])) if k is not None else (None, "")


def columns(offices, parties, names):
    """[(column, office, party or None, label)] for every named column. A block of columns starts under an office heading
    or after an empty column; its party is the one party heading printed over the block, None when the sheet leaves it
    out (the caller settles that, never by guessing)."""
    blocks, cur = [], None
    for j in range(1, len(names)):
        label = txt(names[j])
        if j < len(offices) and txt(offices[j]):
            cur = None                                       # a new office heading starts a new block
        if not label:
            cur = None
            continue
        if cur is None:
            _oc, office = left_of(offices, j)
            if not OFFICE_ANY.match(office):
                fail(f"column {j + 1} of a results sheet sits under {office!r}, not an office heading")
            cur = {"office": CONT.sub("", office), "cols": []}
            blocks.append(cur)
        cur["cols"].append((j, label))
    out = []
    for b in blocks:
        lo, hi = b["cols"][0][0], b["cols"][-1][0]
        heads = [txt(parties[k]) for k in range(lo, min(hi, len(parties) - 1) + 1) if txt(parties[k])]
        if len(heads) > 1:
            fail(f"two party headings over one block of {b['office']} ({heads})")
        out += [(j, b["office"], heads[0] if heads else None, label) for j, label in b["cols"]]
    return out


def lev(a, b):
    """Edit distance between two short words."""
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def near(a, b):
    """Two spellings of one candidate within one race's block: the names fit, or the family names differ by one letter and
    the given names fit (Marilyn Connoly for Marilyn Connolly)."""
    ga, fa = name_parts(a)
    gb, fb = name_parts(b)
    return fits((ga, fa), (gb, fb)) or (bool(fa) and bool(fb) and lev(fa, fb) <= 1 and fits((ga, fb), (gb, fb)))


def summary_blocks(wb):
    """{sheet: {(office, party): {"labels": [...], "votes": {label: n}, "withdrawn": (name, county) or None}}} for every sheet
    whose title begins "Statewide " (the ballots-cast sheet aside), and the titles of the sheets not read."""
    out, skipped = OrderedDict(), []
    for ws in wb.worksheets:
        if not ws.title.startswith("Statewide ") or ws.title == "Statewide Total Ballots Cast":
            skipped.append(ws.title)
            continue
        g = [list(r) for r in ws.iter_rows(values_only=True)]
        oi = next((i for i, r in enumerate(g) if any(OFFICE_ANY.match(txt(c)) for c in r)), None)
        if oi is None:
            fail(f"the summary sheet {ws.title!r} has no office headings")
        ti = next((i for i in range(oi + 3, len(g)) if g[i] and txt(g[i][0]) == "Total"), None)
        if ti is None:
            fail(f"the summary sheet {ws.title!r} has no Total row")
        offices, parties, names, totals = g[oi], g[oi + 1], g[oi + 2], g[ti]
        notes = [(j, txt(c)) for r in g[ti + 1:] for j, c in enumerate(r) if txt(c).startswith("*")]
        blocks = OrderedDict()
        for j, office, party, label in columns(offices, parties, names):
            if party is None:
                fail(f"the summary sheet {ws.title!r} has no party heading over a block of {office}")
            b = blocks.setdefault((office, party), {"labels": [], "votes": {}, "withdrawn": None})
            if label in b["votes"]:
                fail(f"{ws.title}: {office} {party} lists {label!r} twice")
            v = num(totals[j]) if j < len(totals) else None
            if v is None:
                fail(f"{ws.title}: {office} {party} has no Total for column {j + 1}")
            b["labels"].append(label)
            b["votes"][label] = v
        for (office, party), b in blocks.items():
            if WITHDRAWN_LABEL in b["votes"]:
                said = [WITHDREW.match(n) for k, n in notes if CONT.sub("", left_of(offices, k)[1]) == office]
                said = [m for m in said if m]
                if len(said) == 1:
                    b["withdrawn"] = (said[0].group("name").strip(), said[0].group("county"))
        out[ws.title] = blocks
    return out, skipped


def precinct_cells(wb):
    """[(county, precinct, office, party or None, label, votes)] for every figure in every county's precinct rows, each
    column checked against the county's own Total row. Returns (cells, problems)."""
    out, problems = [], []
    for ws in wb.worksheets:
        county = ws.title
        g = [list(r) for r in ws.iter_rows(values_only=True)]
        heads = [i for i, r in enumerate(g) if r and txt(r[0]) == "Precinct"]
        if not heads:
            fail(f"the precinct results for {county} have no Precinct heading")
        for i in heads:
            cols = columns(g[i - 2], g[i - 1], g[i])
            colsum = Counter()
            total = None
            for r in g[i + 1:]:
                a = txt(r[0]) if r else ""
                if a == "Total":
                    total = r
                    break
                if a == "" or a.startswith("Precincts Continue") or a == "Precinct":
                    continue
                p = pkey(r[0])
                for j, office, party, label in cols:
                    v = num(r[j]) if j < len(r) else None
                    if v is None:
                        continue
                    out.append((county, p, office, party, label, v))
                    colsum[j] += v
            if total is None:
                fail(f"the precinct results for {county} have no Total row")
            for j, office, party, label in cols:
                tv = num(total[j]) if j < len(total) else None
                if (tv or 0) != colsum[j]:
                    problems.append(f"{county}, {office} {party} {label}: precincts add to {colsum[j]:,}, the county Total says {tv}")
    return out, problems


def settle_parties(cells, sheets):
    """A block a county's precinct sheet prints without a party heading is given the one party the summary has for that
    office that the same county's sheet does not label; the sums checked afterwards must confirm it. Returns (cells,
    [what was settled], [what could not be])."""
    in_summary = defaultdict(set)
    for blocks in sheets.values():
        for office, party in blocks:
            in_summary[office].add(party)
    labelled = defaultdict(set)
    for county, _p, office, party, _l, _v in cells:
        if party:
            labelled[(county, office)].add(party)
    settled, unsettled, pick = [], [], {}
    for county, office in sorted({(c, o) for c, _p, o, pa, _l, _v in cells if pa is None}):
        left = in_summary[office] - labelled[(county, office)]
        if len(left) == 1:
            pick[(county, office)] = left.pop()
            settled.append(f"{county} County's precinct sheet prints one block of {office} without a party heading; read as "
                           f"{pick[(county, office)]}, the one party the sheet does not otherwise label there")
        else:
            unsettled.append(f"{county}, {office}: a block with no party heading")
    return ([(c, p, o, pa or pick[(c, o)], l, v) for c, p, o, pa, l, v in cells if pa or (c, o) in pick], settled, unsettled)


def label_map(pre_labels, sum_labels):
    """{precinct label: summary label}, allowing another spelling of one candidate in a county's sheet; None when a label
    fits no summary label, or more than one."""
    out = {}
    for lab in pre_labels:
        if lab in sum_labels:
            out[lab] = lab
            continue
        if lab in TALLY or lab == WITHDRAWN_LABEL:
            return None
        got = [s for s in sum_labels if s not in TALLY and s != WITHDRAWN_LABEL and near(lab, s)]
        if len(got) != 1:
            return None
        out[lab] = got[0]
    return out


def ballots_cast(wb):
    """{county: {precinct: {party: ballots}}} from the precinct ballots-cast book, each county checked against its Total row."""
    out, problems = {}, []
    for ws in wb.worksheets:
        county = ws.title
        g = [list(r) for r in ws.iter_rows(values_only=True)]
        hi = next((i for i, r in enumerate(g) if r and txt(r[0]) == "Precinct"), None)
        if hi is None:
            fail(f"the ballots-cast sheet for {county} has no Precinct heading")
        head = [txt(c) for c in g[hi]]
        want = [h for h in head[1:] if h]
        rows, sums = {}, Counter()
        for r in g[hi + 1:]:
            a = txt(r[0]) if r else ""
            if a == "Total":
                for h in want:
                    if (num(r[head.index(h)]) or 0) != sums[h]:
                        problems.append(f"{county} ballots cast, {h}: precincts add to {sums[h]:,}, the Total says {num(r[head.index(h)])}")
                break
            if a == "" or a.startswith("Precincts Continue"):
                continue
            p = pkey(r[0])
            if p in rows:
                fail(f"the ballots-cast sheet for {county} lists precinct {p} twice")
            rows[p] = {h: num(r[head.index(h)]) or 0 for h in want}
            for h in want:
                sums[h] += rows[p][h]
        out[county] = rows
    return out, problems


def statewide_ballots(wb):
    g = [list(r) for r in wb["Statewide Total Ballots Cast"].iter_rows(values_only=True)]
    head = next(r for r in g if "Republican" in [txt(c) for c in r])
    total = next(r for r in g if r and txt(r[0]) == "Total")
    return {txt(c): num(total[j]) or 0 for j, c in enumerate(head) if txt(c)}


# ---------- the roster of sitting members: names, parties, districts, ids and term starts only ----------

def roster():
    con = sqlite3.connect(ROSTER_DB)
    seats = defaultdict(list)
    for bid, first, last, full, party, district, chamber, start in con.execute(
            "SELECT bioguide_id, first_name, last_name, official_full, party_name, district, chamber, term_start "
            "FROM legislators WHERE is_current = 1"):
        seats[(chamber, str(district))].append({"id": bid, "first": first or "", "last": last or "", "full": full or f"{first} {last}",
                                                "party": party, "chamber": chamber, "district": str(district), "start": start or ""})
    offices = {}
    for bid, first, last, full, office, label, party in con.execute(
            "SELECT bioguide_id, first_name, last_name, official_full, office, office_label, party_name FROM officials"):
        offices[office] = {"id": bid, "first": first or "", "last": last or "", "full": full or f"{first} {last}", "party": party,
                           "label": label, "chamber": None, "district": None}
    as_of = (con.execute("SELECT MAX(updated_at) FROM legislators").fetchone()[0] or "")[:10]
    con.close()
    return seats, offices, as_of


def counties():
    """{folded county name: (GEOID, 'Albany County')} for Wyoming from the Census Bureau's cartographic county file."""
    import shapefile                                   # pyshp
    z = zipfile.ZipFile(COUNTY_ZIP)
    base = next(n[:-4] for n in z.namelist() if n.endswith(".dbf"))
    rdr = shapefile.Reader(dbf=io.BytesIO(z.read(base + ".dbf")))
    fields = [f[0] for f in rdr.fields[1:]]
    out = {}
    for rec in rdr.iterRecords():
        rec = dict(zip(fields, rec))
        if str(rec.get("STATEFP")) == FIPS:
            out[fold(str(rec["NAME"]))] = (str(rec["GEOID"]), str(rec["NAMELSAD"]))
    if len(out) != 23:
        fail(f"the county file gives {len(out)} Wyoming counties, not 23")
    return out


def forms(parts):
    """A name's (given names, family name), and the same with initials written together (J.D. Williams as JD Williams)."""
    g, f = parts
    out = [(g, f)]
    if len(g) > 1 and all(len(w) == 1 for w in g):
        out.append((["".join(g)], f))
    return out


def person_fits(name, p):
    for cand in forms(name_parts(name)):
        for reg0 in ((fold(p["first"]).split(), fold(p["last"])), name_parts(p["full"])):
            for reg in forms(reg0):
                if reg[1] and fits(cand, reg) and not initials_clash(name, reg):
                    return True
    return False


def same_person(a, b):
    return fold(a) == fold(b) or fits(name_parts(a), name_parts(b))


def long_date(iso_date):
    d = dt.date.fromisoformat(iso_date)
    return f"{d.day} {d.strftime('%B')} {d.year}"


# =====================================================================================================================
# The county and local level: judges standing for retention, the county clerks' sample ballots and rosters
# =====================================================================================================================

LOCAL_SUB = "local"                                            # ballot_cache/wy/local/
CSV_URL = "https://sos.wyo.gov/Elections/Docs/2026/2026_WY_General_Election_Candidates.csv"
CLERKS_URL = "https://sos.wyo.gov/Elections/Docs/WYCountyClerks.pdf"      # the Secretary of State's directory of the 23 county clerks
TITLE5_URL = "https://wyoleg.gov/statutes/compress/title05.pdf"
TITLE22_URL = "https://wyoleg.gov/statutes/compress/title22.pdf"
PLACE_URL = "https://www2.census.gov/geo/docs/reference/codes2020/place/st56_wy_place2020.txt"
PLACE_HEAD = ["STATE", "STATEFP", "PLACEFP", "PLACENS", "PLACENAME", "TYPE", "CLASSFP", "FUNCSTAT", "COUNTIES"]
NONPARTISAN = "Nonpartisan office"
CAPITALS = "The ballot prints names in capitals; they are shown here in ordinary capitals."

SRC_CSV, SRC_TITLE5, SRC_PLACES = "wy-sos-2026-general-candidates-csv", "wy-legislature-statutes-title-5", "wy-census-2020-places"

# Each county clerk's election page, from the Secretary of State's directory of county clerks (its link for each county).
COUNTY_PAGES = {
    "Albany": "https://www.albanycountywy.gov/164/Elections",
    "Big Horn": "https://www.bighorncountywy.gov/departments/clerk/elections",
    "Campbell": "https://www.campbellcountywy.gov/867/Elections",
    "Carbon": "https://carboncountywy.gov/940/Elections",
    "Converse": "https://www.conversecountywy.gov/187/Elections",
    "Crook": "https://www.crookcounty.wy.gov/elected_officials/clerk/election/index.php",
    "Fremont": "https://fremontcountywy.gov/government/elections___voting.php",
    "Goshen": "https://www.goshencountywy.gov/189/Elections",
    "Hot Springs": "https://hscounty.com/elections",
    "Johnson": "https://www.johnsoncowy.gov/departments/%20elections",
    "Laramie": "https://www.laramiecountywy.gov/County-Government/Elected-Officials/County-Clerk/Elections",
    "Lincoln": "https://www.lincolncountywy.gov/government/clerk/elections_voting_information/index.php",
    "Natrona": "https://www.natronacounty-wy.gov/128/Elections",
    "Niobrara": "https://www.niobraracounty.org/_departments/_county_clerk/elections.asp",
    "Park": "https://parkcounty-wy.gov/county-elections/",
    "Platte": "https://www.plattecountywyoming.com/departments/Elections",
    "Sheridan": "https://www.sheridancountywy.gov/departments/elections/index.php",
    "Sublette": "https://www.sublettecountywy.gov/110/Election-Information",
    "Sweetwater": "https://www.sweetwatercountywy.gov/departments/county_clerk/elections.php",
    "Teton": "https://www.tetoncountywy.gov/268/Elections",
    "Uinta": "https://uintacountywy.gov/26/Elections",
    "Washakie": "https://www.washakiecountywy.gov/196/Elections",
    "Weston": "https://www.westongov.com/county-clerk/elections/",
}

# Where each clerk posts the November sample ballots: the page, the words its links to them carry and, where the words
# alone would also fit another file, what the link's address must hold. Only links on the page's own site are followed.
BALLOT_PAGES = OrderedDict([
    ("Sweetwater", {"page": "https://www.sweetwatercountywy.gov/departments/county_clerk/election_information/index.php", "pick": r"^2026 Sample Ballots$"}),
    ("Albany", {"page": "https://www.albanycountywy.gov/1705/Sample-Ballots", "pick": r"^\d+-\d+(-[A-Z ]+)?$"}),
    ("Park", {"page": "https://parkcounty-wy.gov/county-elections/2026-general-sample-ballots/", "pick": r"^Precinct \d+(-\d+)+$"}),
    ("Teton", {"page": "https://www.tetoncountywy.gov/3108/2026-Sample-Ballots", "pick": r"^\d+-\d+(-\d+)? \S"}),
    ("Uinta", {"page": "https://uintacountywy.gov/26/Elections", "pick": r"^General Election Sample Ballots$"}),
    ("Lincoln", {"page": "https://www.lincolncountywy.gov/government/clerk/elections_voting_information/election_results.php",
                 "pick": r"^\d+-\d+ [A-Z0-9/ ]+$", "href": r"/2026/(?!OFFICIAL)[^/]+\.pdf"}),
    ("Carbon", {"page": "https://carboncountywy.gov/940/Elections", "pick": r"^Sample Ballots For Current Elections$"}),
    ("Converse", {"page": "https://www.conversecountywy.gov/558/2026-General-Election-Sample-Ballots", "pick": r"^\d+(-\d+)+ .* SAMPLE$"}),
    ("Goshen", {"page": "https://www.goshencountywy.gov/376/General-Election-Sample-Ballots", "pick": r"^\d+(-\d+)+ .* 2026 General Ballot \(PDF\)$"}),
    ("Sublette", {"page": "https://www.sublettecountywy.gov/110/Election-Information", "pick": r"^Sample Ballots 2026 General$"}),
    ("Platte", {"page": "https://www.plattecountywyoming.com/departments/Elections/candidateinformation", "pick": r"^2026 General Election Sample Ballots$"}),
    ("Washakie", {"page": "https://www.washakiecountywy.gov/196/Elections", "pick": r"^2026 GENERAL SAMPLE BALLOTS$"}),
    ("Crook", {"page": "https://www.crookcounty.wy.gov/elected_officials/clerk/election/sample_ballots.php", "pick": r"^Sample Ballots$", "href": r"\.pdf"}),
    ("Weston", {"page": "https://www.westongov.com/county-clerk/elections/", "pick": r"^\d+-\d+$", "href": r"-gen\.pdf$"}),
    ("Hot Springs", {"page": "https://hscounty.com/elections", "pick": r"^2026 General Election SAMPLE BALLOTS$"}),
    # Johnson County's page links a shared folder; the general election ballots are the folder's files named "General ... Sample.pdf"
    ("Johnson", {"page": COUNTY_PAGES["Johnson"], "pick": r"^2026 Election Sample Ballots$", "folder_files": r"^General .* Sample\.pdf$"}),
    # The clerks below had posted no November sample ballot when this was written. Their pages are looked at for any link
    # that says "sample ballot"; a file that is not a November 3, 2026 general election ballot is set aside.
    ("Campbell", {"page": COUNTY_PAGES["Campbell"], "pick": r"(?i)\bsample ballots?\b"}),
    ("Fremont", {"page": COUNTY_PAGES["Fremont"], "pick": r"(?i)\bsample ballots?\b"}),
    ("Sheridan", {"page": COUNTY_PAGES["Sheridan"], "pick": r"(?i)\bsample ballots?\b"}),
    ("Big Horn", {"page": COUNTY_PAGES["Big Horn"], "pick": r"(?i)\bsample ballots?\b"}),
    ("Niobrara", {"page": COUNTY_PAGES["Niobrara"], "pick": r"(?i)\bsample ballots?\b"}),
])
# Laramie County's general election rosters: the page, and the three workbooks' link words (the links that say "general").
LARAMIE_PAGE = "https://www.laramiecountywy.gov/County-Government/Elected-Officials/County-Clerk/Elections/Candidates"
LARAMIE_BOOKS = OrderedDict([("county", r"^County Offices\(XLSX"), ("municipal", r"^Municipal Offices\(XLSX"),
                             ("districts", r"^School Board, Conservation District, & Fire District\(XLSX")])
LARAMIE_KEEP = ("Municipality", "Office", "Party", "Ballot Name", "Withdrawn Date")     # the only columns of a roster ever read
# What is known of the counties whose lists are not loaded, for the reason given in sl_gaps (the clerk's page is the address).
NOT_POSTED = "The county clerk publishes this county's list of November 3 candidates. "
WHY_NOT = {
    "Campbell": NOT_POSTED + "The clerk had posted a list of special district candidates without column headings, and no sample ballot or full "
                "November roster, when this was built, so the county is not loaded yet.",
    "Fremont": NOT_POSTED + "The clerk had posted a list of school and special district filings and no November sample ballot when this was built, "
               "so the county is not loaded yet.",
    "Sheridan": NOT_POSTED + "The clerk had posted a list of general election filings in a layout this loader does not read yet, and no sample "
                "ballot, when this was built, so the county is not loaded yet.",
    "Big Horn": NOT_POSTED + "The clerk had posted its general election candidate list in a layout this loader does not read yet, and no November "
                "sample ballot, when this was built, so the county is not loaded yet.",
    "Niobrara": NOT_POSTED + "The clerk had posted its general election roster in a layout this loader does not read yet, and no November sample "
                "ballot, when this was built, so the county is not loaded yet.",
}

NUMBER_WORDS = {"ONE": 1, "TWO": 2, "THREE": 3, "FOUR": 4, "FIVE": 5, "SIX": 6, "SEVEN": 7, "EIGHT": 8, "NINE": 9, "TEN": 10}
NUMBER_NAMES = {v: k.lower() for k, v in NUMBER_WORDS.items()}
ORDINAL_WORDS = ["first", "second", "third", "fourth", "fifth", "sixth", "seventh", "eighth", "ninth"]
# the page builder's own last check is a little wider than the trial check's (it also stops at Court and Place)
BUILDER_STREET = re.compile(r"\b\d{1,6}\s+(?:[NSEW]\.?\s+)?[A-Za-z0-9.' -]{1,40}?\s(?:St|Street|Ave|Avenue|Rd|Road|Blvd|Boulevard|Dr|Drive|Ln|Lane|Way|Ct|Court|"
                            r"Cir|Circle|Pkwy|Parkway|Hwy|Highway|Trl|Trail|Pl|Place|Ter|Terrace)\b\.?", re.I)


def lslug(text):
    """Letters, digits and hyphens, for ids: "Little Snake River Conservation" -> little-snake-river-conservation."""
    t = unicodedata.normalize("NFKD", text or "")
    return re.sub(r"[^a-z0-9]+", "-", "".join(ch for ch in t if not unicodedata.combining(ch)).lower()).strip("-")


def lfail(msg):
    raise SystemExit(f"Wyoming (local races): {msg}")


def squash(text):
    """Letters and digits only, lower case, accents folded: for comparing names of places, never for showing them."""
    t = unicodedata.normalize("NFKD", text or "")
    return re.sub(r"[^a-z0-9]", "", "".join(ch for ch in t if not unicodedata.combining(ch)).lower())


def and_names(items):
    items = list(items)
    if not items:
        return ""
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]


def tidy_url(url):
    """An address with its spaces written as %20 (several county sites print them raw)."""
    return quote(url, safe=":/?=&%#+,;@()!'*~$-._")


def polite_get(url, accept="*/*", timeout=180):
    """One request through states/net.py (its honest User-Agent, its plain Accept header), a second after the last one;
    asked again twice, more slowly, if the line drops or the site refuses, and then not again (a site that keeps
    refusing is left alone)."""
    last = None
    for attempt in range(3):
        time.sleep(1.0 + 4.0 * attempt)
        try:
            return net.get(tidy_url(url), timeout=timeout, accept=accept)
        except (HTTPError, URLError, OSError) as e:
            last = e
    raise last


SMALL_CAPS = {"WM": "Wm", "ST": "St", "ST.": "St.", "DR": "Dr", "DR.": "Dr.", "JR": "Jr", "JR.": "Jr.", "SR": "Sr", "SR.": "Sr."}


def ordinary(caps):
    """Words a ballot prints in capitals, in ordinary capitals: JON R CONRAD -> Jon R Conrad; MAUREEN "MO" MURPHY ->
    Maureen "Mo" Murphy; SARAH MCKERGOW -> Sarah McKergow; BRUNSKI-DECKERT -> Brunski-Deckert; JOHN (JACK) COBB -> John
    (Jack) Cobb. A word that is not all capitals is kept as printed; two letters with no vowel are initials run together
    and stay in capitals (TJ, JD), as do initials written with their stops (J.W.); II, III and IV stay."""
    out = []
    for i, w in enumerate((caps or "").split()):
        core = re.sub(r"[^A-Za-z]", "", w)
        if w.upper() in SMALL_CAPS and (i or w.upper() in ("ST", "ST.", "WM", "DR", "DR.")):
            out.append(SMALL_CAPS[w.upper()])
        elif w != w.upper() or not core:
            out.append(w)
        elif i and re.fullmatch(r"[IVX]+,?", w) and w.rstrip(",") in ("II", "III", "IV", "V"):
            out.append(w)
        elif re.fullmatch(r"(?:[A-Z]\.){2,}|(?:[A-Z]\.)+[A-Z]", w):
            out.append(w)                                          # initials written with their stops: J.W. stays J.W.
        elif len(core) == 2 and not re.search(r"[AEIOUY]", core) and not re.search(r"[^A-Za-z]", w.strip("\"'()")):
            out.append(w)
        else:
            parts = []
            for p in w.split("-"):
                p2 = p.lower()
                p2 = re.sub(r"[a-z]", lambda m: m.group(0).upper(), p2, count=1)
                p2 = re.sub(r"^(\W*)Mc([a-z])", lambda m: m.group(1) + "Mc" + m.group(2).upper(), p2)
                p2 = re.sub(r"^(\W*[A-Z])'([a-z])", lambda m: m.group(1) + "'" + m.group(2).upper(), p2)
                p2 = re.sub(r"/([a-z])", lambda m: "/" + m.group(1).upper(), p2)
                parts.append(p2)
            out.append("-".join(parts))
    return " ".join(out)


def alike(a, b):
    """Two lists' spellings of one person: the same letters (MICHAEL R. MANN and MICHAEL R MANN), or the same family name
    and the same first given name, or an initial for it (TANYA S. EVANS and TANYA EVANS). A middle initial alone is not
    enough: Jason M. and James M. are two names."""
    if fold(a) == fold(b):
        return True
    (ga, la), (gb, lb) = name_parts(a), name_parts(b)
    if la != lb or not la or not ga or not gb:
        return False
    return ga[0] == gb[0] or (1 in (len(ga[0]), len(gb[0])) and ga[0][0] == gb[0][0])


def same_names(a, b):
    """Two lists' spellings of one set of candidates, paired one to one. Returns the pairs, or None when the two sets
    are not the same people."""
    a, b = list(a), list(b)
    if len(a) != len(b):
        return None
    pairs, left = [], list(b)
    for x in a:
        hit = [y for y in left if fold(x) == fold(y)] or [y for y in left if alike(x, y)]
        if len(hit) != 1:
            return None
        pairs.append((x, hit[0]))
        left.remove(hit[0])
    return pairs


# ---------- the judicial districts (W.S. 5-3-101), for the counties that vote on a district or circuit judge ----------

def judicial_districts(folder, cmap, say=print, refresh=False):
    """({district number: [county GEOIDs]}, path) from W.S. 5-3-101 in the Legislature's own file of Title 5 (a law
    text with no personal details; kept whole). The section names every county once, in nine districts; anything else
    stops the loader."""
    path = os.path.join(folder, "wy_statutes_title05.pdf")
    if refresh or not os.path.exists(path):
        try:
            data = polite_get(TITLE5_URL)
            if not data.startswith(b"%PDF"):
                raise ValueError("the address answered without a PDF")
            os.makedirs(folder, exist_ok=True)
            with open(path + ".part", "wb") as fh:
                fh.write(data)
            os.replace(path + ".part", path)
        except Exception as e:  # noqa: BLE001
            if not os.path.exists(path):
                lfail(f"the Legislature's file of Title 5 (the judicial districts) could not be read ({type(e).__name__}); nothing was changed")
            say(f"    Wyoming (local): Title 5 could not be read again ({type(e).__name__}); using the copy on disk")
    pdf = PDF(open(path, "rb").read())
    lines = [join(rs) for page, res in pdf.pages() for _y, rs in pdf_rows(pdf, page, res)]
    lines = [t for t in lines if t]
    start = next((i for i, t in enumerate(lines) if t.startswith("5-3-101.")), None)
    if start is None:
        lfail("W.S. 5-3-101 is not in the Legislature's file of Title 5")
    end = next((i for i in range(start + 1, len(lines)) if re.match(r"\(b\)|5-3-10[2-9]\.", lines[i])), len(lines))
    text = " ".join(lines[start:end])
    out, used = {}, []
    for m in re.finditer(r"(?:The count(?:y|ies) of ([A-Z][A-Za-z ,]+?)|([A-Z][a-z]+(?: [A-Z][a-z]+)?) county) (?:is|are) the (\w+) judicial district", text):
        names = [n.strip() for n in re.split(r",| and ", m.group(1) or m.group(2)) if n.strip()]
        if m.group(3) not in ORDINAL_WORDS:
            lfail(f"W.S. 5-3-101 names a judicial district this loader cannot number ({m.group(3)})")
        ids = []
        for n in names:
            if fold(n) not in cmap:
                lfail("W.S. 5-3-101 names a county the Census county file does not have")
            ids.append(cmap[fold(n)][0])
        out[ORDINAL_WORDS.index(m.group(3)) + 1] = sorted(ids)
        used += ids
    if sorted(out) != list(range(1, 10)) or sorted(used) != sorted(g for g, _n in cmap.values()):
        lfail(f"W.S. 5-3-101 was not read as nine judicial districts holding each of the 23 counties once ({len(out)} districts, {len(used)} counties)")
    return out, path


# ---------- the Secretary of State's candidates file: five columns of it, for the judges standing for retention ----------

CSV_KEEP = ("Election", "Office Sought", "Party Affiliation", "Ballot Name", "Date Withdrawn")
JUDGE_OFFICES = (
    (re.compile(r"^SC-\d+ - JUSTICE OF THE SUPREME COURT$"), "supreme"),
    (re.compile(r"^CHC-\d+ - JUDGE OF THE CHANCERY COURT$"), "chancery"),
    (re.compile(r"^DC-\d+ - JUDGE ([A-Z]) OF THE DISTRICT COURT OF THE ([A-Z]+) JUDICIAL DISTRICT$"), "district"),
    (re.compile(r"^CC-\d+ - ([A-Z]+) JUDICIAL DISTRICT, CIRCUIT COURT JUDGE$"), "circuit"),
)


def sos_csv(folder, say=print, refresh=False):
    """The Secretary of State's general election candidates file, cut down in memory to five columns chosen by their
    headings (the election, the office, the party, the ballot name and a withdrawal date) and kept that way, as JSON,
    with the day it was fetched and the fingerprint of the file as it came. Its mailing address, city, telephone,
    e-mail and web address columns are never read into anything. A kept cell that reads like contact details is blanked
    and counted. Asked for again only with refresh."""
    path = os.path.join(folder, "wy_2026_general_candidates_csv.json")
    old = json.load(open(path, encoding="utf-8")) if os.path.exists(path) else None
    if old and not refresh:
        return old
    try:
        raw = polite_get(CSV_URL)
    except Exception as e:  # noqa: BLE001
        if old:
            say(f"    Wyoming (local): the candidates file could not be read again ({type(e).__name__}); using the columns kept on {old['fetched']}")
            return old
        lfail(f"the Secretary of State's candidates file could not be read ({type(e).__name__}); nothing was changed")
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", "replace")
    if text.lstrip()[:1] == "<":                                   # a page, not the file (the site answers some requests that way)
        if old:
            say(f"    Wyoming (local): the candidates file's address answered with a page, not the file; using the columns kept on {old['fetched']}")
            return old
        lfail("the Secretary of State's candidates file's address answered with a page, not the file; nothing was changed. Re-run later.")
    table = list(csv.reader(io.StringIO(text)))
    head = [h.strip() for h in table[0]] if table else []
    if any(head.count(k) != 1 for k in CSV_KEEP):
        lfail("the Secretary of State's candidates file no longer has the column headings this loader reads; stopped (no row is printed)")
    idx = [head.index(k) for k in CSV_KEEP]
    rows, blanked = [], 0
    for n, r in enumerate(table[1:], start=2):
        if not any(c.strip() for c in r):
            continue
        if len(r) != len(head):
            lfail(f"line {n} of the Secretary of State's candidates file has {len(r)} cells, not {len(head)}; stopped (the line is not printed)")
        cut = [r[i].strip() for i in idx]
        for j, v in enumerate(cut):
            if v and contact_like(v, True):
                cut[j], blanked = "", blanked + 1
        rows.append(cut)
    kept = {"title": "2026 General Election Candidates (data file)", "url": CSV_URL, "fetched": dt.date.today().isoformat(),
            "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw), "columns_in_file": len(head), "columns": list(CSV_KEEP),
            "blanked": blanked, "rows": rows}
    del raw, text, table
    os.makedirs(folder, exist_ok=True)
    with open(path + ".part", "w", encoding="utf-8") as fh:
        json.dump(kept, fh, ensure_ascii=False, indent=0)
    os.replace(path + ".part", path)
    return kept


def retention_rows(kept, districts, cmap, say=print):
    """The judges standing for retention as rows ready to write: (races, candidates, places, the judges by court for the
    check against the county ballots, withdrawn left off, rows that are not judges)."""
    full = {g: n for g, n in cmap.values()}
    races, cands, places, judges, gone, others = [], [], {}, [], 0, 0
    seen = set()
    for n, (election, office, party, name, withdrawn) in enumerate(kept["rows"], start=2):
        if election != "2026 GENERAL ELECTION":
            lfail(f"line {n} of the candidates file is not for the 2026 general election")
        hit = next(((m, kind) for rx, kind in JUDGE_OFFICES for m in [rx.match(office)] if m), None)
        if not hit:
            if not party:
                lfail(f"line {n} of the candidates file is an office without a party that this loader does not know; stopped")
            others += 1
            continue
        if party:
            lfail(f"line {n} of the candidates file gives a party to a judge standing for retention; stopped")
        if not name:
            lfail(f"line {n} of the candidates file has no ballot name for a judge; stopped")
        if withdrawn:
            gone += 1
            continue
        m, kind = hit
        family = re.sub(r"[^A-Z]", "", name_parts(name)[1].upper())
        if kind == "supreme":
            rid, okind, shown = f"2026-{STATE}-SCRET-{family}", "supreme_court_retention", "Justice of the Supreme Court (retention vote)"
            jur, jid, cids, dist, seat = NAME, FIPS, None, None, None
            note = ("A retention vote: voters answer Yes or No on keeping this justice in office. No one runs against the justice and no party "
                    "is printed. Every voter in the state votes on it (W.S. 22-2-105).")
            who = "Standing for retention as a sitting justice."
        elif kind == "chancery":
            rid, okind, shown = f"2026-{STATE}-CHCRET-{family}", "chancery_court_retention", "Judge of the Chancery Court (retention vote)"
            jur, jid, cids, dist, seat = NAME, FIPS, None, None, None
            note = ("A retention vote: voters answer Yes or No on keeping this judge in office. No one runs against the judge and no party is "
                    "printed. Every voter in the state votes on it (W.S. 5-13-106).")
            who = "Standing for retention as the sitting judge."
        else:
            word = (m.group(2) if kind == "district" else m.group(1)).lower()
            if word not in ORDINAL_WORDS:
                lfail(f"line {n} of the candidates file names a judicial district this loader cannot number; stopped")
            d = ORDINAL_WORDS.index(word) + 1
            jur, jid, cids, dist = f"{word.title()} Judicial District", f"{STATE}-JD{d}", districts[d], str(d)
            where = and_names(full[g].replace(" County", "") for g in cids) + (" County" if len(cids) == 1 else " counties")
            places[("judicial", jid)] = ("judicial", jid, jur, json.dumps(cids), SRC_TITLE5)
            if kind == "district":
                rid, okind, shown = f"2026-{STATE}-DCRET{d}-{family}", "district_court_retention", "Judge of the District Court (retention vote)"
                seat = f"Judge {m.group(1)}"
                note = ("A retention vote: voters answer Yes or No on keeping this judge in office. No one runs against the judge and no party "
                        f"is printed. Voted on in {where}, which {'is' if len(cids) == 1 else 'are'} the {jur} (W.S. 5-3-101 and 22-2-105).")
            else:
                rid, okind, shown, seat = f"2026-{STATE}-CCRET{d}-{family}", "circuit_court_retention", "Circuit Court Judge (retention vote)", None
                note = ("A retention vote: voters answer Yes or No on keeping this judge in office. No one runs against the judge and no party "
                        f"is printed. A circuit judge is voted on throughout the circuit, which has the lines of the {jur}: {where} "
                        "(W.S. 5-9-102 and 22-2-105).")
            who = "Standing for retention as the sitting judge."
        if rid in seen:
            lfail(f"two judges share the race id {rid}; stopped")
        seen.add(rid)
        races.append((rid, STATE, "court", okind, shown, jur, jid, json.dumps(cids) if cids else None, dist, seat, 0, 0, None, name, None, GENERAL, note))
        cands.append((rid, "general", GENERAL, name, NONPARTISAN, "N", None, 1, 0, None, None, None, None, SRC_CSV, who))
        judges.append({"name": name, "kind": kind, "counties": cids})
    return races, cands, list(places.values()), judges, gone, others


# ---------- the Census Bureau's place codes: Wyoming's incorporated cities and towns ----------

def census_places(folder, cmap, say=print):
    """({squashed bare name: [(place code, name with its kind word, [county GEOIDs])]}, path, rows read) for Wyoming's
    incorporated cities and towns, from the Census Bureau's 2020 place codes file (names, codes and counties only; kept
    whole and asked for once)."""
    path = os.path.join(folder, "st56_wy_place2020.txt")
    try:
        net.download(PLACE_URL, path, 3650, tries=3, say=say)
    except Exception as e:  # noqa: BLE001
        lfail(f"the Census Bureau's place codes file could not be fetched ({type(e).__name__}); nothing was changed")
    by_full = {fold(full): geoid for geoid, full in cmap.values()}
    out, n = {}, 0
    with open(path, encoding="utf-8", errors="replace") as fh:
        if fh.readline().rstrip("\r\n").split("|") != PLACE_HEAD:
            lfail("st56_wy_place2020.txt: the header is not the one this loader was checked against")
        for ln, line in enumerate(fh, start=2):
            if not line.strip():
                continue
            f = line.rstrip("\r\n").split("|")
            if len(f) != len(PLACE_HEAD) or f[1] != FIPS or not re.fullmatch(r"\d{5}", f[2]):
                lfail(f"st56_wy_place2020.txt: line {ln} does not fit the header")
            n += 1
            if f[5] != "INCORPORATED PLACE":
                continue
            cids = [by_full.get(fold(c)) for c in f[8].split("~~~")]
            if not all(cids) or not re.search(r" (city|town)$", f[4]):
                lfail(f"st56_wy_place2020.txt: line {ln} names a county the county file does not have, or a place that is not a city or town")
            out.setdefault(squash(re.sub(r" (city|town)$", "", f[4])), []).append((f[2], f[4], sorted(cids)))
    return out, path, n


# ---------- reading a county's sample ballots (the counties' ballot-printing system lays them all out one way) ----------

FIRST_SECTION = re.compile(r"^FEDERAL( OFFICES?)?$")
FIRST_TITLE = re.compile(r"^(FOR )?UNITE[DS] STATES SENATOR$")       # where a ballot printed without section headings begins
CHOICE = re.compile(r"^[\"“]?(YES|NO|FOR|AGAINST)\b")
VOTE_FOR = re.compile(r"^VOTE FOR\b")
TERM_WORD = re.compile(r"^(UNEXPIRED|TERM|YEARS?|YR|REGULAR|FOR|A|OF|TO|BE|DETERMINED|ONE|TWO|THREE|FOUR|FIVE|SIX|EIGHT|\(\d\)|\d)$")
COURT_WORD = re.compile(r"\b(COURT|JUDGES?|JUSTICES?|JUSTICE\(S\)|JUDICIAL)\b")
RETAIN = re.compile(r"SHALL (?:(JUSTICE|JUDGE) )?(.+?) BE RETAINED\b", re.I)     # ... IN OFFICE? / ... AS JUDGE OF THE CHANCERY COURT?
BALLOT_TITLE = re.compile(r"^(OFFICIAL|SAMPLE) GENERAL ELECTION BALLOT\b")
FOOT = re.compile(r"^(VOTE BOTH SIDES|END OF BALLOT|TURN BALLOT OVER|CONTINUE VOTING|BALLOT CONTINUES)")
MACHINE = re.compile(r"^Typ:\d+ Seq:\d+ Spl:\d+$")
BALLOT_PARTIES = {"REPUBLICAN": "Republican", "DEMOCRATIC": "Democratic", "LIBERTARIAN": "Libertarian", "CONSTITUTION": "Constitution",
                  "INDEPENDENT": "Independent"}
HEAD_SKIP = re.compile(r"^((OFFICIAL|SAMPLE) GENERAL ELECTION BALLOT|.+ COUNTY(, WYOMING)?|NOVEMBER 3, 2026|ELECTION JUDGE INITIALS.*|"
                       r"INSTRUCTIONS?( TO VOTERS?)?)$", re.I)
INSTRUCTIONS = re.compile(r"^INSTRUCTIONS?( TO VOTERS?)?$")
PITCH = 181.5                                                        # points from one column of the full-size ballot to the next


class BallotError(Exception):
    """A sample ballot this reader cannot read; the message names the file and the check, never a line of the file.
    why: "scan" (a picture with no text), "layout" (text this reader cannot follow), "not-general" (no general election
    ballot in the file), "other-election" (a general election ballot of another year), "not-pdf" or "contact"."""

    def __init__(self, msg, why="layout"):
        super().__init__(msg)
        self.why = why


def is_term(line):
    """A line made only of the words of a term of office: FOUR (4) YEAR TERM; UNEXPIRED TERM OF TWO (2) YEARS;
    UNEXPIRED; TERM TO BE DETERMINED."""
    words = re.sub(r"[,.-]", " ", line).split()
    return bool(words) and all(TERM_WORD.match(w) for w in words) and (any(w.startswith(("YEAR", "TERM")) for w in words) or words == ["UNEXPIRED"])


def glyph_runs(pdf, page, res):
    """Every character on a page as its own run (x, y, size, character, x_end), spaces included: pdftext.page_runs taken
    apart, for a file that draws its text more than once (see page_texts). The font, its size, the character and word
    spacing, the horizontal scale, the leading and the rise belong to the graphics state, so q saves them and Q puts
    them back: a file that redraws pieces of its text inside saved states (under a SAMPLE mark laid across the page)
    reads out of step without that."""
    from ballot.pdftext import Font, Ref, _mul, _ops
    fonts, runs = {}, []
    fres = pdf.get((pdf.get(res) or {}).get("Font")) or {}
    contents = pdf.get(page.get("Contents"))
    parts = contents if isinstance(contents, list) else [page.get("Contents")]
    data = b"\n".join(pdf.stream(p) or b"" for p in parts if isinstance(p, Ref))
    ctm, saved = [1, 0, 0, 1, 0, 0], []
    tm = tlm = [1, 0, 0, 1, 0, 0]
    font, size, tc, tw, th, tl, rise = None, 1.0, 0.0, 0.0, 1.0, 0.0, 0.0

    def show(s):
        nonlocal tm
        if font is None or not isinstance(s, (bytes, bytearray)):
            return
        codes = [int.from_bytes(s[k:k + 2], "big") for k in range(0, len(s) - 1, 2)] if font.two else list(s)
        for code in codes:
            ch = font.map.get(code, "") if font.two else (font.map.get(code) or bytes([code & 0xFF]).decode("cp1252", "replace"))
            trm = _mul([size * th, 0, 0, size, 0, rise], _mul(tm, ctm))
            adv = (font.width(code) * size + tc + (tw if (not font.two and code == 32) else 0)) * th
            tm = _mul([1, 0, 0, 1, adv, 0], tm)
            x1 = _mul([size * th, 0, 0, size, 0, rise], _mul(tm, ctm))[4]
            if ch:
                runs.append((trm[4], trm[5], abs(trm[3]) or size, " " if not ch.strip() else ch, x1))

    for op, a in _ops(data):
        if op == "q":
            saved.append((ctm[:], font, size, tc, tw, th, tl, rise))
        elif op == "Q":
            ctm, font, size, tc, tw, th, tl, rise = saved.pop() if saved else ([1, 0, 0, 1, 0, 0], font, size, tc, tw, th, tl, rise)
        elif op == "cm" and len(a) == 6:
            ctm = _mul([float(x) for x in a], ctm)
        elif op == "BT":
            tm = tlm = [1, 0, 0, 1, 0, 0]
        elif op == "Tf" and len(a) == 2:
            name = str(a[0])
            if name not in fonts:
                fonts[name] = Font(pdf, fres.get(name))
            font, size = fonts[name], float(a[1])
        elif op == "Tc" and a:
            tc = float(a[0])
        elif op == "Tw" and a:
            tw = float(a[0])
        elif op == "Tz" and a:
            th = float(a[0]) / 100
        elif op == "TL" and a:
            tl = float(a[0])
        elif op == "Ts" and a:
            rise = float(a[0])
        elif op in ("Td", "TD") and len(a) == 2:
            tx, ty = float(a[0]), float(a[1])
            if op == "TD":
                tl = -ty
            tlm = _mul([1, 0, 0, 1, tx, ty], tlm)
            tm = tlm
        elif op == "Tm" and len(a) == 6:
            tm = tlm = [float(x) for x in a]
        elif op == "T*":
            tlm = _mul([1, 0, 0, 1, 0, -tl], tlm)
            tm = tlm
        elif op == "Tj" and a:
            show(a[-1])
        elif op in ("'", '"') and a:
            tlm = _mul([1, 0, 0, 1, 0, -tl], tlm)
            tm = tlm
            if op == '"' and len(a) == 3:
                tw, tc = float(a[0]), float(a[1])
            show(a[-1])
        elif op == "TJ" and a and isinstance(a[-1], list):
            for item in a[-1]:
                if isinstance(item, (bytes, bytearray)):
                    show(item)
                elif isinstance(item, (int, float)):
                    tm = _mul([1, 0, 0, 1, -float(item) / 1000.0 * size * th, 0], tm)
    return runs


SAME_SPOT = 0.3                                                      # points: a letter drawn again this near the same letter is that letter


def letters_once(runs):
    """(the characters of a page with each kept once, how many were drawn again on the spot of the same character). A
    letter far larger than the page's type (a SAMPLE mark across the page) is dropped."""
    kept, grid, again = [], defaultdict(list), 0
    for r in runs:
        gx, gy = int(r[0] // 2), int(r[1] // 2)
        if any(abs(o[0] - r[0]) <= SAME_SPOT and abs(o[1] - r[1]) <= SAME_SPOT
               for dx in (-1, 0, 1) for dy in (-1, 0, 1) for o in grid[(gx + dx, gy + dy, r[3])]):
            again += r[3] != " "
            continue
        grid[(gx, gy, r[3])].append(r)
        kept.append(r)
    sizes = Counter(round(r[2], 1) for r in kept if r[3] != " ")
    main = sizes.most_common(1)[0][0] if sizes else 0
    return [r for r in kept if r[2] <= 2.5 * main], again


def page_texts(pdf, pages):
    """([the runs of each page], drawn_twice). Some clerks post a ballot printed to a file by a program that draws its
    letters more than once on the same spot (every letter two or three times, or again wherever a SAMPLE mark lies
    across them); such a file is read letter by letter and each letter kept once, with its own spaces. A file is read
    that way when a twentieth or more of its letters are drawn again on the spot of the same letter."""
    by_letter = [letters_once(glyph_runs(pdf, page, res)) for page, res in pages]
    inked = sum(sum(1 for r in kept if r[3] != " ") for kept, _again in by_letter)
    again = sum(a for _kept, a in by_letter)
    if not inked or again < 0.05 * inked:
        return [[r for r in page_runs(pdf, page, res) if r[3].strip()] for page, res in pages], False
    return [kept for kept, _again in by_letter], True


def baseline_lines(runs, gap=None):
    """Runs grouped by baseline and type size, left to right: [(y, x0, size, text, x1)]. With gap set, a baseline is cut
    into pieces wherever two runs stand further apart than gap. A line of spaces alone is no line."""
    lines = []
    for r in sorted(runs, key=lambda r: (-r[1], round(r[2], 1), r[0])):
        if lines and abs(lines[-1][0] - r[1]) <= 0.8 and abs(lines[-1][1] - r[2]) < 0.6 and (gap is None or r[0] - lines[-1][3] <= gap):
            lines[-1][2].append(r)
            lines[-1][3] = max(lines[-1][3], r[4])
        else:
            lines.append([r[1], r[2], [r], r[4]])
    out = []
    for y, size, rs, end in lines:
        inked = [q for q in rs if q[3].strip()]
        if inked:
            out.append((y, min(q[0] for q in inked), size, join(rs), end))
    return out


def column_edges(texts):
    """[(names, write-ins, choices)]: for each column of the ballot, the left edge of the candidates' names, of the
    WRITE-IN lines and of the YES/NO and FOR/AGAINST words, over the whole file. Counties set these a point or two apart
    in different ways, so each is measured: the names' edge is where the party lines start, the choices' edge where a
    bare YES, NO, FOR or AGAINST starts. A column with no party line anywhere takes the offset the other columns show."""
    wx, px, cx = Counter(), Counter(), Counter()
    for runs in texts:
        for _y, x0, _size, text, _x1 in baseline_lines(runs, gap=4.0):
            if text == "WRITE-IN":
                wx[round(x0, 1)] += 1
            elif text in BALLOT_PARTIES:
                px[round(x0, 1)] += 1
            elif text in ("YES", "NO", "FOR", "AGAINST"):
                cx[round(x0, 1)] += 1
    cols = []
    for x in sorted(wx):
        if not cols or x - cols[-1] >= 20:
            cols.append(x)

    def near(counter, x):
        got = [p for p in counter if abs(p - x) <= 3.0]
        return max(got, key=lambda p: counter[p]) if got else None

    names = [near(px, w) for w in cols]
    choices = [near(cx, w) for w in cols]
    dn = [round(n - w, 1) for n, w in zip(names, cols) if n is not None]
    dc = [round(c - w, 1) for c, w in zip(choices, cols) if c is not None]
    if (dn and max(dn) - min(dn) > 0.25) or (dc and max(dc) - min(dc) > 0.25):
        raise BallotError("the columns do not set their names or their choices the same way")
    scale = (cols[1] - cols[0]) / PITCH if len(cols) > 1 else 1.0
    return [(w + (Counter(dn).most_common(1)[0][0] if dn else 0.0), w, w + (Counter(dc).most_common(1)[0][0] if dc else -1.8 * scale)) for w in cols]


def ballot_pages(texts, edges, s):
    """Every page as {"n", "titled", "head": [lines], "cols": [[lines] per column], "closing": [y of the page's closing
    words]}; a line is (y, x0, size, text, x1). The head is everything above the FEDERAL OFFICES heading on a page that
    carries the ballot's title (above the first contest, on a ballot printed without section headings). s is the size
    of the file's ballot against the full-size one."""
    left = [min(e) - 25.5 * s for e in edges]
    out = []
    for n, all_runs in enumerate(texts, 1):
        runs = [r for r in all_runs if r[0] >= left[0] - 3.0 * s]    # the timing marks' numbers sit in the margin
        cols = [[] for _ in edges]
        for r in runs:
            k = max([i for i, lft in enumerate(left) if r[0] >= lft], default=0)
            cols[k].append(r)
        cols = [baseline_lines(rs, gap=12 * s) for rs in cols]
        whole = baseline_lines(runs)
        titled = any(BALLOT_TITLE.search(ln[3]) for ln in whole)
        top = None
        if titled:
            ys = [ln[0] for ln in cols[0] if FIRST_SECTION.match(ln[3])] or [ln[0] for ln in cols[0] if FIRST_TITLE.match(ln[3])]
            if not ys:
                raise BallotError(f"page {n}: a ballot title but no FEDERAL OFFICES heading, and no first contest, in the first column")
            top = max(ys) + 1.0 * s
        head = [ln for ln in whole if top is not None and ln[0] > top]
        body = [[ln for ln in c if top is None or ln[0] <= top] for c in cols]
        out.append({"n": n, "titled": titled, "head": head, "cols": body, "closing": [ln[0] for c in body for ln in c if FOOT.match(ln[3])]})
    return out


def read_stream(stream, edges, where, s=1.0):
    """The contests of one ballot from its lines in reading order, each (column id, y, x0, size, text); the edges of a
    line's column are edges[column id % 10]. Type sizes differ from county to county, so lines are told apart by their
    words and where they start: a contest is its headings, a VOTE FOR line, then party lines and names at the names'
    edge, then its WRITE-IN lines; a question is its headings and then YES and NO, or FOR and AGAINST, at the choices'
    edge. Returns (contests, headings left over at the end, section headings seen)."""
    out, sections, head = [], [], []
    section = cur = pending_party = last_choice = None
    state = "head"

    def split_head():
        groups = []
        for col, y, text in head:
            if groups and groups[-1][-1][0] == col and groups[-1][-1][1] - y <= 14.0 * s:
                groups[-1].append((col, y, text))
            else:
                groups.append([(col, y, text)])
        return [[t for _c, _y, t in g] for g in groups]

    def set_section(text):
        nonlocal section
        section = text
        if text not in sections:
            sections.append(text)

    def cut_section(lines):
        """A block that opens with a section heading already seen on this ballot is cut after it."""
        for k in (2, 1):
            if len(lines) > k and " ".join(lines[:k]) in sections:
                set_section(" ".join(lines[:k]))
                return lines[k:]
        return lines

    def close():
        nonlocal cur, pending_party
        if cur is not None:
            if pending_party:
                raise BallotError(f"{where}: a party line with no name under it")
            parties = [p for p, _n in cur["cands"]]
            if any(parties) and not all(parties):
                raise BallotError(f"{where}: a contest in which some names have a party line and some have none")
            out.append(cur)
        cur, pending_party = None, None

    for col, y, x0, _size, text in stream:
        N, W, C = edges[col % 10]
        at_name = abs(x0 - N) <= 0.6 or (text == "WRITE-IN" and abs(x0 - W) <= 0.6)
        is_choice = abs(x0 - C) <= 0.7 and bool(CHOICE.match(text)) and not is_term(text) and not COURT_WORD.search(text)
        if state == "cands":
            if at_name and text in BALLOT_PARTIES and col == cur["col"]:
                if cur["write_ins"] or pending_party:
                    raise BallotError(f"{where}: a party line out of place")
                pending_party = text
                continue
            if at_name and text == "WRITE-IN" and col == cur["col"]:
                cur["write_ins"] += 1
                continue
            if at_name and not cur["write_ins"] and col == cur["col"]:
                if cur["cands"] and pending_party is None and cur["last"] - y <= 11.2 * s:
                    p, nm = cur["cands"][-1]                      # a long name running on to a second line
                    cur["cands"][-1] = (p, nm + " " + text)
                    cur["wrapped"] = True
                else:
                    cur["cands"].append((pending_party, text))
                pending_party = None
                cur["last"] = y
                continue
            if not cur["cands"] and not cur["write_ins"] and pending_party is None and col == cur["col"] and cur["vf_y"] - y <= 12.5 * s \
                    and re.fullmatch(r"(\w+ )?\(\d+\)", text):
                cur["vote_for"] += " " + text                     # VOTE FOR NOT MORE THAN THREE / (3)
                continue
            if not cur["write_ins"]:
                raise BallotError(f"{where}: a line that is neither a name nor a write-in line before a contest's write-in line")
            close()
            state = "head"
        if state == "choice":
            if col == last_choice[0] and (is_choice or last_choice[1] - y <= 11.5 * s):
                last_choice = (col, y)                            # the other choice, or a choice's own words running on
                continue
            state = "head"
        if VOTE_FOR.match(text):
            groups = split_head()
            head = []
            if not groups:
                raise BallotError(f"{where}: a vote-for line with no heading over it")
            for g in groups[:-1]:
                set_section(" ".join(g))
            g = cut_section(groups[-1])
            title = [t for t in g if not is_term(t)]
            if not title:
                raise BallotError(f"{where}: a vote-for line under a heading with no title")
            cur = {"kind": "office", "section": section, "title": title, "term": [t for t in g if is_term(t)], "vote_for": text, "vf_y": y,
                   "cands": [], "write_ins": 0, "col": col, "last": None, "wrapped": False}
            state = "cands"
            continue
        if is_choice and head:
            groups = split_head()
            head = []
            m = RETAIN.search(" ".join(groups[-1]))
            if m and text == "YES":
                last = groups[-1]
                k = next(i for i, t in enumerate(last) if t.upper().startswith("SHALL "))
                office, before = (last[:k], groups[:-1]) if k else (groups[-2] if len(groups) > 1 else [], groups[:-2])
                if not office and out and out[-1]["kind"] == "retention":
                    office = out[-1]["title"]                     # a second justice under one heading
                if not office:
                    raise BallotError(f"{where}: a retention question without an office over it")
                for g in before:
                    set_section(" ".join(g))
                out.append({"kind": "retention", "section": section, "title": cut_section(office), "judge": m.group(2), "rank": (m.group(1) or "").upper()})
            else:
                for g in groups[:-1]:
                    if " ".join(g) in sections or (g is groups[0] and len(g) == 1 and re.search(r"(PROPOSITIONS?|AMENDMENTS?|QUESTIONS?|MEASURES?)$", g[0])):
                        set_section(" ".join(g))
                out.append({"kind": "measure", "section": section})   # a ballot question: counted, its wording never kept
            state, last_choice = "choice", (col, y)
            continue
        head.append((col, y, text))
    if state == "cands":
        if not cur["write_ins"]:
            raise BallotError(f"{where}: the ballot ends inside a contest, before its write-in line")
        close()
    return out, split_head(), sections


def kept_texts(doc):
    """Every piece of text of a read ballot that could ever be stored or shown (the wording of ballot questions is not
    among them: it is never kept)."""
    out = list(doc["head"]) + list(doc["foot"]) + list(doc["sections"])
    for c in doc["contests"]:
        out += c.get("title", [])
        if c["kind"] == "office":
            out += c["term"] + [c["vote_for"]] + [p for p, _n in c["cands"] if p] + [n for _p, n in c["cands"]]
        elif c["kind"] == "retention":
            out.append(c["judge"])
    return out


def read_ballots(data, label):
    """[{"style", "pages", "contests", "head", "foot", "sections", "drawn_twice"}], one for each ballot in a PDF of sample
    ballots (a precinct's ballot is two pages; some clerks post one file a precinct, some one file for the county).
    Only a general election ballot dated November 3, 2026 is read."""
    pdf = PDF(data)
    pages = list(pdf.pages())
    texts, twice = page_texts(pdf, pages)
    if not any(r[3].strip() for runs in texts for r in runs):
        raise BallotError(f"{label}: a scanned picture with no text to read", why="scan")
    if not any(BALLOT_TITLE.search(ln[3]) for runs in texts for ln in baseline_lines(runs)):
        raise BallotError(f"{label}: no general election ballot in the file", why="not-general")
    edges = column_edges(texts)
    if not edges:
        raise BallotError(f"{label}: no WRITE-IN line found; not a ballot this reader knows")
    if len(edges) > 9:
        raise BallotError(f"{label}: {len(edges)} columns")
    s = (edges[1][1] - edges[0][1]) / PITCH if len(edges) > 1 else 1.0
    ballots = []
    for p in ballot_pages(texts, edges, s):
        if p["titled"]:
            instr = [ln[0] for ln in p["head"] if INSTRUCTIONS.match(ln[3])]
            lines = sorted([ln for ln in p["head"] if not instr or ln[0] > instr[0]], key=lambda l: (-l[0], l[1]))
            style = " ".join(ln[3] for ln in lines if not HEAD_SKIP.match(ln[3]))
            if not any(re.fullmatch(r"NOVEMBER 3,? 2026", ln[3], re.I) for ln in lines):
                raise BallotError(f"{label}: page {p['n']}: a general election ballot that is not dated November 3, 2026", why="other-election")
            ballots.append({"style": style, "pages": [], "stream": [], "head": [ln[3] for ln in lines], "foot": []})
        if not ballots:
            raise BallotError(f"{label}: page {p['n']} comes before any ballot's title")
        b = ballots[-1]
        b["pages"].append(p["n"])
        hi = max(p["closing"]) + 18.0 * s if p["closing"] else None      # the foot of the page: the closing words and the ballot's own label
        lo = min(p["closing"]) - 5.0 * s if p["closing"] else None
        for k, lines in enumerate(p["cols"]):
            for (y, x0, size, text, _x1) in lines:
                if hi is not None and y < hi:
                    if FOOT.match(text) or MACHINE.match(text) or (b["style"] and (text in b["style"] or b["style"] in text)):
                        b["foot"].append(text)
                        continue
                    if y < lo:
                        raise BallotError(f"{label}: page {p['n']}: a line under the page's closing words that is not the ballot's own label")
                b["stream"].append((len(b["pages"]) * 10 + k, y, x0, size, text))
    out = []
    for k, b in enumerate(ballots, 1):
        contests, leftover, sections = read_stream(b["stream"], edges, f"{label}, ballot {k}", s)
        if leftover and not (contests and contests[-1]["kind"] == "measure"):
            raise BallotError(f"{label}: headings at the end of a ballot with no contest under them")
        if twice and any(" " not in n.strip() or "WRITE" in n for c in contests if c["kind"] == "office" for _p, n in c["cands"]):
            raise BallotError(f"{label}: a name that is one word, or a write-in line read as a name, in a file that draws its text twice")
        out.append({"style": b["style"], "pages": b["pages"], "contests": contests, "head": b["head"], "foot": b["foot"], "sections": sections,
                    "drawn_twice": twice})
    return out


class PageLinks(HTMLParser):
    """The links of a page: their own words and their addresses. Nothing else of the page is kept."""

    def __init__(self):
        super().__init__()
        self.links, self.cur, self.buf, self.base = [], None, [], None

    def handle_starttag(self, tag, attrs):
        d = dict(attrs)
        if tag == "base" and d.get("href"):
            self.base = d["href"]
        if tag == "a" and d.get("href"):
            self.cur, self.buf = d["href"], []

    def handle_data(self, data):
        if self.cur is not None:
            self.buf.append(data)

    def handle_endtag(self, tag):
        if tag == "a" and self.cur is not None:
            self.links.append((re.sub(r"\s+", " ", "".join(self.buf)).strip(), self.cur))
            self.cur = None


def picked_links(page_url, pick, href=None, other_host=None):
    """([(link words, address)], fingerprint of the page, its size) for the links of a page whose own words fit `pick`
    (and whose address holds `href`), on the page's own site only, or on the one other site named. The page is held in
    memory and not kept."""
    raw = polite_get(page_url)
    p = PageLinks()
    p.feed(raw.decode("utf-8", "replace"))
    base = urljoin(page_url, p.base) if p.base else page_url
    hosts = {urlparse(page_url).netloc.lower()} | ({other_host} if other_host else set())
    out, seen = [], set()
    for text, h in p.links:
        text = html.unescape(text)
        if not re.search(pick, text) or h.startswith(("mailto:", "tel:", "javascript:", "#")):
            continue
        full = urljoin(base, h)
        if urlparse(full).netloc.lower() not in hosts or (href and not re.search(href, full, re.I)):
            continue
        if full not in seen:
            seen.add(full)
            out.append((text, full))
    return out, hashlib.sha256(raw).hexdigest(), len(raw)


DRIVE = "drive.google.com"


def drive_folder_files(folder_url, pick):
    """[(file name, address to fetch it from)] for the files of a public shared folder whose names fit `pick`. The
    clerk's page links the folder; its plain listing page names each file, and nothing else of it is kept."""
    m = re.search(r"/folders/([A-Za-z0-9_-]{10,})", folder_url)
    if not m:
        raise ValueError("the shared folder's address is not one this loader knows")
    page = polite_get(f"https://{DRIVE}/embeddedfolderview?id={m.group(1)}").decode("utf-8", "replace")
    out = []
    for fid, title in re.findall(r'<div class="flip-entry" id="entry-([A-Za-z0-9_-]+)".*?<div class="flip-entry-title">(.*?)</div>', page, re.S):
        title = html.unescape(re.sub(r"\s+", " ", title)).strip()
        if re.search(pick, title):
            out.append((title, f"https://{DRIVE}/uc?export=download&id={fid}"))
    return out


def fetch_ballot(url, label):
    """A sample ballot's bytes, checked in memory before it is kept. BallotError when it is not a text PDF, not a general
    election ballot this reader can read, or when a line among those that could ever be kept reads like contact details.
    (The wording of ballot questions is not looked at: it is never kept, and its figures can look like telephone
    numbers.)"""
    b = polite_get(url)
    if b[:4] != b"%PDF":
        raise BallotError(f"{label}: the address answered without a PDF", why="not-pdf")
    docs = read_ballots(b, label)
    bad = sum(1 for d in docs for t in kept_texts(d) if contact_like(t, True))
    if bad:
        raise BallotError(f"{label}: {bad} lines read like contact details", why="contact")
    return b


def county_ballots(name, fips3, cfg, folder, say=print, refresh=False):
    """The manifest of one county's sample ballots in ballot_cache/wy/local/: the page they were linked from (its
    address, fingerprint and the day it was read), each ballot file kept (its link's words, address, file, size,
    fingerprint and day), and each file refused, with the reason. Nothing is asked for when the manifest is there,
    unless refresh; a file already kept under its address is never fetched twice. None when the page cannot be reached
    and nothing is on disk."""
    stem = f"{fips3}_{re.sub(r'[^a-z]+', '_', name.lower())}"
    mpath = os.path.join(folder, f"{stem}_sample_ballots.json")
    old = json.load(open(mpath, encoding="utf-8")) if os.path.exists(mpath) else None
    have = {d["url"]: d for d in (old or {}).get("docs", []) if os.path.exists(os.path.join(folder, d["file"]))}
    if old and not refresh and old.get("page_sha256") and len(have) == len(old.get("docs", [])):
        return old
    try:
        links, page_sha, page_bytes = picked_links(cfg["page"], cfg["pick"], cfg.get("href"), DRIVE if cfg.get("folder_files") else None)
        if cfg.get("folder_files"):                                   # the page links a shared folder; the ballots are the folder's files
            if len(links) != 1:
                raise ValueError("the page does not link exactly one folder of sample ballots")
            links = drive_folder_files(links[0][1], cfg["folder_files"])
    except Exception as e:  # noqa: BLE001
        if old:
            say(f"    Wyoming (local): {name} County's page did not answer ({type(e).__name__}); using what is on disk")
            return old
        return None
    refused_before = {d["url"]: d for d in (old or {}).get("refused", [])}
    taken = {d["file"] for d in have.values()}
    kept, refused = [], []
    for i, (text, url) in enumerate(links, 1):
        if url in have:
            kept.append(have[url])
            continue
        if url in refused_before and not refresh:
            refused.append(refused_before[url])
            continue
        fname = next(f for f in (f"{stem}_sample_ballot_{k:02d}.pdf" for k in range(i, i + 500)) if f not in taken)
        try:
            b = fetch_ballot(url, f"{name} County, file {i}")
        except BallotError as e:
            refused.append({"words": text, "url": url, "why": e.why, "detail": str(e)[:200]})
            continue
        except Exception as e:  # noqa: BLE001  the file did not come: said so, and asked for again on a refresh
            refused.append({"words": text, "url": url, "why": "unreached", "detail": type(e).__name__})
            continue
        with open(os.path.join(folder, fname), "wb") as fh:
            fh.write(b)
        taken.add(fname)
        kept.append({"words": text, "url": url, "file": fname, "bytes": len(b), "sha256": hashlib.sha256(b).hexdigest(),
                     "fetched": dt.date.today().isoformat()})
    m = {"county": name, "page": cfg["page"], "page_sha256": page_sha, "page_bytes": page_bytes, "page_fetched": dt.date.today().isoformat(),
         "found": len(links), "docs": kept, "refused": refused}
    with open(mpath + ".part", "w", encoding="utf-8") as fh:
        json.dump(m, fh, ensure_ascii=False, indent=1)
    os.replace(mpath + ".part", mpath)
    return m


# ---------- Laramie County's rosters: three workbooks, five columns of each ----------

def laramie_rosters(folder, say=print, refresh=False):
    """The manifest of Laramie County's three general election rosters, each cut down in memory to the columns named in
    LARAMIE_KEEP (chosen by their headings) and kept that way as JSON; the mailing address, city, phone, e-mail and
    website columns are never read. None when the page cannot be reached and nothing is on disk."""
    mpath = os.path.join(folder, "021_laramie_rosters.json")
    old = json.load(open(mpath, encoding="utf-8")) if os.path.exists(mpath) else None
    if old and not refresh and all(os.path.exists(os.path.join(folder, d["file"])) for d in old["docs"]) and len(old["docs"]) == len(LARAMIE_BOOKS):
        return old
    try:
        # the 2026 general election workbooks say so in their addresses (the page also links the primary's)
        links, page_sha, page_bytes = picked_links(LARAMIE_PAGE, r"\(XLSX", r"/2026/[^/]*general")
        books = {}
        for key, pick in LARAMIE_BOOKS.items():
            hit = [(t, u) for t, u in links if re.search(pick, t)]
            if len(hit) != 1:
                raise ValueError(f"the page does not link exactly one general election workbook for {key}")
            books[key] = hit[0]
        docs = []
        for key in LARAMIE_BOOKS:
            text, url = books[key]
            raw = polite_get(url)
            if raw[:2] != b"PK":
                raise ValueError("a roster's address answered without a workbook")
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            if len(wb.worksheets) != 1:
                raise ValueError("a roster workbook does not hold exactly one sheet")
            grid = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
            cell = lambda c: c.date().isoformat() if isinstance(c, dt.datetime) else re.sub(r"\s+", " ", str(c)).strip() if c is not None else ""  # noqa: E731
            hi = next((i for i, r in enumerate(grid) if {"Office", "Ballot Name"} <= {cell(c) for c in r}), None)
            if hi is None:
                raise ValueError("a roster workbook has no row of headings with Office and Ballot Name")
            head = [cell(c) for c in grid[hi]]
            cols = [k for k in LARAMIE_KEEP if k in head]
            if any(head.count(k) != 1 for k in cols):
                raise ValueError("a roster workbook repeats a heading")
            idx = [head.index(k) for k in cols]
            rows, blanked = [], 0
            for r in grid[hi + 1:]:
                cut = [cell(r[i]) if i < len(r) else "" for i in idx]
                if not any(cut):
                    continue
                for j, v in enumerate(cut):
                    if v and contact_like(v, True):
                        cut[j], blanked = "", blanked + 1
                rows.append(cut)
            fname = f"021_laramie_roster_{key}.json"
            with open(os.path.join(folder, fname), "w", encoding="utf-8") as fh:
                json.dump({"columns": cols, "rows": rows}, fh, ensure_ascii=False, indent=0)
            docs.append({"key": key, "words": text, "url": url, "file": fname, "bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest(),
                         "fetched": dt.date.today().isoformat(), "columns_in_file": len([h for h in head if h]), "columns": cols,
                         "rows": len(rows), "blanked": blanked})
            del raw, grid, wb
    except Exception as e:  # noqa: BLE001
        if old:
            say(f"    Wyoming (local): Laramie County's rosters could not be read again ({type(e).__name__}); using what is on disk")
            return old
        say(f"    Wyoming (local): Laramie County's rosters could not be read ({type(e).__name__}: {str(e)[:120]})")
        return None
    m = {"county": "Laramie", "page": LARAMIE_PAGE, "page_sha256": page_sha, "page_bytes": page_bytes, "page_fetched": dt.date.today().isoformat(),
         "docs": docs}
    with open(mpath + ".part", "w", encoding="utf-8") as fh:
        json.dump(m, fh, ensure_ascii=False, indent=1)
    os.replace(mpath + ".part", mpath)
    return m


# ---------- two county clerks' lists that are PDF tables with contact columns: a few columns of each ----------

NATRONA_PAGE = "https://www.natronacounty-wy.gov/131/Candidate-Information"
NATRONA_PICK = r"^2026 General Election Candidates \(PDF\)$"
NATRONA_HEADS = ("Office", "Party", "Ballot Name", "Mailing Address")      # the headings the column edges are taken from
FREMONT_PAGE = COUNTY_PAGES["Fremont"]
FREMONT_PICK = r"^Candidate Filings - Schools and Special Districts$"
FREMONT_HEADS = ("District Name", "Office", "Ballot Name", "Residential Address")
ROSTER_PARTY = {"REP": "Republican", "DEM": "Democratic", "LIB": "Libertarian", "LBR": "Libertarian", "CON": "Constitution", "CT": "Constitution",
                "IND": "Independent"}


def tight_join(runs):
    """Runs on one printed row as text. These tables set their words closer than the kit's own join allows for, so a
    gap wider than a twentieth of the type is a space."""
    text, end = "", None
    for x0, _y, size, t, x1 in sorted(runs, key=lambda r: r[0]):
        if end is not None and x0 - end > 0.05 * size and not text.endswith(" ") and not t.startswith(" "):
            text += " "
        text += t
        end = max(end or x1, x1)
    return re.sub(r"\s+", " ", text).strip()


def table_rows(pdf, page, res):
    """A page's runs grouped into printed rows, top to bottom: [(y, [runs])]."""
    rows = []
    for r in sorted([r for r in page_runs(pdf, page, res) if r[3].strip()], key=lambda r: (-r[1], r[0])):
        if rows and abs(rows[-1][0] - r[1]) <= 2.0:
            rows[-1][1].append(r)
        else:
            rows.append([r[1], [r]])
    return rows


def heading_starts(rs, wanted):
    """{heading: the x its column's cells start at} when the row is the table's row of headings (it holds every heading
    in `wanted`), else None. Only a row of headings is ever read this way."""
    cells = []
    for r in sorted(rs, key=lambda r: r[0]):
        if cells and r[0] - cells[-1][2] <= 6:
            cells[-1][1].append(r)
            cells[-1][2] = max(cells[-1][2], r[4])
        else:
            cells.append([r[0], [r], r[4]])
    heads = {tight_join(c[1]): c[0] for c in cells}
    return heads if all(h in heads for h in wanted) else None


def cell_chains(rs, starts):
    """The cells of one row as [start x, end x, pieces]. A cell's text can run on under the next column, so text is put
    together by where it starts, not by where it lies: the pieces of one cell follow one another with no more than a
    word's space between them, and a piece that starts where a column starts begins that column's cell unless it
    follows the piece before it with no gap at all (the file cuts a long label in two at the column's edge)."""
    chains = []
    for r in sorted(rs, key=lambda r: r[0]):
        at_start = any(abs(r[0] - s) <= 0.6 for s in starts)
        hit = next((c for c in chains if -0.5 <= r[0] - c[1] <= (0.15 if at_start else 2.0)), None)
        if hit:
            hit[1] = max(hit[1], r[4])
            hit[2].append(r)
        else:
            chains.append([r[0], r[4], [r]])
    return chains


def natrona_rows(data):
    """[[whose office, office, party, name]] for every candidate row of Natrona County's roster, and the number of
    columns its heading row names. The table's first column (it has no heading) names whose office it is and the
    second the office ("Bar Nunn" and "Bar Nunn Town Council"); for a district the first column holds the whole label
    ("Casper Mountain Fire District Director") and the second is empty: such a row comes back with the label as its
    office and nothing in the first cell. A label is printed once, on the first row of its group, and holds for the
    rows under it. Only cells that start in the first four columns are ever put together (see cell_chains); a piece
    that starts at or beyond the Mailing Address column is never joined into text. A row that starts a group (it stands
    apart from the row above) must carry a label, and a row inside a group must not: anything else stops the loader."""
    pdf = PDF(data)
    out, jur, office, ncols = [], None, None, 0
    for pn, (page, res) in enumerate(pdf.pages(), 1):
        starts, lines = None, []
        for y, rs in table_rows(pdf, page, res):
            if starts is None:
                heads = heading_starts(rs, NATRONA_HEADS)
                if heads:
                    starts = sorted(heads.values())                   # where each headed column's cells begin
                    office_x, party_x, name_x, addr_x = (heads[h] for h in NATRONA_HEADS)
                    ncols = max(ncols, len(heads) + 1)
                continue
            cell = {"j": "", "o": "", "p": "", "n": ""}
            for x0, x1, pieces in cell_chains(rs, starts):
                key = "j" if x0 < office_x - 4 else "o" if x0 < party_x - 4 else "p" if x0 < name_x - 4 else "n" if x0 < addr_x - 4 else None
                if key is None:
                    continue                                          # the contact columns: never put together
                if cell[key]:
                    lfail(f"Natrona County's roster, page {pn}: two pieces of text start in one cell; stopped (the row is not printed)")
                if key == "n" and x1 > addr_x - 1:
                    lfail(f"Natrona County's roster, page {pn}: a name runs into the address column; stopped (the row is not printed)")
                cell[key] = tight_join(pieces)
            if cell["n"]:                                             # a row with no name is the date at the foot of the page
                lines.append((y, cell["j"], cell["o"], cell["p"], cell["n"]))
        if starts is None:
            lfail(f"Natrona County's roster, page {pn}: no row of headings with Office, Party, Ballot Name and Mailing Address; stopped")
        gaps_ = [a[0] - b[0] for a, b in zip(lines, lines[1:])]
        pitch = min(gaps_) if gaps_ else None
        for i, (y, j, o, p, name) in enumerate(lines):
            begins = i == 0 or gaps_[i - 1] > 1.5 * pitch
            if i and begins != bool(j or o):
                lfail(f"Natrona County's roster, page {pn}: a row's label does not fit its place in the table; stopped (the row is not printed)")
            if j and o:
                jur, office = j, o
            elif j:
                jur, office = "", j                                   # a district: the whole label sits in the first column
            elif o:
                if not jur:
                    lfail(f"Natrona County's roster, page {pn}: an office with no place named over it; stopped")
                office = o
            if office is None:
                lfail(f"Natrona County's roster, page {pn}: a candidate row before any office is named; stopped")
            out.append([jur, office, p, name])
    return out, ncols


def fremont_rows(data):
    """[[district, office, name]] for every row of Fremont County's list of general election filings, the number of
    columns its heading row names, and whether the lines over the table say 2026. Each row carries its own district and
    office. Three columns are put together, by where their cells start (see cell_chains): District Name, Office and
    Ballot Name. The party and filing date between them are not read, and a piece that starts at or beyond the
    Residential Address column is never joined into text."""
    pdf = PDF(data)
    out, ncols, dated, starts = [], 0, False, None
    for pn, (page, res) in enumerate(pdf.pages(), 1):
        rows = table_rows(pdf, page, res)
        at = next((i for i, (_y, rs) in enumerate(rows) if heading_starts(rs, FREMONT_HEADS)), None)
        if at is None and starts is None:
            lfail(f"Fremont County's filing list, page {pn}: no row of headings with District Name, Office and Ballot Name; stopped")
        if at is not None:
            heads = heading_starts(rows[at][1], FREMONT_HEADS)
            starts = sorted(heads.values())
            _dist_x, office_x, name_x, addr_x = (heads[h] for h in FREMONT_HEADS)
            after_office = min(v for v in heads.values() if v > office_x)
            ncols = max(ncols, len(heads) + 1)                     # the party and the filing date share one heading cell
            # the lines over the table: only looked at for the year the filing closed (nothing of them is kept)
            dated = dated or any(re.search(r"\b(2026|\d\d?/\d\d?/26)\b", tight_join(rs)) for _y, rs in rows[:at])
        for y, rs in rows[(at + 1) if at is not None else 0:]:     # a page without headings of its own goes on under the page before
            cell = {"d": "", "o": "", "n": ""}
            for x0, x1, pieces in cell_chains(rs, starts):
                key = "d" if x0 < office_x - 4 else "o" if x0 < after_office - 4 else None if x0 < name_x - 4 else "n" if x0 < addr_x - 4 else None
                if key is None:
                    continue                                          # the party, the filing date and the contact columns
                if cell[key]:
                    lfail(f"Fremont County's filing list, page {pn}: two pieces of text start in one cell; stopped (the row is not printed)")
                if key == "n" and x1 > addr_x - 1:
                    lfail(f"Fremont County's filing list, page {pn}: a name runs into the address column; stopped (the row is not printed)")
                cell[key] = tight_join(pieces)
            if not (cell["d"] or cell["o"] or cell["n"]):
                continue
            if not (cell["d"] and cell["o"] and cell["n"]):
                lfail(f"Fremont County's filing list, page {pn}: a row without a district, an office or a name; stopped (the row is not printed)")
            out.append([cell["d"], cell["o"], cell["n"]])
    return out, ncols, dated


def clerk_table(folder, key, county, page, pick, reader, columns, say=print, refresh=False):
    """The manifest of one county clerk's PDF table (Natrona's roster, Fremont's filing list), cut down in memory by
    `reader` to the columns named and kept that way as JSON; the contact columns are never read. None when the page or
    the file cannot be reached and nothing is on disk."""
    mpath = os.path.join(folder, key)
    old = json.load(open(mpath, encoding="utf-8")) if os.path.exists(mpath) else None
    if old and not refresh:
        return old
    try:
        links, page_sha, page_bytes = picked_links(page, pick)
        if len(links) != 1:
            raise ValueError("the page does not link exactly one such list")
        text, url = links[0]
        raw = polite_get(url)
        if raw[:4] != b"%PDF":
            raise ValueError("the list's address answered without a PDF")
        rows, ncols, *more = reader(raw)
        if more and not more[0]:
            raise ValueError("the list does not say 2026 over its table")
        blanked = 0
        for row in rows:
            for k, v in enumerate(row):
                if v and contact_like(v, True):
                    row[k], blanked = "", blanked + 1
        m = {"county": county, "page": page, "page_sha256": page_sha, "page_bytes": page_bytes, "page_fetched": dt.date.today().isoformat(),
             "words": text, "url": tidy_url(url), "bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest(), "fetched": dt.date.today().isoformat(),
             "columns_in_file": ncols, "columns": columns, "blanked": blanked, "rows": rows}
        del raw
    except SystemExit:
        raise
    except Exception as e:  # noqa: BLE001
        if old:
            say(f"    Wyoming (local): {county} County's list could not be read again ({type(e).__name__}); using what is on disk")
            return old
        say(f"    Wyoming (local): {county} County's list could not be read ({type(e).__name__}: {str(e)[:120]})")
        return None
    with open(mpath + ".part", "w", encoding="utf-8") as fh:
        json.dump(m, fh, ensure_ascii=False, indent=0)
    os.replace(mpath + ".part", mpath)
    return m


def natrona_roster(folder, say=print, refresh=False):
    return clerk_table(folder, "025_natrona_roster.json", "Natrona", NATRONA_PAGE, NATRONA_PICK, natrona_rows,
                       ["whose office (the unheaded first column)", "Office", "Party", "Ballot Name"], say, refresh)


def fremont_filings(folder, say=print, refresh=False):
    return clerk_table(folder, "013_fremont_filings.json", "Fremont", FREMONT_PAGE, FREMONT_PICK, fremont_rows,
                       ["District Name", "Office", "Ballot Name"], say, refresh)


def natrona_contests(manifest):
    """Natrona County's roster as contests like a ballot's: ({(title, term): {...}}, rows read, the judges' names it
    lists under Judicial Retention). The first column and the office are put together into a title the way a ballot
    words one: "Casper Ward 1" and "Casper City Council" make CASPER CITY COUNCIL WARD 1; "Evansville" and "Mayor" make
    EVANSVILLE MAYOR. An office that says "(2 year unexpired)" is for the rest of a term. The roster prints a party
    beside every name; it is kept only for the county's own offices, which are the ones on the party ballot."""
    out, judges_listed = OrderedDict(), []
    for n, (jur, office, party, name) in enumerate(manifest["rows"], start=1):
        if not (office and name):
            lfail(f"row {n} of Natrona County's roster has an empty cell where a name or an office is expected; stopped")
        if jur.upper() == "JUDICIAL RETENTION":
            judges_listed.append(name)
            continue
        office = re.sub(r"\bCoundil\b", "Council", office)             # one misprint in the roster's office column
        term = []
        m = re.search(r"\s*\((\d) year unexpired\)$", office, re.I)
        if m:
            office, term = office[:m.start()], [f"{NUMBER_NAMES[int(m.group(1))].upper()} ({m.group(1)}) YEAR UNEXPIRED TERM"]
        j = re.sub(r"\bWard ?(\d+)$", r"Ward \1", jur)
        ward = re.search(r"\s*\bWard \d+$", j)
        place = j[:ward.start()] if ward else j
        if jur.upper() == "NATRONA COUNTY":
            section, title = "COUNTY OFFICES", office.upper()
        elif not jur:
            section, title = "NONPARTISAN OFFICES", office.upper()   # a district: the label is whole
        else:
            section = "NONPARTISAN OFFICES"
            rest = office[len(place):].strip() if office.upper().startswith(place.upper()) else office
            title = f"{place} {rest}{ward.group(0) if ward else ''}".upper()
        e = out.setdefault((title, " ".join(term)), {"title": [title], "term": term, "section": section, "cands": [], "county_office": section == "COUNTY OFFICES"})
        if e["county_office"]:
            if party not in ROSTER_PARTY:
                lfail(f"row {n} of Natrona County's roster: a county office with a party this loader does not know; stopped")
            e["cands"].append((ROSTER_PARTY[party], name))
        else:
            e["cands"].append((None, name))
    return out, len(manifest["rows"]), judges_listed


FREMONT_SHORT = (("DIST", "DISTRICT"), ("SR", "SENIOR"), ("SVC", "SERVICE"), ("SUBDIST", "SUBDISTRICT"))      # the list's own abbreviations


def fremont_contests(manifest):
    """Fremont County's filing list as contests like a ballot's: ({(title, term): {...}}, rows read, withdrawn left off,
    rows for the county's own offices left to the county's gap, {name as shown: name as the list prints it} for the
    names whose quotation marks the list misplaces, names with a quotation mark kept as printed). The clerk writes the
    districts short: "FC SCHOOL#1 (LANDER)" is Fremont County School District #1, "DIST" is district, "SR CITIZEN SVC
    DIST" a senior citizen service district; an office that says "2 YR UNEXPIRED" (written five ways) is for the rest
    of a term, and one that says WITHDRAWN is a candidate who withdrew. The list has no column for how many are
    elected. A nickname's quotation marks come out of the clerk's system one word late (GIVEN NICK" FAMILY", the mark
    of a quoted field read back without its doubled quotes): with exactly one word before the nickname they are put
    back (GIVEN "NICK" FAMILY) and the race's candidate says so; any other name with a quotation mark is kept as
    printed and reported."""
    out, gone, county_rows, requoted, odd_quotes = OrderedDict(), 0, 0, {}, 0
    for n, (district, office, name) in enumerate(manifest["rows"], start=1):
        d, o = district.upper(), office.upper()
        if re.match(r"COMMISSIONER\b", d):
            county_rows += 1                                          # the county's own offices are not on this list but for a late filing
            continue
        if "WITHDRAWN" in o:
            gone += 1
            continue
        if '"' in name:
            q = re.fullmatch(r'([^\s"]+) ([^\s"]+)" ([^"]+)"', name)
            if q:
                printed, name = name, f'{q.group(1)} "{q.group(2)}" {q.group(3)}'
                requoted[name] = printed
            elif not re.fullmatch(r'[^"]+ "[^"]+" [^"]+', name):
                odd_quotes += 1
        cut = re.search(r"\s*[-(]?\s*(\d ?YR\b|UNEXP)", o)
        base, tail = (o[:cut.start()], o[cut.start():]) if cut else (o, "")
        years = re.search(r"(\d) ?YR\b", tail)
        term = [] if not tail else [f"{NUMBER_NAMES[int(years.group(1))].upper()} ({years.group(1)}) YEAR UNEXPIRED TERM"] if years else ["UNEXPIRED"]
        if tail and "UNEXP" not in tail:
            lfail(f"row {n} of Fremont County's filing list: an office with a term this loader does not know; stopped")
        said = None
        m = re.fullmatch(r"FC SCHOOL ?#(\d+) \((.+)\)", d)
        if m:
            d, said = f"FREMONT COUNTY SCHOOL DISTRICT #{m.group(1)}", district
        for short, long_ in FREMONT_SHORT:
            d = re.sub(rf"\b{short}\b", long_, d)
            base = re.sub(rf"\b{short}\b", long_, base)
        base = re.sub(r"^CWC\s+", "", base)                           # "CWC SUBDISTRICT 2" under CENTRAL WYOMING COLLEGE
        base = re.sub(r"^.*\bSENIOR CIT\s+(?=TRUSTEE)", "", base)     # "LANDER SR CIT TRUSTEE" under LANDER SR CITIZEN SVC DIST
        title = f"{d} {base}".strip()
        e = out.setdefault((title, " ".join(term)), {"title": [title], "term": term, "section": "NONPARTISAN OFFICES", "cands": [], "said": said})
        e["cands"].append((None, name))
    return out, len(manifest["rows"]), gone, county_rows, requoted, odd_quotes


# ---------- Big Horn County's roster: nine tables whose cells are centred under their headings ----------

BIGHORN_PAGE = COUNTY_PAGES["Big Horn"]
BIGHORN_PICK = r"^General Election Candidate Filings$"
BIGHORN_TABLES = ("OFFICE", "TOWNS", "SCHOOL DISTRICTS", "CEMETERY DISTRICTS", "CONSERVATION DISTRICTS", "FIRE DISTRICTS", "HOSPITAL DISTRICTS",
                  "SENIOR CITIZEN DISTRICTS", "RURAL HEALTH DISTRICT")       # each table's first heading
CENTRED = 2.5                                                    # points: how far a cell's middle may lie from its heading's middle


def gap_cells(rs, gap=6.0):
    """The cells of one printed row of a table whose cells stand apart: [start x, end x, pieces], a new cell wherever
    two pieces stand more than gap apart. Nothing is put together into text here."""
    cells = []
    for r in sorted(rs, key=lambda r: r[0]):
        if cells and r[0] - cells[-1][1] <= gap:
            cells[-1][1] = max(cells[-1][1], r[4])
            cells[-1][2].append(r)
        else:
            cells.append([r[0], r[4], [r]])
    return cells


def bighorn_rows(data):
    """([[kind, table, first column, middle column, name column]], the most headings a table has, whether the lines over
    the tables say 2026 and general) for Big Horn County's roster. kind is "L" for a row of labels (an office, a town,
    a district, a seat: with how many seats and the term) and "C" for a candidate. The roster is nine tables, each
    under its own row of headings; cells are centred, so a column is known by its heading: a cell is on the first
    column's side when it ends before the second heading starts, in the middle column (Party, TERM) when it lies between
    the second heading's start and the Candidate Name heading's start, and a name only when its middle sits under the
    middle of the Candidate Name heading. Everything that starts to the right of that middle and is not a name belongs
    to the mailing address, phone and e-mail columns and is never put together into text. A cell that fits nowhere
    stops the loader (it is not printed)."""
    pdf = PDF(data)
    out, dated, ncols = [], False, 0
    table = second_x = name_x = name_c = None
    has_mid = False
    for pn, (page, res) in enumerate(pdf.pages(), 1):
        for _y, rs in table_rows(pdf, page, res):
            cells = gap_cells(rs)
            if table is None:                                         # the lines over the first table: the roster's own title
                firsts = [cells[0]]
                if tight_join(cells[0][2]) not in BIGHORN_TABLES:
                    if len(cells) != 1:
                        lfail(f"Big Horn County's roster, page {pn}: a row of several cells before any table; stopped (the row is not printed)")
                    text = tight_join(cells[0][2])
                    dated = dated or bool(re.search(r"\b2026\b", text) and re.search(r"(?i)\bgeneral\b", text))
                    continue
            else:
                firsts = [c for c in cells if c[1] < second_x - 2.0]
            if firsts and tight_join(firsts[0][2]) in BIGHORN_TABLES:    # a row of headings: every cell of it is a heading
                heads = [(tight_join(c[2]), c[0], c[1]) for c in cells]
                names = [h[0] for h in heads]
                k = names.index("Candidate Name") if "Candidate Name" in names else None
                if k not in (1, 2) or "Mailing Address" not in names[k + 1:]:
                    lfail(f"Big Horn County's roster, page {pn}: a row of headings without Candidate Name and Mailing Address where they are expected; "
                          "stopped")
                table, second_x, has_mid = heads[0][0], heads[1][1], k == 2
                name_x, name_c = heads[k][1], (heads[k][1] + heads[k][2]) / 2
                ncols = max(ncols, len(heads))
                continue
            mid, name, contact, wide = [], [], 0, []
            for c in cells:
                if any(c is f for f in firsts):
                    continue
                if has_mid and c[0] >= second_x - 2.0 and c[1] < name_x - 2.0:
                    mid.append(c)
                elif abs((c[0] + c[1]) / 2 - name_c) <= CENTRED:
                    name.append(c)
                elif c[0] > name_c:
                    contact += 1                                      # the address, phone and e-mail columns: never put together
                else:
                    wide.append(c)
            if wide:
                if len(cells) != 1:
                    lfail(f"Big Horn County's roster, page {pn}: a cell that fits no column; stopped (the row is not printed)")
                continue                                              # a line across the page: the roster's title, printed again over a later page
            joined = [" ".join(tight_join(c[2]) for c in part) for part in (firsts, mid, name)]
            if firsts:
                if contact:
                    lfail(f"Big Horn County's roster, page {pn}: a row of labels with a cell in the contact columns; stopped (the row is not printed)")
                out.append(["L", table] + joined)
            else:
                if len(name) != 1 or len(mid) > 1:
                    lfail(f"Big Horn County's roster, page {pn}: a candidate's row without exactly one name; stopped (the row is not printed)")
                out.append(["C", table] + joined)
    return out, ncols, dated


def bighorn_roster(folder, say=print, refresh=False):
    return clerk_table(folder, "003_big_horn_roster.json", "Big Horn", BIGHORN_PAGE, BIGHORN_PICK, bighorn_rows,
                       ["L for a row of labels, C for a candidate", "the table (its first heading)", "the table's first column", "Party or TERM",
                        "Candidate Name"], say, refresh)


def bighorn_contests(manifest):
    """Big Horn County's roster as contests like a ballot's: ({key: {"title", "term", "section", "cands", "seats",
    "county_office", "said"}}, candidate rows read). A county office's label can carry how many are elected ("County
    Commissioner (2)"). A town's name stands alone on its row; under it "Mayor (4 Year Term)" or "Town Council (2 Year
    Term)" with "(1 SEAT)" beside it. In the district tables a row names the district ("BHC SCHOOL DIST # 1
    (Cowley-Burlington)", Big Horn County School District #1) and the rows under it its seats ("AT LARGE (2 SEATS)",
    "4 YR"); the cemetery and fire tables have one row for both ("BURLINGTON CEMETERY (3 SEATS)", "FIRE # 1 LOVELL
    (1 SEAT)"), and "Park County Fire Dist # 1" names its seat in the name column ("Director 1"). A contest with no
    candidate under it is kept, empty. A label this reader does not know stops the loader."""
    out, town, district, said, cur, n_cands = OrderedDict(), None, None, None, None, 0
    last_table = None

    def contest(title, years, seats, section, county_office, said_):
        term = [f"{NUMBER_NAMES[years].upper()} ({years}) YEAR TERM"] if years else []
        key = (title, " ".join(term))
        if key in out:
            lfail("Big Horn County's roster names one contest twice; stopped")
        out[key] = {"title": [title], "term": term, "section": section, "cands": [], "seats": seats, "county_office": county_office, "said": said_}
        return out[key]

    for n, (kind, table, first, mid, name) in enumerate(manifest["rows"], start=1):
        where = f"row {n} of Big Horn County's roster"
        if table != last_table:
            last_table, town, district, said, cur = table, None, None, None, None
        if kind == "C":
            if cur is None or not name:
                lfail(f"{where}: a candidate with no contest over it, or without a name; stopped")
            if (table == "OFFICE") != bool(mid) or (mid and mid not in ROSTER_PARTY):
                lfail(f"{where}: a party where none is expected, none where one is, or one this loader does not know; stopped")
            cur["cands"].append((ROSTER_PARTY[mid] if mid else None, name))
            n_cands += 1
            continue
        if table == "OFFICE":
            m = re.fullmatch(r"(County [A-Za-z ]+?)(?: \((\d)\))?", first)
            if not m or mid or name:
                lfail(f"{where}: a county office label this loader does not know; stopped")
            cur = contest(m.group(1).upper(), None, int(m.group(2) or 1), "COUNTY OFFICES", True, None)
        elif table == "TOWNS":
            m = re.fullmatch(r"(Mayor|Town Council) \((\d) Year Term\)", first)
            seats = re.fullmatch(r"\((\d) SEATS?\)", name)
            if not m and not name and re.fullmatch(r"[A-Z][A-Z .'-]+", first):
                town, cur = first, None                               # a town's own row
            elif m and seats and town:
                cur = contest(f"{town} {m.group(1).upper()}", int(m.group(2)), int(seats.group(1)), "MUNICIPAL OFFICES", False, None)
            else:
                lfail(f"{where}: a town label this loader does not know; stopped")
        else:
            seats = re.search(r"\s*\((\d) SEATS?\)", first)
            label = (first[:seats.start()] + first[seats.end():]).strip() if seats else first
            years = re.fullmatch(r"(\d) ?YR", mid)
            if bool(seats) != bool(years) or (mid and not years):
                lfail(f"{where}: a district label with its seats but no term, or its term but no seats; stopped")
            if not seats:                                             # a district's own row; its seats follow
                school = re.fullmatch(r"BHC SCHOOL DIST # ?(\d+)(?: \([^()]+\))?", label)
                if table == "SCHOOL DISTRICTS":
                    if not school:
                        lfail(f"{where}: a school district label this loader does not know; stopped")
                    district, said, cur = f"BIG HORN COUNTY SCHOOL DISTRICT #{school.group(1)}", label, None
                else:
                    district, said, cur = re.sub(r"# (?=\d)", "#", label.upper()), None, None
                    if "DISTRICT" not in district.split():
                        district = re.sub(r"^(.*?)(?=( #\d+)?$)", lambda g: g.group(1) + " DISTRICT", district, count=1)
                continue
            seat = name if re.fullmatch(r"(Director|Trustee|Supervisor|Seat|Position) \w+", name) else ""    # else the label, printed again
            if district is None or table in ("CEMETERY DISTRICTS", "FIRE DISTRICTS"):
                # one row for the district and its seat
                fire = re.fullmatch(r"FIRE # ?(\d+) (.+)", label)
                if fire:
                    whole, said_ = f"FIRE DISTRICT #{fire.group(1)}", label
                else:
                    whole, said_ = re.sub(r"# (?=\d)", "#", re.sub(r"\bDIST\b", "DISTRICT", label.upper())), None
                    if "DISTRICT" not in whole.split():
                        whole += " DISTRICT"
                title = f"{whole} {seat.upper()}".strip()
            else:
                title, said_ = f"{district} {label.upper()}", said
            cur = contest(title, int(years.group(1)), int(seats.group(1)), "NONPARTISAN OFFICES", False, said_)
    return out, n_cands


# ---------- what a contest's title means ----------

STATE_TITLE = re.compile(r"^(UNITED STATES (SENATOR|REPRESENTATIVE)|GOVERNOR|SECRETARY OF STATE|STATE AUDITOR|STATE TREASURER|"
                         r"(STATE )?SUPERINTENDENT OF PUBLIC INSTRUCTION|STATE SENAT(E|OR)\b.*|STATE REPRESENTATIVE\b.*|(SENATE|HOUSE) DISTRICT \d+)$")
COUNTY_OFFICES = (      # the title once "X COUNTY" is set aside -> (office_kind, the office in plain words)
    (r"COMMISSIONERS?", "county_commissioner", "County Commissioner"),
    (r"CORONER", "coroner", "County Coroner"),
    (r"(AND PROSECUTING )?ATTORNEY", "county_attorney", "County Attorney"),
    (r"(\w+ JUDICIAL )?DISTRICT ATTORNEY", "district_attorney", "District Attorney"),
    (r"SHERIFF", "sheriff", "County Sheriff"),
    (r"CLERK OF (THE )?DISTRICT COURT", "clerk_of_court", "Clerk of District Court"),
    (r"CLERK", "county_clerk", "County Clerk"),
    (r"TREASURER", "county_treasurer", "County Treasurer"),
    (r"ASSESSOR", "county_assessor", "County Assessor"),
)
MUNI_WORD = re.compile(r"\b(MAYOR|COUNCIL|COUNCIL ?PERSONS?|COUNCIL ?MEMBERS?|COUNCILMAN|COUNCILWOMAN)\b")
MUNI_STRIP = re.compile(r"\b(MAYOR|(?:CITY |TOWN )?COUNCIL(?: ?MEMBERS?| ?PERSONS?|MAN|WOMAN)?|MEMBERS? OF|CITY OF|TOWN OF|CITY|TOWN|MEMBERS?)\b")
# a special district's kind, by a word of its title: (pattern, level, office_kind, the board in plain words, the kind in plain words)
SPECIAL_KINDS = (
    (r"\bCONSERVATION\b|\bNATURAL RESOURCE\b", "soil_water", "soil_water", "Conservation District", "conservation district"),
    (r"\bRURAL HEALTH( ?CARE)?\b", "other", "rural_health_board", "Rural Health Care District", "rural health care district"),
    (r"\bHOSPITAL\b|\bMEDICAL SERVICES? DISTRICT\b", "hospital", "hospital_board", "Hospital District", "hospital district"),
    (r"\bEMERGENCY MEDICAL\b|\(EMS\)", "other", "ems_board", "Emergency Medical Services District", "emergency medical services district"),
    (r"\bFIRE\b", "other", "fire_board", "Fire District", "fire district"),
    (r"\bCEMETERY\b", "other", "cemetery_board", "Cemetery District", "cemetery district"),
    (r"\bMUSEUM\b|\bMUSUEM\b", "other", "museum_board", "Museum District", "museum district"),
    (r"\bSENIOR CITIZENS?\b", "other", "senior_citizens_board", "Senior Citizens Service District", "senior citizens service district"),
    (r"\bWATER (AND|&) SEWER\b", "other", "water_board", "Water and Sewer District", "water and sewer district"),
    (r"\bSEWER\b", "other", "sanitary_board", "Sewer District", "sewer district"),
    (r"\bIMPROVEMENT\b", "other", "improvement_board", "Improvement District", "improvement district"),
)
ONE_HOLDER = {"coroner", "county_attorney", "district_attorney", "sheriff", "clerk_of_court", "county_clerk", "county_treasurer", "county_assessor", "mayor"}
ROLE_WORDS = re.compile(r"\b(BOARD OF TRUSTEES|SUPERVISORS?|TRUSTEES?|DIRECTORS?|POSITION|HOSPITAL DISTRICT|CONSERVATION DISTRICT|SERVICE DISTRICT|DISTRICT|"
                        r"BOARD|MEMBERS?)\b")
GENERIC_WORDS = {"SOIL", "RESOURCE", "NATURAL", "CONSERVATION", "DISTRICT", "HOSPITAL", "RURAL", "FIRE", "PROTECTION", "CONTROL", "SENIOR", "CITIZEN",
                 "CITIZENS", "SERVICE", "SERVICES", "CEMETERY", "MUSEUM", "WATER", "AND", "SEWER", "IMPROVEMENT", "HEALTH", "CARE", "HEALTHCARE",
                 "EMERGENCY", "MEDICAL", "(EMS)", "NO.", "NO", "COMMUNITY", "COLLEGE"}
SMALL_WORDS = {"And": "and", "Of": "of", "The": "the", "For": "for"}
SPELLED = {"MUSUEM": "MUSEUM"}                                  # a word one clerk's ballot misspells; the race's note says so


class Unread(Exception):
    """A contest whose title this loader does not know: it is named in sl_gaps, never guessed at."""


def title_words(caps):
    """A district's name in ordinary capitals: WHITE MOUNTAIN WATER AND SEWER DISTRICT -> White Mountain Water and Sewer
    District; initials in brackets stay as printed: (EMS)."""
    printed = (caps or "").split()
    words = ordinary(caps).split()
    return " ".join(p if re.fullmatch(r"\([A-Z]{2,5}\)", p) else SMALL_WORDS.get(w, w) if i else w for i, (w, p) in enumerate(zip(words, printed)))


def term_of(term_lines, title):
    """{"years", "unexpired", "undetermined"} from a contest's term lines (and a title that says 2YR UNEXPIRED)."""
    text = " ".join(term_lines).upper()
    m = re.search(r"\((\d)\)", text)
    years = int(m.group(1)) if m else next((NUMBER_WORDS[w] for w in re.findall(r"[A-Z]+", text) if w in NUMBER_WORDS), None)
    return {"years": years, "unexpired": "UNEXPIRED" in text or bool(re.search(r"\bUNEXPIRED\b", title)), "undetermined": "DETERMINED" in text}


def seat_text(rest):
    """What is left of a title once the office's and the district's own words are taken out, in ordinary capitals."""
    s = re.sub(r"[,;]", " ", rest or "")
    s = re.sub(r"\bAT[- ]LARGE\b", "At Large", s, flags=re.I)
    s = re.sub(r"\s+", " ", s).strip(" -")
    return " ".join(w if w in ("At", "Large") else ordinary(w) for w in s.split()) or None


def district_and_seat(text):
    """("Area E", "At Large") from "Area E At Large"; (None, "Rural") from "Rural"; ("Hulett District", None)."""
    if not text:
        return None, None
    at = "At Large" if re.search(r"\bAt Large\b", text) else None
    rest = re.sub(r"\s+", " ", re.sub(r"\bAt Large\b", " ", text)).strip(" -")
    if rest in ("", "District"):
        return None, at
    if re.search(r"\b(Area|District|Subdistrict|Ward)\b", rest):
        return rest, at
    return None, rest + (" At Large" if at else "")


def classify(title_lines, term_lines, section, county, county_names, places, partisan=False, labels=None):
    """What one contest title of a county's list is. `county` is the county whose list it is on ("Hot Springs");
    county_names maps every county's name in capitals to its GEOID; places is the Census list of cities and towns;
    partisan says the list prints parties beside the contest's names (only the county's own offices, and the state's,
    are on the party ballot). Returns {"what": "state"} for an office the state part of this loader (or the federal
    loader) carries, else a description of the office and of whose office it is; Unread when the title is not one this
    loader knows."""
    raw = re.sub(r"\s+", " ", " ".join(title_lines)).strip()
    t = re.sub(r"^FOR ", "", raw.upper())                           # some ballots word a title "FOR COUNTY COMMISSIONERS"
    term = term_of(term_lines, t)
    t = re.sub(r"\s*-?\s*\b2 ?YR\b", " ", t)
    t = re.sub(r"\bUNEXPIRED\b", " ", t)
    t = re.sub(r"\s+", " ", t).strip(" -,")
    if STATE_TITLE.match(t) or re.match(r"(FEDERAL|STATE)\b", (section or "").upper()):
        return {"what": "state", "title": t}                        # the ballot's own FEDERAL OFFICES and STATE OFFICES sections, however a title is worded
    out = {"term": term, "raw": raw}
    cnames = "|".join(re.escape(n) for n in sorted(county_names, key=len, reverse=True))
    here = county_names[county.upper()]

    # the county's own offices
    bare = re.sub(rf"^(?:{re.escape(county.upper())} )?COUNTY ", "", t)
    for pat, kind, shown in COUNTY_OFFICES:
        if re.fullmatch(pat, bare) and (bare != t or "COUNTY" in (section or "").upper() or partisan):
            return dict(out, what="county", level="county", office_kind=kind, office=shown, district=None, seat=None, partisan=1, jkey=("county", here))

    # a city's or town's mayor and council
    if "SCHOOL" not in t and "COLLEGE" not in t and MUNI_WORD.search(t):
        ward = re.search(r"\bWARD (\w+)\b", t)
        at = re.search(r"\bAT[- ]LARGE\b", t)
        rest = MUNI_STRIP.sub(" ", re.sub(r"\bWARD \w+\b|\bAT[- ]LARGE\b", " ", t))
        hit, by_label = places.get(squash(rest), []), None
        if not squash(rest) and labels:
            # The title names no city or town ("FOR MAYOR"). It is the one place of this county whose name stands in the
            # label of every precinct ballot that carries the contest, when there is exactly one such place.
            hit = [p for ps in places.values() for p in ps if here in p[2]
                   and all(re.search(rf"\b{re.escape(p[1].rsplit(' ', 1)[0].upper())}\b", lab.upper()) for lab in labels)]
            by_label = hit[0][1].rsplit(" ", 1)[0] if len(hit) == 1 else None
        if len(hit) != 1:
            raise Unread("a mayor or council contest whose city or town is not found once in the Census Bureau's list of places")
        mayor = bool(re.search(r"\bMAYOR\b", t))
        return dict(out, what="city", level="city", office_kind="mayor" if mayor else "council", office="Mayor" if mayor else "Council Member",
                    place=hit[0], district=f"Ward {ward.group(1)}" if ward else None, seat="At Large" if at else None, partisan=0,
                    jkey=("city", hit[0][0]), by_label=by_label)

    # a community college district's trustees
    if "COLLEGE" in t:
        m = re.search(r"^(.*?\bCOLLEGE)(?: DISTRICT)?\b", t)
        name = m.group(0).strip()
        rest = re.sub(r"\b(BOARD OF TRUSTEES|TRUSTEES?|BOARD|MEMBERS?)\b", " ", t[m.end():])
        district, seat = district_and_seat(seat_text(rest))
        generic = all(w in GENERIC_WORDS for w in name.split())
        return dict(out, what="college", level="other", office_kind="college_board", office="Community College Trustee", district=district, seat=seat,
                    partisan=0, name=name, generic=generic, kind_words="community college district", named=None,
                    jkey=("special", lslug(re.sub(r"\bDISTRICT\b", " ", name)), here if generic else None))

    # a school district's trustees
    if "SCHOOL" in t:
        m = re.search(rf"\b({cnames}) COUNTY\b", t) or re.search(rf"^({cnames}) (?=SCHOOL DISTRICT)", t)
        home = county_names[m.group(1)] if m else here
        rest = (t[:m.start()] + " " + t[m.end():]) if m else t
        short = re.search(r"\b([A-Z]{1,3})CSD\b", rest)             # a county's own ballots write its districts short: CCSD #1
        if short:
            if m or short.group(1) != "".join(w[0] for w in county.upper().split()):
                raise Unread("a school district written short whose county the title does not settle")
            rest = rest[:short.start()] + " " + rest[short.end():]
        num = re.search(r"(?:#|\bNO\.?|\bNUMBER)\s*(\d+)\b", rest) or re.search(r"\bDISTRICT (\d+)\b", rest)
        n = int(num.group(1)) if num else None                      # a county with one school district may leave the number off its ballot
        if num:
            rest = re.sub(rf"\bDISTRICT\s*(?:#|NO\.?|NUMBER)?\s*{n}\b", " ", rest, count=1)
            if re.search(rf"(?:#|\bNO\.?|\bNUMBER)\s*{n}\b", rest):
                rest = re.sub(rf"(?:#|\bNO\.?|\bNUMBER)\s*{n}\b", " ", rest, count=1)
                rest = re.sub(r"\bDISTRICT\b", " ", rest, count=1)
        else:
            rest = re.sub(r"\bDISTRICT\b", " ", rest, count=1)
        rest = re.sub(r"\b(SCHOOL|TRUSTEES?|BOARD|MEMBERS?|OF)\b", " ", rest)
        district, seat = district_and_seat(seat_text(rest))
        return dict(out, what="school", level="school", office_kind="school_board", office="School Board Trustee", district=district, seat=seat,
                    partisan=0, home=home, number=n, jkey=("school", home, n))

    # every other district
    for pat, level, kind, board, words in SPECIAL_KINDS:
        k = re.search(pat, t)
        if not k:
            continue
        t2 = re.sub(r"^(TRUSTEES?|SUPERVISORS?|DIRECTORS?)\s*,\s*", "", t)
        m = re.search(r"\bDISTRICT\b(\s+(?:#\s*\d+|NO\.?\s*\d+|\d+[A-Z]?)(?=\s|$))?", t2)
        if m:
            name, rest = t2[:m.end()].strip(), t2[m.end():]
        else:
            k2 = re.search(pat, t2)
            name, rest = t2[:k2.end()].strip(), t2[k2.end():]
        # what the title calls a member of the board; where it does not say, a conservation district's are supervisors
        # (the word every list that names them uses) and any other board's are simply its members
        given = "Supervisor" if re.search(r"\bSUPERVISORS?\b", t) else "Trustee" if re.search(r"\bTRUSTEES?\b", t) else \
            "Director" if re.search(r"\bDIRECTORS?\b", t) else None
        role = given or ("Supervisor" if kind == "soil_water" else "Board Member")
        left = seat_text(ROLE_WORDS.sub(" ", rest))
        if left and re.fullmatch(r"\d+", left):
            left = f"{role} {left}"
        district, seat = district_and_seat(left)
        tokens = [SPELLED.get(w, w) for w in name.split()]
        is_generic = lambda ws: all(w in GENERIC_WORDS or re.fullmatch(r"#?\d+[A-Z]?", w) for w in ws)   # noqa: E731
        generic = is_generic(tokens)
        # The key a district is filed under. "X COUNTY FIRE DISTRICT #1" on any list, "X FIRE DISTRICT #1" on X County's
        # own and a bare "FIRE DISTRICT #1" on X County's own are one district, so the county's name goes into the key's
        # county and not into its words; "PROTECTION" is left out of the words too (one list writes FIRE PROTECTION
        # DISTRICT where another writes FIRE DISTRICT).
        named = re.match(rf"^({cnames}) COUNTY\b", name)
        own = county.upper().split()
        if named:
            whose, core = county_names[named.group(1)], [SPELLED.get(w, w) for w in name[named.end():].split()]
        elif tokens[:len(own)] == own and len(tokens) > len(own) and is_generic(tokens[len(own):]):
            whose, core = here, tokens[len(own):]
        else:
            whose, core = (here if generic else None), tokens
        canon = " ".join(w for w in core if w not in ("COUNTY", "DISTRICT", "PROTECTION"))
        canon = re.sub(r"\bHEALTH CARE\b", "HEALTHCARE", canon).replace("&", "AND").replace("CITIZENS", "CITIZEN").replace("SERVICES", "SERVICE")
        canon = re.sub(r"\bNO\.?\s*(?=\d)", "", canon)
        if not lslug(canon):
            raise Unread("a district whose title gives it no words of its own")
        return dict(out, what="special", level=level, office_kind=kind, office=f"{board} {role}", role_given=bool(given), district=district, seat=seat,
                    partisan=0, name=name, generic=generic, kind_words=words, named=county_names[named.group(1)] if named else None,
                    jkey=("special", lslug(canon), whose))
    raise Unread("a contest whose title this loader does not know")


# ---------- the counties' lists into rows ----------

def csv_office(title):
    """A state or federal contest's title on a county ballot, as the Secretary of State's candidates file words the
    office (the clerks word these titles in several ways: STATE GOVERNOR; STATE SENATOR SENATE DISTRICT #6); None when
    the title is none of them."""
    federal = re.match(r"UNITE[DS] STATES\b", title)
    m = re.search(r"#?0*(\d+)$", title)                            # ... HOUSE DISTRICT 49; ... SENATE DISTRICT #6; STATE REPRESENTATIVE 20
    if m and not federal:
        chamber = "SENATOR" if re.search(r"\bSENAT(E|OR)\b", title) else "REPRESENTATIVE" if re.search(r"\b(HOUSE|REPRESENTATIVES?)\b", title) else None
        return f"STATE {chamber} {int(m.group(1)):02d}" if chamber else None
    for word, office in (("SENATOR", "UNITED STATES SENATOR"), ("REPRESENTATIVE", "UNITED STATES REPRESENTATIVE")):
        if federal and word in title:
            return office
    for word, office in (("GOVERNOR", "GOVERNOR"), ("SECRETARY OF STATE", "SECRETARY OF STATE"), ("AUDITOR", "STATE AUDITOR"), ("TREASURER", "STATE TREASURER"),
                         ("SUPERINTENDENT", "SUPERINTENDENT OF PUBLIC INSTRUCTION")):
        if word in title:
            return office
    return None


def county_list_from_ballots(name, geoid, manifest, folder):
    """One county's sample ballots read and merged: {"contests": {(title, term): {...}}, "retained": {judge as printed:
    ballots}, "styles", "files", "names_read", "questions", "per_file": {file: names read}}. BallotError when a kept
    file cannot be read after all, or holds a line that reads like contact details."""
    merged, retained, per_file = OrderedDict(), Counter(), {}
    styles = names_read = questions = 0
    twice = False
    for d in manifest["docs"]:
        ballots = read_ballots(open(os.path.join(folder, d["file"]), "rb").read(), d["file"])
        per_file[d["file"]] = 0
        for b in ballots:
            if any(contact_like(t, True) for t in kept_texts(b)):
                raise BallotError(f"{d['file']}: a line that reads like contact details")
            styles += 1
            twice = twice or b["drawn_twice"]
            for c in b["contests"]:
                if c["kind"] == "measure":
                    questions += 1
                    continue
                if c["kind"] == "retention":
                    retained[re.sub(r"\s+", " ", c["judge"]).strip().upper()] += 1
                    continue
                key = (" / ".join(c["title"]), " ".join(c["term"]))
                e = merged.setdefault(key, {"title": c["title"], "term": c["term"], "section": c["section"], "votes": Counter(), "sets": Counter(),
                                            "styles": 0, "file": d["file"], "wrapped": False, "labels": []})
                e["votes"][c["vote_for"]] += 1
                e["sets"][frozenset(c["cands"])] += 1
                e["styles"] += 1
                e["labels"].append(b["style"])
                e["wrapped"] = e["wrapped"] or c["wrapped"]
                names_read += len(c["cands"])
                per_file[d["file"]] += len(c["cands"])
    return {"contests": merged, "retained": retained, "styles": styles, "files": len(manifest["docs"]), "names_read": names_read,
            "questions": questions, "per_file": per_file, "drawn_twice": twice}


def laramie_contests(manifest, folder):
    """Laramie County's three rosters as contests like a ballot's: ([(section, title lines, term lines, [(party, name)])],
    rows read, withdrawn left off, the source file of each contest)."""
    out, rows_in, gone = OrderedDict(), 0, 0
    for d in manifest["docs"]:
        kept = json.load(open(os.path.join(folder, d["file"]), encoding="utf-8"))
        cols = kept["columns"]
        for n, row in enumerate(kept["rows"], start=1):
            r = dict(zip(cols, row))
            rows_in += 1
            office, name = r.get("Office", ""), r.get("Ballot Name", "")
            if not office or not name:
                lfail(f"row {n} of Laramie County's {d['key']} roster has no office or no ballot name; stopped (the row is not printed)")
            if r.get("Withdrawn Date"):
                gone += 1
                continue
            term = []
            if d["key"] == "county":
                section, title = "COUNTY OFFICES", office.upper()
            elif d["key"] == "municipal":
                if not r.get("Municipality"):
                    lfail(f"row {n} of Laramie County's municipal roster names no municipality; stopped")
                section, title = "MUNICIPAL OFFICES", f"{r['Municipality']} {office}".upper()
            else:
                section = "DISTRICTS"
                m = re.fullmatch(r"(.+?),\s*(\d)\s*-\s*Year Term", office)
                if m:
                    office, term = m.group(1), [f"{NUMBER_NAMES[int(m.group(2))].upper()} ({m.group(2)}) YEAR TERM"]
                title = re.sub(r"^LCSD ?#(\d+)", r"LARAMIE COUNTY SCHOOL DISTRICT #\1", office.upper())
            if bool(r.get("Party")) != (d["key"] == "county"):
                lfail(f"row {n} of Laramie County's {d['key']} roster: a party where none is expected, or none where one is; stopped")
            e = out.setdefault((title, " ".join(term)), {"title": [title], "term": term, "section": section, "cands": [], "file": d["file"], "key": d["key"]})
            e["cands"].append((r.get("Party") or None, name))
    return out, rows_in, gone


def local_level(folder, cmap, say=print, refresh=False):
    """Wyoming's judges standing for retention and its county and local contests as rows ready to write (races,
    candidates, places, sources, gaps, notes) with the counts behind them. Nothing here touches the database."""
    os.makedirs(folder, exist_ok=True)
    full = {g: n for g, n in cmap.values()}                                  # "56017" -> "Hot Springs County"
    short = {g: n[:-len(" County")] for g, n in full.items()}                # "56017" -> "Hot Springs"
    county_names = {n.upper(): g for g, n in short.items()}
    problems, gaps, sources, counts = [], [], [], Counter()
    name_notes = {}                                                          # (source id, name as read) -> a note for that candidate

    districts, t5path = judicial_districts(folder, cmap, say, refresh)
    kept = sos_csv(folder, say, refresh)
    j_races, j_cands, j_places, judges, j_gone, csv_others = retention_rows(kept, districts, cmap, say)
    places, ppath, p_rows = census_places(folder, cmap, say)
    csv_by_office = defaultdict(list)
    for _election, office, party, name, withdrawn in kept["rows"]:
        if party and not withdrawn:
            csv_by_office[office].append(name)

    def county_gap(geoid, reason, url):
        gaps.append((STATE, "county", geoid, full[geoid], "county, city, school and special district races", reason, url))

    def race_gap(geoid, what, reason, url, names):
        counts["unread"] += 1
        counts["names_in_gaps"] += names
        gaps.append((STATE, "race", f"2026-{STATE}-{geoid}-unread-{counts['unread']}", full[geoid], what, reason, url))

    instances, loaded = [], OrderedDict()

    # ---- the counties whose clerks post sample ballots
    for name, cfg in BALLOT_PAGES.items():
        geoid = county_names[name.upper()]
        f3 = geoid[2:]
        m = county_ballots(name, f3, cfg, folder, say, refresh)
        if m is None:
            county_gap(geoid, NOT_POSTED + "The clerk's page could not be reached when this was built, so the county is not loaded yet.", cfg["page"])
            continue
        # a linked file that is not a November 3, 2026 general election ballot (the primary's sample, say) is simply not one of the ballots
        blocking = [r for r in m["refused"] if r["why"] not in ("not-general", "other-election")]
        if blocking or not m["docs"]:
            whys = {r["why"] for r in blocking}
            if not blocking:
                reason = WHY_NOT.get(name, NOT_POSTED + "No November sample ballot was linked from the clerk's page when this was built, so the county is "
                                     "not loaded yet.")
            elif whys == {"scan"}:
                reason = ("The county clerk posts this county's November 3 sample ballots as scanned pictures, which have no text a program can "
                          "read, so its contests are not loaded yet.")
            elif "unreached" in whys:
                reason = NOT_POSTED + "Some of its sample ballot files could not be fetched when this was built, so the county is not loaded yet."
            elif whys == {"not-pdf"}:
                reason = NOT_POSTED + ("A link on the clerk's page that says sample ballot did not lead to a ballot file this loader can read when "
                                       "this was built, so the county is not loaded yet.")
            else:
                reason = ("The county clerk posts this county's November 3 sample ballots, but the text inside the files as posted is printed over "
                          "itself or laid out in a way this loader cannot read line by line, so its contests are not loaded yet.")
            county_gap(geoid, reason, cfg["page"])
            counts["files_refused"] += len(blocking)
            continue
        try:
            got = county_list_from_ballots(name, geoid, m, folder)
        except BallotError as e:
            county_gap(geoid, NOT_POSTED + "One of its sample ballot files could not be read by this loader, so the county is not loaded yet.", cfg["page"])
            problems.append(f"{name} County: {str(e)[:200]}")
            continue
        src_of = {d["file"]: f"wy-{f3}-clerk-2026-sample-ballot-{d['file'][-6:-4]}" for d in m["docs"]}
        mine, mine_problems, mine_gaps, state_bad, local_names = [], [], [], 0, 0
        # the judges asked about on the county's ballots, against the Secretary of State's file
        expect = [j["name"] for j in judges if j["counties"] is None or geoid in j["counties"]]
        if same_names(expect, list(got["retained"])) is None:
            mine_problems.append(f"{name} County's sample ballots ask about {len(got['retained'])} judges and the Secretary of State's file has "
                                 f"{len(expect)} for this county, or a name is printed differently; the judges stored are the Secretary of State's")
        if any(k != got["styles"] for k in got["retained"].values()):
            mine_problems.append(f"{name} County: a retention question is not on every one of its {got['styles']} ballots")
        for (_tkey, _termkey), e in got["contests"].items():
            sets = list(e["sets"])
            partisan = any(p for s in sets for p, _n in s)
            try:
                c = classify(e["title"], e["term"], e["section"], name, county_names, places, partisan, e["labels"])
            except Unread as u:
                mine_gaps.append(("one contest on the county's sample ballots", f"The county's sample ballots carry {u}; it is not shown rather than "
                                  "guessed at.", max(len(s) for s in sets)))
                mine_problems.append(f"{name} County: a contest title not known to the loader (on {e['styles']} ballots)")
                continue
            if c["what"] == "state":
                counts["state_contests_checked"] += 1
                want = csv_by_office.get(csv_office(c["title"]) or "")
                if want is None:
                    state_bad += 1
                    mine_problems.append(f"{name} County: a contest in the ballot's federal or state section is not one of the offices in the Secretary "
                                         f"of State's file ({csv_office(c['title']) or 'title not known'})")
                elif len(sets) != 1 or same_names(want, [n for _p, n in sets[0]]) is None:
                    state_bad += 1
                    mine_problems.append(f"{name} County's ballots and the Secretary of State's file do not list the same candidates for "
                                         f"{csv_office(c['title']).title()}")
                continue
            if len(sets) != 1:
                mine_gaps.append((f"{c['office']} contest", "The county's precinct ballots print different sets of names under one contest title, so the "
                                  "contests could not be told apart; they are not shown rather than guessed at.", sum(len(s) for s in sets)))
                mine_problems.append(f"{name} County: different names under one contest title")
                continue
            cands = sorted(sets[0], key=lambda pn: pn[1])
            if len(e["votes"]) != 1:
                mine_problems.append(f"{name} County: one contest is printed with different vote-for lines on different ballots")
            mv = re.search(r"\((\d+)\)", e["votes"].most_common(1)[0][0])
            if cands and bool(c["partisan"]) != partisan:
                mine_gaps.append((f"{c['office']} contest", "The ballot prints this contest with parties where the office has none, or without them where "
                                  "it has; it is not shown rather than guessed at.", len(cands)))
                mine_problems.append(f"{name} County: a contest whose party lines do not fit its office")
                continue
            mine.append(dict(c, county=geoid, src=src_of[e["file"]], seats=int(mv.group(1)) if mv else None,
                             cands=[(BALLOT_PARTIES[p] if p else None, n) for p, n in cands], caps=True, wrapped=e["wrapped"], listed="ballot"))
            local_names += len(cands)
        if got["drawn_twice"] and state_bad:
            # a file that draws its text twice is trusted only when every state and federal name on it reads exactly as the Secretary of State's file has it
            county_gap(geoid, "The county clerk posts this county's November 3 sample ballots in files that draw their text more than once, and the "
                       "names read from them did not all agree with the Secretary of State's own list, so its contests are not loaded yet.", cfg["page"])
            problems.append(f"{name} County: {state_bad} state or federal contests on its ballots (files that draw their text twice) did not read as "
                            "the Secretary of State's file has them; the county is not loaded")
            continue
        instances += mine
        problems += mine_problems
        for what, reason, n in mine_gaps:
            race_gap(geoid, what, reason, cfg["page"], n)
        loaded[geoid] = {"how": "sample ballots", "styles": got["styles"], "files": got["files"], "names": local_names, "questions": got["questions"],
                         "names_read": got["names_read"]}
        counts["ballots"] += got["styles"]
        counts["ballot_files"] += got["files"]
        counts["names_read"] += got["names_read"]
        counts["files_set_aside"] += len(m["refused"]) - len(blocking)
        sources.append((f"wy-{f3}-clerk-2026-sample-ballot-page", STATE, "official page", f"{name} County Clerk",
                        f"The {name} County Clerk's page that links the November 3, 2026 sample ballots", m["page"], "", m["page_fetched"], m["page_sha256"],
                        m["found"], "Read only for the links to the sample ballot files: each link's own words and address. The page is not kept."
                        + (" The page links a shared folder, whose list of files is read for the general election ballots' names." if cfg.get("folder_files") else "")))
        for k, d in enumerate(m["docs"], start=1):
            # the clerk's own words for the file name a precinct, and a precinct can be named after a road: such words stay out of the texts
            said = d["words"] if not (contact_like(d["words"], True) or BUILDER_STREET.search(d["words"])) else None
            sources.append((src_of[d["file"]], STATE, "official sample ballot", f"{name} County Clerk",
                            f"2026 General Election sample ballot, {name} County" + (f" (file {k} of {len(m['docs'])})" if len(m["docs"]) > 1 else ""),
                            tidy_url(d["url"]), "", d["fetched"], d["sha256"], got["per_file"][d["file"]],
                            "The ballot as the county clerk printed it for one precinct, or for every precinct of the county. Section headings, contest "
                            "titles, the term and vote-for lines, party lines, names and write-in lines are read; the wording of ballot questions is "
                            "not. A ballot has no contact columns. The row count is the names printed on it, state and federal offices included. Names "
                            "rotate by precinct, so the order they are printed in is not kept."
                            + (" The file draws each letter more than once; each is read once." if got["drawn_twice"] else "")
                            + (f" The clerk's link to it reads \"{said}\"." if said else "")))

    # ---- Laramie County: the clerk's three rosters
    lgeo = county_names["LARAMIE"]
    lm = laramie_rosters(folder, say, refresh)
    if lm is None:
        county_gap(lgeo, NOT_POSTED + "Its candidate rosters could not be read when this was built, so the county is not loaded yet.", LARAMIE_PAGE)
    else:
        contests, rows_in, gone = laramie_contests(lm, folder)
        placed = unplaced = 0
        for (_title, _term), e in contests.items():
            try:
                c = classify(e["title"], e["term"], e["section"], "Laramie", county_names, places, e["key"] == "county")
            except Unread as u:
                race_gap(lgeo, "one contest on the county clerk's roster", f"The county clerk's roster carries {u}; it is not shown rather than guessed at.",
                         LARAMIE_PAGE, len(e["cands"]))
                problems.append("Laramie County: an office on the roster not known to the loader")
                unplaced += len(e["cands"])
                continue
            if c["what"] == "state" or bool(c["partisan"]) != any(p for p, _n in e["cands"]):
                lfail("an office on Laramie County's rosters is a state office, or its party column does not fit the office; stopped")
            instances.append(dict(c, county=lgeo, src=f"wy-021-clerk-2026-general-roster-{e['key']}", seats=None,
                                  cands=sorted(e["cands"], key=lambda pn: pn[1]), caps=False, wrapped=False, listed="candidate roster"))
            placed += len(e["cands"])
        if placed + unplaced + gone != rows_in:
            lfail(f"Laramie County's rosters: {rows_in} rows read, {placed} placed, {unplaced} in contests named among the gaps and {gone} withdrawn; stopped")
        loaded[lgeo] ={"how": "candidate rosters", "styles": 0, "files": len(lm["docs"]), "names": placed, "questions": 0, "rows_in": rows_in,
                        "withdrawn": gone}
        counts["roster_rows"] += rows_in
        counts["withdrawn"] += gone
        sources.append(("wy-021-clerk-2026-general-roster-page", STATE, "official page", "Laramie County Clerk",
                        "The Laramie County Clerk's Candidates page, which links the 2026 general election rosters", lm["page"], "", lm["page_fetched"],
                        lm["page_sha256"], len(lm["docs"]), "Read only for the links to the three general election workbooks. The page is not kept."))
        titles = {"county": "County Offices", "municipal": "Municipal Offices", "districts": "School Board, Conservation District, & Fire District"}
        for d in lm["docs"]:
            sources.append((f"wy-021-clerk-2026-general-roster-{d['key']}", STATE, "official candidate list", "Laramie County Clerk",
                            f"2026 General Election candidate roster, Laramie County: {titles[d['key']]} (XLSX)", tidy_url(d["url"]), "", d["fetched"], d["sha256"],
                            d["rows"],
                            f"A workbook of {d['columns_in_file']} columns. Read by their headings: {', '.join(d['columns'])}. The mailing address, city and "
                            "ZIP, phone, e-mail and website columns are never read, and only the columns named are kept on disk; the fingerprint is of "
                            "the workbook as it came. A row with a withdrawal date is left off. The roster does not say how many are to be elected and "
                            "gives no ballot order."))

    # ---- Natrona County: the clerk's roster
    ngeo = county_names["NATRONA"]
    nm = natrona_roster(folder, say, refresh)
    if nm is None:
        county_gap(ngeo, NOT_POSTED + "Its candidate roster could not be read when this was built, so the county is not loaded yet.", NATRONA_PAGE)
    else:
        contests, rows_in, listed = natrona_contests(nm)
        if same_names([j["name"] for j in judges if j["counties"] is None or ngeo in j["counties"]], listed) is None:
            problems.append("Natrona County's roster and the Secretary of State's file do not list the same judges standing for retention; the judges "
                            "stored are the Secretary of State's")
        placed = unplaced = 0
        for (_title, _term), e in contests.items():
            try:
                c = classify(e["title"], e["term"], e["section"], "Natrona", county_names, places, e["county_office"])
            except Unread as u:
                race_gap(ngeo, "one contest on the county clerk's roster", f"The county clerk's roster carries {u}; it is not shown rather than guessed at.",
                         NATRONA_PAGE, len(e["cands"]))
                problems.append("Natrona County: an office on the roster not known to the loader")
                unplaced += len(e["cands"])
                continue
            if c["what"] == "state" or bool(c["partisan"]) != e["county_office"]:
                lfail("an office on Natrona County's roster is a state office, or is filed with the county's offices and is not one; stopped")
            instances.append(dict(c, county=ngeo, src="wy-025-clerk-2026-general-roster", seats=None, cands=sorted(e["cands"], key=lambda pn: pn[1]),
                                  caps=False, wrapped=False, listed="candidate roster"))
            placed += len(e["cands"])
        if placed + unplaced + len(listed) != rows_in:
            lfail(f"Natrona County's roster: {rows_in} rows read, {placed} placed, {unplaced} in contests named among the gaps and {len(listed)} judges; stopped")
        loaded[ngeo] = {"how": "candidate rosters", "styles": 0, "files": 1, "names": placed, "questions": 0, "rows_in": rows_in, "withdrawn": 0}
        counts["roster_rows"] += rows_in
        counts["roster_judges"] += len(listed)
        sources.append(("wy-025-clerk-2026-general-roster-page", STATE, "official page", "Natrona County Clerk",
                        "The Natrona County Clerk's Candidate Information page, which links the 2026 general election roster", nm["page"], "",
                        nm["page_fetched"], nm["page_sha256"], 1, "Read only for the link to the roster. The page is not kept."))
        sources.append(("wy-025-clerk-2026-general-roster", STATE, "official candidate list", "Natrona County Clerk",
                        "Natrona County 2026 General Candidate Roster (PDF)", tidy_url(nm["url"]), "", nm["fetched"], nm["sha256"], rows_in,
                        f"A table of {nm['columns_in_file']} columns. Four are read, their edges taken from each page's own row of headings: whose office "
                        "it is (the first column, which has no heading), Office, Party and Ballot Name. The mailing address, phone number and e-mail "
                        "columns are never read, and only the four columns are kept on disk; the fingerprint is of the file as it came. The party "
                        "column is kept only for the county's own offices: the roster prints a party beside every name, and city, school and "
                        f"district offices carry none on the ballot. {len(listed)} rows are judges standing for retention, taken from the Secretary of "
                        "State's file instead. The roster has no column for withdrawals, does not say how many are to be elected and gives no "
                        "ballot order."))

    # ---- Fremont County: the clerk's list of school, college and special district filings (the county's own offices and
    # the cities' are not on it, and stay a gap); skipped once the clerk posts sample ballots this loader can read
    fgeo = county_names["FREMONT"]
    fm = None if fgeo in loaded else fremont_filings(folder, say, refresh)
    if fm is not None:
        contests, rows_in, gone, county_rows, requoted, odd_quotes = fremont_contests(fm)
        for fixed, printed in requoted.items():
            name_notes[("wy-013-clerk-2026-general-filings", fixed)] = (f"The clerk's list prints this name as {printed}; the quotation marks are put "
                                                                        "back around the nickname here.")
        if odd_quotes:
            problems.append(f"Fremont County: {odd_quotes} names on the list of filings carry a quotation mark in an unusual place; kept as printed")
        placed = unplaced = 0
        for (_title, _term), e in contests.items():
            try:
                c = classify(e["title"], e["term"], e["section"], "Fremont", county_names, places, False)
            except Unread as u:
                race_gap(fgeo, "one contest on the county clerk's list of filings", f"The county clerk's list of filings carries {u}; it is not shown "
                         "rather than guessed at.", FREMONT_PAGE, len(e["cands"]))
                problems.append("Fremont County: an office on the list of filings not known to the loader")
                unplaced += len(e["cands"])
                continue
            if c["what"] not in ("school", "college", "special"):
                lfail("an office on Fremont County's list of filings is not a school, college or special district office; stopped")
            instances.append(dict(c, county=fgeo, src="wy-013-clerk-2026-general-filings", seats=None, cands=sorted(e["cands"], key=lambda pn: pn[1]),
                                  caps=all(n == n.upper() for _p, n in e["cands"]), wrapped=False, listed="list of filings", said=e["said"]))
            placed += len(e["cands"])
        if placed + unplaced + gone + county_rows != rows_in:
            lfail(f"Fremont County's list of filings: {rows_in} rows read, {placed} placed, {unplaced} in contests named among the gaps, {gone} withdrawn "
                  f"and {county_rows} for the county's own offices; stopped")
        for g in [g for g in gaps if g[1] == "county" and g[2] == fgeo]:
            gaps.remove(g)
        gaps.append((STATE, "county", fgeo, full[fgeo], "county and city offices",
                     NOT_POSTED + "The clerk had posted its list of school, college and special district filings, which is loaded, and no November sample "
                     "ballot or list of the candidates for the county's own offices and for mayor and council when this was built, so those offices "
                     "are not loaded yet.", FREMONT_PAGE))
        loaded[fgeo] = {"how": "list of filings", "styles": 0, "files": 1, "names": placed, "questions": 0, "rows_in": rows_in, "withdrawn": gone}
        counts["roster_rows"] += rows_in
        counts["withdrawn"] += gone
        counts["rows_for_county_offices"] += county_rows
        sources.append(("wy-013-clerk-2026-general-filings-page", STATE, "official page", "Fremont County Clerk",
                        "The Fremont County Clerk's Elections and Voting page, which links the list of general election filings", fm["page"], "",
                        fm["page_fetched"], fm["page_sha256"], 1, "Read only for the link to the list. The page is not kept."))
        sources.append(("wy-013-clerk-2026-general-filings", STATE, "official candidate list", "Fremont County Clerk",
                        "Candidate Filings - Schools and Special Districts, 2026 General Election, Fremont County (PDF)", fm["url"], "", fm["fetched"],
                        fm["sha256"], rows_in,
                        f"A table of {fm['columns_in_file']} columns. Three are read, their edges taken from each page's own row of headings: District "
                        "Name, Office and Ballot Name. The party and filing date, the residential and mailing address, the phone number and the e-mail "
                        "columns are never read, and only the three columns are kept on disk; the fingerprint is of the file as it came. "
                        f"Rows whose office says withdrawn are left off ({gone}), and rows for one of the county's own offices are not used "
                        f"({county_rows}): the list covers school, college and special district offices only. It does not say how many are to be "
                        "elected and gives no ballot order."
                        + (f" The list prints a nickname's quotation marks one word late; they are put back in {len(requoted)} names, and each of "
                           "those candidates carries a note with the name as printed." if requoted else "")))

    # ---- Big Horn County: the clerk's roster (nine tables); skipped once the clerk posts sample ballots this loader can read
    bgeo = county_names["BIG HORN"]
    bm = None if bgeo in loaded else bighorn_roster(folder, say, refresh)
    if bm is not None:
        contests, rows_in = bighorn_contests(bm)
        placed = unplaced = 0
        for _key, e in contests.items():
            try:
                c = classify(e["title"], e["term"], e["section"], "Big Horn", county_names, places, e["county_office"])
            except Unread as u:
                race_gap(bgeo, "one contest on the county clerk's roster", f"The county clerk's roster carries {u}; it is not shown rather than guessed at.",
                         BIGHORN_PAGE, len(e["cands"]))
                problems.append("Big Horn County: an office on the roster not known to the loader")
                unplaced += len(e["cands"])
                continue
            if c["what"] == "state" or bool(c["partisan"]) != e["county_office"]:
                lfail("an office on Big Horn County's roster is a state office, or is filed with the county's offices and is not one; stopped")
            instances.append(dict(c, county=bgeo, src="wy-003-clerk-2026-general-roster", seats=e["seats"], cands=sorted(e["cands"], key=lambda pn: pn[1]),
                                  caps=bool(e["cands"]) and all(n == n.upper() for _p, n in e["cands"]), wrapped=False, listed="candidate roster",
                                  said=e["said"]))
            placed += len(e["cands"])
        if placed + unplaced != rows_in:
            lfail(f"Big Horn County's roster: {rows_in} candidates read, {placed} placed and {unplaced} in contests named among the gaps; stopped")
        for g in [g for g in gaps if g[1] == "county" and g[2] == bgeo]:
            gaps.remove(g)
        loaded[bgeo] = {"how": "candidate rosters", "styles": 0, "files": 1, "names": placed, "questions": 0, "rows_in": rows_in, "withdrawn": 0}
        counts["roster_rows"] += rows_in
        sources.append(("wy-003-clerk-2026-general-roster-page", STATE, "official page", "Big Horn County Clerk",
                        "The Big Horn County Clerk's Elections page, which links the 2026 general election roster", bm["page"], "", bm["page_fetched"],
                        bm["page_sha256"], 1, "Read only for the link to the roster. The page is not kept."))
        sources.append(("wy-003-clerk-2026-general-roster", STATE, "official candidate list", "Big Horn County Clerk",
                        "Big Horn County 2026 General Candidate Roster (PDF)", tidy_url(bm["url"]), "", bm["fetched"], bm["sha256"], rows_in,
                        f"Nine tables (the county's offices, its towns, and school, cemetery, conservation, fire, hospital, senior citizen and rural "
                        f"health districts) of up to {bm['columns_in_file']} columns, their cells centred under their headings. Read: each table's "
                        "first column (the office, town or district, how many seats and the term), Party, TERM and Candidate Name; a cell is taken "
                        "for a name only when it sits under the middle of the Candidate Name heading. The mailing address, campaign phone and e-mail "
                        "columns are never read, and only the columns named are kept on disk; the fingerprint is of the file as it came. The row "
                        f"count is the candidates; {len(contests)} contests, {sum(1 for e in contests.values() if not e['cands'])} of them listed with "
                        "no candidate. The roster says how many are to be elected; it has no column for withdrawals and gives no ballot order."))

    # ---- every other county: named in sl_gaps
    for geoid in sorted(full):
        if geoid not in loaded and not any(g[1] == "county" and g[2] == geoid for g in gaps):
            county_gap(geoid, WHY_NOT.get(short[geoid], NOT_POSTED + "It is not loaded yet."), COUNTY_PAGES[short[geoid]])

    # ---- whose office each contest is, and one race for a contest printed on two counties' lists
    # A district is one place however its lists word its name: where every list that carries the name is one county's
    # (or names that county), a title that says "X COUNTY ..." and one that leaves the county out are the same district.
    homes = defaultdict(set)
    for inst in instances:
        if inst["jkey"][0] == "special":
            homes[inst["jkey"][1]].add(inst["jkey"][2] or inst["county"])
    for inst in instances:
        if inst["jkey"][0] == "special" and inst["jkey"][2] is None and len(homes[inst["jkey"][1]]) == 1:
            inst["jkey"] = ("special", inst["jkey"][1], inst["county"])
    by_j = OrderedDict()
    for inst in instances:
        by_j.setdefault(inst["jkey"], []).append(inst)
    races, cands, place_rows, seen_ids = [], [], OrderedDict(), set()
    for jkey, insts in by_j.items():
        on = sorted({i["county"] for i in insts})
        first = insts[0]
        spelled = None
        if first["what"] == "county":
            home, jid, jname, cids, pkind = jkey[1], jkey[1], full[jkey[1]], [jkey[1]], None
        elif first["what"] == "city":
            code, pname, pcounties = first["place"]
            home, jid, jname, cids, pkind, psrc = pcounties[0], f"{STATE}-M-{code}", pname, sorted(set(pcounties) | set(on)), "mcd", SRC_PLACES
            if not set(on) <= set(pcounties):
                problems.append(f"{pname}: on the list of a county the Census Bureau's place file does not put it in")
        elif first["what"] == "school":
            home, n = jkey[1], jkey[2]
            jid, jname = (f"{STATE}-S-{home[2:]}-{n}", f"{full[home]} School District #{n}") if n else \
                (f"{STATE}-S-{home[2:]}-school-district", f"School district ({full[home]})")
            cids, pkind, psrc = sorted(set(on) | {home}), "school", first["src"]
        else:
            named = jkey[2]
            home = named or (on[0] if len(on) == 1 else None)
            # the fullest of the names the lists give the district (one list writes FIRE DISTRICT #5 where another writes it out)
            best = sorted({i["name"] for i in insts}, key=lambda v: ("DISTRICT" not in v.split(), -len(v), v))[0]
            fixed = " ".join(SPELLED.get(w, w) for w in best.split())
            spelled = title_words(best) if fixed != best else None
            nameless = all(w in GENERIC_WORDS or re.fullmatch(r"#?\d+[A-Z]?", w) for w in fixed.split())
            if first["what"] == "college" and nameless:
                jname = f"Community college district ({full[home]})"
            elif nameless:
                jname = f"{title_words(fixed)} ({full[home]})"
            elif "DISTRICT" not in fixed.split():
                jname = f"{title_words(fixed)} ({first['kind_words']})"
            else:
                jname = title_words(fixed)
            jid = f"{STATE}-{'H' if first['level'] == 'hospital' else 'X'}-{(home[2:] + '-') if home else ''}{jkey[1]}"
            cids, pkind, psrc = sorted(set(on) | ({named} if named else set())), "hospital" if first["level"] == "hospital" else "special", first["src"]
        if len({(i["level"], i["what"]) for i in insts}) != 1:
            lfail(f"contests of two kinds are filed under one place ({jid}); stopped")
        if pkind:
            place_rows[(pkind, jid)] = (pkind, jid, jname, json.dumps(cids), psrc)

        merged = []
        for inst in sorted(insts, key=lambda i: (i["county"] != home, i["county"])):
            target, pairs = None, None
            for r in merged:
                ya, yb = r["term"]["years"], inst["term"]["years"]         # a roster that gives no term fits any
                if r["office_kind"] != inst["office_kind"] or (ya != yb and None not in (ya, yb)) or inst["county"] in r["counties"]:
                    continue
                # a ballot says when a term is unexpired; a roster may print the same contest as a "2 Year Term" and no more,
                # so only two ballots that differ on it are two contests for certain
                differ = r["term"]["unexpired"] != inst["term"]["unexpired"]
                if differ and r["listed"] == "ballot" and inst["listed"] == "ballot":
                    continue
                pairs = same_names([n for _p, n in r["cands"]], [n for _p, n in inst["cands"]])
                same_seat = (r["district"], r["seat"]) == (inst["district"], inst["seat"])
                # the same people on two counties' lists for one board are one contest, however each list words the seat;
                # a contest with no candidate is matched by its seat alone
                if pairs is not None and (r["cands"] or (same_seat and not differ)):
                    target = r
                    break
            if target is None:
                merged.append(dict(inst, counties=[inst["county"]], also=[]))
                continue
            target["counties"].append(inst["county"])
            target["also"].append({"county": inst["county"], "seats": inst["seats"],
                                   "respelled": sum(1 for a, b in pairs if re.sub(r"[^A-Z ]", "", a.upper()).split() != re.sub(r"[^A-Z ]", "", b.upper()).split())})
            if not (target["district"] or target["seat"]):
                target["district"], target["seat"] = inst["district"], inst["seat"]
            if inst["term"]["unexpired"] and not target["term"]["unexpired"]:
                target["unexpired_said_by"] = (inst["county"], inst["listed"])    # that list says so; the one the race is read from does not
            if target["term"]["years"] is None:
                target["term"] = dict(inst["term"], unexpired=inst["term"]["unexpired"] or target["term"]["unexpired"])
            elif inst["term"]["unexpired"] and not target["term"]["unexpired"]:
                target["term"] = dict(target["term"], unexpired=True)
            if target["seats"] is None:
                target["seats"] = inst["seats"]                       # a ballot says how many are elected; not every roster does
            if not target.get("role_given") and inst.get("role_given"):
                target["office"], target["role_given"] = inst["office"], True      # the list that says what the board's members are called
            counts["second_printing"] += len(inst["cands"])
            counts["contests_on_two_lists"] += 1

        for r in merged:
            parts = [f"2026-{STATE}", jid[len(STATE) + 1:] if jid.startswith(STATE + "-") else jid, r["office_kind"].replace("_", "-")]
            parts += [lslug(x) for x in (r["district"], r["seat"]) if x]
            y, unexpired = r["term"]["years"], r["term"]["unexpired"]
            if not unexpired and y not in (None, 4):
                parts.append(f"{y}yr")
            rid = "-".join(parts) + ("-S" if unexpired else "")
            if rid in seen_ids:
                lfail(f"two different contests share the race id {rid}; stopped")
            seen_ids.add(rid)
            roster = r["listed"] != "ballot"
            word = "ballot" if not roster else "roster" if r["listed"] == "candidate roster" else "list"
            notes = []
            if r["seats"] and r["seats"] > 1:
                notes.append(f"Voters choose {r['seats']}.")
            if unexpired:
                yrs = f"{NUMBER_NAMES[y]}-year " if y in NUMBER_NAMES else ""
                if r.get("unexpired_said_by"):
                    by, how = r["unexpired_said_by"]
                    notes.append(f"For the rest of a term: {short[by]} County's {'ballots call' if how == 'ballot' else how + ' calls'} it an unexpired "
                                 f"{yrs}term; the {word} this race is read from "
                                 + (f"prints a {yrs}term and no more." if yrs else "gives no term."))
                else:
                    notes.append(f"For the rest of a term: the {word} calls it an unexpired {yrs}term.")
            elif y not in (None, 4):
                notes.append(f"A {NUMBER_NAMES.get(y, y)}-year term, as the {word} prints it.")
            if r["term"]["undetermined"]:
                notes.append("The ballot says the term is to be determined.")
            if roster and r["seats"] is None and r["office_kind"] not in ONE_HOLDER:
                notes.append(f"From the county clerk's {r['listed']}, which does not say how many are to be elected.")
            if r.get("said"):
                notes.append(f"The clerk's list writes the district \"{r['said']}\".")
            if not r["cands"]:
                notes.append("The ballot prints this contest with no candidate named; a write-in line is the only choice." if not roster else
                             f"The county clerk's {r['listed']} lists this contest with no candidate.")
            if len(r["counties"]) > 1:
                notes.append(f"On the lists of {and_names(short[g] for g in sorted(r['counties']))} counties: "
                             + (f"the {jname.rsplit(' ', 1)[-1]} lies in more than one county." if first["what"] == "city" else
                                "the district crosses the county line."))
                for a in r["also"]:
                    if a["seats"] and r["seats"] and a["seats"] != r["seats"]:
                        notes.append(f"{short[r['counties'][0]]} County's ballots say to vote for {NUMBER_NAMES.get(r['seats'], r['seats'])}; "
                                     f"{short[a['county']]} County's say {NUMBER_NAMES.get(a['seats'], a['seats'])}.")
                        problems.append(f"{rid}: two counties' ballots give different vote-for numbers")
                    if a["respelled"]:
                        notes.append(f"{short[a['county']]} County's ballots print {NUMBER_NAMES.get(a['respelled'], a['respelled'])} of the names another way.")
            if home and home not in r["counties"] and first["what"] != "city":
                notes.append(f"Read from {and_names(short[g] for g in sorted(r['counties']))} County's ballots, which carry this district's contests for the "
                             f"precincts inside it; {full[home]}'s own list is not loaded yet.")
            if (first["what"] in ("special", "college") and nameless and not re.search(r"\d", best)) or (first["what"] == "school" and not jkey[2]):
                notes.append(f"The {word}'s title gives the district no name of its own.")
            if r.get("by_label"):
                notes.append(f"The ballot's title does not name the city or town; every precinct ballot that carries this contest is labelled "
                             f"{r['by_label']}.")
            if spelled:
                notes.append(f"The ballot spells the district's name \"{spelled}\".")
            if r["caps"] and r["cands"]:
                notes.append(CAPITALS if not roster else "The clerk's list prints names in capitals; they are shown here in ordinary capitals.")
            races.append((rid, STATE, r["level"], r["office_kind"], r["office"], jname, jid, json.dumps(cids), r["district"], r["seat"],
                          1 if unexpired else 0, 1 if r["partisan"] else 0, None, None, None, GENERAL, " ".join(notes) or None))
            names_here = set()
            for party, nm in r["cands"]:
                shown = ordinary(nm) if r["caps"] else nm
                if shown in names_here:
                    lfail(f"a name is printed twice in {rid}; stopped")
                names_here.add(shown)
                cands.append((rid, "general", GENERAL, shown, party if r["partisan"] else NONPARTISAN, party_code(party) if r["partisan"] else "N",
                              None, 0, 0, None, None, None, None, r["src"], name_notes.get((r["src"], nm))))
            if r["wrapped"]:
                problems.append(f"{rid}: a name that runs on to a second line on the ballot; read it against the ballot")

    # ---- counts: every name read for a local contest is stored once, was a second printing of a stored contest, or is in a gap
    names_local = sum(len(i["cands"]) for i in instances)
    if names_local != len(cands) + counts["second_printing"]:
        lfail(f"{names_local} names read for local contests, {len(cands)} stored and {counts['second_printing']} on a second county's list; stopped")
    if len({(c[0], c[3]) for c in cands}) != len(cands):
        lfail("a candidate is stored twice in one local race; stopped")

    by_level, by_kind = Counter(r[2] for r in races), Counter(r[3] for r in races)
    reached = sorted({g for r in races for g in json.loads(r[7])})
    n_empty = sum(1 for r in races if not any(c[0] == r[0] for c in cands))
    ballot_counties = [short[g] for g, v in loaded.items() if v["how"] == "sample ballots"]
    roster_counties = [short[g] for g, v in loaded.items() if v["how"] == "candidate rosters"]
    part_counties = [short[g] for g, v in loaded.items() if v["how"] == "list of filings"]
    not_loaded = [short[g[2]] for g in gaps if g[1] == "county" and g[2] not in loaded]
    questions = sum(v["questions"] for v in loaded.values())
    read_from = f"the county clerks' November sample ballots ({and_names(ballot_counties) or 'no county'})"
    if roster_counties:
        read_from += f"{' and' if not part_counties else ','} the clerks' candidate rosters ({and_names(roster_counties)})"
    if part_counties:
        read_from += f" and, for school, college and special district offices only, the clerk's list of filings ({and_names(part_counties)})"

    notes = [
        (STATE, "local_calendar",
         "On November 3, 2026 every Wyoming county elects its clerk, treasurer, assessor, coroner, sheriff, clerk of district court and county attorney "
         "or district attorney for four years, and the county commissioners whose terms are up, all with parties on the ballot; school district and "
         "community college trustees are elected the same day without parties. Cities and towns elect mayors and council members that day too, unless "
         "a town has chosen by charter ordinance to vote in May, and each special district (conservation, fire, hospital, cemetery and others) elects "
         "its directors in March, May or November as the district has decided, so not every district is on this ballot. Circuit and district judges "
         "and Supreme Court justices do not run against anyone: voters are asked whether to keep them.",
         "W.S. 22-2-105 (county offices and judges), 22-22-102 (school and college trustees), 22-23-202 (towns that vote in May), 22-29-112 (special "
         "districts), 22-6-117 and 22-6-125 (which offices carry parties)", TITLE22_URL),
        (STATE, "local_coverage",
         f"Wyoming has no statewide list of local candidates: each county clerk publishes its own. Loaded here: {len(races)} county, city, school and "
         f"special district contests with {len(cands)} candidates in {len(loaded)} of the 23 counties, read from {read_from}; "
         f"{n_empty} of the contests have no candidate on their list. The {len(j_races)} judges standing for retention come from the Secretary of "
         f"State's own file. Left out: ballot questions and constitutional amendments, and the {len(not_loaded)} counties named among the gaps, whose "
         "clerks had posted their sample ballots as scanned pictures, or no November sample ballot and only a list this loader does not read yet"
         + (f"; in {and_names(part_counties)} County the county's own offices and the cities' are among the gaps too. " if part_counties else ". ")
         + "A school or special district that crosses a county line is "
         "shown once, with every county whose ballots carry it. Names rotate from precinct to precinct, so no ballot order is given; a candidate who "
         "withdrew is simply not printed on a ballot and cannot be counted there.",
         "County clerks' 2026 general election sample ballots and rosters; Wyoming Secretary of State, 2026 General Election Candidates; W.S. 22-6-122 "
         "(names rotated)", CLERKS_URL),
    ]

    sources += [
        (SRC_CSV, STATE, "official candidate list", "Wyoming Secretary of State, Elections Division",
         "2026 General Election Candidates (data file): the judges standing for retention", kept["url"], "", kept["fetched"], kept["sha256"], len(kept["rows"]),
         f"A file of {kept['columns_in_file']} columns and {len(kept['rows'])} rows. Five columns are read, by their headings: {', '.join(kept['columns'])}. "
         "The mailing address, city, state and ZIP, telephone, e-mail and web address columns are never read, and only the five columns are kept on "
         f"disk; the fingerprint is of the file as it came. {len(j_races) + j_gone} rows are judges standing for retention ({j_gone} withdrawn, left "
         f"off); the other {csv_others} rows are the federal, state and legislative candidates the roster carries, used here only to check the names "
         "on the county ballots. The file gives no ballot order."),
        (SRC_TITLE5, STATE, "statute", "Wyoming Legislature", "Wyoming Statutes, Title 5 (Courts): W.S. 5-3-101, the judicial districts", TITLE5_URL, "",
         mdate(t5path), sha_of(t5path), len(districts),
         "Read for the counties of each of the nine judicial districts, which are the counties that vote on a district judge and, the circuit court "
         "having the same lines, on a circuit judge. A law text with no personal details; kept whole."),
        (SRC_PLACES, STATE, "official place codes", "U.S. Census Bureau", "2020 place codes, Wyoming (st56_wy_place2020.txt)", PLACE_URL, "", mdate(ppath),
         sha_of(ppath), p_rows, "Names, codes and counties of Wyoming's incorporated cities and towns, to which a ballot's mayor and council contests are "
         "matched by name. No personal details."),
    ]

    # ---- the last look before anything is written: nothing that reads like contact details, by the trial check's own test and the page builder's
    all_races, all_cands, all_places = j_races + races, j_cands + cands, j_places + list(place_rows.values())
    for table, strict, items in (("sl_races", True, [(r[0], (r[4], r[5], r[8], r[9], r[16], r[13])) for r in all_races]),
                                 ("sl_candidates", True, [(c[0], (c[3], c[4], c[14])) for c in all_cands]),
                                 ("sl_places", True, [(p[1], (p[2],)) for p in all_places]),
                                 ("sl_gaps", False, [(g[2], (g[3], g[4], g[5])) for g in gaps]),
                                 ("sl_notes", False, [(x[1], (x[2], x[3])) for x in notes]),
                                 ("sl_sources", False, [(s[0], (s[3], s[4], s[10])) for s in sources])):
        for ident, texts in items:
            if any(t and (contact_like(t, strict) or (strict and BUILDER_STREET.search(str(t)))) for t in texts):
                lfail(f"a text for {table} ({ident}) reads like contact details; stopped (the text is not printed)")

    return {"races": all_races, "cands": all_cands, "places": all_places, "sources": sources, "gaps": gaps, "notes": notes, "problems": problems,
            "counts": counts, "loaded": loaded, "by_level": by_level, "by_kind": by_kind, "reached": reached, "empty": n_empty, "judges": len(j_races),
            "judges_withdrawn": j_gone, "local_races": len(races), "local_cands": len(cands), "names_local": names_local, "questions": questions,
            "not_loaded": not_loaded, "in_part": part_counties, "csv_rows": len(kept["rows"]), "csv_others": csv_others, "csv_fetched": kept["fetched"],
            "short": short}


# ---------- the load ----------

def race_of_roster(office):
    """(race_id, info) for a state office as the roster writes it."""
    if office in STATEWIDE:
        key, kind, shown, roster_office, sheet = STATEWIDE[office]
        return f"2026-{STATE}-{key}", {"level": "statewide", "office_kind": kind, "office": shown, "district": None, "chamber": None,
                                       "roster": roster_office, "results": sheet}
    m = LEG_ROSTER.fullmatch(office or "")
    if m:
        chamber = "Senate" if m.group(1) == "SENATOR" else "House"
        key, kind, shown = CHAMBER[chamber]
        d = str(int(m.group(2)))
        return f"2026-{STATE}-{key}{d}", {"level": "legislature", "office_kind": kind, "office": shown, "district": d, "chamber": chamber,
                                          "roster": None, "results": f"{chamber} District {d}"}
    fail(f"an office the loader does not know on the roster: {office!r}")


def race_of_results(office):
    """race_id for an office heading in the results, or None for the federal offices (the federal loader's)."""
    if office.startswith("United States "):
        return None
    for key, kind, shown, roster_office, sheet in STATEWIDE.values():
        if office == sheet:
            return f"2026-{STATE}-{key}"
    m = LEG_RESULTS.fullmatch(office)
    if m:
        return f"2026-{STATE}-{CHAMBER[m.group(1)][0]}{int(m.group(2))}"
    fail(f"an office heading the loader does not know in the results: {office!r}")


def load(db_path, say=print, cache=CACHE, local_cache=None, refresh=False):
    """Wyoming's rows into the database at db_path: the state races, then the judges standing for retention and the
    county and local races. local_cache is the folder for the county-level files (ballot_cache/wy/local/ unless told
    otherwise); refresh=True asks the county pages, the candidates file and the statute again."""
    net.patient_lookups()
    seats, offices, as_of = roster()
    cmap = counties()
    checks = []

    # ---- the November roster
    rrows, printed, r_fetched, r_sha, r_how, r_all = read_roster(cache, say)
    races = OrderedDict()
    for office, party, name, wd in rrows:
        rid, info = race_of_roster(office)
        races.setdefault(rid, info)

    # ---- the primary: the summary, checked against the precinct files
    zpath = os.path.join(cache, "wy_2026_primary_results.zip")
    if not os.path.exists(zpath):
        net.download(RESULTS_URL, zpath, max_age_days=30, say=say)
    with zipfile.ZipFile(zpath) as z:
        names_in = z.namelist()
        for need in (BOOK, PRECINCT_BOOK, BALLOTS_BOOK):
            if need not in names_in:
                fail(f"{need} is not in the primary results zip")
        wb_sum = openpyxl.load_workbook(io.BytesIO(z.read(BOOK)), read_only=True, data_only=True)
        wb_pre = openpyxl.load_workbook(io.BytesIO(z.read(PRECINCT_BOOK)), read_only=True, data_only=True)
        wb_bal = openpyxl.load_workbook(io.BytesIO(z.read(BALLOTS_BOOK)), read_only=True, data_only=True)
    sheets, not_read = summary_blocks(wb_sum)
    state_ballots = statewide_ballots(wb_sum)
    raw_cells, pre_problems = precinct_cells(wb_pre)
    bal, bal_problems = ballots_cast(wb_bal)
    pcells, settled, unsettled = settle_parties(raw_cells, sheets)
    checks += pre_problems + bal_problems + unsettled
    county_names = {ws.title for ws in wb_pre.worksheets}
    unknown_c = sorted(c for c in county_names if fold(c) not in cmap)
    if unknown_c or len(county_names) != 23:
        fail(f"the precinct results' county sheets do not match the Census file's 23 counties ({unknown_c})")
    missing_p = sorted({f"{c} {p}" for c, p, _o, _pa, _l, _v in pcells if p not in bal.get(c, {})})
    if missing_p:
        checks.append(f"precincts in the results not in the ballots-cast file: {', '.join(missing_p[:12])}")
    pre = defaultdict(lambda: {"votes": defaultdict(Counter), "precincts": set(), "counties": set()})
    for c, p, office, party, lab, v in pcells:
        b = pre[(office, party)]
        b["votes"][lab][c] += v
        b["precincts"].add((c, p))
        b["counties"].add(c)

    fields_src = {}                                          # (race, party) -> summary block confirmed by the precincts
    set_aside = defaultdict(list)                            # sheet -> blocks the precinct files do not confirm
    spellings = []
    for sheet, blocks in sheets.items():
        for (office, party), b in blocks.items():
            rid = race_of_results(office)
            if rid is None:
                continue
            p = pre.get((office, party))
            m = label_map(list(p["votes"]), b["labels"]) if p else None
            sums = Counter()
            for lab, by_county in (p["votes"].items() if m else []):
                sums[m[lab]] += sum(by_county.values())
            if not m or set(sums) != set(b["votes"]) or any(sums[k] != v for k, v in b["votes"].items()):
                set_aside[sheet].append(f"{office} {party}")
                continue
            if (rid, party) in fields_src:
                fail(f"{office} {party} is confirmed twice in the summary")
            for lab, s_lab in m.items():
                if lab != s_lab:
                    spellings.append(f"{office} {party}: \"{lab}\" in {', '.join(sorted(p['votes'][lab]))} County's sheet for "
                                     f"\"{s_lab}\"")
            fields_src[(rid, party)] = dict(b, sheet=sheet, office=office, precincts=p["precincts"], counties=p["counties"])
    wanted = {(race_of_results(o), pa) for (o, pa) in pre if race_of_results(o)}
    unconfirmed = sorted(f"{r} {pa}" for r, pa in wanted - set(fields_src))
    if unconfirmed:
        checks.append(f"precinct-file blocks with no matching summary block: {', '.join(unconfirmed)}")

    # precinct by precinct, each chamber's votes (with write-ins, overvotes and undervotes) against the party's ballots cast:
    # every voter has one House race and every statewide race; the Senate is only half the state this year
    per = defaultdict(int)
    house_of = defaultdict(set)
    for c, p, office, party, _lab, v in pcells:
        if office.startswith("United States "):
            continue
        kind = office.split(" District")[0] if " District " in office else office
        per[(c, p, kind, party)] += v
        if kind == "House":
            house_of[(c, p)].add(office)
    split = {k for k, v in house_of.items() if len(v) > 1}
    unreconciled, senate_short, senate_short_split = [], 0, 0
    for (c, p, kind, party), v in sorted(per.items()):
        cast = bal.get(c, {}).get(p, {}).get(party)
        if cast is None:
            continue
        if kind == "Senate":
            if v > cast:
                unreconciled.append(f"{c} {p} Senate {party} ({v:,} against {cast:,} ballots)")
            elif v < cast:
                senate_short += 1
                senate_short_split += (c, p) in split
        elif v != cast:
            unreconciled.append(f"{c} {p} {kind} {party} ({v:,} against {cast:,} ballots)")
    if senate_short != senate_short_split:
        checks.append(f"{senate_short - senate_short_split} precinct Senate totals short of the ballots cast in precincts the "
                      "results do not split between House districts")
    for (rid, party), f in sorted(fields_src.items()):
        if races.get(rid, {}).get("level") == "statewide" and state_ballots.get(party) != sum(f["votes"].values()):
            unreconciled.append(f"{rid} {party} ({sum(f['votes'].values()):,} against the statewide sheet's "
                                f"{state_ballots.get(party)} ballots)")
    if unreconciled:
        checks.append("primary figures that do not add up to the ballots cast: " + "; ".join(unreconciled))
    n_prec = len({(c, p) for c, p, *_ in pcells})

    # which races the primary had, and the roster's seats against them
    primary_races = {rid for rid, _pa in fields_src}
    not_on_roster = sorted(primary_races - set(races))
    if not_on_roster:
        checks.append(f"on the August 18 primary ballot but no candidate on the November roster: {', '.join(not_on_roster)}")
    no_primary = sorted(rid for rid in races if rid not in primary_races)
    if no_primary:
        checks.append(f"on the November roster but not in the August 18 primary results: {', '.join(no_primary)}")
    house = sorted(int(i["district"]) for i in races.values() if i["chamber"] == "House")
    if house != list(range(1, 63)):
        checks.append(f"State House districts on the roster are not 1 to 62: missing {sorted(set(range(1, 63)) - set(house))}")
    sen = sorted(int(i["district"]) for i in races.values() if i["chamber"] == "Senate")
    parity = Counter(d % 2 for d in sen).most_common(1)[0][0] if sen else 1
    regular = [d for d in sen if d % 2 == parity]
    if regular != list(range(2 - parity, 32, 2)):
        checks.append(f"the regular State Senate seats on the roster are not every {'odd' if parity else 'even'} district: {regular}")
    early = [d for d in sen if d % 2 != parity]
    for (ch, d), hs in seats.items():
        if len(hs) != 1:
            checks.append(f"the roster lists {len(hs)} sitting members for {ch} District {d}")

    # ---- holders and notes
    holders, notes, special, county_ids = {}, {}, {}, {}
    for rid, i in races.items():
        if i["level"] == "legislature":
            hs = seats.get((i["chamber"], i["district"]), [])
            holders[rid] = hs
            if not hs:
                notes[rid] = f"The Open States roster ({as_of}) lists no sitting member for this seat."
            got = set()
            for party in ("Republican", "Democratic"):
                got |= fields_src.get((rid, party), {}).get("counties", set())
            county_ids[rid] = ",".join(sorted(cmap[fold(c)][0] for c in got)) or None
            if not got:
                checks.append(f"{rid}: no precinct rows in the primary results, so no counties")
        else:
            h = offices.get(i["roster"]) if i["roster"] else None
            holders[rid] = [h] if h else []
            if not h:
                notes[rid] = "The Open States roster this site uses does not carry this office, so today's holder is not shown."
    for d in early:
        rid = f"2026-{STATE}-SS{d}"
        special[rid] = 1
        hs = holders.get(rid, [])
        since = f" The Open States roster shows today's member serving since {long_date(hs[0]['start'])}." if hs and hs[0]["start"] else ""
        notes[rid] = ((notes[rid] + " ") if rid in notes else "") + (
            f"Senate District {d} is the one {'even' if parity else 'odd'}-numbered Senate district on this year's roster; "
            f"the other Senate seats up this year are {'odd' if parity else 'even'}-numbered, and Wyoming senators serve four "
            f"years.{since} It is treated here as an election for the rest of the term; the roster itself does not say so in words.")
        checks.append(f"{rid} marked special (off-cycle seat on the roster){since}")

    # other-chamber districts that share a precinct with each legislative race, from the primary results
    prec_of = defaultdict(set)
    for (rid, _party), f in fields_src.items():
        prec_of[rid] |= f["precincts"]
    everyone = [p for ps in seats.values() for p in ps] + list(offices.values())
    # every name each race carries this year (the roster and the primary), and each roster person's own race on the ballot
    names_in = defaultdict(set)
    for office, _party, name, _wd in rrows:
        names_in[race_of_roster(office)[0]].add(name)
    for (rid, _party), f in fields_src.items():
        names_in[rid] |= {lab for lab in f["labels"] if lab not in TALLY and lab != WITHDRAWN_LABEL}
        if f["withdrawn"]:
            names_in[rid].add(f["withdrawn"][0])
    own_race = {}
    for rid, i in races.items():
        for h in holders[rid]:
            own_race[h["id"]] = rid

    def running_at_home(p, rid):
        """True when the person's own seat or office is on this year's ballot with a candidate of that name: then a candidate
        of the same name in another race is taken to be someone else."""
        home = own_race.get(p["id"])
        return bool(home and home != rid and any(person_fits(n, p) for n in names_in[home]))

    def sitting(rid, names):
        i, out = races[rid], {}
        hs = holders[rid]
        pairs = [(n, h) for n in names for h in hs if person_fits(n, h)]
        for n, h in pairs:
            if sum(1 for a, _ in pairs if a == n) == 1 and sum(1 for _, b in pairs if b["id"] == h["id"]) == 1:
                out[n] = (h["id"], 1, None)
        for n in names:
            if n in out:
                continue
            if i["level"] == "legislature":
                other = "House" if i["chamber"] == "Senate" else "Senate"
                touching = {r2 for r2 in prec_of if r2 != rid and races.get(r2, {}).get("chamber") == other and prec_of[r2] & prec_of[rid]}
                pool = [p for r2 in touching for p in seats.get((other, races[r2]["district"]), [])]
            else:
                pool = [p for p in everyone if not any(p["id"] == h["id"] for h in hs)]
            got = [p for p in pool if person_fits(n, p)]
            if len(got) == 1 and not running_at_home(got[0], rid):
                p = got[0]
                out[n] = (p["id"], 0, f"Serves today in {CHAMBER_WORDS[p['chamber']]}, District {p['district']}." if p["chamber"]
                          else f"Serves today as {p['label']}.")
        return out

    cand = []

    # ---- the November ballot
    on_ballot, shown, withdrawn = defaultdict(list), defaultdict(list), []
    for office, party, name, wd in rrows:
        rid, _ = race_of_roster(office)
        shown[(rid, party)].append(name)
        if wd:
            withdrawn.append(rid)
            continue
        on_ballot[rid].append((party, name))
    for rid, rows in on_ballot.items():
        fit = sitting(rid, [n for _p, n in rows])
        for party, name in rows:
            mid, inc, note = fit.get(name, (None, 0, None))
            n = [note] if note else []
            f = fields_src.get((rid, party))
            if party in FIELD_CODE and not any(same_person(name, x) for x in (f["labels"] if f else []) + (
                    [f["withdrawn"][0]] if f and f["withdrawn"] else [])):
                n.append(f"Not printed on the August 18 {party} primary ballot for this seat; the roster does not say how the "
                         "nomination was made.")
            if inc and holders[rid]:
                h = next(h for h in holders[rid] if h["id"] == mid)
                if h["party"] and h["party"] != party:
                    checks.append(f"{rid}: sitting member listed as {h['party']}, on the roster as {party}")
            cand.append((rid, "general", GENERAL, name, party, party_code(party), None, inc, 0, None, None, None, mid, SRC_GENERAL,
                         " ".join(n) or None))

    # ---- the August 18 primary fields
    fields = Counter()
    upset, open_ = [], []
    for (rid, party), f in sorted(fields_src.items()):
        if rid not in races:
            continue
        people = [lab for lab in f["labels"] if lab not in TALLY]
        if len(people) < 2:
            continue
        if party not in FIELD_CODE:
            fail(f"a primary field for a party the loader does not know: {party}")
        entries = []
        for lab in people:
            if lab == WITHDRAWN_LABEL:
                if not f["withdrawn"]:
                    checks.append(f"{rid} {party}: a withdrawn candidate's votes with no name in the footnote; left out of the field")
                    continue
                who, where = f["withdrawn"]
                note = (f"Withdrew after {where} County's primary ballots had been printed; the official summary counts the votes cast "
                        f"for this candidate in {where} County." if where else
                        "Withdrew after the primary ballots had been printed; the official summary still counts the votes cast for this "
                        "candidate.")
                entries.append((who, f["votes"][lab], note))
            else:
                entries.append((lab, f["votes"][lab], None))
        counted = sum(v for _n, v, _x in entries) + f["votes"].get("Write-Ins", 0)
        noms = [x for x in shown.get((rid, party), [])]
        won = [e for e in entries if any(same_person(e[0], x) for x in noms)]
        top = sorted(entries, key=lambda e: -e[1])
        tie = len(top) > 1 and top[0][1] == top[1][1]
        if len(won) == 1:
            winner = won[0][0]
            if tie or top[0][0] != winner:
                upset.append(f"{rid} {party}")
        elif not won and not tie:
            winner = top[0][0]
            open_.append(f"{rid} {party} (the most votes: {winner}; not on the November roster)")
        else:
            winner = None
            open_.append(f"{rid} {party} ({len(won)} names fit the November roster)")
        fields[races[rid]["office_kind"] if races[rid]["level"] == "legislature" else "statewide"] += 1
        fit = sitting(rid, [e[0] for e in entries])
        for who, votes, note in entries:
            mid, inc, snote = fit.get(who, (None, 0, None))
            extra = []
            if winner and who == winner and not won:
                extra.append("Won the most votes in the primary; not on the November roster (the roster does not say why).")
            cand.append((rid, f"primary-{FIELD_CODE[party]}", PRIMARY, who, party, party_code(party), None, inc, 0, votes,
                         round(100 * votes / counted, 1) if counted else None,
                         None if winner is None else ("advanced" if who == winner else "lost"), mid, SRC_PRIMARY,
                         " ".join(x for x in [snote, note] + extra if x) or None))

    # ---- checks on the whole
    keys = Counter((c[0], c[1], c[3]) for c in cand)
    dup = [k for k, v in keys.items() if v > 1]
    if dup:
        fail(f"the same name twice in one election: {dup}")
    gen_by_race = Counter(c[0] for c in cand if c[1] == "general")
    if gen_by_race.total() + len(withdrawn) != len(rrows):
        fail("November rows written plus withdrawn do not equal the roster's state rows")
    empty = sorted(rid for rid in races if gen_by_race[rid] == 0)

    race_rows = []
    for rid, i in races.items():
        hs = holders[rid]
        jur = f"{i['chamber']} District {i['district']}" if i["level"] == "legislature" else NAME
        race_rows.append((rid, STATE, i["level"], i["office_kind"], i["office"], jur, FIPS, county_ids.get(rid), i["district"], None,
                          special.get(rid, 0), 1, ",".join(h["id"] for h in hs) or None, " and ".join(h["full"] for h in hs) or None,
                          ", ".join(dict.fromkeys(h["party"] for h in hs if h["party"])) or None, GENERAL, notes.get(rid)))
    place_rows = [("county", geoid, full, json.dumps([geoid]), SRC_COUNTIES) for geoid, full in sorted(cmap.values())]

    # ---- the county and local level, and the judges standing for retention (nothing above this line is changed by it)
    local = local_level(local_cache or os.path.join(cache, LOCAL_SUB), cmap, say, refresh)
    if local["csv_others"] != r_all:
        checks.append(f"the Secretary of State's candidates file (fetched {local['csv_fetched']}) has {local['csv_others']} federal, state and "
                      f"legislative rows; the roster read for the state races has {r_all}")

    n_field_rows = sum(1 for c in cand if c[1] != "general")
    aside = "; ".join(f"{s}: {len(v)} blocks" for s, v in set_aside.items())
    src = [
        (SRC_GENERAL, STATE, "official candidate list", "Wyoming Secretary of State, Elections Division",
         "2026 General Election Candidate Roster" + (f" (printed {printed})" if printed else "") + ": state offices and the Legislature",
         LIST_URL, iso(printed), r_fetched, r_sha, len(rrows),
         f"Every office on the roster ({r_all} entries, read from {r_how}); the {len(rrows)} for state offices and the Legislature kept "
         "(the two federal races are the federal loader's). Office, party and name only, and a date under Date Withdrawn; the "
         "mailing address, city, state and ZIP, telephone and e-mail columns are never read. The kept rows, those columns only, are "
         "ballot_cache/wy/wy_2026_state_roster.json; the fingerprint is of the PDF. The roster gives no ballot order (county clerks "
         f"print the ballots), so none is stored. Withdrawn, left off the November ballot: {len(withdrawn)}."),
        (SRC_PRIMARY, STATE, "official results", "Wyoming Secretary of State, Elections Division",
         "2026 Official Primary Election Results (August 18, 2026): Statewide Candidates, Statewide Senate and Statewide House Official "
         "Summaries, from the zip file of data files (\"2026 Primary Results Summaries - OFFICIAL.xlsx\")", RESULTS_URL, "", mdate(zpath),
         sha_of(zpath), n_field_rows,
         "Votes from each office's Total row; a party primary is a field when two or more names were printed on its ballot. Write-ins "
         "count in the share but are not listed; over- and undervotes are left out. A candidate who withdrew after the ballots were "
         "printed is kept with the votes the summary counts. Every block kept equals the county precinct-by-precinct results, "
         "candidate by candidate. Set aside, because the precinct results do not add up to them: "
         + (aside or "nothing") + ("; the sheet \"Statewide Senate Even\" names districts 2 to 30, even, none of which but District 6 is "
         "on the 2026 roster, none of which is in the precinct results, and its District 6 block differs from the one the precinct "
         "results confirm in \"Statewide Senate Odd\"" if "Statewide Senate Even" in set_aside else "") + ". Not read: "
         + ", ".join(not_read) + ". "
         + ("Every figure adds up to the party ballots cast, precinct by precinct." if not unreconciled
            else "Did not add up: " + "; ".join(unreconciled) + ".")),
        (SRC_PRECINCTS, STATE, "official results (used as a check)", "Wyoming Secretary of State, Elections Division",
         "2026 Primary County PbP Results - OFFICIAL.xlsx and 2026 Primary County PbP Total Ballots Cast - OFFICIAL.xlsx, from the same zip",
         RESULTS_URL, "", mdate(zpath), sha_of(zpath), sum(len(v) for v in bal.values()),
         f"Every county's precinct rows ({n_prec} precincts) summed and checked against the county's Total row, and each summary "
         "block confirmed from them. Precinct by precinct, the votes, write-ins, overvotes and undervotes in each statewide race and "
         "in the House races together equal that party's ballots cast; in the Senate, only half the state votes this year, and "
         f"{senate_short} precinct totals fall short of the ballots cast, {'all' if senate_short == senate_short_split else senate_short_split} "
         "of them in precincts the results split between House districts. A district's counties (county_ids) are those where its "
         "race was on the ballot. Precinct numbers the spreadsheet stored as dates are read as month-day. "
         + (" ".join(s_ + "." for s_ in settled) + " " if settled else "")
         + (("Another spelling of a candidate in one county's sheet, counted with the summary's: " + "; ".join(spellings) + ". ")
            if spellings else "")
         + (f"Problems: {len(pre_problems) + len(bal_problems) + len(unsettled)}." if pre_problems or bal_problems or unsettled
            else "No problems.")),
        (SRC_ROSTER, STATE, "roster", "Open States people project (CC0)",
         "Wyoming legislators and statewide officials, as loaded into state_wy.sqlite", "https://github.com/openstates/people",
         as_of, as_of, "", sum(len(v) for v in seats.values()) + len(offices),
         "Today's holder of each seat and office (names, parties, districts, ids and term starts only). The roster does not carry the "
         "Auditor, the Treasurer or the Superintendent of Public Instruction."),
        (SRC_COUNTIES, STATE, "official boundaries", "U.S. Census Bureau",
         "Cartographic boundary file, counties, 2024, 1:500,000 (cb_2024_us_county_500k)", COUNTY_URL, "", mdate(COUNTY_ZIP),
         sha_of(COUNTY_ZIP), len(place_rows), "Wyoming's 23 counties: names and GEOIDs only."),
    ]

    # Wyoming's rows only, in one transaction: its races and candidates by state and race id, its sources, gaps and notes
    # by state, its places by their wy- source ids and its counties by their codes
    con = sqlite3.connect(db_path)
    try:
        con.executescript(SCHEMA)
        con.executescript(EXTRA_SCHEMA)
        with con:
            con.execute("DELETE FROM sl_candidates WHERE race_id IN (SELECT race_id FROM sl_races WHERE state = ?) OR race_id LIKE ?",
                        (STATE, f"2026-{STATE}-%"))
            con.execute("DELETE FROM sl_races WHERE state = ?", (STATE,))
            con.execute("DELETE FROM sl_places WHERE source_id LIKE 'wy-%' OR (kind = 'county' AND id GLOB '56[0-9][0-9][0-9]')")
            con.execute("DELETE FROM sl_sources WHERE state = ?", (STATE,))
            con.execute("DELETE FROM sl_gaps WHERE state = ?", (STATE,))
            con.execute("DELETE FROM sl_notes WHERE state = ?", (STATE,))
            con.executemany("INSERT INTO sl_races VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", race_rows + local["races"])
            con.executemany("INSERT INTO sl_candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", cand + local["cands"])
            con.executemany("INSERT INTO sl_places VALUES (?,?,?,?,?)", place_rows + local["places"])
            con.executemany("INSERT INTO sl_sources VALUES (?,?,?,?,?,?,?,?,?,?,?)", src + local["sources"])
            con.executemany("INSERT INTO sl_gaps VALUES (?,?,?,?,?,?,?)", local["gaps"])
            con.executemany("INSERT INTO sl_notes VALUES (?,?,?,?,?)", local["notes"])
    finally:
        con.close()

    # ---- the report: counts only
    kind_of = lambda rid: races[rid]["office_kind"] if races[rid]["level"] == "legislature" else "statewide"
    by = Counter(kind_of(rid) for rid in races)
    gen = Counter(kind_of(c[0]) for c in cand if c[1] == "general")
    alone = Counter(kind_of(rid) for rid, v in gen_by_race.items() if v == 1)
    inc = Counter(kind_of(c[0]) for c in cand if c[1] == "general" and c[7])
    say(f"    Wyoming (state races): {by['state_senate']} Senate seats ({len(early)} off-cycle), {by['state_house']} House seats, "
        f"{by['statewide']} statewide offices; {gen.total()} candidates on the November ballot (Senate {gen['state_senate']}, House "
        f"{gen['state_house']}, statewide {gen['statewide']}; {len(withdrawn)} withdrawn left off; unopposed: Senate "
        f"{alone['state_senate']}, House {alone['state_house']}, statewide {alone['statewide']}); sitting member on the ballot: Senate "
        f"{inc['state_senate']}, House {inc['state_house']}, statewide {inc['statewide']}; primary fields: Senate "
        f"{fields['state_senate']}, House {fields['state_house']}, statewide {fields['statewide']} ({n_field_rows} rows, official "
        f"votes, each confirmed from the precinct results)")
    for label, items in (("no candidate on the November roster", empty),
                         ("summary blocks set aside (the precinct results do not add up to them)",
                          [f"{s} ({len(v)})" for s, v in set_aside.items()]),
                         ("the November nominee is not the primary's top vote-getter", upset), ("primary fields left open", open_)):
        if items:
            say(f"    CHECK Wyoming (state races): {label}: {', '.join(items)}")
    for c in checks:
        say(f"    CHECK Wyoming (state races): {c}")

    # ---- the report on the county and local level: counts only
    lv, lc = local["by_level"], local["counts"]
    how = Counter(v["how"] for v in local["loaded"].values())
    say(f"    Wyoming (local): {local['judges']} judges standing for retention ({local['judges_withdrawn']} withdrawn left off); "
        f"{local['local_races']} county and local races in {len(local['loaded'])} of 23 counties (county {lv.get('county', 0)}, city {lv.get('city', 0)}, "
        f"school {lv.get('school', 0)}, conservation {lv.get('soil_water', 0)}, hospital {lv.get('hospital', 0)}, other districts {lv.get('other', 0)}), "
        f"{local['local_cands']} candidates, {local['empty']} contests with no candidate; {how.get('sample ballots', 0)} counties from sample ballots "
        f"({lc['ballot_files']} files, {lc['ballots']} ballots), {how.get('candidate rosters', 0)} from the clerk's rosters, "
        f"{how.get('list of filings', 0)} in part from the clerk's list of filings ({', '.join(local['in_part']) or 'none'}: school, college and "
        f"special districts only); {local['questions']} ballot questions seen and left out; not loaded: {', '.join(local['not_loaded']) or 'none'}")
    say(f"      check: {local['names_local']} names read for local contests = {local['local_cands']} stored + {lc['second_printing']} listed a second time "
        f"by a neighbouring county ({lc['contests_on_two_lists']} contests on two counties' lists); {lc['names_in_gaps']} more in {lc['unread']} contests "
        f"named in sl_gaps; {lc['names_read']} names read on all ballots, state and federal offices included, and {lc['state_contests_checked']} "
        f"state and federal contests on the county ballots checked against the Secretary of State's file; rosters and lists: {lc['roster_rows']} rows, "
        f"{lc['withdrawn']} withdrawn left off, {lc['roster_judges']} judges (stored from the Secretary of State's file), "
        f"{lc['rows_for_county_offices']} for a county office on a list that does not cover them")
    for p in local["problems"]:
        say(f"    CHECK Wyoming (local): {p}")
    return len(cand) + len(local["cands"])


if __name__ == "__main__":
    args = sys.argv[1:]
    cache, local_cache, refresh = CACHE, None, False
    if "--refresh" in args:
        args.remove("--refresh")
        refresh = True
    for flag in ("--cache", "--local-cache"):
        if flag in args:
            i = args.index(flag)
            cache, local_cache = (args[i + 1], local_cache) if flag == "--cache" else (cache, args[i + 1])
            del args[i:i + 2]
    if len(args) != 1:
        raise SystemExit("usage: python ballot/state_local_wy.py <database file> [--cache <folder>] [--local-cache <folder>] [--refresh]")
    load(args[0], cache=cache, local_cache=local_cache, refresh=refresh)
