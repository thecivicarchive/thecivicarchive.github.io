"""
ballot/ia_place_votes.py - how each Iowa place voted in past partisan general elections, added up from the Iowa
Secretary of State's official precinct results, so a page can show the record of a county, a legislative district or
the precinct a township votes in without anyone labelling a candidate. The Iowa twin of ballot/mn_place_votes.py; the
output has the same shape (the Democratic count is "dem" here, as in Wisconsin's file, where Minnesota's is "dfl").

    python ballot/ia_place_votes.py               reads (or downloads) the files, writes ballot/lean/ia_place_votes.json
    python ballot/ia_place_votes.py --refresh     asks for everything again even when the cached copies are there
    python ballot/ia_place_votes.py --out x.json  writes somewhere else; --cache DIR keeps the downloads somewhere else
    python ballot/ia_place_votes.py --selftest    the arithmetic and the readers on made-up tables; downloads nothing

What it is, and is not
----------------------
For President in 2024, Governor and U.S. Senator in 2022, and President and U.S. Senator in 2020 (Iowa elected no
senator in 2024): the votes for the Democratic ticket, the Republican ticket, everyone else together (write-ins
included) and the total, for every county; for the 2022 and 2024 contests, every state Senate and House district of
the 2021 plan (the lines the 2026 Legislature is elected on); and, for the townships our pages use, the precinct or
precincts the township's voters vote in. It is how the people of a place voted then. It is not a prediction, it says
nothing about any candidate on a later ballot or about any voter, and it turns no nonpartisan office into a partisan
one. No database is opened for writing; ballot_local_2026.sqlite is opened read-only, only to see which townships our
pages use and to count how many of our places are covered.

Where the numbers come from
---------------------------
The Secretary of State's "Election Results & Statistics" page (sos.iowa.gov/election-results-statistics) links, for
each general election, "Precinct Results by County":
  - 2024: the Secretary's results system (electionresults.iowa.gov/IA/122322), one "detailxml" report per county, the
    county's canvassed results by precinct, absentee votes included in each precinct.
  - 2022 and 2020: one workbook per county on sos.iowa.gov (the same report, as a spreadsheet: one sheet per contest,
    one row per precinct, a candidate's Election Day, Absentee and Total Votes).
All 297 files are small (about 100 MB together), fetched one polite request at a time and kept in
states_cache/ia_local/results/. The files name every candidate; this loader reads only the contests named above and
the precinct lists of the contests for State Representative and State Senator, and keeps only the two tickets typed
in CONTESTS, to say which election this was.

  - Legislative districts. An Iowa precinct lies in one House district (Iowa Code 49.3), and the county's file lists,
    under "State Representative District N", exactly the precincts that voted in that contest: that is how a precinct
    is placed in its district, with no map. A Senate district of the 2021 plan is two House districts (district n is
    House districts 2n-1 and 2n); every "State Senator District" contest in the 2022 and 2024 files is checked against
    that rule, and the loader stops if one disagrees. Districts are given for 2022 and 2024 only: 2020 was held on the
    plan of 2011.
  - Townships. Iowa counts no votes by township: a rural precinct is one township, or several, often with a small
    city. So a township's figures are those of the precinct or precincts its voters outside the larger cities vote
    in, and the place's name says exactly that ("The precinct Runnells, where Camp township votes with the city of
    Runnells"; "Douglas township (the precinct Douglas 1)"). Which precinct that is comes from two official maps: the
    Legislative Services Agency's precinct layer (the 2022 precinct plans) and the Census Bureau's township and city
    lines; see place_townships(). The map's name for a precinct and the results' name for it are often spelled
    differently ("074 Allen 1" and "ALLEN"; "34 MOVILLE/ARLINGTON-MOVILLE-WOLF CREEK" and "34 Moville"); see
    link_names(). A township whose precinct cannot be told apart by name in the results is left out, never guessed,
    and a pairing by likeness is accepted only in the ten counties where every one was read (LIKENESS_READ).
  - 2020. The Secretary posts Scott County's 2020 precinct results only as a PDF. Counties are the only kind of
    place given for 2020, so Scott County's 2020 figures are its own row of the State Canvass Summary, and the file
    says so; the other 98 counties are added up from their precincts and compared with the Summary.
  - Not read: the files' registered-voter and turnout figures, and every contest but those named.

The control
-----------
Nothing is written unless all of this holds: in every county file each choice's precinct rows add up to the file's
own total for the choice; counties, House districts and Senate districts each add up to the statewide sum; every
county's sum equals that county's Total row in the State Canvass Summary (the abstract the state board of canvassers
signed; sos.iowa.gov/elections/pdf/<year>/general/canvsummary.pdf, and govcanvsummary.pdf for Governor 2022), for
all 99 counties and every contest; and the statewide sum equals the Summary's statewide Total row. The Summary's own
"Total" column counts under votes and over votes as well; the total here is the votes cast for candidates, and the
file gives both. The figures this loader was checked against on 2026-10-02 are typed below (CHECKED); they are used,
and the file says so, when a document cannot be read again.
"""

import argparse
import collections
import io
import json
import os
import pathlib
import re
import sqlite3
import sys
import time
import zipfile
import xml.etree.ElementTree as ET
from urllib.parse import quote

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from states import net  # noqa: E402
from ballot import mn_place_votes as M  # noqa: E402  the shared small things: caching, fingerprints, the file writer

Stop = M.Stop
METHOD = "1.0"                 # change how anything is added up or matched, and this goes up
CACHE = os.path.join(HERE, "states_cache", "ia_local")
OUT = os.path.join(HERE, "ballot", "lean", "ia_place_votes.json")
DB = os.path.join(HERE, "ballot_local_2026.sqlite")
STATE_FIPS = "19"
FEW = 20
SIDES = ("dem", "rep", "other", "total")

AGENCY = "Iowa Secretary of State"
STATS_PAGE = "https://sos.iowa.gov/election-results-statistics"
RESULTS = {
    2024: {"kind": "clarity", "title": "2024 General Election: Precinct Results by County (the Secretary of State's results system, one report per county)",
           "url": "https://electionresults.iowa.gov/IA/122322/web.345435/#/reporting",
           "index": "https://electionresults.iowa.gov/IA/122322/{version}/json/en/electionsettings.json",
           "version": "https://electionresults.iowa.gov/IA/122322/current_ver.txt",
           "file": "https://electionresults.iowa.gov/IA/{county}/{eid}/{version}/reports/detailxml.zip", "ext": "zip"},
    2022: {"kind": "workbook", "title": "2022 General Election: Precinct Results by County - Excel format",
           "url": "https://sos.iowa.gov/precinct-results-county-2022-general-0",
           "file": "https://sos.iowa.gov/elections/pdf/precinctresults/2022general/{county}.xls", "ext": "xls"},
    2020: {"kind": "workbook", "title": "2020 General Election: Precinct Results by County - Excel format",
           "url": "https://sos.iowa.gov/precinct-results-county-2020-general",
           "file": "https://sos.iowa.gov/elections/pdf/precinctresults/2020general/{county}.xlsx", "ext": "xlsx"},
}

# Iowa's 99 counties in the order of their FIPS codes (19001, 19003, ... 19197: the odd numbers, alphabetically).
COUNTIES = ("Adair Adams Allamakee Appanoose Audubon Benton Black_Hawk Boone Bremer Buchanan Buena_Vista Butler Calhoun Carroll Cass Cedar "
            "Cerro_Gordo Cherokee Chickasaw Clarke Clay Clayton Clinton Crawford Dallas Davis Decatur Delaware Des_Moines Dickinson Dubuque Emmet "
            "Fayette Floyd Franklin Fremont Greene Grundy Guthrie Hamilton Hancock Hardin Harrison Henry Howard Humboldt Ida Iowa Jackson Jasper "
            "Jefferson Johnson Jones Keokuk Kossuth Lee Linn Louisa Lucas Lyon Madison Mahaska Marion Marshall Mills Mitchell Monona Monroe "
            "Montgomery Muscatine O'Brien Osceola Page Palo_Alto Plymouth Pocahontas Polk Pottawattamie Poweshiek Ringgold Sac Scott Shelby Sioux "
            "Story Tama Taylor Union Van_Buren Wapello Warren Washington Wayne Webster Winnebago Winneshiek Woodbury Worth Wright").split()
COUNTIES = [c.replace("_", " ") for c in COUNTIES]
COUNTY_FIPS = {name: f"{STATE_FIPS}{2 * i + 1:03d}" for i, name in enumerate(COUNTIES)}

CONTESTS = [
    {"id": "2024-president", "year": 2024, "date": "2024-11-05", "office": "President and Vice President of the United States",
     "prefix": "pres", "title": "President and Vice President", "dem": "Harris/Walz", "rep": "Trump/Vance", "find": ("Harris", "Trump"),
     "official": "ia-canvass-2024", "columns": (0, 1)},
    {"id": "2022-governor", "year": 2022, "date": "2022-11-08", "office": "Governor and Lieutenant Governor",
     "prefix": "gov", "title": "Governor and Lt. Governor", "dem": "DeJear/Van Lancker", "rep": "Reynolds/Gregg", "find": ("DeJear", "Reynolds"),
     "official": "ia-canvass-2022-governor", "columns": (1, 0)},
    {"id": "2022-us-senate", "year": 2022, "date": "2022-11-08", "office": "United States Senator",
     "prefix": "ussen", "title": "United States Senator", "dem": "Franken", "rep": "Grassley", "find": ("Franken", "Grassley"),
     "official": "ia-canvass-2022", "columns": (1, 0)},
    {"id": "2020-president", "year": 2020, "date": "2020-11-03", "office": "President and Vice President of the United States",
     "prefix": "pres", "title": "President and Vice President", "dem": "Biden/Harris", "rep": "Trump/Pence", "find": ("Biden", "Trump"),
     "official": "ia-canvass-2020", "columns": (1, 0)},
    {"id": "2020-us-senate", "year": 2020, "date": "2020-11-03", "office": "United States Senator",
     "prefix": "ussen", "title": "United States Senator", "dem": "Greenfield", "rep": "Ernst", "find": ("Greenfield", "Ernst"),
     "official": "ia-canvass-2020", "columns": (1, 0)},
]
# "columns": which of the Canvass Summary's columns are the Democratic and the Republican ticket (the Summary prints the
# candidates in ballot order, with their names broken over several lines, so the order is typed here; a wrong order
# cannot pass the control, which compares all 99 counties).

# The official statewide totals this loader was checked against on 2026-10-02, read from the documents in OFFICIAL_DOCS.
CHECKED = {
    "2024-president": {"dem": 707278, "rep": 927019, "other": 29209, "total": 1663506},
    "2022-governor": {"dem": 482950, "rep": 709198, "other": 29716, "total": 1221864},
    "2022-us-senate": {"dem": 533330, "rep": 681501, "other": 1815, "total": 1216646},
    "2020-president": {"dem": 759061, "rep": 897672, "other": 34138, "total": 1690871},
    "2020-us-senate": {"dem": 754859, "rep": 864997, "other": 51972, "total": 1671828},
}
PDF = "https://sos.iowa.gov/elections/pdf/{year}/general/{name}.pdf"
OFFICIAL_DOCS = {
    "ia-canvass-2024": {"year": 2024, "name": "canvsummary", "title": "State of Iowa Election Canvass Summary, 2024 General Election held on November 5, 2024 "
                        "(Official Canvass by County)", "canvass": "2024-12-02"},
    "ia-canvass-2022": {"year": 2022, "name": "canvsummary", "title": "2022 General Election Canvass Summary (Official Canvass by County)", "canvass": None},
    "ia-canvass-2022-governor": {"year": 2022, "name": "govcanvsummary", "title": "2022 General Election Canvass Summary: Official Canvass by County for "
                                 "Governor/Lt. Governor", "canvass": None},
    "ia-canvass-2020": {"year": 2020, "name": "canvsummary", "title": "State of Iowa Election Canvass Summary, General Election held on November 3, 2020 "
                        "(Official Canvass by County)", "canvass": "2020-11-30"},
}

# The kinds of place: (kind, the election years it is given for, covers the whole state, what it is, what its key is)
KINDS = [
    ("county", (2020, 2022, 2024), True, "Counties", "five-digit county FIPS code, as sl_places kind county"),
    ("mcd", (2022, 2024), False, "Townships, as the precinct or precincts each township's voters vote in",
     "IA-M- and the five-digit Census county-subdivision code, as sl_places kind mcd"),
    ("senate", (2022, 2024), True, "State Senate districts of the 2021 plan", "district number, as the district of our Senate races"),
    ("house", (2022, 2024), True, "State House districts of the 2021 plan", "district number, as the district of our House races"),
]
KIND_NAMES = [k[0] for k in KINDS]
REDRAWN = ("The 2020 election was held on the district plan of 2011 and on the precincts of that time; the 2021 plan and the precincts drawn "
           "after it were first used in 2022, so 2020 is not given.")

WHAT = ("How each Iowa place voted in five contests of the 2020, 2022 and 2024 general elections: the votes for the Democratic ticket, the "
        "Republican ticket, everyone else together (write-ins included) and the total, added up from the Iowa Secretary of State's official "
        "precinct results.")
NOTE = ("What this is: how the people of a place voted in that election, in the Iowa Secretary of State's official precinct results, "
        "absentee votes included, added up here by county and by legislative district. What this is not: it is not a prediction of any "
        "election; it says nothing about any candidate on a later ballot or about any voter; a nonpartisan office stays nonpartisan; and a "
        "place is not its lines for ever: the figures are for the precinct and district lines of that year. Iowa counts no votes by "
        "township: for a township the figures are those of the precinct or precincts its voters vote in, named beside them, and such a "
        "precinct often takes in other townships or a small city as well. The tickets are named only to say which election this was. In a "
        "place with very few voters the split would come close to saying how particular people voted; those contests are listed in the "
        "place's too_few, and a page should leave the split out.")
HOW_TO_READ = ("Each place's votes are four counts for each contest: dem (the Democratic ticket), rep (the Republican ticket), other (every "
               "other candidate and all write-ins) and total (the votes cast for candidates; under votes and over votes are not counted). "
               "Minnesota's file calls the first count dfl. A contest a place does not have was not held on its lines.")
TOO_FEW = (f"A place's too_few lists the contests in which it had fewer than {FEW} votes, or in which every vote went the same way. There the "
           "split comes close to saying how particular people voted, so a page should leave it out. The counts are the official record all the "
           "same, and are kept here so that the sums can be checked.")


# ---------------------------------------------------------------- the county files

def slug(county):
    return re.sub(r"[^a-z]+", "_", county.lower()).strip("_")


def county_file(year, county, cache):
    return os.path.join(cache, "results", f"{year}_{slug(county)}.{RESULTS[year]['ext']}")


def _get(url):
    b = net.get(url)
    if b[:2] == b"\x1f\x8b":
        import gzip
        b = gzip.decompress(b)
    return b


def clarity_counties(cache, refresh, say):
    """{county: (election id, version)} for 2024, from the results system's own list of the counties' reports."""
    path = os.path.join(cache, "results", "clarity_2024_counties.json")
    if os.path.exists(path) and not refresh:
        return M._load(path)
    src = RESULTS[2024]
    version = _get(src["version"]).decode("ascii").strip()
    time.sleep(1.0)
    listed = json.loads(_get(src["index"].format(version=version)))["settings"]["electiondetails"]["participatingcounties"]
    doc = {p.split("|")[0].replace("_", " "): p.split("|")[1:3] for p in listed}
    if sorted(doc) != sorted(COUNTIES):
        raise OSError(f"the results system lists {len(doc)} counties, not Iowa's 99 by name")
    M._save(path, doc)
    return doc


def workbook_links(year, cache, refresh, say):
    """{county: address of its file} from the Secretary's page of links for one election. Only the links are kept (the
    page's address and the county's name on it), never the page. A county whose link is a PDF has no workbook."""
    import html
    path = os.path.join(cache, "results", f"sos_precinct_results_{year}_links.json")
    if os.path.exists(path) and not refresh:
        return M._load(path)
    page = _get(RESULTS[year]["url"]).decode("utf-8", "replace")
    doc = {}
    for href, text in re.findall(r'<a[^>]+href="([^"]+)"[^>]*>(.*?)</a>', page, flags=re.S):
        name = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", "", text))).strip()
        if name in COUNTY_FIPS and f"/precinctresults/{year}general/" in href.lower():
            doc[name] = "https://sos.iowa.gov" + href if href.startswith("/") else html.unescape(href)
    del page
    if sorted(doc) != sorted(COUNTIES):
        raise OSError(f"the page of {year} links names {len(doc)} counties, not Iowa's 99")
    M._save(path, doc)
    return doc


def fetch_files(year, cache, refresh, say):
    """Every county's file for one election, fetched once and kept (a canvassed report does not change). Returns
    ({county: path}, [counties whose results are posted only as a PDF])."""
    src, out, todo, urls, pdf_only = RESULTS[year], {}, [], {}, []
    os.makedirs(os.path.join(cache, "results"), exist_ok=True)
    if src["kind"] == "workbook":
        try:
            links = workbook_links(year, cache, refresh, say)
        except OSError as e:
            raise Stop(f"    {year}: the Secretary of State's page of county files could not be read ({e}). Wait a few minutes and run this again.")
        pdf_only = sorted(c for c, u in links.items() if u.lower().endswith(".pdf"))
        urls = {c: re.sub(r"\s+", " ", u).replace(" ", "%20") for c, u in links.items() if c not in pdf_only}
    for county in COUNTIES:
        if county in pdf_only:
            continue
        out[county] = county_file(year, county, cache)
        if refresh or not os.path.exists(out[county]) or not os.path.getsize(out[county]):
            todo.append(county)
    if todo:
        say(f"      {year}: fetching {len(todo)} county file{'s' if len(todo) != 1 else ''}, one a second")
        ids = clarity_counties(cache, refresh, say) if src["kind"] == "clarity" else {}
    for county in todo:
        if src["kind"] == "clarity":
            url = src["file"].format(county=quote(county.replace(" ", "_")), eid=ids[county][0], version=ids[county][1])
        else:
            url = urls[county]
        for attempt in range(4):
            try:
                data = _get(url)
                break
            except OSError as e:
                if attempt == 3:
                    raise Stop(f"    {year}, {county} County: the file could not be fetched ({e}). Wait a few minutes and run this again; "
                               "the files already fetched are kept.")
                time.sleep(5 * (attempt + 1))
        with open(out[county] + ".part", "wb") as fh:
            fh.write(data)
        os.replace(out[county] + ".part", out[county])
        time.sleep(1.0)
    return out, pdf_only


def _n(v, what):
    """A count of votes from a cell: 12, 12.0, '12' or '1,234'."""
    if isinstance(v, str):
        v = v.replace(",", "").strip()
        if not re.fullmatch(r"\d+(\.0+)?", v):
            raise Stop(f"    {what} is {v!r}, not a count of votes; stopping")
        return int(float(v))
    if isinstance(v, bool) or v is None or not isinstance(v, (int, float)) or v != int(v) or v < 0:
        raise Stop(f"    {what} is {v!r}, not a count of votes; stopping")
    return int(v)


def title_of(text):
    """'United States Senator ( 1)' -> 'United States Senator': a contest's name without the number to vote for."""
    return re.sub(r"\s*\(\s*\d+\s*\)\s*$", "", (text or "").strip())


def read_clarity(data, want):
    """One county's 2024 report (detail.xml inside the zip): {contest title: {"choices": [(text, party, {precinct: votes})],
    "precincts": [names in the contest]}} for the titles `want` accepts. A choice's precinct votes are its vote
    types (Election Day, Absentee) added together, and must add up to the report's own total for the choice."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = ET.fromstring(z.read(z.namelist()[0]))
    out = {}
    for c in root.findall("Contest"):
        title = title_of(c.get("text"))
        if not want(title):
            continue
        if title in out:
            raise Stop(f"    {root.findtext('Region')} County, 2024: the contest {title} appears twice; stopping")
        choices, listed = [], []
        for ch in c.findall("Choice"):
            votes = collections.OrderedDict()
            for vt in ch.findall("VoteType"):
                for p in vt.findall("Precinct"):
                    votes[p.get("name")] = votes.get(p.get("name"), 0) + _n(p.get("votes"), f"{title}, precinct {p.get('name')}")
            if sum(votes.values()) != _n(ch.get("totalVotes"), f"{title}, a choice's total"):
                raise Stop(f"    {root.findtext('Region')} County, 2024, {title}: a choice's precincts add up to {sum(votes.values()):,} and the "
                           f"report's own total is {ch.get('totalVotes')}; stopping")
            choices.append((ch.get("text") or "", ch.get("party") or "", votes))
            listed = listed or list(votes)
            if list(votes) != listed:
                raise Stop(f"    {root.findtext('Region')} County, 2024, {title}: the choices do not list the same precincts; stopping")
        out[title] = {"choices": choices, "precincts": listed}
    return out, [p.get("name") for p in root.find("VoterTurnout/Precincts")]


SS = "{urn:schemas-microsoft-com:office:spreadsheet}"


def sheets_xml(data):
    """The sheets of a workbook saved as XML (the 2022 files, named .xls): [[row, ...], ...], a merged cell followed by blanks."""
    out = []
    for ws in ET.fromstring(data).iter(SS + "Worksheet"):
        rows = []
        for r in ws.iter(SS + "Row"):
            row = []
            for c in r.findall(SS + "Cell"):
                if c.get(SS + "Index"):
                    row += [None] * (int(c.get(SS + "Index")) - 1 - len(row))
                d = c.find(SS + "Data")
                row.append(d.text if d is not None else None)
                row += [None] * int(c.get(SS + "MergeAcross") or 0)
            rows.append(row)
        out.append(rows)
    return out


def sheets_xlsx(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    try:
        return [[list(r) for r in wb[n].iter_rows(values_only=True)] for n in wb.sheetnames]
    finally:
        wb.close()


def read_sheet(rows, where):
    """One contest's sheet: (title, [(choice, {precinct: votes})], [precincts]). The sheet is the title, a row of
    choices, a row of column headings (each choice has Election Day, Absentee, Total Votes), a row per precinct and a
    row 'Total:'. Only each choice's Total Votes column is read, and it must add up to the 'Total:' row."""
    blank = lambda v: v is None or str(v).strip() == ""                                             # noqa: E731
    head = next((i for i, r in enumerate(rows[:6]) if any(str(v).strip() == "Total Votes" for v in r if v is not None)), None)
    if head is None or head < 2 or blank(rows[0][0] if rows[0] else None):
        return None
    title = title_of(str(rows[0][0]))
    names = [(i, str(v).strip()) for i, v in enumerate(rows[head - 1]) if not blank(v)]
    cols = []
    for k, (i, name) in enumerate(names):
        end = names[k + 1][0] if k + 1 < len(names) else len(rows[head])
        tv = [j for j in range(i, end) if j < len(rows[head]) and str(rows[head][j]).strip() == "Total Votes"]
        if len(tv) != 1:
            raise Stop(f"    {where}, {title}: a choice has {len(tv)} Total Votes columns, not one; stopping")
        cols.append((name, tv[0]))
    choices, precincts, total = [(name, collections.OrderedDict()) for name, _j in cols], [], None
    for r in rows[head + 1:]:
        if not r or blank(r[0]):
            continue
        label = str(r[0]).strip()
        if label == "Total:":
            total = [_n(r[j], f"{where}, {title}, the Total row") for _name, j in cols]
            break
        if label in precincts:
            raise Stop(f"    {where}, {title}: the precinct {label} appears twice; stopping")
        precincts.append(label)
        for (name, j), (_n2, votes) in zip(cols, choices):
            votes[label] = _n(r[j], f"{where}, {title}, precinct {label}")
    if total is None or total != [sum(v.values()) for _name, v in choices]:
        raise Stop(f"    {where}, {title}: the precinct rows do not add up to the sheet's own Total row; stopping")
    return title, choices, precincts


def read_workbook(data, year, want, where):
    """One county's 2022 or 2020 workbook, in the shape read_clarity gives (a choice's party is not in these files)."""
    sheets = sheets_xml(data) if data.lstrip(b"\xef\xbb\xbf")[:5] == b"<?xml" else sheets_xlsx(data)
    out, everyone = {}, []
    for rows in sheets:
        if not rows or not rows[0] or not want(title_of(str(rows[0][0] or ""))):
            continue
        got = read_sheet(rows, where)
        if got is None:
            raise Stop(f"    {where}: the sheet {title_of(str(rows[0][0]))} is not laid out as a contest; stopping")
        title, choices, precincts = got
        if title in out:
            raise Stop(f"    {where}: the contest {title} appears twice; stopping")
        out[title] = {"choices": [(name, "", votes) for name, votes in choices], "precincts": precincts}
    return out


DISTRICT_YEARS = (2022, 2024)      # the elections held on the 2021 plan
DISTRICT = re.compile(r"^State (Representative|Senator) District (\d+)$")


def norm(name):
    """A precinct's name for comparing only: capitals and single spaces."""
    return re.sub(r"\s+", " ", (name or "").strip()).upper()


def county_rows(year, county, data, contests):
    """The precinct rows of one county in one election, and what was seen on the way."""
    titles = {c["title"] for c in contests}
    want = lambda t: t in titles or bool(DISTRICT.match(t))                                          # noqa: E731
    where = f"{county} County, {year}"
    if RESULTS[year]["kind"] == "clarity":
        found, _all = read_clarity(data, want)
    else:
        found = read_workbook(data, year, want, where)
    rows, seen = collections.OrderedDict(), {"senate": []}
    for c in contests:
        con = found.get(c["title"])
        if con is None:
            raise Stop(f"    {where}: no contest {c['title']} in the file; stopping")
        side = {}
        for k, (text, party, _votes) in enumerate(con["choices"]):
            for s, surname, code in (("dem", c["find"][0], "DEM"), ("rep", c["find"][1], "REP")):
                if re.search(rf"\b{re.escape(surname)}\b", text) and (not party or party == code):
                    if s in side:
                        raise Stop(f"    {where}, {c['title']}: two choices fit the {s} ticket; stopping")
                    side[s] = k
        if len(side) != 2 or side["dem"] == side["rep"]:
            raise Stop(f"    {where}, {c['title']}: the tickets {c['dem']} and {c['rep']} were not both found among the choices; stopping")
        for p in con["precincts"]:
            r = rows.setdefault(norm(p), {"county": COUNTY_FIPS[county], "countyname": county, "precinct": re.sub(r"\s+", " ", p.strip()), "house": None})
            v = [ch[2][p] for ch in con["choices"]]
            r[c["prefix"] + "_dem"], r[c["prefix"] + "_rep"] = v[side["dem"]], v[side["rep"]]
            r[c["prefix"] + "_other"] = sum(v) - v[side["dem"]] - v[side["rep"]]
            r[c["prefix"] + "_total"] = sum(v)
        seen[c["prefix"] + "_others"] = len(con["choices"]) - 2
    for title, con in found.items():
        m = DISTRICT.match(title)
        if not m:
            continue
        for p in con["precincts"]:
            r = rows.get(norm(p))
            if r is None:
                raise Stop(f"    {where}: the precinct {p} of {title} is in no statewide contest; stopping")
            if m.group(1) == "Representative":
                r.setdefault("houses", {})[m.group(2)] = sum(ch[2][p] for ch in con["choices"])
            else:
                seen["senate"].append((norm(p), m.group(2)))
    listed = {}
    for key, r in rows.items():
        hs = listed[key] = r.pop("houses", {})
        for c in contests:
            if c["prefix"] + "_total" not in r:
                raise Stop(f"    {where}: the precinct {r['precinct']} is not in the contest {c['title']}; stopping")
        voted = {d: n for d, n in hs.items() if n} or hs               # a district the precinct is listed under and cast no vote in does not count
        if len(voted) > 1:
            r["house_votes"] = voted                                    # a precinct in two districts: filed where most of it voted (house_of)
        r["house"] = max(voted, key=lambda d: (voted[d], -int(d))) if voted else None
    if year in DISTRICT_YEARS:
        seen["unplaced"] = [r["precinct"] for r in rows.values() if r["house"] is None]
        for key, sd in seen["senate"]:
            if not any((int(h) + 1) // 2 == int(sd) for h in listed[key]):
                raise Stop(f"    {where}: the precinct {rows[key]['precinct']} voted for State Representative District {', '.join(listed[key])} and "
                           f"State Senator District {sd}, which is not the 2021 plan's rule (Senate district n is House districts 2n-1 and 2n); stopping")
    else:
        for r in rows.values():                                         # the plan of 2011: no district is given
            r["house"] = None
            r.pop("house_votes", None)
    seen["senate"] = sorted({sd for _key, sd in seen["senate"]}, key=int)
    return list(rows.values()), seen


def results_table(year, contests, cache, refresh, say):
    """Every precinct row of one election: read from the county files (fetched if need be) and kept as one small table."""
    path = os.path.join(cache, f"sos_precinct_results_{year}.json")
    mine = [c for c in contests if c["year"] == year]
    fields = ["county", "countyname", "precinct", "house"] + [c["prefix"] + "_" + s for c in mine for s in SIDES]
    if os.path.exists(path) and not refresh:
        doc = M._load(path)
        if doc.get("fields") == fields and doc.get("method") == METHOD:
            return doc
    files, pdf_only = fetch_files(year, cache, refresh, say)
    rows, shas, senate, unplaced, others = [], {}, set(), [], {}
    for county in COUNTIES:
        if county in pdf_only:
            continue
        with open(files[county], "rb") as fh:
            data = fh.read()
        shas[county] = M._sha(data)
        got, seen = county_rows(year, county, data, mine)
        rows += got
        senate |= set(seen["senate"])
        unplaced += [f"{county} County: {p}" for p in seen.get("unplaced", [])]
        for c in mine:
            others.setdefault(c["prefix"], set()).add(seen[c["prefix"] + "_others"])
    doc = {"year": year, "method": METHOD, "fields": fields, "fetched": max(M._day(p) for p in files.values()), "layer_name": RESULTS[year]["title"],
           "files": len(shas), "files_sha256": M._sha(json.dumps(shas, sort_keys=True).encode("utf-8")),
           "pdf_only": pdf_only, "senate_checked": sorted(senate, key=int) if year in DISTRICT_YEARS else [], "unplaced": unplaced, "other_choices": {p: sorted(v) for p, v in others.items()}, "rows": rows}
    M._save(path, doc)
    say(f"      {year}: {len(rows):,} precinct rows read from {len(shas)} county files")
    return doc


# ---------------------------------------------------------------- which precinct a township votes in

LSA = {"title": "Iowa Precincts (the precinct plans in effect from 2022)", "agency": "Iowa Legislative Services Agency, on the State of Iowa's ArcGIS site",
       "service": "https://services.arcgis.com/vPD5PVLI6sfkZ5E4/arcgis/rest/services/Iowa_Precincts/FeatureServer/0",
       "file": "lsa_precincts_2022_geometry_4326.json.gz"}
TIGER = {"cousub": {"url": "https://www2.census.gov/geo/tiger/TIGER2025/COUSUB/tl_2025_19_cousub.zip", "file": "tl_2025_19_cousub.zip",
                    "title": "TIGER/Line 2025, county subdivisions, Iowa (the township lines)"},
         "place": {"url": "https://www2.census.gov/geo/tiger/TIGER2025/PLACE/tl_2025_19_place.zip", "file": "tl_2025_19_place.zip",
                   "title": "TIGER/Line 2025, places, Iowa (the city lines)"}}
T_GRID = 24                    # a township is sampled on a lattice of this many points a side
P_GRID = 9                     # a precinct and a city, on this many
CITY_SHARE = 0.5               # a precinct with this share of its points inside one city is a city precinct
PART = 0.05                    # a precinct holding less than this share of a township's rural points is a sliver of a line, not a part


def _shapes():
    from ballot import wi_place_votes as W                     # the polygon test and the lattice are Wisconsin's
    return W.Shape, W.lattice


def precinct_shapes(cache, refresh, say):
    """[(county code, official precinct name, rings)] of the Legislative Services Agency's precinct layer. The copy the
    map builder keeps (lsa_precincts_2022_geometry_4326.json.gz) is used when it is there; else the layer is read
    from the Agency's service, a thousand precincts a request, and kept under the same name."""
    import gzip
    path = os.path.join(cache, LSA["file"])
    if not os.path.exists(path):
        rows, offset = [], 0
        while True:
            j = json.loads(_get(f"{LSA['service']}/query?where={quote('1=1')}&outFields=CoFIPS,CoName,PctNameOfficial&returnGeometry=true&outSR=4326"
                                f"&geometryPrecision=7&orderByFields=OBJECTID&resultOffset={offset}&resultRecordCount=1000&f=json"))
            if "error" in j:
                raise OSError(str(j["error"]))
            feats = j.get("features", [])
            rows += [[f["attributes"], (f.get("geometry") or {}).get("rings") or []] for f in feats]
            offset += len(feats)
            if not feats or not j.get("exceededTransferLimit"):
                break
            time.sleep(1.0)
        with gzip.open(path + ".part", "wt", encoding="utf-8") as fh:
            json.dump({"service": LSA["service"], "fields": "CoFIPS,CoName,PctNameOfficial", "fetched": M._now(), "rows": rows}, fh)
        os.replace(path + ".part", path)
    with gzip.open(path, "rt", encoding="utf-8") as fh:
        doc = json.load(fh)
    return [(STATE_FIPS + r[0]["CoFIPS"], (r[0]["PctNameOfficial"] or "").strip(), r[1]) for r in doc["rows"] if r[1]], doc.get("fetched"), M._sha_file(path)


def tiger_shapes(kind, cache, say):
    """The Census Bureau's township or city lines: [(record, rings)]; only names, codes and classes are read."""
    import shapefile
    path = os.path.join(cache, TIGER[kind]["file"])
    net.download(TIGER[kind]["url"], path, 3650, tries=3, say=say)
    out = []
    with shapefile.Reader(path) as r:
        names = [f[0] for f in r.fields[1:]]
        for sr in r.iterShapeRecords():
            rec = dict(zip(names, sr.record))
            pts, parts = sr.shape.points, list(sr.shape.parts) + [len(sr.shape.points)]
            out.append((rec, [[list(p) for p in pts[parts[i]:parts[i + 1]]] for i in range(len(parts) - 1)]))
    return out, M._day(path), M._sha_file(path)


def place_townships(precincts, townships, cities, t_grid=T_GRID):
    """The geometry, on plain lists: precincts [(county, name, rings)], townships [(county, code, name, rings)],
    cities [(name, rings)]. Returns {county: {"precincts": {name: {"city": the city it lies in, or None, "cities":
    {city: "all" or "part"}}}, "townships": {code: {"name", "precincts": [names], "share": {name: share of the
    township's rural points}}}}}.

    A precinct is a city precinct when at least half of a lattice of points inside it lies in one city. A township's
    precincts are the others that hold at least a twentieth of the lattice points of the township that lie in no city
    precinct: where its voters outside the larger cities vote. A city is in a precinct when at least a tenth of the
    city's own lattice points are, or a twentieth of the precinct's are in the city: all of it when nine tenths of the
    city's points are, else part. Where two shapes of the map carry one name, or a shape carries none, neither can
    be found in the results by name: each is kept under a name that says so, and a township that votes there is left
    out."""
    Shape, lattice = _shapes()
    city_shapes = [(name, Shape(rings), rings) for name, rings in cities]
    out = {}
    by_county = collections.defaultdict(list)
    count = collections.Counter((county, name) for county, name, _rings in precincts)
    for county, name, rings in precincts:
        if count[(county, name)] > 1 or not name:
            name = f"{name or 'a precinct the map does not name'} (shape {len(by_county[county]) + 1} of the county's map; the name is not its alone)"
        by_county[county].append((name, rings))
    for county, named in sorted(by_county.items()):
        plist = [(name, Shape(rings), rings) for name, rings in named]
        box = (min(s.box[0] for _n, s, _r in plist), min(s.box[1] for _n, s, _r in plist), max(s.box[2] for _n, s, _r in plist), max(s.box[3] for _n, s, _r in plist))
        near = [(n, s, r) for n, s, r in city_shapes if not (s.box[2] < box[0] or s.box[0] > box[2] or s.box[3] < box[1] or s.box[1] > box[3])]
        pinfo = {}
        for name, _shape, rings in plist:
            pts = lattice(rings, P_GRID)
            hits = collections.Counter(cn for x, y in pts for cn, cs, _cr in near if cs.has(x, y))
            top = hits.most_common(1)
            pinfo[name] = {"city": top[0][0] if top and pts and top[0][1] >= CITY_SHARE * len(pts) else None,
                           "cities": {cn: "part" for cn, n in hits.items() if n >= PART * len(pts)}}

        def precinct_at(x, y):
            for name, shape, _rings in plist:
                if shape.has(x, y):
                    return name
            return None
        for cn, _cs, crings in near:
            pts = lattice(crings, P_GRID)
            hits = collections.Counter(precinct_at(x, y) for x, y in pts)
            for name, n in hits.items():
                if name is None or not pts:
                    continue
                if n >= 0.9 * len(pts):
                    pinfo[name]["cities"][cn] = "all"
                elif n >= 0.1 * len(pts):
                    pinfo[name]["cities"].setdefault(cn, "part")
        for info in pinfo.values():
            info["cities"] = dict(sorted(info["cities"].items()))
        towns = {}
        for c2, code, tname, rings in townships:
            if c2 != county:
                continue
            pts = lattice(rings, t_grid)
            hits = collections.Counter(precinct_at(x, y) for x, y in pts)
            rural = {n: k for n, k in hits.items() if n is not None and pinfo[n]["city"] is None}
            total = sum(rural.values())
            mine = sorted(n for n, k in rural.items() if total and k >= PART * total)
            towns[code] = {"name": tname, "precincts": mine, "share": {n: round(rural[n] / total, 3) for n in mine},
                           "rural_points": total, "points": len(pts)}
        out[county] = {"precincts": pinfo, "townships": towns}
    return out


def township_precincts(cache, refresh, say):
    """Which precinct or precincts each township's voters outside the larger cities vote in, from the Legislative
    Services Agency's precinct layer and the Census Bureau's township and city lines; kept as one small table."""
    path = os.path.join(cache, "ia_township_precincts.json")
    if os.path.exists(path) and not refresh:
        doc = M._load(path)
        if doc.get("method") == METHOD:
            return doc
    try:
        precincts, p_day, p_sha = precinct_shapes(cache, refresh, say)
        subs, s_day, s_sha = tiger_shapes("cousub", cache, say)
        places, c_day, c_sha = tiger_shapes("place", cache, say)
    except Exception as e:  # noqa: BLE001  without the maps the townships are left out, and the file says so
        if os.path.exists(path):
            say(f"      the precinct and township maps could not be read again ({e}); using the table of {M._day(path)}")
            return M._load(path)
        say(f"      the precinct and township maps could not be read ({e}); townships are left out")
        return None
    say(f"      placing the townships among {len(precincts):,} precincts of the map (a few minutes)")
    townships = [(STATE_FIPS + r["COUNTYFP"], r["COUSUBFP"], r["NAMELSAD"], rings) for r, rings in subs if r["NAMELSAD"].endswith(" township")]
    cities = [(r["NAMELSAD"], rings) for r, rings in places if str(r["CLASSFP"]).startswith("C")]
    doc = {"what": "Which precinct or precincts each Iowa township's voters outside the larger cities vote in.", "method": METHOD, "made": M._now(),
           "grid": {"township": T_GRID, "precinct": P_GRID, "city_share": CITY_SHARE, "part": PART},
           "sources": {"precincts": {"fetched": p_day, "sha256": p_sha}, "cousub": {"fetched": s_day, "sha256": s_sha}, "place": {"fetched": c_day, "sha256": c_sha}},
           "counties": place_townships(precincts, townships, cities)}
    M._save(path, doc)
    return doc


CODE = r"^\s*(\(\w+\)|\d{3}(?=\s))\s*"         # a precinct's code in front of its name: '074 ' or '(AG) '


def canon(name, loose=False):
    """A precinct's name for matching one list's spelling to another's: capitals, no code in front, no punctuation,
    TWP as TOWNSHIP, numbers without leading zeros; loosely, also without 'TOWNSHIP', 'CITY', 'OF', 'PRECINCT'."""
    s = re.sub(CODE, "", (name or "").upper())
    found = ["TOWNSHIP" if w in ("TWP", "TWSHP", "TWNSHP") else re.sub(r"^0+(?=\d)", "", w) for w in re.findall(r"[A-Z]+|\d+", s)]
    return " ".join(w for w in found if not (loose and w in ("TOWNSHIP", "CITY", "OF", "PRECINCT", "PCT")))


# The counties whose pairings by likeness were read one by one on 2026-10-02 (the counties whose townships our pages
# use): Black Hawk, Dallas, Johnson, Linn, Marshall, Polk, Scott, Story, Warren, Woodbury. In any other county a
# precinct is paired only when the map and the results write the same name, so a township there is given or left out,
# never guessed. Read a county's pairings (coverage.precinct_names_matched_by_likeness) before adding it here.
LIKENESS_READ = ("19013", "19049", "19103", "19113", "19127", "19153", "19163", "19169", "19181", "19193")


def link_names(official, results, likeness=True):
    """{the map's precinct name: (the results' name for it, "the same name" or "a like name")} within one county. Names
    equal once written one way (canon), then once written loosely, are the same name, when only one on each side is
    written so. With `likeness`, what is left is paired when one name is the beginning of exactly one other ('OX' and
    'OX/OC'), or by likeness when each is the other's best fit by a clear margin; anything else stays unpaired, and a
    township that needs it is left out."""
    import difflib
    out, left_o, left_r = {}, list(official), list(results)
    for loose in (False, True):
        by_o, by_r = collections.defaultdict(list), collections.defaultdict(list)
        for o in left_o:
            by_o[canon(o, loose)].append(o)
        for r in left_r:
            by_r[canon(r, loose)].append(r)
        for key, os_ in by_o.items():
            if key and len(os_) == 1 and len(by_r.get(key, [])) == 1:
                out[os_[0]] = (by_r[key][0], "the same name")
        left_o = [o for o in left_o if o not in out]
        used = {r for r, _how in out.values()}
        left_r = [r for r in left_r if r not in used]
    if not likeness:
        return out
    starts = lambda o, r: canon(o) and canon(r).split()[:len(canon(o).split())] == canon(o).split()            # noqa: E731
    for o in list(left_o):
        fits = [r for r in left_r if starts(o, r)]
        if len(fits) == 1 and sum(1 for o2 in left_o if starts(o2, fits[0])) == 1:
            out[o] = (fits[0], "a like name")
    left_o = [o for o in left_o if o not in out]
    used = {r for r, _how in out.values()}
    left_r = [r for r in left_r if r not in used]

    def like(a, b):
        ca, cb = canon(a), canon(b)
        if not ca or not cb:
            return 0.0
        score = difflib.SequenceMatcher(None, ca, cb).ratio()
        ta, tb = ca.split(), cb.split()
        if len(ta) > 1 and len(tb) > 1 and ta[0] == tb[0] and ta[0].isdigit() and ta[1][:3] == tb[1][:3]:      # the same number and first word
            score = max(score, 0.9)
        return score
    best_o = {o: sorted(((like(o, r), r) for r in left_r), reverse=True)[:2] for o in left_o}
    best_r = {r: sorted(((like(o, r), o) for o in left_o), reverse=True)[:2] for r in left_r}
    for o, fits in best_o.items():
        if not fits or fits[0][0] < 0.72 or (len(fits) > 1 and fits[0][0] - fits[1][0] < 0.08):
            continue
        r = fits[0][1]
        back = best_r[r]
        if back[0][1] == o and (len(back) == 1 or back[0][0] - back[1][0] >= 0.08):
            out[o] = (r, "a like name")
    return out


LIST_UP_TO = 6                 # a township's name lists up to this many others it votes with; more are counted, and listed in shared_with
NAME_FROM = 10                 # a district shared precincts move by this many votes or more says so in its name


def words(items):
    items = list(items)
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]


def township_places(geo, tables, contests):
    """The township records: for each township of the geometry table, the results rows of its precinct or precincts in
    each year, and a name that says exactly what the figures are. Returns (places, what was left out and why, the
    precinct names paired by likeness)."""
    places, left, likes = {}, {}, []
    years = sorted(y for y in tables if any(k[0] == "mcd" and y in k[1] for k in KINDS))
    rows_of = {y: collections.defaultdict(dict) for y in years}
    for y in years:
        for r in tables[y]["rows"]:
            if r.get("precinct"):
                rows_of[y][r["county"]][r["precinct"]] = r
    for county, g in sorted((geo or {}).get("counties", {}).items()):
        links = {y: link_names(list(g["precincts"]), list(rows_of[y][county]), likeness=county in LIKENESS_READ) for y in years}
        towns_of = collections.defaultdict(list)
        for code, t in g["townships"].items():
            for p in t["precincts"]:
                towns_of[p].append(code)
        for code, t in sorted(g["townships"].items()):
            key = "IA-M-" + code
            if not t["precincts"]:
                left[key] = f"{t['name']}: the precinct map shows no precinct outside a city for it"
                continue
            votes, named = {}, {}
            for y in years:
                got = [links[y].get(p) for p in t["precincts"]]
                if any(x is None for x in got):
                    continue
                named[str(y)] = [x[0] for x in got]
                likes += [{"county": county, "year": y, "map": p, "results": x[0]} for p, x in zip(t["precincts"], got)
                          if x[1] != "the same name" and not any(k["county"] == county and k["year"] == y and k["map"] == p for k in likes)]
                for c in contests:
                    if c["year"] != y:
                        continue
                    v = [0, 0, 0, 0]
                    for res, _how in got:
                        r = rows_of[y][county][res]
                        for i, s in enumerate(SIDES):
                            v[i] += r[c["prefix"] + "_" + s]
                    votes[c["id"]] = dict(zip(SIDES, v))
            if not votes:
                left[key] = f"{t['name']}: its precinct on the map ({words(t['precincts'])}) could not be told apart by name in the results"
                continue
            mine = set(t["precincts"])
            others = sorted({(g["townships"][c2]["name"], set(g["townships"][c2]["precincts"]) <= mine) for p in mine for c2 in towns_of[p] if c2 != code})
            cities = sorted({(re.sub(r" city$", "", cn), how == "all") for p in mine for cn, how in g["precincts"][p]["cities"].items()})
            whole = {cn for cn, all_ in cities if all_}
            with_ = [n if all_ else f"part of {n}" for n, all_ in others] + \
                    [f"the city of {cn}" if cn in whole else f"part of the city of {cn}" for cn in sorted({cn for cn, _a in cities})]
            # the precinct as the results name it; where the results write it in capitals, as the map names it
            shown = [re.sub(r"\s*-$", "", re.sub(CODE, "", r if re.search(r"[a-z]", r) else p)).strip() for p, r in zip(t["precincts"], named[max(named)])]
            pword = "precinct" if len(shown) == 1 else "precincts"
            if len(with_) > LIST_UP_TO:
                nt, nc = len(others), len({cn for cn, _a in cities})
                name = (f"The {pword} {words(shown)}, where {t['name']} votes with {nt} other township{'s' if nt != 1 else ''}"
                        + (f" and {nc} cit{'ies' if nc != 1 else 'y'}" if nc else "") + ", in whole or in part")
            elif with_:
                name = f"The {pword} {words(shown)}, where {t['name']} votes with {words(with_)}"
            else:
                name = f"{t['name']} (the {pword} {words(shown)})"
            rec = {"name": name, "type": "township", "counties": [county], "township": t["name"], "precincts": named, "votes": votes}
            if with_:
                rec["shared_with"] = with_
            places[key] = rec
    return places, left, likes


# ---------------------------------------------------------------- the official figures

def read_canvass(path, title):
    """One contest of the State Canvass Summary: {"counties": {name: [the Total row's numbers]}, "state": [numbers],
    "page": where it begins}. Each county is three rows (Election Day, Absentee, Total); only the Total row is read.
    The numbers run: each candidate in ballot order, write-in, under votes, over votes, total."""
    from ballot import pdftext
    lines = pdftext.lines(path)
    start = next((i for i, (_p, _y, t) in enumerate(lines) if t.strip() == title), None)
    if start is None:
        raise ValueError(f"no contest {title}")
    end = next((i for i in range(start + 1, len(lines)) if lines[i][2].strip() == "IOWA SECRETARY OF STATE"), len(lines))
    return dict(parse_canvass([t for _p, _y, t in lines[start + 1:end]], title), page=lines[start][0])


def parse_canvass(texts, title):
    counties, state, current = {}, None, None
    for t in texts:
        m = re.match(r"^([A-Za-z'. ]+?) ((?:[\d,]+ ){3,}[\d,]+)$", t.strip())
        if not m:
            continue
        label, nums = m.group(1).strip(), [int(x.replace(",", "")) for x in m.group(2).split()]
        if label == "Absentee":
            continue
        if label == "Total":
            if sum(nums[:-1]) != nums[-1]:
                raise ValueError(f"{title}: a Total row does not add up to its last column")
            if current is None or current == "TOTAL":
                state = nums
            else:
                counties[current] = nums
            current = None
            continue
        current = re.sub(r"\s*Election(?: Day)?$", "", label)             # 'Adair', 'Adair Election Day', 'Emmet Election', 'PottawattamieElection Day'
    if sorted(counties) != sorted(COUNTIES):
        raise ValueError(f"{title}: {len(counties)} county Total rows, not Iowa's 99 by name")
    if state is None or state != [sum(c[i] for c in counties.values()) for i in range(len(state))]:
        raise ValueError(f"{title}: the county rows do not add up to the statewide Total row")
    return {"counties": counties, "state": state}


def canvass_split(nums, columns):
    """[dem, rep, other, total] and [under votes, over votes, the Summary's own total] of one Total row."""
    cast = nums[:-3]
    dem, rep = cast[columns[0]], cast[columns[1]]
    return [dem, rep, sum(cast) - dem - rep, sum(cast)], nums[-3:]


def official_totals(contests, cache, refresh, say):
    """({contest id: {dem, rep, other, total, under_votes, over_votes, source, where, read, counties: {name: [4]}}},
    the documents' records). A contest whose document cannot be read stands on the figures typed in CHECKED, without
    the county rows."""
    docs, out = {}, {}
    os.makedirs(os.path.join(cache, "results"), exist_ok=True)
    for sid, src in OFFICIAL_DOCS.items():
        if not any(c["official"] == sid for c in contests):
            continue
        url = PDF.format(year=src["year"], name=src["name"])
        path = os.path.join(cache, "results", f"sos_{src['year']}_general_{src['name']}.pdf")
        rec = {"id": sid, "kind": "official canvass by county", "agency": AGENCY, "title": src["title"], "url": url}
        if src.get("canvass"):
            rec["canvassed"] = src["canvass"]
        try:
            net.download(url, path, 0 if refresh else 3650, tries=3, say=say)
            rec.update(fetched=M._day(path), sha256=M._sha_file(path))
            docs[sid] = (rec, path)
        except Exception as e:  # noqa: BLE001  the control then runs on the typed figures, and the file says so
            say(f"      {src['title'][:60]}: could not be fetched ({e}); the control uses the figures typed in on 2026-10-02")
            rec.update(unread=f"could not be read on {M._now()}: {e}")
            docs[sid] = (rec, None)
    for c in contests:
        rec, path = docs[c["official"]]
        typed, got = CHECKED.get(c["id"]), None
        if path:
            try:
                parsed = read_canvass(path, c["title"])
                state, extra = canvass_split(parsed["state"], c["columns"])
                got = dict(zip(SIDES, state), under_votes=extra[0], over_votes=extra[1], summary_total=extra[2],
                           counties={n: canvass_split(v, c["columns"])[0] for n, v in parsed["counties"].items()},
                           where=f"{c['title']}, from page {parsed['page']} of the PDF: the Total row of each county and of the state",
                           read=f"from the document on {rec.get('fetched')}")
            except Exception as e:  # noqa: BLE001
                say(f"      {rec['title'][:60]}: {e}; the control uses the figures typed in on 2026-10-02")
        if got is None:
            if typed is None:
                raise Stop(f"    {c['id']}: no official total could be read and none is typed in this loader")
            got = dict(typed, counties=None, where=c["title"],
                       read="typed into the loader from the document on 2026-10-02; the document could not be read again today")
        elif typed is not None and {s: got[s] for s in SIDES} != typed:
            say(f"      {c['id']}: the document now reads {dict((s, got[s]) for s in SIDES)}; this loader was checked against {typed}")
        out[c["id"]] = dict(got, source=c["official"])
    return out, [rec for rec, _path in docs.values()]


# ---------------------------------------------------------------- adding up

def house_of(row):
    """(the House district a precinct is filed in, the votes for State Representative it cast in other districts)."""
    hv = row.get("house_votes")
    if not hv:
        return row.get("house"), {}
    top = max(hv, key=lambda d: (hv[d], -int(d)))
    return top, {d: n for d, n in hv.items() if d != top}


def tally(tables, contests, official, geo):
    """Everything the tables say, added up: (places, statewide sums, per-kind sums, the contests each kind is given
    for, notes)."""
    years_of = {k[0]: k[1] for k in KINDS}
    given = {k: [c["id"] for c in contests if c["year"] in years_of[k]] for k in KIND_NAMES}
    votes = {k: collections.defaultdict(dict) for k in ("county", "senate", "house")}
    names, state = {"county": {}, "senate": {}, "house": {}}, {}
    info = {"from_summary": {}, "split": [], "unplaced": []}
    for year in sorted(tables, reverse=True):
        t = tables[year]
        mine = [c for c in contests if c["year"] == year]
        have = {r["countyname"] for r in t["rows"]}
        rows = list(t["rows"])
        for county in COUNTIES:
            if county in have:
                continue
            row = {"county": COUNTY_FIPS[county], "countyname": county, "precinct": None, "house": None}
            for c in mine:                                             # a county with no precinct file: its row of the Canvass Summary
                off = (official[c["id"]].get("counties") or {}).get(county)
                if off is None:
                    raise Stop(f"    {year}, {county} County: no precinct file and no row of the Canvass Summary; stopping")
                row.update({c["prefix"] + "_" + s: off[i] for i, s in enumerate(SIDES)})
            info["from_summary"].setdefault(str(year), []).append(county)
            rows.append(row)
        districts = year in years_of["house"]
        for r in rows:
            keys = {"county": r["county"]}
            names["county"].setdefault(r["county"], f"{r['countyname']} County")
            if districts:
                h, elsewhere = house_of(r)
                if h is None:
                    info["unplaced"].append(f"{year}, {r['countyname']} County, {r['precinct']}")
                else:
                    keys["house"], keys["senate"] = h, str((int(h) + 1) // 2)
                    names["house"].setdefault(h, f"House District {h}")
                    names["senate"].setdefault(keys["senate"], f"Senate District {keys['senate']}")
                    for d, n in sorted(elsewhere.items(), key=lambda kv: int(kv[0])):
                        if n:
                            info["split"].append({"year": year, "county": f"{r['countyname']} County", "precinct": r["precinct"], "filed_in": h,
                                                  "also_in": d, "house_votes_there": n, "house_votes_here": r["house_votes"][h]})
            for c in mine:
                v = [r[c["prefix"] + "_" + s] for s in SIDES]
                if v[0] + v[1] + v[2] != v[3] or min(v) < 0:
                    raise Stop(f"    {year}, {r['countyname']} County, {r['precinct']}: the counts do not add up to the total; stopping")
                s = state.setdefault(c["id"], [0, 0, 0, 0])
                for i in range(4):
                    s[i] += v[i]
                for kind, key in keys.items():
                    if c["id"] not in given[kind]:
                        continue
                    cur = votes[kind][key].setdefault(c["id"], [0, 0, 0, 0])
                    for i in range(4):
                        cur[i] += v[i]
    pack = lambda d: {cid: dict(zip(SIDES, d[cid])) for cid in (c["id"] for c in contests) if cid in d}             # noqa: E731
    places = {k: {} for k in KIND_NAMES}
    for kind in ("county", "senate", "house"):
        for key in sorted(votes[kind], key=M.sort_key):
            places[kind][key] = {"name": names[kind][key], "votes": pack(votes[kind][key])}
    for a, p in places["house"].items():
        p["senate"] = str((int(a) + 1) // 2)
    for s, p in places["senate"].items():
        p["house"] = [str(2 * int(s) - 1), str(2 * int(s))]
    # a precinct that voted in two House districts is counted whole where most of it voted: both districts say so
    for kind in ("house", "senate"):
        shared = collections.defaultdict(list)
        for x in info["split"]:
            a, b = (x["filed_in"], x["also_in"]) if kind == "house" else (str((int(x["filed_in"]) + 1) // 2), str((int(x["also_in"]) + 1) // 2))
            if a != b:
                shared[a].append(dict(x, counted="here, whole"))
                shared[b].append(dict(x, counted="in the other district"))
        for d, xs in shared.items():
            if d not in places[kind]:
                continue
            most = max(sum(x["house_votes_there"] for x in xs if x["year"] == y) for y in {x["year"] for x in xs})
            n = len({(x["county"], x["precinct"]) for x in xs})
            places[kind][d]["approximate"] = {"precincts": xs, "about_votes": most}
            if most >= NAME_FROM:
                places[kind][d]["name"] += (f" (approximate, by about {most:,} votes: {n} precinct{'s' if n != 1 else ''} shared with a "
                                            "neighbouring district, each counted whole where most of it voted)")
    mcd, left, likes = township_places(geo, tables, contests) if geo else ({}, {}, [])
    for key in sorted(mcd, key=M.sort_key):
        places["mcd"][key] = mcd[key]
    for kind in KIND_NAMES:
        for rec in places[kind].values():
            hold = [cid for cid, v in rec["votes"].items() if v["total"] < FEW or max(v["dem"], v["rep"], v["other"]) == v["total"]]
            if hold:
                rec["too_few"] = hold
    sums = {kind: {cid: [sum(p["votes"][cid][x] for p in places[kind].values() if cid in p["votes"]) for x in SIDES] for cid in given[kind]}
            for kind in KIND_NAMES}
    info.update(townships_left=left, names_by_likeness=likes)
    return places, state, sums, given, info


def control(state, sums, official, contests, tables, places, info):
    failed, per = [], {}
    for c in contests:
        mine = dict(zip(SIDES, state[c["id"]]))
        off = official[c["id"]]
        theirs = {n: off[n] for n in SIDES}
        diff = {n: mine[n] - theirs[n] for n in SIDES if mine[n] != theirs[n]}
        per[c["id"]] = {"sum_of_precincts": mine, "official": theirs, "equal": not diff, "source": off["source"], "where": off["where"], "read": off["read"]}
        if "summary_total" in off:
            per[c["id"]]["the_summary_also_counts"] = {"under_votes": off["under_votes"], "over_votes": off["over_votes"], "its_total": off["summary_total"]}
        if diff:
            per[c["id"]]["difference"] = diff
            failed.append(f"{c['id']}: the precincts add up to {mine} and the official totals ({off['source']}) are {theirs}; precincts minus official: {diff}")
        if off.get("counties"):
            taken = set(info["from_summary"].get(str(c["year"]), []))
            bad = [n for n, v in off["counties"].items() if [places["county"][COUNTY_FIPS[n]]["votes"][c["id"]][s] for s in SIDES] != v]
            per[c["id"]]["counties"] = {"compared": len(off["counties"]) - len(taken), "equal": len(off["counties"]) - len(taken) - len(bad), "differ": bad}
            if taken:
                per[c["id"]]["counties"]["taken_from_the_summary"] = sorted(taken)
            failed += [f"{c['id']}, {n} County: the precincts add up to {places['county'][COUNTY_FIPS[n]]['votes'][c['id']]} and the Canvass Summary's row is "
                       f"{dict(zip(SIDES, off['counties'][n]))}" for n in bad]
    kinds = {}
    for kind, _years, covers, _what, _key in KINDS:
        if not covers:
            continue
        bad = [cid for cid, s in sums[kind].items() if s != state[cid]]
        kinds[kind] = "equal" if not bad else f"differs for {', '.join(bad)}"
        failed += [f"{kind}: its places add up to {sums[kind][cid]} for {cid}, and the precincts to {state[cid]}" for cid in bad]
    failed += [f"a precinct in no House district: {x}" for x in info["unplaced"]]
    rec = {"result": "equal" if not failed else "differs",
           "statement": ("For every contest the precinct rows add up to the State Canvass Summary's statewide totals and, county by county, to "
                         "its row for each of the 99 counties; in every county file each choice's precincts add up to the file's own total; "
                         "and counties, House districts and Senate districts each add up to the statewide sum."
                         if not failed else "The sums do not all agree; see the differences."),
           "precincts": {str(y): sum(1 for r in t["rows"] if r.get("precinct")) for y, t in sorted(tables.items())}, "contests": per, "kinds": kinds,
           "notes": []}
    for y, v in sorted(info["from_summary"].items()):
        rec["notes"].append(f"{y}: the Secretary of State posts the precinct results of {words(v)} {'County' if len(v) == 1 else 'counties'} only as "
                            f"PDF, so {'that county is' if len(v) == 1 else 'those counties are'} taken whole from the Canvass Summary's own "
                            f"{'row' if len(v) == 1 else 'rows'}; the other {99 - len(v)} are added up from their precincts and compared with it.")
    checked = {str(y): len(t.get("senate_checked") or []) for y, t in sorted(tables.items()) if t.get("senate_checked")}
    if checked:
        rec["districts"] = {"statement": "A precinct is in the House district whose contest for State Representative the county's file lists it "
                                         "under; a Senate district is two House districts (n is 2n-1 and 2n). Every precinct listed under a contest for "
                                         "State Senator was checked against that rule. A precinct listed under two House contests is counted whole in "
                                         "the district where most of its votes for State Representative were cast, and both districts are marked "
                                         "approximate by the number of votes cast in the other.",
                            "senate_contests_checked": checked, "precincts_in_two_house_districts": info["split"]}
    return rec, failed


# ---------------------------------------------------------------- the file

def source_records(tables, geo, official_docs):
    out = []
    for y in sorted(tables, reverse=True):
        t, src = tables[y], RESULTS.get(y, {})
        out.append({"id": f"ia-sos-precinct-results-{y}", "kind": "official results by precinct", "agency": AGENCY, "title": src.get("title"),
                    "url": src.get("url"), "listed_on": STATS_PAGE, "county_files": t.get("files"), "files_sha256": t.get("files_sha256"),
                    "precincts": sum(1 for r in t["rows"] if r.get("precinct")), "fetched": t.get("fetched"), "rows_sha256": M.rows_fingerprint(t["rows"]),
                    "read": "Of each county's file: the precinct rows of the contests named here (each choice's total, absentee votes included) and "
                            "the lists of precincts under the contests for State Representative and State Senator. Of the choices only the two "
                            "tickets named here are kept apart; every other choice is added into other."})
    if geo:
        g = geo.get("sources", {})
        out.append({"id": "ia-lsa-precincts", "kind": "official precinct lines", "agency": LSA["agency"], "title": LSA["title"], "url": LSA["service"],
                    "fetched": g.get("precincts", {}).get("fetched"), "sha256": g.get("precincts", {}).get("sha256"),
                    "read": "Each precinct's county, official name and lines, to see which precinct a township's voters vote in."})
        for kind in ("cousub", "place"):
            out.append({"id": f"census-tiger-2025-{kind}-19", "kind": "official boundary file", "agency": "U.S. Census Bureau", "title": TIGER[kind]["title"],
                        "url": TIGER[kind]["url"], "fetched": g.get(kind, {}).get("fetched"), "sha256": g.get(kind, {}).get("sha256"),
                        "read": "Names, codes and lines only."})
    return out + official_docs


def build(tables, official, official_docs, geo, contests=CONTESTS, strict=True):
    places, state, sums, given, info = tally(tables, contests, official, geo)
    ctl, failed = control(state, sums, official, contests, tables, places, info)
    if strict:
        if len(places["county"]) != 99:
            failed.append(f"{len(places['county'])} counties, not 99")
        if any(c["year"] in DISTRICT_YEARS for c in contests) and (len(places["senate"]) != 50 or len(places["house"]) != 100):
            failed.append(f"{len(places['senate'])} Senate and {len(places['house'])} House districts, not 50 and 100")
        failed += [f"{len(t['rows'])} precinct rows in {y}, not about 1,650" for y, t in tables.items() if not 1400 <= len(t["rows"]) <= 1900]
    records = []
    for c in contests:
        t = tables[c["year"]]
        records.append({"id": c["id"], "date": c["date"], "office": c["office"], "table": f"ia-sos-precinct-results-{c['year']}",
                        "kinds": [kind for kind in KIND_NAMES if c["id"] in given[kind]],
                        "dem": {"party": "Democratic", "ticket": c["dem"], "column": f"the choice naming {c['find'][0]}"},
                        "rep": {"party": "Republican", "ticket": c["rep"], "column": f"the choice naming {c['find'][1]}"},
                        "other": {"what": "every other candidate and all write-ins, together", "choices": (t.get("other_choices") or {}).get(c["prefix"])},
                        "total": {"what": "the votes cast for candidates; under votes and over votes are not counted"},
                        "statewide": dict(zip(SIDES, state[c["id"]])), "official_source": official[c["id"]]["source"]})
    kinds = {}
    for kind, years, covers, what, key in KINDS:
        kinds[kind] = {"what": what, "key": key, "places": len(places[kind]), "contests": given[kind], "covers_the_state": covers,
                       "places_with_a_contest_too_few": sum(1 for p in places[kind].values() if p.get("too_few"))}
        if 2020 not in years and any(c["year"] == 2020 for c in contests):
            kinds[kind]["why_not_2020"] = REDRAWN
    for kind in ("house", "senate"):
        kinds[kind]["approximate"] = sorted((d for d, p in places[kind].items() if p.get("approximate")), key=int)
        kinds[kind]["note"] = ("A precinct is filed in the House district whose contest for State Representative the county's results list it "
                               "under. A precinct listed under two districts is counted whole where most of its votes for State Representative "
                               "were cast; both districts are marked approximate, by the number of votes for State Representative cast in the other, "
                               f"and say so in their names when that is {NAME_FROM} votes or more.")
    kinds["mcd"]["note"] = ("Iowa counts no votes by township. Each record is the precinct or precincts the township's voters outside the larger "
                            "cities vote in, by the Legislative Services Agency's precinct map of the 2022 plans, and its name says which precinct "
"that is and who else votes there. A precinct redrawn since 2022 under the same name would not be noticed. A township "
                            "whose precinct could not be told apart by name in the results is left out: outside the ten counties whose pairings "
                            "were read one by one, that is every township whose precinct the map and the results spell differently.")
    doc = {"what": WHAT, "note": NOTE, "state": "IA", "generated": M._now(), "method_version": METHOD, "how_to_read": HOW_TO_READ,
           "vote_keys": list(SIDES), "too_few": {"fewer_than": FEW, "why": TOO_FEW}, "contests": records, "kinds": kinds, "control": ctl,
           "coverage": {"townships_given": len(places["mcd"]), "townships_left_out": info["townships_left"],
                        "precinct_names_matched_by_likeness": info["names_by_likeness"],
                        "counties_taken_from_the_canvass_summary": info["from_summary"],
                        "not_given": "Cities, county supervisor districts and school districts: Iowa's precinct results do not say which a precinct "
                                     "is in, and a precinct can lie in more than one. Townships and legislative districts for 2020: other precincts and "
                                     "another district plan."},
           "sources": source_records(tables, geo, official_docs), "places": places}
    return doc, failed


def our_places(db):
    """{kind: ids} of our Iowa places and legislative races, read only; None when the database is not there."""
    if not db or not os.path.exists(db):
        return None
    con = sqlite3.connect(pathlib.Path(os.path.abspath(db)).as_uri() + "?mode=ro", uri=True)
    try:
        out = collections.defaultdict(set)
        for kind, pid in con.execute("SELECT kind, id FROM sl_places WHERE source_id LIKE 'ia-%' AND kind IN ('county', 'mcd')"):
            out[kind].add(pid)
        for office, district in con.execute("SELECT office_kind, district FROM sl_races WHERE state = 'IA' AND office_kind IN ('state_senate', 'state_house')"):
            out["senate" if office == "state_senate" else "house"].add(str(district))
        return dict(out)
    except sqlite3.Error:
        return None
    finally:
        con.close()


def load(say=print, out=OUT, cache=CACHE, db=DB, refresh=False):
    say("    Iowa place votes: the Secretary of State's official precinct results, county file by county file")
    net.patient_lookups()
    os.makedirs(cache, exist_ok=True)
    tables = {}
    for year in sorted({c["year"] for c in CONTESTS}, reverse=True):
        tables[year] = results_table(year, CONTESTS, cache, refresh, say)
        t = tables[year]
        say(f"      {year}: {len(t['rows']):,} precincts in {t.get('files')} county files (copy of {t.get('fetched')})"
            + (f"; {words(t['pdf_only'])} posted only as PDF" if t.get("pdf_only") else ""))
    say("    the official figures: the State Canvass Summary, county by county")
    official, official_docs = official_totals(CONTESTS, cache, refresh, say)
    geo = township_precincts(cache, refresh, say)
    doc, failed = build(tables, official, official_docs, geo)

    p = doc["places"]
    say(f"    added up: {len(p['county'])} counties; {len(p['senate'])} Senate and {len(p['house'])} House districts; {len(p['mcd']):,} townships by "
        f"their precincts ({len(doc['coverage']['townships_left_out']):,} left out)")
    for c in doc["contests"]:
        ctl, s = doc["control"]["contests"][c["id"]], c["statewide"]
        cs = ctl.get("counties") or {}
        say(f"    control, {c['id']}: Democratic {s['dem']:,}, Republican {s['rep']:,}, other {s['other']:,}, total {s['total']:,}: "
            + ("equal to" if ctl["equal"] else "DIFFERS from") + f" {ctl['source']} ({ctl['read']})"
            + (f"; {cs['equal']} of {cs['compared']} counties equal" if cs else ""))
    for kind, result in doc["control"]["kinds"].items():
        if result != "equal":
            say(f"    control, {kind}: {result}")
    for kind in ("house", "senate"):
        if doc["kinds"][kind]["approximate"]:
            say(f"      {kind} districts marked approximate (a precinct shared with a neighbour): {', '.join(doc['kinds'][kind]['approximate'])}")
    if failed:
        for line in failed:
            say(f"    CHECK: {line}")
        raise Stop(f"    Iowa place votes: the control did not hold; {os.path.basename(out)} was not written. Nothing was changed to make it fit.")
    ours = our_places(db)
    if ours:
        covered = {}
        for kind in KIND_NAMES:
            ids = ours.get(kind, set())
            without = sorted(i for i in ids if not p[kind].get(i, {}).get("votes"))
            covered[kind] = {"places": len(ids), "with_votes": len(ids) - len(without), "without": without}
            say(f"      our {kind} places with votes: {len(ids) - len(without):,} of {len(ids):,}" + (f" (none for {', '.join(without[:12])})" if without else ""))
        doc["coverage"]["our_places"] = dict(covered, read=f"sl_places and sl_races in {os.path.basename(db)}, opened read-only on {M._now()}")
    M.write(out, doc)
    say(f"    Iowa place votes: {len(doc['contests'])} contests, {sum(len(v) for v in p.values()):,} places, control {doc['control']['result']}; "
        f"{os.path.relpath(out, HERE)} ({os.path.getsize(out) / 1e6:.1f} MB)")
    return doc


# ---------------------------------------------------------------- the arithmetic and the readers on made-up tables

def selftest(say=print):
    checks = []

    def check(what, got, want):
        checks.append(got == want)
        say(f"      {'ok  ' if got == want else 'FAIL'} {what}" + ("" if got == want else f": got {got!r}, expected {want!r}"))

    def row(county, precinct, house, d, r, o, hv=None, p="pres"):
        out = {"county": COUNTY_FIPS[county], "countyname": county, "precinct": precinct, "house": house,
               p + "_dem": d, p + "_rep": r, p + "_other": o, p + "_total": d + r + o}
        if hv:
            out["house_votes"] = hv
        return out
    contests = [{"id": "2024-president", "year": 2024, "date": "2024-11-05", "office": "President", "prefix": "pres", "title": "President",
                 "dem": "A/B", "rep": "C/D", "find": ("A", "C"), "official": "x", "columns": (0, 1)},
                {"id": "2020-president", "year": 2020, "date": "2020-11-03", "office": "President", "prefix": "pres", "title": "President",
                 "dem": "E/F", "rep": "G/H", "find": ("E", "G"), "official": "x", "columns": (1, 0)}]
    t24 = [row("Adair", "Alder Twp", "1", 20, 10, 1), row("Adair", "Birch/Cedar", "2", 30, 40, 2, {"2": 60, "3": 9}),         # Adair is not in LIKENESS_READ
           row("Adams", "Elm 1", "3", 5, 15, 0), row("Adams", "Fir", "4", 0, 12, 0)]
    t20 = [row("Adair", "Old 1", None, 11, 9, 0)]
    tables = {2024: {"rows": t24, "files": 2, "fetched": "-", "senate_checked": ["1"], "other_choices": {"pres": [1]}},
              2020: {"rows": t20, "files": 1, "fetched": "-", "pdf_only": ["Adams"]}}
    two = {c: [0, 0, 0, 0] for c in COUNTIES}
    official = {"2024-president": {"dem": 55, "rep": 77, "other": 3, "total": 135, "source": "x", "where": "-", "read": "-",
                                   "counties": dict(two, Adair=[50, 50, 3, 103], Adams=[5, 27, 0, 32])},
                "2020-president": {"dem": 18, "rep": 14, "other": 1, "total": 33, "source": "x", "where": "-", "read": "-",
                                   "counties": dict(two, Adair=[11, 9, 0, 20], Adams=[7, 5, 1, 13])}}
    geo = {"counties": {"19001": {"precincts": {"Alder Township": {"city": None, "cities": {}}, "Birch Cedar": {"city": None, "cities": {"Gum city": "all", "Oak city": "part"}}},
                                  "townships": {"90001": {"name": "Alder township", "precincts": ["Alder Township"]},
                                                "90002": {"name": "Birch township", "precincts": ["Birch Cedar"]},
                                                "90003": {"name": "Cedar township", "precincts": ["Birch Cedar", "Missing"]},
                                                "90004": {"name": "Town township", "precincts": []}}}}}
    # the made-up state has two counties with votes; the other 97 have none, so the per-county comparison is cut to the two
    for o in official.values():
        o["counties"] = {k: v for k, v in o["counties"].items() if k in ("Adair", "Adams")}
    fill = [row(c, "none", "1", 0, 0, 0) for c in COUNTIES if c not in ("Adair", "Adams")]
    tables[2024]["rows"] = t24 + fill
    tables[2020]["rows"] = t20 + [dict(r, house=None) for r in fill]
    doc, failed = build(tables, official, [], geo, contests, strict=False)
    p = doc["places"]
    check("the control holds on the made-up tables", (doc["control"]["result"], failed), ("equal", []))
    check("a county adds up its precincts", p["county"]["19001"]["votes"]["2024-president"], {"dem": 50, "rep": 50, "other": 3, "total": 103})
    check("a county with no precinct file is taken from the Canvass Summary, and the file says so", (p["county"]["19003"]["votes"]["2020-president"],
          doc["coverage"]["counties_taken_from_the_canvass_summary"]), ({"dem": 7, "rep": 5, "other": 1, "total": 13}, {"2020": ["Adams"]}))
    check("a precinct is filed in its House district, and the Senate district is two House districts", (p["house"]["3"]["votes"]["2024-president"]["total"],
          p["senate"]["2"]["votes"]["2024-president"]["total"], p["senate"]["2"]["house"]), (20, 32, ["3", "4"]))
    check("a precinct in two House districts is counted whole where most of it voted", p["house"]["2"]["votes"]["2024-president"]["total"], 72)
    check("and both districts say so", (p["house"]["2"]["approximate"]["about_votes"], p["house"]["3"]["approximate"]["about_votes"]), (9, 9))
    check("and so do the Senate districts they lie in", sorted(d for d, x in p["senate"].items() if x.get("approximate")), ["1", "2"])
    check("districts are not given for 2020", list(p["house"]["1"]["votes"]), ["2024-president"])
    check("a township that is one precinct is named with it", p["mcd"]["IA-M-90001"]["name"], "Alder township (the precinct Alder Twp)")
    check("a township in a shared precinct says who else votes there", p["mcd"]["IA-M-90002"]["name"],
          "The precinct Birch/Cedar, where Birch township votes with part of Cedar township, the city of Gum and part of the city of Oak")
    check("and carries that precinct's votes", p["mcd"]["IA-M-90002"]["votes"]["2024-president"], {"dem": 30, "rep": 40, "other": 2, "total": 72})
    check("a township whose precinct is not in the results by name is left out", ("IA-M-90003" in p["mcd"], "IA-M-90003" in doc["coverage"]["townships_left_out"]), (False, True))
    check("and so is one with no precinct outside a city", "IA-M-90004" in doc["coverage"]["townships_left_out"], True)
    check("a contest where every vote went the same way is marked too_few", p["house"]["4"].get("too_few"), ["2024-president"])
    wrong = dict(official, **{"2024-president": dict(official["2024-president"], dem=56, total=136)})
    check("a wrong official total is caught", build(tables, wrong, [], geo, contests, strict=False)[0]["control"]["result"], "differs")
    wrong = dict(official, **{"2024-president": dict(official["2024-president"], counties={"Adair": [49, 51, 3, 103], "Adams": [5, 27, 0, 32]})})
    check("a county that differs from its row of the Summary is caught", build(tables, wrong, [], geo, contests, strict=False)[1][0][:30], "2024-president, Adair County: ")
    sheet = [["United States Senator ( 1)", None], [None, None, "A One", None, None, "B Two", None, None, "Write-in", None, None, None],
             [None, "Registered Voters", "Election Day", "Absentee", "Total Votes", "Election Day", "Absentee", "Total Votes", "Election Day", "Absentee", "Total Votes", "Total"],
             ["P 1", 0.0, "3", "4", 7.0, "5", "1", 6.0, "0", "0", 0.0, "13"], ["P 2", 0.0, "1", "1", 2.0, "9", "1", 10.0, "1", "0", 1.0, "13"],
             ["Total:", "0", "4", "5", "9", "14", "2", "16", "1", "0", "1", "26"]]
    title, choices, precincts = read_sheet(sheet, "made up")
    check("a contest's sheet is read by its Total Votes columns", (title, [(n, dict(v)) for n, v in choices], precincts),
          ("United States Senator", [("A One", {"P 1": 7, "P 2": 2}), ("B Two", {"P 1": 6, "P 2": 10}), ("Write-in", {"P 1": 0, "P 2": 1})], ["P 1", "P 2"]))
    try:
        read_sheet(sheet[:-1] + [["Total:", "0", "4", "5", "10", "14", "2", "16", "1", "0", "1", "27"]], "made up")
        caught = False
    except Stop:
        caught = True
    check("a sheet whose precincts do not add up to its Total row stops the loader", caught, True)
    xml = ('<ElectionResult><Region>Adair</Region><VoterTurnout><Precincts><Precinct name="P 1"/></Precincts></VoterTurnout>'
           '<Contest text="State Representative District 7"><Choice text="X" party="DEM" totalVotes="5"><VoteType name="Election Day"><Precinct name="P 1" votes="2"/>'
           '</VoteType><VoteType name="Absentee"><Precinct name="P 1" votes="3"/></VoteType></Choice></Contest></ElectionResult>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("detail.xml", xml)
    found, _all = read_clarity(buf.getvalue(), lambda t: True)
    check("a county's 2024 report is read, vote types added together", (found["State Representative District 7"]["choices"][0][2]["P 1"],
          found["State Representative District 7"]["precincts"]), (5, ["P 1"]))
    lines = ["A One, B Two,", "Write-in Under Votes Over Votes Total", "REP DEM"]
    for i, c in enumerate(COUNTIES):
        lines += [f"{c} Election Day 1 2 0 1 0 4", "Absentee 1,000 0 1 0 0 1,001", "Total 1,001 2 1 1 0 1,005"] if i % 2 else \
                 ["Election", f"{c} 1 2 0 1 0 4", "Day", "Absentee 1,000 0 1 0 0 1,001", "Total 1,001 2 1 1 0 1,005"]
    lines += ["TOTAL Election Day 99 198 0 99 0 396", "Absentee 99,000 0 99 0 0 99,099", "Total 99,099 198 99 99 0 99,495", "Page 3 of 9"]
    got = parse_canvass(lines, "made up")
    check("the Canvass Summary's Total rows are read, in both of its layouts", (len(got["counties"]), got["counties"]["Black Hawk"], got["state"]),
          (99, [1001, 2, 1, 1, 0, 1005], [99099, 198, 99, 99, 0, 99495]))
    check("and split into the tickets, everyone else and the votes cast", canvass_split(got["state"], (1, 0)), ([198, 99099, 99, 99396], [99, 0, 99495]))
    check("county codes are the odd numbers in the alphabet's order", (COUNTY_FIPS["Adair"], COUNTY_FIPS["O'Brien"], COUNTY_FIPS["Polk"], COUNTY_FIPS["Wright"]),
          ("19001", "19141", "19153", "19197"))
    links = link_names(["074 Allen 1", "LeClaire Twp", "LeClaire City 1", "Barc/Benn/Lester/Dunkerton", "34 MOVILLE/ARLINGTON-WOLF CREEK", "MD"],
                       ["ALLEN 1", "(LCT) LeClaire Township", "(LC1) City of LeClaire", "Barclay/Benn/Lester/Dunkerton", "34 Moville -", "MD/NL.ANX", "NL07/MD.85"])
    check("precinct names are matched when written the same way", (links["074 Allen 1"], links["LeClaire Twp"]),
          (("ALLEN 1", "the same name"), ("(LCT) LeClaire Township", "the same name")))
    check("or by a clear likeness, and said so", (links["Barc/Benn/Lester/Dunkerton"][1], links["34 MOVILLE/ARLINGTON-WOLF CREEK"]),
          ("a like name", ("34 Moville -", "a like name")))
    check("or when one name begins the other, and only one does", links.get("MD"), ("MD/NL.ANX", "a like name"))
    check("and left unmatched when it is not clear", "LeClaire City 1" in link_names(["LeClaire City 1", "LeClaire City 2"], ["LC 1", "LC 2"]), False)
    check("outside the counties read one by one, only the same name counts", link_names(["Barc/Benn"], ["Barclay/Benn"], likeness=False), {})
    Shape, _lattice = _shapes()
    sq = lambda x0, y0, x1, y1: [[[x0, y0], [x1, y0], [x1, y1], [x0, y1], [x0, y0]]]                # noqa: E731
    g = place_townships([("19001", "Rural", sq(0, 0, 10, 10)), ("19001", "Town 1", sq(10, 0, 12, 2)), ("19001", "East", sq(10, 2, 20, 10) + sq(12, 0, 20, 2))],
                        [("19001", "90001", "West township", sq(0, 0, 10, 10)), ("19001", "90002", "East township", sq(10, 0, 20, 10))],
                        [("Town city", sq(10, 0, 12, 2)), ("Dot city", sq(4, 4, 5, 5))], t_grid=10)["19001"]
    check("a precinct inside a city is a city precinct", (g["precincts"]["Town 1"]["city"], g["precincts"]["Rural"]["city"]), ("Town city", None))
    check("a township's precincts are those outside the cities", (g["townships"]["90001"]["precincts"], g["townships"]["90002"]["precincts"]), (["Rural"], ["East"]))
    check("a small city inside a rural precinct is named with it", g["precincts"]["Rural"]["cities"], {"Dot city": "all"})
    say(f"    self-test: {sum(checks)} of {len(checks)} checks passed")
    return all(checks)


def main(argv=None):
    ap = argparse.ArgumentParser(description="How each Iowa place voted in past partisan general elections -> ballot/lean/ia_place_votes.json")
    ap.add_argument("--out", default=OUT, help="the file to write (default: ballot/lean/ia_place_votes.json)")
    ap.add_argument("--cache", default=CACHE, help="where downloads are kept (default: states_cache/ia_local)")
    ap.add_argument("--db", default=DB, help="the ballot database to count our places in, opened read-only (default: ballot_local_2026.sqlite)")
    ap.add_argument("--refresh", action="store_true", help="ask for every file and document again, even when the cached copies are there")
    ap.add_argument("--selftest", action="store_true", help="check the arithmetic and the readers on made-up tables; downloads and writes nothing")
    a = ap.parse_args(argv)
    if a.selftest:
        raise SystemExit(0 if selftest() else 1)
    load(out=a.out, cache=a.cache, db=a.db, refresh=a.refresh)


if __name__ == "__main__":
    main()
