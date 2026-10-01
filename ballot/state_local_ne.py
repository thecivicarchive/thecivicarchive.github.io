"""
ballot/state_local_ne.py - Nebraska's state races on the November 3, 2026 ballot, into ballot_local_2026.sqlite (never
ballot_2026.sqlite): the seats of the one-house Legislature up this year (the 24 even-numbered districts, and District
41 for a two-year term), Governor with Lieutenant Governor (one ticket, one vote), Secretary of State, State Treasurer,
Attorney General, Auditor of Public Accounts, the Public Service Commission seat, the State Board of Education seats and
the University of Nebraska Board of Regents seats on this year's list, with the May 12 primaries that chose the nominees.
And the local level, into the same database: the district boards and the judges' retention votes on the Secretary of
State's filing list, and the county, city, village, school and township contests of the counties whose own notice
or sample ballot can be read ("The local level", below).

Sources, the Secretary of State's own, the same two files the federal loader (ballot/lists/ne.py) reads:

  - Final Statewide General Candidate List, November 3, 2026 General Election (PDF, linked from sos.nebraska.gov/
    elections). A table printed on its side: each candidate is a strip across the page and the column headings sit in a
    strip of their own on the left. Only the bands under Office, District Name, Term, Vote For, Party, Candidate Name and
    Incumbency Status are ever turned into text; City of Residence, Mailing Address and Phone/Email stay inside the file
    and are never read, printed, logged or stored. The federal loader's cached copy is read when it is fresh; otherwise
    the list is fetched into memory and never written to disk whole. What was read (the allowed cells of the state
    offices' strips) is kept as a small JSON extract with the file's SHA-256. The list also carries local district
    boards (natural resources, public power, community colleges, educational service units and the like); the state
    part counts them and the local part, below, loads them from the same list as a workbook.
  - Official Report of the Board of State Canvassers, Primary Election, May 12, 2026 (the canvass book, PDF, linked as
    "Primary Election Official Results"). Names and vote counts by county only, so it is read from the federal
    loader's cached copy (fetched if missing).

What the record says, and what the loader does with it:
  - The Legislature is nonpartisan: candidates carry no party on either ballot, so party is "Nonpartisan office"
    (code N). The primary is one nonpartisan contest for all voters and the top two advance (the canvass book's own
    rule); a contest appears on the primary ballot however many filed, so every seat has official primary figures. A
    primary becomes a field (election "primary-NP") when two or more candidates were on it. The State Board of Education
    and the Board of Regents work the same way. The seats are stored with
    office_kind state_senate and race ids 2026-NE-SS<district> (the roster files them under the chamber "Legislature").
  - Governor, Secretary of State, Treasurer, Attorney General, Auditor and the Public Service Commission are partisan:
    each party's primary nominates its top vote-getter (primary-REP, primary-DEM, primary-LMN ...). The nominee for
    Governor then chose a running mate; the November list names each ticket, the candidate for Governor first.
  - Who advanced is the one the canvass book checks; every county column is added up and must equal the printed Total,
    and the checks must fall on the top vote-getters. A write-in candidate the report names is kept with write_in 1;
    "Write-In Scatterings" are not a candidate, but they are votes in the contest, so percentages are of every vote
    the report counts in it.
  - The list prints no ballot order. In partisan races candidates are numbered in the list's own order (by party, as
    the federal loader numbers them); nonpartisan candidates get no number, since the list's order is not a stated
    ballot position.
  - Today's holders come from state_ne.sqlite (the Open States roster): senators by district, and the officials table
    for Governor, Attorney General and Secretary of State. The roster does not carry the Treasurer, the Auditor, the
    Public Service Commission, the State Board of Education or the Regents; for those the list's own "Incumbent" mark
    is kept (incumbent 1, no roster id) and the race shows no holder. Only names, parties, districts and ids are read
    from the roster; its e-mail, phone and address columns never are.
  - A candidate is the sitting member (incumbent 1, state_member_id) only when the name fits the seat's holder one to
    one; the list's incumbency mark must agree. A candidate who holds another seat or office in the roster is given
    that id with incumbent 0 and a note, again only when exactly one roster person fits.

The local level (county, city, village, school, township and district boards), and the judges
---------------------------------------------------------------------------------------------
Nebraska has no one list of local candidates. Two kinds of source are read, and every local row says which it came from.

  1. The Secretary of State's Statewide Candidate Filing List, as a workbook (linked from the Information for
     Candidates page). Its General Election Candidates sheet holds, besides the state offices above, every candidate
     for the boards that file with the Secretary of State: natural resources districts, public power (and irrigation)
     districts, reclamation districts, community colleges, educational service units, the Learning Community of
     Douglas and Sarpy Counties and the Metropolitan Utilities District of Omaha. The workbook is fetched into memory
     and never saved whole: seven columns are read, each on its own by its heading (Office, District Name, Term, Vote
     For, Party, Candidate Name, Incumbency Status), and City of Residence, Mailing Address and Phone/Email are never
     turned into values. The board rows are checked against the printed list the state part reads (the same rows by
     two routes). A seat listed for a shorter term than its kind of board usually has is a special race ("-S").
     Its Judicial Retention sheet (no contact columns) names the judges who stand for a yes-or-no retention vote.
     They are stored under level court, as the other states' retention votes are, each district and county judge with
     the counties of the district as sections 24-301.02 and 24-503 of the statutes list them; every county file read
     below is checked to name exactly the judges those districts give it.
     The list names a board and a seat but no county. Where a seat is voted on comes from the Secretary of State's own
     "State Level Contests Primary 2026" workbook (sheet State Contests by County); for the boards that file only for
     November (the educational service units and most public power districts) it comes from the counties whose rows
     stand in the seat's table in the Board of State Canvassers' report of the last general election that had it
     (2024, else 2022, else 2020), and the race's note says so; a county's own file adds its county. Only that list
     of counties is read from those reports and kept: no name, no figure. A seat no list places is shown for the
     board's other counties, says so, and is a gap.
  2. County by county, the county election office's own Notice of Election (Neb. Rev. Stat. 32-802) or the
     "publication ballot" it prints as a sample ballot, for the counties in COUNTY_FILES (ten in this pass, largest
     first, each address checked by hand). These carry the county offices (on the partisan ballot), cities, villages,
     school boards, townships, airport authorities and the like. None of these files has a contact column; each layout
     has a reader of its own, and what a reader keeps is the office headings, the names and the parties. The district
     boards a county prints are matched to the filing list (board, seat, term and names): that is how a county is known
     to vote on a seat this year, and a seat a county prints with no name is kept as a contest with no candidate. The
     federal and state offices in a county's file are left to the state part; ballot questions are not read. A
     contest two counties print is one race, each candidate once. A county whose file is a scanned picture, or is not
     posted, or is laid out in a way no reader here knows, is recorded in sl_gaps with its election office's page
     (from the Secretary of State's directory of county election offices, of which only each county's name and
     Website link are read). If a county's site refuses a program, the file saved by hand in that county's folder of
     the local cache is read instead. The counties' candidate filing lists are never read: they carry addresses and
     telephone numbers.

Places: a city or village is given its Census Bureau place code when its name fits exactly one entry of the 2020
place file, a township its county subdivision code, a school district its Census LEA code when its name without its
number and kind words fits exactly one district; otherwise the county's three-digit code and the name. A district
board is one place under its own name, with every county any of its seats is voted in.

What the lists do not give: a ballot order (names rotate by precinct), write-in candidates, and withdrawn candidates
(a notice simply has no line for them). Douglas County's notice does not say how many are elected to a board.

Everything fetched for the local level is kept cut down in ballot_cache/ne/local/ (JSON: allowed cells only, with the
day fetched and the SHA-256 of the file as it came), so a second run downloads nothing for a week; the Census files
have no contact data and are kept whole. --refresh asks again.

Usage: python ballot/state_local_ne.py <database file> [--cache <folder>] [--keep <folder for the list extract>]
                                       [--local <folder for the local extracts>] [--refresh]
"""

import datetime as dt
import hashlib
import html as H
import io
import json
import os
import re
import sqlite3
import sys
import time
import unicodedata
import warnings
import zipfile
from collections import Counter, defaultdict
from urllib.error import HTTPError

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from ballot.check_local import EXTRA_SCHEMA, contact_like      # noqa: E402
from ballot.common import fold, name_parts, party_code          # noqa: E402
from ballot.match import fits                                   # noqa: E402
from ballot.pdftext import PDF, join, page_runs, rows as pdf_rows   # noqa: E402
from states import net                                          # noqa: E402

STATE, FIPS, NAME = "NE", "31", "Nebraska"
GENERAL, PRIMARY = "2026-11-03", "2026-05-12"
PAGE = "https://sos.nebraska.gov/elections"
CACHE = os.path.join(HERE, "ballot_cache")
ROSTER_DB = os.path.join(HERE, "state_ne.sqlite")
COUNTY_ZIP = os.path.join(HERE, "states_cache", "census", "cb_2024_us_county_500k.zip")
LIST_FILE = "ne_2026_general_candidate_filing_list.pdf"
BOOK_FILE = "ne_2026_primary_canvass_book.pdf"
EXTRACT_FILE = "sl_ne_general_list_extract.json"

SRC_LIST = "ne-sos-2026-sl-general-list"
SRC_BOOK = "ne-sos-2026-sl-primary-canvass"
SRC_ROSTER = "ne-openstates-roster-2026"
SRC_COUNTIES = "ne-census-cb-2024-county"

# the only bands of the list ever turned into text; City of Residence, Mailing Address and Phone/Email never are
LIST_KEEP = ("Office", "District Name (if applicable)", "Term", "Vote For", "Party (if applicable)", "Candidate Name",
             "Incumbency Status")
NONPARTISAN = "Nonpartisan office"
CODES = {"Republican": "REP", "Democratic": "DEM", "Libertarian": "LIB", "Legal Marijuana NOW": "LMN"}
CHECK = "✓"

# office -> (race key, level, office_kind, office shown, partisan, by district, roster office, usual term in years)
OFFICES = {
    "governor": ("GOV", "statewide", "governor", "Governor and Lieutenant Governor", 1, False, "governor", 4),
    "sos": ("SOS", "statewide", "secretary_of_state", "Secretary of State", 1, False, "secretary of state", 4),
    "treas": ("TREAS", "statewide", "state_treasurer", "State Treasurer", 1, False, None, 4),
    "ag": ("AG", "statewide", "attorney_general", "Attorney General", 1, False, "attorney general", 4),
    "aud": ("AUD", "statewide", "state_auditor", "Auditor of Public Accounts", 1, False, None, 4),
    "psc": ("PSC", "statewide", "public_service_commissioner", "Public Service Commissioner", 1, True, None, 6),
    "leg": ("SS", "legislature", "state_senate", "Member of the Legislature", 0, True, None, 4),
    "sboe": ("SBOE", "statewide", "state_board_of_education", "Member of the State Board of Education", 0, True, None, 4),
    "regent": ("REG", "statewide", "university_board", "Member of the University of Nebraska Board of Regents", 0, True, None, 6),
}
# the office as the list writes it, and as the canvass book heads it
LIST_OFFICE = {"For Governor and Lt. Governor": "governor", "For Secretary of State": "sos", "For State Treasurer": "treas",
               "For Attorney General": "ag", "For Auditor of Public Accounts": "aud",
               "For Public Service Commissioner": "psc", "For Member of the Legislature": "leg",
               "For Member of the State Board of Education": "sboe", "For University of Nebraska Board of Regents": "regent"}
BOOK_OFFICE = {"Governor": "governor", "Secretary of State": "sos", "State Treasurer": "treas", "Attorney General": "ag",
               "Auditor of Public Accounts": "aud", "Public Service Commissioner": "psc", "Member of the Legislature": "leg",
               "Member of the State Board of Education": "sboe",
               "Member of the Board of Regents of the University of Nebraska": "regent"}
FEDERAL = {"For United States Senator", "For Representative in Congress"}

TITLE = re.compile(r"^(?P<party>.+?)\s*Party\s*Nomination\s*(?:\((?P<contest>[^)]*)\))?\s*(?P<cont>[—–-]\s*continued)?$")
NONE = re.compile(r"^(?P<party>.+?)\s*Party\s*did not make a nomination", re.I)
DISTRICT = re.compile(r"^(?:Legislative\s*)?District\s*(?P<d>\d+)\s*(?:[—–-]\s*(?P<term>.+))?$")
NUMBER = re.compile(r"^\d{1,3}(?:,\d{3})*$")
TERM_WORDS = {"Two": 2, "Four": 4, "Six": 6}

SCHEMA = """
CREATE TABLE IF NOT EXISTS sl_races (race_id TEXT PRIMARY KEY, state TEXT NOT NULL, level TEXT NOT NULL, office_kind TEXT NOT NULL, office TEXT NOT NULL, jurisdiction TEXT, jurisdiction_id TEXT, county_ids TEXT, district TEXT, seat TEXT, special INTEGER NOT NULL DEFAULT 0, partisan INTEGER NOT NULL, holder_id TEXT, holder_name TEXT, holder_party TEXT, election_date TEXT NOT NULL, note TEXT);
CREATE TABLE IF NOT EXISTS sl_candidates (race_id TEXT NOT NULL, election TEXT NOT NULL, election_date TEXT, name TEXT NOT NULL, party TEXT, party_code TEXT, ballot_order INTEGER, incumbent INTEGER NOT NULL DEFAULT 0, write_in INTEGER NOT NULL DEFAULT 0, votes INTEGER, pct REAL, outcome TEXT, state_member_id TEXT, source_id TEXT, note TEXT, PRIMARY KEY (race_id, election, name));
CREATE TABLE IF NOT EXISTS sl_sources (source_id TEXT PRIMARY KEY, state TEXT, kind TEXT, agency TEXT, title TEXT, url TEXT, published TEXT, fetched TEXT, sha256 TEXT, rows INTEGER, note TEXT);
CREATE TABLE IF NOT EXISTS sl_places (kind TEXT NOT NULL, id TEXT NOT NULL, name TEXT NOT NULL, county_ids TEXT, source_id TEXT, PRIMARY KEY (kind, id));
"""


def fail(msg):
    raise SystemExit(f"Nebraska (state races): {msg}")


def race_id(office, district=None):
    key = OFFICES[office][0]
    return f"2026-{STATE}-{key}{district or ''}"


def cells(runs):
    """A printed row's pieces grouped into table cells, [(x0, runs, text, x1)]: a gap wider than six points starts the
    next cell (the federal loader's rule)."""
    out = []
    for r in sorted(runs, key=lambda r: r[0]):
        if out and r[0] - out[-1][2] <= 6:
            out[-1][1].append(r)
            out[-1][2] = max(out[-1][2], r[4])
        else:
            out.append([r[0], [r], r[4]])
    return [(x, rs, join(rs), end) for x, rs, end in out]


# ---------- the November list: allowed bands only ----------

def read_list(data):
    """(printed date, [record]) for every candidate strip of the list; a record holds only the LIST_KEEP bands. The
    City of Residence, Mailing Address and Phone/Email bands are never joined into text."""
    pdf = PDF(data)
    out, printed, titled = [], "", False
    for page, res in pdf.pages():
        runs = page_runs(pdf, page, res)
        heads = {}
        for x0, y0, _s, t, _x1 in runs:
            if 55 <= x0 < 76 and t.strip():
                heads.setdefault(round(y0), []).append((x0, t.strip()))
            elif x0 < 55 and re.fullmatch(r"\d{1,2}/\d{1,2}/20\d\d", t.strip()) and not printed:
                mo, d, y = t.strip().split("/")
                printed = f"{y}-{int(mo):02d}-{int(d):02d}"
            elif x0 < 55 and "November 3, 2026 General Election" in t:
                titled = True
        fields = sorted((y, " ".join(t for _x, t in sorted(v))) for y, v in heads.items())
        names = [n for _y, n in fields]
        if not all(k in names for k in LIST_KEEP):
            fail("the candidate list's column headings changed")
        bands = [(y - 3, (fields[i + 1][0] - 3) if i + 1 < len(fields) else 1e9, n) for i, (y, n) in enumerate(fields)]
        band = lambda y: next((n for lo, hi, n in bands if lo <= y < hi), None)
        starts = sorted({round(x0, 1) for x0, y0, _s, _t, _x1 in runs if x0 >= 76 and band(y0) == "Vote For"})
        for i, sx in enumerate(starts):
            ex = starts[i + 1] if i + 1 < len(starts) else sx + 21
            got = {}
            for x0, y0, _s, t, _x1 in sorted(runs, key=lambda r: (round(r[0], 1), r[1])):
                if sx - 0.5 <= x0 < ex - 0.5:
                    k = band(y0)
                    if k in LIST_KEEP:
                        got.setdefault(k, []).append(t.strip())
            out.append({k: re.sub(r"\s+", " ", " ".join(v)).strip() for k, v in got.items()})
    if not titled:
        fail("the candidate list is no longer the November 3, 2026 General Election")
    return printed, out


def list_links(say):
    """The list's and the canvass book's addresses from the Elections page (the federal loader's own lookup)."""
    from ballot.lists import ne as fed
    for attempt in range(3):
        try:
            return fed.links()
        except (Exception, SystemExit) as e:
            if attempt == 2:
                say(f"    Nebraska (state races): could not read {PAGE} ({e}); the Elections page is given as the address")
                return {}
            time.sleep(3)


def get_list(cache, keep, url, max_age_days, say):
    """The November list's state records: from the federal loader's cached copy when fresh, else fetched into memory
    (never saved whole), else from the saved extract. Returns (printed, all strips counted, state records, local strips,
    sha256, fetched, how)."""
    cached = os.path.join(cache, "ne", LIST_FILE)
    extract = os.path.join(keep, EXTRACT_FILE)
    data = fetched = how = None
    if os.path.exists(cached) and time.time() - os.path.getmtime(cached) < max_age_days * 86400:
        data = open(cached, "rb").read()
        fetched, how = dt.date.fromtimestamp(os.path.getmtime(cached)).isoformat(), "the federal loader's cached copy"
    elif url:
        try:
            data = net.get(url)
            fetched, how = dt.date.today().isoformat(), "fetched into memory, not saved"
        except OSError as e:
            say(f"    Nebraska (state races): the list could not be fetched ({e})")
    if data is None and os.path.exists(cached):
        data = open(cached, "rb").read()
        fetched, how = dt.date.fromtimestamp(os.path.getmtime(cached)).isoformat(), "the federal loader's cached copy (older)"
    if data is None:
        if not os.path.exists(extract):
            fail("the November candidate list could not be read, and no extract of it is kept")
        ex = json.load(open(extract, encoding="utf-8"))
        say(f"    Nebraska (state races): using the saved extract of the list ({ex['fetched']})")
        return ex["printed"], ex["strips"], ex["rows"], ex["local"], ex["sha256"], ex["fetched"], "the saved extract"
    printed, recs = read_list(data)
    sha = hashlib.sha256(data).hexdigest()
    del data
    state, local, unknown = [], 0, []
    for r in recs:
        office = r.get("Office", "")
        if office in LIST_OFFICE:
            state.append({k: r.get(k, "") for k in LIST_KEEP})
        elif office in FEDERAL:
            continue
        elif office.startswith("For "):
            unknown.append(office)
        else:
            local += 1                  # a local district board: left for the local pages
    if unknown:
        fail(f"offices on the list the loader does not know: {sorted(set(unknown))}")
    os.makedirs(keep, exist_ok=True)
    with open(extract, "w", encoding="utf-8") as fh:       # the allowed cells of the state strips only
        json.dump({"url": url or PAGE, "sha256": sha, "fetched": fetched, "printed": printed, "strips": len(recs),
                   "local": local, "columns": list(LIST_KEEP), "rows": state}, fh, ensure_ascii=False, indent=0)
    return printed, len(recs), state, local, sha, fetched, how


# ---------- the canvass book ----------

def book_names(heads, first, lo, hi, where):
    """[(name, checked, kind)] for a table's heading cells; given names sit on the line above where they wrap. kind is
    'cand', 'writein' (a write-in candidate the report names) or 'scatter' (Write-In Scatterings)."""
    pool = [f for f in (first or []) if lo <= f[0] < hi]
    used, out = set(), []
    for c in heads:
        given = ""
        over = [(min(c[3], f[3]) - max(c[0], f[0]), i) for i, f in enumerate(pool)]
        over = [o for o in over if o[0] > 0]
        if over:
            i = max(over)[1]
            if i in used:
                fail(f"a given name in the canvass book sits over two candidates ({where})")
            used.add(i)
            given = pool[i][2]
        family = c[2]
        checked = family.endswith(CHECK)
        family = family.rstrip(CHECK).strip()
        if family == "Scatterings":
            out.append(("Write-in scatterings", False, "scatter"))
        elif family == "(Write-In)":
            out.append((given, checked, "writein"))
        else:
            name = given + family if given.endswith("-") else f"{given} {family}"
            out.append((re.sub(r"\s+", " ", name).strip(), checked, "cand"))
    if len(used) != len(pool):
        fail(f"a given name in the canvass book belongs to no candidate ({where})")
    return out


def canvass(path):
    """({(race, party or None): table}, {race: term words}, {race: set of county names}) for the state offices in the
    canvass book. A table is {"names": [(name, checked, kind)], "total": [...], "sum": [...]}."""
    pdf = PDF(open(path, "rb").read())
    tables, terms, places = {}, {}, defaultdict(set)
    started, office, ctx_district = False, None, None
    for n, (page, res) in enumerate(pdf.pages(), start=1):
        rows = [(y, cells(runs)) for y, runs in pdf_rows(pdf, page, res)]
        texts = [" ".join(c[2] for c in cs) for _y, cs in rows]
        if not started:
            started = any(t == "Statewide Constitutional Offices" for t in texts)
            if not started:
                continue
        if any(t.startswith("Member of Board of Governors") for t in texts):
            break                                   # community colleges and after: local offices
        if any("Auto-advance Rules" in t for t in texts):
            continue                                # the pages that explain each office's rules
        cols = []
        for (y, cs), text in zip(rows, texts):
            where = f"page {n}"
            if re.search(r"Page \| ?\d+$", text):
                continue
            if len(cs) == 1 and text in BOOK_OFFICE:
                office, ctx_district, cols = BOOK_OFFICE[text], None, []
                continue
            kinds = []
            for c in cs:
                t = c[2]
                if NONE.match(t):
                    kinds.append(("none", None))
                elif TITLE.match(t):
                    kinds.append(("title", TITLE.match(t)))
                elif DISTRICT.match(t):
                    kinds.append(("district", DISTRICT.match(t)))
                else:
                    kinds.append((None, None))
            if any(k for k, _m in kinds):
                if not all(k for k, _m in kinds):
                    fail(f"a heading row in the canvass book could not be read ({where}: {text!r})")
                if any(k == "none" for k, _m in kinds):
                    cols = []
                    continue
                new = []
                for i, ((kind, m), c) in enumerate(zip(kinds, cs)):
                    lo = (cs[i - 1][3] + c[0]) / 2 if i else -1e9
                    hi = (c[3] + cs[i + 1][0]) / 2 if i + 1 < len(cs) else 1e9
                    if kind == "district":
                        if office is None:
                            fail(f"a district heading before any office ({where})")
                        d = str(int(m.group("d")))
                        rid = race_id(office, d)
                        if m.group("term"):
                            terms[rid] = m.group("term").strip()
                        if OFFICES[office][4]:                  # partisan: the party tables below carry the district
                            if len(cs) != 1:
                                fail(f"two district headings side by side for a partisan office ({where})")
                            ctx_district = d
                            continue
                        new.append({"lo": lo, "hi": hi, "key": (rid, None), "subs": None, "first": None})
                    else:
                        party = m.group("party").strip()
                        if m.group("contest"):
                            contest = m.group("contest").strip()
                            if contest not in BOOK_OFFICE:
                                fail(f"an unknown contest in the canvass book ({contest!r}, {where})")
                            office, ctx_district = BOOK_OFFICE[contest], None
                        if office is None or not OFFICES[office][4]:
                            fail(f"a party's table for an office that is not partisan ({where}: {text!r})")
                        if OFFICES[office][5] and not ctx_district:
                            fail(f"a party's table with no district heading above it ({where}: {text!r})")
                        new.append({"lo": lo, "hi": hi, "key": (race_id(office, ctx_district if OFFICES[office][5] else None), party),
                                    "subs": None, "first": None})
                cols = new
                continue
            for col in cols:
                part = [c for c in cs if col["lo"] <= c[0] < col["hi"]]
                if not part:
                    continue
                key = col["key"]
                labels = [c for c in part if c[2] == "County"]
                if labels:
                    xs = [c[0] for c in labels]
                    subs = []
                    for i, x in enumerate(xs):
                        lo, hi = x - 3, (xs[i + 1] - 3) if i + 1 < len(xs) else col["hi"]
                        heads = [c for c in part if lo <= c[0] < hi and c[2] != "County"]
                        names = book_names(heads, col["first"], lo, hi, f"{where}, {key[0]}")
                        subs.append((lo, hi, names))
                        t = tables.setdefault(key, {"names": names, "total": None, "sum": [0] * len(names)})
                        if t["names"] != names:
                            fail(f"the canvass table for {key} names different candidates in different places")
                    col["subs"], col["first"] = subs, None
                    continue
                if col["subs"] is None:
                    if any(NUMBER.match(c[2]) for c in part):
                        fail(f"figures before a table's headings in the canvass book ({where}, {key[0]})")
                    col["first"] = part             # given names, on the line above the family names
                    continue
                for lo, hi, names in col["subs"]:
                    sub = [c for c in part if lo <= c[0] < hi]
                    if not sub:
                        continue
                    label, nums = sub[0][2], [c[2] for c in sub[1:]]
                    if NUMBER.match(label) or len(nums) != len(names) or not all(NUMBER.match(v) for v in nums):
                        fail(f"a canvass row for {key} does not line up with its candidates ({where}, {len(nums)} figures)")
                    vals = [int(v.replace(",", "")) for v in nums]
                    t = tables[key]
                    if label == "Total":
                        if t["total"] is not None:
                            fail(f"two Total rows for {key}")
                        t["total"] = vals
                    else:
                        t["sum"] = [a + b for a, b in zip(t["sum"], vals)]
                        places[key[0]].add(label)
    if not tables:
        fail("no state tables were found in the canvass book")
    for key, t in tables.items():
        if t["total"] is None:
            fail(f"no Total row for {key}")
        if t["total"] != t["sum"]:
            fail(f"the counties for {key} add up to {t['sum']}, the report's Total is {t['total']}")
        cand = [(v, i) for i, ((_n, _c, kind), v) in enumerate(zip(t["names"], t["total"])) if kind != "scatter"]
        marks = {i for i, (_n, c, _k) in enumerate(t["names"]) if c}
        k = 1 if key[1] else 2                      # a party nominates one; a nonpartisan primary advances two
        want = min(k, len(cand))
        top = sorted(cand, reverse=True)
        if len(marks) != want or any(v < top[want - 1][0] for v, i in cand if i in marks):
            fail(f"the check marks for {key} are not on the top {want} vote-getter(s)")
    return tables, terms, places


# ---------- the roster: names, parties, districts and ids only ----------

def roster():
    con = sqlite3.connect(ROSTER_DB)
    seats = defaultdict(list)
    for bid, first, last, full, party, district in con.execute(
            "SELECT bioguide_id, first_name, last_name, official_full, party_name, district FROM legislators "
            "WHERE is_current = 1 AND chamber = 'Legislature'"):
        seats[str(district)].append({"id": bid, "first": first or "", "last": last or "", "full": full or f"{first} {last}",
                                     "party": party, "district": str(district), "label": None})
    offices = {}
    for bid, first, last, full, office, label, party in con.execute(
            "SELECT bioguide_id, first_name, last_name, official_full, office, office_label, party_name FROM officials"):
        offices[office] = {"id": bid, "first": first or "", "last": last or "", "full": full or f"{first} {last}",
                           "party": party, "district": None, "label": label}
    as_of = (con.execute("SELECT MAX(updated_at) FROM legislators").fetchone()[0] or "")[:10]
    con.close()
    return seats, offices, as_of


def person_fits(name, p):
    cand = name_parts(name)
    return fits(cand, (fold(p["first"]).split(), fold(p["last"]))) or fits(cand, name_parts(p["full"]))


def same_person(a, b):
    return fold(a) == fold(b) or fits(name_parts(a), name_parts(b))


def ticket_head(name):
    """'Jim Pillen and Joe Kelly' -> 'Jim Pillen' (the list names the candidate for Governor first)."""
    return re.split(r"\s+(?:&|and)\s+", name, maxsplit=1)[0]


def counties(path=COUNTY_ZIP):
    """{folded county name: (GEOID, "Adams County")} for Nebraska from the Census Bureau's cartographic county file."""
    import shapefile                                   # pyshp
    z = zipfile.ZipFile(path)
    base = next(n[:-4] for n in z.namelist() if n.endswith(".dbf"))
    rdr = shapefile.Reader(dbf=io.BytesIO(z.read(base + ".dbf")))
    fields = [f[0] for f in rdr.fields[1:]]
    out = {}
    for rec in rdr.iterRecords():
        rec = dict(zip(fields, rec))
        if str(rec.get("STATEFP")) == FIPS:
            out[fold(str(rec["NAME"]))] = (str(rec["GEOID"]), str(rec["NAMELSAD"]))
    if len(out) != 93:
        fail(f"the county file gives {len(out)} Nebraska counties, not 93")
    return out


def sha_of(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest() if path and os.path.exists(path) else ""


def mdate(path):
    return dt.date.fromtimestamp(os.path.getmtime(path)).isoformat() if os.path.exists(path) else ""


# ================================================================================================================
# The local level (see the docstring): the district boards and the judges on the Secretary of State's filing list,
# and the county, city, village, school and township contests of the counties whose own notice or sample ballot is
# a text file a program can read.
# ================================================================================================================

LOCAL_FRESH_DAYS = 7                                    # the filing list and the counties' files are asked for again after a week
WORKBOOK_URL = "https://sos.nebraska.gov/sites/default/files/doc/elections/2026/Statewide_Candidate_Filing_List.xlsx"
BY_COUNTY_URL = "https://sos.nebraska.gov/sites/default/files/doc/elections/2026/State_Level_Contests_PR26.xlsx"
CANVASS_BOOKS = (       # the last three general elections' official reports, newest first
    ("2024", "November 5, 2024", "https://sos.nebraska.gov/sites/default/files/doc/elections/2024/2024%20General%20Canvass%20Book.pdf"),
    ("2022", "November 8, 2022", "https://sos.nebraska.gov/sites/default/files/doc/elections/2022/2022%20General%20Canvass%20Book.pdf"),
    ("2020", "November 3, 2020", "https://sos.nebraska.gov/sites/default/files/doc/elections/2020/2020-General-Canvass-Book.pdf"))
STATUTE_URL = "https://nebraskalegislature.gov/laws/statutes.php?statute={}"
STATUTES_PAGE = "https://nebraskalegislature.gov/laws/browse-chapters.php?chapter=32"
DIRECTORY_URL = "https://sos.nebraska.gov/election-officials-contact-information"
PLACE_URL = "https://www2.census.gov/geo/docs/reference/codes2020/place/st31_ne_place2020.txt"
COUSUB_URL = "https://www2.census.gov/geo/docs/reference/codes2020/cousub/st31_ne_cousub2020.txt"
UNSD_URL = "https://www2.census.gov/geo/tiger/GENZ2024/shp/cb_2024_31_unsd_500k.zip"
PLACE_HEAD = "STATE|STATEFP|PLACEFP|PLACENS|PLACENAME|TYPE|CLASSFP|FUNCSTAT|COUNTIES"
COUSUB_HEAD = "STATE|STATEFP|COUNTYFP|COUNTYNAME|COUSUBFP|COUSUBNS|COUSUBNAME|CLASSFP|FUNCSTAT"

SRC_WB, SRC_BYCO, SRC_DIR = "ne-sos-2026-local-filing-list", "ne-sos-2026-primary-contests-by-county", "ne-sos-county-election-offices"
SRC_PLACE, SRC_COUSUB, SRC_UNSD = "ne-census-2020-places", "ne-census-2020-county-subdivisions", "ne-census-cb-2024-unsd"

# the only columns of the filing list ever turned into values; City of Residence, Mailing Address and Phone/Email never are
WB_KEEP = ("Office", "District Name (if applicable)", "Term", "Vote For", "Party (if applicable)", "Candidate Name", "Incumbency Status")
RET_KEEP = ("Office", "District (if applicable)", "Term", "Judge (Ballot Name)")
BYCO_KEEP = ("County", "Contest", "District")

# The ten counties read in this pass, largest first: each county election office's own posting for November 3, 2026,
# found on its page in the Secretary of State's directory of county election offices and checked by hand on
# 2026-10-01. (fips, county, reader, kind of file, title, address, the file's name if it has to be saved by hand)
COUNTY_FILES = (
    ("31055", "Douglas", "douglas", "notice of election", "Notice of General Election: offices and candidates appearing on the general ballot",
     "https://www.votedouglascounty-ne.gov/elections/2026/General/GN26Candidates.pdf", "GN26Candidates.pdf"),
    ("31109", "Lancaster", "lancaster", "notice of election", "Notice of General Election 2026",
     "https://www.lancaster.ne.gov/DocumentCenter/View/33610", "Notice-of-Election.pdf"),
    ("31153", "Sarpy", "sarpy", "notice of election", "Notice of General Election",
     "https://www.sarpy.gov/DocumentCenter/View/9683", "Notice-of-Election.pdf"),
    ("31019", "Buffalo", "ess", "sample ballot", "Sample Ballot, General Election, November 3, 2026 (publication ballot)",
     "https://buffalocounty.ne.gov/Portals/0/adam/Content/nTqk3h2tJUCIAb7MwVtHkA/Ballot/26GNEBUF_3_PUB.pdf", "26GNEBUF_3_PUB.pdf"),
    ("31119", "Madison", "ess", "sample ballot", "Sample Ballot, General Election, November 3, 2026 (publication ballot)",
     "https://madisoncountyne.gov/wp-content/uploads/2026/09/GN26-NDN-Sample-Ballot-Publication.pdf", "GN26-NDN-Sample-Ballot-Publication.pdf"),
    ("31001", "Adams", "ess", "sample ballot", "Sample Ballot, General Election, November 3, 2026 (publication ballot)",
     "https://adamscountyne.gov/images/PDFS/Clerks/Sample_Ballot.pdf", "Sample_Ballot.pdf"),
    ("31177", "Washington", "washington", "notice of election", "Notice of General Election",
     "https://www.washingtoncountyne.gov/_files/ugd/ac3a10_b59f7937249b4f4c9ec781f90dece45e.pdf", "Notice-of-General-Election.pdf"),
    ("31185", "York", "york", "notice of election", "Notice of General Election",
     "https://www.yorkcounty.ne.gov/uploads/1/4/5/5/145507130/york_county_42_day_election_notice_gn26.docx.pdf", "york_county_42_day_election_notice_gn26.docx.pdf"),
    ("31137", "Phelps", "ess", "sample ballot", "Sample Ballot, General Election, November 3, 2026 (publication ballot)",
     "https://phelpscounty.ne.gov/pdfs/election/2026/General-26%20publication%20ballot.pdf", "General-26 publication ballot.pdf"),
    ("31031", "Cherry", "ess", "sample ballot", "Sample Ballot, General Election, November 3, 2026 (publication ballot)",
     "https://cherrycountyne.gov/wp-content/uploads/sites/51/2026/09/26GNECHE_PUB.pdf", "26GNECHE_PUB.pdf"),
)
# What a hand check on 2026-10-01 found at five other election offices, largest first; said in those counties' gaps
COUNTY_SEEN = {
    "31079": "When this was loaded the county election office had posted its notice for the May primary and a list of filings, and no notice or sample "
             "ballot for November 3.",
    "31053": "When this was loaded the county's Notice of Election was posted as a scanned picture, which a program cannot read as text.",
    "31141": "When this was loaded the county's Notice of Election was posted as a scanned picture, which a program cannot read as text.",
    "31081": "When this was loaded the county's notice and sample ballot were posted as scanned pictures, which a program cannot read as text.",
    "31167": "When this was loaded the county's publication ballot was posted as a scanned picture, which a program cannot read as text.",
}

PARTIES = ("Republican", "Democratic", "Democrat", "Libertarian", "Legal Marijuana NOW", "Nebraska Working People", "America First", "By Petition")
PARTY_WORDS = {"Democrat": "Democratic"}                # the words the Secretary of State's list uses
VOTE = re.compile(r"^Vote for (?:up to )?(ONE|TWO|THREE|FOUR|FIVE|SIX|SEVEN|EIGHT|NINE|TEN)\b", re.I)
NUM = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10}
SHALL = re.compile(r"^Shall Judge (?P<judge>.+?) be retained in office\?")
NOBODY = re.compile(r"^(?:no candidates? (?:have |has )?filed\b.*|no filings?\b.*|vacan(?:t|cy))\.?$", re.I)      # a line that says nobody filed
NAMELIKE = re.compile(r"^[A-Za-zÀ-ÿ][A-Za-zÀ-ÿ.,'’\"() -]*$")
OFFICE_WORDS = re.compile(r"\b(Board|Council|Trustees|District|Ticket|Court|Vote|Term|School|Schools|County|Village|City|Township|Authority|Mayor|"
                          r"Member|Election|Ballot|Shall)\b")
# the page builder's own last check is a little wider than the trial check's (it also stops at Court and Place)
BUILDER_STREET = re.compile(r"\b\d{1,6}\s+(?:[NSEW]\.?\s+)?[A-Za-z0-9.' -]{1,40}?\s(?:St|Street|Ave|Avenue|Rd|Road|Blvd|Boulevard|Dr|Drive|Ln|Lane|Way|Ct|Court|"
                            r"Cir|Circle|Pkwy|Parkway|Hwy|Highway|Trl|Trail|Pl|Place|Ter|Terrace)\b\.?", re.I)


class Layout(Exception):
    """A file whose layout is not the one its reader was written for. The message names the page and the check that
    failed, never the line itself."""


def stop(msg):
    """The local part stops before anything is written; the message names the file, the row number and the check, never the row."""
    raise SystemExit(f"Nebraska (local): {msg}")


def slug(text):
    """Letters, digits and hyphens, for ids: "Papio-Missouri River Natural Resources District" -> papio-missouri-river-..."""
    t = unicodedata.normalize("NFKD", text or "")
    return re.sub(r"[^a-z0-9]+", "-", "".join(ch for ch in t if not unicodedata.combining(ch)).lower()).strip("-")


def nkey(name):
    """A name for comparing one printing with another: letters only, one space between words ("Casey J. Foster" and
    "Casey J Foster" are one person)."""
    return " ".join(fold(name).split())


def and_names(items):
    items = list(items)
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]


def read_json(path):
    try:
        with open(path, encoding="utf-8") as fh:
            return json.load(fh)
    except (OSError, ValueError):
        return None


def keep_json(path, obj):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path + ".part", "w", encoding="utf-8") as fh:
        json.dump(obj, fh, ensure_ascii=False, indent=0)
    os.replace(path + ".part", path)


def fetch(url):
    """One file into memory, a second after the last request. A site that answers 403, 429 or 503 is asked twice more,
    a quarter of a minute apart, and then left alone."""
    for attempt in range(3):
        time.sleep(1.0)
        try:
            return net.get(url)
        except HTTPError as e:
            if e.code in (403, 429, 503) and attempt < 2:
                time.sleep(15)
                continue
            raise


def extract(path, build, say, what, max_age=None, refresh=False):
    """A cut-down copy of a source, kept as JSON: the copy on disk while it is fresh (for good when max_age is None),
    else built afresh by build(); when the source cannot be reached and an older copy exists, the older copy, with a
    line saying so. Only the cut-down copy is ever written."""
    old = read_json(path)
    if old and not refresh and (max_age is None or (dt.date.today() - dt.date.fromisoformat(old["fetched"])).days < max_age):
        return old
    try:
        new = build()
    except (SystemExit, Layout):
        raise
    except Exception as err:  # noqa: BLE001  unreachable: fall back to the copy on disk, and say so
        if old:
            say(f"      Nebraska (local): {what} could not be fetched ({type(err).__name__}); using the copy of {old['fetched']} on disk")
            return old
        raise
    keep_json(path, new)
    return new


# ---------------------------------------------------------------- the Secretary of State's filing list (a workbook)

def column_cells(ws, wanted, where):
    """[row as {heading: text}] for the wanted headings of one sheet. Each wanted column is read on its own, so that no
    cell of any other column (a city of residence, a mailing address, a telephone number) is ever turned into a value."""
    head = [str(c).strip() if c is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    cols = {}
    for h in wanted:
        if head.count(h) != 1:
            stop(f"the {where} sheet no longer has exactly one column headed \"{h}\"")
        i = head.index(h) + 1
        vals = []
        for (v,) in ws.iter_rows(min_row=2, min_col=i, max_col=i, values_only=True):
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            vals.append(re.sub(r"\s+", " ", str(v)).strip() if v is not None else "")
        cols[h] = vals
    if len({len(v) for v in cols.values()}) != 1:
        stop(f"the {where} sheet's columns are not all the same length")
    rows = [dict(zip(wanted, vals)) for vals in zip(*(cols[h] for h in wanted))]
    return [r for r in rows if any(r.values())]


def filing_list(folder, say, refresh=False):
    """The filing list's General Election Candidates and Judicial Retention sheets, cut down in memory to the allowed
    columns and kept as JSON with the day fetched and the fingerprint of the workbook as it came."""
    def build():
        import openpyxl
        raw = fetch(WORKBOOK_URL)
        if not raw.startswith(b"PK"):
            raise ValueError("the address did not return a workbook")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            for name in ("General Election Candidates", "Judicial Retention"):
                if name not in wb.sheetnames:
                    stop(f"the filing list no longer has a sheet named \"{name}\"")
            gen = column_cells(wb["General Election Candidates"], WB_KEEP, "General Election Candidates")
            ret = column_cells(wb["Judicial Retention"], RET_KEEP, "Judicial Retention")
        blanked = 0
        for r in gen + ret:                                 # a kept cell that reads like contact details is blanked and counted
            for k, v in r.items():
                if v and contact_like(v, True):
                    r[k], blanked = "", blanked + 1
        return {"title": "Statewide Candidate Filing List (workbook)", "url": WORKBOOK_URL, "fetched": dt.date.today().isoformat(),
                "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw), "columns": list(WB_KEEP), "retention_columns": list(RET_KEEP),
                "blanked": blanked, "rows": gen, "retention": ret}
    return extract(os.path.join(folder, "sos_statewide_candidate_filing_list.json"), build, say, "the Secretary of State's filing list",
                   LOCAL_FRESH_DAYS, refresh)


def contests_by_county(folder, say):
    """The Secretary of State's "State Level Contests Primary 2026" workbook, sheet "State Contests by County": which
    state-filed contests were on which county's ballot this year. It has no contact columns; three are read."""
    def build():
        import openpyxl
        raw = fetch(BY_COUNTY_URL)
        if not raw.startswith(b"PK"):
            raise ValueError("the address did not return a workbook")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            name = next((n for n in wb.sheetnames if n.startswith("State Contests by County")), None)
            if not name:
                stop("the State Level Contests workbook no longer has a \"State Contests by County\" sheet")
            rows = column_cells(wb[name], BYCO_KEEP, name)
        return {"title": "State Level Contests Primary 2026 (workbook), sheet \"" + name + "\"", "url": BY_COUNTY_URL, "fetched": dt.date.today().isoformat(),
                "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw), "columns": list(BYCO_KEEP), "rows": rows}
    return extract(os.path.join(folder, "sos_state_contests_by_county_pr26.json"), build, say, "the State Level Contests workbook")


# ---------------------------------------------------------------- naming a board and a seat the same way everywhere

WORD_NUM = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10, "eleven": 11, "twelve": 12,
            "thirteen": 13, "fourteen": 14, "fifteen": 15, "sixteen": 16, "seventeen": 17, "eighteen": 18, "nineteen": 19, "twenty": 20, "thirty": 30,
            "forty": 40}
_UNITS = "One|Two|Three|Four|Five|Six|Seven|Eight|Nine"
NUMBERED = re.compile(r"\b(District|Subdistrict|Subdivision|Subcouncil|Ward|Number|Township|Precinct)\s+((?:Twenty|Thirty|Forty)(?:[- ](?:" + _UNITS + r"))?|"
                      + _UNITS + r"|Ten|Eleven|Twelve|Thirteen|Fourteen|Fifteen|Sixteen|Seventeen|Eighteen|Nineteen)\b", re.I)
YEARS = re.compile(r"\b(\d+|Two|Four|Six)[- ]Year[- ]Term\b", re.I)
BOARD_CLASSES = (("nrd", ("natural", "resources", "district")), ("ppid", ("public", "power", "and", "irrigation", "district")),
                 ("ppd", ("public", "power")), ("mud", ("metropolitan", "utilities", "district")), ("recl", ("reclamation", "district")),
                 ("cc", ("community", "college")), ("esu", ("educational", "service", "unit")), ("lc", ("learning", "community")))


def number_words(text):
    """"District Two", "Ward Thirty-Six" -> "District 2", "Ward 36": only a number written after a word that numbers a seat."""
    def conv(m):
        return f"{m.group(1)} {sum(WORD_NUM[w] for w in re.split(r'[- ]', m.group(2).lower()))}"
    return NUMBERED.sub(conv, text or "")


def plain(text):
    """One spelling for comparing headings: dashes alike, words that ran together parted, numbers as digits."""
    t = re.sub(r"[–—]", "-", text or "").replace("’", "'")
    t = re.sub(r"([a-z])(Resources|District)\b", r"\1 \2", t)
    return re.sub(r"\s+", " ", number_words(t)).strip()


def tokens(text):
    return re.findall(r"[a-z0-9]+", plain(text).lower().replace("&", " and "))


def board_key(text):
    """(class, stem) of the board a heading names, or None: ("nrd", "central platte"), ("ppd", "nebraska"),
    ("esu", "10"), ("cc", "southeast"), ("lc", ""). The stem is every word standing before the board's kind words."""
    tk = tokens(text)
    for cls, seq in BOARD_CLASSES:
        for p in range(len(tk) - len(seq) + 1):
            if tuple(tk[p:p + len(seq)]) != seq:
                continue
            if cls == "esu":
                rest = [t for t in tk[p + len(seq):] if t not in ("no", "number")]
                return (cls, str(int(rest[0]))) if rest and rest[0].isdigit() else None
            if cls in ("mud", "lc"):
                return (cls, "")
            if cls == "ppid" or cls == "ppd":
                stem = [t for t in tk[:p] if t != "rural"]
                return ("ppid" if cls == "ppid" else "ppd", " ".join(stem))
            return (cls, " ".join(tk[:p]))
    return None


def match_board(text, known):
    """The key of the board on the Secretary of State's list that a county's heading names, or None. A county writes
    more before the name ("For Board Of Directors - Omaha Public Power District"), so the list's own stems are tried
    against the words standing just before the kind words, the longest first."""
    tk = tokens(text)
    for cls, seq in BOARD_CLASSES:
        for p in range(len(tk) - len(seq) + 1):
            if tuple(tk[p:p + len(seq)]) != seq:
                continue
            if cls in ("esu", "mud", "lc"):
                key = board_key(text)
                return key if key in known else None
            before = [t for t in tk[:p] if t != "rural"]
            fits = [k for k in known if k[0] == cls and k[1] and before[-len(k[1].split()):] == k[1].split()]
            return max(fits, key=lambda k: len(k[1])) if fits else None
    return None


def seat_key(text):
    """A seat's words as one key: "Subdistrict 03" -> "3", "At Large (Two Elected)" -> "al", "Gosper County Subdivision"
    -> "gosper", "" -> "". The kind word (district, subdistrict, subdivision, subcouncil) is set aside: a board uses one."""
    t = YEARS.sub(" ", re.sub(r"\(.*?\)", " ", plain(text))).lower()
    m = re.search(r"\b(?:sub ?district|subdivision|subcouncil|district)\s*(?:no\.?\s*)?0*(\d+)\b", t)
    if m:
        return m.group(1)
    if re.search(r"\bat[ -]large\b", t):
        return "al"
    m = re.search(r"([a-z][a-z ]*?)\s+(?:county\s+)?subdivision\b", t)
    return " ".join(m.group(1).split()) if m else ""


BOARD_OFFICE = re.compile(r"\b(?:for\s+)?(?:board of (?:directors|governors)|member of the board|board member|coordinating council)\b", re.I)


def heading_seat(text):
    """The seat a county's heading names after the board's own name: "Dawson Public Power District For Board of
    Directors Buffalo Subdivision - Six Year Term" -> "buffalo"."""
    t = plain(text)
    low = t.lower()
    for cls, seq in BOARD_CLASSES:
        m = re.search(r"\b" + r"\W+".join(seq) + r"\b", low)
        if m:
            rest = t[m.end():]
            if cls == "esu":
                rest = re.sub(r"^\s*(?:No\.?|Number)?\s*\d+", "", rest, flags=re.I)
            rest = re.sub(r"^\s*District\b(?!\s*(?:\d|at\b))", " ", rest, flags=re.I)
            return seat_key(BOARD_OFFICE.sub(" ", rest))
    return ""


def term_of(text):
    m = YEARS.search(plain(text))
    if not m:
        return None
    return int(m.group(1)) if m.group(1).isdigit() else WORD_NUM[m.group(1).lower()]


# ---------------------------------------------------------------- where each board's seats are voted on

CANVASS_BOARD = re.compile(r"Natural\s*Resources\s*District|Public Power|Irrigation\s*District|Community College|Educational Service Unit|"
                           r"Reclamation District|Learning Community|Metropolitan Utilities District", re.I)
CANVASS_SEAT = re.compile(r"^(?:(?:Sub)?district|Subdivision|Subcouncil)\s*\d+\b|^At[ -]Large\b|\bSubdivision\b", re.I)
CANVASS_END = re.compile(r"^(Proposed Amendment|Initiative Measure|Referendum Measure|Historical Tabulation)")


def canvass_counties(data, cmap):
    """{"class|stem|seat": [county GEOIDs]} for the district boards in a general election canvass book: the county
    labels standing under each table's "County" heading, with the seat and the board named above the table. No name
    and no figure is read."""
    names = {full[:-len(" County")]: geoid for geoid, full in cmap.values()}
    pdf = PDF(data)
    out, started, last_board, skipped = {}, False, None, 0
    for n, (page, res) in enumerate(pdf.pages(), start=1):
        allc = [(x0, y, text, x1) for y, runs in pdf_rows(pdf, page, res) for x0, _rs, text, x1 in cells(runs)]
        texts = [c[2] for c in allc]
        if not started:
            started = "Community Colleges" in texts          # the section's own title; the contents page prints it with dots
            if not started:
                continue
        if any(CANVASS_END.match(t) for t in texts):
            break
        heads = [c for c in allc if c[2] == "County"]
        boards = [c for c in allc if CANVASS_BOARD.search(c[2]) and not NUMBER.match(c[2])]
        seats = [c for c in allc if CANVASS_SEAT.search(c[2]) and not CANVASS_BOARD.search(c[2])]
        for h in sorted(heads, key=lambda c: (-c[1], c[0])):
            labels = []
            for c in sorted((c for c in allc if abs(c[0] - h[0]) <= 3.5 and c[1] < h[1] - 0.5), key=lambda c: -c[1]):
                if c[2] != "Total" and c[2] not in names:
                    break                                   # the end of this table's own rows
                labels.append(c[2])
            if not labels or labels[0] != "Total" or len(labels) < 2:
                skipped += 1
                continue
            same = sorted((c for c in heads if abs(c[1] - h[1]) <= 2 and c[0] > h[0]), key=lambda c: c[0])
            right = same[0][0] - 8 if same else 1e9
            span_end = max(c[3] for c in allc if abs(c[1] - h[1]) <= 2 and h[0] <= c[0] < right)

            def over(c):                                    # a centred heading standing over this table
                return h[0] - 10 <= (c[0] + c[3]) / 2 <= max(span_end, h[0] + 120) + 40 and c[0] < right
            d = [c for c in seats if 0 < c[1] - h[1] < 60 and (abs(c[0] - (h[0] - 5.9)) <= 3 or over(c))]
            d = min(d, key=lambda c: c[1] - h[1]) if d else None
            b = [c for c in boards if c[1] > h[1] and (abs((c[0] + c[3]) / 2 - 306) < 60 or (c[0] < 300) == (h[0] < 300))]
            b = min(b, key=lambda c: c[1] - h[1])[2] if b else last_board
            key = board_key(re.sub(r"^\(?\s*(?:Director|Member of the)\s+", "", re.sub(r"\s*[–—-]\s*Continued\)?$", "", b or "")))
            if not key:
                skipped += 1
                continue
            out.setdefault("|".join(key + (seat_key(d[2]) if d else "",)), set()).update(names[x] for x in labels[1:])
        if boards:
            last_board = min(boards, key=lambda c: c[1])[2]  # the lowest heading on a page carries over to the next
    if not out:
        raise Layout("no board's table was found in the canvass book")
    return {k: sorted(v) for k, v in out.items()}, skipped


def canvass_book(folder, year, url, cmap, say):
    def build():
        raw = fetch(url)
        if not raw.startswith(b"%PDF"):
            raise ValueError("the address did not return a PDF")
        tables, skipped = canvass_counties(raw, cmap)
        return {"year": year, "url": url, "fetched": dt.date.today().isoformat(), "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw),
                "kept": "for each district board's seat, the counties whose rows stand in its table; no name, no figure", "skipped": skipped,
                "tables": tables}
    return extract(os.path.join(folder, f"sos_general_canvass_{year}_board_counties.json"), build, say, f"the {year} general election canvass book")


def judicial_districts(folder, cmap, say):
    """{"24-301.02": {district number: [county GEOIDs]}, "24-503": {...}}: the counties of each district court judicial
    district and each county judge district, as the two statutes list them."""
    names = {fold(full[:-len(" County")]): geoid for geoid, full in cmap.values()}

    def build():
        out = {"fetched": dt.date.today().isoformat(), "statutes": {}}
        for sec in ("24-301.02", "24-503"):
            raw = fetch(STATUTE_URL.format(sec))
            page = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", raw.decode("utf-8", "replace"))
            text = re.sub(r"\s+", " ", H.unescape(re.sub(r"<[^>]+>", " ", page)))
            found = {}
            for m in re.finditer(r"District No\. (\d+) shall contain the count(?:y|ies) of ([^;.]+)[;.]", text):
                cs = [c.strip() for c in re.split(r",|\band\b", m.group(2)) if c.strip()]
                bad = [c for c in cs if fold(c) not in names]
                if bad:
                    raise Layout(f"section {sec} names a county the Census Bureau's county file does not have")
                found[m.group(1)] = sorted(names[fold(c)] for c in cs)
            if sorted(found, key=int) != [str(i) for i in range(1, 13)] or len({c for v in found.values() for c in v}) != 93 \
                    or sum(len(v) for v in found.values()) != 93:
                raise Layout(f"section {sec} was not read as twelve districts holding each of the 93 counties once")
            out["statutes"][sec] = {"url": STATUTE_URL.format(sec), "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw), "districts": found}
        return out
    return extract(os.path.join(folder, "statutes_judicial_districts.json"), build, say, "the statutes on judicial districts")


def election_offices(folder, cmap, say):
    """{county GEOID: the address of its election office's page} from the Secretary of State's directory of county
    election offices. Two things are read from each entry, the county's name and the Website link; the names,
    addresses, telephone numbers and e-mail of the officials are never read or kept."""
    names = {fold(full[:-len(" County")]): geoid for geoid, full in cmap.values()}

    def build():
        raw = fetch(DIRECTORY_URL)
        page = raw.decode("utf-8", "replace")
        out = {}
        for part in page.split("<strong>County:")[1:]:
            m = re.match(r"\s*</strong>\s*([A-Za-z .]+?)\s*\(\d+\)", part)
            w = re.search(r"<strong>Website:(?:\s|&nbsp;)*</strong>\s*<a[^>]+href=\"(https?://[^\"@]+)\"", part)
            if m and w and fold(m.group(1)) in names:
                out[names[fold(m.group(1))]] = H.unescape(w.group(1))
        if len(out) < 80:
            raise Layout("the directory of county election offices was not read as one entry a county")
        return {"url": DIRECTORY_URL, "fetched": dt.date.today().isoformat(), "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw),
                "kept": "each county's name and the address of its election office's page; nothing else", "offices": out}
    return extract(os.path.join(folder, "sos_county_election_offices.json"), build, say, "the directory of county election offices", 30)


# ---------------------------------------------------------------- the Census Bureau's names and codes

def census_places(folder, cmap, say):
    """({(folded name, "city" or "village"): [(code, name, [GEOIDs])]}, {folded bare name: [...]}) for Nebraska's
    incorporated places, and {(county GEOID, name in lower case): (code, name)} for its townships. The two files hold names,
    codes and counties only and are kept whole."""
    by_full = {full: geoid for geoid, full in cmap.values()}
    ppath, spath = os.path.join(folder, "st31_ne_place2020.txt"), os.path.join(folder, "st31_ne_cousub2020.txt")
    net.download(PLACE_URL, ppath, 3650, say=say)
    net.download(COUSUB_URL, spath, 3650, say=say)
    exact, bare, towns = defaultdict(list), defaultdict(list), {}
    with open(ppath, encoding="utf-8", errors="replace") as fh:
        if fh.readline().strip() != PLACE_HEAD:
            stop("st31_ne_place2020.txt: the header is not the one this loader was checked against")
        for ln, line in enumerate(fh, start=2):
            f = line.rstrip("\r\n").split("|")
            if len(f) != 9 or f[1] != FIPS:
                stop(f"st31_ne_place2020.txt: line {ln} does not fit the header")
            m = re.fullmatch(r"(.+) (city|village)", f[4])
            if f[5] != "INCORPORATED PLACE" or not m:
                continue
            cids = [by_full.get(c) for c in f[8].split("~~~")]
            if not all(cids):
                stop(f"st31_ne_place2020.txt: line {ln} names a county the county file does not have")
            rec = (f[2], f[4], sorted(cids))
            exact[(fold(m.group(1)), m.group(2))].append(rec)
            bare[fold(m.group(1))].append(rec)
    with open(spath, encoding="utf-8", errors="replace") as fh:
        if fh.readline().strip() != COUSUB_HEAD:
            stop("st31_ne_cousub2020.txt: the header is not the one this loader was checked against")
        for ln, line in enumerate(fh, start=2):
            f = line.rstrip("\r\n").split("|")
            if len(f) != 9 or f[1] != FIPS or f[3] not in by_full:
                stop(f"st31_ne_cousub2020.txt: line {ln} does not fit the header")
            if f[7] == "T1":                                # an active township
                towns[(by_full[f[3]], f[6].lower())] = (f[4], f[6])
    return exact, bare, towns, ppath, spath


SCHOOL_DROP = {"public", "school", "schools", "district", "community", "area", "dist"}


def school_core(name):
    """A school district's name without its number and its kind words, for matching one printing to another:
    "Millard Public Schools #17", "Millard School District 17" and "Millard Public Schools" are all "millard"."""
    t = plain(name).replace("/", " ")
    t = re.sub(r"\s*#\s*[\w-]+$", "", t)
    t = re.sub(r"\s+\d\d-\d{4}$", "", t)
    t = re.sub(r"\s*(?:Public\s+)?(?:School\s+)?(?:District|Dist\.?)\s*#?\s*(?:OR-)?[\dR -]*$", "", t, flags=re.I)
    return " ".join(w for w in re.findall(r"[a-z0-9]+", t.lower()) if w not in SCHOOL_DROP)


def census_schools(folder, say):
    """{core name: [(LEA code, name)]} for Nebraska's school districts, from the attribute table of the Census Bureau's
    unified school district file (names and codes only; no shape is read)."""
    path = os.path.join(folder, "cb_2024_31_unsd_500k.zip")
    net.download(UNSD_URL, path, 3650, say=say)
    import shapefile                                        # pyshp
    z = zipfile.ZipFile(path)
    dbf = [n for n in z.namelist() if n.lower().endswith(".dbf")]
    if len(dbf) != 1:
        stop("the Census school district file does not hold exactly one attribute table")
    rdr = shapefile.Reader(dbf=io.BytesIO(z.read(dbf[0])))
    fields = [f[0] for f in rdr.fields[1:]]
    if not all(k in fields for k in ("STATEFP", "UNSDLEA", "NAME")):
        stop("the Census school district file has no STATEFP, UNSDLEA and NAME columns")
    out = defaultdict(list)
    for rec in rdr.iterRecords():
        rec = dict(zip(fields, rec))
        if str(rec["STATEFP"]) == FIPS and re.fullmatch(r"\d{5}", str(rec["UNSDLEA"])):
            out[school_core(str(rec["NAME"]))].append((str(rec["UNSDLEA"]), str(rec["NAME"])))
    if len(out) < 200:
        stop(f"the Census school district file gives {len(out)} Nebraska districts; about 244 were expected")
    return out, path


# ---------------------------------------------------------------- the counties' own files: one reader a layout
#
# Every reader returns {"contests": [{"head": [heading lines], "vote_for": n or None, "cands": [[name, party or None]],
# "none": the file says nobody filed, "where": page and column}], "retention": [[court, judge]], "other": {counts}}.
# None of these files has a contact column; a reader that meets a layout it does not know raises Layout, naming the
# page and the check, and the county is then left as a gap rather than half read.

def tight(runs):
    """Runs on one printed line as text: a gap wider than a tenth of the type is a space (these files set type tightly)."""
    text, end = "", None
    for x0, _y, size, t, x1 in sorted(runs, key=lambda r: r[0]):
        if end is not None and x0 - end > 0.1 * size and not text.endswith(" ") and not t.startswith(" "):
            text += " "
        text += t
        end = max(end or x1, x1)
    return re.sub(r"\s+", " ", text).strip()


def page_cells(pdf, page, res, gap=8.0):
    """[(y, x0, x1, size, text)] of a page: runs grouped into printed rows, split where the gap is wider than `gap`."""
    rows = []
    for r in sorted((r for r in page_runs(pdf, page, res) if r[3].strip()), key=lambda r: (-r[1], r[0])):
        if rows and abs(rows[-1][0] - r[1]) <= max(1.5, 0.35 * r[2]):
            rows[-1][1].append(r)
        else:
            rows.append([r[1], [r]])
    out = []
    for y, rs in rows:
        cur = []
        for r in sorted(rs, key=lambda r: r[0]):
            if cur and r[0] - max(c[4] for c in cur) > gap:
                out.append((y, cur[0][0], max(c[4] for c in cur), max(c[2] for c in cur), tight(cur)))
                cur = []
            cur.append(r)
        if cur:
            out.append((y, cur[0][0], max(c[4] for c in cur), max(c[2] for c in cur), tight(cur)))
    return out


def pages_of(data):
    """A county file's pages as cells, after three checks: it is a PDF, it has a text layer, and it says it is for the
    2026 general election."""
    if not data.startswith(b"%PDF"):
        raise Layout("the file is not a PDF")
    pdf = PDF(data)
    out = [page_cells(pdf, page, res) for page, res in pdf.pages()]
    if not any(out):
        raise Layout("the PDF has no text layer (a scanned picture)")
    first = " ".join(c[4] for c in out[0]) if out[0] else ""
    if not re.search(r"general\s+election", first, re.I) or "2026" not in first:
        raise Layout("the first page does not say it is for the 2026 general election")
    return out


def rows_of(page, lo=0, hi=1e9):
    """A page's cells whose left edge lies in [lo, hi), as printed rows from the top: [(y, [cells left to right])]."""
    rows = []
    for c in sorted((c for c in page if lo <= c[1] < hi), key=lambda c: (-c[0], c[1])):
        if rows and abs(rows[-1][0] - c[0]) <= 1.5:
            rows[-1][1].append(c)
        else:
            rows.append([c[0], [c]])
    return rows


def contest(head, vote_for, where):
    return {"head": list(head), "vote_for": vote_for, "cands": [], "none": False, "where": where}


def read_ess(data):
    """The "publication ballot" county election offices print from the state's ballot system: three columns; a contest
    is a few centred heading lines ending in "Vote for ...", then the names at the column's left edge, each with its
    party in small type underneath on a partisan ticket."""
    paged = pages_of(data)
    starts = Counter(round(c[1], 1) for cells_ in paged for c in cells_ if 9.5 <= c[3] <= 10.5)
    top, best = [x for x, _k in starts.most_common(12)], None
    for a in top:
        for b in top:
            for c in top:
                if 150 <= b - a <= 230 and abs((c - b) - (b - a)) <= 1.0 and (best is None or starts[a] + starts[b] + starts[c] > best[0]):
                    best = (starts[a] + starts[b] + starts[c], [a, b, c])
    if best is None:
        raise Layout("the three name columns of the publication ballot were not found")
    names = best[1]
    centre = [x + 66.3 for x in names]
    contests, retention, other = [], [], Counter()
    for n, cells_ in enumerate(paged, start=1):
        for col in range(3):
            lo, hi = names[col] - 30, (names[col + 1] - 30 if col < 2 else 1e9)
            head, cur, last, carry, ticket2 = [], None, None, None, False      # last: the height of the Vote-for line or of the last name
            for y, x0, x1, size, t in sorted((c for c in cells_ if lo <= c[1] < hi), key=lambda c: -c[0]):
                at_name, centred = abs(x0 - names[col]) <= 0.7, abs((x0 + x1) / 2 - centre[col]) <= 1.6
                if carry:                                   # a judge's question that runs over two lines
                    t, carry = carry + " " + t, None
                elif t.startswith("Shall Judge ") and not SHALL.match(t):
                    carry = t
                    continue
                if size > 12:
                    continue                                # the ballot's title
                if centred and re.fullmatch(r"[A-Z .&'/-]+ TICKET", t):
                    head, cur = [], None
                    continue
                if centred and VOTE.match(t):
                    if not head:
                        raise Layout(f"page {n} column {col + 1}: a Vote-for line with no office above it")
                    cur = contest(head, NUM[VOTE.match(t).group(1).lower()], f"page {n} column {col + 1}")
                    contests.append(cur)
                    ticket2 = head[0].startswith("For Governor")      # two names to a ticket, set in a way of their own: the state part reads them
                    head, last = [], y
                    continue
                # a name sits within 27 points of the Vote-for line and 22 of the name before it; a heading starts lower
                if cur is not None and at_name and not head and last - y <= (27 if not cur["cands"] else 22):
                    if size < 9:                            # the party, in small type under the name
                        if not cur["cands"] or cur["cands"][-1][1] is not None:
                            raise Layout(f"page {n} column {col + 1}: a party line with no name above it")
                        cur["cands"][-1][1] = t
                    else:
                        cur["cands"].append([t, None])
                        last = y
                    continue
                m = SHALL.match(t)
                if m:
                    retention.append([" ".join(head), m.group("judge")])
                    head, cur = [], None
                    continue
                if centred:
                    cur, ticket2 = None, False
                    head.append(t)
                    continue
                if at_name and size >= 9.5 and not ticket2 and NAMELIKE.match(t) and len(t) <= 40:
                    raise Layout(f"page {n} column {col + 1}: a name with no contest above it")
                other["other text"] += 1                    # instructions, Yes and No, the text of a measure
                head, cur = [], None
    return {"contests": contests, "retention": retention, "other": dict(other)}


def read_douglas(data):
    """Douglas County's notice: one column; "For <office> (Partisan)" or "(Nonpartisan)", then a line a candidate,
    "Name (Party)" on the partisan ones. A contest that runs over a page is headed again with "(Cont.)". The list ends
    where the issues begin."""
    contests, retention = [], []
    cur, on, judges = None, False, None
    for n, cells_ in enumerate(pages_of(data), start=1):
        for _y, x0, _x1, _size, t in sorted(cells_, key=lambda c: (-c[0], c[1])):
            if re.fullmatch(r"\d{1,2}", t) and x0 > 400:
                continue                                    # the page number
            if not on:
                on = t == "Offices and Candidates Appearing on the General Ballot"
                continue
            if t.startswith("Issue(s) Appearing on the General Ballot"):
                return {"contests": contests, "retention": retention, "other": {}}
            if abs(x0 - 54.0) > 1.5:
                raise Layout(f"page {n}: a line that does not start at the left margin")
            m = re.fullmatch(r"For (?P<office>.+) \((?P<kind>Partisan|Nonpartisan)\)(?P<cont> \(Cont\.\))?", t)
            if m and m.group("cont"):                       # the same contest, carried over the page
                if (judges or (cur and cur["head"][0][4:])) != m.group("office"):
                    raise Layout(f"page {n}: a continued heading that is not the contest before it")
                continue
            if m:
                judges = m.group("office") if m.group("office").startswith("Judge Of ") else None
                cur = None
                if not judges:
                    cur = contest(["For " + m.group("office")], None, f"page {n}")
                    cur["partisan"] = m.group("kind") == "Partisan"
                    contests.append(cur)
                continue
            if t.startswith("For ") or "(Partisan)" in t or "(Nonpartisan)" in t:
                raise Layout(f"page {n}: a heading that runs over two lines")
            if judges:
                retention.append([judges, t])
            elif cur is None:
                raise Layout(f"page {n}: a name before any office heading")
            elif cur["partisan"]:
                m = re.fullmatch(r"(?P<name>.+?) \((?P<party>[^()]+)\)(?: - (?:Governor|Lieutenant Governor))?", t)
                if not m:
                    raise Layout(f"page {n}: a partisan candidate's line without a party in brackets")
                cur["cands"].append([m.group("name"), m.group("party")])
            else:
                cur["cands"].append([t, None])
    raise Layout("the notice has no line saying where the issues begin")


def read_sarpy(data):
    """Sarpy County's notice: three columns; a contest is a block of lines closed by a blank line: centred heading
    lines down to "Vote for ...", then a row a candidate, the party to the left of the name on the partisan ones."""
    contests, retention, other = [], [], Counter()
    for n, cells_ in enumerate(pages_of(data), start=1):
        cells_ = [c for c in cells_ if c[3] <= 11]         # the title and the opening paragraph are in larger type
        for col, (lo, hi) in enumerate(((0, 208), (208, 405), (405, 1e9)), start=1):
            blocks = []
            for y, cs in rows_of(cells_, lo, hi):
                if blocks and blocks[-1][-1][0] - y <= 18:
                    blocks[-1].append((y, cs))
                else:
                    blocks.append([(y, cs)])
            for block in blocks:
                texts = [[c[4] for c in cs] for _y, cs in block]
                at = next((i for i, ts in enumerate(texts) if len(ts) == 1 and VOTE.match(ts[0])), None)
                yn = next((i for i, ts in enumerate(texts) if ts == ["Vote YES or NO"]), None)
                if at is None and yn and texts[0][0].startswith("Judge of ") and all(len(ts) == 1 for ts in texts):
                    retention += [[" ".join(ts[0] for ts in texts[:yn]), ts[0]] for ts in texts[yn + 1:]]      # a court, then its judges
                    continue
                if at is None:
                    other["block without a Vote-for line"] += 1      # a ticket's name, the judges, the notice's own words
                    continue
                if at == 0 or any(len(ts) != 1 for ts in texts[:at]):
                    raise Layout(f"page {n} column {col}: a heading that is not one line to a row")
                c = contest([ts[0] for ts in texts[:at]], NUM[VOTE.match(texts[at][0]).group(1).lower()], f"page {n} column {col}")
                pending = None                              # a party's name that runs over two lines
                for ts in texts[at + 1:]:
                    if len(ts) == 1 and ts[0] in ("Nebraska Working", "Legal Marijuana"):
                        pending = ts[0]
                    elif len(ts) == 1:
                        c["cands"].append([ts[0], None])
                    elif len(ts) == 2:
                        c["cands"].append([ts[1], (pending + " " + ts[0]) if pending else ts[0]])
                        pending = None
                    else:
                        raise Layout(f"page {n} column {col}: a candidate row with {len(ts)} cells")
                contests.append(c)
    return {"contests": contests, "retention": retention, "other": dict(other)}


LANC_BOARD = re.compile(r"(Natural Resources District|Public Power District|Power and Irrigation District|Community College|Educational Service Unit No\. \d+|"
                        r"Learning Community.*|Utilities District.*|Reclamation District)$")
LANC_SEAT = re.compile(r"^(?:(?:District|Subdistrict|Subdivision|Subcouncil|Ward|Precinct) [0-9IVX]+\b.*|At[ -]Large|Statewide\b.*)$")
LANC_TERM = re.compile(r"^[-–—]?\s*\d+ Year Term$", re.I)
LANC_MORE = re.compile(r"^(School Board Member|Board of (Education|Trustees|Governors|Directors)|City Council|Village Board of Trustees|Board Member)$")


def read_lancaster(data):
    """Lancaster County's notice: two columns, read one after the other; inside a column a contest's seats or its
    names can stand two abreast. A heading starts with "For" or names a board; "District 2", "Ward 4", "At Large" open
    a seat under it; a party stands to the right of its candidate on the partisan tickets. The notice says a voter
    votes for one in each contest unless it says otherwise."""
    contests, retention = [], []
    by_key, st = {}, {"group": None, "sub": [None, None], "fresh": False, "skip": False, "vote": None, "judge": False}

    def target(side, where):
        sub = st["sub"][side] or st["sub"][0]
        key = (tuple(st["group"]), sub)
        if key not in by_key:
            by_key[key] = contest(st["group"] + ([sub] if sub else []), st["vote"] or 1, where)
            contests.append(by_key[key])
        return by_key[key]

    started = False
    for n, cells_ in enumerate(pages_of(data), start=1):
        for main, (lo, hi) in enumerate(((0, 320), (320, 1e9))):
            where = f"page {n} column {main + 1}"
            for _y, cs in rows_of(cells_, lo, hi):
                if not started:
                    started = any(c[4].startswith("The certified list of candidates appears below") for c in cs)
                    continue
                for _yy, x0, _x1, size, t in cs:
                    side = 1 if x0 - (27.0 if main == 0 else 333.0) > 60 else 0
                    if size >= 10.5:                        # a ticket's name
                        if t == "SPECIAL ISSUES":
                            return {"contests": contests, "retention": retention, "other": {}}
                        st.update(group=None, sub=[None, None], fresh=False, skip=False, vote=None, judge=False)
                        continue
                    if t in PARTIES and side == 1:
                        if st["skip"]:
                            continue
                        tgt = target(0, where) if st["group"] else None
                        if not tgt or not tgt["cands"] or tgt["cands"][-1][1] is not None:
                            raise Layout(f"{where}: a party with no candidate to its left")
                        tgt["cands"][-1][1] = t
                        continue
                    if side == 0 and (t.startswith("For ") or LANC_BOARD.search(t) or t.startswith("Judge of ")):
                        if (st["fresh"] and st["group"] and len(st["group"]) == 1 and t.startswith("For ")
                                and not st["group"][0].startswith(("For ", "Judge of "))):
                            st["group"].append(t)           # a board's name, then "For Board of ..." on the next line
                        else:
                            st.update(group=[t], sub=[None, None], vote=None)
                        st.update(fresh=True, skip=t.startswith("For Governor"), judge=t.startswith("Judge of "))
                        continue
                    if st["skip"]:
                        continue
                    if st["group"] is None:
                        raise Layout(f"{where}: text before any office heading")
                    if st["fresh"] and side == 0 and LANC_MORE.match(t):
                        st["group"].append(t)
                        continue
                    m = SHALL.match(t)
                    if m:
                        retention.append([" ".join(st["group"] + ([st["sub"][0]] if st["sub"][0] else [])), m.group("judge")])
                        st["fresh"] = False
                        continue
                    if st["judge"]:
                        if LANC_SEAT.match(t) or LANC_TERM.match(t):
                            st["sub"][0] = ((st["sub"][0] + " ") if st["sub"][0] and LANC_TERM.match(t) else "") + t
                            continue
                        raise Layout(f"{where}: a line under a judge's heading that is not understood")
                    if LANC_SEAT.match(t):
                        st["sub"][side], st["fresh"] = t, False
                        continue
                    if LANC_TERM.match(t):
                        t2 = t.lstrip("-–— ").strip()
                        if st["sub"][side]:
                            st["sub"][side] += " - " + t2
                        elif st["fresh"]:
                            st["group"].append(t2)
                        else:
                            raise Layout(f"{where}: a term line that belongs to nothing")
                        continue
                    if VOTE.match(t):
                        st["vote"], st["fresh"] = NUM[VOTE.match(t).group(1).lower()], False
                        continue
                    if re.fullmatch(r"No Candidates? Filed", t):
                        target(side, where)["none"] = True
                        st["fresh"] = False
                        continue
                    if not NAMELIKE.match(t) or OFFICE_WORDS.search(t):
                        raise Layout(f"{where}: a line that is neither a heading nor a name")
                    target(side, where)["cands"].append([t, None])
                    st["fresh"] = False
    raise Layout("the notice has no SPECIAL ISSUES heading to end the list of candidates")


YORK_VOTE = re.compile(r"^Vote for (?:up to )?(ONE|TWO|THREE|FOUR|FIVE|SIX|SEVEN|EIGHT|NINE|TEN)(?:, (\d+)-Year Term)?$", re.I)


def read_york(data):
    """York County's notice: three columns that run on from one to the next; a contest is a block closed by a blank
    line: heading lines, "Vote for ONE, 4-Year Term", then a line a candidate, the party on the line above the name on
    the partisan ticket ("Vacancy" where nobody filed). A board's subdistricts follow one another in one block, each
    indented, with its own Vote-for line."""
    contests, retention, other = [], [], Counter()
    stream = []                                             # (where, text, gap above: True, False, None at a column's top, indented)
    for n, cells_ in enumerate(pages_of(data), start=1):
        for col, (lo, hi) in enumerate(((40, 280), (280, 520), (520, 1e9)), start=1):
            lines = sorted((c for c in cells_ if lo <= c[1] < hi), key=lambda c: (-c[0], c[1]))
            if not lines:
                continue
            base = min(c[1] for c in lines)
            for i, c in enumerate(lines):
                stream.append((f"page {n} column {col}", c[4], None if i == 0 else lines[i - 1][0] - c[0] > 18, c[1] - base > 40))
    group, pending, cur, party = [], [], None, None

    def close(where):
        """The end of a block. Lines that never reached a Vote-for line are a ticket's name, a judge's question or the
        notice's own words."""
        nonlocal group, pending, cur, party
        if party:
            raise Layout(f"{where}: a block that ends with a party and no name")
        if cur is not None and pending:
            raise Layout(f"{where}: a subdistrict with no Vote-for line under it")
        texts = [p for p, _ind in pending]
        for i, t in enumerate(texts):
            m = SHALL.match(t)
            if m:
                retention.append([" ".join(texts[:i]), m.group("judge")])
        if pending:
            other["block without a Vote-for line"] += 1
        group, pending, cur, party = [], [], None, None

    for k, (where, t, gap, indented) in enumerate(stream):
        if gap is None and cur is not None and not indented and t not in PARTIES and t != "Vacancy":
            # the top of a column: more names of the contest above, or the next contest? A contest has a Vote-for line
            # before its block ends
            span = []
            for _w, t2, gap2, _i in stream[k:]:
                if span and gap2 is not False:
                    break
                span.append(t2)
            if any(YORK_VOTE.match(x) for x in span) and not YORK_VOTE.match(t):
                gap = True
        if gap is True:
            close(where)
        if re.fullmatch(r"[A-Za-z ]+ Ticket", t):
            continue
        m = YORK_VOTE.match(t)
        if m:
            if not pending:
                raise Layout(f"{where}: a Vote-for line with no office or subdistrict above it")
            if cur is None:
                group, subs = [p for p, ind in pending if not ind], [p for p, ind in pending if ind]
            else:
                subs = [p for p, _ind in pending]
            cur = contest(group + subs + ([f"{m.group(2)} Year Term"] if m.group(2) else []), NUM[m.group(1).lower()], where)
            contests.append(cur)
            pending, party = [], None
            continue
        if cur is None:
            pending.append((t, indented))
        elif indented:                                      # the next subdistrict of the same board
            pending = [(t, True)]
        elif pending:
            raise Layout(f"{where}: a subdistrict with no Vote-for line under it")
        elif t in PARTIES:
            if party:
                raise Layout(f"{where}: two party lines with no name between them")
            party = t
        elif t == "Vacancy":
            cur["none"] = True
        else:
            cur["cands"].append([t, party])
            party = None
    close("the last page")
    return {"contests": contests, "retention": retention, "other": dict(other)}


WASH_VOTE = re.compile(r"^(?P<head>.*?)\s*[-–—,]\s*Vote for (?:up to )?(?P<n>ONE|TWO|THREE|FOUR|FIVE|SIX|SEVEN|EIGHT|NINE|TEN)$", re.I)
WASH_SUB = re.compile(r"^(Ward|District|Subdistrict|Precinct) \d+$")
WASH_SKIP = ("UNITED STATES SENATORIAL TICKET", "CONGRESSIONAL TICKET", "STATE TICKET")
WASH_GROUP = re.compile(r"\b(City|Village|School|District|Council|Board|Authority)\b")


def read_washington(data):
    """Washington County's notice: one column; "<office> – Vote for ONE", then the names, one to a line; on the county
    ticket a row of parties and under it the row of names; a council's or a school board's wards, and the townships,
    stand two abreast. The federal and state tickets at its head are not read."""
    contests, retention, other = [], [], Counter()
    rows = [(n, cs) for n, cells_ in enumerate(pages_of(data), start=1) for _y, cs in rows_of(cells_)]
    skip, group, cols, parties, judge = True, None, [], None, None      # cols: [(left edge, contest)] of the open contest(s)
    for i, (n, cs) in enumerate(rows):
        texts = [c[4] for c in cs]
        flat, where = " ".join(texts), f"page {n}"
        if len(cs) == 1 and re.fullmatch(r"[A-Z ]+ TICKET", flat):
            skip, group, cols, parties, judge = flat in WASH_SKIP, None, [], None, None
            continue
        if skip:
            continue
        m = SHALL.match(flat)
        if m:
            if not judge:
                raise Layout(f"{where}: a retention question with no court named above it")
            retention.append([judge, m.group("judge")])
            cols, parties = [], None
            continue
        if flat.startswith("Judge of "):
            judge, group, cols, parties = flat, None, [], None
            continue
        votes = [WASH_VOTE.match(t) for t in texts]
        if all(votes):
            subs = all(WASH_SUB.match(v.group("head")) for v in votes)
            if subs and not group:
                raise Layout(f"{where}: wards with no office named above them")
            if not subs:
                group = None
            cols, parties, judge = [], None, None
            for c, v in zip(cs, votes):
                con = contest((group or []) + [v.group("head")], NUM[v.group("n").lower()], where)
                contests.append(con)
                cols.append((c[1], con))
            continue
        if any(votes):                                      # one column's next contest beside another column's last name
            if parties is not None or not cols:
                raise Layout(f"{where}: a row that mixes a heading with something else")
            for c, v in zip(cs, votes):
                near = [k for k, (x, _con) in enumerate(cols) if abs(x - c[1]) <= 3]
                if len(near) != 1:
                    raise Layout(f"{where}: a row that mixes a heading with something else")
                old = cols[near[0]][1]
                if v and bool(WASH_SUB.match(v.group("head"))) == bool(WASH_SUB.match(old["head"][-1])) \
                        and v.group("head").split()[0] == old["head"][-1].split()[0]:
                    con = contest(old["head"][:-1] + [v.group("head")], NUM[v.group("n").lower()], where)
                    contests.append(con)
                    cols[near[0]] = (c[1], con)
                elif not v and NAMELIKE.match(c[4]) and len(c[4]) <= 40:
                    old["cands"].append([c[4], None])
                else:
                    raise Layout(f"{where}: a row that mixes a heading with something else")
            continue
        if all(t in PARTIES for t in texts) and cols:
            parties = [(c[1], c[4]) for c in cs]
            continue
        nxt_votes = [WASH_VOTE.match(c[4]) for c in (rows[i + 1][1] if i + 1 < len(rows) else [])]
        if len(cs) == 1 and nxt_votes and all(nxt_votes) and all(WASH_SUB.match(v.group("head")) for v in nxt_votes):
            if not WASH_GROUP.search(flat):
                raise Layout(f"{where}: wards under a line that does not name an office")
            group, cols, parties, judge = [flat], [], None, None
            continue
        if not cols or not all(NAMELIKE.match(t) and len(t) <= 40 and not re.search(r"[A-Z]{3,} [A-Z]{3,}", t) for t in texts):
            cols, parties, group = [], None, None           # the words of a question, or a line between contests
            other["other line"] += 1
            continue
        if parties is not None:
            if len(parties) != len(cs) or len(cols) != 1 or any(abs(c[1] - p[0]) > 3 for c, p in zip(cs, parties)):
                raise Layout(f"{where}: a row of names that does not sit under its row of parties")
            cols[0][1]["cands"] += [[c[4], p[1]] for c, p in zip(cs, parties)]
            parties = None
            continue
        for c in cs:
            near = [con for x, con in cols if abs(x - c[1]) <= 3]
            if len(near) != 1:
                raise Layout(f"{where}: a name that sits under no contest")
            near[0]["cands"].append([c[4], None])
    return {"contests": contests, "retention": retention, "other": dict(other)}


READERS = {"ess": read_ess, "douglas": read_douglas, "sarpy": read_sarpy, "lancaster": read_lancaster, "york": read_york, "washington": read_washington}


def county_file(folder, spec, say, refresh=False):
    """One county's notice or sample ballot, read into contests and kept as JSON (headings, names and parties only)
    with the day fetched and the fingerprint of the file as it came. A file saved by hand in the county's folder is read
    when the site cannot be reached. Returns (what was read, or None; why not, in plain words)."""
    fips, county, reader, _kind, title, url, name = spec
    sub = os.path.join(folder, slug(county))
    path = os.path.join(sub, "contests.json")
    old = read_json(path)
    if old and not refresh and (dt.date.today() - dt.date.fromisoformat(old["fetched"])).days < LOCAL_FRESH_DAYS:
        return old, None
    saved, how, why = os.path.join(sub, name), "fetched", None
    try:
        raw = fetch(url)
        if not raw.startswith(b"%PDF"):
            raise ValueError("it answered with a page, not the file")
    except Exception as err:  # noqa: BLE001  the site did not answer: a file saved by hand, else the older copy, else a gap
        raw, why = None, ("the county's site did not give the file to a program ("
                          + (f"HTTP {err.code}" if isinstance(err, HTTPError) else str(err) if isinstance(err, ValueError) else type(err).__name__) + ")")
        if os.path.exists(saved):
            raw, how = open(saved, "rb").read(), "saved by hand"
    if raw is None:
        if old:
            say(f"      Nebraska (local): {county} County: {why}; using the copy of {old['fetched']} on disk")
            return old, None
        return None, why
    try:
        got = READERS[reader](raw)
    except Layout as err:
        if old:
            say(f"      Nebraska (local): {county} County: the file's layout has changed ({err}); using the copy of {old['fetched']} on disk")
            return old, None
        return None, f"the file's layout is not the one this loader reads ({err})"
    blanked = 0
    for c in got["contests"]:                               # nothing that reads like contact details is kept
        keep = []
        c["head"] = [h.replace("’", "'") for h in c["head"]]
        for nm, party in c["cands"]:
            nm = nm.replace("’", "'")                   # a typewriter apostrophe, as the Secretary of State's list writes names
            if NOBODY.match(nm):                            # "No candidates filed for this office."
                c["none"] = True
                continue
            if contact_like(nm, True) or (party and contact_like(party, True)):
                blanked += 1
            else:
                keep.append([nm, party])
        c["cands"] = keep
        if any(contact_like(h, True) for h in c["head"]):
            return None, "a heading in the file reads like contact details"
    got["retention"] = [[court.replace("’", "'"), judge.replace("’", "'")] for court, judge in got["retention"]]
    new = {"county": county, "fips": fips, "title": title, "url": url, "how": how, "fetched": dt.date.today().isoformat(),
           "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw), "blanked": blanked, "kept": "office headings, names and parties",
           "contests": got["contests"], "retention": got["retention"], "other": got["other"]}
    keep_json(path, new)
    return new, None


# ---------------------------------------------------------------- what a heading names

STATE_OFFICE = re.compile(r"\b(?:united states senator|u\.s\. house|representative in congress|governor|secretary of state|state treasurer|attorney general|"
                          r"auditor of public accounts|public service commissioner|legislature|state board of education|board of regents|"
                          r"university of nebraska)\b", re.I)      # "governor", not a college's Board of Governors
COUNTY_OFFICE = (       # (the heading without "For" and the county's own name, office kind, office as shown)
    (r"county assessor ?[/-] ?register of deeds", "county_assessor_register_of_deeds", "County Assessor/Register of Deeds"),
    (r"county clerk[ /-]+register of deeds", "county_clerk_register_of_deeds", "County Clerk/Register of Deeds"),
    (r"(?:county )?clerk of (?:the )?district court", "clerk_of_court", "Clerk of the District Court"),
    (r"county (?:board of )?commissioners?", "county_commissioner", "County Commissioner"),
    (r"county board of supervisors", "county_commissioner", "County Supervisor"),
    (r"county assessor", "county_assessor", "County Assessor"),
    (r"county attorney", "county_attorney", "County Attorney"),
    (r"county clerk", "county_clerk", "County Clerk"),
    (r"county engineer", "county_engineer", "County Engineer"),
    (r"county sheriff", "sheriff", "County Sheriff"),
    (r"county surveyor", "county_surveyor", "County Surveyor"),
    (r"county treasurer", "county_treasurer", "County Treasurer"),
    (r"(?:county )?public defender", "public_defender", "Public Defender"),
    (r"(?:county )?register of deeds", "register_of_deeds", "Register of Deeds"),
)
SEAT = r"(?: (?P<d>ward \S+|at[ -]large))?"
MUNICIPAL = (           # (pattern, office kind, office as shown, the place's kind word)
    (r"mayor of (?:the )?(?:city of )?(?P<p>.+)", "mayor", "Mayor", "city"),
    (r"(?P<p>.+?) city mayor(?: at large)?", "mayor", "Mayor", "city"),
    (r"city of (?P<p>.+?) mayor", "mayor", "Mayor", "city"),
    (r"member of (?P<p>.+?) city council" + SEAT, "council", "City Council Member", "city"),
    (r"city of (?P<p>.+?) (?:for )?city council board member", "council", "City Council Member", "city"),
    (r"city of (?P<p>.+?) council member" + SEAT, "council", "City Council Member", "city"),
    (r"(?P<p>.+?) city council(?: member)?" + SEAT, "council", "City Council Member", "city"),
    (r"(?P<p>.+?) village (?:for )?board of trustees", "council", "Village Board of Trustees Member", "village"),
    (r"board of trustees village of (?P<p>.+)", "council", "Village Board of Trustees Member", "village"),
    (r"village of (?P<p>.+?)(?: (?:for )?board (?:of trustees|member))?", "council", "Village Board of Trustees Member", "village"),
)
MANY = ("council", "school_board", "town_supervisor", "airport_board", "sanitary_board")     # boards that seat more than one at a time
USUAL_TERM = {"airport_board": 6}                       # years; four for the other county-filed offices


def lower_small(text):
    """Douglas County prints "Of" and "The" with capitals; they are shown in the usual way."""
    return re.sub(r"(?<=\w )(Of|The|And|For)(?= \w)", lambda m: m.group(1).lower(), text)


def school_parts(t):
    """(name as printed, bare name, number, seat) of a school board heading."""
    s = re.sub(r"^Board of Education\s+", "", t, flags=re.I)
    seat = None
    m = re.search(r"\s+(Ward|Subdistrict)\s+(\S+)$", s, re.I)
    if m:
        seat, s = f"{m.group(1).title()} {m.group(2)}", s[:m.start()]
    s = re.sub(r"\s*(?:For\s+)?(?:Board of Education|Board Member|School Board)$", "", s, flags=re.I).strip()
    # "Norris School District 160 School", "Holdrege District 44 School": that last word belongs to "School Board Member";
    # in "Madison Public School" it is the district's own
    if re.search(r"\s+School$", s, re.I) and re.search(r"\b(?:Schools?|District|Dist\.?)\b|\d\d-\d{4}", s[:-len(" School")], re.I):
        s = s[:-len(" School")].strip()
    shown, number = s, None
    for pat in (r"\s*#\s*([\w-]+)$", r"\s+(\d\d-\d{4})$", r"\s*(?:School\s+)?(?:District|Dist\.?)\s*#?\s*((?:OR-)?[\w-]+)$"):
        m = re.search(pat, s, re.I)
        if m:
            number, s = m.group(1), s[:m.start()].strip()
            break
    return lower_small(shown), s, number, seat


def classify(head, county):
    """What one heading of a county's file names: {"cat": "state"} for a federal or state office (the state part reads
    those), {"cat": "board"} for a district board on the Secretary of State's list, or {"cat": "local", ...} with the
    level, the office, the place and the seat. None when the heading is not understood."""
    text = plain(" ".join(head))
    term = term_of(text)
    t = re.sub(r"\s+", " ", YEARS.sub(" ", text)).strip()
    t = re.sub(r"(?:\s*[-,/]\s*)+$", "", t)
    t = re.sub(r"^For\s+", "", t, flags=re.I)
    if STATE_OFFICE.search(t):
        return {"cat": "state"}
    low = t.lower()
    if any(re.search(r"\b" + r"\W+".join(seq) + r"\b", low) for _cls, seq in BOARD_CLASSES):
        return {"cat": "board", "text": text, "term": term}
    t = re.sub(r"\s*[-,]\s+|\s+[-,]\s*", " ", t)
    t = re.sub(r"\s+", " ", re.sub(r"^" + re.escape(county) + r" (?=County\b)", "", t, flags=re.I)).strip()
    out = {"cat": "local", "term": term, "district": None, "seat": None}
    for pat, kind, office in COUNTY_OFFICE:
        m = re.fullmatch(pat + r"(?: district (\d+))?", t, re.I)
        if m:
            return dict(out, level="county", kind=kind, office=office, jur=("county",), district=f"District {int(m.group(1))}" if m.group(1) else None)
    m = re.fullmatch(r"township (\d+)", t, re.I)
    if m:
        # the three members of a township's board: the kind the other states' township boards are filed under, Nebraska's own title
        return dict(out, level="township", kind="town_supervisor", office="Township Board Member", jur=("township", f"Township {int(m.group(1))}"))
    m = re.fullmatch(r"board of directors (regional metropolitan transit board of omaha) district (\d+)", t, re.I)
    if m:
        return dict(out, level="other", kind="transit_board", office="Regional Metropolitan Transit Board Director",
                    jur=("special", lower_small(m.group(1))), district=f"District {int(m.group(2))}")
    m = re.fullmatch(r"trustee (.+? sanitary district)", t, re.I)
    if m:
        return dict(out, level="other", kind="sanitary_board", office="Sanitary District Trustee", jur=("special", lower_small(m.group(1))))
    m = re.fullmatch(r"city of (.+?) member airport authority(?: (at large))?", t, re.I) or re.fullmatch(r"(.+?) airport authority(?: (at large))?", t, re.I)
    if m:
        return dict(out, level="other", kind="airport_board", office="Airport Authority Board Member", jur=("special", lower_small(m.group(1)) + " Airport Authority"),
                    seat="At Large" if m.group(2) else None)
    for pat, kind, office, word in MUNICIPAL:
        m = re.fullmatch(pat, t, re.I)
        if m:
            d = (m.groupdict().get("d") or "").strip()
            ward = re.fullmatch(r"ward (\S+)", d, re.I)
            return dict(out, level="city", kind=kind, office=office, jur=("place", lower_small(m.group("p")).strip(), word),
                        district=f"Ward {ward.group(1).upper()}" if ward else None, seat="At Large" if d and not ward else None)
    if re.search(r"\bschools?\b|board of education", t, re.I):
        shown, bare_name, number, seat = school_parts(t)
        if bare_name:
            return dict(out, level="school", kind="school_board", office="School Board Member", jur=("school", shown, bare_name, number), district=seat)
    return None


def race_id_for(key, kind, seat=None, special=0):
    return "-".join(x for x in (f"2026-{STATE}", key, kind.replace("_", "-"), slug(seat or "")) if x) + ("-S" if special else "")


BOARD_KINDS = (         # (class, level, office kind, office as shown)
    ("nrd", "other", "natural_resources_board", "Natural Resources District Director"),
    ("ppid", "other", "utility_board", "Public Power and Irrigation District Director"),
    ("ppd", "other", "utility_board", "Public Power District Director"),
    ("mud", "other", "utility_board", "Metropolitan Utilities District Director"),
    ("recl", "other", "water_board", "Reclamation District Director"),
    ("cc", "other", "college_board", "Community College Board of Governors Member"),
    ("esu", "other", "educational_service_board", "Educational Service Unit Board Member"),
    ("lc", "other", "learning_community_council", "Learning Community Coordinating Council Member"),
)
RETENTION = {           # the Judicial Retention sheet's Office -> (code, office kind, the statute that lists its districts' counties)
    "Judge of the Appeals Court": ("COA", "court_of_appeals_retention", None),
    "Judge of the Nebraska Workers' Compensation Court": ("WCC", "workers_compensation_court_retention", None),
    "Judge of the District Court": ("DC", "district_court_retention", "24-301.02"),
    "Judge of the County Court": ("CC", "county_court_retention", "24-503"),
    "Judge of the Separate Juvenile Court": ("JUV", "juvenile_court_retention", None),
}


def local_level(folder, say, cmap, pdf_local=None, refresh=False):
    """Nebraska's local rows, ready to write: the district boards and the judges on the Secretary of State's filing
    list, and the county, city, village, school and township contests of the counties in COUNTY_FILES. Nothing here
    touches the database."""
    net.patient_lookups()
    full = {geoid: name for geoid, name in cmap.values()}           # "31001" -> "Adams County"
    by_county = {name[:-len(" County")]: geoid for geoid, name in full.items()}
    problems = []

    # ---- 1. the filing list: the district boards
    try:
        wb = filing_list(folder, say, refresh)
        byco = contests_by_county(folder, say)
    except Exception as err:  # noqa: BLE001  nothing has been written yet
        stop(f"the Secretary of State's filing list or its contests-by-county sheet could not be fetched ({type(err).__name__}) and no earlier "
             "copy is on disk; nothing was changed. Re-run.")
    rows = wb["rows"]
    counts = Counter()
    boards = {}                                                     # board key -> its name, kind, place id and seats
    for n, r in enumerate(rows, start=2):
        office = r["Office"]
        if office in FEDERAL:
            counts["federal"] += 1
            continue
        if office in LIST_OFFICE:
            counts["state"] += 1
            continue
        if office.startswith("For ") or not office:
            stop(f"row {n} of the filing list names an office this loader does not know; stopping (the row is not printed)")
        counts["board rows"] += 1
        key = board_key(office)
        kind = next((k for k in BOARD_KINDS if key and k[0] == key[0]), None)
        if not kind:
            stop(f"row {n} of the filing list names a board of a kind this loader does not know; stopping (the row is not printed)")
        name = re.sub(r"\s+For Board of (?:Directors|Governors)$", "", office).strip()
        b = boards.setdefault(key, {"name": name, "kind": kind, "id": f"{STATE}-X-{slug(name)}", "seats": {}, "label": Counter()})
        if b["name"] != name:
            stop(f"row {n} of the filing list: two boards read as one; stopping")
        if not r["Candidate Name"]:
            counts["blank names"] += 1
            continue
        if not r["Term"].isdigit() or not r["Vote For"].isdigit():
            stop(f"row {n} of the filing list has a Term or a Vote For that is not a number; stopping")
        if r["Party (if applicable)"]:
            problems.append(f"row {n} of the filing list gives a party to a district board's candidate; the office is nonpartisan and none is shown")
        district = re.sub(r"\b0+(\d)", r"\1", r["District Name (if applicable)"])
        sk = seat_key(district)
        if district and not sk:
            stop(f"row {n} of the filing list words a district in a way this loader does not know; stopping")
        m = re.match(r"(Subdistrict|Subdivision|Subcouncil|District)\b", district)
        if m:
            b["label"][m.group(1)] += 1
        b["seats"].setdefault(sk, {})
        b["seats"][sk].setdefault(int(r["Term"]), {"district": district, "vote_for": int(r["Vote For"]), "names": []})
        s = b["seats"][sk][int(r["Term"])]
        if s["district"] != district or s["vote_for"] != int(r["Vote For"]):
            stop(f"row {n} of the filing list: one seat written two ways; stopping")
        if nkey(r["Candidate Name"]) in [nkey(x) for x in s["names"]]:
            stop(f"row {n} of the filing list repeats a name already in its contest; stopping")
        s["names"].append(r["Candidate Name"])
        counts["board candidates"] += 1
    if counts["blank names"]:
        problems.append(f"{counts['blank names']} board rows of the filing list have no name (blanked because the cell read like contact details, or empty)")
    usual = {}                                                      # the usual term of each kind of board: the commonest on the list
    for key, b in boards.items():
        usual.setdefault(key[0], Counter()).update(t for seat in b["seats"].values() for t in seat)
    usual = {cls: c.most_common(1)[0][0] for cls, c in usual.items()}

    # ---- 2. where each seat is voted on: this year's contests-by-county sheet, then the last three general elections
    where_2026 = defaultdict(set)
    for n, r in enumerate(byco["rows"], start=2):
        if r["Contest"].startswith("For ") or not r["Contest"]:
            continue
        key = board_key(re.sub(r"\s+For Board of (?:Directors|Governors)$", "", r["Contest"]))
        if r["County"] not in by_county:
            stop(f"row {n} of the State Contests by County sheet names a county the Census Bureau's county file does not have; stopping")
        if key:
            where_2026["|".join(key + (seat_key(r["District"]),))].add(by_county[r["County"]])
    books = []
    for year, day, url in CANVASS_BOOKS:
        try:
            books.append((year, day, canvass_book(folder, year, url, cmap, say)))
        except Exception as err:  # noqa: BLE001  a check and a fallback only: the loader goes on without that year
            problems.append(f"the {year} general election canvass book could not be read ({type(err).__name__}: {err}); seats are placed without it")

    def placed(key, sk):
        """(counties, how they are known) for one seat of one board."""
        k = "|".join(key + (sk,))
        if where_2026.get(k):
            return set(where_2026[k]), "2026"
        for year, day, book in books:
            if book["tables"].get(k):
                return set(book["tables"][k]), day
        return set(), None

    # ---- 3. the counties' own files
    try:
        exact, bare, towns, ppath, spath = census_places(folder, cmap, say)
        schools, upath = census_schools(folder, say)
    except OSError as err:
        stop(f"the Census Bureau's place and school district files could not be fetched ({type(err).__name__}); nothing was changed. Re-run.")
    try:
        offices = election_offices(folder, cmap, say)["offices"]
    except Exception as err:  # noqa: BLE001  the gaps then point at the directory itself
        offices = {}
        problems.append(f"the directory of county election offices could not be read ({type(err).__name__}: {err}); each county's gap gives the directory's address")
    known = set(boards)
    sos_index = {(key, sk): seat for key, b in boards.items() for sk, seat in b["seats"].items()}
    confirmed = defaultdict(set)                                    # (board key, seat key, term) -> counties whose own file lists it
    printed_by = defaultdict(set)                                   # the same, whether or not the names agreed
    local, extra_boards = {}, {}                                    # race id -> race (county-filed); empty board contests a county prints
    places, files, county_gap, skipped = {}, {}, {}, {}
    fallback_school = {}
    county_counts = {}

    def place_row(kind, pid, name, cids, src):
        p = places.setdefault((kind, pid), {"name": name, "counties": set(), "src": src})
        p["counties"].update(cids)

    for spec in COUNTY_FILES:
        fips, county, reader, fkind, title, url, _name = spec
        src = f"ne-{slug(county)}-2026-general-" + ("sample-ballot" if fkind == "sample ballot" else "notice")
        got, why = county_file(folder, spec, say, refresh)
        if got is None:
            county_gap[fips] = f"This run could not read the county's {fkind}: {why}."
            problems.append(f"{county} County: not loaded: {why}"
                            + (f". If a browser can open {url}, save the file as {os.path.join(folder, slug(county), _name)} and run again"
                               if "did not give the file" in why else ""))
            continue
        unknown, plan, c_counts = [], [], Counter()
        for c in got["contests"]:
            info = classify(c["head"], county)
            if info is None:
                unknown.append(" | ".join(c["head"]))
                continue
            if info["cat"] != "state":
                bad = [nm for nm, _p in c["cands"] if not NAMELIKE.match(nm) or len(nm) > 60 or len(nm.split()) > 5 or nm.startswith("For ")
                       or re.search(r"\b(Board of|City Council|School District|Public Schools?|Village of|Vote for|Year Term|Ticket|candidates?|filed|"
                                    r"office|vacan\w*)\b", nm, re.I)]
                if bad:
                    unknown.append(" | ".join(c["head"]) + " (a name that reads like a heading)")
                    continue
                plan.append((c, info))
            c_counts[info["cat"]] += 1
            c_counts[info["cat"] + " names"] += len(c["cands"])
        if len(unknown) > 2:                                # more than a stray heading: the file is not what its reader expects
            county_gap[fips] = (f"The county's {fkind} was fetched, but {len(unknown)} of its headings are worded in a way this loader does not read yet, "
                                "so none of its contests is loaded.")
            problems.append(f"{county} County: not loaded: headings not understood: " + "; ".join(unknown[:6]))
            continue
        if unknown:                                         # one or two: the rest is loaded, and the gap names what is not
            skipped[fips] = unknown
            problems.append(f"{county} County: {len(unknown)} contest(s) left out, the heading not understood: " + "; ".join(unknown))
        files[fips] = dict(got, src=src, spec=spec, counts=c_counts)
        for c, info in plan:
            names = [nm for nm, _p in c["cands"]]
            if len({nkey(x) for x in names}) != len(names):
                problems.append(f"{county} County: a name twice in one contest ({' | '.join(c['head'])}); kept once")
            if info["cat"] == "board":
                key, sk = match_board(info["text"], known), heading_seat(info["text"])
                label = " | ".join(c["head"])
                if not key:
                    problems.append(f"{county} County prints a district board the Secretary of State's list does not have: {label}")
                    continue
                seat = sos_index.get((key, sk)) or {}
                terms = [t for t in seat if info["term"] in (None, t)]
                if not terms:
                    if names:
                        problems.append(f"{county} County prints candidates for a seat the Secretary of State's list does not have: {label}")
                    else:                                           # a seat nobody filed for: the county's file is the only list that shows it
                        e = extra_boards.setdefault((key, sk, info["term"]), {"counties": set(), "srcs": set(), "label": label})
                        e["counties"].add(fips)
                        e["srcs"].add(src)
                    continue
                fit = [t for t in terms if {nkey(x) for x in seat[t]["names"]} == {nkey(x) for x in names}]
                term = fit[0] if len(fit) == 1 else terms[0] if len(terms) == 1 else None
                if term is None:
                    problems.append(f"{county} County: which of two terms a heading means could not be told: {label}")
                    continue
                printed_by[(key, sk, term)].add(fips)
                a, b_ = {nkey(x) for x in seat[term]["names"]}, {nkey(x) for x in names}
                if a == b_ or a & b_:
                    confirmed[(key, sk, term)].add(fips)
                if a != b_:
                    twin = [boards[k]["name"] for k in boards if k != key and k[0] == key[0]
                            and any({nkey(x) for x in s2["names"]} == b_ for s2 in boards[k]["seats"].get(sk, {}).values())] if b_ else []
                    problems.append(f"{county} County's {fkind} and the Secretary of State's list name different candidates for {label}: "
                                    f"the county {sorted(names)}, the list {sorted(seat[term]['names'])}"
                                    + ("" if a & b_ else "; no name agrees, so the county is not added to the seat's counties")
                                    + (f" (the county's names are the list's for the same seat of {and_names(twin)}, so the heading is likely misprinted)" if twin else ""))
                continue
            # a county-filed contest
            kind, level = info["kind"], info["level"]
            jur = info["jur"]
            note_place = None
            if jur[0] == "county":
                key, jname, jid, cids = fips, full[fips], fips, {fips}
            elif jur[0] == "place":
                hits = exact.get((fold(jur[1]), jur[2]), [])
                other_kind = [h for h in bare.get(fold(jur[1]), []) if h not in hits]
                if len(hits) == 1:
                    code, cname, ccs = hits[0]
                elif not hits and len(other_kind) == 1:             # the Census Bureau's 2020 list files it under the other kind word
                    code, cname, ccs = other_kind[0]
                    note_place = (f"The county's {fkind} calls {jur[1]} a {jur[2]}; the Census Bureau's 2020 list of places calls it "
                                  f"a {cname.rsplit(' ', 1)[1]}.")
                    cname = f"{cname.rsplit(' ', 1)[0]} {jur[2]}"
                else:
                    code, cname, ccs = None, f"{jur[1]} {jur[2]}", []
                key = f"M-{code}" if code else f"M-{fips[2:]}-{slug(cname)}"
                jname, jid, cids = cname, f"{STATE}-{key}", set(ccs) | {fips}
                place_row("mcd", jid, jname, cids, SRC_PLACE if code else src)
            elif jur[0] == "township":
                hit = towns.get((fips, jur[1].lower()))
                key = f"M-{hit[0]}" if hit else f"M-{fips[2:]}-{slug(jur[1])}"
                jname, jid, cids = f"{jur[1]} ({full[fips]})", f"{STATE}-{key}", {fips}
                place_row("mcd", jid, jname, cids, SRC_COUSUB if hit else src)
            elif jur[0] == "school":
                core = school_core(jur[2])
                hits = schools.get(core, [])
                if len(hits) == 1:
                    key, jname, psrc = f"S-{hits[0][0]}", hits[0][1], SRC_UNSD
                else:
                    key = fallback_school.setdefault(core, f"S-{fips[2:]}-{slug(core or jur[1])}")
                    jname, psrc = jur[1] if re.search(r"school", jur[1], re.I) else f"{jur[1]} (school district)", src
                    if places.get(("school", f"{STATE}-{key}")):
                        jname = places[("school", f"{STATE}-{key}")]["name"]
                jid, cids = f"{STATE}-{key}", {fips}
                place_row("school", jid, jname, cids, psrc)
            else:
                jname = jur[1]
                key = f"X-{fips[2:]}-{slug(jname)}"
                jid, cids = f"{STATE}-{key}", {fips}
                place_row("special", jid, jname, cids, src)
            term = info["term"]
            special = 1 if term and term < USUAL_TERM.get(kind, 4) else 0
            rid = race_id_for(key, kind, info["district"] or info["seat"], special)
            partisan = 1 if level == "county" else 0
            parties = [p for _nm, p in c["cands"]]
            if (partisan and not all(parties)) or (not partisan and any(parties)):
                problems.append(f"{county} County: {' | '.join(c['head'])}: " + ("a candidate with no party on the partisan ballot" if partisan
                                                                                   else "a party printed for a nonpartisan office") + "; the contest is left out")
                c_counts["left out"] += 1
                continue
            r = local.get(rid)
            if r is None:
                r = local[rid] = {"level": level, "kind": kind, "office": info["office"], "jur": jname, "jid": jid, "counties": set(), "district": info["district"],
                                  "seat": info["seat"], "special": special, "partisan": partisan, "term": term, "vote_for": c["vote_for"], "cands": {},
                                  "none": False, "by": {}, "src": src, "place_note": note_place, "fkind": fkind}
            elif (r["office"], r["jid"], r["district"], r["seat"]) != (info["office"], jid, info["district"], info["seat"]):
                problems.append(f"{county} County: two different contests share the race id {rid}; the second is left out")
                c_counts["left out"] += 1
                continue
            if fips in r["by"]:
                problems.append(f"{county} County prints {rid} twice; the second printing is left out")
                c_counts["left out"] += 1
                continue
            r["counties"].update(cids)
            r["by"][fips] = [nkey(x) for x in names]
            r["none"] = r["none"] or c["none"]
            if r["vote_for"] is None:
                r["vote_for"] = c["vote_for"]
            elif c["vote_for"] is not None and c["vote_for"] != r["vote_for"]:
                problems.append(f"{rid}: {county} County says vote for {c['vote_for']}, another county {r['vote_for']}")
            for nm, party in c["cands"]:
                r["cands"].setdefault(nkey(nm), (nm, PARTY_WORDS.get(party, party), src))
            c_counts["loaded"] += 1
            c_counts["loaded names"] += len(c["cands"])
        if (len(got["contests"]) != c_counts["state"] + c_counts["board"] + c_counts["local"] + len(unknown)
                or len(plan) != c_counts["board"] + c_counts["local"] or c_counts["local"] != c_counts["loaded"] + c_counts["left out"]):
            problems.append(f"{county} County: {len(got['contests'])} contests read, {c_counts['state']} state, {c_counts['board']} district board and "
                            f"{c_counts['local']} local ones counted, {c_counts['loaded']} loaded and {c_counts['left out']} left out: the counts do not add up")
        county_counts[fips] = c_counts
    for rid, r in local.items():                                    # a contest two counties print must name the same people
        sets = {tuple(sorted(v)) for v in r["by"].values()}
        if len(sets) > 1:
            problems.append(f"{rid}: the counties that print this contest ({and_names(full[f] for f in sorted(r['by']))}) do not name the same candidates; "
                            "every name any of them prints is kept")

    # ---- 4. rows: the district boards
    race_rows, cand_rows, gaps = [], [], []
    how_counts, board_races = Counter(), 0
    board_counties = defaultdict(set)
    seat_place = {}
    for key, b in boards.items():
        for sk, seat in b["seats"].items():
            for term in seat:
                base, how = placed(key, sk)
                own = confirmed.get((key, sk, term), set())
                seat_place[(key, sk, term)] = (base | own, how, own - base)
                board_counties[key] |= base | own
    for (key, sk, term), e in extra_boards.items():
        board_counties[key] |= e["counties"] | placed(key, sk)[0]
    n_2026 = "this year's contests-by-county sheet"

    def county_words(how, more):
        """The sentence saying how a seat's counties are known, where this year's sheet does not carry the seat."""
        return ("The Secretary of State's candidate list does not say which counties vote on this seat; the counties given are those where it was on the "
                f"ballot at the general election of {how}, by the Board of State Canvassers' report"
                + (f", with {and_names(full[f] for f in sorted(more))} added from the county's own notice or sample ballot this year." if more else "."))

    for key, b in sorted(boards.items(), key=lambda kv: kv[1]["name"]):
        cls, level, kind, office = b["kind"]
        label = b["label"].most_common(1)[0][0] if b["label"] else "District"
        for sk, seat in sorted(b["seats"].items(), key=lambda kv: (not kv[0].isdigit(), int(kv[0]) if kv[0].isdigit() else 0, kv[0])):
            short = [t for t in seat if t < usual[cls]]
            for term, s in sorted(seat.items(), reverse=True):
                special = 1 if term < usual[cls] else 0
                rid = race_id_for(b["id"][len(STATE) + 1:], kind, s["district"], special) + (str(term) if special and len(short) > 1 else "")
                cs, how, more = seat_place[(key, sk, term)]
                notes = []
                if s["vote_for"] > 1:
                    notes.append(f"Voters choose up to {s['vote_for']}.")
                if special:
                    notes.append(f"For a {term}-year term on the Secretary of State's list; seats on boards of this kind are usually filled for {usual[cls]} years.")
                if how == "2026":
                    how_counts[n_2026] += 1
                elif how:
                    how_counts[f"the {how[-4:]} general election's canvass"] += 1
                    notes.append(county_words(how, more))
                elif cs:
                    how_counts["a county's own notice or sample ballot"] += 1
                    notes.append("The Secretary of State's candidate list does not say which counties vote on this seat; the counties given are those whose own "
                                 "notice or sample ballot lists it this year, and there may be others.")
                else:
                    cs = set(board_counties[key])
                    how_counts["the board's other seats"] += 1
                    if cs:
                        notes.append("No official list read here says which counties vote on this seat. It is shown for every county where another seat of this "
                                     "board is voted on, and may not reach all of them.")
                        gaps.append((STATE, "race", rid, b["name"], "which counties vote on this seat",
                                     "The Secretary of State's candidate list names the board and the seat but no county, the seat is not on this year's "
                                     "contests-by-county sheet, and it was not on the ballot at the last three general elections; it is shown for every county "
                                     "where another seat of the board is voted on.", WORKBOOK_URL))
                if not cs:
                    gaps.append((STATE, "race", rid, b["name"], f"the contest for {s['district'] or 'this board'}",
                                 "The Secretary of State's candidate list names this contest but no county, and no other official list read here places the "
                                 "board in any county, so the contest cannot be shown on a county's page yet.", WORKBOOK_URL))
                    counts["board contests with no county"] += 1
                    counts["board candidates with no county"] += len(s["names"])
                    continue
                board_races += 1
                race_rows.append((rid, STATE, level, kind, office, b["name"], b["id"], json.dumps(sorted(cs)), s["district"] or None, None, special, 0,
                                  None, None, None, GENERAL, " ".join(notes) or None))
                cand_rows += [(rid, "general", GENERAL, nm, NONPARTISAN, "N", None, 0, 0, None, None, None, None, SRC_WB, None) for nm in s["names"]]
                counts["board candidates placed"] += len(s["names"])
        for (k2, sk, term), e in sorted(extra_boards.items(), key=lambda kv: str(kv[0])):
            if k2 != key:
                continue
            district = "At Large" if sk == "al" else f"{label} {sk}" if sk.isdigit() else f"{sk.title()} Subdivision" if sk else ""
            special = 1 if term and term < usual[cls] else 0
            rid = race_id_for(b["id"][len(STATE) + 1:], kind, district, special)
            if any(r[0] == rid for r in race_rows):
                problems.append(f"{rid}: a county prints this seat with no candidate, but the Secretary of State's list has candidates for it")
                continue
            base, how = placed(key, sk)
            cs = base | e["counties"]
            where_words = ("" if how == "2026" else " " + county_words(how, e["counties"] - base) if how
                           else " The county given is the one whose own file prints the contest; there may be others.")
            race_rows.append((rid, STATE, level, kind, office, b["name"], b["id"], json.dumps(sorted(cs)), district or None, None, special, 0, None, None, None,
                              GENERAL, f"No candidate is named for this seat: the Secretary of State's candidate list has none, and the sample ballot or notice of "
                                       f"{and_names(full[f] for f in sorted(e['counties']))} prints the contest with no name."
                                       + (f" Listed for a {term}-year term." if special else "") + where_words))
            counts["empty board contests"] += 1
            board_counties[key] |= cs

    # ---- 5. rows: the county-filed contests
    for rid, r in sorted(local.items()):
        notes = []
        many = r["kind"] in MANY and not r["district"]
        if r["vote_for"] and r["vote_for"] > 1:
            notes.append(f"Voters choose up to {r['vote_for']}.")
        elif r["vote_for"] is None and many:
            notes.append(f"The county's {r['fkind']} does not say how many are elected.")
        if r["term"] and r["term"] != USUAL_TERM.get(r["kind"], 4):
            notes.append(f"A {r['term']}-year term, as the county's {r['fkind']} prints it.")
        if not r["cands"]:
            notes.append(f"The county's {r['fkind']} shows this contest with no candidate named.")
        if r["place_note"]:
            notes.append(r["place_note"])
        if len(r["by"]) > 1:
            notes.append(f"Printed in the notices or sample ballots of {and_names(full[f] for f in sorted(r['by']))}.")
        race_rows.append((rid, STATE, r["level"], r["kind"], r["office"], r["jur"], r["jid"], json.dumps(sorted(r["counties"])), r["district"], r["seat"],
                          r["special"], r["partisan"], None, None, None, GENERAL, " ".join(notes) or None))
        for _k, (nm, party, src) in r["cands"].items():
            if r["partisan"]:
                cand_rows.append((rid, "general", GENERAL, nm, party, "I" if party == "By Petition" else party_code(party), None, 0, 0, None, None, None, None, src, None))
            else:
                cand_rows.append((rid, "general", GENERAL, nm, NONPARTISAN, "N", None, 0, 0, None, None, None, None, src, None))

    # ---- 6. rows: the judges' retention votes (level court)
    jd = None
    try:
        jd = judicial_districts(folder, cmap, say)
    except Exception as err:  # noqa: BLE001  the judges are still loaded, without their counties
        problems.append(f"the statutes that list each judicial district's counties could not be read ({type(err).__name__}: {err}); the judges carry no counties")
    court_rows, court_cands = [], []
    judge_counties = defaultdict(set)                               # a judge's name -> the counties whose own files print the retention vote
    for fips, f in files.items():
        for _court, judge in f["retention"]:
            judge_counties[nkey(judge)].add(fips)
    for n, r in enumerate(wb["retention"], start=2):
        court = r["Office"].replace("’", "'")
        if court not in RETENTION or not r["Judge (Ballot Name)"]:
            stop(f"row {n} of the Judicial Retention sheet names a court this loader does not know, or no judge; stopping (the row is not printed)")
        code, kind, statute = RETENTION[court]
        d, judge = r["District (if applicable)"], r["Judge (Ballot Name)"]
        cids, jur, jid, district = None, NAME, FIPS, None
        m = re.fullmatch(r"(\d+)(?: - ([A-Za-z ]+))?", d)
        if code in ("DC", "CC"):
            if not m or m.group(2):
                stop(f"row {n} of the Judicial Retention sheet words a district in a way this loader does not know; stopping")
            district = m.group(1)
            jur, jid = (f"District Court Judicial District {district}" if code == "DC" else f"County Judge District {district}"), f"{STATE}-{code}{district}"
            cids = jd["statutes"][statute]["districts"].get(district) if jd else None
        elif code == "JUV":
            if not m or not m.group(2) or m.group(2).strip() not in by_county:
                stop(f"row {n} of the Judicial Retention sheet names a juvenile court's county in a way this loader does not know; stopping")
            jur, jid, cids = full[by_county[m.group(2).strip()]], by_county[m.group(2).strip()], [by_county[m.group(2).strip()]]
        elif code == "COA":
            if not m or m.group(2):
                stop(f"row {n} of the Judicial Retention sheet words a Court of Appeals district in a way this loader does not know; stopping")
            district = m.group(1)
            jur, jid = f"Court of Appeals District {district}", f"{STATE}-COA{district}"
        elif d != "Statewide":
            stop(f"row {n} of the Judicial Retention sheet gives the Workers' Compensation Court a district; stopping")
        rid = f"2026-{STATE}-RET-{code}{district or ''}-{slug(judge)}"
        notes = ["A retention vote: voters answer Yes or No on keeping this judge in office."]
        if r["Term"] and r["Term"] != "6":
            notes.append(f"Listed for a {r['Term']}-year term.")
        if code == "COA":
            cids = sorted(judge_counties.get(nkey(judge), ())) or None
            notes.append("Only the voters of this Court of Appeals district vote on it. Its lines follow the Supreme Court's judicial districts, which the "
                         "Secretary of State's list does not give by county"
                         + (f"; the county given is the one whose own notice prints this vote ({and_names(full[f] for f in cids)}), and there may be others."
                            if cids else ", so no county is given."))
        court_rows.append((rid, STATE, "court", kind, f"{court} (retention vote)", jur, jid, json.dumps(cids) if cids else None, district, None, 0, 0,
                           None, judge, None, GENERAL, " ".join(notes)))
        court_cands.append((rid, "general", GENERAL, judge, NONPARTISAN, "N", None, 1, 0, None, None, None, None, SRC_WB, "Standing for retention as the sitting judge."))
    if len({r[0] for r in court_rows}) != len(court_rows):
        stop("two judges on the Judicial Retention sheet read as one; stopping")
    # a check: a county's file names exactly the judges of the districts the statutes put the county in
    sheet_judges = {nkey(c[3]) for c in court_cands}
    for fips, f in files.items():
        if not f["retention"]:
            continue
        printed = {nkey(judge) for _court, judge in f["retention"]}
        for _court, judge in f["retention"]:
            if nkey(judge) not in sheet_judges:
                problems.append(f"{full[fips]}'s file names a judge for retention who is not on the Secretary of State's Judicial Retention sheet: {judge}")
        expected = {nkey(r[13]) for r in court_rows if (r[7] and fips in json.loads(r[7])) or r[3] == "workers_compensation_court_retention"}
        if printed & sheet_judges != expected:
            problems.append(f"{full[fips]}'s file names {len(printed & sheet_judges)} of the sheet's judges; the statutes' districts give the county "
                            f"{len(expected)}: only in the file {sorted(printed & sheet_judges - expected)}, only by the districts {sorted(expected - printed)}")

    # ---- 7. two routes to the same total: the filing list read as a workbook and as the printed list
    if pdf_local is not None:
        a = Counter((r["Office"], re.sub(r"\b0+(\d)", r"\1", r["District Name (if applicable)"]), r["Term"], nkey(r["Candidate Name"]))
                    for r in rows if r["Office"] not in FEDERAL and r["Office"] not in LIST_OFFICE)
        b_ = Counter((r.get("Office", ""), re.sub(r"\b0+(\d)", r"\1", r.get("District Name (if applicable)", "")), r.get("Term", ""),
                      nkey(r.get("Candidate Name", ""))) for r in pdf_local)
        if a != b_:
            problems.append(f"the filing list as a workbook ({sum(a.values())} board candidates) and as the printed Final Statewide General Candidate List "
                            f"({sum(b_.values())}) do not hold the same rows: {sum((a - b_).values())} only in the workbook, {sum((b_ - a).values())} only in the "
                            "printed list; the workbook is used")
        counts["checked against the printed list"] = sum((a & b_).values())
    if counts["board candidates"] != counts["board candidates placed"] + counts["board candidates with no county"]:
        stop(f"{counts['board candidates']} board candidates read, {counts['board candidates placed']} placed in a race; stopping")
    if len({(c[0], c[1], c[3]) for c in cand_rows}) != len(cand_rows):
        stop("a candidate is stored twice in one local race; stopping")
    if len({r[0] for r in race_rows}) != len(race_rows):
        stop("two local races share an id; stopping")

    # ---- 8. places, gaps, notes, sources
    place_rows = [("county", geoid, name, json.dumps([geoid]), SRC_COUNTIES) for geoid, name in sorted(full.items())]
    for key, b in sorted(boards.items(), key=lambda kv: kv[1]["id"]):
        if board_counties[key]:
            place_rows.append(("special", b["id"], b["name"], json.dumps(sorted(board_counties[key])), SRC_WB))
    place_rows += [(kind, pid, p["name"], json.dumps(sorted(p["counties"])), p["src"]) for (kind, pid), p in sorted(places.items())
                   if any(r[6] == pid for r in race_rows)]
    loaded = sorted(files)
    what = "county, city, village, school board and township races"
    for geoid, name in sorted(full.items()):
        if geoid in files:
            continue
        reason = county_gap.get(geoid) or ("Its " + what + " are not loaded yet. The county election office publishes them in its Notice of Election and its sample "
                                           f"ballot; {len(files)} counties have been read so far, starting with the largest whose files a program can read as text. "
                                           + COUNTY_SEEN.get(geoid, "")).strip()
        gaps.append((STATE, "county", geoid, name, what, reason, offices.get(geoid) or DIRECTORY_URL))
    for fips, heads in sorted(skipped.items()):
        gaps.append((STATE, "county", fips, full[fips], f"{len(heads)} of the county's contests",
                     f"The county's {files[fips]['spec'][3]} prints {'a contest' if len(heads) == 1 else 'contests'} under a heading this loader does not read yet "
                     f"({'; '.join(heads)}); the county's other contests are loaded.", offices.get(fips) or DIRECTORY_URL))

    by_level, by_kind = Counter(r[2] for r in race_rows), Counter(r[3] for r in race_rows)
    county_names = and_names(full[f][:-len(" County")] for f in loaded) if loaded else "no county"
    n_county_races = len(local)
    n_empty = sum(1 for r in local.values() if not r["cands"]) + counts["empty board contests"]
    notes = [
        (STATE, "local_calendar",
         "On November 3, 2026 Nebraska's counties elect, on the partisan ballot, their county officers (clerk, register of deeds, assessor, sheriff, treasurer, "
         "attorney, surveyor and, where the county elects them, engineer, public defender and clerk of the district court) and the county board seats whose terms "
         "are up; and on the nonpartisan ballot the township boards of counties under township government, village boards of trustees, the mayors and councils of "
         "most cities, school boards, airport authority boards, county weed district boards where they are elected, the Omaha-area transit board, and the boards "
         "of natural resources districts, public power districts, reclamation districts, community colleges, educational service units and the Douglas-Sarpy "
         "Learning Community; judges stand for a yes-or-no retention vote. Omaha and Lincoln elect their city officers, and Lincoln its school board, at city "
         "elections in the spring of odd-numbered years, and local hospital district boards are elected at the May primary, so those offices are not on this "
         "ballot.",
         "Neb. Rev. Stat. 32-512 to 32-551 (each office and when it is elected); 14-201 and 15-301 (Omaha's and Lincoln's city elections); 32-544 (Lincoln's "
         "school board); 32-550 (hospital districts); Secretary of State, 2026 Candidate Filing Guide", STATUTES_PAGE),
        (STATE, "local_coverage",
         f"Loaded for the whole state, from the Secretary of State's Statewide Candidate Filing List: {board_races + counts['empty board contests']} contests for the "
         f"boards of {len([k for k in boards if board_counties[k]])} natural resources, public power, reclamation, community college, educational service unit and other "
         f"districts ({counts['board candidates placed']} candidates), and the {len(court_rows)} judges standing for retention. Loaded county by county, from the county "
         f"election office's own Notice of Election or sample ballot: {n_county_races} county, city, village, school board, township and other contests in "
         f"{len(loaded)} of the 93 counties ({county_names}), {n_empty} of all these contests with no candidate named. The other {93 - len(loaded)} counties' own "
         "contests are not loaded yet; each is named among the gaps. Left out everywhere: ballot questions and measures (bond issues, sales taxes, recalls, "
         "constitutional amendments and initiatives). The lists give no ballot order (names rotate by precinct) and name no write-in candidates, and a county's "
         "notice has no line for a candidate who withdrew, so none can be counted. A school district or a city that reaches into a county not yet read is shown "
         "only for the counties read.",
         "Nebraska Secretary of State, Statewide Candidate Filing List; county election offices' notices of election (Neb. Rev. Stat. 32-802) and sample ballots",
         WORKBOOK_URL),
    ]

    def src_file(path, sid, kind, agency, title, url, n, note):
        return (sid, STATE, kind, agency, title, url, "", mdate(path), sha_of(path), n, note)

    sources = [
        (SRC_WB, STATE, "official candidate list", "Nebraska Secretary of State", "Statewide Candidate Filing List (workbook): sheets General Election Candidates "
         "and Judicial Retention", wb["url"], "", wb["fetched"], wb["sha256"], counts["board rows"] + len(wb["retention"]),
         f"Fetched into memory and never saved whole; the fingerprint is of the workbook as it came ({wb['bytes']:,} bytes). Seven columns of the General Election "
         "Candidates sheet are read, each by its heading (Office, District Name, Term, Vote For, Party, Candidate Name, Incumbency Status); City of Residence, "
         f"Mailing Address and Phone/Email are never read. Of its {len(rows)} rows, {counts['board rows']} are candidates for district boards (loaded here), "
         f"{counts['state']} for state offices and the Legislature and {counts['federal']} for Congress (read by the state and federal loaders from the printed "
         f"list). The Judicial Retention sheet has no contact columns; its {len(wb['retention'])} judges are loaded as retention votes. The Candidate Petitions "
         f"sheet (petition candidates for Congress and Governor) is not read. {wb['blanked']} kept cells read like contact details and were blanked. The list "
         "gives no ballot order, no withdrawn candidates and no write-ins."),
        (SRC_BYCO, STATE, "official contest list", "Nebraska Secretary of State", byco["title"], byco["url"], "", byco["fetched"], byco["sha256"], len(byco["rows"]),
         "Which state-filed contests were on which county's ballot in the 2026 primary cycle: three columns read (County, Contest, District); the workbook has no "
         f"contact columns. It places {how_counts[n_2026]} of the district boards' November contests in their counties."),
    ]
    for year, day, book in books:
        used = how_counts[f"the {year} general election's canvass"]
        sources.append((f"ne-sos-{year}-general-canvass", STATE, "official results", "Nebraska Board of State Canvassers (compiled by the Secretary of State)",
                        f"Official Report of the Board of State Canvassers: General Election, {day}", book["url"], "", book["fetched"], book["sha256"], len(book["tables"]),
                        f"Read only for where each district board's seat was on the ballot: the county names standing in each seat's table ({len(book['tables'])} seats). "
                        f"No candidate's name and no vote is read or kept; only that list of counties is kept on disk. Used for {used} seats that this year's "
                        "contests-by-county sheet does not carry; each such race says so."))
    if jd:
        for sec, s in jd["statutes"].items():
            sources.append((f"ne-statute-{sec.replace('.', '-')}", STATE, "statute", "Nebraska Legislature",
                            f"Neb. Rev. Stat. {sec}: " + ("district court judicial districts" if sec == "24-301.02" else "county judge districts"), s["url"], "",
                            jd["fetched"], s["sha256"], len(s["districts"]), "The counties of each of the twelve districts, for the judges' retention votes."))
    dirx = read_json(os.path.join(folder, "sos_county_election_offices.json")) or {}
    sources += [
        (SRC_DIR, STATE, "directory", "Nebraska Secretary of State", "County election offices (directory)", DIRECTORY_URL, "", dirx.get("fetched", ""), dirx.get("sha256", ""),
         len(offices), "Read for two things only: each county's name and the address of its election office's page, which each county's gap gives. The officials' "
         "names, addresses, telephone numbers and e-mail on the page are never read or kept."),
        src_file(ppath, SRC_PLACE, "official place codes", "U.S. Census Bureau", "2020 place codes, Nebraska (st31_ne_place2020.txt)", PLACE_URL,
                 sum(len(v) for v in bare.values()), "Names, codes and counties of Nebraska's incorporated cities and villages; a city or village a county's file "
                 "names is given this code when the name fits exactly one entry."),
        src_file(spath, SRC_COUSUB, "official place codes", "U.S. Census Bureau", "2020 county subdivision codes, Nebraska (st31_ne_cousub2020.txt)", COUSUB_URL,
                 len(towns), "Names and codes of the townships of counties under township government."),
        src_file(upath, SRC_UNSD, "official boundaries (attributes)", "U.S. Census Bureau", "Cartographic boundary file, unified school districts, Nebraska, 2024 "
                 "(cb_2024_31_unsd_500k)", UNSD_URL, sum(len(v) for v in schools.values()), "School districts' names and codes, from the file's attribute table (no shape "
                 "is read). A district a county's file names is given this code when its name, without its number and kind words, fits exactly one entry."),
    ]
    for fips, f in sorted(files.items(), key=lambda kv: [s[0] for s in COUNTY_FILES].index(kv[0])):
        _f, county, reader, fkind, title, url, _n = f["spec"]
        cc = f["counts"]
        sources.append((f["src"], STATE, "official " + fkind, f"{county} County election office", f"{county} County: {title}", url, "", f["fetched"], f["sha256"],
                        sum(len(c["cands"]) for c in f["contests"]),
                        f"{'Fetched' if f.get('how') != 'saved by hand' else 'Saved by hand and read from disk'}; the file has no address, telephone or e-mail column, and "
                        f"only office headings, names and parties are kept (the fingerprint is of the file as it came, {f['bytes']:,} bytes). Contests read: "
                        f"{len(f['contests'])}. Loaded here: {cc['loaded']} county, city, village, school, township and other local contests "
                        f"({cc['loaded names']} names). Checked against the Secretary of State's list: {cc['board']} district board contests. Left to the state "
                        f"loader: {cc['state']} federal and state contests. Checked against the Judicial Retention sheet: {len(f['retention'])} judges. "
                        + (f"Not loaded, and said in the run's checks: {cc['left out'] + len(skipped.get(fips, []))} contests. "
                           if cc["left out"] or skipped.get(fips) else "")
                        + ("The federal and state tickets at the head of the notice are not read. " if reader == "washington" else "")
                        + ("The notice says a voter votes for one in each contest unless it says otherwise. " if reader == "lancaster" else "")
                        + ("The notice does not say how many are elected to a board. " if reader == "douglas" else "")
                        + "Ballot questions are not read."))

    # the last look before anything is written: nothing that reads like contact details, by the trial check's own test and the page builder's
    for table, strict, items in (("sl_races", True, [(r[0], (r[4], r[5], r[8], r[9], r[16], r[13])) for r in race_rows + court_rows]),
                                 ("sl_candidates", True, [(c[0], (c[3], c[4], c[14])) for c in cand_rows + court_cands]),
                                 ("sl_places", True, [(p[1], (p[2],)) for p in place_rows]),
                                 ("sl_gaps", False, [(g[2], (g[3], g[4], g[5])) for g in gaps]),
                                 ("sl_notes", False, [(x[1], (x[2], x[3])) for x in notes]),
                                 ("sl_sources", False, [(s[0], (s[3], s[4], s[10])) for s in sources])):
        for ident, texts in items:
            if any(t and (contact_like(t, strict) or (strict and BUILDER_STREET.search(str(t)))) for t in texts):
                stop(f"a text for {table} ({ident}) reads like contact details; stopping (the text is not printed)")

    return {"races": race_rows, "cands": cand_rows, "court_races": court_rows, "court_cands": court_cands, "places": place_rows, "sources": sources, "gaps": gaps,
            "notes": notes, "problems": problems, "counts": counts, "by_level": by_level, "by_kind": by_kind, "how": how_counts, "loaded": loaded,
            "county_counts": county_counts, "county_races": n_county_races, "board_races": board_races, "rows": len(rows), "fetched": wb["fetched"],
            "files": {f: (v["spec"][1], len(v["contests"]), sum(len(c["cands"]) for c in v["contests"]), len(v["retention"])) for f, v in files.items()},
            "boards": len(boards), "printed_by": {"|".join(k[0] + (k[1], str(k[2]))): sorted(v) for k, v in printed_by.items()}}


# ---------- the load ----------

def load(db_path, say=print, cache=CACHE, keep=None, local_cache=None, refresh=False):
    """Nebraska's rows into the database at db_path: the state races, then the local ones. local_cache is the folder for the
    cut-down local sources (ballot_cache/ne/local/ unless told otherwise); refresh=True asks for the filing list and
    the counties' files again even when the copies on disk are less than a week old."""
    keep = keep or os.path.join(cache, "ne")
    seats, offices, as_of = roster()
    url = list_links(say)
    book_path = os.path.join(cache, "ne", BOOK_FILE)
    if not os.path.exists(book_path):
        if not url.get("canvass"):
            fail("the canvass book is not cached and its address could not be found")
        net.download(url["canvass"], book_path, max_age_days=30, say=say)
    printed, strips, listed, local, list_sha, list_fetched, list_how = get_list(cache, keep, url.get("list"), 7, say)
    tables, terms, places = canvass(book_path)
    cmap = counties()
    checks = []
    pdf_local = None                                         # the district boards' strips of the printed list, to check the workbook against
    if os.path.exists(os.path.join(cache, "ne", LIST_FILE)):
        try:
            pdf_local = [r for r in read_list(open(os.path.join(cache, "ne", LIST_FILE), "rb").read())[1]
                         if r.get("Office", "") not in LIST_OFFICE and r.get("Office", "") not in FEDERAL]
        except SystemExit:
            pdf_local = None

    # ---- races: every state contest on the November list
    races, gen = {}, defaultdict(list)
    for r in listed:
        office = LIST_OFFICE[r["Office"]]
        key, level, kind, shown, partisan, by_district, roster_office, usual = OFFICES[office]
        d = None
        if by_district:
            m = re.fullmatch(r"District 0*(\d+)", r["District Name (if applicable)"])
            if not m:
                fail(f"a district the loader cannot read for {shown}: {r['District Name (if applicable)']!r}")
            d = m.group(1)
        elif r["District Name (if applicable)"]:
            fail(f"a district on a statewide office's strip ({shown})")
        if r["Vote For"] != "1" or not r["Candidate Name"]:
            fail(f"a candidate strip for {shown} {d or ''} did not read whole")
        if bool(partisan) != bool(r["Party (if applicable)"]):
            fail(f"a strip for {shown} {d or ''} {'has no party' if partisan else 'carries a party'}")
        if r["Incumbency Status"] not in ("Incumbent", "Nonincumbent"):
            fail(f"an incumbency mark the loader does not know ({shown} {d or ''})")
        rid = race_id(office, d)
        term = int(r["Term"]) if r["Term"].isdigit() else None
        info = races.setdefault(rid, {"office": office, "level": level, "office_kind": kind, "shown": shown,
                                      "partisan": partisan, "district": d, "roster": roster_office, "usual": usual,
                                      "term": term})
        if info["term"] != term:
            fail(f"the candidates for {rid} are listed with different terms")
        gen[rid].append(r)

    # ---- holders
    holders, notes = {}, defaultdict(list)
    for rid, i in races.items():
        if i["office"] == "leg":
            holders[rid] = seats.get(i["district"], [])
            if not holders[rid]:
                notes[rid].append(f"The Open States roster ({as_of}) lists no sitting senator for this district.")
        elif i["roster"]:
            h = offices.get(i["roster"])
            holders[rid] = [h] if h else []
            if not h:
                notes[rid].append("The Open States roster lists no holder of this office today.")
        else:
            holders[rid] = []
            notes[rid].append("The Open States roster this site uses does not carry this office, so today's holder is not "
                              "shown; the Secretary of State's list marks which candidate, if any, is the incumbent.")

    # ---- terms and special elections, from the list, checked against the canvass book's headings
    for rid, i in races.items():
        t, book = i["term"], terms.get(rid)
        book_years = TERM_WORDS.get(book.split()[0]) if book else None
        if book and book_years != t:
            checks.append(f"{rid}: the November list gives a {t}-year term, the canvass book's heading says \"{book}\"")
            notes[rid].append(f"The Secretary of State's November list gives this seat a {t}-year term; the canvass book's "
                              f"heading for the May 12 primary says \"{book}\".")
            i["special"] = 0
        else:
            i["special"] = 1 if t and t < i["usual"] else 0
            if i["special"]:
                notes[rid].append(f"An election for a {t}-year term, not the usual {i['usual']} years (the November list and "
                                  "the canvass book both say so).")
    for rid, i in races.items():
        if i["office"] == "leg":
            notes[rid].insert(0, "Nebraska's Legislature has one house, elected on a nonpartisan ballot: candidates carry no "
                                 "party. The top two from the May 12 primary are on the November ballot.")
        elif not i["partisan"]:
            notes[rid].insert(0, "A nonpartisan office: the top two from the May 12 primary are on the November ballot.")
    notes[race_id("governor")].insert(0, "Nebraska elects the Governor and Lieutenant Governor together, on one vote. A "
                                         "party's winner in the May 12 primary then chose a running mate (the canvass book's "
                                         "rule, section 32-619.01); the list names each ticket, the candidate for Governor first.")

    everyone = [p for ps in seats.values() for p in ps] + list(offices.values())

    def where_serves(p):
        return (f"Serves today in the Nebraska Legislature, District {p['district']}." if p["district"]
                else f"Serves today as {p['label']}.")

    def sitting(rid, names):
        """{name: (member id, incumbent, note)}: a one-to-one fit with the seat's holders; else exactly one other roster
        person (another district's senator, or a statewide official)."""
        out, hs = {}, holders[rid]
        heads = {n: ticket_head(n) if rid.endswith("-GOV") else n for n in names}
        pairs = [(n, h) for n in names for h in hs if person_fits(heads[n], h)]
        for n, h in pairs:
            if sum(1 for a, _ in pairs if a == n) == 1 and sum(1 for _, b in pairs if b["id"] == h["id"]) == 1:
                out[n] = (h["id"], 1, None)
        for n in names:
            if n in out:
                continue
            got = [p for p in everyone if not any(p["id"] == h["id"] for h in hs) and person_fits(heads[n], p)]
            if len(got) == 1:
                out[n] = (got[0]["id"], 0, where_serves(got[0]))
        return out

    cand = []

    # ---- the November ballot
    list_inc = {}                                           # race -> the name the list marks Incumbent
    for rid, rows in gen.items():
        i = races[rid]
        names = [r["Candidate Name"] for r in rows]
        if len(set(fold(n) for n in names)) != len(names):
            fail(f"the same name twice on the list for {rid}")
        fit = sitting(rid, names)
        order = 0
        for r in rows:
            name = r["Candidate Name"]
            marked = r["Incumbency Status"] == "Incumbent"
            if marked:
                list_inc[rid] = name
            mid, inc, note = fit.get(name, (None, 0, None))
            n = [note] if note else []
            if holders[rid]:
                if marked != bool(inc):
                    checks.append(f"{rid}: the list marks {name} {'Incumbent' if marked else 'Nonincumbent'}, the roster "
                                  f"{'does not agree' if marked else 'fits the seat holder'}")
            elif marked:
                inc = 1
                n.append("Marked as the incumbent on the Secretary of State's list.")
            if i["partisan"]:
                party = r["Party (if applicable)"]
                code = party_code(party)
                order += 1
                bo = order
                pc = CODES.get(party)
                t = tables.get((rid, party))
                if t is None:
                    n.append(f"There was no {party} primary for this office on May 12; named afterwards (the list does not "
                             "say how).")
                elif not any(c and same_person(ticket_head(name), nm) for nm, c, _k in t["names"]):
                    n.append(f"Not the nominee the canvass book checks in the May 12 {party} primary.")
                    checks.append(f"{rid}: {name} ({party}) on the November list is not the checked nominee of the primary")
            else:
                party, code, bo = NONPARTISAN, "N", None
            cand.append((rid, "general", GENERAL, name, party, code, bo, inc, 0, None, None, None, mid, SRC_LIST,
                         " ".join(n) or None))

    # ---- the May 12 primaries
    unused = sorted(set(k[0] for k in tables) - set(races))
    if unused:
        checks.append(f"the canvass book has state tables for contests not on the November list: {unused}")
    fields = Counter()
    for (rid, party), t in sorted(tables.items(), key=lambda kv: (kv[0][0], kv[0][1] or "")):
        if rid not in races:
            continue
        i = races[rid]
        real = [(nm, c, k, v) for (nm, c, k), v in zip(t["names"], t["total"]) if k != "scatter"]
        total = sum(t["total"])
        # the November list must carry every one who advanced (a partisan nominee for Governor heads a ticket)
        on_list = [r["Candidate Name"] for r in gen[rid] if (not party or r["Party (if applicable)"] == party)]
        for nm, c, k, _v in real:
            there = any(same_person(nm, ticket_head(x)) for x in on_list)
            if c and not there:
                checks.append(f"{rid}: {nm} advanced from the {party or 'nonpartisan'} primary but is not on the November list")
            if not c and there and not party:
                checks.append(f"{rid}: {nm} lost the primary but is on the November list")
        if not party and len(on_list) != sum(1 for _nm, c, _k, _v in real if c):
            checks.append(f"{rid}: {len(on_list)} on the November list, {sum(1 for x in real if x[1])} advanced from the primary")
        if len(real) < 2:
            continue
        fields[i["office_kind"] if i["office"] == "leg" else "statewide"] += 1
        fit = sitting(rid, [nm for nm, *_ in real])
        code = CODES.get(party) or re.sub(r"[^A-Z]", "", (party or "").upper())[:3] if party else "NP"
        for nm, c, k, v in real:
            mid, inc, note = fit.get(nm, (None, 0, None))
            n = [note] if note else []
            if not holders[rid] and rid in list_inc and same_person(nm, list_inc[rid]):
                inc = 1
            if c and party and not any(same_person(nm, ticket_head(x)) for x in on_list):
                n.append("Won the nomination; not on the Secretary of State's final list of candidates for November.")
            cand.append((rid, f"primary-{code}", PRIMARY, nm, party or NONPARTISAN, party_code(party) if party else "N",
                         None, inc, 1 if k == "writein" else 0, v, round(100 * v / total, 1) if total else None,
                         "advanced" if c else "lost", mid, SRC_BOOK, " ".join(n) or None))

    # ---- checks on the whole
    keys = Counter((c[0], c[1], c[3]) for c in cand)
    dup = [k for k, v in keys.items() if v > 1]
    if dup:
        fail(f"the same name twice in one election: {dup}")
    gen_rows = [c for c in cand if c[1] == "general"]
    if len(gen_rows) != len(listed):
        fail(f"{len(gen_rows)} November rows written for {len(listed)} state strips on the list")
    leg = sorted(int(i["district"]) for i in races.values() if i["office"] == "leg" and not i["special"])
    if leg != list(range(2, 49, 2)):
        checks.append(f"the regular Legislature seats on the list are not the 24 even-numbered districts: {leg}")
    extra = sorted(int(i["district"]) for i in races.values() if i["office"] == "leg" and i["special"])
    missing_statewide = sorted(k for k in ("governor", "sos", "treas", "ag", "aud") if race_id(k) not in races)
    if missing_statewide:
        checks.append(f"statewide offices not on the list: {missing_statewide}")
    no_primary = sorted(rid for rid, i in races.items() if not i["partisan"] and (rid, None) not in tables)
    if no_primary:
        checks.append(f"nonpartisan contests with no table in the canvass book: {no_primary}")

    race_rows = []
    for rid, i in sorted(races.items()):
        hs = holders[rid]
        cids = None
        if i["district"]:
            got = places.get(rid, set())
            unknown = sorted(p for p in got if fold(p) not in cmap)
            if unknown:
                fail(f"county names in the canvass book not in the Census file: {unknown}")
            if got:
                cids = ",".join(sorted({cmap[fold(p)][0] for p in got}))
            else:
                checks.append(f"{rid}: no county rows in the canvass book, so no counties for the race")
        hparty = ", ".join(dict.fromkeys(h["party"] for h in hs if h["party"])) or None
        if i["partisan"] and hparty == "Nonpartisan":
            checks.append(f"{rid}: the roster gives the holder's party as Nonpartisan for a partisan office; left empty")
            hparty = None
        race_rows.append((rid, STATE, i["level"], i["office_kind"], i["shown"], NAME, FIPS, cids, i["district"], None,
                          i["special"], i["partisan"], ",".join(h["id"] for h in hs) or None,
                          " and ".join(h["full"] for h in hs) or None, hparty, GENERAL, " ".join(notes[rid]) or None))

    place_rows = [("county", geoid, full, None, SRC_COUNTIES) for geoid, full in sorted(cmap.values())]
    n_book = sum(len([x for x in t["names"] if x[2] != "scatter"]) for k, t in tables.items() if k[0] in races)
    src = [
        (SRC_LIST, STATE, "official candidate list", "Nebraska Secretary of State",
         "Final Statewide General Candidate List, November 3, 2026 General Election", url.get("list") or PAGE, printed,
         list_fetched, list_sha, len(listed),
         f"Read from {list_how}. Of {strips} candidate strips, the {len(listed)} for state offices and the Legislature are kept "
         "(Office, District Name, Term, Vote For, Party, Candidate Name, Incumbency Status only); City of Residence, Mailing "
         f"Address and Phone/Email are never read. {local} strips for local district boards are left for the local pages, and "
         "Congress for the federal pages. The list prints no ballot order: partisan candidates are numbered in the list's "
         "order, nonpartisan ones are not numbered. It has no withdrawn or write-in entries."),
        (SRC_BOOK, STATE, "official results", "Nebraska Board of State Canvassers (compiled by the Secretary of State)",
         "Official Report of the Board of State Canvassers: Primary Election, May 12, 2026", url.get("canvass") or PAGE, "",
         mdate(book_path), sha_of(book_path), n_book,
         f"{sum(1 for k in tables if k[0] in races)} state tables read (statewide offices, the Public Service Commission, the "
         "Legislature, the State Board of Education and the Regents); every county column adds up to the printed Total. The one "
         "the report checks advanced (one per party; two in a nonpartisan primary). Percentages are of every vote the report "
         "counts in the contest, write-in scatterings included. Community college boards and later sections are local and "
         "not read."),
        (SRC_ROSTER, STATE, "roster", "Open States people project (CC0)",
         "Nebraska legislators and statewide officials, as loaded into state_ne.sqlite", "https://github.com/openstates/people",
         as_of, as_of, "", sum(len(v) for v in seats.values()) + len(offices),
         "Today's holder of each seat and office (names, parties, districts and ids only). The roster does not carry the "
         "Treasurer, the Auditor, the Public Service Commission, the State Board of Education or the Regents."),
        (SRC_COUNTIES, STATE, "boundaries", "U.S. Census Bureau",
         "Cartographic boundary file, counties, 2024, 1:500,000 (cb_2024_us_county_500k)",
         "https://www2.census.gov/geo/tiger/GENZ2024/shp/cb_2024_us_county_500k.zip", "", mdate(COUNTY_ZIP),
         sha_of(COUNTY_ZIP), len(place_rows),
         "Nebraska's 93 counties: names and GEOIDs only. A district race's county_ids are the counties the canvass book lists "
         "under it."),
    ]

    # the local level: the district boards and the judges on the filing list, and the counties whose own files can be read
    local = local_level(local_cache or os.path.join(cache, "ne", "local"), say, cmap, pdf_local, refresh)
    if len({r[0] for r in race_rows} & {r[0] for r in local["races"] + local["court_races"]}):
        fail("a local race shares its id with a state race")
    src = src + local["sources"]

    con = sqlite3.connect(db_path)
    try:
        con.executescript(SCHEMA)
        con.executescript(EXTRA_SCHEMA)
        with con:
            con.execute("DELETE FROM sl_candidates WHERE race_id IN (SELECT race_id FROM sl_races WHERE state = ?) OR race_id LIKE ?",
                        (STATE, f"2026-{STATE}-%"))
            con.execute("DELETE FROM sl_races WHERE state = ?", (STATE,))
            con.execute("DELETE FROM sl_sources WHERE state = ?", (STATE,))
            con.execute("DELETE FROM sl_places WHERE source_id LIKE 'ne-%'")
            con.execute("DELETE FROM sl_gaps WHERE state = ?", (STATE,))
            con.execute("DELETE FROM sl_notes WHERE state = ?", (STATE,))
            con.executemany("INSERT INTO sl_races VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", race_rows + local["court_races"] + local["races"])
            con.executemany("INSERT INTO sl_candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", cand + local["court_cands"] + local["cands"])
            con.executemany("INSERT INTO sl_sources VALUES (?,?,?,?,?,?,?,?,?,?,?)", src)
            con.executemany("INSERT INTO sl_places VALUES (?,?,?,?,?)", local["places"])     # the 93 counties, then every district and place a race names
            con.executemany("INSERT INTO sl_gaps VALUES (?,?,?,?,?,?,?)", local["gaps"])
            con.executemany("INSERT INTO sl_notes VALUES (?,?,?,?,?)", local["notes"])
    finally:
        con.close()

    # ---- the report: counts only
    kind_of = lambda rid: races[rid]["office_kind"] if races[rid]["office"] == "leg" else "statewide"
    by = Counter(kind_of(rid) for rid in races)
    g = Counter(kind_of(c[0]) for c in gen_rows)
    per = Counter(c[0] for c in gen_rows)
    one = Counter(kind_of(rid) for rid, v in per.items() if v == 1)
    inc = Counter(kind_of(c[0]) for c in gen_rows if c[7])
    say(f"    Nebraska (state races): {by['state_senate']} Legislature seats (even districts"
        + (f", and District {', '.join(map(str, extra))} for a shorter term" if extra else "")
        + f"), {by['statewide']} statewide and state-board races; {len(gen_rows)} candidates on the November ballot "
        f"(Legislature {g['state_senate']}, statewide and boards {g['statewide']}; one candidate only: Legislature "
        f"{one['state_senate']}, statewide and boards {one['statewide']}); incumbent on the ballot: Legislature "
        f"{inc['state_senate']}, statewide and boards {inc['statewide']}; primary fields: Legislature "
        f"{fields['state_senate']}, statewide and boards {fields['statewide']}, votes from the official canvass")
    for c in checks:
        say(f"    CHECK Nebraska (state races): {c}")
    cc, lv = local["counts"], local["by_level"]
    say(f"    Nebraska (local): {local['board_races'] + cc['empty board contests']} district board contests ({cc['board candidates placed']} candidates) for "
        f"{local['boards']} boards from the Secretary of State's filing list of {local['fetched']}; {local['county_races']} county, city, village, school, township "
        f"and other county-filed contests from {len(local['loaded'])} counties' own files; {len(local['court_races'])} judges' retention votes; local races by level: "
        + ", ".join(f"{k} {v}" for k, v in sorted(lv.items())))
    say(f"      check: the filing list's {local['rows']} rows = {cc['state']} state + {cc['federal']} federal + {cc['board rows']} district board rows; "
        f"{cc['board candidates']} board candidates read, {cc['board candidates placed']} placed, each in one race"
        + (f"; {cc['checked against the printed list']} of them are on the printed list too" if "checked against the printed list" in cc
           else "; the printed list was not on disk to check against"))
    say("      where the boards' seats are voted on: " + ", ".join(f"{v} from {k}" for k, v in sorted(local["how"].items(), key=lambda kv: -kv[1])))
    for fips, (county, n_contests, n_names, n_judges) in local["files"].items():
        c = local["county_counts"][fips]
        say(f"      {county} County: {n_contests} contests and {n_names} names read: loaded {c['loaded']} local contests ({c['loaded names']} names); checked "
            f"against the filing list {c['board']} district board contests; left to the state part {c['state']} federal and state contests; judges {n_judges}")
    say(f"      gaps: {len(local['gaps'])} ({sum(1 for g in local['gaps'] if g[1] == 'county')} counties whose own contests are not loaded yet)")
    for p in local["problems"]:
        say(f"    CHECK Nebraska (local): {p}")
    return len(cand)


if __name__ == "__main__":
    args = sys.argv[1:]
    opts = {}
    refresh = "--refresh" in args
    args = [a for a in args if a != "--refresh"]
    for flag in ("--cache", "--keep", "--local"):
        if flag in args:
            i = args.index(flag)
            opts[flag[2:]] = args[i + 1]
            del args[i:i + 2]
    if len(args) != 1:
        raise SystemExit("usage: python ballot/state_local_ne.py <database file> [--cache <folder>] [--keep <folder>] [--local <folder>] [--refresh]")
    load(args[0], cache=opts.get("cache", CACHE), keep=opts.get("keep"), local_cache=opts.get("local"), refresh=refresh)
