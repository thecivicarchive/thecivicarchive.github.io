"""
ballot/co_geo.py - the geography behind Colorado's ballot map, in the same files and formats ballot/mn_geo.py writes for
Minnesota (it imports that module for the geometry, the topology and the TopoJSON writer, ballot/sd_geo.py for putting
census blocks together into precincts, ballot/ia_geo.py, ballot/wi_geo.py and ballot/wy_geo.py for the few things those
added, and ballot/state_local_co.py for the judicial districts, and changes nothing in any of them), so the same page
and the same reader (ballot/mn_geo_reader.js) read them all.

    python ballot/co_geo.py                 builds ballot_geo/co/ and runs the self-test (the first time: 220 MB of
                                            census blocks, then school districts laid over precincts)
    python ballot/co_geo.py --selftest      runs the self-test on the files already built
    python ballot/co_geo.py --refresh       asks for the Secretary of State's precinct table and RTD's districts again
    python ballot/co_geo.py --out DIR       builds somewhere else (a trial run)

Nothing here writes a database. ballot_local_2026.sqlite is opened read-only, for place names and for the check that
every shape's id is one the races carry.

What Colorado publishes, and what that does to the files
---------------------------------------------------------
Colorado has no statewide map file of its precincts as they stand in 2026 that a script may read: each county clerk
draws the county's own (the State Demography Office has compiled them, but does not publish the file). The only
statewide precinct lines a script may read are the Census Bureau's 2020 voting districts: the precincts the counties
reported for the 2020 census (3,108 of them). The commissions redrew every legislative and congressional district in
2021 and the counties re-precincted to fit, so a shape here is a 2020 voting district cut wherever a 2021-plan
legislative or congressional line, a city or town limit, or an RTD director district line crosses it, and index.json
says so in plain words. What a ballot depends on does not depend on the precinct: the county, the congressional, Senate
and House district, the judicial district, the city or town, the RTD director district and the school district at a
point are each taken from a current official file.

Sources (downloaded with states/net.py, an honest User-Agent, one request at a time, cached in states_cache/co_local/)
----------------------------------------------------------------------------------------------------------------------
  - Census blocks: the Census Bureau's TIGER/Line 2020 tabulation blocks of Colorado (tl_2020_08_tabblock20.zip, 140,345
    blocks). Only each block's number, its inside point and its outline are read.
  - Which voting district a block is in: the Bureau's 2020 Block Assignment File (BlockAssign_ST08_CO.zip, the VTD
    table), and the voting districts' names from tl_2020_08_vtd20.zip (the county's name and the precinct's number).
  - Which House, Senate and congressional district a block is in: the Bureau's 2024 State Legislative District Block
    Equivalency Files (the national sldl24.zip and sldu24.zip ballot/sd_geo.py reads) and its 119th Congress file
    (cd119.zip, as ballot/mo_geo.py and ballot/ne_geo.py read): the plans the Independent Legislative and Congressional
    Redistricting Commissions adopted in 2021 and the Supreme Court approved, 65 House, 35 Senate and 8 congressional
    districts. A block is in exactly one of each, so those lines are exact.
  - Cities and towns: TIGER/Line 2025 places (tl_2025_08_place.zip), the incorporated, active ones only (Colorado's 272
    statutory and home rule cities and towns, Denver and Broomfield among them; census-designated places are not
    governments). A block belongs to the city or town its own inside point lies in.
  - RTD director districts: the Regional Transportation District's own feature service "RTD GIS Boundaries (download
    view)" (layer DirectorDistricts, 15 districts lettered A to O, the lines RTD's board adopted after the 2020 census),
    read in longitude and latitude. A block belongs to the director district its own inside point lies in, so a
    precinct piece lies wholly inside one, or outside the RTD.
  - School districts: TIGER/Line 2025 unified school districts (tl_2025_08_unsd.zip; Colorado has no elementary or
    secondary-only districts in the Bureau's files).
  - County lines and the state's outline for the area check: TIGER/Line 2025 county subdivisions
    (tl_2025_08_cousub.zip), and county names from the Bureau's 2020 list of county codes (st08_co_cou2020.txt).
  - Judicial districts: sections 13-5-102 to 13-5-123.1, C.R.S. (23 districts, each whole counties), as
    ballot/state_local_co.read_judicial reads them from the Office of Legislative Legal Services' file of Title 13.
  - The Secretary of State's monthly voter registration statistics (the newest month's workbook, linked from
    coloradosos.gov/pubs/elections/VoterRegNumbers/VoterRegNumbers.html), sheet "Voter Counts by Precinct": every
    precinct of 2026 by its ten-digit number, which carries the precinct's congressional district (one digit), Senate
    district (two), House district (two), county (two, the county's place in alphabetical order) and number in the county
    (three). Only the county and precinct number cells are read; the voter counts never are. Used as a check: a piece
    is marked listed_2026 where its 2020 number is a 2026 number of its county AND that 2026 precinct's own digits give
    the piece's congressional, Senate and House district.
  - Polling places: Colorado votes by mail; each county clerk designates Voter Service and Polling Centers and drop
    boxes, any of which serves any voter of the county. The Secretary of State sells the statewide list. See POLLING
    PLACES below.

What is built (ballot_geo/co/): index.json, manifest.json, precincts/<county>.json (64), layers/<kind>.json (state,
county, cd, senate, house, judicial, mcd, school), school/<id>.json, polling_places.json and reader.js, each as
ballot/mn_geo.py describes.

Ids (the ballot database's own; a county is its five-digit code)
----------------------------------------------------------------
  state     "CO" (j = "08", the jurisdiction_id of statewide races and of the Supreme Court and Court of Appeals votes)
  county    "08031"                 sl_places county id; jurisdiction_id of county offices and county court judges
  mcd       "CO-M-20000"            CO-M- and the Census Bureau's place code (check_local's convention)
  house     "6"   senate "31"   cd "1" (properties.j "CO-CD1", the jurisdiction_id of the State Board of Education's and
                                the University of Colorado regents' seats; properties.race is 2026-CO-H01)
  judicial  "CO-JD2"                the jurisdiction_id of district court retention votes; d is the number
  school    "CO-S-03360"            CO-S- and the Bureau's five-digit district code (properties.nces is the federal code)
  rtd       "CO-X-RTD-B"            a precinct property, not a layer (the page then finds RTD director races by the
                                    precinct's own word, as Nebraska's natural resources districts are found)
  sboe, regent  "CO-CD1"            precinct properties: the seat on the State Board of Education and of the University
                                    of Colorado regents the precinct votes for (each the congressional district of the
                                    same number, C.R.S. 22-2-105 and 23-20-102)

Not drawn, because no statewide file has the lines: county commissioner districts (each county keeps its own; Colorado's
commissioners are elected by the whole county with a residence district, or by district in a few home rule counties, and
the ballot database has no commissioner race for 2026 yet), city council districts and wards, and special districts
other than RTD. index.json lists every race without a shape under "check", with the reason.

POLLING PLACES: see POLL_HOW. polling_places.json is "waiting" and says so.
"""

import argparse
import collections
import datetime as dt
import hashlib
import io
import json
import math
import os
import re
import shutil
import sqlite3
import sys
import time
import zipfile

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from states import net  # noqa: E402
from ballot import mn_geo as G  # noqa: E402
from ballot import wi_geo as W  # noqa: E402
from ballot import ia_geo as I  # noqa: E402,E741
from ballot import sd_geo as S  # noqa: E402
from ballot import wy_geo as Y  # noqa: E402

GeoError = G.GeoError
STATE, FIPS, STATE_NAME = "CO", "08", "Colorado"
OUT = os.path.join(HERE, "ballot_geo", "co")
CACHE = os.path.join(HERE, "states_cache", "co_local")
JUD_FOLDER = os.path.join(HERE, "ballot_cache", "co", "local")      # where ballot/state_local_co.py keeps the statute's sentences
DB = os.path.join(HERE, "ballot_local_2026.sqlite")
READER_JS = G.READER_JS
ELECTION = "2026-11-03"
FINDER = "https://www.govotecolorado.gov/"      # the Secretary of State's voter lookup (Voter Service and Polling Centers, drop boxes)
N_HOUSE, N_SENATE, N_CD, N_COUNTY, N_JUD = 65, 35, 8, 64, 23

CENSUS = "https://www2.census.gov/"
BLOCK_URL = CENSUS + "geo/tiger/TIGER2020/TABBLOCK20/tl_2020_08_tabblock20.zip"
BAF_URL = CENSUS + "geo/docs/maps-data/data/baf2020/BlockAssign_ST08_CO.zip"
VTD_URL = CENSUS + "geo/tiger/TIGER2020PL/STATE/08_COLORADO/08/tl_2020_08_vtd20.zip"
SLDL_URL = S.SLDL_URL
SLDU_URL = S.SLDU_URL
CD_URL = CENSUS + "programs-surveys/decennial/rdo/mapping-files/2025/119-congressional-district-befs/cd119.zip"
COUSUB_URL = CENSUS + "geo/tiger/TIGER2025/COUSUB/tl_2025_08_cousub.zip"
PLACE_URL = CENSUS + "geo/tiger/TIGER2025/PLACE/tl_2025_08_place.zip"
UNSD_URL = CENSUS + "geo/tiger/TIGER2025/UNSD/tl_2025_08_unsd.zip"
COUNTY_URL = CENSUS + "geo/docs/reference/codes2020/cou/st08_co_cou2020.txt"
SHARED = {"sldl24.zip": "sd_local", "sldu24.zip": "sd_local", "cd119.zip": "ne_local"}      # national files another state's build fetched
RTD_SERVICE = "https://services5.arcgis.com/1fZoXlzLW6FCIUcE/arcgis/rest/services/RTD_GIS_Boundaries_view/FeatureServer/1"
RTD_ITEM = "https://gis-rtd-denver.opendata.arcgis.com/datasets/674d164b39694c79849b7c19465077ba"
SOS_STATS_PAGE = "https://www.coloradosos.gov/pubs/elections/VoterRegNumbers/VoterRegNumbers.html"
SOS_STATS_BASE = "https://www.coloradosos.gov/pubs/elections/VoterRegNumbers/"
STATS_SHEET = "Voter Counts by Precinct"
CD_NOTE = ("The eight districts the Independent Congressional Redistricting Commission adopted in 2021, approved by the Colorado Supreme Court "
           "(the Census Bureau's 119th Congress equivalency file).")

ARC_KINDS = ["county", "mcd", "house", "senate", "cd", "judicial"]      # each precinct piece lies wholly inside one of each (or no city)
AREA_SLACK = 0.02                 # a county's precincts (2020 blocks) and the Bureau's 2025 county may differ in area by this share
PLACE_LSAD = {"25": "city", "43": "town"}
RTD_LETTERS = "ABCDEFGHIJKLMNO"

clean = S.clean
as_lonlat = Y.as_lonlat


def rtd_id(letter):
    return f"{STATE}-X-RTD-{letter}"


def cd_id(n):
    return f"{STATE}-CD{n}"


def county_no(county):
    """Colorado's own number for a county (the two digits in a precinct number): its place in the order of the Census
    Bureau's county codes (Adams 1 for 001, Yuma 63 for 125), and Broomfield, the county made in 2001, 64. Two counties
    (Hinsdale 27, Mineral 40) reported their 2020 voting districts under this number rather than the Bureau's code."""
    return 64 if county == FIPS + "014" else (int(county[2:]) + 1) // 2


# ---------------------------------------------------------------- the Bureau's tables

def read_tables(bafpath, lpath, upath, cdpath, vpath):
    """{block: (county, voting district code, House, Senate, congressional district)} and {(county, code): the voting
    district's name} from the Bureau's files."""
    z = zipfile.ZipFile(bafpath)
    lines = z.read("BlockAssign_ST08_CO_VTD.txt").decode("utf-8").splitlines()
    if lines[0].strip() != "BLOCKID|COUNTYFP|DISTRICT":
        raise GeoError(f"    {os.path.basename(bafpath)}: the voting district table's heading is now {lines[0]!r}; stopping")
    vtd = {}
    for line in lines[1:]:
        b, c, v = line.split("|")
        vtd[b] = (c, v.strip())

    def bef(path, head):
        zz, out = zipfile.ZipFile(path), {}
        member = next(n for n in zz.namelist() if n.lower().startswith("national"))
        with zz.open(member) as fh:
            text = io.TextIOWrapper(fh, "utf-8")
            if next(text).strip() != head:
                raise GeoError(f"    {os.path.basename(path)}: the heading is not {head!r}; stopping")
            for line in text:
                if line.startswith(FIPS):
                    b, d = line.strip().split(",")
                    out[b] = d
        return out

    low, up, cd = bef(lpath, "GEOID,SLDLST"), bef(upath, "GEOID,SLDUST"), bef(cdpath, "GEOID,CDFP")
    if set(low) != set(vtd) or set(up) != set(vtd) or set(cd) != set(vtd):
        raise GeoError("    the block assignment file and the legislative and congressional equivalency files do not list the same blocks; stopping")
    table = {}
    for b, (c, v) in vtd.items():
        h, s, d = re.sub(r"^0+", "", low[b]), re.sub(r"^0+", "", up[b]), re.sub(r"^0+", "", cd[b])
        if not (b[:5] == FIPS + c and re.fullmatch(r"\d{6}", v) and v[:3] in (c, f"{county_no(FIPS + c):03d}") and re.fullmatch(r"\d{1,2}", h) and 1 <= int(h) <= N_HOUSE
                and re.fullmatch(r"\d{1,2}", s) and 1 <= int(s) <= N_SENATE and re.fullmatch(r"[1-8]", d)):
            raise GeoError(f"    block {b}: county {c!r}, voting district {v!r}, House {low[b]!r}, Senate {up[b]!r}, Congress {cd[b]!r} do not fit "
                           "the layout this builder was checked against; stopping")
        table[b] = (FIPS + c, v, h, s, d)
    seen = [{t[i] for t in table.values()} for i in (2, 3, 4)]
    if [len(x) for x in seen] != [N_HOUSE, N_SENATE, N_CD]:
        raise GeoError(f"    the equivalency files give {len(seen[0])} House, {len(seen[1])} Senate and {len(seen[2])} congressional districts; stopping")
    import shapefile
    zz = zipfile.ZipFile(vpath)
    base = next(n for n in zz.namelist() if n.endswith(".dbf"))
    vname = {(FIPS + r["COUNTYFP20"], r["VTDST20"].strip()): r["NAME20"] for r in (x.as_dict() for x in shapefile.Reader(dbf=io.BytesIO(zz.read(base))).iterRecords())}
    missing = {(t[0], t[1]) for t in table.values()} - set(vname)
    if missing:
        raise GeoError(f"    {len(missing)} voting districts of the block assignment file are not in {os.path.basename(vpath)} (e.g. {sorted(missing)[0]}); stopping")
    return table, vname


def county_names(path):
    """{county code: "Adams County"} from the Bureau's list of county codes."""
    import csv
    out = {}
    rows = list(csv.reader(open(path, encoding="utf-8"), delimiter="|"))
    col = {h: i for i, h in enumerate(rows[0])}
    for r in rows[1:]:
        if r and r[col["STATEFP"]] == FIPS:
            out[FIPS + r[col["COUNTYFP"]]] = r[col["COUNTYNAME"]]
    if len(out) != N_COUNTY:
        raise GeoError(f"    {os.path.basename(path)}: {len(out)} counties, not {N_COUNTY}; stopping")
    return out


def short_county(name):
    return re.sub(r"\s+(County|city and county|City and County)$", "", name.strip(), flags=re.I)


def vtd_name(county_short, code, raw):
    """A voting district's name as the page prints it: the Bureau's, which in Colorado is the county's name and the
    precinct's number ("Adams 230"); anything else is named by its number, so no name the rule sets aside can pass."""
    n = re.sub(r"\s+", " ", raw or "").strip()
    if re.fullmatch(re.escape(county_short) + r" \d{1,4}", n, flags=re.I):
        return n
    return "Voting district " + (re.sub(r"^0+(?=\d)", "", code[3:]) or code)


# ---------------------------------------------------------------- cities and towns, and RTD, by each block's inside point

def read_places(path):
    """The incorporated, active cities and towns: [(record, rings as vertex keys)], in code order."""
    rows = I.read_shapefile(path, lambda r: r["STATEFP"] == FIPS and r["LSAD"] in PLACE_LSAD and r["FUNCSTAT"] == "A")
    if not 265 <= len(rows) <= 280 or any(r["CLASSFP"] != "C1" for r, _ in rows):
        raise GeoError(f"    {os.path.basename(path)}: {len(rows)} incorporated cities and towns that do not fit the layout this builder was checked against; stopping")
    return sorted(rows, key=lambda x: x[0]["PLACEFP"])


def read_rtd(doc):
    """RTD's director districts: [(letter, rings as vertex keys)], from the service's rows (longitude, latitude)."""
    out = []
    for a, rings in doc["rows"]:
        letter = clean(str(a.get("BND") or "")).upper()
        if letter not in RTD_LETTERS or len(letter) != 1:
            raise GeoError(f"    RTD director districts: a district is lettered {letter!r}; stopping")
        keys = [k for k in (G.clean_ring([(x, y) for x, y in ring]) for ring in rings) if k]
        out.append((letter, keys))
    if sorted(l for l, _ in out) != list(RTD_LETTERS):
        raise GeoError(f"    RTD director districts: {sorted(l for l, _ in out)}, not the fifteen letters A to O; stopping")
    return sorted(out)


def block_in(blockpath, groups):
    """{group name: {block: code or None}}: for each group of shapes [(code, rings as vertex keys)], the shape whose
    outline holds the block's own inside point."""
    import shapefile
    z = zipfile.ZipFile(blockpath)
    base = next(n for n in z.namelist() if n.endswith(".dbf"))[:-4]
    prepared = {}
    for gname, shapes in groups.items():
        pl = []
        for code, rings in shapes:
            xy = [[G.vxy(k) for k in ring] for ring in rings]
            xs, ys = [p[0] for ring in xy for p in ring], [p[1] for ring in xy for p in ring]
            pl.append((code, (min(xs), min(ys), max(xs), max(ys)), xy))
        prepared[gname] = pl
    out = {g: {} for g in groups}
    for rec in shapefile.Reader(dbf=io.BytesIO(z.read(base + ".dbf"))).iterRecords():
        d = rec.as_dict()
        x, y = int(round(float(d["INTPTLON20"]) * 1e7)), int(round(float(d["INTPTLAT20"]) * 1e7))
        for gname, pl in prepared.items():
            hit = None
            for code, (x0, y0, x1, y1), xy in pl:
                if x0 <= x <= x1 and y0 <= y <= y1 and G.in_rings(x, y, xy):
                    hit = code
                    break
            out[gname][d["GEOID20"]] = hit
    return out


def read_units(blockpath, table, bplace, brtd, vname, cshort, say):
    """The precincts: the blocks of each 2020 voting district put together, one piece for each legislative and
    congressional district, each city or town (or none) and each RTD director district (or none) it reaches. Returns
    (pre, polys): what is said of each, and its rings as vertex keys."""
    import shapefile
    z = zipfile.ZipFile(blockpath)
    base = next(n for n in z.namelist() if n.endswith(".shp"))[:-4]
    r = shapefile.Reader(shp=io.BytesIO(z.read(base + ".shp")), dbf=io.BytesIO(z.read(base + ".dbf")), shx=io.BytesIO(z.read(base + ".shx")))
    groups, seen, twice = collections.defaultdict(set), set(), 0
    for sr in r.iterShapeRecords():
        b = sr.record["GEOID20"]
        if b not in table or b in seen:
            raise GeoError(f"    {os.path.basename(blockpath)}: block {b} is not in the block assignment file, or is there twice; stopping")
        seen.add(b)
        s = groups[table[b] + (bplace.get(b) or "", brtd.get(b) or "")]
        pts, parts = sr.shape.points, list(sr.shape.parts) + [len(sr.shape.points)]
        for i in range(len(parts) - 1):
            ks = G.clean_ring(pts[parts[i]:parts[i + 1]])
            if not ks:
                continue
            a = ks[-1]
            for k in ks:
                if (k, a) in s:
                    s.remove((k, a))
                elif (a, k) in s:
                    twice += 1
                else:
                    s.add((a, k))
                a = k
    if len(seen) != len(table):
        raise GeoError(f"    {os.path.basename(blockpath)}: {len(seen):,} blocks, but the block assignment file lists {len(table):,}; stopping")
    if twice:
        raise GeoError(f"    {os.path.basename(blockpath)}: {twice} edges are claimed by two blocks of one precinct; stopping")
    pieces = collections.Counter(k[:2] for k in groups)
    order = sorted(groups, key=lambda k: (k[0], k[1], G.natkey(k[2]), G.natkey(k[3]), k[4], k[5], k[6]))
    part_no, n_of = {}, collections.Counter()
    for key in order:
        n_of[key[:2]] += 1
        part_no[key] = n_of[key[:2]]
    pre, polys, ids = [], [], set()
    for key in order:
        county, code, house, senate, cd, place, rtd = key
        many = pieces[key[:2]] > 1
        pid = f"{county}.{code}" + (f".{part_no[key]}" if many else "")
        if pid in ids:
            raise GeoError(f"    two precincts would share the id {pid}; stopping")
        ids.add(pid)
        name = vtd_name(cshort[county], code, vname[(county, code)])
        pre.append({"id": pid, "county": county, "vtd": code, "precinct": county + code, "house": house, "senate": senate, "cd": cd,
                    "mcd": f"{STATE}-M-{place}" if place else None, "rtd": rtd_id(rtd) if rtd else None,
                    "name": name + (f", part {part_no[key]}" if many else ""), "vtdname": name, "parts": pieces[key[:2]]})
        polys.append(S.loops(groups[key]))
    say(f"      {len(seen):,} census blocks put together into {len(pieces):,} voting districts of 2020, {len(pre):,} pieces once cut by the 2021-plan "
        f"districts, the cities' and towns' limits and RTD's director districts ({sum(1 for n in pieces.values() if n > 1)} voting districts reach more than one)")
    return pre, polys


# ---------------------------------------------------------------- the Secretary of State's 2026 precinct numbers

def stats_url(say):
    """The newest month's statistics workbook the Secretary's page links to (its link text names the month)."""
    page = net.get(SOS_STATS_PAGE).decode("utf-8", "replace")
    links = re.findall(r'href="((\d{4})/([A-Z][a-z]+)Statistics\2\.xlsx)"', page)
    months = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
    links = [(int(y), months.index(m), rel) for rel, y, m in links if m in months]
    if not links:
        raise GeoError("    the Secretary of State's statistics page links no monthly workbook this builder can read; stopping")
    return SOS_STATS_BASE + max(links)[2]


def read_sos_table(path, cname):
    """{county: {number in the county: ten-digit precinct number}} from the "Voter Counts by Precinct" sheet. Only the
    county and precinct cells are read. Every county's two-digit number must be one and the same: the county's place in
    the order of the Census Bureau's county codes (Adams 01 for 001, Yuma 63 for 125), Broomfield, the county made in
    2001, 64."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    if STATS_SHEET not in wb.sheetnames:
        raise GeoError(f"    {os.path.basename(path)} has no sheet {STATS_SHEET!r}; stopping")
    ws = wb[STATS_SHEET]
    by_name = {S.name_fold(short_county(n)): c for c, n in cname.items()}
    alpha = {c: county_no(c) for c in cname}
    out, head, cno = collections.defaultdict(dict), None, {}
    for row in ws.iter_rows(values_only=True):
        cells = list(row[:2]) if row else []
        if head is None:
            if [str(c or "").strip().upper() for c in cells] == ["COUNTY", "PRECINCT"]:
                head = True
            continue
        county, number = cells + [None] * (2 - len(cells))
        if county is None and number is None:
            continue
        county = str(county or "").strip()
        if re.fullmatch(r"(?i)(grand )?totals?|end of document", county) or number in (None, ""):
            continue
        num = str(number or "").strip()
        num = num[:-2] if num.endswith(".0") else num
        c = by_name.get(S.name_fold(county))
        if c is None or not re.fullmatch(r"\d{10}", num):
            raise GeoError(f"    {os.path.basename(path)}: a precinct row this reader was not checked against ({len(county)} and {len(num)} characters); stopping")
        cno.setdefault(c, set()).add(int(num[5:7]))
        n3 = int(num[7:])
        if n3 in out[c]:
            raise GeoError(f"    {os.path.basename(path)}: {county} lists precinct number {n3} twice; stopping")
        out[c][n3] = num
    if head is None or len(out) != N_COUNTY or sum(len(v) for v in out.values()) < 2500:
        raise GeoError(f"    {os.path.basename(path)} was read as {len(out)} counties, {sum(len(v) for v in out.values())} precincts; stopping")
    bad = [cname[c] for c, ns in cno.items() if ns != {alpha[c]}]
    if bad:
        raise GeoError(f"    {os.path.basename(path)}: the county digits of {', '.join(bad[:4])} are not the county's place in the order of county codes; stopping")
    return out


def code_districts(num):
    """(congressional, Senate, House) district from a ten-digit Colorado precinct number."""
    return str(int(num[0])), str(int(num[1:3])), str(int(num[3:5]))


# ---------------------------------------------------------------- the ballot database (read-only)

def read_db(db):
    info = {"names": {}, "races": [], "found": False}
    if not os.path.exists(db):
        return info
    con = sqlite3.connect("file:" + db.replace("\\", "/") + "?mode=ro", uri=True)
    try:
        for kind, pid, name in con.execute("SELECT kind, id, name FROM sl_places WHERE source_id LIKE 'co-%'"):
            info["names"][(kind, pid)] = name
        info["races"] = list(con.execute("SELECT race_id, level, office_kind, jurisdiction, jurisdiction_id, district, county_ids, office "
                                         "FROM sl_races WHERE state = ?", (STATE,)))
    finally:
        con.close()
    info["found"] = True
    return info


def overlay_cached(name, pre_rings, districts, stamp, refresh, say):
    """mn_geo.school_overlay, its answer kept beside the downloads while they have not changed."""
    path = os.path.join(CACHE, f"co_geo_overlay_{name}.json")
    if not refresh and os.path.exists(path):
        try:
            kept = json.load(open(path, encoding="utf-8"))
            if kept.get("stamp") == stamp and len(kept.get("rows", [])) == len(pre_rings):
                return kept["rows"]
        except (ValueError, OSError):
            pass
    say(f"      laying the {name} lines over the precincts (kept for the next build)")
    res = G.school_overlay(pre_rings, districts, levels=(G.SCHOOL_THICK,), say=None)
    rows = [[[d, round(share, 5), thick[0]] for d, share, thick in row] for row in res]
    with open(path + ".part", "w", encoding="utf-8") as fh:
        json.dump({"stamp": stamp, "rows": rows}, fh, separators=(",", ":"))
    os.replace(path + ".part", path)
    return rows


# ---------------------------------------------------------------- what the precincts say, and ids against the ballot database

SBOE_WORD = "State Board of Education district"
REGENT_WORD = "Regent of the University of Colorado district"


def said_words(word):
    return f"{word} ({', '.join(f'{cd_id(n)} Congressional District {n}' for n in range(1, N_CD + 1))}): the congressional district of the same number"


def check_ids(info, shape_ids, said_ids):
    """Every Colorado race in the ballot database against the shapes, by the rule the page builder uses
    (build_ballot_state_dev.race_shape, then race_said and place_shape). A race without a shape is listed with the
    reason, and with what the precincts say of its district where they name it."""
    if not info["found"]:
        return {"races": 0, "matched": 0, "by_layer": {}, "no_shape": [], "part_of_a_shape": [], "note": "not checked in this build"}
    Sh = shape_ids
    by_layer, missing, matched, said, seats = collections.Counter(), collections.OrderedDict(), 0, collections.Counter(), collections.Counter()
    for rid, level, kind, jur, jid, district, county_ids, office in sorted(info["races"], key=lambda r: r[0]):
        jid = str(jid or "").strip()
        d = str(district).strip() if district not in (None, "") else ""
        court = re.sub(r"_retention$", "", kind or "")
        counties = re.findall(r"\d+", str(county_ids or ""))
        hit, why = None, None
        if level == "statewide" or (court in ("supreme_court", "court_of_appeals") and not d):
            hit = ("state", STATE)
            if level == "statewide" and d:
                prop = "sboe" if kind == "state_board_of_education" else "regent" if kind == "university_board" else None
                if prop and jid in said_ids.get(prop, ()):
                    seats[prop] += 1
        elif level == "court" and not d and not counties and jid in (STATE, FIPS):
            hit = ("state", STATE)
        elif kind == "state_senate":
            hit = ("senate", re.sub(r"^0+(?=\d)", "", d.upper()))
        elif kind == "state_house":
            hit = ("house", re.sub(r"^0+(?=\d)", "", d.upper()))
        elif kind == "district_court" or (level == "court" and jid in Sh.get("judicial", {})):
            hit = ("judicial", jid)
        elif kind in ("county_commissioner", "county_council"):
            hit = ("county", jid) if not d else ("com", f"{jid}|{d}")
            if d:
                why = "no statewide file has the lines of county commissioner districts; each county keeps its own"
        elif level == "county":
            hit = ("county", jid)
        elif level in ("city", "township"):
            hit = ("ward", f"{jid}|{d}") if d else ("mcd", jid)
            if d:
                why = "no statewide file has the lines of city council districts or wards; the city keeps them (the page places the race on its city)"
        elif level == "school":
            hit = ("school", jid)
        elif level == "court" and jid:
            hit = ("mcd", jid) if jid in Sh.get("mcd", {}) else ("county", jid)
        elif level == "other" and jid in Sh.get("county", {}):
            hit = ("county", jid)
        else:
            why = "no statewide file read here has the lines of this kind of district"
        if hit and hit[1] in Sh.get(hit[0], {}):
            matched += 1
            by_layer[hit[0]] += 1
            continue
        named = next((k for k, ids in sorted(said_ids.items()) if jid and jid in ids), None)
        if named:
            said[named] += 1
        key = (level, kind, jur, jid, district)
        e = missing.setdefault(key, {"level": level, "office_kind": kind, "jurisdiction": jur, "jurisdiction_id": jid, "district": district, "races": 0,
                                     "why": ("no layer draws it; every precinct names its own under " + named) if named else (why or "no shape carries this id")})
        if named:
            e["said_by_the_precinct"] = f"{named}:{jid}"
        e["races"] += 1
    return {"races": len(info["races"]), "matched": matched, "by_layer": dict(sorted(by_layer.items())),
            "named_by_the_precinct": dict(sorted(said.items())), "races_named_by_the_precinct": sum(said.values()),
            "statewide_seats_also_named_by_the_precinct": dict(sorted(seats.items())),
            "no_shape": list(missing.values()), "part_of_a_shape": []}


# ---------------------------------------------------------------- polling places

POLL_WAITING = {
    "why": "Colorado votes by mail. Each county clerk designates the county's Voter Service and Polling Centers and drop boxes, and any of them serves "
           "any voter of the county, so they belong to the county rather than to a precinct. The Secretary of State's statewide list of them is sold "
           "as a data request, not published as a file; GoVoteColorado and the county clerk say where a voter can vote or drop off a ballot.",
}
POLL_HOW = ("      polling places: waiting. Colorado votes by mail, and its Voter Service and Polling Centers (VSPCs) and drop boxes serve every voter of their "
            "county, not a precinct. The Secretary of State's 'Statewide Voter Service and Polling Center Locations List (VSPC)' for the general election "
            "is sold on the Elections Division's data request form (coloradosos.gov/pubs/elections/forms/dataRequests.pdf, $50 a report): buying it is "
            "John's decision. GoVoteColorado answers one voter at a time and is not to be scripted. Were a person to obtain that list, the place for it is "
            "states_cache/co_local/sos/vspc/ (the file as received, e.g. VSPC_2026_General.xlsx); a reader for it would still have to be written and "
            "checked against it, keeping only the site's name, address, county, kind (VSPC or drop box) and dates and hours, never a contact.")


def polling_places(put, say):
    doc = {"v": G.VERSION, "election": ELECTION, "finder": FINDER, "status": "waiting", **POLL_WAITING}
    put("polling_places.json", doc)
    say(POLL_HOW)
    return {"file": "polling_places.json", "status": "waiting"}


# ---------------------------------------------------------------- the build

def judicial_of(county_long, say):
    """{county: "CO-JD<n>"} and the statute record, from ballot/state_local_co.read_judicial (its cached sentences)."""
    from ballot import state_local_co as L
    cmap = {L.ckey(short_county(n)): (c, n) for c, n in county_long.items()}
    path, kept = L.read_judicial(JUD_FOLDER, cmap, say)
    jud_of = {}
    for no, e in kept["districts"].items():
        for n in e["counties"]:
            c = cmap[L.ckey(n)][0]
            if c in jud_of:
                raise GeoError(f"    the judicial districts put {n} County in two districts; stopping")
            jud_of[c] = f"{STATE}-JD{int(no)}"
    if sorted(jud_of) != sorted(county_long) or len(set(jud_of.values())) != N_JUD:
        raise GeoError(f"    the judicial districts hold {len(jud_of)} counties in {len(set(jud_of.values()))} districts, not {N_COUNTY} in {N_JUD}; stopping")
    return jud_of, path, kept


def rtd_about():
    """The service's own date of its last data edit, if it says."""
    try:
        j = json.loads(net.get(RTD_SERVICE + "?f=json", accept="application/json"))
        ms = ((j.get("editingInfo") or {}).get("dataLastEditDate"))
        return {"name": j.get("name"), "data_last_edited": dt.datetime.fromtimestamp(ms / 1000, dt.timezone.utc).strftime("%Y-%m-%d") if ms else None}
    except Exception:  # noqa: BLE001
        return {}


def build(out=OUT, db=DB, refresh=False, say=print):
    t0 = time.time()
    say("    Colorado ballot map: 2020 voting districts, 2021-plan congressional and legislative districts, county, city and town, RTD, school and judicial "
        "district lines (the Census Bureau, RTD, C.R.S. 13-5-102 on, the Secretary of State's 2026 precinct numbers)")
    os.makedirs(CACHE, exist_ok=True)
    path = lambda name: os.path.join(CACHE, name)      # noqa: E731
    urls = (BLOCK_URL, BAF_URL, VTD_URL, SLDL_URL, SLDU_URL, CD_URL, COUSUB_URL, PLACE_URL, UNSD_URL, COUNTY_URL)
    bpath, bafpath, vpath, lpath, upath, cdpath, tpath, ppath, unpath, cpath = (path(u.rsplit("/", 1)[1]) for u in urls)
    for name, folder in SHARED.items():
        twin = os.path.join(HERE, "states_cache", folder, name)
        if not os.path.exists(path(name)) and os.path.exists(twin):
            shutil.copyfile(twin, path(name))      # the same national file, already fetched once by another state's build
    for url, p in zip(urls, (bpath, bafpath, vpath, lpath, upath, cdpath, tpath, ppath, unpath, cpath)):
        net.download(url, p, 3650, say=lambda *_a: None)
    statpath = path("co_sos_voter_statistics_latest.xlsx")
    if refresh and os.path.exists(statpath):
        os.utime(statpath, (0, 0))
    stat_url = None
    if not os.path.exists(statpath) or time.time() - os.path.getmtime(statpath) > 30 * 86400:
        stat_url = stats_url(say)
        net.download(stat_url, statpath, 0 if refresh else 30, say=lambda *_a: None)
        with open(statpath + ".url", "w", encoding="utf-8") as fh:
            fh.write(stat_url)
    if os.path.exists(statpath + ".url"):
        stat_url = open(statpath + ".url", encoding="utf-8").read().strip()
    rtd_doc = G.fetch_full(RTD_SERVICE, "BND", path("rtd_director_districts_geometry_4326.json.gz"), 50, refresh, say)
    rtd_meta = rtd_about()

    info = read_db(db)
    names = info["names"]
    if not info["found"]:
        say(f"      {os.path.basename(db)} is not there: places are named from the sources alone")
    county_long = county_names(cpath)
    cname = {c: names.get(("county", c)) or n for c, n in county_long.items()}
    cshort = {c: short_county(n) for c, n in county_long.items()}
    jud_of, law_path, law = judicial_of(county_long, say)
    sos = read_sos_table(statpath, county_long)

    # ---- precincts: the 2020 voting districts, from blocks, cut by the 2021 districts, city and town limits and RTD's districts
    table, vname = read_tables(bafpath, lpath, upath, cdpath, vpath)
    places = read_places(ppath)
    rtd = read_rtd(rtd_doc)
    inside = block_in(bpath, {"place": [(r["PLACEFP"], rings) for r, rings in places], "rtd": rtd})
    pre, polys = read_units(bpath, table, inside["place"], inside["rtd"], vname, cshort, say)
    if sorted({p["county"] for p in pre}) != sorted(cname):
        raise GeoError(f"    the blocks and the Bureau's list of counties do not have the same {N_COUNTY} counties; stopping")
    arcs, sides, rings_of, odd = G.topology(polys)
    if any(not r for r in rings_of):
        raise GeoError("    a precinct has no ring left after cleaning; stopping")
    fine, _restored = G.simplify_arcs(arcs, rings_of, G.TOL_PRECINCT)
    qfine = G.quantise(fine)
    lone = sum(1 for r, l in sides if r < 0 or l < 0)
    say(f"      {len(pre):,} precinct pieces, {sum(len(r) for r in rings_of):,} rings, {len(arcs):,} lines between them ({sum(len(a) for a in arcs):,} points; "
        f"{sum(len(a) for a in qfine if a):,} after generalising to {G.TOL_PRECINCT} m); {lone:,} lines are the state's edge"
        + (f"; odd: {dict(odd)}" if odd else ""))

    # ---- the Census Bureau's county subdivisions: only for the county lines of the area check
    cous = sorted(I.read_shapefile(tpath, lambda r: r["STATEFP"] == FIPS), key=lambda x: (x[0]["COUNTYFP"], x[0]["COUSUBFP"]))
    carcs, csides, _crings, codd = G.topology([rings for _r, rings in cous])
    chains, _pairs, cshapes, _loose = G.dissolve(carcs, csides, [FIPS + r["COUNTYFP"] for r, _ in cous])
    county_rings = {c: [G.ring_xy(ring, chains) for ring, _p in rings] for c, rings in cshapes.items()}
    clist = sorted(county_rings)

    # ---- cities and towns (2025 limits), one fabric for the mcd layer
    mrow = {f"{STATE}-M-{r['PLACEFP']}": r for r, _ in places}
    mkeys, _mpolys, marcs, msides, _mr, modd = I.fabric([(r, as_lonlat(rings)) for r, rings in places], lambda r: f"{STATE}-M-{r['PLACEFP']}")
    mkind = {k: PLACE_LSAD[r["LSAD"]] for k, r in mrow.items()}
    mname = {k: names.get(("mcd", k)) or r["NAMELSAD"] for k, r in mrow.items()}

    # ---- school districts: the unified districts, one fabric
    srows, sname, snces = [], {}, {}
    for r, rings in I.read_shapefile(unpath, lambda r: r["STATEFP"] == FIPS):
        if r["UNSDLEA"] == "99997":
            continue
        sid = f"{STATE}-S-{r['UNSDLEA']}"
        srows.append((dict(r, _id=sid), as_lonlat(rings)))
        sname[sid] = names.get(("school", sid)) or r["NAME"]
        snces[sid] = FIPS + r["UNSDLEA"]
    schkeys, _schpolys, scharcs, schsides, _sr, schodd = I.fabric(srows, lambda r: r["_id"])
    schfine, _ = G.simplify_arcs(scharcs, _sr, G.TOL_SCHOOL)
    say(f"      {len(mkeys)} cities and towns ({len(marcs):,} lines), {len(schkeys)} school districts ({len(scharcs):,} lines), "
        f"{len(rtd)} RTD director districts"
        + "".join(f"; odd in {n}: {dict(o)}" for n, o in (("places", modd), ("school", schodd), ("Census", codd)) if o))

    pre_rings = [[G.ring_xy(refs, fine) for refs in rings] for rings in rings_of]
    stamp = lambda *paths: hashlib.sha256(json.dumps([G.sha_file(p) for p in paths] + [G.TOL_PRECINCT, G.TOL_SCHOOL, G.SCHOOL_THICK, 1]).encode()).hexdigest()   # noqa: E731
    units_stamp = (bpath, bafpath, lpath, upath, cdpath, ppath, path("rtd_director_districts_geometry_4326.json.gz"))
    schdistricts = [[G.ring_xy(refs, schfine) for refs in rings] for rings in _sr]
    o_sch = overlay_cached("school", pre_rings, schdistricts, stamp(*units_stamp, unpath), refresh, say)

    split = none = edges = 0
    for p, got in zip(pre, o_sch):
        rows = [(schkeys[d], share, thick) for d, share, thick in got]
        keep = sorted(((k, s) for k, s, thick in rows if thick), key=lambda x: (-x[1], x[0]))
        if not keep and rows:
            best = max(((k, s) for k, s, _t in rows), key=lambda x: x[1])
            if best[1] >= 0.5:
                keep = [best]
        outside = max(0.0, 1.0 - sum(s for _k, s, _t in rows))
        p["school"] = [k for k, _s in keep]
        p["school_pct"] = [round(100 * s, 1) for _k, s in keep] if len(keep) > 1 or (keep and outside >= 0.03) else None
        p["school_out"] = round(100 * outside, 1) if outside >= 0.03 else None
        p["school_edge"] = sorted({k for k, _s, _t in rows} - set(p["school"]))
        p["jud"] = jud_of[p["county"]]
        split += len(keep) > 1
        none += not keep
        edges += bool(p["school_edge"]) and len(keep) < 2
    say(f"      school districts by precinct: {split:,} precinct pieces are split between two or more districts, {none} lie in none, "
        f"{edges:,} others only brush a neighbouring district along a line; cities and towns: {sum(1 for p in pre if p['mcd']):,} pieces lie inside one; "
        f"RTD: {sum(1 for p in pre if p['rtd']):,} pieces lie inside a director district")

    # ---- the area check: each county's precincts (2020 blocks) against the Bureau's 2025 county
    signed = lambda r: I.ring_area_m2(r) * (1 if G.area2(r) < 0 else -1)      # noqa: E731
    area_p = collections.Counter()
    for p, rings in zip(pre, pre_rings):
        for r in rings:
            area_p[p["county"]] += signed(r)
    area_c = {c: sum(signed(r) for r in county_rings[c]) for c in clist}
    if sorted(area_p) != clist:
        raise GeoError(f"    area check: the blocks and the 2025 county subdivisions do not have the same {N_COUNTY} counties; stopping")
    worst = max(clist, key=lambda c: abs(area_p[c] / area_c[c] - 1))
    worst_pct = 100 * (area_p[worst] / area_c[worst] - 1)
    state_pct = 100 * (sum(area_p.values()) / sum(area_c.values()) - 1)
    off = [c for c in clist if abs(area_p[c] / area_c[c] - 1) > AREA_SLACK]
    if off:
        raise GeoError(f"    area check: the precincts of {', '.join(cname[c] for c in off[:5])} do not cover the county the Census Bureau draws; stopping")
    say(f"      area check: the precincts cover {100 + state_pct:.3f}% of the Census Bureau's 2025 Colorado; the county furthest off is {cname[worst]} ({worst_pct:+.2f}%)")

    # ---- the Secretary's 2026 precinct numbers against the pieces: the same number in the county, and the same three districts
    listed = 0
    for p in pre:
        p["listed"] = None
        num = sos[p["county"]].get(int(p["vtd"][3:]))
        if num and code_districts(num) == (p["cd"], p["senate"], p["house"]):
            p["listed"] = num
            listed += 1
    n_2026 = sum(len(v) for v in sos.values())
    say(f"      the Secretary's 2026 precinct numbers: {n_2026:,} precincts in {len(sos)} counties; {listed:,} of {len(pre):,} pieces carry a 2020 number that is "
        "a 2026 number of their county whose own digits give the piece's congressional, Senate and House district")

    vals = {"county": [p["county"] for p in pre], "mcd": [p["mcd"] for p in pre], "house": [p["house"] for p in pre], "senate": [p["senate"] for p in pre],
            "cd": [p["cd"] for p in pre], "judicial": [p["jud"] for p in pre]}
    mask = []
    for r, l in sides:
        m = 0
        for bit, kind in enumerate(ARC_KINDS):
            if r < 0 or l < 0:
                if vals[kind][max(r, l)] is not None:
                    m |= 1 << bit
            elif vals[kind][r] != vals[kind][l]:
                m |= 1 << bit
        mask.append(m)

    # ---- write (into a folder beside the real one, moved into place only when everything is written)
    final, out = out, out.rstrip("\\/") + ".part"
    if os.path.isdir(out):
        shutil.rmtree(out)
    os.makedirs(out)
    files, today = {}, dt.date.today().isoformat()

    def put(rel, doc):
        size, sha = G.write_json(os.path.join(out, *rel.split("/")), doc)
        files[rel] = {"bytes": size, "sha256": sha}
        return size

    layers, shape_ids = [], {}

    def layer(kind, geoms, qlines, tol, props, source, zoom=None):
        gl = []
        for v, polys_ in sorted(geoms, key=lambda x: G.natkey(x[0])):
            if v is None:
                continue
            pr = props(v)
            gl.append({"id": pr.pop("id"), "properties": pr, "polys": polys_, "label": True})
        doc, _dropped, _pts = G.topo_doc({kind: gl}, qlines)
        doc["kind"], doc["tolerance_m"] = kind, tol
        size = put(f"layers/{kind}.json", doc)
        empty = [g["id"] for g in doc["objects"][kind]["geometries"] if g["type"] is None]
        if empty:
            say(f"      layer {kind}: {len(empty)} shapes with no outline ({', '.join(empty[:5])})")
        shape_ids[kind] = {g["id"]: g["properties"] for g in doc["objects"][kind]["geometries"]}
        layers.append({"kind": kind, "file": f"layers/{kind}.json", "bytes": size, "shapes": len(gl), "tolerance_m": tol,
                       "good_to_zoom": zoom or G.zoom_for(tol), "lines_from": source})

    def own(v, tol):
        g, q, _l, _r = G.build_layer(arcs, sides, [v], tol)
        return g, q

    BLOCKS, BEF, CDB, PLACE, SCH, LAW, VTD, TAB, COUSUB, RTD = (
        "co-census-tiger-2020-blocks", "co-census-2024-legislative-bef", "co-census-cd119-bef", "co-census-tiger-2025-places",
        "co-census-tiger-2025-school-districts", "co-crs-2026-title-13-judicial-districts", "co-census-2020-voting-districts",
        "co-sos-voter-statistics-precincts", "co-census-tiger-2025-cousub", "co-rtd-director-districts")
    jud_counties = collections.defaultdict(list)
    for c in sorted(cname, key=lambda c: cname[c]):
        jud_counties[jud_of[c]].append(cname[c])
    layer("state", *own([STATE] * len(pre), G.TOL_WIDE), G.TOL_WIDE, lambda v: {"id": STATE, "name": STATE_NAME, "j": FIPS, "d": None}, BLOCKS)
    layer("county", *own(vals["county"], G.TOL_WIDE), G.TOL_WIDE, lambda v: {"id": v, "name": cname[v], "j": v, "d": None, "fips": v}, BLOCKS)
    layer("cd", *own(vals["cd"], G.TOL_WIDE), G.TOL_WIDE,
          lambda v: {"id": v, "name": f"Congressional District {v}", "j": cd_id(v), "d": v, "race": f"2026-{STATE}-H{int(v):02d}"}, BLOCKS + ", put together by " + CDB)
    layer("senate", *own(vals["senate"], G.TOL_WIDE), G.TOL_WIDE,
          lambda v: {"id": v, "name": f"Senate District {v}", "j": v, "d": v}, BLOCKS + ", put together by " + BEF)
    layer("house", *own(vals["house"], G.TOL_WIDE), G.TOL_WIDE,
          lambda v: {"id": v, "name": f"House District {v}", "j": v, "d": v}, BLOCKS + ", put together by " + BEF)
    layer("judicial", *own(vals["judicial"], G.TOL_WIDE), G.TOL_WIDE,
          lambda v: {"id": v, "name": names.get(("judicial", v)) or f"{G.ordinal(v[len(STATE) + 3:])} Judicial District", "j": v,
                     "d": v[len(STATE) + 3:], "counties": jud_counties[v]}, BLOCKS + "; which counties, from " + LAW)
    mg, mq, _l, _r = G.build_layer(marcs, msides, [mkeys], Y.TOL_MCD)
    layer("mcd", mg, mq, Y.TOL_MCD, lambda v: {"id": v, "name": mname[v], "j": v, "d": None, "t": mkind[v]}, PLACE, zoom=Y.MCD_ZOOM)
    sprops = lambda v: {"id": v, "name": sname[v], "j": v, "d": None, "nces": snces[v]}      # noqa: E731
    sg, sq, _l, _r = G.build_layer(scharcs, schsides, [schkeys], G.TOL_LOCAL)
    layer("school", sg, sq, G.TOL_LOCAL, sprops, SCH)
    dgeoms, dq, _l, _r = G.build_layer(scharcs, schsides, [schkeys], G.TOL_SCHOOL)
    school_bytes = 0
    for v, polys_ in dgeoms:
        pr = sprops(v)
        pr.pop("id")
        doc, _d, _p = G.topo_doc({"school": [{"id": v, "properties": pr, "polys": polys_, "label": True}]}, dq)
        doc["kind"], doc["tolerance_m"] = "school", G.TOL_SCHOOL
        school_bytes += put(f"school/{v}.json", doc)

    # county files
    by_county = collections.defaultdict(list)
    for i, p in enumerate(pre):
        by_county[p["county"]].append(i)
    counties, boxes, dropped_rings = [], {}, 0
    for county, idxs in sorted(by_county.items()):
        local = {i: n for n, i in enumerate(idxs)}
        geoms, used_names = [], {"mcd": {}, "school": {}}
        for i in idxs:
            p = pre[i]
            pr = {"name": p["name"], "county": county, "precinct": p["precinct"], "house": p["house"], "senate": p["senate"], "cd": p["cd"],
                  "judicial": p["jud"], "school": p["school"], "as_of": 2020, "sboe": cd_id(p["cd"]), "regent": cd_id(p["cd"])}
            if p["mcd"]:
                pr["mcd"] = p["mcd"]
                used_names["mcd"][p["mcd"]] = mname[p["mcd"]]
            if p["rtd"]:
                pr["rtd"] = p["rtd"]
            if p["listed"]:
                pr["listed_2026"], pr["code_2026_said"] = True, p["listed"]
            for k in ("school_pct", "school_out", "school_edge"):
                if p.get(k):
                    pr[k] = p[k]
            for k in p["school"] + p["school_edge"]:
                used_names["school"][k] = sname[k]
            geoms.append({"id": p["id"], "properties": pr, "polys": G.group_polys([(refs, G.ring_xy(refs, fine)) for refs in rings_of[i]]), "label": True})

        def more(used, county=county, local=local, used_names=used_names):
            am, asd = [], []
            for a in used:
                am.append(mask[a])
                for side in sides[a]:
                    asd.append(local[side] if side in local else (-1 if side >= 0 else -2))
            return {"county": county, "fips": county, "name": cname[county], "arcKinds": ARC_KINDS, "arcMask": am, "arcSides": asd,
                    "names": {k: dict(sorted(v.items())) for k, v in used_names.items() if v}}

        doc, dropped, pts = G.topo_doc({"precincts": geoms}, qfine, more)
        dropped_rings += dropped
        bx = []
        for g in geoms:
            polys_pts = pts[("precincts", g["id"])]
            if not polys_pts:
                raise GeoError(f"    precinct {g['id']} has no shape on the grid; stopping")
            xs = [x for poly in polys_pts for x, _y in poly[0]]
            ys = [y for poly in polys_pts for _x, y in poly[0]]
            bx += [min(xs) // G.BOX_STEP, min(ys) // G.BOX_STEP, -(-max(xs) // G.BOX_STEP), -(-max(ys) // G.BOX_STEP)]
        boxes[county] = bx
        rel = f"precincts/{county}.json"
        size = put(rel, doc)
        counties.append({"id": county, "fips": county, "name": cname[county], "bbox": doc["bbox"], "file": rel, "bytes": size, "precincts": len(idxs)})

    polls = polling_places(put, say)
    if os.path.exists(READER_JS):
        shutil.copyfile(READER_JS, os.path.join(out, "reader.js"))
        files["reader.js"] = {"bytes": os.path.getsize(os.path.join(out, "reader.js")), "sha256": G.sha_file(os.path.join(out, "reader.js"))}

    said_ids = {"rtd": {p["rtd"] for p in pre if p["rtd"]}, "sboe": {cd_id(p["cd"]) for p in pre}, "regent": {cd_id(p["cd"]) for p in pre}}
    check = check_ids(info, shape_ids, said_ids)
    check["pieces_matched_to_the_2026_precinct_numbers"] = listed
    with open(os.path.join(out, "layers", "state.json"), encoding="utf-8") as fh:
        state_doc = json.load(fh)
    mtime = lambda p: dt.datetime.fromtimestamp(os.path.getmtime(p)).strftime("%Y-%m-%d")      # noqa: E731
    cb = "U.S. Census Bureau"
    rtd_list = ", ".join(f"{rtd_id(x)} {x}" for x in RTD_LETTERS)
    index = {
        "v": G.VERSION, "state": STATE, "election": ELECTION, "built": today, "unit": "precinct",
        "transform": {"scale": [G.SCALE, G.SCALE], "translate": [G.ORIGIN[0], G.ORIGIN[1]]}, "bbox": state_doc.get("bbox"),
        "sources": [
            {"id": BLOCKS, "agency": cb, "title": "TIGER/Line Shapefiles 2020, tabulation blocks, Colorado (tl_2020_08_tabblock20.zip)",
             "about": "The 2020 census blocks: every line of the map's precincts, counties and congressional and legislative districts is a block's.",
             "url": BLOCK_URL, "fetched": mtime(bpath), "sha256": G.sha_file(bpath), "rows": len(table)},
            {"id": VTD, "agency": cb, "title": "2020 Census Block Assignment File, voting districts (BlockAssign_ST08_CO.zip), and the voting districts' names "
                                                "(tl_2020_08_vtd20.zip)",
             "about": "The precincts Colorado's counties reported to the Bureau for the 2020 census: which voting district each block is in.",
             "url": BAF_URL, "names_url": VTD_URL, "fetched": mtime(bafpath), "sha256": G.sha_file(bafpath), "names_sha256": G.sha_file(vpath),
             "rows": len({(p["county"], p["vtd"]) for p in pre})},
            {"id": BEF, "agency": cb, "title": "2024 State Legislative District Block Equivalency Files (sldl24.zip, sldu24.zip: the national files)",
             "about": "Which House and Senate district each block is in, under the plans the Independent Legislative Redistricting Commission adopted in 2021 "
                      "(65 House and 35 Senate districts).",
             "url": SLDL_URL, "upper_url": SLDU_URL, "fetched": mtime(lpath), "sha256": G.sha_file(lpath), "upper_sha256": G.sha_file(upath), "rows": len(table)},
            {"id": CDB, "agency": cb, "title": "119th Congressional District Block Equivalency File (cd119.zip)",
             "about": "Which congressional district each block is in. " + CD_NOTE,
             "url": CD_URL, "fetched": mtime(cdpath), "sha256": G.sha_file(cdpath), "rows": len(table)},
            {"id": PLACE, "agency": cb, "title": "TIGER/Line Shapefiles 2025, places, Colorado (tl_2025_08_place.zip): the incorporated, active cities and towns",
             "about": "City and town limits as the Bureau had them on January 1, 2025. A block belongs to the city or town its own inside point lies in.",
             "url": PLACE_URL, "fetched": mtime(ppath), "sha256": G.sha_file(ppath), "rows": len(mkeys)},
            {"id": RTD, "agency": "Regional Transportation District", "title": "RTD GIS Boundaries (download view), layer DirectorDistricts",
             "about": "The fifteen director districts (A to O) of the Regional Transportation District's board. A block belongs to the director district its own "
                      "inside point lies in." + (f" The service says its data were last edited on {rtd_meta['data_last_edited']}." if rtd_meta.get("data_last_edited") else ""),
             "url": RTD_SERVICE, "page": RTD_ITEM, "fetched": rtd_doc.get("fetched"),
             "sha256": G.sha_file(path("rtd_director_districts_geometry_4326.json.gz")), "rows": len(rtd)},
            {"id": SCH, "agency": cb, "title": "TIGER/Line Shapefiles 2025, unified school districts, Colorado (tl_2025_08_unsd.zip)",
             "about": "School district lines as the Bureau had them for the 2024-2025 school year, from the State's own reporting.",
             "url": UNSD_URL, "fetched": mtime(unpath), "sha256": G.sha_file(unpath), "rows": len(schkeys)},
            {"id": COUSUB, "agency": cb, "title": "TIGER/Line Shapefiles 2025, county subdivisions, Colorado (tl_2025_08_cousub.zip)",
             "about": "Read only for the county areas the precincts are checked against (Colorado's county subdivisions are census county divisions, not governments).",
             "url": COUSUB_URL, "fetched": mtime(tpath), "sha256": G.sha_file(tpath), "rows": len(cous)},
            {"id": LAW, "agency": "Colorado General Assembly, Office of Legislative Legal Services",
             "title": f"Colorado Revised Statutes {law.get('edition')}, sections 13-5-102 to 13-5-123.1: the judicial districts",
             "url": law.get("url"), "fetched": law.get("read"), "sha256": law.get("sha256"), "rows": N_JUD},
            {"id": TAB, "agency": "Colorado Secretary of State, Elections Division",
             "title": "Voter registration statistics, sheet 'Voter Counts by Precinct' (the newest monthly workbook)",
             "about": "Every 2026 precinct's ten-digit number (congressional, Senate and House district, county, number in the county); only those two "
                      "cells are read. Used as a check (listed_2026).",
             "url": stat_url, "page": SOS_STATS_PAGE, "fetched": mtime(statpath), "sha256": G.sha_file(statpath), "rows": n_2026},
        ],
        "notes": {
            "lines": f"Every precinct, county, congressional and legislative line is a 2020 census block's, generalised by at most {G.TOL_PRECINCT} metres and set "
                     "on a grid of 0.00001 degree (about a metre). A point within a few metres of a line can fall on either side of it.",
            "precincts": "Colorado publishes no statewide map of its precincts as they stand in 2026 that can be read here. Each shape is a voting district as "
                         "the county reported it to the Census Bureau for the 2020 census, cut wherever a congressional or legislative district of the 2021 "
                         "plans, a city or town limit or an RTD director district crosses it. The counties re-precincted after 2021, so the county clerk is "
                         "the authority on a voter's precinct. listed_2026 marks a piece whose 2020 number is a 2026 precinct number of its county in the "
                         "Secretary of State's list, whose own digits give the piece's congressional, Senate and House district (code_2026_said gives "
                         "that ten-digit number). What a ballot depends on (county, congressional, Senate and House district, judicial district, city or "
                         "town, RTD director district, school district) does not depend on the precinct.",
            "places": "Colorado's incorporated cities and towns (Denver and Broomfield are each a city and county) have no townships around them. A "
                      "precinct piece lies wholly inside one city or town (mcd) or outside every one: a piece in unincorporated country names no place. A "
                      "block is given to the city or town its own inside point lies in, by the limits of January 1, 2025, so a recent annexation that cuts a "
                      "block is drawn to the nearest block line.",
            "commissioners": "No statewide file has the lines of county commissioner districts; each county keeps its own (commissioners are elected by the "
                             "whole county with a residence district, or by district in a few home rule counties).",
            "legislative": "Each of the 65 House districts and 35 Senate districts elects one member. Which district a block is in is the Census Bureau's 2024 "
                           "equivalency file's word for the 2021 plans, so a precinct piece's districts are exact.",
            "boards": "The State Board of Education and the University of Colorado's Board of Regents each have one seat elected by each congressional "
                      "district (and members elected statewide); sboe and regent name the seat a precinct votes for.",
            "rtd": "The Regional Transportation District's fifteen directors are elected by director district. rtd names the district a piece lies in; a "
                   "piece outside the RTD has none. RTD is the authority on its districts.",
            "school": "School district lines are the Census Bureau's 2025 file. Which districts a precinct piece lies in is analysis, not an official list: a "
                      f"district counts when its piece of the precinct is at least {G.SCHOOL_THICK:.0f} metres thick somewhere; school_pct is each "
                      "district's share of the piece's area (land and water, not voters). School boards are elected in odd-numbered years.",
            "judicial": "The 23 judicial districts are whole counties (C.R.S. 13-5-102 to 13-5-123.1); district court judges stand for retention in them, "
                        "and county court judges in their county.",
            "wards": "No statewide file has the lines of city council districts or wards: a council district's race is placed on its city, and the city "
                     "clerk says which district an address is in.",
            "authority": "For which precinct an address is in, the county clerk and recorder is the authority; for where to vote, GoVoteColorado and the "
                         "county clerk.",
            "precinct_ids": "A precinct's id is its county's five digits, a full stop and the Census Bureau's 2020 voting district code (08001.001230); a "
                            "voting district cut by a district line, a city limit or an RTD line has a second full stop and the piece's number.",
        },
        "format": {
            "files": "Every file is TopoJSON on one grid (transform above): arcs are delta-encoded lines; a shape's arcs name the lines of "
                     "its rings, a negative number n meaning line ~n backwards; outer rings run counter-clockwise, holes clockwise.",
            "county_files": "precincts/<county>.json, object 'precincts': one shape per precinct piece. arcMask[i] has bit k set when line i is an outline of "
                            "arc_kinds[k]; arcSides[2i] and arcSides[2i+1] are the precincts on the right and left of line i (-1: a precinct in another "
                            "county's file; -2: outside Colorado); names gives the names of the places the file's precincts lie in.",
            "boxes": "boxes[county] holds four whole numbers a precinct, in the file's order: west, south, east, north in steps of box_step degrees "
                     "from the transform's translate, rounded outwards. A point is tried against every precinct whose box (widened by half a step) holds it.",
            "precinct_properties": {"name": "the precinct's name: the 2020 voting district's as the Census Bureau's file gives it (the county and the precinct's "
                                            "number), and the piece's number where a line cuts it",
                                    "county": "county id", "precinct": "the county's five digits and the Bureau's 2020 voting district code",
                                    "as_of": "the year of the precinct's lines (2020)",
                                    "listed_2026": "true where the Secretary of State's 2026 list has a precinct of the county with the same number whose own "
                                                   "digits give the piece's congressional, Senate and House district",
                                    "code_2026_said": "with listed_2026: that 2026 precinct's ten-digit number",
                                    "mcd": "the city or town the piece lies in (absent in unincorporated country)",
                                    "house": "House district", "senate": "Senate district", "cd": "congressional district",
                                    "judicial": "judicial district",
                                    "sboe": said_words(SBOE_WORD),
                                    "regent": said_words(REGENT_WORD),
                                    "rtd": f"Regional Transportation District director district (absent outside the RTD; {rtd_list})",
                                    "school": "list: the school districts the piece lies in, largest share first",
                                    "school_pct": "list, when the piece is in more than one: each district's share of its area, in percent",
                                    "school_out": "percent of the piece's area that lies in no school district (given from 3 percent up)",
                                    "school_edge": "list: neighbouring districts that only brush the piece along a line (try them too when placing a point)",
                                    "c": "a point inside the piece's largest part, for a label: [longitude, latitude]"},
            "ids": {"every layer": "a shape's properties j and d are the jurisdiction id and the district its races carry on the ballot pages",
                    "state": "CO (j is 08)", "county": "the county's five-digit code (08031)",
                    "mcd": "CO-M- and the Census Bureau's place code (CO-M-20000); properties.t says city or town",
                    "house": "the district (6)", "senate": "the district (31)",
                    "cd": "the district (1); properties.j is CO-CD1, the jurisdiction of the State Board of Education's and the regents' seats; properties.race "
                          "is the race for Congress",
                    "judicial": "CO-JD and the district's number (CO-JD2); d is the number",
                    "school": "CO-S- and the Census Bureau's five-digit district code (CO-S-03360); properties.nces is the federal district code",
                    "rtd": "a precinct property only: CO-X-RTD- and the director district's letter (CO-X-RTD-B)",
                    "sboe and regent": "precinct properties only: CO-CD and the congressional district's number"},
        },
        "counties": counties, "box_step": G.BOX_STEP * G.SCALE, "boxes": boxes,
        "layers": layers,
        "school": {"files": "school/<id>.json", "ids": sorted(schkeys, key=G.natkey), "bytes": school_bytes, "tolerance_m": G.TOL_SCHOOL},
        "tolerance_m": {"precincts": G.TOL_PRECINCT, "school": G.TOL_SCHOOL},
        "arc_kinds": ARC_KINDS,
        "supervisor_plans": {},
        "polling_places": polls,
        "counts": {"precincts": len(pre), "voting_districts_2020": len({(p["county"], p["vtd"]) for p in pre}), "counties": len(counties), "census_blocks": len(table),
                   "pieces_in_a_city_or_town": sum(1 for p in pre if p["mcd"]), "pieces_in_the_rtd": sum(1 for p in pre if p["rtd"]),
                   "pieces_matched_to_the_2026_precinct_numbers": listed, "precincts_in_the_2026_list": n_2026,
                   "split_between_school_districts": split, "rings_too_small_for_the_grid": dropped_rings, "lines_with_a_precinct_on_one_side": lone,
                   "area_against_census_counties_percent": {"state": round(state_pct, 3), "furthest_county": worst, "its_difference": round(worst_pct, 2)}},
        "check": check,
    }
    put("index.json", index)
    manifest = {"v": G.VERSION, "built": today, "files": dict(sorted(files.items()))}
    G.write_json(os.path.join(out, "manifest.json"), manifest)
    total = sum(f["bytes"] for f in files.values()) + os.path.getsize(os.path.join(out, "manifest.json"))
    os.makedirs(final, exist_ok=True)
    G.clear_out(final)
    for rel in list(files) + ["manifest.json"]:
        dst = os.path.join(final, *rel.split("/"))
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        os.replace(os.path.join(out, *rel.split("/")), dst)
    shutil.rmtree(out)
    say(f"      {len(counties)} county files ({sum(c['bytes'] for c in counties) / 1e6:.1f} MB), {len(layers)} layers "
        f"({sum(l['bytes'] for l in layers) / 1e6:.1f} MB), {len(schkeys)} school district files ({school_bytes / 1e6:.1f} MB), "
        f"index {files['index.json']['bytes'] / 1e3:.0f} KB; {total / 1e6:.1f} MB in all, in {final}")
    say(f"      ids against the ballot database: {check['matched']:,} of {check['races']:,} Colorado races have a shape"
        + (f"; {check['races_named_by_the_precinct']} more are named by the precincts ({check['named_by_the_precinct']})" if check.get("races_named_by_the_precinct") else "")
        + (f"; {len(check['no_shape'])} kinds of place without a shape (listed in index.json)" if check["no_shape"] else ""))
    say(f"    Colorado ballot map: built in {time.time() - t0:.0f} s")
    return index


# ---------------------------------------------------------------- the self-test: what a device does, done here

TEST_POINTS = [
    # name, longitude, latitude, what the point's precinct lies in, the city or town at the point (None outside every one),
    # the federal code (NCES) of its school district, and its RTD director district (None outside the RTD). County, place,
    # congressional district, 2026 legislative districts (upper and lower) and school district are the Census Bureau's
    # geocoder's answer for those coordinates (asked 2026-10-02: geocoding.geo.census.gov, "geographies/coordinates",
    # Current), an answer that owes nothing to the files tested here; the RTD letter is RTD's own feature service's answer
    # to a point query. The judicial district is the county's (C.R.S. 13-5-102 on).
    ("the State Capitol, Denver", -104.9848, 39.7393, {"county": "08031", "cd": "1", "senate": "31", "house": "6", "judicial": "CO-JD2"}, "CO-M-20000", "0803360", "A"),
    ("Denver, Washington Park", -104.9700, 39.7000, {"county": "08031", "cd": "1", "senate": "31", "house": "2", "judicial": "CO-JD2"}, "CO-M-20000", "0803360", "A"),
    ("Colorado Springs City Hall", -104.8226, 38.8339, {"county": "08041", "cd": "5", "senate": "11", "house": "18", "judicial": "CO-JD4"}, "CO-M-16000", "0803060", None),
    ("downtown Boulder, Pearl Street", -105.2790, 40.0176, {"county": "08013", "cd": "2", "senate": "18", "house": "10", "judicial": "CO-JD20"}, "CO-M-07850", "0802490", "O"),
    ("Old Town Fort Collins", -105.0770, 40.5875, {"county": "08069", "cd": "2", "senate": "14", "house": "53", "judicial": "CO-JD8"}, "CO-M-27425", "0803990", None),
    ("downtown Pueblo", -104.6091, 38.2683, {"county": "08101", "cd": "3", "senate": "3", "house": "46", "judicial": "CO-JD10"}, "CO-M-62000", "0806120", None),
    ("downtown Grand Junction", -108.5640, 39.0685, {"county": "08077", "cd": "3", "senate": "7", "house": "55", "judicial": "CO-JD21"}, "CO-M-31660", "0804350", None),
    ("downtown Greeley", -104.6910, 40.4233, {"county": "08123", "cd": "8", "senate": "13", "house": "50", "judicial": "CO-JD19"}, "CO-M-32155", "0804410", None),
    ("Aurora Municipal Center", -104.7615, 39.7110, {"county": "08005", "cd": "6", "senate": "28", "house": "36", "judicial": "CO-JD18"}, "CO-M-04000", "0802340", "F"),
    ("Lakewood Civic Center", -105.0815, 39.7091, {"county": "08059", "cd": "7", "senate": "22", "house": "30", "judicial": "CO-JD1"}, "CO-M-43000", "0804800", "M"),
    ("downtown Durango", -107.8801, 37.2753, {"county": "08067", "cd": "3", "senate": "6", "house": "59", "judicial": "CO-JD6"}, "CO-M-22035", "0803480", None),
    ("downtown Steamboat Springs", -106.8317, 40.4850, {"county": "08107", "cd": "2", "senate": "8", "house": "26", "judicial": "CO-JD14"}, "CO-M-73825", "0806660", None),
    ("downtown Castle Rock", -104.8561, 39.3722, {"county": "08035", "cd": "4", "senate": "2", "house": "45", "judicial": "CO-JD23"}, "CO-M-12415", "0803450", None),
    ("downtown Brighton", -104.8205, 39.9853, {"county": "08001", "cd": "8", "senate": "13", "house": "48", "judicial": "CO-JD17"}, "CO-M-08675", "0802580", "K"),
    ("Broomfield city hall", -105.0526, 39.9205, {"county": "08014", "cd": "7", "senate": "25", "house": "33", "judicial": "CO-JD17"}, "CO-M-09280", "0806900", "I"),
    ("open plains east of Limon", -103.50, 39.00, {"county": "08073", "cd": "4", "senate": "35", "house": "56", "judicial": "CO-JD23"}, None, "0804740", None),
    ("downtown Alamosa", -105.8700, 37.4695, {"county": "08003", "cd": "3", "senate": "6", "house": "62", "judicial": "CO-JD12"}, "CO-M-01090", "0802070", None),
    ("downtown Glenwood Springs", -107.3248, 39.5505, {"county": "08045", "cd": "3", "senate": "5", "house": "57", "judicial": "CO-JD9"}, "CO-M-30780", "0804260", None),
    ("downtown Sterling", -103.2077, 40.6255, {"county": "08075", "cd": "4", "senate": "1", "house": "63", "judicial": "CO-JD13"}, "CO-M-73935", "0806690", None),
    ("Highlands Ranch, Douglas County", -104.9690, 39.5530, {"county": "08035", "cd": "4", "senate": "30", "house": "43", "judicial": "CO-JD23"}, None, "0803450", "H"),
    ("downtown Leadville", -106.2925, 39.2508, {"county": "08065", "cd": "7", "senate": "4", "house": "13", "judicial": "CO-JD5"}, "CO-M-44320", "0805190", None),
    ("Westminster city hall", -105.0372, 39.8367, {"county": "08001", "cd": "8", "senate": "21", "house": "35", "judicial": "CO-JD17"}, "CO-M-83835", "0807230", "J"),
    ("Arvada, Olde Town", -105.0811, 39.8028, {"county": "08059", "cd": "7", "senate": "19", "house": "24", "judicial": "CO-JD1"}, "CO-M-03455", "0804800", "L"),
]
LINE_POINTS = [("the Denver-Jefferson county line on Sheridan Boulevard", -105.0532, 39.7280, ("08031", "08059")),
               ("the Arapahoe-Douglas county line on County Line Road", -104.8800, 39.5660, ("08005", "08035"))]
LAYER_SLACK = 0.005


def selftest(out=OUT, say=print):
    """Reads only the built files. Returns True when everything holds."""
    files = G.Files(out)
    index, fails = files.index, []
    locate = W.locate

    def check(ok, what):
        if not ok:
            fails.append(what)
            say(f"      FAIL: {what}")
        return ok

    def own(geom, kind):
        v = geom["properties"].get(kind)
        return v[0] if isinstance(v, list) else v

    # 1. the files are all there and say what the index says; lines know their two sides and their outline marks
    total = bad_side = bad_mask = lines_seen = edge_lines = 0
    ids = set()
    for c in index["counties"]:
        doc, lines = files.topo(c["file"])
        geoms = doc["objects"]["precincts"]["geometries"]
        total += len(geoms)
        ids |= {g["id"] for g in geoms}
        ok = (len(geoms) == c["precincts"] == len(index["boxes"][c["id"]]) // 4 and len(doc["arcMask"]) == len(doc["arcs"])
              and len(doc["arcSides"]) == 2 * len(doc["arcs"]) and all(g["type"] in ("Polygon", "MultiPolygon") for g in geoms)
              and doc["arcKinds"] == index["arc_kinds"])
        if not check(ok, f"{c['file']} does not fit the index"):
            continue
        sides_, masks = doc["arcSides"], doc["arcMask"]
        for gi, g in enumerate(geoms):
            for poly in (g["arcs"] if g["type"] == "MultiPolygon" else [g["arcs"]]):
                for refs in poly:
                    for r in refs:
                        bad_side += sides_[2 * r + 1 if r >= 0 else 2 * ~r] != gi
        for a in range(len(lines)):
            r, l = sides_[2 * a], sides_[2 * a + 1]
            lines_seen += 1
            edge_lines += min(r, l) < 0
            for bit, kind in enumerate(doc["arcKinds"]):
                if r >= 0 and l >= 0:
                    want = own(geoms[r], kind) != own(geoms[l], kind)
                elif min(r, l) == -2:
                    want = own(geoms[max(r, l)], kind) is not None
                else:
                    continue
                bad_mask += bool(masks[a] & (1 << bit)) != want
    check(total == index["counts"]["precincts"] == len(ids), "the county files do not hold as many precincts as the index says, each with its own id")
    shapes_seen, layer_ids = 0, {}
    for L in index["layers"]:
        gs = files.topo(L["file"])[0]["objects"][L["kind"]]["geometries"]
        shapes_seen += len(gs)
        layer_ids[L["kind"]] = {g["id"] for g in gs}
        empty = [g["id"] for g in gs if g["type"] not in ("Polygon", "MultiPolygon")]
        check(len(gs) == L["shapes"] and not empty and len({g["id"] for g in gs}) == len(gs), f"layer {L['kind']}: {len(empty)} shapes without an outline ({', '.join(empty[:4])}), or ids repeated")
    empty = [i for i in index["school"]["ids"] if files.topo(f"school/{i}.json")[0]["objects"]["school"]["geometries"][0]["type"] not in ("Polygon", "MultiPolygon")]
    check(not empty, f"{len(empty)} school district files have no outline ({', '.join(empty[:4])})")
    check(len(index["counties"]) == N_COUNTY, f"there are not {N_COUNTY} county files")
    check(bad_side == 0, f"{bad_side} rings use a line whose sides do not name them")
    check(bad_mask == 0, f"{bad_mask} outline marks disagree with the precincts on the two sides of a line")
    check(layer_ids.get("senate") == {str(n) for n in range(1, N_SENATE + 1)} and layer_ids.get("house") == {str(n) for n in range(1, N_HOUSE + 1)}
          and layer_ids.get("judicial") == {f"{STATE}-JD{n}" for n in range(1, N_JUD + 1)} and layer_ids.get("cd") == {str(n) for n in range(1, N_CD + 1)}
          and len(layer_ids.get("county", ())) == N_COUNTY and 265 <= len(layer_ids.get("mcd", ())) <= 280 and len(index["school"]["ids"]) >= 170,
          "there are not 35 Senate, 65 House, 23 judicial and 8 congressional districts, 64 counties, about 272 cities and towns and about 178 school districts")
    say(f"      self-test: {len(index['counties'])} county files, {total:,} precinct pieces, every one with a shape; {lines_seen:,} lines "
        f"({edge_lines:,} on a file's edge), sides and outline marks all in order; {len(index['layers'])} layers, {shapes_seen:,} shapes and "
        f"{len(index['school']['ids'])} school district files, every one with an outline")

    # 2. every precinct is found again from a point inside it; every id it carries is a shape of that layer; the layers agree with it
    wrong, tested, agree, skipped, no_shape = 0, 0, 0, 0, collections.Counter()
    disagree = collections.defaultdict(list)
    asked = collections.Counter()
    school, place, rtd_seen, board_bad = collections.Counter(), collections.Counter(), set(), 0
    layer_for = {"county": "county", "house": "house", "senate": "senate", "cd": "cd", "judicial": "judicial"}
    tol_of = {l["kind"]: l["tolerance_m"] for l in index["layers"]}
    for c in index["counties"]:
        doc, lines = files.topo(c["file"])
        for i, g in enumerate(doc["objects"]["precincts"]["geometries"]):
            polys = G.geom_polys(lines, g)
            px, py = G.label_point(polys[0])
            lon, lat = px * G.SCALE + G.ORIGIN[0], py * G.SCALE + G.ORIGIN[1]
            found, _near = locate(files, lon, lat)
            tested += 1
            if not found or found["geometry"]["id"] != g["id"]:
                wrong += 1
            pr = g["properties"]
            board_bad += not (pr.get("sboe") == pr.get("regent") == f"{STATE}-CD{pr.get('cd')}")
            if pr.get("rtd"):
                rtd_seen.add(pr["rtd"])
            for prop, kind in list(layer_for.items()) + [("mcd", "mcd"), ("school", "school")]:
                v = pr.get(prop)
                for one in (v if isinstance(v, list) else [v] if v else []):
                    no_shape[kind] += one not in layer_ids.get(kind, ())
            sid = G.school_at(files, g, lon, lat)
            school["in a district the precinct lies in" if sid in pr["school"] else "in a district the precinct only brushes" if sid else
                   "in none, in a precinct partly outside every district" if pr.get("school_out") else "in none"] += 1
            shape, edge = G.shape_at(files, "mcd", lon, lat)
            place["the piece's own city or town" if shape and shape["id"] == pr.get("mcd") else
                  "unincorporated, as the piece says" if not shape and not pr.get("mcd") else
                  "within a city's 2025 limits, on a block whose inside point is not" if shape and not pr.get("mcd") else
                  "outside the 2025 limits of the city the piece names" if not shape else "in another city"] += 1
            for prop, kind in layer_for.items():
                if i % 3 and kind not in ("house",):
                    continue
                shape, edge = G.shape_at(files, kind, lon, lat)
                if shape is not None and edge <= tol_of[kind] + 10:
                    skipped += 1
                    continue
                asked[kind] += 1
                if shape is not None and shape["id"] == pr.get(prop):
                    agree += 1
                else:
                    disagree[kind].append((g["id"], shape and shape["id"], pr.get(prop)))
    check(wrong <= 3, f"{wrong} of {tested} precincts are not found again from a point inside them")
    check(not sum(no_shape.values()), f"precincts carry ids that are no shape of their layer: {dict(no_shape)}")
    check(board_bad == 0, f"{board_bad} precincts name a State Board of Education or regent seat that is not their congressional district's")
    check(rtd_seen == {f"{STATE}-X-RTD-{x}" for x in RTD_LETTERS}, f"the precincts name {len(rtd_seen)} RTD director districts, not the fifteen")
    for kind, bad in disagree.items():
        check(len(bad) <= max(2, int(LAYER_SLACK * asked[kind])), f"layer {kind}: {len(bad)} of {asked[kind]} shapes disagree with the precinct at a point well inside it, e.g. {bad[:3]}")
    check(school["in none"] <= 0.005 * tested, f"{school['in none']} precincts' own points fall in no school district though the precinct is said to lie in one")
    check(place["the piece's own city or town"] + place["unincorporated, as the piece says"] >= 0.97 * tested, f"the city or town at a piece's own point: {dict(place)}")
    say(f"      self-test: {tested - wrong:,} of {tested:,} precinct pieces found again from a point inside them; every id a piece carries is a shape; layers agree at {agree:,} of "
        f"{sum(asked.values()):,} points ({skipped} too near a line to ask"
        + "".join(f"; {kind}: {len(bad)} of {asked[kind]} differ" for kind, bad in sorted(disagree.items())) + "); school district at each piece's point: "
        + "; ".join(f"{n:,} {what}" for what, n in sorted(school.items(), key=lambda x: -x[1])) + "; city or town: "
        + "; ".join(f"{n:,} {what}" for what, n in sorted(place.items(), key=lambda x: -x[1])) + f"; {len(rtd_seen)} RTD director districts named")

    # 3. known points
    rel = {c["id"]: c["file"] for c in index["counties"]}
    nces = {g["id"]: g["properties"].get("nces") for g in files.topo("layers/school.json")[0]["objects"]["school"]["geometries"]}
    for name, lon, lat, want, mcd, code, rtd in TEST_POINTS:
        found, _near = locate(files, lon, lat)
        if not check(found is not None, f"{name}: in no precinct"):
            continue
        pr = found["geometry"]["properties"]
        got = {k: pr.get(k) for k in want}
        ok = check(got == want, f"{name}: lands in {got}, not {want}")
        sid = G.school_at(files, found["geometry"], lon, lat)
        ok &= check(sid is not None and sid in pr["school"] and nces.get(sid) == code, f"{name}: the school district at the point is {sid} ({nces.get(sid)}); the geocoder says {code}")
        shape, _edge = G.shape_at(files, "mcd", lon, lat)
        ok &= check((shape and shape["id"]) == mcd and pr.get("mcd") == mcd, f"{name}: the mcd layer gives {shape and shape['id']} and the piece {pr.get('mcd')}; the place is {mcd}")
        ok &= check(pr.get("rtd") == (f"{STATE}-X-RTD-{rtd}" if rtd else None), f"{name}: the piece's RTD director district is {pr.get('rtd')}; RTD's own service says {rtd}")
        for prop, kind in layer_for.items():
            shape2, _edge = G.shape_at(files, kind, lon, lat)
            ok &= check(shape2 is not None and shape2["id"] == pr.get(prop), f"{name}: layer {kind} gives {shape2 and shape2['id']}, the piece says {pr.get(prop)}")
        fdoc, _l = files.topo(rel[found["county"]])
        say(f"      self-test: {name}: {'ok' if ok else 'FAIL'}: {pr['name']} ({found['geometry']['id']}), {fdoc['names'].get('mcd', {}).get(pr.get('mcd'), 'unincorporated')}, "
            f"{fdoc['name']}, CD {pr['cd']}, House {pr['house']}, Senate {pr['senate']}, {pr['judicial']}, RTD {pr.get('rtd') or 'none'}, "
            f"{fdoc['names']['school'].get(sid, sid)}; {found['edge']:.0f} m from the piece's line")

    # 4. county lines: the spot on one county's own drawing of the line nearest to a chosen point
    for name, lon, lat, pair in LINE_POINTS:
        doc, lines = files.topo(rel[pair[0]])
        bit = 1 << doc["arcKinds"].index("county")
        x, y = G.to_grid(lon, lat)
        cs = math.cos(math.radians(lat))
        best = None
        for a, pts in enumerate(lines):
            if doc["arcMask"][a] & bit and -1 in (doc["arcSides"][2 * a], doc["arcSides"][2 * a + 1]):
                for (ax, ay), (bx, by) in zip(pts, pts[1:]):
                    d = G.seg_dist2(x * cs, y, ax * cs, ay, bx * cs, by)
                    if best is None or d < best[0]:
                        best = (d, (ax + bx) / 2.0, (ay + by) / 2.0)
        on = (best[1] * G.SCALE + G.ORIGIN[0], best[2] * G.SCALE + G.ORIGIN[1])
        f2, n2 = locate(files, on[0], on[1])
        sides_ = ({f2["county"]} | {h["county"] for h in n2}) if f2 else set()
        ok = check(f2 is not None and f2["edge"] < 1.0 and sides_ == set(pair),
                   f"a point on {pair[0]}'s drawing of the {pair[0]}/{pair[1]} line lands in {f2 and f2['county']} with {sorted(sides_)} at hand, {f2 and round(f2['edge'], 2)} m from the line")
        if f2:
            say(f"      self-test: {name}: {'ok' if ok else 'FAIL'}: the spot on the line ({on[0]:.5f}, {on[1]:.5f}) is given to "
                f"{f2['geometry']['properties']['name']} with {', '.join(sorted(h['geometry']['properties']['name'] for h in n2))} at hand")

    # 5. polling places say what they are
    polls = files.load("polling_places.json")
    check(polls.get("status") in ("waiting", "unchecked", "loaded"), "polling_places.json has no status")
    check(not re.search(r"\(\d{3}\)\s*\d{3}|\d{3}-\d{3}-\d{4}|@", json.dumps(polls)), "polling_places.json carries something that reads like a phone number or an e-mail address")
    say(f"      self-test: polling places: {polls.get('status')}")

    # 6. every precinct name is the county's name and a number (or "Voting district" and a number): no other name reaches a file
    odd_names = 0
    for c in index["counties"]:
        doc, _l = files.topo(c["file"])
        short = short_county(doc["name"])
        for g in doc["objects"]["precincts"]["geometries"]:
            odd_names += not re.fullmatch(rf"(?:{re.escape(short)}|Voting district) \d+(?:, part \d+)?", g["properties"]["name"], flags=re.I)
    check(odd_names == 0, f"{odd_names} precinct names are not a county's name or 'Voting district' and a number")
    say("    self-test: " + ("PASS" if not fails else f"FAIL ({len(fails)})"))
    return not fails


def main(argv=None):
    ap = argparse.ArgumentParser(description="The geography behind Colorado's ballot map -> ballot_geo/co/")
    ap.add_argument("--out", default=OUT, help="folder to build (default: ballot_geo/co)")
    ap.add_argument("--db", default=DB, help="ballot database to read names from, read-only")
    ap.add_argument("--refresh", action="store_true", help="ask for the Secretary of State's precinct numbers and RTD's districts again even when the cached copies are fresh")
    ap.add_argument("--selftest", action="store_true", help="only run the self-test on the files already built")
    a = ap.parse_args(argv)
    out = os.path.abspath(a.out)
    if not a.selftest:
        build(out=out, db=os.path.abspath(a.db), refresh=a.refresh)
    if not selftest(out):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
