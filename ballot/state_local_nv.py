"""
ballot/state_local_nv.py - Nevada's state races on the November 3, 2026 ballot, and the county and local races of
Clark and Washoe counties (see "County and local races" below), into ballot_local_2026.sqlite (the federal
ballot_2026.sqlite is never opened here). The state races:

  - the six statewide offices, each elected on its own, every four years, at the governor's election: Governor,
    Lieutenant Governor, Secretary of State, Attorney General, State Treasurer and State Controller (Nevada
    Constitution, Article 5, Sections 2, 17 and 19, read from the Legislature's copy);
  - the State Senate seats whose four-year terms end in 2026. Nevada's Senate follows no odd-and-even rule, so the
    seats are read from the Legislature's own list of current senators, which gives each one's "Term Ends" year
    (2026-09-30: districts 2, 8, 9, 10, 12, 13, 14, 16, 17, 20 and 21, eleven of the twenty-one);
  - all 42 seats of the Assembly (two-year terms; the Legislature's list gives every one as ending in 2026);
  - Supreme Court and Court of Appeals seats wherever the certified list carries them (judges are elected without
    party labels; stored as "Nonpartisan office"). Any other office on a saved list (district courts, the Board of
    Regents, county and city offices) is counted by its title in the report, never stored and never dropped silently;
with each party's June 9, 2026 closed primary field and its official votes.

Sources, the State of Nevada's own:
  - The Secretary of State's Elections Division: the certified candidate list for the 2026 General Election and the
    2026 Official Statewide Primary Election Results. www.nvsos.gov and silverstateelection.nv.gov answer scripts with
    an Incapsula challenge (ballot/lists/nv.py), which is never worked around, so nothing is requested from them: this
    loader reads what a browser saves into ballot_cache/nv/, the folder and file names the federal loader uses
    (nv_2026_general_candidates.*, nv_2026_primary_results.*, optionally nv_2026_candidate_filings.*, each .html,
    .htm, .md, .csv or .xlsx), plus any further copy named with a suffix (nv_2026_primary_results_legislature.html and
    the like), since the certified list and the results carry every office, not only Congress. A saved copy of the
    wall's own page is refused; a PDF is reported, not read. A results file that does not call itself official, or
    calls itself unofficial, is refused: only canvassed figures are stored.
  - The Nevada Legislature (Legislative Counsel Bureau), www.leg.state.nv.us: its lists of current senators and
    Assembly members (district, party, the counties each district reaches, and "Term Ends"), and the Constitution.
    These answer scripts; they are asked once a week at most, and only what is named here is kept, as
    ballot_cache/nv/sl_nv_legislature.json with each page's SHA-256.
  - Today's holders from the Open States roster in state_nv.sqlite (legislators with is_current = 1 by chamber and
    district; the officials table for the Governor, Lieutenant Governor, Secretary of State and Attorney General). It
    does not carry the Treasurer or the Controller, so no holder is shown for those. Only ids, names, parties,
    chambers, districts and term ends are selected; its e-mail, phone and address columns never are.
  - County codes from the Census Bureau's 2024 county file (states_cache/census/), matched by name to the
    Legislature's county list for each district.

Privacy. From any list, only the office (or contest), district, candidate's name, party, ballot order and status
cells are ever turned into text, and from results the names and votes. Columns are found by their headings (a row of
<th> cells, or one of a table's first rows); every other cell, addresses, cities, ZIP codes, telephones, websites,
e-mail and treasurers among them, is never turned into text, printed, logged, cached or stored, and no list is copied
or saved. An error names the file and the check, never a row. From the Legislature's pages only the name, party,
district and county cells of each member's row and the "Term Ends" field are read. Before anything is written, every
stored text is checked for anything that looks like a contact detail, and the load stops (without showing it) if one
does.

Matching. A candidate is marked as the seat's holder (incumbent 1, state_member_id) only when the name fits the
roster's holder of that same seat (or the Legislature's spelling of that member) and no other candidate in the race
fits. A candidate who serves today in another seat or office gets state_member_id with incumbent 0 and a note, only
when exactly one roster person of the same party fits.

Primaries. Nevada's major parties nominate in closed primaries; a candidate alone in a party's contest is not printed
on the primary ballot, so a field is two or more names. Minor parties nominate otherwise, and Nevada allows no
write-in votes. Shares are of the contest's total, which for a statewide office includes "None of these candidates"
(a choice Nevada prints for statewide offices; it is not a candidate and is not stored as one, and its votes are named
in the race's note). Who advanced is the party's candidate on the November list; until that list is saved, the most
votes (Nevada nominates by plurality), said so in a note. A judicial primary keeps the two the November list names.

Ballot order: stored only if a list prints an order column; otherwise left empty.

County and local races (a narrow pass, 2026-10-01: Clark and Washoe counties, about nine in ten Nevadans)
---------------------------------------------------------------------------------------------------------
The Secretary of State's statewide list is behind the bot wall, but each county's election office posts its own list
of November contests and candidates. Two are read here, each fetched with states/net.py's honest User-Agent, one
request at a time, and kept in ballot_cache/nv/local/ for a week (kept.json says where each file came from and when):

  - Clark County Election Department, "All Contests and Candidates in the November 3, 2026, General Election in Clark
    County, Nevada" (a ten-page PDF linked as "Candidates and Contests"). A ruled table of five columns: CONTEST,
    CANDIDATE NAME ON BALLOT (Last Name; First, Middle, and/or Nickname), PARTY and TERM (Years). The contest and term
    cells are centred beside their candidates, so the table is read by its own ruled lines: a rule across the whole
    table ends a contest, a rule across the name and party columns ends a candidate. The list says it is "in order of
    appearance on the ballot, subject to change", so its order is stored as the ballot order.
  - Washoe County Registrar of Voters, "Washoe County 2026 General Election Contests" (a workbook of one sheet):
    Name on Ballot, Party, Office Sought, Office Type, Term Length, Vote For, found by those headings; ballot questions
    follow in a second table and are not read. It states no ballot order, so none is stored.

Neither file has a contact column of any kind (no address, telephone, e-mail or website), so each is kept whole. Only
the contest, the name on the ballot, the party, the term and the vote-for number are read. A name cell that looks like
a contact detail would be left out and counted, never printed. A layout that no longer fits stops the loader, which
names the file, the page or row and the check, never the text.

What is written for them (ballot/check_local.py states the conventions):
  - county offices (level county): commissioners and the partisan row offices as each list prints them (district
    attorney, assessor, clerk, recorder, treasurer, public administrator), and constables, who are elected by the
    voters of a township (NRS 258.010) and are filed under their county with the township as the district;
  - judges (level court, county_ids filled in): district court judges by department, justices of the peace by township
    (NRS 4.020) and municipal court judges under their city;
  - city offices (level city) under the Census Bureau's 2020 place code and name ("Henderson city"), school district
    trustees (level school; a county school district has its county's name, NRS 386.010 and 386.030), and water,
    general improvement and other districts and a town advisory board (level other);
  - partisan or not is what each list prints (Clark's section headings and its NP code; Washoe's Office Type and its
    empty Party cell), and the party words are the ones the state rows use;
  - sl_places: Nevada's seventeen counties (names from the Census county file) and every place the races use;
    sl_gaps: a row for each of the fifteen counties whose list is not read yet, and for any contest this loader could
    not place; sl_notes: local_calendar (from the statutes, read on the Legislature's site on 2026-10-01) and
    local_coverage.
The federal and state offices on the two lists (Congress, the statewide offices, the Legislature, the Supreme Court
and the Board of Regents) and Washoe's ballot questions are counted and left out: the state rows above still wait for
the Secretary of State's own list, and are written exactly as before. Each list is read twice, by two routes (Clark's
by its ruled lines and again line by line; Washoe's through openpyxl and again from the workbook's own XML), and the
two readings must agree; every name read is placed in exactly one contest. Neither list prints totals of its own.
Checked by hand on 2026-10-01: each county's loaded names against a second document of the same office (Clark's Legal
Notice of General Election, Washoe's sample ballot booklet), all 208 found.

When a county's file cannot be fetched and no copy is kept, that county gets a gap and the rest loads. Clark and
Washoe counties do not touch, so no district is on both lists; when a neighbouring county is added, a district on two
counties' lists must become one place and one race (a key without a county code, every county in county_ids).

    python -m ballot.state_local_nv <path to a test database> [--cache <folder holding nv/>]
"""

import csv
import datetime as dt
import glob
import hashlib
import html as H
import io
import json
import os
import re
import sqlite3
import sys
import time
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter, defaultdict
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlparse

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from ballot import pdftext  # noqa: E402
from ballot.check_local import EXTRA_SCHEMA, contact_like  # noqa: E402
from ballot.common import CACHE, SUFFIXES, fold, name_parts  # noqa: E402
from ballot.lists import nv as NVL  # noqa: E402  the saved-file conventions: STEMS, KINDS, clean(), official(), party_of(), turned()
from ballot.lists.tx import proper  # noqa: E402
from ballot.match import fits  # noqa: E402
from states import net  # noqa: E402

STATE, FIPS, NAME = "NV", "32", "Nevada"
GENERAL, PRIMARY = "2026-11-03", "2026-06-09"
ROSTER = os.path.join(HERE, "state_nv.sqlite")
COUNTY_ZIP = os.path.join(HERE, "states_cache", "census", "cb_2024_us_county_500k.zip")
COUNTY_URL = "https://www2.census.gov/geo/tiger/GENZ2024/shp/cb_2024_us_county_500k.zip"
LEG_PAGES = {"Senate": "https://www.leg.state.nv.us/App/Legislator/A/Senate/Current",
             "Assembly": "https://www.leg.state.nv.us/App/Legislator/A/Assembly/Current"}
CONST_URL = "https://www.leg.state.nv.us/Const/NvConst.html"
KEPT = "sl_nv_legislature.json"
KEEP_DAYS = 7
SEATS = {"Senate": 21, "Assembly": 42}

SRC_GENERAL, SRC_PRIMARY, SRC_FILINGS = "nv-sos-2026-sl-general-list", "nv-sos-2026-sl-primary-results", "nv-sos-2026-sl-filings"
SRC_SENATE, SRC_ASSEMBLY, SRC_CONST = "nv-leg-2026-senate-members", "nv-leg-2026-assembly-members", "nv-leg-constitution-art5"
SRC_ROSTER, SRC_COUNTY = "nv-openstates-roster", "nv-census-2024-county-codes"

SCHEMA = """
CREATE TABLE IF NOT EXISTS sl_races (race_id TEXT PRIMARY KEY, state TEXT NOT NULL, level TEXT NOT NULL, office_kind TEXT NOT NULL,
  office TEXT NOT NULL, jurisdiction TEXT, jurisdiction_id TEXT, county_ids TEXT, district TEXT, seat TEXT, special INTEGER NOT NULL DEFAULT 0,
  partisan INTEGER NOT NULL, holder_id TEXT, holder_name TEXT, holder_party TEXT, election_date TEXT NOT NULL, note TEXT);
CREATE TABLE IF NOT EXISTS sl_candidates (race_id TEXT NOT NULL, election TEXT NOT NULL, election_date TEXT, name TEXT NOT NULL, party TEXT,
  party_code TEXT, ballot_order INTEGER, incumbent INTEGER NOT NULL DEFAULT 0, write_in INTEGER NOT NULL DEFAULT 0, votes INTEGER, pct REAL,
  outcome TEXT, state_member_id TEXT, source_id TEXT, note TEXT, PRIMARY KEY (race_id, election, name));
CREATE TABLE IF NOT EXISTS sl_sources (source_id TEXT PRIMARY KEY, state TEXT, kind TEXT, agency TEXT, title TEXT, url TEXT, published TEXT,
  fetched TEXT, sha256 TEXT, rows INTEGER, note TEXT);
CREATE TABLE IF NOT EXISTS sl_places (kind TEXT NOT NULL, id TEXT NOT NULL, name TEXT NOT NULL, county_ids TEXT, source_id TEXT,
  PRIMARY KEY (kind, id));
"""

# The statewide offices on the 2026 ballot: (race key, office_kind, office shown, the office as a list writes it, roster officials.office)
STATEWIDE = [
    ("GOV", "governor", "Governor", r"governor", "governor"),
    ("LTG", "lieutenant_governor", "Lieutenant Governor", r"(?:lieutenant|lt\.?) governor", "lt_governor"),
    ("SOS", "secretary_of_state", "Secretary of State", r"secretary of state", "secretary of state"),
    ("AG", "attorney_general", "Attorney General", r"attorney general", "attorney general"),
    ("TREAS", "state_treasurer", "State Treasurer", r"(?:state )?treasurer", None),
    ("CTRL", "state_controller", "State Controller", r"(?:state )?controller", None),
]
STATEWIDE_NOTE = {
    "GOV": "Nevada elects its governor every four years (Nevada Constitution, Article 5, Section 2); the lieutenant governor is "
           "elected separately, under 2026-NV-LTG.",
    "LTG": "Elected on its own, at the same time and for the same term as the governor (Nevada Constitution, Article 5, Section 17).",
}
OTHER_STATEWIDE_NOTE = "Elected at the same time as the governor, for the same four-year term (Nevada Constitution, Article 5, Section 19)."
OFFICE_WORDS = {"governor": "Governor", "lt_governor": "Lieutenant Governor", "secretary of state": "Secretary of State",
                "attorney general": "Attorney General"}
NO_HOLDER_NOTE ="The Open States roster does not carry this office, so no holder is shown."
CONST_TEXT = {
    "Section 2": "The Governor shall be elected by the qualified electors at the time and places of voting for members of the Legislature",
    "Section 17": "A Lieutenant Governor shall be elected at the same time and places and in the same manner as the Governor",
    "Section 19": "A Secretary of State, a Treasurer, a Controller, and an Attorney General, shall be elected at the same time and places",
}
WAITING_NOTE = ("The November candidates are not loaded yet: the Secretary of State's website turns away automated requests, so its "
                "certified list waits to be saved from a browser.")
COURT_NOTE = "Nevada elects its judges without party labels."
PRIMARY_CODE = {"Democratic": "DEM", "Republican": "REP"}
NONPARTISAN = "Nonpartisan office"
CAPS = "Nevada's list prints names in capitals; they are shown here in ordinary capitals."
CAPS_RESULTS = "Nevada's results print names in capitals; they are shown here in ordinary capitals."

FEDERAL = re.compile(r"representative in congress|u\.?\s?s\.?\s+(?:representative|senat)|united states (?:representative|senat)|congressional|"
                     r"president", re.I)
LOCAL_OR_COURT = re.compile(r"\b(?:county|city|township|town|justice of the peace|constable|district (?:court|judge|attorney)|judge|"
                            r"department|family|municipal|regent|board|trustee|school|commission|council|mayor|sheriff|assessor|clerk|"
                            r"recorder|public administrator|question|ward|hospital|library|fire|water|general improvement)\b", re.I)
OFFICEISH = re.compile(r"governor|secretary of state|attorney general|treasurer|controller|senat|assembly|congress|court|justice|judge|"
                       r"regent|board|trustee|commission|council|mayor|sheriff|assessor|clerk|recorder|constable|district|seat|department|"
                       r"public administrator|president|question", re.I)
PARTY_TAIL = re.compile(r"\s*(?:[-–—,:]\s*)?\(?\b(?:DEM|REP|LPN|LIB|IAP|NP|NPP|GRN|Democratic|Republican|Libertarian|"
                        r"Independent American|Nonpartisan|Non-Partisan)\b(?:\s+Party)?\)?\s*$", re.I)
VOTE_FOR = re.compile(r"\s*\(?\bvote for (?:not more than )?(?:\d+|one|two)\)?\s*", re.I)
PARTY_ONLY = re.compile(r"(?:democratic|republican)(?: party)?|dem|rep|nonpartisan|non-partisan", re.I)
NOTC = re.compile(r"^none of these candidates$", re.I)
HEADS = {"name": r"(candidate(?:'s)?|ballot|choice)?\s*name(?: on ballot| as it appears on ballot)?|candidate|choice",
         "office": r"(office|contest|race)(?: name| title| sought)?",
         "district": r"district(?: name| no\.?| number)?|seat|department",
         "jurisdiction": r"(?:office )?jurisdiction|office level|office type",
         "party": r"party(?: affiliation| preference)?", "status": r"(filing |candidate )?status",
         "order": r"ballot order|order|position", "votes": r"(total )?votes(?: cast)?|vote count"}
STATEWIDE_JURISDICTION = re.compile(r"^(?:statewide|state|state of nevada|nevada|legislative|state legislature|judicial|supreme court|"
                                    r"court of appeals)?$", re.I)

# ---- county and local races: the counties whose own lists are read, and the words written about them
LOCAL_DIR, KEPT_LOCAL, LOCAL_KEEP_DAYS = "local", "kept.json", 7
PLACE_URL = "https://www2.census.gov/geo/docs/reference/codes2020/place/st32_nv_place2020.txt"
PLACE_FILE = "census_st32_nv_place2020.txt"
PLACE_HEAD = ["STATE", "STATEFP", "PLACEFP", "PLACENS", "PLACENAME", "TYPE", "CLASSFP", "FUNCSTAT", "COUNTIES"]
CLARK_PAGE = "https://www.clarkcountynv.gov/government/departments/elections/index.php"
CLARK_URL = ("https://www.clarkcountynv.gov/adobe/assets/urn:aaid:aem:1db09a57-4c1b-47e1-94d9-e7782c442571/original/as/"
             "officesup-26g.pdf")
CLARK_FILE = "clark_officesup-26g.pdf"
WASHOE_PAGE = "https://www.washoecounty.gov/voters/information/index.php"
WASHOE_URL = "https://www.washoecounty.gov/voters/files/Washoe%20County%202026%20General%20Election%20Contests.xlsx"
WASHOE_FILE = "washoe_2026_general_election_contests.xlsx"
NRS_URL = "https://www.leg.state.nv.us/nrs/NRS-293.html"
SRC_CLARK, SRC_WASHOE, SRC_PLACES = "nv-clark-2026-general-contests", "nv-washoe-2026-general-contests", "nv-census-2020-places"
LOCAL_PARTIES = {"DEM", "REP", "LPN", "IAP", "NPP"}                     # the codes both lists print for a partisan office
# Clark County's section headings: (whose offices they are, whether the list prints them with parties)
CLARK_SECTIONS = {"FEDERAL PARTISAN": ("state", 1), "STATE PARTISAN": ("state", 1), "COUNTY PARTISAN": ("county", 1),
                  "STATE AND DISTRICT NONPARTISAN": ("district", 0), "COUNTY NONPARTISAN": ("county", 0),
                  "CITY (ALL NONPARTISAN)": ("city", 0), "TOWNSHIP NONPARTISAN": ("township", 0)}
CLARK_HEAD = ("CONTEST", "CANDIDATE NAME ON BALLOT Last Name", "First, Middle, and/or Nickname", "PARTY", "TERM (Years)")
WASHOE_HEAD = ["Name on Ballot", "Party", "Office Sought", "Office Type", "Term Length", "Vote For"]
WASHOE_STATE_TYPES = {"Federal District Partisan", "Statewide Partisan", "District Partisan", "Statewide Nonpartisan"}
# a county's own offices, as either list words them: (office_kind, the office in plain words)
COUNTY_ROW = {"district attorney": ("district_attorney", "District Attorney"), "county district attorney": ("district_attorney", "District Attorney"),
              "county assessor": ("county_assessor", "County Assessor"), "county clerk": ("county_clerk", "County Clerk"),
              "county recorder": ("county_recorder", "County Recorder"), "county treasurer": ("county_treasurer", "County Treasurer"),
              "public administrator": ("public_administrator", "Public Administrator"),
              "county public administrator": ("public_administrator", "Public Administrator"),
              "sheriff": ("sheriff", "Sheriff"), "county sheriff": ("sheriff", "Sheriff")}
COUNCIL_WORDS = {"council member": "Council Member", "city council": "City Council Member", "councilman": "Councilman",
                 "councilwoman": "Councilwoman", "councilmember": "Councilmember"}
CITY_ROW = {"city attorney": ("city_attorney", "City Attorney"), "city clerk": ("city_clerk", "City Clerk"),
            "city treasurer": ("city_treasurer", "City Treasurer")}
STATE_ON_COUNTY_LIST = re.compile(r"(?:Nevada )?Board of Regents\b|State Board of Education\b|Justice of the Supreme Court\b|Supreme Court\b|"
                                  r"Judge of the Court of Appeals\b|Court of Appeals\b")
GID_NOTE = ("The county's list prints only the district's name for this contest; the office is trustee of the district's board "
            "(NRS 318.095).")
NAME_ONLY_NOTE = "The county's list prints only the district's name for this contest."
LOCAL_CAPS = "The county's list prints this name in capitals; it is shown here in ordinary capitals."
LOCAL_WHAT = "county, township, city, school and district races"
CALENDAR = ("Nevada holds its general election in November of each even-numbered year, and county, township, city, school district and "
            "general improvement district offices are filled at it, so those whose terms are ending are on the November 3, 2026 ballot. A "
            "nonpartisan office is missing from it when the June 9 primary settled it: a candidate with no opponent, or one who won a majority "
            "there, was declared elected (a judge of the district court in that position goes on to November alone). Irrigation districts hold "
            "their own elections in April.")
CALENDAR_SOURCE = "Nevada Revised Statutes 293.12755, 293.260, 293C.140, 318.095 and 539.115 (Nevada Legislature)"

# the last check on every text stored (the site builder's own patterns)
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[A-Za-z]{2,}")
PHONE = re.compile(r"\(?\b\d{3}\)?[\s.-]\d{3}[\s.-]\d{4}\b")
WEB = re.compile(r"https?://|\bwww\.|\b[\w-]+\.(?:com|org|net|us|gov|edu|info|biz)\b", re.I)
STREET = re.compile(r"\b\d{1,6}\s+(?:[NSEW]\.?\s+)?[A-Za-z0-9.' -]{1,40}?\s(?:St|Street|Ave|Avenue|Rd|Road|Blvd|Boulevard|Dr|Drive|Ln|Lane|"
                    r"Way|Ct|Court|Cir|Circle|Pkwy|Parkway|Hwy|Highway|Trl|Trail|Pl|Place|Ter|Terrace)\b\.?", re.I)
ZIP = re.compile(r"\b[A-Z]{2}\s+\d{5}\b|\b\d{5}-\d{4}\b")


def fail(msg):
    raise SystemExit(f"Nevada (state races): {msg}")


def sha_bytes(b):
    return hashlib.sha256(b).hexdigest()


def sha_file(path):
    return sha_bytes(open(path, "rb").read()) if path and os.path.exists(path) else ""


def mdate(path):
    return dt.date.fromtimestamp(os.path.getmtime(path)).isoformat() if path and os.path.exists(path) else ""


# ------------------------------------------------------------------------------------------ the Legislature's pages

class Blocked(Exception):
    pass


def fetch(url, say=print):
    """One request, asked at most twice more on a refusal or a server error; a challenge page is never answered."""
    for attempt in range(3):
        try:
            raw = net.get(url, timeout=90)
        except HTTPError as e:
            if e.code not in (403, 429, 500, 502, 503, 504) or attempt == 2:
                raise
            say(f"      {url}: HTTP {e.code}; asking again in {10 * (attempt + 1)} s")
            time.sleep(10 * (attempt + 1))
            continue
        if re.search(rb"_Incapsula_Resource|Request unsuccessful|cf-chl|challenge-platform|Just a moment", raw[:8000]):
            raise Blocked(f"{url} answered with a bot check")
        return raw
    raise Blocked(url)


def leg_members(raw, chamber):
    """{district: {name, party, counties, term_ends}} from the Legislature's list of current members. Only each member row's
    name, party, district and county cells (their data-order values) and the "Term Ends" field are read."""
    page = raw.decode("utf-8", "replace")
    heads = [re.sub(r"\s+", " ", H.unescape(re.sub(r"<[^>]+>", "", h))).strip() for h in re.findall(r"<th\b[^>]*>(.*?)</th>", page, re.S | re.I)]
    want = ["Legislator", "Party", "District", "County"]
    if [h for h in heads if h] [:4] != want:
        fail(f"the Legislature's {chamber} page no longer heads its columns {want} (it heads them {[h for h in heads if h][:6]})")
    out, cur = {}, None
    for m in re.finditer(r"<tr\b[^>]*>(.*?)</tr>", page, re.S | re.I):
        body = m.group(1)
        orders = [H.unescape(o).strip() for o in re.findall(r'<td\b[^>]*\bdata-order="([^"]*)"', body)]
        at = [i for i, o in enumerate(orders) if re.fullmatch(r"No\. \d+", o)]
        if at:
            i = at[0]
            if i < 2 or i + 1 >= len(orders):
                fail(f"a member row on the Legislature's {chamber} page does not line up with its headings")
            d = int(orders[i].split()[1])
            if d in out:
                fail(f"the Legislature's {chamber} page lists district {d} twice")
            counties = [re.sub(r"\s*\(Part\)\s*$", "", c).strip() for c in orders[i + 1].split(",") if c.strip()]
            out[d] = {"district": str(d), "name": orders[i - 2], "party": orders[i - 1], "counties": counties,
                      "part": [c for c in orders[i + 1].split(",") if "(Part)" in c and c.strip()] != [], "term_ends": None}
            cur = d
            continue
        v = re.search(r'Term Ends:</span>(?:\s|&nbsp;)*<span class="field">\s*(\d{4})', body)
        if v and cur is not None:
            out[cur]["term_ends"] = v.group(1)
    if sorted(out) != list(range(1, SEATS[chamber] + 1)):
        fail(f"the Legislature's {chamber} page lists districts {sorted(out)}, not 1 to {SEATS[chamber]}")
    missing = [d for d, r in out.items() if not r["term_ends"]]
    if missing:
        fail(f"the Legislature's {chamber} page gives no Term Ends year for districts {missing}")
    for r in out.values():
        del r["part"]
    return {str(d): r for d, r in sorted(out.items())}


def constitution(raw):
    text = re.sub(r"\s+", " ", H.unescape(re.sub(r"<[^>]+>", " ", raw.decode("windows-1252", "replace"))))
    return {sec: (words in text) for sec, words in CONST_TEXT.items()}


def legislature(folder, say=print, refresh=False):
    """The Legislature's member lists and the Constitution's check, kept a week in ballot_cache/nv/sl_nv_legislature.json."""
    path = os.path.join(folder, KEPT)
    if os.path.exists(path) and not refresh and time.time() - os.path.getmtime(path) < KEEP_DAYS * 86400:
        return json.load(open(path, encoding="utf-8")), "kept copy"
    try:
        data = {"fetched": dt.date.today().isoformat(), "chambers": {}}
        for chamber, url in LEG_PAGES.items():
            raw = fetch(url, say)
            data["chambers"][chamber] = {"url": url, "sha256": sha_bytes(raw), "members": leg_members(raw, chamber)}
        raw = fetch(CONST_URL, say)
        data["constitution"] = {"url": CONST_URL, "sha256": sha_bytes(raw), "found": constitution(raw)}
    except (HTTPError, URLError, OSError, Blocked) as e:
        if os.path.exists(path):
            say(f"      the Legislature's pages could not be read ({type(e).__name__}); using the copy kept on {mdate(path)}")
            return json.load(open(path, encoding="utf-8")), "kept copy (the site could not be reached)"
        fail(f"the Legislature's pages could not be read ({type(e).__name__}: {e}) and no copy is kept")
    os.makedirs(folder, exist_ok=True)
    json.dump(data, open(path, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    return data, "read afresh"


# ------------------------------------------------------------------------------------------------- saved lists

def saved(folder, kind):
    """(readable files, problems) a browser saved for one kind: the federal loader's own name and any copy named with a suffix."""
    stem = NVL.STEMS[kind]
    files, problems = [], []
    for p in sorted(set(glob.glob(os.path.join(folder, stem + ".*")) + glob.glob(os.path.join(folder, stem + "_*.*")))):
        low = p.lower()
        if low.endswith(".pdf"):
            problems.append(f"{os.path.basename(p)} is a PDF; this loader reads a list as a table (HTML, CSV or a workbook)")
            continue
        if not low.endswith(NVL.KINDS):
            continue
        head = open(p, "rb").read(4096)
        if low.endswith(".xlsx") and not head.startswith(b"PK"):
            problems.append(f"{os.path.basename(p)} is not a workbook")
            continue
        if b"_Incapsula_Resource" in head or b"Request unsuccessful" in head:
            problems.append(f"{os.path.basename(p)} is the bot wall's page, not the file; it needs saving again from a browser that shows the page")
            continue
        files.append(p)
    return files, problems


def blank(fmt, cell):
    if fmt == "html":
        return not re.sub(r"<[^>]+>|&nbsp;|&#160;|\s", "", cell)
    return not str(cell).strip()


def text_of(fmt, cell):
    return NVL.clean(cell) if fmt == "html" else re.sub(r"\s+", " ", str(cell)).strip()


def raw_rows(path):
    """None where a new table begins, then (format, raw cells, whether it is a heading row) for each row. No cell is turned
    into text here: a caller picks the cells it may read."""
    low = path.lower()
    if low.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            yield None
            for k, r in enumerate(ws.iter_rows(values_only=True)):
                yield "plain", ["" if c is None else c for c in r], k < 10
        wb.close()
        return
    text = open(path, encoding="utf-8-sig", errors="replace").read()
    if text.lstrip().startswith("<") and not low.endswith(".csv"):
        for table in re.findall(r"<table\b(?:(?!<table\b).)*?</table>", text, re.S | re.I):
            yield None
            for k, tr in enumerate(re.findall(r"<tr\b[^>]*>(.*?)</tr>", table, re.S | re.I)):
                yield "html", re.findall(r"<t[hd]\b[^>]*>(.*?)</t[hd]>", tr, re.S | re.I), k < 3 or bool(re.search(r"<th\b", tr, re.I))
        return
    lines = text.splitlines()
    if sum(1 for ln in lines[:60] if ln.lstrip().startswith("|")) >= 2:        # a Markdown table
        k = 0
        yield None
        for ln in lines:
            s = ln.strip()
            if not s.startswith("|"):
                if s:
                    k = 0
                    yield None
                continue
            if re.fullmatch(r"\|?[\s:|-]+\|?", s):
                continue
            yield "plain", s.strip("|").split("|"), k < 3
            k += 1
        return
    delim = ";" if text[:4000].count(";") > text[:4000].count(",") else ","
    yield None
    for k, r in enumerate(csv.reader(io.StringIO(text), delimiter=delim)):
        yield "plain", r, k < 10


def head_key(text):
    t = re.sub(r"[:*]+$", "", text or "").strip().lower()
    return next((k for k, pat in HEADS.items() if re.fullmatch(pat, t)), None)


def allowed_rows(path):
    """([dicts of the allowed columns only, with the office heading above each row], [allowed headings found per table], short rows).
    A table's heading row is a heading-like row naming a candidate column and at least one other allowed column."""
    out, seen, cols, group, short = [], [], None, "", 0
    for row in raw_rows(path):
        if row is None:
            cols, group = None, ""
            continue
        fmt, cells, heading_like = row
        if cols is None:
            if not heading_like:
                continue
            heads = [text_of(fmt, c) for c in cells]                           # headings, not people
            keys = [head_key(h) for h in heads]
            if "name" in keys and sum(1 for k in keys if k) >= 2:
                cols = {k: i for i, k in reversed(list(enumerate(keys))) if k}
                seen.append([h for h, k in zip(heads, keys) if k])
            continue
        full = [i for i, c in enumerate(cells) if not blank(fmt, c)]
        if not full:
            continue
        if len(full) == 1 and sum(1 for i in cols.values() if i < len(cells) and not blank(fmt, cells[i])) < 2:
            g = text_of(fmt, cells[full[0]])
            if len(g) < 160 and OFFICEISH.search(g) and not re.search(r"\d{3}.*\d{4}|@", g):
                group = g                                                   # an office heading across the table; anything else is left unread
            continue
        if len(cells) <= max(cols.values()):
            short += 1
            continue
        out.append({**{k: text_of(fmt, cells[i]) for k, i in cols.items()}, "_group": group})
    return out, seen, short


# ------------------------------------------------------------------------------------------------ offices

def office_text(t):
    t = VOTE_FOR.sub(" ", re.sub(r"\s+", " ", t or "")).strip().strip(":").strip()
    prev = None
    while prev != t:
        prev, t = t, PARTY_TAIL.sub("", t).strip().rstrip(",-:").strip()
    return t


def district_no(text, column=""):
    m = re.search(r"district\s*(?:no\.?\s*)?(\d{1,2})\b|\b(\d{1,2})(?:st|nd|rd|th)\s+district", text or "", re.I)
    if m:
        return str(int(m.group(1) or m.group(2)))
    m = re.fullmatch(r"(?:(?:state\s+)?(?:senate|assembly)\s*)?(?:district\s*|dist\.?\s*)?(?:no\.?\s*)?0*(\d{1,2})", (column or "").strip(), re.I)
    return str(int(m.group(1))) if m else None


def classify(office, district="", group="", jurisdiction=""):
    """('federal', None) | ('other', title) | ('statewide', key) | ('senate', d) | ('assembly', d) | ('supreme', seat) |
    ('appeals', seat) for an office as a list or a results page writes it."""
    t = office_text(office or group)
    low = t.lower()
    if not low:
        return "other", ""
    if FEDERAL.search(low):
        return "federal", None
    if jurisdiction and not STATEWIDE_JURISDICTION.match(jurisdiction.strip()) and not re.search(r"senate|assembly|district", jurisdiction, re.I):
        return "other", f"{t} ({jurisdiction})"
    if "supreme court" in low:
        m = re.search(r"seat\s*([A-G])\b", f"{t} {district}", re.I) or re.fullmatch(r"\s*([A-G])\s*", district or "", re.I)
        return ("supreme", m.group(1).upper()) if m else ("other", t)
    if "court of appeals" in low:
        m = re.search(r"(?:seat|department|dept\.?)\s*([A-Z0-9]{1,2})\b", f"{t} {district}", re.I) or re.fullmatch(r"\s*([A-Z0-9])\s*", district or "")
        return ("appeals", m.group(1).upper()) if m else ("other", t)
    if LOCAL_OR_COURT.search(low):
        return "other", t
    for key, _k, _s, pat, _r in STATEWIDE:
        if re.fullmatch(pat, low):
            return "statewide", key
    if re.search(r"\bsenat", low):
        d = district_no(t, district)
        if d is None or not 1 <= int(d) <= SEATS["Senate"]:
            fail(f"a State Senate office whose district is not read ({t!r}, district {district!r})")
        return "senate", d
    if re.search(r"\bassembly", low):
        d = district_no(t, district)
        if d is None or not 1 <= int(d) <= SEATS["Assembly"]:
            fail(f"an Assembly office whose district is not read ({t!r}, district {district!r})")
        return "assembly", d
    return "other", t


def race_id(c):
    kind, x = c
    return {"statewide": f"2026-{STATE}-{x}", "senate": f"2026-{STATE}-SS{x}", "assembly": f"2026-{STATE}-SH{x}",
            "supreme": f"2026-{STATE}-SC-{x}", "appeals": f"2026-{STATE}-COA-{x}"}.get(kind)


def party_in(text):
    codes = re.findall(r"\b(DEM|REP)\b", text or "")
    if codes:
        return {"DEM": "Democratic", "REP": "Republican"}[codes[-1]]
    for w, p in (("democratic", "Democratic"), ("republican", "Republican")):
        if re.search(rf"\b{w}\b", text or "", re.I):
            return p
    return ""


# ---------------------------------------------------------------------------------------------------- people

def roster(path):
    con = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    legs = [dict(zip(("id", "first", "last", "full", "other", "party", "chamber", "district"), r)) for r in con.execute(
        "SELECT bioguide_id, first_name, last_name, official_full, other_names, party_name, chamber, district FROM legislators WHERE is_current = 1")]
    offs = {r[1]: dict(zip(("id", "office", "full", "first", "last", "party", "term_end"), r)) for r in con.execute(
        "SELECT bioguide_id, office, official_full, first_name, last_name, party_name, term_end FROM officials")}
    con.close()
    for p in legs:
        p["chamber"] = "Assembly" if p["chamber"] == "House" else p["chamber"]
        p["district"] = str(int(p["district"])) if str(p["district"] or "").isdigit() else str(p["district"] or "")
    return legs, offs


def forms(p, extra=()):
    """Every (given names, family name) the record writes this person under."""
    out = [(fold(p.get("first") or "").split(), " ".join(fold(p.get("last") or "").split()))]
    others = [o.strip() for o in (p.get("other") or "").split(";") if o.strip() and "." not in o]      # initials-only forms left out
    for f in [p.get("full") or ""] + others + list(extra):
        if f:
            out.append(name_parts(NVL.turned(f)))
    return [f for f in out if f[1]]


def person_fits(name, p, extra=()):
    cand = name_parts(name)
    return bool(cand[1]) and any(fits(cand, f) for f in forms(p, extra))


def show(raw, people=()):
    """(name to show, whether it was printed in capitals). A name in capitals is put in ordinary capitals; when it spells a
    roster person's name letter for letter, the roster's own capitals (and accents) are used."""
    name = re.sub(r"\s+", " ", NVL.turned(raw)).strip()
    if not (re.search(r"[A-Z]", name) and name == name.upper()):
        return name, False
    fixed = re.sub(r"(['\"(“])([a-z])", lambda m: m.group(1) + m.group(2).upper(), proper(name))
    for p in people:
        for f in (p.get("full"), f"{p.get('first') or ''} {p.get('last') or ''}".strip()):
            if f and fold(f) == fold(fixed):
                return f, True
    return fixed, True


# ------------------------------------------------------------------------------------------------ the lists

def general_list(paths, report):
    """([(race kind tuple, name as printed, party label, order or None)], [(race, name) left off], {office: count} not loaded,
    headings seen, rows read) from the saved November lists."""
    on, off, other, seen_all, read, keys = [], [], Counter(), [], 0, set()
    for path in paths:
        rows, seen, short = allowed_rows(path)
        seen_all += seen
        if short:
            report.append(f"{os.path.basename(path)}: {short} rows shorter than their headings were left unread")
        if not seen:
            fail(f"no table with a candidate column was found in {os.path.basename(path)}")
        for r in rows:
            if not r.get("name"):
                continue
            c = classify(r.get("office", ""), r.get("district", ""), "" if r.get("office") else r["_group"], r.get("jurisdiction", ""))
            if c[0] == "federal":
                continue
            read += 1
            if c[0] == "other":
                other[c[1] or "(office not named)"] += 1
                continue
            status = r.get("status", "")
            if status and NVL.OFF.search(status):
                off.append((c, r["name"]))
                continue
            if status and not NVL.ON.match(status):
                fail(f"a status on the November list that is not read ({status!r}, {race_id(c)})")
            raw = r["name"]
            party = NVL.party_of(r.get("party", ""))
            m = re.search(r"\s*\(([A-Z]{2,3})\)\s*$", raw)
            if m:
                party = party or NVL.party_of(m.group(1))
                raw = raw[:m.start()]
            if c[0] in ("supreme", "appeals"):
                party = NONPARTISAN
            elif not party:
                fail(f"a candidate for {race_id(c)} with no party on the November list ({os.path.basename(path)})")
            k = (race_id(c), fold(NVL.turned(raw)))
            if k in keys:
                continue                                                    # the same row in a second saved copy
            keys.add(k)
            order = r.get("order", "")
            on.append((c, raw, party, int(order) if order.isdigit() else None))
    return on, off, other, seen_all, read


def primary_results(path, report):
    """{(race kind tuple, party): {"names": [(name as printed, votes)], "total": printed or None, "notc": votes or None}}."""
    if not NVL.official(path):
        fail(f"{os.path.basename(path)} does not call itself the official results (or calls itself unofficial); only canvassed figures are stored")
    out = {}

    def add(c, party, label, votes):
        f = out.setdefault((c, party), {"names": [], "total": None, "notc": None})
        if re.match(r"total\b", label.strip(), re.I):
            f["total"] = votes
        elif NOTC.match(label.strip()):
            f["notc"] = (f["notc"] or 0) + votes
        else:
            f["names"].append((label, votes))

    if path.lower().endswith((".csv", ".xlsx")) or not open(path, encoding="utf-8-sig", errors="replace").read(200).lstrip().startswith("<"):
        rows, seen, _short = allowed_rows(path)
        for r in rows:
            c = classify(r.get("office", ""), r.get("district", ""), r["_group"], r.get("jurisdiction", ""))
            if race_id(c) is None or not r.get("name") or not NVL.NUMBER.match(r.get("votes", "")):
                continue
            party = NONPARTISAN if c[0] in ("supreme", "appeals") else (NVL.party_of(r.get("party", "")) or party_in(r.get("office", "") + " " + r["_group"]))
            if not party:
                fail(f"no party read for a contest under {race_id(c)} in {os.path.basename(path)}")
            label = re.sub(r"\s*\([A-Z]{2,3}\)\s*$", "", r["name"])
            add(c, party, label, int(r["votes"].replace(",", "")))
        return out
    page = open(path, encoding="utf-8-sig", errors="replace").read()
    token = re.compile(r"<tr\b[^>]*>(.*?)</tr>|>([^<>]*\S[^<>]*)<", re.S | re.I)
    c, party, cols = None, "", None
    for m in token.finditer(page):
        if m.group(2) is not None:                                          # a heading outside a table
            t = NVL.clean(m.group(2))
            if len(t) < 160 and OFFICEISH.search(t) and not NOTC.match(t):
                k = classify(t)
                c = k if race_id(k) else None                               # any other contest's rows are not ours
                party, cols = (NONPARTISAN if c and c[0] in ("supreme", "appeals") else party_in(t)), None
            elif c and PARTY_ONLY.fullmatch(t):
                party = party_in(t) or party
            continue
        cells = re.findall(r"<t[hd]\b[^>]*>(.*?)</t[hd]>", m.group(1), re.S | re.I)
        texts = [NVL.clean(x) for x in cells]                               # results carry names and votes only
        filled = [x for x in texts if x]
        joined = " ".join(filled)
        if 1 <= len(filled) <= 2 and len(joined) < 160 and OFFICEISH.search(joined) and not any(NVL.NUMBER.match(x) for x in filled):
            k = classify(joined)
            c = k if race_id(k) else None
            party, cols = (NONPARTISAN if c and c[0] in ("supreme", "appeals") else party_in(joined)), None
            continue
        if not c:
            continue
        if len(filled) == 1 and PARTY_ONLY.fullmatch(filled[0]):
            party = party_in(filled[0]) or party
            continue
        keys = [head_key(x) for x in texts]
        if "votes" in keys:
            cols = {k: i for i, k in reversed(list(enumerate(keys))) if k}
            cols.setdefault("name", 0)
            continue
        if not cols or len(texts) <= max(cols.values()):
            continue
        votes = texts[cols["votes"]]
        if not NVL.NUMBER.match(votes):
            continue
        label = texts[cols["name"]]
        p = NVL.party_of(texts[cols["party"]]) if "party" in cols else ""
        suffix = re.search(r"\s*\(([A-Z]{2,3})\)\s*$", label)
        if suffix:
            label = label[:suffix.start()]
        p = NONPARTISAN if c[0] in ("supreme", "appeals") else (p or (NVL.party_of(suffix.group(1)) if suffix else "") or party)
        if not p:
            fail(f"no party read for a contest under {race_id(c)} in {os.path.basename(path)}")
        add(c, p, label, int(votes.replace(",", "")))
    if not out:
        fail(f"no state contest read from {os.path.basename(path)}; read the saved page again")
    return out


def filings(paths, report):
    """{(race kind tuple, party): [name as printed]} of who filed and did not withdraw, from Aurora's list for the cycle."""
    out = defaultdict(list)
    for path in paths:
        rows, seen, _short = allowed_rows(path)
        for r in rows:
            if not r.get("name") or NVL.OFF.search(r.get("status", "")):
                continue
            c = classify(r.get("office", ""), r.get("district", ""), "" if r.get("office") else r["_group"], r.get("jurisdiction", ""))
            if race_id(c) is None:
                continue
            party = NONPARTISAN if c[0] in ("supreme", "appeals") else NVL.party_of(r.get("party", ""))
            if party in PRIMARY_CODE or party == NONPARTISAN:
                nm = r["name"]
                if fold(NVL.turned(nm)) not in {fold(NVL.turned(x)) for x in out[(c, party)]}:
                    out[(c, party)].append(nm)
    return out


# ---------------------------------------------------------------------------------------------------- loading

def census_counties(path):
    import shapefile                                                        # pyshp, in the kit's environment
    z = zipfile.ZipFile(path)
    base = next(n for n in z.namelist() if n.endswith(".dbf"))
    rdr = shapefile.Reader(dbf=io.BytesIO(z.read(base)))
    return {fold(r["NAME"]): (r["GEOID"], r["NAME"]) for r in (x.as_dict() for x in rdr.iterRecords()) if r["STATEFP"] == FIPS}


def guard(races, cands):
    """Stop, without showing it, if any stored text looks like a contact detail."""
    for r in races.values():
        for f in ("holder_name", "note", "office", "jurisdiction"):
            v = r.get(f) or ""
            if any(p.search(v) for p in (EMAIL, PHONE, WEB, STREET, ZIP)):
                fail(f"the {f} of {r['race_id']} looks like it holds a contact detail; nothing was written")
    for c in cands:
        for i, f in ((3, "name"), (4, "party"), (14, "note")):
            v = c[i] or ""
            if any(p.search(v) for p in (EMAIL, PHONE, WEB, STREET, ZIP)):
                fail(f"a candidate {f} in {c[0]} ({c[1]}) looks like it holds a contact detail; nothing was written")


# ============================================================================================ county and local races

class Unplaced(Exception):
    """A contest on a county's list that this loader does not know how to file: it becomes a gap, never a guess."""


def lfail(msg):
    raise SystemExit(f"Nevada (county and local races): {msg}")


def slug(text):
    t = unicodedata.normalize("NFKD", text or "")
    return re.sub(r"[^a-z0-9]+", "-", "".join(c for c in t if not unicodedata.combining(c)).lower()).strip("-")


def squeeze(text):
    return re.sub(r"\s+", " ", str(text or "")).strip()


def plain_caps(text):
    """A title printed in capitals, in ordinary capitals: VERDI TV DISTRICT -> Verdi TV District; DISTRICT F AT-LARGE ->
    District F At-Large. A single letter, a Roman numeral and TV stay as they are."""
    out = []
    for i, w in enumerate(squeeze(text).split(" ")):
        core = w.strip(",.")
        if len(core) <= 1 or core in ("TV", "II", "III", "IV", "VI", "VII"):
            out.append(w)
        elif i and core.lower() in ("of", "the", "and", "for"):
            out.append(w.lower())
        else:
            out.append("-".join(p[:1].upper() + p[1:].lower() for p in w.split("-")))
    return " ".join(out)


def census_county_names(path):
    """{GEOID: the county's name as the Census Bureau writes it} for Nevada: 'Clark County', and 'Carson City' as it is."""
    import shapefile
    z = zipfile.ZipFile(path)
    base = next(n for n in z.namelist() if n.endswith(".dbf"))
    rdr = shapefile.Reader(dbf=io.BytesIO(z.read(base)))
    return {r["GEOID"]: r["NAMELSAD"] for r in (x.as_dict() for x in rdr.iterRecords()) if r["STATEFP"] == FIPS}


def census_places(folder, by_full, say):
    """(path, {folded name without its kind word: [(place code, 'Henderson city', [county GEOIDs])]}, rows read) for
    Nevada's incorporated places, from the Census Bureau's 2020 place codes file (places only, no people; kept whole and
    asked for once)."""
    path = os.path.join(folder, PLACE_FILE)
    try:
        net.download(PLACE_URL, path, max_age_days=3650, tries=3, say=say)
    except (HTTPError, URLError, OSError) as e:
        lfail(f"the Census Bureau's place codes file could not be fetched ({type(e).__name__}); nothing was changed. Run again.")
    out, n = defaultdict(list), 0
    with open(path, encoding="utf-8", errors="replace") as fh:
        if fh.readline().rstrip("\r\n").split("|") != PLACE_HEAD:
            lfail(f"{PLACE_FILE}: the header is not the one this loader was checked against")
        for ln, line in enumerate(fh, start=2):
            if not line.strip():
                continue
            f = line.rstrip("\r\n").split("|")
            if len(f) != len(PLACE_HEAD) or f[1] != FIPS or not re.fullmatch(r"\d{5}", f[2]):
                lfail(f"{PLACE_FILE}: line {ln} does not fit the header")
            n += 1
            if f[5] != "INCORPORATED PLACE":
                continue
            cids = [by_full.get(fold(c)) for c in f[8].split("~~~")]
            if not all(cids):
                lfail(f"{PLACE_FILE}: line {ln} names a county the Census county file does not have")
            out[fold(re.sub(r"\s+(?:city|town)$", "", f[4]))].append((f[2], f[4], sorted(cids)))
    return path, out, n


def kept(folder, name, url, magic, say, find=None, days=LOCAL_KEEP_DAYS):
    """(path, the address it came from, the day it was fetched, how) for one county's file; path None when it could not
    be had. A copy kept for less than `days` is used as it is, so a second run asks for nothing. Otherwise one request
    (asked at most twice more on a refusal or a server error; a bot check is never answered), and the answer is kept
    only if it is the kind of file expected. When the address has moved (404), `find` reads the county's own page for
    the new one. If the county cannot be reached, an older copy is used when there is one."""
    path, side_path = os.path.join(folder, name), os.path.join(folder, KEPT_LOCAL)
    side = {}
    if os.path.exists(side_path):
        try:
            side = json.load(open(side_path, encoding="utf-8"))
        except ValueError:
            side = {}
    info = side.get(name) or {}
    have = os.path.exists(path) and os.path.getsize(path) > 0
    if have and time.time() - os.path.getmtime(path) < days * 86400:
        return path, info.get("url") or url, info.get("fetched") or mdate(path), "kept copy"
    used = url
    try:
        try:
            raw = fetch(url, say)
        except HTTPError as e:
            if e.code != 404 or find is None:
                raise
            time.sleep(1.5)
            used = find(say)
            time.sleep(1.5)
            raw = fetch(used, say)
        time.sleep(1.5)
        if not raw.startswith(magic):
            raise Blocked("the answer is not the file")
    except (HTTPError, URLError, OSError, Blocked) as e:
        if have:
            say(f"      {name} could not be read afresh ({type(e).__name__}); using the copy kept on {mdate(path)}")
            return path, info.get("url") or url, info.get("fetched") or mdate(path), "kept copy (the county's site could not be reached)"
        say(f"      {name} could not be fetched ({type(e).__name__}), and no copy is kept")
        return None, url, "", f"not fetched ({type(e).__name__})"
    os.makedirs(folder, exist_ok=True)
    with open(path + ".part", "wb") as fh:
        fh.write(raw)
    os.replace(path + ".part", path)
    side[name] = {"url": used, "fetched": dt.date.today().isoformat(), "sha256": sha_bytes(raw), "bytes": len(raw)}
    json.dump(side, open(side_path, "w", encoding="utf-8"), indent=1)
    return path, used, side[name]["fetched"], "read afresh"


def clark_link(say):
    """Where Clark County's elections page links the list today (the county's own host only), for when the address
    kept here has moved. Only the page's links are looked at, and only one of them is returned."""
    page = fetch(CLARK_PAGE, say).decode("utf-8", "replace")
    for m in re.finditer(r'href\s*=\s*["\']([^"\']+)["\']', page, re.I):
        full = urljoin(CLARK_PAGE, H.unescape(m.group(1)).strip()).split("?")[0]
        u = urlparse(full)
        if (u.hostname or "").lower() == "www.clarkcountynv.gov" and u.path.lower().endswith("/officesup-26g.pdf"):
            return full
    raise Blocked("the elections page no longer links the list")


# ---------------------------------------------------------------------------------- what an office on a county list is

def district_board(name, area=None, named_only=False):
    """A district's board contest, filed by what the district's own name says it is."""
    low = name.lower()
    out = dict(level="other", kind="special_district_board", office="Board Member", place=("special", name), district=area,
               note=NAME_ONLY_NOTE if named_only else None)
    if "general improvement district" in low:
        out.update(kind="improvement_board")
        if named_only:
            out.update(office="General Improvement District Trustee", note=GID_NOTE)
    elif "hospital" in low:
        out.update(level="hospital", kind="hospital_board", place=("hospital", name))
    elif "water" in low:
        out.update(kind="water_board")
    elif "fire" in low:
        out.update(kind="fire_board")
    elif "library" in low:
        out.update(kind="library_board")
    return out


def city_office(bare, rest):
    """A city's contest, from the words that follow the city's name on the list."""
    r = squeeze(rest).strip(" ,")
    low = r.lower()
    if low in ("mayor", "city mayor"):
        return dict(level="city", kind="mayor", office="Mayor", place=("city", bare))
    m = re.fullmatch(r"(council member|city council|councilman|councilwoman|councilmember),? ?(ward \w+|district \w+|seat \w+|at[ -]large)?", low)
    if m:
        part = r[m.start(2):m.end(2)] if m.group(2) else None
        out = dict(level="city", kind="council", office=COUNCIL_WORDS[m.group(1)], place=("city", bare))
        if part and re.match(r"(?:ward|district) ", part, re.I):
            out["district"] = part
        elif part:
            out["seat"] = "At Large" if re.fullmatch(r"at[ -]large", part, re.I) else part
        return out
    m = re.fullmatch(r"municipal court judge,? ?(department \w+)?", low)
    if m:
        return dict(level="court", kind="municipal_court", office="Municipal Court Judge", place=("city", bare),
                    seat=r[m.start(1):m.end(1)] if m.group(1) else None)
    if low in CITY_ROW:
        k, o = CITY_ROW[low]
        return dict(level="city", kind=k, office=o, place=("city", bare))
    raise Unplaced()


def township_office(text):
    """A constable or a justice of the peace, each elected by the voters of one township; None for anything else. A
    Nevada township has no government of its own, so both are filed under the county with the township as the district."""
    m = re.fullmatch(r"Constable,? (.+? Township)", text)
    if m:
        return dict(level="county", kind="constable", office="Constable", place=("county",), district=m.group(1),
                    note=f"Elected by the voters of {m.group(1)} only (NRS 258.010).")
    m = re.fullmatch(r"Justice of the Peace,? (.+? Township)(?:,? (Department \w+))?", text)
    if m:
        return dict(level="court", kind="justice_of_the_peace", office="Justice of the Peace", place=("county",), district=m.group(1),
                    seat=m.group(2), note=f"Elected by the voters of {m.group(1)} only (NRS 4.020).")
    return None


def clark_office(section, text, city, cname):
    """What one contest of Clark County's list is: None for a federal or state office (left to the state part of this
    loader), a dict for a county, township, city, school, district or local court contest; Unplaced for any other."""
    scope = CLARK_SECTIONS[section][0]
    if scope == "state":
        return None
    if scope == "district":
        m = re.fullmatch(r"District Court Judge ((?:Family Division )?Department \w+)", text)
        if m:
            return dict(level="court", kind="district_court", office="District Court Judge", place=("county",), seat=m.group(1))
        if STATE_ON_COUNTY_LIST.match(text):
            return None
        raise Unplaced()
    if scope == "county":
        m = re.fullmatch(r"County Commission(?:er)? District (\w+)", text)
        if m:
            return dict(level="county", kind="county_commissioner", office="County Commissioner", place=("county",), district=f"District {m.group(1)}")
        if text.lower() in COUNTY_ROW:
            k, o = COUNTY_ROW[text.lower()]
            return dict(level="county", kind=k, office=o, place=("county",))
        m = re.fullmatch(rf"Trustee ({re.escape(cname)} School District) (\w+)", text)
        if m:
            return dict(level="school", kind="school_board", office="School District Trustee", place=("school", m.group(1)),
                        district=f"District {m.group(2)}")
        m = re.fullmatch(r"Board Member (.+? District)(?: (?!No\b|#|\d)(.+))?", text)
        if m:
            return district_board(m.group(1), m.group(2))
        m = re.fullmatch(r"(.+) Town Advisory Board", text)
        if m:
            return dict(level="other", kind="town_advisory_board", office="Town Advisory Board Member",
                        place=("special", f"{m.group(1)} (unincorporated town)"))
        raise Unplaced()
    if scope == "city":
        if not city or not text.startswith(city + " "):
            raise Unplaced()
        return city_office(city[len("City of "):], text[len(city) + 1:])
    t = township_office(text)
    if t is None:
        raise Unplaced()
    return t


def washoe_office(office, otype, cname):
    """What one contest of Washoe County's workbook is, from its Office Sought and Office Type (as clark_office)."""
    text = plain_caps(office)                                               # the workbook prints contests in capitals
    if otype in WASHOE_STATE_TYPES:
        return None
    if otype == "County Partisan":
        m = re.fullmatch(r"County Commissioner, District (\w+)", text)
        if m:
            return dict(level="county", kind="county_commissioner", office="County Commissioner", place=("county",), district=f"District {m.group(1)}")
        if text.lower() in COUNTY_ROW:
            k, o = COUNTY_ROW[text.lower()]
            return dict(level="county", kind=k, office=o, place=("county",))
        raise Unplaced()
    if otype in ("District Nonpartisan", "County Nonpartisan", "Township Nonpartisan"):
        m = re.fullmatch(r"District Court Judge, (Department \w+(?:, Family Court)?)", text)
        if m:
            return dict(level="court", kind="district_court", office="District Court Judge", place=("county",), seat=m.group(1))
        if STATE_ON_COUNTY_LIST.match(text):
            return None
        m = re.fullmatch(rf"{re.escape(cname)} School Board, (District \w+(?: At-Large)?)", text)
        if m:
            return dict(level="school", kind="school_board", office="School District Trustee", place=("school", f"{cname} School District"),
                        district=m.group(1))
        if text.lower() in COUNTY_ROW:
            k, o = COUNTY_ROW[text.lower()]
            return dict(level="county", kind=k, office=o, place=("county",))
        t = township_office(text)
        if t:
            return t
        if re.fullmatch(r".+ District", text):
            return district_board(text, None, named_only=True)
        raise Unplaced()
    m = re.fullmatch(r"City of (.+)", otype)
    if m and text.lower().startswith(m.group(1).lower() + " "):
        return city_office(m.group(1), text[len(m.group(1)) + 1:])
    raise Unplaced()


# ---------------------------------------------------------------------------------------- Clark County's list (a PDF)

def page_rects(pdf, page):
    """Every rectangle a page draws, in page space: (x, y, width, height). The list's table is ruled with thin ones."""
    contents = pdf.get(page.get("Contents"))
    parts = contents if isinstance(contents, list) else [page.get("Contents")]
    data = b"\n".join(pdf.stream(p) or b"" for p in parts if isinstance(p, pdftext.Ref))
    ctm, saved, out = [1, 0, 0, 1, 0, 0], [], []
    for op, a in pdftext._ops(data):
        if op == "q":
            saved.append(ctm[:])
        elif op == "Q":
            ctm = saved.pop() if saved else [1, 0, 0, 1, 0, 0]
        elif op == "cm" and len(a) == 6:
            ctm = pdftext._mul([float(x) for x in a], ctm)
        elif op == "re" and len(a) == 4:
            x, y, w, h = (float(v) for v in a)
            x0, y0 = ctm[0] * x + ctm[2] * y + ctm[4], ctm[1] * x + ctm[3] * y + ctm[5]
            x1, y1 = ctm[0] * (x + w) + ctm[2] * (y + h) + ctm[4], ctm[1] * (x + w) + ctm[3] * (y + h) + ctm[5]
            out.append((min(x0, x1), min(y0, y1), abs(x1 - x0), abs(y1 - y0)))
    return out


def apart(values, gap=1.5):
    """Sorted values, a rule drawn twice counted once."""
    out = []
    for v in sorted(values):
        if not out or v - out[-1] > gap:
            out.append(v)
    return out


def cell_lines(runs):
    """The text runs of one cell as its printed lines, top to bottom."""
    lines = []
    for r in sorted(runs, key=lambda r: (-r[1], r[0])):
        if lines and abs(lines[-1][0] - r[1]) <= 2.0:
            lines[-1][1].append(r)
        else:
            lines.append([r[1], [r]])
    return [t for t in (pdftext.join(rs) for _y, rs in lines) if t]


def clark_list(path):
    """Clark County's table, read by its own ruled lines. A rule across the whole table ends a block (the headings, a
    section heading, a city's heading or one contest); inside a block a rule across the name and party columns ends a
    candidate. Returns the blocks in order, each {page, k, label: the CONTEST cell's lines, term, cells: [(last name,
    first names, party, lines)]}, and the words printed outside the table on page 1. Read a second time line by line,
    without the rules: every printed line with a party code must be one cell of the first reading."""
    name = os.path.basename(path)
    pdf = pdftext.PDF(open(path, "rb").read())
    blocks, by_line, outside, foot = [], Counter(), [], []
    for n, (page, res) in enumerate(pdf.pages(), 1):
        runs = pdftext.page_runs(pdf, page, res)
        rects = page_rects(pdf, page)
        edges = apart(x for x, _y, w, h in rects if w <= 2.0 and h > 20)
        if len(edges) != 6:
            lfail(f"{name}: page {n} is ruled into {max(len(edges) - 1, 0)} columns, not the five this loader was checked against")
        width = edges[-1] - edges[0]
        full = apart(y for _x, y, w, h in rects if h <= 2.0 and w > width - 10)[::-1]
        part = apart(y for x, y, w, h in rects if h <= 2.0 and 20 < w <= width - 10 and x > edges[1] - 3)[::-1]
        if len(full) < 2:
            lfail(f"{name}: page {n} has no ruled table")

        def band(x):
            for i, b in enumerate("ABCDE"):
                if edges[i] - 1 <= x < edges[i + 1] - 1:
                    return b
            return None

        if n == 1:
            outside = [r for r in runs if r[1] > full[0]]
            foot = [r for r in runs if r[1] < full[-1]]
        for k, (top, bottom) in enumerate(zip(full, full[1:]), 1):
            inside = [r for r in runs if bottom < r[1] < top]
            if not inside:
                continue
            if any(band(r[0]) is None for r in inside):
                lfail(f"{name}: page {n}, block {k} has text outside the table's five columns")
            cuts = [top] + [p for p in part if bottom < p < top] + [bottom]
            cells = []
            for ct, cb in zip(cuts, cuts[1:]):
                got = {b: cell_lines([r for r in inside if band(r[0]) == b and cb < r[1] < ct]) for b in "BCD"}
                cells.append((" ".join(got["B"]), " ".join(got["C"]), " ".join(got["D"]), max(len(v) for v in got.values())))
            blocks.append({"page": n, "k": k, "label": cell_lines([r for r in inside if band(r[0]) == "A"]),
                           "term": " ".join(cell_lines([r for r in inside if band(r[0]) == "E"])), "cells": cells})
        for _y, rs in pdftext.rows(pdf, page, res):                         # the second reading: printed lines, no rules
            rs = [r for r in rs if full[-1] < r[1] < full[0]]
            code = pdftext.join([r for r in rs if band(r[0]) == "D"])
            if code in LOCAL_PARTIES or code == "NP":
                by_line[(pdftext.join([r for r in rs if band(r[0]) == "B"]), pdftext.join([r for r in rs if band(r[0]) == "C"]), code)] += 1
    by_cell = Counter((c[0], c[1], c[2]) for b in blocks for c in b["cells"] if c[2] in LOCAL_PARTIES or c[2] == "NP")
    wrapped = sum(1 for b in blocks for c in b["cells"] if (c[2] in LOCAL_PARTIES or c[2] == "NP") and c[3] > 1)
    one_line = Counter((c[0], c[1], c[2]) for b in blocks for c in b["cells"] if (c[2] in LOCAL_PARTIES or c[2] == "NP") and c[3] == 1)
    if sum(by_line.values()) != sum(by_cell.values()) or Counter(k[2] for k in by_line.elements()) != Counter(k[2] for k in by_cell.elements()) \
            or sum((one_line - by_line).values()) or sum((by_line - one_line).values()) != wrapped:
        lfail(f"{name}: read by its ruled lines the table has {sum(by_cell.values())} names with a party code, and read line by line "
              f"{sum(by_line.values())}; the two readings do not agree")
    head = re.sub(r"[,\s]+", " ", " ".join(cell_lines(outside)))
    stamp = re.search(r"(\d{1,2})/ ?(\d{1,2})/ ?(20\d\d)", pdftext.join(foot))
    return {"blocks": blocks, "names": sum(by_cell.values()), "wrapped": wrapped,
            "titled": bool(re.search(r"All Contests and Candidates in the November 3 2026 General Election in Clark County Nevada", head)),
            "ordered": "in order of appearance on the ballot" in head,
            "published": f"{stamp.group(3)}-{int(stamp.group(1)):02d}-{int(stamp.group(2)):02d}" if stamp else ""}


def clark_contests(path, cname):
    """Clark County's contests in the list's order. Comes back as {published, ordered: whether the list says it is in
    ballot order, contests: [{where, label, c: what the office is (clark_office), partisan, names: [(name, party code)],
    vote: how many to choose, special, term, wider}], unplaced: [(contest, names)] for contests this loader does not
    know, rows: names read in all, notc: None of These Candidates lines, left and left_rows: federal and state contests
    and their names, questions, check}. Every name with a party code is in exactly one contest."""
    name = os.path.basename(path)
    data = clark_list(path)
    if not data["titled"]:
        lfail(f"{name} is not headed as Clark County's list of the November 3, 2026 general election")
    out = {"published": data["published"], "ordered": data["ordered"], "contests": [], "unplaced": [], "rows": 0, "notc": 0, "left": 0,
           "left_rows": 0, "left_titles": [], "questions": 0, "check": []}
    section = city = None
    heads = 0
    for b in data["blocks"]:
        where = f"page {b['page']}, block {b['k']}"
        label = squeeze(" ".join(b["label"]))
        filled = [c for c in b["cells"] if c[0] or c[1] or c[2]]
        if label == CLARK_HEAD[0]:
            got = (label, squeeze(" ".join(c[0] for c in b["cells"])), squeeze(" ".join(c[1] for c in b["cells"])),
                   squeeze(" ".join(c[2] for c in b["cells"])), b["term"])
            if got != CLARK_HEAD:
                lfail(f"{name}: {where}: the table's headings are not the five this loader was checked against")
            heads += 1
            continue
        if not filled and not b["term"]:
            if label in CLARK_SECTIONS:
                section, city = label, None
            elif section and CLARK_SECTIONS[section][0] == "city" and re.fullmatch(r"City of .+", label):
                city = label
            elif label:
                lfail(f"{name}: {where} is a heading this loader does not know")
            continue
        if section is None:
            lfail(f"{name}: {where} is a contest before any section heading")
        if not label:
            lfail(f"{name}: {where} has names but no contest beside them (a contest split over two pages?)")
        if len(filled) != len(b["cells"]) and filled:
            lfail(f"{name}: {where} has an empty row between its names")
        text, vote = label, 1
        m = re.search(r"\s*\bVote for not more than (\d+)\b", text)
        if m:
            vote, text = int(m.group(1)), squeeze(text[:m.start()] + " " + text[m.end():])
        m = re.search(r"\s*NOTE: This office covers areas in multiple Counties\b.*$", text)
        wider = bool(m)
        if m:
            text = text[:m.start()].strip()
        text = squeeze(re.sub(r"\s*\(See [^)]*\)", "", text))
        partisan = CLARK_SECTIONS[section][1]
        names = []
        for i, (last, first, code, _lines) in enumerate(filled, 1):
            if not code and NOTC.match(squeeze(f"{last} {first}")):
                out["notc"] += 1
                continue
            if not last or not first or (code not in LOCAL_PARTIES and code != "NP"):
                lfail(f"{name}: {where}, name {i}: a name cell is empty or the party code is not one of the list's own")
            if (code == "NP") == bool(partisan):
                lfail(f"{name}: {where}, name {i}: the party code does not fit the section the contest is printed under")
            family, _, suffix = last.partition(",")
            if suffix.strip() and fold(suffix) not in SUFFIXES:
                lfail(f"{name}: {where}, name {i}: the Last Name cell has a comma that is not before Jr., II and the like")
            names.append((squeeze(f"{first} {family} {suffix.strip()}"), code))
        out["rows"] += len(names)
        try:
            c = clark_office(section, text, city, cname)
        except Unplaced:
            out["unplaced"].append((text, len(names)))
            continue
        if c is None:
            out["left"] += 1
            out["left_rows"] += len(names)
            out["left_titles"].append(text)
            continue
        m = re.fullmatch(r"(\d+)-Year Unexpired Term", b["term"])
        if not m and not re.fullmatch(r"\d+", b["term"]):
            lfail(f"{name}: {where}: the TERM cell is not a number of years or an unexpired term")
        if wider:
            out["check"].append(f"{text}: the list says this office reaches other counties too; only {cname} is recorded")
        out["contests"].append({"where": where, "label": text, "c": c, "partisan": partisan, "names": names, "vote": vote,
                                "special": bool(m), "term": b["term"], "wider": wider})
    if heads != len({b["page"] for b in data["blocks"]}):
        lfail(f"{name}: not every page begins with the table's headings")
    if out["rows"] != data["names"]:
        lfail(f"{name}: {data['names']} names carry a party code but {out['rows']} were placed in contests")
    return out


# --------------------------------------------------------------------------------- Washoe County's list (a workbook)

def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = squeeze(v)
    return t.split(".")[0] if re.fullmatch(r"-?\d+\.0+", t) else t


def xlsx_cells(path):
    """A workbook's first sheet read without openpyxl, from the sheet's own XML and the shared strings: {(row, column
    counted from 0): text}. The second route by which Washoe County's list is read."""
    xns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    z = zipfile.ZipFile(path)
    sheet = ET.fromstring(z.read("xl/workbook.xml")).find(f"{xns}sheets/{xns}sheet")
    target = next(r.get("Target") for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels")) if r.get("Id") == sheet.get(f"{rns}id"))
    target = target.lstrip("/") if target.startswith("/") else "xl/" + target
    shared = []
    if "xl/sharedStrings.xml" in z.namelist():
        for si in ET.fromstring(z.read("xl/sharedStrings.xml")).findall(f"{xns}si"):
            shared.append("".join(t.text or "" for t in si.findall(f"{xns}t")) + "".join(t.text or "" for r in si.findall(f"{xns}r") for t in r.findall(f"{xns}t")))
    out = {}
    for c in ET.fromstring(z.read(target)).iter(f"{xns}c"):
        m = re.fullmatch(r"([A-Z]+)(\d+)", c.get("r") or "")
        if not m:
            continue
        col = 0
        for ch in m.group(1):
            col = col * 26 + ord(ch) - 64
        v = c.find(f"{xns}v")
        if c.get("t") == "s" and v is not None:
            text = shared[int(v.text)]
        elif c.get("t") == "inlineStr":
            text = "".join(t.text or "" for t in c.iter(f"{xns}t"))
        else:
            text = v.text if v is not None else ""
        text = cell_text(text)
        if text:
            out[(int(m.group(2)), col - 1)] = text
    return out


def washoe_contests(path, cname):
    """Washoe County's contests in the workbook's order, in the shape clark_contests() gives. The candidates' table is
    read by its headings; the ballot questions' table under it is counted and not read."""
    import openpyxl
    name = os.path.basename(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    if len(wb.worksheets) != 1:
        lfail(f"{name} has {len(wb.worksheets)} sheets, not the one this loader was checked against")
    grid = [[cell_text(v) for v in row] for row in wb.worksheets[0].iter_rows(values_only=True)]
    modified = wb.properties.modified
    wb.close()
    title = next((v for row in grid for v in row if v), "")
    if title != "2026 General Election":
        lfail(f"{name} is not headed 2026 General Election")
    head = next((i for i, row in enumerate(grid) if all(h in row for h in WASHOE_HEAD)), None)
    if head is None:
        lfail(f"{name}: no row carries the headings {WASHOE_HEAD}")
    col = {h: grid[head].index(h) for h in WASHOE_HEAD}
    rows, questions, asked, last = [], 0, False, head
    for i in range(head + 1, len(grid)):
        if not any(grid[i]):
            continue
        nm, party, office, otype, term, vote = (grid[i][col[h]] for h in WASHOE_HEAD)
        if asked:
            questions += 1
        elif nm == "Question" and party == "Text":
            asked = True                                                    # the ballot questions' own table: counted, not read
        elif not (nm and office and otype):
            lfail(f"{name}: row {i + 1} lacks a name, an office or an office type")
        else:
            rows.append((i + 1, nm, party, office, otype, term, vote))
            last = i
    other = xlsx_cells(path)                                                # the second reading, from the workbook's own XML
    for i in range(head, last + 1):
        for j, v in enumerate(grid[i]):
            if v != other.get((i + 1, j), ""):
                lfail(f"{name}: row {i + 1} reads differently by the two routes (openpyxl, and the workbook's own XML)")
    if sum(1 for (r, _c) in other if head < r <= last + 1) != sum(1 for i in range(head, last + 1) for v in grid[i] if v):
        lfail(f"{name}: the two routes (openpyxl, and the workbook's own XML) do not find the same number of filled cells")
    out = {"published": modified.date().isoformat() if modified else "", "ordered": False, "contests": [], "unplaced": [], "rows": 0, "notc": 0,
           "left": 0, "left_rows": 0, "left_titles": [], "questions": questions, "check": []}
    groups, seen = [], set()
    for r in rows:
        key = r[3:]
        if groups and groups[-1][0] == key:
            groups[-1][1].append(r)
            continue
        if key in seen:
            lfail(f"{name}: row {r[0]}: a contest is printed in two places")
        seen.add(key)
        groups.append((key, [r]))
    for (office, otype, term, vote), rs in groups:
        where = f"rows {rs[0][0]}-{rs[-1][0]}"
        partisan = 1 if otype.endswith(" Partisan") else 0
        names = []
        for n, nm, party, *_rest in rs:
            if NOTC.match(nm):
                if party:
                    lfail(f"{name}: row {n}: a party beside None of these candidates")
                out["notc"] += 1
                continue
            if (party in LOCAL_PARTIES) != bool(partisan) or (not partisan and party not in ("", "NP")):
                lfail(f"{name}: row {n}: the Party cell does not fit the contest's Office Type")
            names.append((squeeze(NVL.turned(nm)), party if partisan else "NP"))
        out["rows"] += len(names)
        try:
            c = washoe_office(office, otype, cname)
        except Unplaced:
            out["unplaced"].append((plain_caps(office), len(names)))
            continue
        if c is None:
            out["left"] += 1
            out["left_rows"] += len(names)
            out["left_titles"].append(plain_caps(office))
            continue
        if not re.fullmatch(r"\d+ Years?", term) and "unexpired" not in term.lower():
            lfail(f"{name}: {where}: the Term Length cell is not a number of years")
        if not vote.isdigit() or int(vote) < 1:
            lfail(f"{name}: {where}: the Vote For cell is not a number")
        out["contests"].append({"where": where, "label": plain_caps(office), "c": c, "partisan": partisan, "names": names, "vote": int(vote),
                                "special": "unexpired" in term.lower(), "term": term, "wider": False})
    if out["rows"] + out["notc"] != len(rows):
        lfail(f"{name}: {len(rows)} name rows but {out['rows'] + out['notc']} were placed in contests")
    return out


LOCAL_COUNTIES = [
    {"fips": "32003", "src": SRC_CLARK, "agency": "Clark County Election Department", "file": CLARK_FILE, "url": CLARK_URL, "page": CLARK_PAGE,
     "magic": b"%PDF-", "read": clark_contests, "find": clark_link,
     "title": "All Contests and Candidates in the November 3, 2026, General Election in Clark County, Nevada",
     "how": "A ruled table of five columns, read by its own lines: contest, name on the ballot (last name; first, middle or nickname), party "
            "and term. Read twice, by the ruled lines and again line by line, and the two readings agree."},
    {"fips": "32031", "src": SRC_WASHOE, "agency": "Washoe County Registrar of Voters", "file": WASHOE_FILE, "url": WASHOE_URL, "page": WASHOE_PAGE,
     "magic": b"PK", "read": washoe_contests, "find": None,
     "title": "Washoe County 2026 General Election Contests (workbook)",
     "how": "A workbook of one sheet, its columns found by their headings: Name on Ballot, Party, Office Sought, Office Type, Term Length and "
            "Vote For. Read twice, through a workbook reader and again from the workbook's own XML, and the two readings agree."},
]


def local_rows(folder, county_zip, say, refresh=False):
    """Clark and Washoe counties' county, township, city, school, district and local court contests on the November
    ballot: rows for sl_races, sl_candidates, sl_places, sl_sources, sl_gaps and sl_notes, and the lines of the report.
    Each county's list is read whole, and every name on it is placed in exactly one contest or counted as left out
    (a federal or state office, or None of these candidates). Nothing is placed by guess: a contest this loader does
    not know, or a county whose list could not be had, becomes a gap."""
    os.makedirs(folder, exist_ok=True)
    cnames = census_county_names(county_zip)
    if len(cnames) != 17:
        lfail(f"the Census county file gives Nevada {len(cnames)} counties, not 17")
    ppath, cplaces, n_places = census_places(folder, {fold(n): g for g, n in cnames.items()}, say)
    out = {"races": [], "cands": [], "places": [], "src": [], "gaps": [], "notes": [], "report": []}
    places, seen_ids, loaded, uncoded = {}, set(), [], []
    for cfg in LOCAL_COUNTIES:
        fips, cname = cfg["fips"], cnames[cfg["fips"]]
        path, url, fetched, how = kept(folder, cfg["file"], cfg["url"], cfg["magic"], say, cfg["find"], 0 if refresh else LOCAL_KEEP_DAYS)
        if path is None:
            out["gaps"].append((STATE, "county", fips, cname, LOCAL_WHAT,
                                f"{cname}'s election office publishes its own list of November contests and candidates, but the file could not be "
                                "fetched when this was loaded and no earlier copy was on hand; loading again will fetch it.", cfg["page"]))
            out["report"].append(f"    CHECK Nevada (county and local): {cname}'s list could not be fetched ({how}); a gap is recorded")
            continue
        data = cfg["read"](path, cname)
        n_races = n_names = blanked = 0
        for con_ in data["contests"]:
            c = con_["c"]
            p = c["place"]
            if p[0] == "county":
                jname, jid, cids = cname, fips, [fips]
            elif p[0] == "city":
                hits = [h for h in cplaces.get(fold(p[1]), []) if fips in h[2]]
                if len(hits) == 1:
                    jname, jid, cids = hits[0][1], f"{STATE}-M-{hits[0][0]}", sorted(set(hits[0][2]) | {fips})
                    places.setdefault(jid, ["mcd", jid, jname, set(), SRC_PLACES])
                else:                                                       # no single Census place of that name in the county: named as the list does
                    jname, jid, cids = f"{p[1]} city", f"{STATE}-M-{fips[2:]}-{slug(p[1])}-city", [fips]
                    places.setdefault(jid, ["mcd", jid, jname, set(), cfg["src"]])
                    uncoded.append(jname)
            else:
                letter, kind = {"school": ("S", "school"), "hospital": ("H", "hospital"), "special": ("X", "special")}[p[0]]
                jname, jid, cids = p[1], f"{STATE}-{letter}-{fips[2:]}-{slug(p[1])}", [fips]
                places.setdefault(jid, [kind, jid, jname, set(), cfg["src"]])
            if jid in places:
                places[jid][3].update(cids)
            special = 1 if con_["special"] else 0
            rid = f"2026-{STATE}-" + "-".join([jid if jid.isdigit() else jid[len(STATE) + 1:], c["kind"].replace("_", "-")]
                                             + [slug(x) for x in (c.get("district"), c.get("seat")) if x]) + ("-S" if special else "")
            if rid in seen_ids:
                lfail(f"{cfg['file']}: {con_['where']}: two contests come to the same race id ({rid})")
            seen_ids.add(rid)
            notes = [c.get("note")]
            if con_["vote"] > 1:
                notes.append(f"Voters choose up to {con_['vote']}.")
            if special:
                notes.append(f"An election for the rest of an unexpired term (the county's list gives the term as {con_['term']}).")
            if con_["wider"]:
                notes.append("The county's list says this office also covers parts of other counties.")
            placed = []
            for i, (nm, code) in enumerate(con_["names"], 1):
                caps = bool(re.search(r"[A-Z]", nm)) and nm == nm.upper()
                shown = re.sub(r"(['’\"(“])([a-z])", lambda m: m.group(1) + m.group(2).upper(), proper(nm)) if caps else nm
                if not shown or contact_like(shown, True):
                    blanked += 1                                            # counted, never shown
                    continue
                party = NVL.party_of(code) if con_["partisan"] else NONPARTISAN
                placed.append([rid, "general", GENERAL, shown, party, NVL.colour(party) if con_["partisan"] else "N",
                               i if data["ordered"] else None, 0, 0, None, None, None, None, cfg["src"], LOCAL_CAPS if caps else None])
            if len({x[3] for x in placed}) != len(placed):
                lfail(f"{cfg['file']}: {con_['where']}: the same name twice in one contest")
            if not placed:
                notes.append("The county's list shows no candidate for this contest.")
            out["races"].append((rid, STATE, c["level"], c["kind"], c["office"], jname, jid, json.dumps(cids), c.get("district"), c.get("seat"),
                                 special, con_["partisan"], None, None, None, GENERAL, " ".join(x for x in notes if x) or None))
            out["cands"] += placed
            n_races += 1
            n_names += len(placed)
        if n_names + blanked + data["left_rows"] + sum(k for _t, k in data["unplaced"]) != data["rows"]:
            lfail(f"{cfg['file']}: {data['rows']} names read, but those placed, left to the state offices and left unplaced do not add up to that")
        for title, k in data["unplaced"]:
            out["gaps"].append((STATE, "county", fips, cname, f"a contest not placed yet: {title}",
                                f"{cname}'s list carries this contest, with {k} name{'s' if k != 1 else ''}, but this loader does not yet know how "
                                "to file the office, so it is left out rather than placed by guess.", url))
        loaded.append({"cfg": cfg, "cname": cname, "races": n_races, "names": n_names, "data": data})
        out["src"].append((
            cfg["src"], STATE, "official candidate list", cfg["agency"], cfg["title"], url, data["published"], fetched, sha_file(path),
            data["rows"] + data["notc"],
            f"{cname}'s own list of every contest and candidate on its November ballot. {cfg['how']} The file has no address, "
            f"telephone, e-mail or website column of any kind, so nothing of that sort is read or kept. {data['rows']} names in all: {n_names} in "
            f"{n_races} county, township, city, school, district and local court contests are loaded; {data['left_rows']} in {data['left']} "
            "federal and state contests (Congress, the statewide offices, the Legislature, the Supreme Court and the Board of Regents) are left "
            f"to the state's own list; {data['notc']} lines reading None of these candidates are not candidates"
            + (f"; {sum(k for _t, k in data['unplaced'])} names in {len(data['unplaced'])} contests this loader does not know are recorded as gaps"
               if data["unplaced"] else "")
            + (f"; {blanked} name cells that looked like contact details were left out" if blanked else "")
            + (f"; {data['questions']} ballot questions in the table below the candidates are not read" if data["questions"] else "") + ". "
            + ("The list says it is in order of appearance on the ballot, subject to change, and that order is stored. " if data["ordered"]
               else "The list states no ballot order, so none is stored. ")
            + "It has no status column: a candidate who withdrew is simply not on it."))
        out["report"].append(
            f"    Nevada (county and local): {cname}: {data['rows']} names read ({how}); {n_races} contests and {n_names} names loaded; "
            f"{data['left_rows']} names in {data['left']} federal and state contests left to the state's own list; {data['notc']} None of these "
            f"candidates lines; {data['questions']} ballot questions not read; published {data['published'] or 'not dated'}")
        for line in data["check"]:
            out["report"].append(f"    CHECK Nevada (county and local): {cname}: {line}")
        for title, k in data["unplaced"]:
            out["report"].append(f"    CHECK Nevada (county and local): {cname}: a contest this loader cannot place, recorded as a gap: {title} ({k} names)")
        if blanked:
            out["report"].append(f"    CHECK Nevada (county and local): {cname}: {blanked} name cells looked like contact details and were left out")
    for name in sorted(set(uncoded)):
        out["report"].append(f"    CHECK Nevada (county and local): {name} is not one incorporated place in the Census Bureau's list; filed by its name")

    # every other county: its own list is not read yet
    have = {x["cfg"]["fips"] for x in loaded} | {g[2] for g in out["gaps"] if g[1] == "county" and g[4] == LOCAL_WHAT}
    so_far = ("so far only " + " and ".join(f"{x['cname']}'s" for x in loaded) + (" lists have" if len(loaded) > 1 else " list has") + " been read"
              if loaded else "so far no county's list has been read")
    for geoid, cname in sorted(cnames.items()):
        if geoid not in have:
            clerk = "Carson City's clerk publishes its" if not cname.endswith("County") else "Each Nevada county's clerk publishes the county's"
            out["gaps"].append((STATE, "county", geoid, cname, LOCAL_WHAT,
                                f"Not loaded yet. {clerk} own list of November contests and candidates; {so_far}, and {cname}'s waits its turn.",
                                None))

    # the state offices the counties' lists also print and nothing else loads yet: said, so that they are not dropped silently
    titles = [t for x in loaded for t in x["data"]["left_titles"]]
    regents = sorted({int(m.group(1)) for t in titles for m in [re.search(r"Board of Regents,? District (\d+)", t)] if m})
    seats = sorted({m.group(1) for t in titles for m in [re.search(r"Supreme Court\b.*\bSeat ([A-Z])\b", t)] if m})
    listed = lambda xs: ", ".join(map(str, xs[:-1])) + (" and " if len(xs) > 1 else "") + str(xs[-1])
    if regents or seats:
        what = " and ".join(x for x in ("Board of Regents" if regents else "", "Supreme Court" if seats else "") if x) + " seats"
        which = "; ".join(x for x in (f"Board of Regents district{'s' if len(regents) > 1 else ''} {listed(regents)}" if regents else "",
                                      f"Supreme Court seat{'s' if len(seats) > 1 else ''} {listed(seats)}" if seats else "") if x)
        out["gaps"].append((STATE, "state", STATE, NAME, what,
                            f"Not loaded yet. These are state offices that the counties' lists print too ({which}), but only county and local "
                            "offices are read from those lists, and the Secretary of State's own list, which covers the whole state, has not been "
                            "read: its website turns away automated requests.", None))

    out["places"] = [(k, pid, pname, json.dumps(sorted(cids)), src) for k, pid, pname, cids, src in sorted(places.values(), key=lambda p: p[1])]
    out["src"].append((SRC_PLACES, STATE, "official place codes", "U.S. Census Bureau", "2020 place codes, Nevada (st32_nv_place2020.txt)", PLACE_URL,
                       "2020", mdate(ppath), sha_file(ppath), n_places,
                       f"The code and the name of each incorporated city on the counties' lists ({sum(1 for p in places.values() if p[4] == SRC_PLACES)} "
                       "used here): Henderson city, Reno city and so on. Places only; the file names no person."))
    lv = Counter(r[2] for r in out["races"])
    n_races, n_names = len(out["races"]), len(out["cands"])
    who = " and ".join(x["cname"].replace(" County", "") for x in loaded)
    out["notes"].append((STATE, "local_calendar", CALENDAR, CALENDAR_SOURCE, NRS_URL))
    if loaded:
        word = lambda n: ("no one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen "
                          "seventeen").split()[n] if 0 <= n <= 17 else str(n)
        coverage = (f"Loaded for {word(len(loaded))} of Nevada's {word(len(cnames))} counties, {who}, each from its own election office's list of "
                    f"November contests and candidates: {n_races} contests and {n_names} names (county offices and constables; district court "
                    "judges, justices of the peace and municipal court judges; city offices; school district trustees; water, improvement and "
                    f"other district boards). The other {word(len(cnames) - len(loaded))} counties are not loaded yet. Not loaded from these "
                    "lists: ballot questions, and the federal and state offices they also carry (Congress, the statewide offices, the Legislature, "
                    "the Supreme Court and the Board of Regents). An office settled at the June primary is on neither list, because it is not on "
                    "the November ballot.")
    else:
        coverage = ("No county or local race is loaded yet: neither Clark County's nor Washoe County's list of November contests and candidates "
                    "could be fetched when this was loaded. Ballot questions are not loaded in any case.")
    out["notes"].append((STATE, "local_coverage", coverage,
                         "Clark County Election Department and Washoe County Registrar of Voters: each county's own list of November 3, 2026 "
                         "contests and candidates", CLARK_PAGE))
    for table, rows, fields in (("race", out["races"], (4, 5, 8, 9, 16)), ("place", out["places"], (2,)), ("gap", out["gaps"], (3, 4, 5)),
                                ("note", out["notes"], (2, 3)), ("source", out["src"], (3, 4, 10))):
        for row in rows:
            if any(row[i] and (contact_like(row[i], table in ("race", "place")) or (table == "race" and i == 16 and STREET.search(row[i])))
                   for i in fields):
                lfail(f"a {table} text of {row[0] if table != 'place' else row[1]} looks like it holds a contact detail; nothing was written")
    out["report"].append(
        f"    Nevada (county and local): {n_races} contests and {n_names} names in {who or 'no county'} (" + ", ".join(
            f"{k} {lv[k]}" for k in ("county", "city", "school", "hospital", "other", "court") if lv[k])
        + f"); {len(out['places'])} places; gaps: " + ", ".join(f"{k} {v}" for k, v in sorted(Counter(g[1] for g in out["gaps"]).items()))
        + "; every name read is placed in exactly one contest or counted as left out")
    return out


def load(db_path, say=print, cache=CACHE, roster_db=ROSTER, county_zip=COUNTY_ZIP, refresh=False):
    folder = os.path.join(cache, "nv")
    report = []
    leg, leg_how = legislature(folder, say, refresh)
    legs, offs = roster(roster_db)
    counties = census_counties(county_zip)
    gpaths, gprob = saved(folder, "general")
    ppaths, pprob = saved(folder, "primary")
    fpaths, fprob = saved(folder, "filings")
    report += gprob + pprob + fprob
    for sec, ok in leg["constitution"]["found"].items():
        if not ok:
            report.append(f"the Legislature's copy of the Constitution no longer carries Article 5, {sec} as read ({CONST_TEXT[sec][:50]}...)")

    # ---- races: the statewide offices, the Senate seats whose terms end in 2026, every Assembly seat
    races = {}
    by_seat = {"Senate": {}, "Assembly": {}}
    for p in legs:
        by_seat.setdefault(p["chamber"], {}).setdefault(p["district"], []).append(p)
    up = {ch: sorted((int(d) for d, r in leg["chambers"][ch]["members"].items() if r["term_ends"] == "2026")) for ch in SEATS}
    if len(up["Senate"]) not in (10, 11):
        report.append(f"the Legislature's list gives {len(up['Senate'])} Senate terms ending in 2026, not half of 21")
    if up["Assembly"] != list(range(1, 43)):
        report.append(f"the Legislature's list gives {len(up['Assembly'])} Assembly terms ending in 2026, not all 42")

    def county_ids(names):
        bad = [n for n in names if fold(n) not in counties]
        if bad:
            fail(f"county names on the Legislature's page not in the Census file: {bad}")
        return json.dumps(sorted(counties[fold(n)][0] for n in names)) if names else None

    for key, okind, shown, _pat, rkey in STATEWIDE:
        rid = f"2026-{STATE}-{key}"
        h = offs.get(rkey) if rkey else None
        if rkey and h is None:
            report.append(f"{rid}: the roster carries no {shown}")
        if h and not str(h.get("term_end") or "").startswith("2027-01"):
            report.append(f"{rid}: the roster's term for {h['full']} ends {h.get('term_end') or '(no date)'}, not January 2027")
        note = [STATEWIDE_NOTE.get(key, OTHER_STATEWIDE_NOTE)]
        if rkey is None:
            note.append(NO_HOLDER_NOTE)
        races[rid] = dict(race_id=rid, state=STATE, level="statewide", office_kind=okind, office=shown, jurisdiction=NAME, jurisdiction_id=FIPS,
                          county_ids=None, district=None, seat=None, special=0, partisan=1, holder_id=h["id"] if h else None,
                          holder_name=h["full"] if h else None, holder_party=h["party"] if h else None, election_date=GENERAL,
                          _note=note, _holder=h, _extra=())
    for ch, prefix, okind, office in (("Senate", "SS", "state_senate", "State Senator"), ("Assembly", "SH", "state_house", "Member of the Assembly")):
        n_up = len(up[ch])
        for d in up[ch]:
            d = str(d)
            rid = f"2026-{STATE}-{prefix}{d}"
            lm = leg["chambers"][ch]["members"][d]
            hs = by_seat.get(ch, {}).get(d, [])
            h = hs[0] if len(hs) == 1 else None
            note = []
            if ch == "Senate":
                note.append(f"Nevada senators serve four-year terms, about half the Senate elected every two years; this seat is one of the "
                            f"{n_up} of 21 whose terms end in 2026 (the Legislature's list of current senators).")
            if h is None:
                report.append(f"{rid}: {len(hs)} sitting members in the roster for this seat")
                note.append("The roster shows no single sitting member for this seat.")
            elif re.search(r"vacan", lm["name"], re.I):
                report.append(f"{rid}: the Legislature lists the seat as vacant; the roster names {h['full']}")
            else:
                if not person_fits(NVL.turned(lm["name"]), h):
                    report.append(f"{rid}: the Legislature lists {NVL.turned(lm['name'])}; the roster's holder is {h['full']}")
                if (lm["party"] or "") != (h["party"] or ""):
                    report.append(f"{rid}: the Legislature gives {h['full']}'s party as {lm['party']}; the roster as {h['party']}")
            races[rid] = dict(race_id=rid, state=STATE, level="legislature", office_kind=okind, office=office,
                              jurisdiction=f"{'Senate' if ch == 'Senate' else 'Assembly'} District {d}", jurisdiction_id=d,
                              county_ids=county_ids(lm["counties"]), district=d, seat=None, special=0, partisan=1,
                              holder_id=h["id"] if h else None, holder_name=h["full"] if h else None, holder_party=h["party"] if h else None,
                              election_date=GENERAL, _note=note, _holder=h, _extra=(NVL.turned(lm["name"]),))

    everyone = legs + list(offs.values())

    def court_race(c):
        rid = race_id(c)
        if rid not in races:
            sc = c[0] == "supreme"
            races[rid] = dict(race_id=rid, state=STATE, level="court", office_kind="supreme_court" if sc else "court_of_appeals",
                              office="Justice of the Supreme Court" if sc else "Judge of the Court of Appeals", jurisdiction=NAME,
                              jurisdiction_id=FIPS, county_ids=None, district=None, seat=f"{'Seat' if sc else 'Department'} {c[1]}", special=0,
                              partisan=0, holder_id=None, holder_name=None, holder_party=None, election_date=GENERAL,
                              _note=[COURT_NOTE, "The Open States roster does not carry judges, so no holder is shown."], _holder=None, _extra=())
        return rid

    def identify(rid, name, party, names_in_race):
        """(incumbent, state_member_id, note) for one name in one race."""
        race = races[rid]
        h = race["_holder"]
        if h is not None and person_fits(name, h, race["_extra"]):
            rivals = [n for n in names_in_race if n != name and person_fits(n, h, race["_extra"])]
            if not rivals:
                return 1, h["id"], None
            report.append(f"{rid}: more than one name fits the holder {h['full']}; none is marked")
            return 0, None, None
        if race["level"] == "court":
            return 0, None, None
        pool = [p for p in everyone if person_fits(name, p) and (p.get("party") or "") == (party or "") and not (h and p["id"] == h["id"])]
        if len(pool) == 1:
            p = pool[0]
            where = (f"the Nevada {'Senate' if p['chamber'] == 'Senate' else 'Assembly'}, District {p['district']}" if "chamber" in p
                     else f"the office of {OFFICE_WORDS.get(p.get('office'), p.get('office') or 'a statewide officer')}")
            return 0, p["id"], f"Serves today in {where}."
        return 0, None, None

    # ---- the November list
    cands, general_loaded, off, other, seen, read = [], bool(gpaths), [], Counter(), [], 0
    nominee = defaultdict(list)                                              # (race, party) -> names on the November list
    if gpaths:
        on, off, other, seen, read = general_list(gpaths, report)
        names_by_race = defaultdict(list)
        for c, raw, party, order in on:
            if c[0] in ("supreme", "appeals"):
                court_race(c)
            rid = race_id(c)
            if rid not in races:
                report.append(f"{rid}: on the November list but not a seat whose term ends in 2026; left out")
                continue
            names_by_race[rid].append((raw, party, order))
        for rid, listed in names_by_race.items():
            people = [p for p in (races[rid]["_holder"],) if p] + everyone
            shown_names = [(show(raw, people), party, order) for raw, party, order in listed]
            plain = [n for (n, _caps), _p, _o in shown_names]
            per_party = Counter(p for _n, p, _o in shown_names if p in PRIMARY_CODE)
            for p, k in per_party.items():
                if k > 1:
                    fail(f"the November list reads {k} {p} candidates for {rid}; a local office may have been read as a state one")
            for (name, caps), party, order in shown_names:
                inc, mid, note = identify(rid, name, party, plain)
                if inc and races[rid]["holder_party"] and party not in (races[rid]["holder_party"], NONPARTISAN):
                    report.append(f"{rid}: the November list gives {name}'s party as {party}; the roster gives the holder's as "
                                  f"{races[rid]['holder_party']} (the roster's figure is stored as the holder's party)")
                code = "N" if party == NONPARTISAN else NVL.colour(party)
                cands.append([rid, "general", GENERAL, name, party, code, order, inc, 0, None, None, None, mid, SRC_GENERAL,
                              " ".join(x for x in (CAPS if caps else None, note) if x) or None])
                nominee[(rid, party)].append(name)
    else:
        for r in races.values():
            r["_note"].append(WAITING_NOTE)

    # ---- the June 9 primary: official votes, or the filing list's fields without votes
    fields, used_results, used_filings = 0, [], []
    counted = {}
    if ppaths:
        merged = {}
        for path in ppaths:
            for key, f in primary_results(path, report).items():
                if key in merged:
                    report.append(f"{race_id(key[0])} {key[1]}: in two saved results files; the first is kept")
                    continue
                merged[key] = f
            used_results.append(path)
        for (c, party), f in sorted(merged.items(), key=lambda kv: (race_id(kv[0][0]), kv[0][1])):
            rid = race_id(c)
            if c[0] in ("supreme", "appeals") and general_loaded:
                court_race(c)
            if rid not in races:
                report.append(f"{rid} {party} primary: in the results but not a race loaded here (a court seat waits for the November list)")
                continue
            names_sum = sum(v for _n, v in f["names"])
            total = names_sum + (f["notc"] or 0)
            if f["total"] is not None and f["total"] not in (total, names_sum):
                fail(f"the {rid} {party} primary's candidates add up to {total:,}; the results print a total of {f['total']:,}")
            counted[(rid, party)] = (len(f["names"]), total)
            if len(f["names"]) < 2:
                continue
            fields += 1
            code = "NP" if party == NONPARTISAN else PRIMARY_CODE.get(party, party[:3].upper())
            if f["notc"]:
                races[rid]["_note"].append(f"In the June 9 {'nonpartisan' if party == NONPARTISAN else party} primary, \"None of these "
                                           f"candidates\" had {f['notc']:,} votes (counted in the shares; it is not a candidate).")
            people = [p for p in (races[rid]["_holder"],) if p] + everyone
            shown_names = [(show(n, people), v) for n, v in f["names"]]
            plain = [n for (n, _c), _v in shown_names]
            listed = nominee.get((rid, party), []) if party != NONPARTISAN else [c2[3] for c2 in cands if c2[0] == rid and c2[1] == "general"]
            won = {n for n in plain if any(person_fits(n, {"first": " ".join(name_parts(x)[0]), "last": name_parts(x)[1]}) for x in listed)}
            top = max(shown_names, key=lambda s: s[1])
            tie = sum(1 for s in shown_names if s[1] == top[1]) > 1
            for (name, caps), votes in sorted(shown_names, key=lambda s: -s[1]):
                note = [CAPS_RESULTS if caps else None]
                if general_loaded and listed:
                    outcome = "advanced" if name in won else "lost"
                elif general_loaded:
                    outcome = None
                    if name == top[0][0]:
                        note.append("The November list names no candidate of this party for this race.")
                elif party != NONPARTISAN and not tie:
                    outcome = "advanced" if name == top[0][0] else "lost"
                    if name == top[0][0]:
                        note.append("Most votes (Nevada nominates by plurality); the November list is not loaded yet to confirm it.")
                else:
                    outcome = None
                inc, mid, n2 = identify(rid, name, party, plain)
                note.append(n2)
                cands.append([rid, f"primary-{code}", PRIMARY, name, party, "N" if party == NONPARTISAN else NVL.colour(party), None, inc, 0,
                              votes, round(100 * votes / total, 1) if total else None, outcome, mid, SRC_PRIMARY,
                              " ".join(x for x in note if x) or None])
            if general_loaded and listed and party != NONPARTISAN:
                if not won:
                    report.append(f"{rid} {party} primary: the November candidate ({', '.join(listed)}) is not among the primary's names")
                elif top[0][0] not in won:
                    report.append(f"{rid} {party} primary: {top[0][0]} had the most votes; the November list names {', '.join(listed)}")
    elif fpaths:
        used_filings = fpaths
        for (c, party), names in sorted(filings(fpaths, report).items(), key=lambda kv: (race_id(kv[0][0]), kv[0][1])):
            rid = race_id(c)
            if rid not in races or len(names) < 2:
                continue
            fields += 1
            code = "NP" if party == NONPARTISAN else PRIMARY_CODE[party]
            people = [p for p in (races[rid]["_holder"],) if p] + everyone
            shown_names = [show(n, people) for n in names]
            plain = [n for n, _c in shown_names]
            listed = nominee.get((rid, party), [])
            for name, caps in shown_names:
                inc, mid, n2 = identify(rid, name, party, plain)
                won = any(person_fits(name, {"first": " ".join(name_parts(x)[0]), "last": name_parts(x)[1]}) for x in listed)
                cands.append([rid, f"primary-{code}", PRIMARY, name, party, "N" if party == NONPARTISAN else NVL.colour(party), None, inc, 0,
                              None, None, ("advanced" if won else "lost") if listed else None, mid, SRC_FILINGS,
                              " ".join(x for x in (CAPS if caps else None, "The primary's official vote counts are not loaded yet.", n2) if x)])

    # ---- checks
    keyset = set()
    for c in cands:
        k = (c[0], c[1], c[3])
        if k in keyset:
            fail(f"{c[3]} is listed twice in {c[0]} {c[1]}")
        keyset.add(k)
    if general_loaded:
        for rid, r in races.items():
            if not any(c[0] == rid and c[1] == "general" for c in cands):
                report.append(f"{rid}: no candidates read from the November list")
        for (rid, party), (n, _t) in counted.items():
            if party in PRIMARY_CODE and not nominee.get((rid, party)) and n:
                report.append(f"{rid}: a {party} primary was held but the November list names no {party} candidate")
    if other:
        report.append("offices on the saved list not loaded here: " + "; ".join(f"{t} ({k})" for t, k in sorted(other.items())[:40])
                      + (" ..." if len(other) > 40 else ""))

    for r in races.values():
        r["note"] = " ".join(x for x in r["_note"] if x) or None
    guard(races, cands)
    cols = ("race_id", "state", "level", "office_kind", "office", "jurisdiction", "jurisdiction_id", "county_ids", "district", "seat", "special",
            "partisan", "holder_id", "holder_name", "holder_party", "election_date", "note")
    general = [c for c in cands if c[1] == "general"]
    prim = [c for c in cands if c[1] != "general"]

    # ---- the county and local races, from the counties' own lists (the state rows above are not touched by them)
    local = local_rows(os.path.join(folder, LOCAL_DIR), county_zip, say, refresh)
    clash = sorted({r[0] for r in local["races"]} & set(races))
    if clash:
        fail(f"a local race id is also a state race id ({clash[0]}); nothing was written")
    county_names = census_county_names(county_zip)

    src = [
        (SRC_SENATE, STATE, "legislature roster", "Nevada Legislature (Legislative Counsel Bureau)",
         "Senate: current members, with district, party, counties and Term Ends", LEG_PAGES["Senate"], "", leg["fetched"],
         leg["chambers"]["Senate"]["sha256"], SEATS["Senate"],
         f"Which Senate seats are on the 2026 ballot: the {len(up['Senate'])} whose Term Ends year is 2026 (districts "
         f"{', '.join(map(str, up['Senate']))}), and the counties each district reaches. Only each member row's name, party, district and "
         "county cells and the Term Ends field are read; the page's contact details are never read or kept."),
        (SRC_ASSEMBLY, STATE, "legislature roster", "Nevada Legislature (Legislative Counsel Bureau)",
         "Assembly: current members, with district, party, counties and Term Ends", LEG_PAGES["Assembly"], "", leg["fetched"],
         leg["chambers"]["Assembly"]["sha256"], SEATS["Assembly"],
         f"Every Assembly term ends in 2026 ({len(up['Assembly'])} of 42 as read). Only name, party, district, county and Term Ends are read."),
        (SRC_CONST, STATE, "constitution", "Nevada Legislature (Legislative Counsel Bureau)",
         "Constitution of the State of Nevada, Article 5, Sections 2, 17 and 19", CONST_URL, "", leg["fetched"], leg["constitution"]["sha256"], 3,
         "The governor is elected every four years at the legislative elections; the lieutenant governor, secretary of state, treasurer, "
         "controller and attorney general at the same time. Each section's words checked on the page: "
         + ", ".join(f"{s} {'found' if ok else 'NOT found'}" for s, ok in leg["constitution"]["found"].items()) + "."),
        (SRC_ROSTER, STATE, "roster", "Open States people project (CC0), as loaded into state_nv.sqlite", "Sitting legislators and statewide officials",
         "https://github.com/openstates/people", "", mdate(roster_db), sha_file(roster_db), len(legs) + len(offs),
         "Who holds each seat today (ids, names, parties, districts and term ends only); it does not carry the Treasurer, the Controller or judges."),
        (SRC_COUNTY, STATE, "county codes", "U.S. Census Bureau", "Cartographic boundary file, counties, 2024 (1:500,000)", COUNTY_URL, "2024",
         mdate(county_zip), sha_file(county_zip), len(counties),
         "Five-digit county codes for the counties the Legislature lists under each district, and the names of Nevada's seventeen counties "
         "(Carson City among them) for the county pages."),
    ]
    if gpaths:
        src.append((SRC_GENERAL, STATE, "official candidate list", "Nevada Secretary of State, Elections Division",
                    "Candidate Filing List, 2026 General Election (November 3, 2026): state offices", NVL.LIST_PAGE, "",
                    max(mdate(p) for p in gpaths), ",".join(sha_file(p) for p in gpaths), read,
                    f"Saved from a browser ({', '.join(os.path.basename(p) for p in gpaths)}): the Secretary of State's site answers scripts with "
                    "an Incapsula challenge. Name, office, district, party, status and ballot order read by their headings; addresses, "
                    f"telephones, e-mail and websites never read. State-office rows read: {read}; stored: {len(general)}; withdrawn or removed, "
                    f"left off: {len(off)}; other offices, not loaded here: {sum(other.values())}. Nevada allows no write-in votes."))
    if used_results:
        src.append((SRC_PRIMARY, STATE, "official results", "Nevada Secretary of State, Elections Division",
                    "2026 Official Statewide Primary Election Results (June 9, 2026): state offices", NVL.RESULTS_PAGE, "",
                    max(mdate(p) for p in used_results), ",".join(sha_file(p) for p in used_results), len(prim),
                    f"Saved from a browser ({', '.join(os.path.basename(p) for p in used_results)}). Each printed Total row matches the "
                    "candidates' votes (with None of these candidates, where a statewide contest offers it)."))
    if used_filings:
        src.append((SRC_FILINGS, STATE, "official candidate list", "Nevada Secretary of State, Elections Division",
                    "Candidate Filing List, 2026 election cycle: state offices", NVL.LIST_PAGE, "", max(mdate(p) for p in used_filings),
                    ",".join(sha_file(p) for p in used_filings), len(prim),
                    "Saved from a browser. Each party primary's field; the vote counts are not loaded. Contact columns never read."))

    con = sqlite3.connect(db_path)
    con.executescript(SCHEMA + EXTRA_SCHEMA)
    with con:      # Nevada's rows only, in one transaction
        con.execute("DELETE FROM sl_candidates WHERE race_id IN (SELECT race_id FROM sl_races WHERE state = ?)", (STATE,))
        con.execute("DELETE FROM sl_candidates WHERE race_id LIKE ?", (f"2026-{STATE}-%",))
        con.execute("DELETE FROM sl_races WHERE state = ?", (STATE,))
        con.execute("DELETE FROM sl_places WHERE source_id LIKE ?", (STATE.lower() + "-%",))
        con.execute("DELETE FROM sl_sources WHERE state = ?", (STATE,))
        con.execute("DELETE FROM sl_gaps WHERE state = ?", (STATE,))
        con.execute("DELETE FROM sl_notes WHERE state = ?", (STATE,))
        con.executemany(f"INSERT INTO sl_races VALUES ({','.join('?' * len(cols))})", [tuple(r[k] for k in cols) for r in races.values()] + local["races"])
        con.executemany("INSERT INTO sl_candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", cands + local["cands"])
        con.executemany("INSERT INTO sl_places VALUES (?,?,?,?,?)",
                        [("county", g, n, json.dumps([g]), SRC_COUNTY) for g, n in sorted(county_names.items())] + local["places"])
        con.executemany("INSERT INTO sl_sources VALUES (?,?,?,?,?,?,?,?,?,?,?)", src + local["src"])
        con.executemany("INSERT INTO sl_gaps VALUES (?,?,?,?,?,?,?)", local["gaps"])
        con.executemany("INSERT INTO sl_notes VALUES (?,?,?,?,?)", local["notes"])
    con.close()

    kind = lambda rid: races[rid]["office_kind"] if races[rid]["level"] == "legislature" else races[rid]["level"]
    n = Counter(kind(r) for r in races)
    g = Counter(kind(c[0]) for c in general)
    head = (f"    Nevada (state races): {len(races)} races (Senate {n['state_senate']}, Assembly {n['state_house']}, statewide {n['statewide']}"
            f"{', courts ' + str(n['court']) if n['court'] else ''}); Legislature's lists {leg_how}")
    if general_loaded:
        head += (f"; {len(general)} candidates on the November list (Senate {g['state_senate']}, Assembly {g['state_house']}, statewide "
                 f"{g['statewide']}{', courts ' + str(g['court']) if g['court'] else ''}; {len(off)} withdrawn or removed left off)")
    else:
        head += "; no November candidates stored"
    head += f"; {fields} primary fields, {len(prim)} primary rows" + (" (official votes)" if used_results else " (filing list, no votes)" if used_filings else "")
    say(head)
    if not general_loaded:
        say(f"      waiting: the Secretary of State's certified list for the 2026 General Election, saved from a browser into {folder} as "
            f"{NVL.STEMS['general']}.html (Aurora's Candidate Filing List, {NVL.LIST_PAGE}, 2026 General Election chosen); www.nvsos.gov "
            "answers scripts with an Incapsula challenge")
    if not used_results:
        say(f"      waiting: the official results of the June 9 primary (\"2026 Official Statewide Primary Election Results\", {NVL.RESULTS_PAGE}), "
            f"every state contest, saved into {folder} as {NVL.STEMS['primary']}.html (or {NVL.STEMS['primary']}_<part>.html for further pages)")
    for c in cands:
        if c[12]:
            say(f"      matched: {c[0]} {c[1]}: {c[3]} -> {c[12]}{' (holds this seat)' if c[7] else ''}")
    for line in report:
        say(f"      check: {line}")
    for line in local["report"]:
        say(line)
    return len(general) + len(local["cands"])


if __name__ == "__main__":
    args = sys.argv[1:]
    cache = CACHE
    if "--cache" in args:
        i = args.index("--cache")
        cache = args[i + 1]
        del args[i:i + 2]
    if len(args) != 1:
        raise SystemExit("usage: python -m ballot.state_local_nv <database> [--cache <folder holding nv/>]")
    load(args[0], cache=cache)
