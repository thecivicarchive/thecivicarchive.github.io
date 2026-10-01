"""
ballot/state_local_al.py - Alabama's state races, and its county and county school races, on the November 3, 2026
ballot, into ballot_local_2026.sqlite.

What is on the ballot. Alabama elects its whole Legislature every four years, in the governor's year: all 35 Senate
seats and all 105 House seats are up in 2026 (four-year terms; states/places.py records the same). The statewide
executive offices on the ballot are Governor, Lieutenant Governor (elected separately), Attorney General, Secretary of
State, State Treasurer, State Auditor, the Commissioner of Agriculture and Industries and two Public Service Commission
places; the State Board of Education's districts up this year, the Supreme Court, Court of Civil Appeals and Court of
Criminal Appeals places up this year, and the circuit and district judgeships up this year are on it too. Every one of
these offices is partisan in Alabama. Which of them are up is read from the ballots themselves, not assumed (in the
September 2026 ballots: State Board of Education districts 2, 4, 6 and 8; Supreme Court places 7 and 8; places 4 and 5
of each appeals court; 22 circuit and 34 district judgeships). U.S. Senator and Representative are left to the federal
pages. The federal ballot database (ballot_2026.sqlite) is never opened here.

Race keys: 2026-AL-SS<n> and -SH<n> for the Legislature; -GOV, -LTG, -AG, -SOS, -TREAS, -AUD, -AGR; -PSC<place>,
-SBOE<district>, -SC<place> (Supreme Court), -CIVA<place> and -CRIMA<place> (the appeals courts), -CC<circuit>-<place>
(circuit courts) and -DC<county FIPS>[-<place>] (district courts).

The county and local part (2026-09-30). The same 67 sample ballots are the statewide official source for county
candidates: a county office's candidates are certified to the county's judge of probate, not to the Secretary of State
(the Secretary's 2026 Administrative Calendar, citing Code of Alabama 17-13-5 and 17-13-18). Every contest a ballot
prints for a county office or a county board of education is read (the section "the county and local part" below):
sheriff, coroner, revenue commissioner (Bibb's tax assessor/collector and Shelby's property tax commissioner are filed
under the same kind), tax assessor, tax collector, Jefferson's assistant tax assessor and collector for the Bessemer
Division, license commissioner (Madison's license director), Macon's racing commissioner, members and chairmen of county
commissions, members of county boards of education and county superintendents of education. All are partisan: the
ballot prints a party under every name. Cities and towns hold their own municipal elections (August 26, 2025 and
August 25, 2026 in the Secretary's municipal election calendars), so no mayor, council or city board is on this
ballot; no judge of probate, circuit clerk, district attorney or constable is on it either. Ballot questions are not
loaded.
  - Level county: jurisdiction the Census Bureau's county name, jurisdiction_id the 5-digit county code; race ids
    2026-AL-<county code>-<office kind>[-d<district>][-p<place>]. Level school: jurisdiction "<County> County Board of
    Education" as the ballot names it, jurisdiction_id AL-S-<3-digit county code>-<the name, hyphenated> (the ballot
    carries no district number, so none is guessed); race ids 2026-AL-S-<the same key>-<office kind>[-d..][-p..].
  - A title is read only if it is one of the shapes local_office() knows and names the county whose ballot it is on;
    anything else is left out and written to sl_gaps, never guessed. Every ballot style of the county must print a
    contest's candidates the same way, or the candidates are left out and the race is a gap. A contest printed with a
    Write-in line and no name (Etowah's commission district 6 in September 2026) is kept, with a note.
  - A county's sample ballot draws every version of its ballot (486 November versions in the 67 files). A contest
    drawn on fewer versions than the county has says so in its note ("Not on every ballot in the county"); nothing is
    claimed about a contest on all of them, and no district lines are read. Names are shown as shown_local() writes
    them: ordinary capitals, with the ballot's own small letters kept ("DJ Skegee", J.T., DeRamus).
  - Two routes to the same count: the contests found from their "FOR ..." titles, and a plain count of the ballots'
    "(Vote for" lines and party lines (the straight party box's lines apart). Every vote-for line must head a contest
    that was read (one drawn without the word FOR is named by what it is: Washington County's U.S. House contest in
    September 2026), and every party line must belong to a candidate that was read.
  - Kept cells (ballot_cache/al/local/): each ballot's county and school contests (title, names, parties, the Write-in
    mark, how many styles draw them) as JSON by the file's SHA-256, and what the Secretary's calendars and filing guide
    say, as findings only. One sl_sources row per county's ballot, with the file's own address and SHA-256.
  - sl_gaps and sl_notes (ballot/check_local.py's EXTRA_SCHEMA) are rewritten for Alabama on every run.

Checks, all printed and kept in the source notes: every seat of both chambers has a race, and every race whose ballots
agree has its candidates or says why not; statewide contests are printed in all 67 counties; each primary field's
winner (or runoff pair) must be the party's candidate on the November ballot, and a winner printed for another seat
instead is named (Senate District 25 in 2026: the Democratic primary was voted in Crenshaw and Montgomery counties and
won by Kirk Hatcher, while the November ballots print the seat in Crenshaw, Elmore and Montgomery with Phadra Carson
Foster, and Hatcher for District 26); the counties a primary was voted in are compared with the counties whose
November ballots print the seat; each statewide contest's votes (with over and under votes) are compared with the
party's BALLOTS CAST line county by county; and the parties' typeset certifications are compared contest by contest.

Sources: the Secretary of State's own files, the same ones ballot/lists/al.py reads for the federal races, cached in
ballot_cache/al/.
  - November ballot: the "2026 General Election Sample Ballots", one typeset PDF per county, 67 in all
    (sos.alabama.gov/alabama-votes/2026-general-election-sample-ballots). Each draws every ballot style of the county
    one over another, so the text is read in drawing order (ballot/lists/al.ballot_lines), a style at a time: the date
    heading, each "FOR <office>" contest, its "(Vote for One)" line, each candidate's name in capitals with the party in
    small type beneath, and the Write-in line. Only styles dated NOVEMBER 3, 2026 count (a few counties also draw a
    2022 style beneath). Every 2026 style of every county must print a contest identically (letters compared, spacing
    and case aside), or the race is left out and named. Ballot order is the ballot's own. The State Certifications of
    candidates are scans with no text, so the typeset sample ballots are read instead; a ballot shows no withdrawn
    candidates and no declared write-ins, so none are loaded.
  - Primary fields: the Secretary's precinct results files for the May 19 primary and the June 16 runoff
    (2026_Primary_Election.zip and 2026_PRIMARY_RUNOFF_ELECTION.zip on Elections Data Downloads): one .xls per county
    (Contest Title, Party, Candidate, a column per precinct, ABSENTEE, PROVISIONAL), summed over the counties; the Over
    Votes and Under Votes lines are kept out of the share. A nominee needs a majority; without one the top two go to
    the runoff. They are checked against the parties' certified vote totals where those are typeset (the Democratic
    workbooks for the primary and the runoff, the Republican workbook for the primary; the Republican runoff's
    certification is a scan), and every difference is named in the source note. Alabama prints only contested
    primaries, so every state contest in these files is a field.
  - Who holds each seat: the Open States roster in state_al.sqlite (legislators with is_current = 1 by chamber and
    district; the officials table for the six executive offices it carries). Names, party and ids only.
  - County codes: the Census Bureau's 2024 county file (states_cache/census/cb_2024_us_county_500k.zip). The counties
    a legislative district, a judicial circuit or a State Board of Education district reaches are the counties whose
    sample ballots print its contest (derived, and said so).

Privacy. None of these files carries an address, telephone, e-mail, website or treasurer: the sample ballots print
offices, names, parties and ballot questions; the results print contest titles, names, parties and vote columns. Even
so, only the office, district, name, party, order and vote cells are kept (as JSON of those cells alone), nothing is
printed but counts and checks, and before anything is written every stored text is checked for anything that looks
like a contact detail. The roster is read for ids, names, party, chamber and district only. The Secretary's calendars
and filing guide (agency documents with no candidate in them) are searched for a few sentences and never printed.

    python -m ballot.state_local_al <path to a test database> [--keep <folder for the state part's kept-cells JSON>]
                                    [--local-keep <folder for the local part's kept cells, default ballot_cache/al/local>]
"""

import argparse
import collections
import datetime as dt
import hashlib
import io
import json
import os
import re
import sqlite3
import sys
import zipfile

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from urllib.error import HTTPError, URLError  # noqa: E402

from ballot.check_local import EXTRA_SCHEMA, contact_like  # noqa: E402
from ballot.common import CACHE, fold, name_parts, party_code  # noqa: E402
from ballot.lists.al import (DATA_PAGE, FILES, GENERAL, INFO_PAGE, MONTHS, PRIMARY, REP_RUNOFF_SCAN, RUNOFF, SAMPLE_PAGES, SOS,  # noqa: E402
                             ballot_lines, fetch, number, same_person, sample_links, sha256, xls_sheets)
from ballot.lists.tx import proper  # noqa: E402
from ballot.match import fits  # noqa: E402
from ballot.pdftext import lines as pdf_lines  # noqa: E402
from states import net  # noqa: E402

STATE, FIPS, NAME = "AL", "01", "Alabama"
FOLDER = os.path.join(CACHE, "al")
ROSTER = os.path.join(HERE, "state_al.sqlite")
COUNTY_ZIP = os.path.join(HERE, "states_cache", "census", "cb_2024_us_county_500k.zip")
COUNTY_URL = "https://www2.census.gov/geo/tiger/GENZ2024/shp/cb_2024_us_county_500k.zip"
SENATE_SEATS, HOUSE_SEATS, COUNTIES = 35, 105, 67
AGENCY = "Alabama Secretary of State, Elections Division"
KEPT = "al_2026_general_sample_ballots_state.json"
KEPT_VERSION = 1

SRC_GENERAL = "al-sos-2026-sl-general-sample-ballots"
SRC_PRIMARY = "al-sos-2026-sl-primary-precinct-results"
SRC_RUNOFF = "al-sos-2026-sl-runoff-precinct-results"
SRC_DEM_P, SRC_REP_P, SRC_DEM_R = "al-dem-2026-sl-primary-certified-totals", "al-gop-2026-sl-primary-certified-totals", "al-dem-2026-sl-runoff-certified-totals"
SRC_ROSTER, SRC_COUNTY = "al-openstates-roster", "al-census-2024-county-codes"

PARTY = {"REP": "Republican", "DEM": "Democrat"}                      # as the ballot prints them
PARTY_OF = {"REP": "Republican", "DEM": "Democratic"}                  # "the Democratic primary"
PRINTED = {"democrat": "Democrat", "democratic": "Democrat", "republican": "Republican", "libertarian": "Libertarian",
           "independent": "Independent"}
DATE = re.compile(r"^(%s) (\d{1,2}), (20\d\d)$" % "|".join(MONTHS))
VOTE_FOR = re.compile(r"^\(\s*Vote\s*for\s*(\w+)\s*\)", re.I)
NAME_LINE = re.compile(r"^[\"'“‘]?[^\W\d_][^\d]*$")
CONTACT = re.compile(r"@|https?:|www\.|\.(?:com|org|net|gov|us)\b|\d{3}|P\.?\s?O\.?\s+Box|\bSuite\b", re.I)
# The page's own guard reads a number followed soon after by a street word as an address, Court, Pl and Place included
# (ballot/check_local.py's does not have those three); every note and gap the local part writes is held to the wider one.
LOOKS_LIKE_STREET = re.compile(r"\b\d{1,6}\s+(?:[NSEW]\.?\s+)?[A-Za-z0-9.' -]{1,40}?\s(?:St|Street|Ave|Avenue|Rd|Road|Blvd|Boulevard|Dr|Drive|Ln|"
                               r"Lane|Way|Ct|Court|Cir|Circle|Pkwy|Parkway|Hwy|Highway|Trl|Trail|Pl|Place|Ter|Terrace)\b\.?", re.I)

# ---------- offices ----------
# key -> (level, office_kind, office as shown, the roster's officials.office or None)
EXEC = {"GOV": ("statewide", "governor", "Governor", "governor"),
        "LTG": ("statewide", "lieutenant_governor", "Lieutenant Governor", "lt_governor"),
        "AG": ("statewide", "attorney_general", "Attorney General", "attorney general"),
        "SOS": ("statewide", "secretary_of_state", "Secretary of State", "secretary of state"),
        "TREAS": ("statewide", "state_treasurer", "State Treasurer", "treasurer"),
        "AUD": ("statewide", "state_auditor", "State Auditor", "auditor"),
        "AGR": ("statewide", "agriculture_commissioner", "Commissioner of Agriculture and Industries", None)}
PLACED = {"PSC": ("statewide", "public_service_commissioner", "Public Service Commissioner"),
          "SC": ("court", "supreme_court", "Associate Justice of the Supreme Court"),
          "CIVA": ("court", "court_of_civil_appeals", "Judge of the Court of Civil Appeals"),
          "CRIMA": ("court", "court_of_criminal_appeals", "Judge of the Court of Criminal Appeals")}

# on the letters and digits of a contest's title alone, "FOR" and any party mark taken off
OFFICE_PATTERNS = [
    (re.compile(r"(?:GOVENOR|GOVERNOR)"), "GOV"),
    (re.compile(r"(?:LIEUTENANT|LT)(?:GOVENOR|GOVERNOR)"), "LTG"),
    (re.compile(r"ATTORNEYGENERAL"), "AG"),
    (re.compile(r"SECRETARYOFSTATE"), "SOS"),
    (re.compile(r"(?:STATE)?TREASURER"), "TREAS"),
    (re.compile(r"(?:STATE)?AUDITOR"), "AUD"),
    (re.compile(r"COMMISSIONEROFAGRICULTURE(?:AND)?INDUSTRIES"), "AGR"),
    (re.compile(r"PUBLICSERVICECOMMISSION(?:ER)?(?:PLACE)?(?:NO)?(?P<a>\d{1,2})"), "PSC"),
    (re.compile(r"(?:MEMBER)?STATEBOARDOFEDUCATION(?:MEMBER)?DISTRICT(?:NO)?(?P<a>\d{1,2})"), "SBOE"),
    (re.compile(r"STATESENATOR(?:DISTRICT)?(?:NO)?(?P<a>\d{1,2})"), "SS"),
    (re.compile(r"STATEREPRESENTATIVE(?:DISTRICT)?(?:NO)?(?P<a>\d{1,3})"), "SH"),
    (re.compile(r"ASSOCIATEJUSTICEOFTHESUPREMECOURTPLACE(?:NO)?(?P<a>\d{1,2})"), "SC"),
    (re.compile(r"COURTOFCIVILAPPEALS(?:JUDGE)?PLACE(?:NO)?(?P<a>\d{1,2})"), "CIVA"),
    (re.compile(r"COURTOFCRIMINALAPPEALS(?:JUDGE)?PLACE(?:NO)?(?P<a>\d{1,2})"), "CRIMA"),
    (re.compile(r"CIRCUITCOURTJUDGE(?P<a>\d{1,2})(?:ST|ND|RD|TH)JUDICIALCIRCUITPLACE(?:NO)?(?P<b>\d{1,2})"), "CC"),
    (re.compile(r"DISTRICTCOURTJUDGE(?P<c>[A-Z]+?)(?:COUNTY)?(?:PLACE(?:NO)?(?P<b>\d{1,2}))?"), "DC"),
]
# contests that are not this loader's, by what they are (counted for the report)
COURT_UNREAD = "court contest not understood"
OTHER_PATTERNS = [
    (re.compile(r"^UNITEDSTATES|^USREPRESENTATIVE|^USSENATOR|CONGRESSIONAL"), "federal (the federal pages)"),
    (re.compile(r"EXEC(?:UTIVE)?COMM|REPUBEXE|DEMEXEC|EXECCOMM|GOPEXECUTIVE|REPEXECCOMM"), "party executive committee"),
    (re.compile(r"DISTRICTATTORNEY"), "district attorney (local phase)"),
    (re.compile(r"PROBATE"), "county office (local phase)"),
    (re.compile(r"CHIEFJUSTICE|SUPREMECOURT|APPEALS|JUDGE"), COURT_UNREAD),
    (re.compile(r"BOARDOFEDUCATION|BOARDOFEDU|SUPERINTENDENT"), "county or city school office (local phase)"),
    (re.compile(r"COUNTY|COMMISSION|SHERIFF|CORONER|REVENUE|TAXASSESSOR|TAXCOLLECTOR|LICENSE|BESSEMER"), "county office (local phase)"),
    (re.compile(r"STATEWIDEAMENDMENT"), "statewide ballot question (not an office)"),
    (re.compile(r"TAX|MILL|REFERENDUM|AMENDMENT|FIREDISTRICT|ELECTION|PROPOSED"), "local question (local phase)"),
    (re.compile(r"^BALLOTSCAST|^REGISTEREDVOTERS"), "turnout line"),
]
KIND_OF = {"SS": "state_senate", "SH": "state_house", "SBOE": "state_board_of_education", "CC": "circuit_court",
           "DC": "district_court", **{k: v[1] for k, v in EXEC.items()}, **{k: v[1] for k, v in PLACED.items()}}

LEG_NOTE = ("Alabama elects all 35 senators and all 105 representatives every four years, in the governor's year; every seat "
            "is on the 2026 ballot.")
NO_HOLDER = "The roster this site reads does not carry who holds this office today."
VACANT = "The roster shows no one holding this seat today."
CAPS = "Alabama's sample ballots print names in capitals; they are shown here in ordinary capitals."
RUNOFF_NOTE = "No candidate had a majority; the top two went to the June 16 runoff."
NO_RUNOFF = "No candidate had a majority, and the Secretary of State's runoff results have no runoff for this contest."
NOT_ON_BALLOT = "Won the {party} nomination but is not on the November sample ballots."
NOT_PRINTED = ("No candidate for this seat is printed on any county's November sample ballot (no one qualified, or every "
               "nominee withdrew).")
SPLIT_NOTE = "The counties' sample ballots do not print this race the same way; its candidates are left out until they agree."
DERIVED = "The counties listed are those whose sample ballots print this contest (derived from the ballots)."
DERIVED_RESULTS = "The counties listed are those whose primary results carry this contest (derived from the results)."


def shown(name):
    """A name printed in capitals, in ordinary capitals (ballot.lists.tx.proper), a quoted nickname included:
    ANGELO "Doc" MANCUSO -> Angelo "Doc" Mancuso."""
    raw = name.split()
    out = proper(" ".join(raw)).split()
    if len(out) != len(raw):
        return proper(name)
    for i, w in enumerate(raw):
        m = re.fullmatch(r"([\"'“‘(]+)([^\W\d_][^\"'”’)]*)([\"'”’)]*)(,?)", w)
        if m:
            out[i] = m.group(1) + proper(m.group(2).upper()) + m.group(3) + m.group(4)
    return " ".join(out)


def ordinal(n):
    n = int(n)
    return f"{n}{'th' if 10 <= n % 100 <= 20 else {1: 'st', 2: 'nd', 3: 'rd'}.get(n % 10, 'th')}"


def squash(text):
    """Letters and digits only, in capitals: how contest titles are compared, whatever their spacing."""
    return re.sub(r"[^A-Z0-9]", "", (text or "").upper())


def county_key(name):
    return re.sub(r"[^a-z]", "", fold((name or "").replace("_", " ")))


def classify(title):
    """(key, groups) for one of this loader's offices, ("other", why) for anything else, (None, None) if not understood."""
    s = squash(title)
    s = re.sub(r"^FOR", "", s)
    s = re.sub(r"(?:REP|DEM)$", "", s) if re.search(r"\((?:REP|DEM)\)\s*$", title or "") else s
    for pat, key in OFFICE_PATTERNS:
        m = pat.fullmatch(s)
        if m:
            return key, m.groupdict()
    for pat, why in OTHER_PATTERNS:
        if pat.search(s):
            return "other", why
    return None, None


# ---------- the counties ----------

def census_counties(path=COUNTY_ZIP):
    """{GEOID: "Autauga County"} for Alabama's 67 counties, from the Census Bureau's file."""
    import shapefile                                   # pyshp
    if not os.path.exists(path):
        raise SystemExit(f"Alabama (state races): the Census county file is missing ({path}); the county pages' stage fetches it")
    z = zipfile.ZipFile(path)
    base = next(n[:-4] for n in z.namelist() if n.endswith(".dbf"))
    rdr = shapefile.Reader(dbf=io.BytesIO(z.read(base + ".dbf")))
    fields = [f[0] for f in rdr.fields[1:]]
    out = {}
    for rec in rdr.iterRecords():
        rec = dict(zip(fields, rec))
        if str(rec.get("STATEFP")) == FIPS:
            out[str(rec["GEOID"])] = str(rec["NAMELSAD"])
    if len(out) != COUNTIES:
        raise SystemExit(f"Alabama (state races): the county file gives {len(out)} Alabama counties, not {COUNTIES}")
    return out


# ---------- the races ----------

def race_of(key, g, geoid_of):
    """The race a classified contest belongs to: (race_id, fields), or (None, why)."""
    base = dict(district=None, seat=None, jurisdiction=NAME, jurisdiction_id=FIPS)
    if key in EXEC:
        level, kind, office, _h = EXEC[key]
        return f"2026-{STATE}-{key}", dict(base, level=level, office_kind=kind, office=office)
    if key in PLACED:
        level, kind, office = PLACED[key]
        n = int(g["a"])
        return f"2026-{STATE}-{key}{n}", dict(base, level=level, office_kind=kind, office=office, seat=f"Place {n}")
    if key == "SBOE":
        d = str(int(g["a"]))
        return f"2026-{STATE}-SBOE{d}", dict(base, level="statewide", office_kind="state_board_of_education",
                                              office="Member of the State Board of Education", district=d,
                                              jurisdiction=f"State Board of Education District {d}", jurisdiction_id=f"SBOE{d}")
    if key in ("SS", "SH"):
        d = str(int(g["a"]))
        if not 1 <= int(d) <= (SENATE_SEATS if key == "SS" else HOUSE_SEATS):
            return None, f"a {'Senate' if key == 'SS' else 'House'} district that does not exist ({d})"
        word = "Senate" if key == "SS" else "House"
        return f"2026-{STATE}-{key}{d}", dict(base, level="legislature", office_kind=KIND_OF[key],
                                               office="State Senator" if key == "SS" else "State Representative", district=d,
                                               jurisdiction=f"{word} District {d}", jurisdiction_id=d)
    if key == "CC":
        c, p = int(g["a"]), int(g["b"])
        return f"2026-{STATE}-CC{c}-{p}", dict(base, level="court", office_kind="circuit_court", office="Circuit Court Judge",
                                                district=str(c), seat=f"Place {p}", jurisdiction=f"{ordinal(c)} Judicial Circuit",
                                                jurisdiction_id=f"JC{c}")
    if key == "DC":
        geoid = geoid_of.get(county_key(g["c"]))
        if not geoid:
            return None, f"a district court whose county could not be read ({g['c']!r})"
        p = int(g["b"]) if g.get("b") else None
        return (f"2026-{STATE}-DC{geoid[2:]}" + (f"-{p}" if p else ""),
                dict(base, level="court", office_kind="district_court", office="District Court Judge", seat=f"Place {p}" if p else None,
                     jurisdiction=geoid_of["_name"][geoid], jurisdiction_id=geoid))
    return None, f"an office this loader does not know ({key})"


# ---------- the sample ballots ----------

_LINES = {}


def lines_of(path):
    """One ballot's drawn lines, read once in a run however many parts of the loader ask for them."""
    if path not in _LINES:
        _LINES[path] = ballot_lines(path)
    return _LINES[path]


def contests(path, lines=None, where=None):
    """Every contest drawn on one county's sample ballot, with the date heading of the style it belongs to:
    [(date, title, [[name, party], ...], write-in line, vote for)]. Text below 2 points (a tagged copy of a heading
    that some counties' files carry) is not ballot text and is skipped. `lines`, when given, are the ballot's lines
    already read; `where`, when given, is filled with the place of each contest's "(Vote for" line among the lines of
    2 points and more."""
    L = [ln for ln in (ballot_lines(path) if lines is None else lines) if ln[2] >= 2]
    out, when, i = [], None, 0
    while i < len(L):
        t = L[i][3]
        m = DATE.match(t)
        if m:
            when = f"{m.group(3)}-{MONTHS.index(m.group(1)) + 1:02d}-{int(m.group(2)):02d}"
        if not t.startswith("FOR ") or VOTE_FOR.search(t):
            i += 1
            continue
        head, j = [t], i + 1
        while j < len(L) and not VOTE_FOR.match(L[j][3]) and j - i < 6:
            head.append(L[j][3])
            j += 1
        if j >= len(L) or not VOTE_FOR.match(L[j][3]):
            i += 1                                          # not a contest ("FOR PROPOSED TAXATION" and the like)
            continue
        vote_for = VOTE_FOR.match(L[j][3]).group(1).lower()
        if where is not None:
            where.append(j)
        cands, x0, size0 = [], None, None
        j += 1
        while j < len(L) and not L[j][3].startswith("Write-in"):
            x, _y, size, name = L[j]
            if x0 is None:
                x0, size0 = x, size
            elif abs(x - x0) > 1.5 or abs(size - size0) > 0.3:
                break
            if not NAME_LINE.match(name) or name.startswith("SAMPLE") or name.startswith("FOR ") or DATE.match(name):
                break
            party = None
            if j + 1 < len(L) and abs(L[j + 1][0] - x0) <= 1.5 and L[j + 1][2] < 0.8 * size0 and not L[j + 1][3].startswith("Write-in"):
                party = L[j + 1][3]
                j += 1
            cands.append([name, party])
            j += 1
        out.append((when, re.sub(r"\s+", " ", " ".join(head).replace(" ,", ",")).strip(), cands,
                    j < len(L) and L[j][3].startswith("Write-in"), vote_for))
        i = j
    return out


def manifest(folder, say):
    """[(county, file name, address)] for the 67 general election sample ballots: from the federal loader's record
    of them when it is there, else from the Secretary's page."""
    fed = os.path.join(folder, "al_2026_general_sample_ballots_federal.json")
    if os.path.exists(fed):
        data = json.load(open(fed, encoding="utf-8"))
        return [(f["county"], f["file"], f["url"], f.get("sha256")) for f in data["files"]], data.get("read")
    say("      reading the Secretary's list of sample ballots")
    return [(c, f, u, None) for c, _p, f, u in sample_links("general")], dt.date.today().isoformat()


def read_ballots(folder, keep, geoid_of, say):
    """The kept cells of every county's sample ballot (this loader's contests only, and counts of the rest), cached as
    JSON in `keep` by each file's SHA-256, so a ballot is read again only when its file changes."""
    files, listed = manifest(folder, say)
    if len(files) != COUNTIES:
        raise SystemExit(f"Alabama (state races): {len(files)} county sample ballots are listed, not {COUNTIES}")
    pdfs = os.path.join(folder, "sample_general")
    os.makedirs(pdfs, exist_ok=True)
    path = os.path.join(keep, KEPT)
    old = {}
    if os.path.exists(path):
        prev = json.load(open(path, encoding="utf-8"))
        if prev.get("version") == KEPT_VERSION:
            old = {f["sha256"]: f for f in prev["files"]}
    out, changed, warn = [], 0, []
    for county, file, url, fed_sha in files:
        pdf = os.path.join(pdfs, file)
        if not os.path.exists(pdf):
            fetch(url, pdf, b"%PDF-", f"the {county} County sample ballot", say, days=3650)
        sha = sha256(pdf)
        if fed_sha and sha != fed_sha:
            warn.append(f"{county} County's sample ballot on disk is not the one the federal loader read (it was replaced since)")
        if sha in old and old[sha]["file"] == file:
            out.append(old[sha])
            continue
        changed += 1
        kept, other, unknown = collections.defaultdict(int), collections.Counter(), collections.Counter()
        for when, title, cands, write_in, vote_for in contests(pdf, lines=lines_of(pdf)):
            key, g = classify(title)
            if key == "other":
                other[f"{g}|{when}"] += 1
                continue
            if key is None:
                unknown[when] += 1
                continue
            rid, info = race_of(key, g, geoid_of)
            if rid is None:
                unknown[when] += 1
                continue
            kept[json.dumps([when, rid, title, cands, bool(write_in), vote_for])] += 1
        out.append({"county": county, "file": file, "url": url, "sha256": sha, "bytes": os.path.getsize(pdf),
                    "contests": [dict(zip(("date", "race", "title", "cands", "write_in", "vote_for"), json.loads(k)), drawn=n)
                                 for k, n in sorted(kept.items())],
                    "other": dict(sorted(other.items())), "unknown": dict(sorted(unknown.items(), key=str))})
    if changed or not os.path.exists(path):
        os.makedirs(keep, exist_ok=True)
        json.dump({"version": KEPT_VERSION, "page": SAMPLE_PAGES["general"], "listed": listed, "read": dt.date.today().isoformat(),
                   "files": out}, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    say(f"    Alabama (state races): {len(out)} county sample ballots ({changed} read afresh, the rest from the kept cells)")
    return out, path, warn


def name_forms(name):
    """A name as printed, and with a quoted nickname in place of the given names (J.T. "Jabo" Waggoner -> Jabo Waggoner)."""
    out = [name]
    m = re.search(r"[\"“]([^\"”]+)[\"”]", name)
    if m:
        rest = (name[:m.start()] + name[m.end():]).split()
        family = [w for w in rest if fold(w) and fold(w) not in ("jr", "sr", "ii", "iii", "iv")]
        if family:
            out.append(f"{m.group(1)} {family[-1]}")
        out.append(re.sub(r"\s+", " ", name[:m.start()] + name[m.end():]).strip())
    return out


def alike(a, b):
    """Two printings of one person: the same letters, or the same family name and a given name (or nickname) that fits."""
    if same_person(a, b):
        return True
    return any(fits(name_parts(x), name_parts(y)) for x in name_forms(a) for y in name_forms(b))


def on_every_ballot(r):
    """An office every voter in the state votes for."""
    return (r["level"] == "statewide" and r["office_kind"] != "state_board_of_education") or r["office_kind"] in (
        "supreme_court", "court_of_civil_appeals", "court_of_criminal_appeals")


def agreed(files, geoid_of, problems, checks):
    """{race: ([[name, party], ...], write-in, {counties}, styles drawn)} for every race whose 2026 styles, in every
    county, print it the same way (letters compared: a space or a capital aside); the races whose counties disagree;
    styles of other dates set aside. Where counties space or capitalize a name differently, the printing of the most
    counties (then the most styles) is kept, and the check says so."""
    seen = collections.defaultdict(lambda: collections.defaultdict(lambda: collections.defaultdict(set)))
    styles = collections.Counter()
    counties, fields, older, drawn = collections.defaultdict(set), {}, collections.Counter(), collections.Counter()
    for f in files:
        geoid = geoid_of[county_key(f["county"])]
        for k in f["contests"]:
            if k["date"] != GENERAL:
                older[(f["county"], k["date"])] += k["drawn"]
                continue
            if k["vote_for"] != "one":
                problems.append(f"{k['race']}: {f['county']} County's ballot says 'Vote for {k['vote_for']}'")
            cands = []
            for name, printed in k["cands"]:
                if printed is not None:
                    if printed.strip().lower() not in PRINTED:
                        raise SystemExit(f"Alabama (state races): {f['county']} County's sample ballot prints a party this loader "
                                         f"does not know for {k['race']}")
                    printed = PRINTED[printed.strip().lower()]
                cands.append((name, printed))
            letters = (tuple((squash(n), p) for n, p in cands), bool(k["write_in"]))
            seen[k["race"]][letters][tuple(cands)].add(f["county"])
            styles[(k["race"], tuple(cands))] += k["drawn"]
            counties[k["race"]].add(geoid)
            drawn[k["race"]] += k["drawn"]
            if k["race"] not in fields:
                fields[k["race"]] = k["title"]
    races, split = {}, {}
    for race, versions in seen.items():
        if len(versions) > 1:
            split[race] = {"; ".join(f"{n} ({p})" for n, p in next(iter(v))): sorted(set().union(*v.values()))
                           for _letters, v in versions.items()}
            continue
        (letters, printings), = versions.items()
        best = max(printings, key=lambda c: (len(printings[c]), styles[(race, c)], c))
        if len(printings) > 1:
            names = sorted({n for c in printings for n, _p in c} - {n for n, _p in best})
            where = sorted(set().union(*(v for c, v in printings.items() if c != best)))
            checks.append(f"{race}: {'; '.join(names)} is spaced or capitalized differently on some styles in "
                          f"{', '.join(where)} (county ballots); the printing used by {len(printings[best])} "
                          f"{'county' if len(printings[best]) == 1 else 'counties'} is kept")
        races[race] = ([list(c) for c in best], letters[1], counties[race], drawn[race])
    return races, split, older


# ---------- the results ----------

def precinct_results(path, geoid_of, problems, variants, what):
    """This loader's contests in one precinct results zip: {(race, code): {candidate: votes}}, the over and under
    votes {(race, code, kind): votes}, the counties each race was voted in, the titles set aside by kind, and the
    counties read. A name printed with different punctuation or capitals in different county files is one candidate
    (the printing of the most files is shown), and `variants` says so."""
    tally = collections.defaultdict(lambda: collections.defaultdict(int))
    printed = collections.defaultdict(lambda: collections.defaultdict(collections.Counter))
    blanks = collections.defaultdict(int)
    where, other, counties = collections.defaultdict(set), collections.Counter(), []
    cast, by_county = {}, collections.defaultdict(int)
    with zipfile.ZipFile(path) as z:
        for n in sorted(z.namelist()):
            m = re.fullmatch(r"2026_PRIMARY(?:_RUNOFF)?_ELECTION-(.+)\.xls", n)
            if not m:
                raise SystemExit(f"Alabama (state races): an unexpected file in {os.path.basename(path)}: {n}")
            geoid = geoid_of.get(county_key(m.group(1)))
            if not geoid:
                raise SystemExit(f"Alabama (state races): {n}: the county could not be read")
            counties.append(geoid)
            for sheet, rows in xls_sheets(z.read(n)).items():
                if not rows:
                    continue
                if [str(c or "").strip() for c in rows[0][:3]] != ["Contest Title", "Party", "Candidate"]:
                    raise SystemExit(f"Alabama (state races): {n} ({sheet}) does not begin Contest Title, Party, Candidate")
                for r in rows[1:]:
                    title = re.sub(r"\s+", " ", str(r[0] or "")).strip()
                    if not title:
                        continue
                    m = re.fullmatch(r"BALLOTS CAST - (REPUBLICAN|DEMOCRAT) \((REP|DEM)\)", title)
                    if m:
                        cast[(geoid, m.group(2))] = number(sum(v for v in r[3:] if isinstance(v, float)))
                    key, g = classify(title)
                    if key == "other":
                        other[g] += 1
                        continue
                    if key is None:
                        other["not understood"] += 1
                        problems.append(f"{what}: a contest title in {n} was not understood (row {rows.index(r) + 1})")
                        continue
                    rid, info = race_of(key, g, geoid_of)
                    if rid is None:
                        problems.append(f"{what}: {n} row {rows.index(r) + 1}: {info}")
                        continue
                    code = str(r[1] or "").strip()
                    mark = re.search(r"\((REP|DEM)\)\s*$", title)
                    if code not in PARTY or (mark and mark.group(1) != code):
                        raise SystemExit(f"Alabama (state races): {n}: {rid} lists a candidate of an unexpected party")
                    cand = re.sub(r"\s+", " ", str(r[2] or "")).strip()
                    if any(v is not None and not isinstance(v, float) and str(v).strip() for v in r[3:]):
                        problems.append(f"{what}: {n}: a vote cell of {rid} is not a number (the row's other cells are counted)")
                    votes = number(sum(v for v in r[3:] if isinstance(v, float)))
                    where[rid].add(geoid)
                    if info["level"] == "statewide" and info["office_kind"] != "state_board_of_education":
                        by_county[(rid, code, geoid)] += votes
                    if cand in ("Over Votes", "Under Votes"):
                        blanks[(rid, code, cand)] += votes
                        continue
                    if not cand:
                        raise SystemExit(f"Alabama (state races): {n}: a row of {rid} has no candidate name")
                    k = fold(cand).replace(" ", "")                     # O'Hara-Grant in one county's file, O`hara-Grant in another's
                    tally[(rid, code)][k] += votes
                    printed[(rid, code)][k][cand] += 1
    totals = {}
    for key, per in tally.items():
        totals[key] = {}
        for k, votes in per.items():
            forms = printed[key][k]
            name = max(forms, key=lambda f: ("`" not in f, forms[f], f))    # a backtick for an apostrophe is a typing slip
            if len(forms) > 1:
                variants.append(f"{key[0]} ({PARTY_OF[key[1]]} {what}): {' / '.join(sorted(forms))} printed in different county files; "
                                f"counted as one, shown as {name}")
            totals[key][name] = votes
        names = sorted(totals[key])
        for i, a in enumerate(names):
            for b in names[i + 1:]:
                if alike(a, b):
                    variants.append(f"{key[0]} ({PARTY_OF[key[1]]} {what}): {a} and {b} may be one person printed two ways; kept apart")
    # a statewide contest's votes, with its over and under votes, against the party's ballots cast in each county
    diffs = [abs(n - cast[(g, c)]) for (rid, c, g), n in by_county.items() if (g, c) in cast]
    recon = (sum(1 for d in diffs if d == 0), len(diffs), max(diffs, default=0), len(by_county) - len(diffs))
    return totals, blanks, where, other, counties, recon


def dem_workbook(path, geoid_of):
    """The Democratic Party's certified totals (sheet Candidates: Office, District/Jurisdiction, Place, Ballot Name, a
    column per county, TOTAL) for this loader's offices: {(race, name): total}."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = wb["Candidates"].iter_rows(values_only=True)
    head = [str(c or "").strip() for c in next(rows)]
    if head[:4] != ["Office", "District/Jurisdiction", "Place", "Ballot Name"] or head[-1] != "TOTAL":
        raise SystemExit(f"Alabama (state races): {os.path.basename(path)} does not have the columns it had")
    out = {}
    for r in rows:
        office = str(r[0] or "").strip()
        if not office:
            continue
        district, place = (str(int(v)) if isinstance(v, (int, float)) else str(v or "").strip() for v in (r[1], r[2]))
        title = office + (f" District {district}" if district and not re.search(r"[A-Za-z]", district) else f" {district}")
        if place:
            title += f" Place {place}"
        key, g = classify(title)
        if key in (None, "other"):
            continue
        rid, _info = race_of(key, g, geoid_of)
        if rid:
            out[(rid, str(r[3] or "").strip())] = number(r[len(head) - 1])
    wb.close()
    return out


def rep_workbook(path, geoid_of):
    """The Republican Party's certified totals, from its Summary sheet (blocks: the office with "Votes" and
    "Percentage", a row per candidate, a Total row): {(race, name): total}."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out, race = {}, None
    for r in wb["Summary"].iter_rows(values_only=True):
        a = str(r[0] or "").strip() if r else ""
        b = str(r[1] or "").strip() if len(r) > 1 else ""
        if b == "Votes":
            key, g = classify(a)
            race = race_of(key, g, geoid_of)[0] if key not in (None, "other") else None
            continue
        if a == "Total":
            race = None
            continue
        if race and a:
            out[(race, a)] = number(b)
    wb.close()
    return out


def compare(totals, cert, code, stage):
    """Where a party's certified workbook differs from the Secretary's precinct files, in words; and how many of the
    Secretary's contests the workbook carries."""
    out, carried, of = [], 0, 0
    races = {r for r, _n in cert}
    for (race, c), cands in sorted(totals.items()):
        if c != code:
            continue
        of += 1
        if race not in races:
            continue
        carried += 1
        diffs, used = [], set()
        names = [n for r, n in cert if r == race]
        for name, votes in sorted(cands.items()):
            hit = ([n for n in names if fold(n) == fold(name)] or [n for n in names if alike(n, name)]
                   or [n for n in names if name_parts(n)[1] == name_parts(name)[1]])
            if len(hit) != 1:
                diffs.append(f"{name} is not in the workbook")
                continue
            used.add(hit[0])
            spelled = "" if fold(hit[0]) == fold(name) else f" (spelled {hit[0]} there)"
            theirs = cert[(race, hit[0])]
            if theirs != votes:
                diffs.append(f"{name}{spelled}: the workbook gives {'no figure' if theirs is None else format(theirs, ',')}, the precinct "
                             f"files {votes:,}")
            elif spelled:
                diffs.append(f"{name}{spelled}: the same votes")
        extra = [n for n in names if n not in used]
        if extra:
            diffs.append("the workbook also names " + ", ".join(extra))
        if diffs:
            out.append(f"{race} ({PARTY_OF[code]} {stage}): {'; '.join(diffs)}")
    return out, carried, of


# ---------- the roster ----------

def roster(path):
    """Sitting legislators {(chamber, district): [member]} and the officials {office: member}: ids, names, party only."""
    if not os.path.exists(path):
        return {}, {}, ""
    con = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    seats = collections.defaultdict(list)
    for r in con.execute("SELECT bioguide_id, first_name, last_name, official_full, other_names, party_name, chamber, district "
                         "FROM legislators WHERE is_current = 1"):
        seats[(r[6], str(r[7]).strip())].append({"id": r[0], "first": r[1], "last": r[2], "full": r[3], "other": r[4], "party": r[5]})
    offs = {}
    for r in con.execute("SELECT bioguide_id, first_name, last_name, official_full, party_name, office FROM officials"):
        offs[r[5]] = {"id": r[0], "first": r[1], "last": r[2], "full": r[3], "other": None, "party": r[4]}
    as_of = (con.execute("SELECT max(updated_at) FROM legislators").fetchone()[0] or "")[:10]
    con.close()
    return seats, offs, as_of


def member_fits(name, m):
    """The roster's own full name, or its given and family names (its "other names" are initials, too loose to use)."""
    forms = [m["full"], f"{m['first']} {m['last']}"]
    return any(fits(name_parts(x), name_parts(f)) for x in name_forms(name) for f in forms if f)


# ---------- sources ----------

def mdate(path):
    return dt.date.fromtimestamp(os.path.getmtime(path)).isoformat() if os.path.exists(path) else ""


def local_or_fetch(folder, key, say):
    """A results file from the cache the federal loader fills; fetched from the Secretary's site only if missing."""
    url, name = FILES[key]
    path = os.path.join(folder, name)
    if not os.path.exists(path):
        fetch(SOS + url, path, b"PK", f"the file {name}", say, days=3650)
    return path


# ---------- the county and local part: the same ballots' county and county school contests ----------

LOCAL_KEPT = "al_2026_general_sample_ballots_local.json"
LOCAL_KEPT_VERSION = 2
CAL_KEPT = "al_2026_calendar_findings.json"
CAL_KEPT_VERSION = 1
PUBLISHED = "2026-09-10"                                  # the day the Secretary of State posted the general election sample ballots
LOCAL_WHY = ("county office (local phase)", "county or city school office (local phase)", "district attorney (local phase)",
             "local question (local phase)")
SRC_BALLOT = "al-sos-2026-sl-sample-ballot-{}"            # one source per county's ballot, by county code
LOCAL_LEVELS = ("county", "school")
# office kinds the pages or other states' loaders already use; anything else this loader writes is named in its report
KNOWN_KINDS = {"county_commissioner", "sheriff", "county_attorney", "county_auditor", "county_treasurer", "county_auditor_treasurer",
               "county_recorder", "county_surveyor", "county_park", "soil_water", "mayor", "council", "city_clerk", "city_treasurer",
               "city_clerk_treasurer", "town_supervisor", "town_clerk", "town_treasurer", "town_clerk_treasurer", "school_board",
               "hospital_board", "utility_board", "sanitary_board", "coroner", "county_assessor", "tax_collector", "county_board_chair"}

# each office kind in a few plain words, for the coverage note's counts
KIND_WORDS = {"sheriff": "sheriff", "coroner": "coroner", "revenue_commissioner": "revenue commissioner", "county_assessor": "tax assessor",
              "tax_collector": "tax collector", "assistant_tax_assessor": "assistant tax assessor",
              "assistant_tax_collector": "assistant tax collector", "license_commissioner": "license commissioner",
              "racing_commissioner": "racing commissioner", "county_commissioner": "county commission seats",
              "county_board_chair": "county commission chairmen", "school_board": "county board of education seats",
              "school_superintendent": "county superintendents of education"}

NO_NAME_NOTE = "The county's sample ballot prints this contest with a Write-in line and no candidate's name."
SOME_STYLES = ("Not on every ballot in the county: the county's sample ballot draws {m} versions of the ballot, and this contest "
               "is on {n} of them.")
LOCAL_SPLIT_NOTE = ("The versions of the ballot in the county's sample ballot do not print this contest's candidates the same way, "
                    "so they are left out here.")
LOCAL_UNREAD_NOTE = "This contest's candidates could not be read with confidence from the county's sample ballot, so they are left out here."

_S = r"[\s,]*"                                            # between two words of a title: spaces, a comma, or nothing at all
_CN = r"(?P<c>[A-Z][A-Z\s.']*?)"                          # the county's name, as the title writes it
_N = r"(\d{1,2})"
_DIST = "DISTRICT" + _S + r"(?:NO\.?" + _S + ")?" + _N
_PLACE = "PLACE" + _S + r"(?:NO\.?" + _S + ")?" + _N


def _p(words):
    """A title's words as a pattern: any spacing between them (a ballot may run two words together), commas optional."""
    return _S.join(words.replace("{C}", _CN).split(" "))


# a county officer elected by the whole county: (title, office kind, office as shown)
OFFICE_PATS = [(re.compile(_p(w)), kind, office) for w, kind, office in (
    ("{C} COUNTY SHERIFF", "sheriff", "Sheriff"),
    ("{C} COUNTY CORONER", "coroner", "Coroner"),
    ("{C} COUNTY REVENUE COMMISS?IONER", "revenue_commissioner", "Revenue Commissioner"),
    ("{C} COUNTY PROPERTY TAX COMMISS?IONER", "revenue_commissioner", "Property Tax Commissioner"),
    ("{C} COUNTY TAX ASSESSOR / COLLECTOR", "revenue_commissioner", "Tax Assessor/Collector"),
    ("{C} COUNTY TAX ASSESSOR", "county_assessor", "Tax Assessor"),
    ("TAX ASSESSOR {C} COUNTY", "county_assessor", "Tax Assessor"),
    ("{C} COUNTY TAX COLLECTOR", "tax_collector", "Tax Collector"),
    ("TAX COLLECTOR {C} COUNTY", "tax_collector", "Tax Collector"),
    ("{C} COUNTY LICENSE COMMISS?IONER", "license_commissioner", "License Commissioner"),
    ("{C} COUNTY LICENSE DIRECTOR", "license_commissioner", "License Director"),
    ("{C} COUNTY RACING COMMISS?IONER", "racing_commissioner", "Racing Commissioner"))]
DIVISION_PAT = re.compile(_p("ASSISTANT TAX (?P<what>ASSESSOR|COLLECTOR) (?P<div>[A-Z][A-Z.'-]*) DIVISION OF {C} COUNTY"))
CHAIR_PAT = re.compile(_p("(?P<chair>CHAIRMAN|PRESIDENT) {C} COUNTY COMMISSION"))
SCHOOL_HEAD_PAT = re.compile(_p("(?P<head>SUPERINTENDENT|CHAIRMAN) {C} COUNTY BOARD OF EDUCATION"))
MEMBER_COMMISSION_PAT = re.compile(_p("MEMBER {C} COUNTY COMMISSION") + r"(?!ER)(?P<rest>.*)")
MEMBER_BOARD_PAT = re.compile(_p("MEMBER {C} COUNTY BOARD OF EDUCATION") + r"(?P<rest>.*)")
# offices a reader may look for that no November 2026 ballot carries: (words, what a title of that office would hold)
ABSENT_OFFICES = (("judge of probate", re.compile(r"PROBATE")), ("circuit clerk", re.compile(r"CIRCUITCLERK|CLERKOFTHECIRCUIT|CIRCUITCOURTCLERK")),
                  ("district attorney", re.compile(r"DISTRICTATTORNEY")), ("constable", re.compile(r"CONSTABLE")))


SMALL_PREFIX = re.compile(r"(Mc|Mac|De|Du|Di|La|Le|St\.)([A-Z][A-Z'-]*)(,?)")


def shown_local(name):
    """A county candidate's name in ordinary capitals, as shown() writes the state rows, but keeping what the ballot's own
    small letters say: a nickname in quotes or brackets exactly as printed ("DJ Skegee", "EB", (Danny)), initials
    written together (J.T.), and a prefix printed small (DeRAMUS-COLEMAN -> DeRamus-Coleman, St.CLAIR -> St.Clair)."""
    raw, out = name.split(), shown(name).split()
    if len(out) != len(raw):
        return shown(name)
    inside = False
    for i, w in enumerate(raw):
        body = w.rstrip(",")
        if body[:1] in "\"“(" and len(body) > 1:
            inside = True
        if inside:
            out[i] = w
            if body[-1:] in "\"”)" and len(body) > 1:
                inside = False
            continue
        m = SMALL_PREFIX.fullmatch(w)
        if m:
            out[i] = m.group(1) + "-".join(p.capitalize() for p in m.group(2).split("-")) + m.group(3)
        elif re.fullmatch(r"(?:[A-Z]\.){2,},?", w):
            out[i] = w
    return " ".join(out)


def slug(text):
    return re.sub(r"[^a-z0-9]+", "-", fold_keep_digits(text)).strip("-")


def fold_keep_digits(text):
    return re.sub(r"[^a-z0-9]+", " ", (text or "").lower()).strip()


def place_words(text):
    """A community's name printed in capitals, in ordinary capitals: ASHVILLE/STEELE -> Ashville/Steele."""
    return "/".join(proper(re.sub(r"\s+", " ", p).strip()) for p in text.split("/") if p.strip())


def _seat_words(rest, school):
    """The district, place or community words that follow a commission's or a board's name:
    {district, seat, tag[, associate]}, or None when they are not words this loader knows."""
    r = re.sub(r"\s+", " ", rest).strip(" ,")
    if not r:
        return dict(district=None, seat=None, tag="")
    m = re.fullmatch(_DIST, r)
    if m:
        return dict(district=str(int(m.group(1))), seat=None, tag=f"d{int(m.group(1))}")
    m = re.fullmatch(_DIST + _S + _PLACE, r)
    if m:
        return dict(district=str(int(m.group(1))), seat=f"Place {int(m.group(2))}", tag=f"d{int(m.group(1))}-p{int(m.group(2))}")
    if re.fullmatch(r"AT[\s-]*LARGE", r):
        return dict(district=None, seat="At Large", tag="at-large")
    if school:
        m = re.fullmatch(_PLACE + r"(?:\s*\(\s*(?P<comm>[A-Z][A-Z\s.'/-]*?)\s*\))?", r)
        if m:
            comm = m.group("comm")
            return dict(district=None, seat=f"Place {int(m.group(1))}" + (f" ({place_words(comm)})" if comm else ""), tag=f"p{int(m.group(1))}")
        m = re.fullmatch(r"(?P<comm>[A-Z][A-Z\s.'-]*?)\s*/\s*" + _DIST, r)
        if m:
            return dict(district=str(int(m.group(2))), seat=place_words(m.group("comm")), tag=f"d{int(m.group(2))}")
        return None
    m = re.fullmatch(_PLACE, r)
    if m:
        return dict(district=None, seat=f"Place {int(m.group(1))}", tag=f"p{int(m.group(1))}")
    m = re.fullmatch("ASSOCIATE" + _S + "COMMISS?IONER" + _S + _DIST, r)
    if m:
        return dict(district=str(int(m.group(1))), seat=None, tag=f"d{int(m.group(1))}", associate=True)
    return None


def local_office(title, county):
    """One county's contest title -> {level, office_kind, office, district, seat, tag}, or None when the title is not one
    of the shapes this loader knows or names another county (never guessed). `county` is the Census Bureau's name
    without "County" ("St. Clair"); the words are matched whatever their spacing (CLEBURNECOUNTY, BOARDOF EDUCATION)."""
    t = re.sub(r"\s+", " ", (title or "").upper()).strip()
    t = re.sub(r"^FOR\b[\s,]*", "", t)
    here = county_key(county)

    def mine(m):
        return county_key(m.group("c")) == here

    for pat, kind, office in OFFICE_PATS:
        m = pat.fullmatch(t)
        if m and mine(m):
            return dict(level="county", office_kind=kind, office=office, district=None, seat=None, tag="")
    m = DIVISION_PAT.fullmatch(t)
    if m and mine(m):
        what, div = m.group("what").title(), proper(m.group("div")) + " Division"
        return dict(level="county", office_kind=f"assistant_tax_{what.lower()}", office=f"Assistant Tax {what}", district=div, seat=None,
                    tag=slug(div))
    m = CHAIR_PAT.fullmatch(t)
    if m and mine(m):
        return dict(level="county", office_kind="county_board_chair", office=f"{m.group('chair').title()} of the County Commission",
                    district=None, seat=None, tag="")
    m = SCHOOL_HEAD_PAT.fullmatch(t)
    if m and mine(m):
        if m.group("head") == "SUPERINTENDENT":
            return dict(level="school", office_kind="school_superintendent", office="Superintendent of the County Board of Education",
                        district=None, seat=None, tag="")
        return dict(level="school", office_kind="school_board", office="Chairman of the County Board of Education", district=None,
                    seat=None, tag="chairman")
    for pat, school in ((MEMBER_COMMISSION_PAT, False), (MEMBER_BOARD_PAT, True)):
        m = pat.fullmatch(t)
        if m and mine(m):
            r = _seat_words(m.group("rest"), school)
            if r is None:
                return None
            if school:
                return dict(level="school", office_kind="school_board", office="Member of the County Board of Education", **r)
            office = "Member of the County Commission" + (" (Associate Commissioner)" if r.pop("associate", False) else "")
            return dict(level="county", office_kind="county_commissioner", office=office, **r)
    return None


def _title_above(L, i):
    """The lines drawn just above a "(Vote for" line that are its contest's title: the same size of type, five at most,
    stopping at the Write-in line of the contest before."""
    size, out, j = L[i][2], [], i - 1
    while j >= 0 and len(out) < 5 and abs(L[j][2] - size) <= 0.3 and not L[j][3].startswith("Write-in") and not VOTE_FOR.match(L[j][3]):
        out.insert(0, L[j][3])
        j -= 1
    return re.sub(r"\s+", " ", " ".join(out).replace(" ,", ",")).strip()


def _class_of(key, g):
    """Whose a contest is: "local" (kept here), "state" (the state part's), "federal", or what else it is."""
    if key is None:
        return "local"                                    # not understood: kept, so that it is named as a gap and never dropped
    if key != "other":
        return "state"
    if g in LOCAL_WHY:
        return "local"
    return "federal" if g.startswith("federal") else g


def read_local_ballots(folder, keep, say):
    """The kept cells of every county's sample ballot for the county and local part: each county or school contest
    (date of its style, title, names and parties in the order printed, the Write-in mark, how many styles draw it), how
    many versions of the ballot the file draws, and two plain counts of the whole ballot ("(Vote for" lines and party
    lines) against what the contests found account for. Cached as JSON in `keep` by each file's SHA-256, so a ballot
    is read again only when its file changes. A sample ballot has no address, telephone, e-mail or website on it; even
    so, a name cell that reads like a contact detail is blanked before it is kept, and counted."""
    files, listed = manifest(folder, say)
    if len(files) != COUNTIES:
        raise SystemExit(f"Alabama (county and school races): {len(files)} county sample ballots are listed, not {COUNTIES}")
    pdfs = os.path.join(folder, "sample_general")
    os.makedirs(pdfs, exist_ok=True)
    path = os.path.join(keep, LOCAL_KEPT)
    old = {}
    if os.path.exists(path):
        prev = json.load(open(path, encoding="utf-8"))
        if prev.get("version") == LOCAL_KEPT_VERSION:
            old = {f["sha256"]: f for f in prev["files"]}
    out, changed = [], 0
    for county, file, url, _fed_sha in files:
        pdf = os.path.join(pdfs, file)
        if not os.path.exists(pdf):
            fetch(url, pdf, b"%PDF-", f"the {county} County sample ballot", say, days=3650)
        sha = sha256(pdf)
        if sha in old and old[sha]["file"] == file:
            out.append(old[sha])
            continue
        changed += 1
        L = [ln for ln in lines_of(pdf) if ln[2] >= 2]
        where = []
        found = contests(pdf, lines=lines_of(pdf), where=where)
        kept, counts, styles, blanked = collections.defaultdict(int), collections.defaultdict(lambda: [0, 0]), collections.Counter(), 0
        for when, title, cands, write_in, vote_for in found:
            key, g = classify(title)
            if key == "GOV":
                styles[str(when)] += 1
            cls = _class_of(key, g)
            counts[f"{cls}|{when}"][0] += 1
            counts[f"{cls}|{when}"][1] += sum(1 for _n, p in cands if p is not None)
            if cls != "local":
                continue
            cells = []
            for name, party in cands:
                if contact_like(name, True) or CONTACT.search(name):
                    name, blanked = "", blanked + 1
                cells.append([name, party])
            kept[json.dumps([when, title, cells, bool(write_in), vote_for])] += 1
        # the second route: every "(Vote for" line and every party line on the ballot, counted plainly. The straight
        # party box at the head of each version names the parties too (ALABAMA / DEMOCRATIC / PARTY); those lines are
        # counted apart, from its heading to the first contest.
        claimed = set(where)
        without_for, loose_parties, party_lines, straight_lines, in_box = collections.Counter(), 0, 0, 0, False
        for i, ln in enumerate(L):
            if re.match(r"STRAIGHT\s*PARTY", ln[3], re.I):
                in_box = True
            elif ln[3].startswith("FOR ") or VOTE_FOR.match(ln[3]) or DATE.match(ln[3]):
                in_box = False
            if ln[3].strip().lower() in PRINTED:
                party_lines, straight_lines = party_lines + (not in_box), straight_lines + in_box
            if not VOTE_FOR.match(ln[3]) or i in claimed:
                continue
            title = _title_above(L, i)
            without_for[_class_of(*classify(title if title.startswith("FOR ") else "FOR " + title))] += 1
            j = i + 1
            while j < len(L) and j - i < 40 and not L[j][3].startswith("Write-in") and not VOTE_FOR.match(L[j][3]) and not L[j][3].startswith("FOR "):
                loose_parties += L[j][3].strip().lower() in PRINTED
                j += 1
        out.append({"county": county, "file": file, "url": url, "sha256": sha, "bytes": os.path.getsize(pdf),
                    "styles": dict(sorted(styles.items())), "vote_for_lines": sum(1 for ln in L if VOTE_FOR.match(ln[3])),
                    "contests_read": len(found), "party_lines": party_lines, "straight_party_lines": straight_lines,
                    "candidate_lines": sum(1 for c in found for _n, p in c[2] if p is not None),
                    "no_party": sum(1 for c in found for _n, p in c[2] if p is None),
                    "without_for": dict(sorted(without_for.items())), "without_for_parties": loose_parties,
                    "counts": {k: v for k, v in sorted(counts.items())}, "blanked": blanked,
                    "local": [dict(zip(("date", "title", "cands", "write_in", "vote_for"), json.loads(k)), drawn=n)
                              for k, n in sorted(kept.items())]})
    if changed or not os.path.exists(path):
        os.makedirs(keep, exist_ok=True)
        json.dump({"version": LOCAL_KEPT_VERSION, "page": SAMPLE_PAGES["general"], "listed": listed, "read": dt.date.today().isoformat(),
                   "files": out}, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    say(f"    Alabama (county and school races): {len(out)} county sample ballots ({changed} read afresh, the rest from the kept cells)")
    return out, path


# What the notes say about when local offices are elected rests on four documents of the Secretary of State. Each is
# fetched once, searched for the sentences below, and only the findings are kept (never the text).
CALENDARS = (
    ("muni2026", SOS + "/sites/default/files/election-2026/2026MunicipalElectionFCPACalendar.pdf", "al_2026_municipal_fcpa_calendar.pdf",
     "al-sos-2026-municipal-fcpa-calendar", "official calendar", "2026 Municipal Election FCPA Filing Calendar (Municipalities Election)",
     {"aug2026": r"general election [-–—] tuesday, august 25, 2026"},
     "Read for one line, the date of the 2026 municipal election (General Election, Tuesday, August 25, 2026)."),
    ("muni2025", SOS + "/sites/default/files/election-2025/FCPA%20Calendar%202025%20Municipal%20Election%20(ex%20Dothan%20and%20Tuscaloosa).pdf",
     "al_2025_municipal_fcpa_calendar.pdf", "al-sos-2025-municipal-fcpa-calendar", "official calendar",
     "2025 Municipal Election FCPA Filing Calendar (Municipalities Election)",
     {"aug2025": r"general election [-–—] tuesday, august 26, 2025"},
     "Read for one line, the date of the 2025 regular municipal elections (General Election, Tuesday, August 26, 2025)."),
    ("guide", SOS + "/sites/default/files/election-2026/2026CandidateFilingGuide.pdf", "al_2026_candidate_filing_guide.pdf",
     "al-sos-2026-candidate-filing-guide", "official guide", "2026 Candidate Filing Guide (FCPA Candidate Filing Guide, nineteenth edition)",
     {"act2021": r"act 2021-157 changes the date of regular municipal elections that were originally scheduled to occur in 2024 to 2025"},
     "Read for one sentence: Act 2021-157 moved the regular municipal elections scheduled for 2024 to 2025, except in municipalities "
     "whose election date is set by another statute."),
    ("admin", SOS + "/sites/default/files/election-2026/AdminCalendar%20-2026.pdf", "al_2026_admin_calendar.pdf",
     "al-sos-2026-administrative-calendar", "official calendar", "2026 Administrative Calendar: statewide primary, primary runoff and general election",
     {"general": r"general election [-–—] november 3, 2026",
      "probate": r"county party chairman must certify names of primary candidates for county office to",
      "nominees": r"for county offices to probate judge not later than noon",
      "cite_5": r"17-13-5\(b\)", "cite_18": r"17-13-18\(d\)",
      "withdraw": r"last day candidates can withdraw their name from ballot is 71 days before"},
     "Read for a few lines: the general election's date; that county party chairmen certify candidates for county office, and the "
     "parties their nominees for county office, to the judge of probate (Code of Alabama 17-13-5 and 17-13-18, as the calendar "
     "cites them); and the last day a candidate could withdraw from the ballot, 71 days before the general election."))


def calendar_findings(keep, say):
    """{key: True} for each sentence found in the Secretary's calendars and filing guide, and the documents read:
    [(source id, kind, title, address, published, fetched, sha256, sentences found, note)]. A document that cannot be
    fetched (two more tries at most) is left out, and the notes then say less. The findings are kept as JSON by each
    file's SHA-256; the files hold dates and rules, no candidate."""
    path = os.path.join(keep, CAL_KEPT)
    old = {}
    if os.path.exists(path):
        prev = json.load(open(path, encoding="utf-8"))
        if prev.get("version") == CAL_KEPT_VERSION:
            old = prev["files"]
    facts, docs, new, missing = {}, [], {}, []
    for key, url, name, sid, kind, title, wanted, note in CALENDARS:
        pdf = os.path.join(keep, name)
        try:
            net.download(url, pdf, max_age_days=3650, tries=3, say=say)
        except (HTTPError, URLError, OSError):
            missing.append(title)
            continue
        with open(pdf, "rb") as fh:
            head = fh.read(5)
        if head != b"%PDF-":                              # a notice or a challenge page instead of the document: not kept, not asked again this run
            os.remove(pdf)
            missing.append(title)
            continue
        sha = sha256(pdf)
        rec = old.get(name)
        if not rec or rec.get("sha256") != sha or set(rec.get("found", {})) != set(wanted):
            text = " ".join(re.sub(r"\s+", " ", t) for _pg, _y, t in pdf_lines(pdf)).lower()
            rev = re.search(r"revised (\d{1,2})/(\d{1,2})/(20\d\d)", text)
            rec = {"sha256": sha, "bytes": os.path.getsize(pdf), "found": {k: bool(re.search(p, text)) for k, p in wanted.items()},
                   "revised": f"{rev.group(3)}-{int(rev.group(1)):02d}-{int(rev.group(2)):02d}" if rev else ""}
        new[name] = rec
        facts.update({k: v for k, v in rec["found"].items() if v})
        docs.append((sid, kind, title, url, rec["revised"], mdate(pdf), sha, sum(rec["found"].values()), note
                     + ("" if all(rec["found"].values()) else " Not every line was found in the copy read; the notes say only what was found.")))
    if new != old or not os.path.exists(path):
        os.makedirs(keep, exist_ok=True)
        json.dump({"version": CAL_KEPT_VERSION, "files": new}, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return facts, docs, missing


def _all(n):
    return f"all {COUNTIES}" if n == COUNTIES else f"{n} of the {COUNTIES}"


def _list(words):
    words = list(words)
    return words[0] if len(words) == 1 else ", ".join(words[:-1]) + " or " + words[-1]


def local_part(files, counties, geoid_of, facts, docs, missing, folder):
    """The county and county school races of the November ballot, from the kept cells: the rows for sl_races,
    sl_candidates, sl_places, sl_sources, sl_gaps and sl_notes, and what to report."""
    races, cands, places, sources, gaps, checks = [], [], [], [], [], []
    stats = collections.Counter()
    by_level, by_kind, kind_counties, offices = collections.Counter(), collections.Counter(), collections.defaultdict(set), collections.Counter()
    older, titles26 = collections.Counter(), []
    reached = set()
    for f in files:
        geoid = geoid_of[county_key(f["county"])]
        cname = counties[geoid]                           # "St. Clair County", as the Census Bureau writes it
        bare = re.sub(r"\s+County$", "", cname)
        src = SRC_BALLOT.format(geoid)
        styles = int(f["styles"].get(GENERAL, 0))
        group, unread = {}, collections.Counter()
        for k in f["local"]:
            if k["date"] != GENERAL:
                older[str(k["date"])] += k["drawn"]
                stats["older_drawn"] += k["drawn"]
                continue
            stats["drawn"] += k["drawn"]
            stats["lines"] += k["drawn"] * len(k["cands"])
            titles26.append(squash(k["title"]))
            o = local_office(k["title"], bare)
            if o is None:
                unread[squash(k["title"])] += k["drawn"]
                stats["unread_drawn"] += k["drawn"]
                stats["unread_lines"] += k["drawn"] * len(k["cands"])
                continue
            if o["level"] == "school":
                jname = f"{cname} Board of Education"
                jid = f"{STATE}-S-{geoid[2:]}-{slug(jname)}"
                rid = f"2026-{jid}-{o['office_kind'].replace('_', '-')}"
            else:
                jname, jid = cname, geoid
                rid = f"2026-{STATE}-{geoid}-{o['office_kind'].replace('_', '-')}"
            if o["tag"]:
                rid += "-" + o["tag"]
            g = group.setdefault(rid, dict(o, jurisdiction=jname, jurisdiction_id=jid, drawn=0, lines=0, vote_for=set(), same=set(),
                                           versions=collections.defaultdict(collections.Counter)))
            g["same"].add((o["office"], o["district"], o["seat"]))
            g["drawn"] += k["drawn"]
            g["lines"] += k["drawn"] * len(k["cands"])
            g["vote_for"].add(k["vote_for"])
            letters = (tuple((squash(n), (p or "").strip().lower()) for n, p in k["cands"]), bool(k["write_in"]))
            g["versions"][letters][tuple((n, p) for n, p in k["cands"])] += k["drawn"]
        if unread:
            n = len(unread)
            gaps.append((STATE, "county", geoid, cname, "a county contest on the sample ballot" if n == 1 else "county contests on the sample ballot",
                         f"{cname}'s sample ballot prints " + ("a contest whose title" if n == 1 else f"{n} contests whose titles")
                         + " this loader does not know how to read, so " + ("it is" if n == 1 else "they are")
                         + " left out here rather than guessed; the county's sample ballot shows " + ("it." if n == 1 else "them."), f["url"]))
            checks.append(f"{cname}: {n} contest title{'s' if n != 1 else ''} not understood ({sum(unread.values())} drawings), left out and listed as a gap")
            stats["unread_titles"] += n
        loose = sum(n for cls, n in f["without_for"].items() if cls in ("local", "state"))
        if loose:
            gaps.append((STATE, "county", geoid, cname, "contests drawn without a heading this loader reads",
                         f"{cname}'s sample ballot draws {loose} contest" + ("" if loose == 1 else "s") + " whose title does not begin with the "
                         "word FOR, as every other contest's does, so " + ("it was" if loose == 1 else "they were") + " not read; the "
                         "county's sample ballot shows " + ("it." if loose == 1 else "them."), f["url"]))
            checks.append(f"{cname}: state or county contests drawn without FOR in the title and not read: {loose}")
        n_file = 0
        for rid, g in sorted(group.items()):
            notes, problem, rows = [], None, []
            if len(g["same"]) > 1:
                problem = "two different contests on the ballot come to one race"
            elif len(g["vote_for"]) > 1:
                problem = "the versions of the ballot disagree on how many to vote for"
            elif len(g["versions"]) > 1:
                problem = "split"
            else:
                (letters, printings), = g["versions"].items()
                # where the versions space a name differently (PAMELAB. for PAMELA B.), the printing drawn most is kept, then the fuller one
                best = max(printings, key=lambda c: (printings[c], sum(len(n) for n, _p in c), c))
                if len(printings) > 1:
                    stats["respaced"] += 1
                if not letters[1]:
                    checks.append(f"{rid}: no Write-in line on the sample ballot")
                seen_p, seen_n = set(), set()
                for order, (name, printed) in enumerate(best, start=1):
                    display = shown_local(name) if name else ""
                    party = PRINTED.get((printed or "").strip().lower())
                    if not display or contact_like(display, True) or CONTACT.search(display):
                        problem, stats["dropped_names"] = "a name cell that does not hold a name", stats["dropped_names"] + 1
                    elif not party:
                        problem = "a candidate with no party, or a party this loader does not know, beneath the name"
                    elif (party in seen_p and party != "Independent") or display in seen_n:
                        problem = "two candidates of one party, or one name twice"
                    if problem:
                        break
                    seen_p.add(party)
                    seen_n.add(display)
                    rows.append([rid, "general", GENERAL, display, party, party_code(party), order, 0, 0, None, None, None, None, src, CAPS])
                if problem:
                    rows = []
            if problem == "split":
                notes.append(LOCAL_SPLIT_NOTE)
            elif problem:
                notes.append(LOCAL_UNREAD_NOTE)
            elif not rows:
                notes.append(NO_NAME_NOTE)
                stats["no_name"] += 1
            if problem:
                gaps.append((STATE, "race", rid, f"{g['jurisdiction']}: {g['office']}", "the candidates",
                             ("The versions of the ballot drawn in the county's sample ballot do not print this contest's candidates the "
                              "same way" if problem == "split" else "The county's sample ballot could not be read with confidence for this "
                              "contest (" + problem + ")") + ", so the race is shown without candidates; the county's sample ballot "
                             "shows them.", f["url"]))
                checks.append(f"{rid}: {problem}; candidates left out ({g['lines']} candidate lines drawn)")
                stats["problem_races"] += 1
                stats["problem_lines"] += g["lines"]
            word = next(iter(g["vote_for"])) if len(g["vote_for"]) == 1 else None
            if word and word != "one":
                notes.append(f"Voters choose {word}.")
            if styles and g["drawn"] < styles:
                notes.append(SOME_STYLES.format(n=g["drawn"], m=styles))
            elif g["drawn"] > styles:
                checks.append(f"{rid}: drawn {g['drawn']} times on a ballot of {styles} versions")
            races.append((rid, STATE, g["level"], g["office_kind"], g["office"], g["jurisdiction"], g["jurisdiction_id"], json.dumps([geoid]),
                          g["district"], g["seat"], 0, 1, None, None, None, GENERAL, " ".join(notes) or None))
            cands += rows
            n_file += len(rows)
            reached.add(geoid)
            by_level[g["level"]] += 1
            by_kind[g["office_kind"]] += 1
            offices[(g["office_kind"], g["office"])] += 1
            kind_counties[g["office_kind"]].add(geoid)
            if not problem:
                stats["placed_lines"] += g["lines"]
            if g["level"] == "school" and not any(p[1] == g["jurisdiction_id"] for p in places):
                places.append(("school", g["jurisdiction_id"], g["jurisdiction"], json.dumps([geoid]), src))
        # the county's ballot as a source, with the two counts side by side
        f26 = {cls: f["counts"].get(f"{cls}|{GENERAL}", [0, 0]) for cls in ("federal", "state", "local")}
        rest26 = [v for k, v in f["counts"].items() if k.endswith("|" + GENERAL) and k.split("|")[0] not in f26]
        old_n = sum(v[0] for k, v in f["counts"].items() if not k.endswith("|" + GENERAL))
        wf = sum(f["without_for"].values())
        ok_votes = f["vote_for_lines"] == f["contests_read"] + wf
        ok_party = f["party_lines"] == f["candidate_lines"] + f["without_for_parties"]
        stats["vote_for_lines"] += f["vote_for_lines"]
        stats["contests_read"] += f["contests_read"]
        stats["party_lines"] += f["party_lines"]
        stats["straight_party_lines"] += f.get("straight_party_lines", 0)
        stats["candidate_lines"] += f["candidate_lines"]
        stats["without_for"] += wf
        stats["without_for_parties"] += f["without_for_parties"]
        stats["blanked"] += f.get("blanked", 0)
        stats["no_party"] += f.get("no_party", 0)
        if not ok_votes or not ok_party:
            stats["unreconciled"] += 1
            checks.append(f"{cname}: the ballot's own lines and the contests read do not add up ({f['vote_for_lines']} vote-for lines "
                          f"against {f['contests_read'] + wf} contests; {f['party_lines']} party lines against "
                          f"{f['candidate_lines'] + f['without_for_parties']} candidates read)")
        for cls, n in sorted(f["without_for"].items()):
            checks.append(f"{cname}: {n} {cls.split(' (')[0]} contest{'s' if n != 1 else ''} drawn without the word FOR in the title "
                          "(counted, " + ("not this loader's" if cls not in ("local", "state") else "not read") + ")")
        pdf = os.path.join(folder, "sample_general", f["file"])
        sources.append((src, STATE, "official sample ballot", AGENCY, f"2026 General Election Sample Ballot, {cname} (county and school offices)",
                        f["url"], PUBLISHED, mdate(pdf), f["sha256"], n_file,
                        f"{cname}'s sample ballot (PDF), which draws {styles} version{'' if styles == 1 else 's'} of the November 3 ballot one "
                        "over another" + (f" and {old_n} contests of an older ballot beneath, set aside" if old_n else "") + ". Read here: the "
                        "title of each county and county school contest, the names under it in the order printed, the party beneath each name "
                        f"and the Write-in line: {len(group)} contests, {n_file} candidates, each version printing a contest the same way"
                        + ("" if not any(len(g["versions"]) > 1 for g in group.values()) else " except where a gap says otherwise")
                        + f". Two counts of the whole ballot agree: its {f['vote_for_lines']:,} vote-for lines head {f26['federal'][0]:,} "
                        f"federal, {f26['state'][0]:,} state and {f26['local'][0]:,} county and school contests"
                        + (f", {sum(v[0] for v in rest26):,} other contests" if rest26 else "") + (f", {old_n:,} on the older ballot" if old_n else "")
                        + (f" and {wf} drawn without the word FOR" if wf else "") + f", and its {f['party_lines']:,} party lines are the "
                        f"{f['candidate_lines'] + f['without_for_parties']:,} candidates under them"
                        + ("" if ok_votes and ok_party else " (they do not add up; see the loader's report)")
                        + ". A sample ballot carries no address, telephone, e-mail or website, shows no one who withdrew and names no "
                        "write-in candidate. Names printed in capitals are shown in ordinary capitals."))

    # the notes: when Alabama elects its local officers, and what is loaded here
    def nc(*kinds):
        return len(set().union(*(kind_counties.get(k, set()) for k in kinds)))

    absent = [words for words, pat in ABSENT_OFFICES if not any(pat.search(t) for t in titles26)]
    calendar = (f"On November 3, 2026 {_all(nc('sheriff'))} Alabama counties elect a sheriff, {nc('coroner')} a coroner and "
                f"{nc('revenue_commissioner', 'county_assessor', 'tax_collector')} a revenue commissioner or a tax assessor and a tax collector; "
                f"{nc('county_commissioner', 'county_board_chair')} elect members of the county commission, {nc('school_board')} members of the "
                f"county board of education and {nc('school_superintendent')} the county superintendent of education.")
    if absent:
        calendar += f" No county's November ballot has a contest for {_list(absent)}."
    days = [d for k, d in (("aug2025", "August 26, 2025 for the regular municipal elections"), ("aug2026", "August 25, 2026 for those held this year"))
            if facts.get(k)]
    if len(days) == 2:
        calendar += (" Cities and towns vote at their own municipal elections, not in November: the Secretary of State's calendars "
                     f"give {days[0]}" + (", which Act 2021-157 moved from 2024," if facts.get("act2021") else "") + f" and {days[1]}.")
    else:
        checks.append("the Secretary's municipal election calendars could not be read, so the calendar note says nothing about cities and towns")
    kinds_line = ", ".join(f"{KIND_WORDS.get(k, k.replace('_', ' '))} {n}" for k, n in sorted(by_kind.items(), key=lambda x: (-x[1], x[0])))
    coverage = (f"Loaded from the {len(files)} county sample ballots the Secretary of State posted for November 3, 2026: every county office and "
                f"county board of education contest printed on them, {len(races):,} contests and {len(cands):,} candidates in "
                f"{_all(len(reached))} counties ({kinds_line}). Every one is a partisan office, and names, parties and ballot order are as "
                "printed. A sample ballot shows no one who withdrew and names no write-in candidates"
                + ("; candidates for county office are certified to each county's judge of probate rather than to the Secretary of State, so "
                   "these ballots are the statewide official source for them"
                   if all(facts.get(k) for k in ("probate", "nominees", "cite_5", "cite_18")) else "")
                + ". Not loaded: constitutional amendments and other ballot questions, the May 19 primary and June 16 runoff for county "
                "offices, and city and town offices, which are not on this ballot."
                + ((" One contest is" if stats["no_name"] == 1 else f" {stats['no_name']} contests are") + " printed with a Write-in line "
                   "and no name." if stats["no_name"] else "")
                + (f" {len(gaps)} gap{'' if len(gaps) == 1 else 's'} list what could not be read." if gaps else ""))
    cal_docs = [d[2].split(" (")[0] for d in docs if d[0] != "al-sos-2026-administrative-calendar"]
    notes = [(STATE, "local_calendar", calendar,
              f"{AGENCY}: the {len(files)} county sample ballots for November 3, 2026"
              + ("; " + "; ".join(cal_docs) if cal_docs else ""), INFO_PAGE),
             (STATE, "local_coverage", coverage, f"{AGENCY}: 2026 General Election Sample Ballots, one for each county"
              + ("; 2026 Administrative Calendar" if facts.get("probate") else ""), SAMPLE_PAGES["general"])]
    for sid, kind, title, url, revised, fetched, sha, found, note in docs:
        sources.append((sid, STATE, kind, AGENCY, title, url, revised, fetched, sha, found, note + " The document holds dates and rules; "
                        "no candidate is named in it."))
    for title in missing:
        checks.append(f"could not be fetched, so the notes rest on less: {title}")

    # the last look at every text the local part stores: nothing that reads like a contact detail
    for row in races:
        if any(v and (contact_like(v, True) or LOOKS_LIKE_STREET.search(str(v))) for v in (row[4], row[5], row[8], row[9], row[16])):
            raise SystemExit(f"Alabama (county and school races): a stored cell of {row[0]} failed the contact-detail check (not shown)")
    for c in cands:
        if contact_like(c[3], True) or CONTACT.search(c[3]) or contact_like(c[14] or "", True):
            raise SystemExit(f"Alabama (county and school races): a stored name in {c[0]} failed the contact-detail check (not shown)")
    for p in places:
        if contact_like(p[2], True):
            raise SystemExit(f"Alabama (county and school races): the name of place {p[1]} failed the contact-detail check (not shown)")
    for g in gaps:
        if any(contact_like(v, True) or LOOKS_LIKE_STREET.search(v) for v in (g[3], g[4], g[5]) if v):
            raise SystemExit(f"Alabama (county and school races): a gap's words for {g[2]} failed the contact-detail check (not shown)")
    for n in notes:
        if contact_like(n[2], True) or LOOKS_LIKE_STREET.search(n[2]) or contact_like(n[3], False):
            raise SystemExit(f"Alabama (county and school races): the note {n[1]} failed the contact-detail check (not shown)")
    keys = collections.Counter((c[0], c[1], c[3]) for c in cands)
    if any(n > 1 for n in keys.values()) or len({r[0] for r in races}) != len(races):
        raise SystemExit("Alabama (county and school races): two rows share a key; stopping")

    stats["races"], stats["cands"], stats["reached"], stats["files"] = len(races), len(cands), len(reached), len(files)
    return dict(races=races, cands=cands, places=places, sources=sources, gaps=gaps, notes=notes, checks=checks, stats=dict(stats),
                by_level=dict(by_level), by_kind=dict(by_kind), offices={f"{k}: {o}": n for (k, o), n in sorted(offices.items())},
                older=dict(older), new_kinds=sorted(set(by_kind) - KNOWN_KINDS), absent=absent,
                contested=sum(1 for n in collections.Counter(c[0] for c in cands).values() if n > 1),
                single=sum(1 for n in collections.Counter(c[0] for c in cands).values() if n == 1))


# ---------- the load ----------

SCHEMA = """
CREATE TABLE IF NOT EXISTS sl_races (race_id TEXT PRIMARY KEY, state TEXT NOT NULL, level TEXT NOT NULL, office_kind TEXT NOT NULL,
  office TEXT NOT NULL, jurisdiction TEXT, jurisdiction_id TEXT, county_ids TEXT, district TEXT, seat TEXT, special INTEGER NOT NULL DEFAULT 0,
  partisan INTEGER NOT NULL, holder_id TEXT, holder_name TEXT, holder_party TEXT, election_date TEXT NOT NULL, note TEXT);
CREATE TABLE IF NOT EXISTS sl_candidates (race_id TEXT NOT NULL, election TEXT NOT NULL, election_date TEXT, name TEXT NOT NULL, party TEXT,
  party_code TEXT, ballot_order INTEGER, incumbent INTEGER NOT NULL DEFAULT 0, write_in INTEGER NOT NULL DEFAULT 0, votes INTEGER, pct REAL,
  outcome TEXT, state_member_id TEXT, source_id TEXT, note TEXT, PRIMARY KEY (race_id, election, name));
CREATE TABLE IF NOT EXISTS sl_sources (source_id TEXT PRIMARY KEY, state TEXT, kind TEXT, agency TEXT, title TEXT, url TEXT, published TEXT,
  fetched TEXT, sha256 TEXT, rows INTEGER, note TEXT);
CREATE TABLE IF NOT EXISTS sl_places (kind TEXT NOT NULL, id TEXT NOT NULL, name TEXT NOT NULL, county_ids TEXT, source_id TEXT,
  PRIMARY KEY (kind, id));
"""


def load(db_path, say=print, folder=FOLDER, roster_db=ROSTER, county_zip=COUNTY_ZIP, keep=None, local_keep=None):
    keep = keep or folder
    local_keep = local_keep or os.path.join(folder, "local")
    net.patient_lookups()
    problems, checks = [], []
    counties = census_counties(county_zip)
    geoid_of = {county_key(n.replace(" County", "")): g for g, n in counties.items()}
    geoid_of["_name"] = counties

    # 1. the November ballot
    files, kept_path, warn = read_ballots(folder, keep, geoid_of, say)
    problems += warn
    general, split, older = agreed(files, geoid_of, problems, checks)
    other_general = collections.Counter()
    unknown_general = 0
    for f in files:
        for k, n in f["other"].items():
            why, when = k.rsplit("|", 1)
            if when == GENERAL:
                other_general[why] += n
        unknown_general += sum(n for when, n in f["unknown"].items() if when == GENERAL)
    if unknown_general:
        problems.append(f"{unknown_general} contests drawn on the November styles were not understood (not loaded)")
    if other_general.get(COURT_UNREAD):
        problems.append(f"{other_general[COURT_UNREAD]} court contests drawn on the November styles were not understood (not loaded)")

    races = {}

    def add_race(rid, info, why=None):
        if rid in races:
            return races[rid]
        r = dict(info, race_id=rid, note=[], county_set=set(), prim_set=set(), holder=None)
        races[rid] = r
        return r

    for d in range(1, SENATE_SEATS + 1):
        add_race(*race_of("SS", {"a": str(d)}, geoid_of))
    for d in range(1, HOUSE_SEATS + 1):
        add_race(*race_of("SH", {"a": str(d)}, geoid_of))
    titles = {k["race"]: k["title"] for f in files for k in f["contests"]}
    for rid in sorted(set(general) | set(split)):
        added = race_of(*classify(titles[rid]), geoid_of)
        if added[0] != rid:
            raise SystemExit(f"Alabama (state races): the kept title of {rid} no longer gives the same race")
        add_race(*added)

    # 2. who holds each seat
    seats, offs, as_of = roster(roster_db)
    if not seats:
        problems.append("the roster (state_al.sqlite) is missing; no holders or incumbents are marked")
    for rid, r in races.items():
        h = None
        if r["office_kind"] in ("state_senate", "state_house"):
            r["note"].append(LEG_NOTE)
            hs = seats.get(("Senate" if r["office_kind"] == "state_senate" else "House", r["district"]), [])
            if len(hs) == 1:
                h = hs[0]
            elif not hs and seats:
                r["note"].append(VACANT)
                checks.append(f"{rid}: no sitting member in the roster")
            elif hs:
                checks.append(f"{rid}: the roster lists {len(hs)} sitting members for one seat; none is shown")
        else:
            key = re.match(rf"2026-{STATE}-([A-Z]+)", rid).group(1)
            office = EXEC.get(key, (None, None, None, None))[3]
            h = offs.get(office) if office else None
            if not h:
                r["note"].append(NO_HOLDER)
        r["holder"] = h

    # 3. the November candidates
    cands = []                 # [race, election, date, name, party, code, order, incumbent, write_in, votes, pct, outcome, member, source, note]
    nominees = collections.defaultdict(dict)
    for rid in sorted(split):
        races[rid]["note"].append(SPLIT_NOTE)
        problems.append(f"{rid}: the counties' sample ballots disagree: {split[rid]}")
    for rid, (clist, write_in, where, _drawn) in sorted(general.items()):
        r = races[rid]
        r["county_set"] |= where
        if not write_in:
            problems.append(f"{rid}: no Write-in line on the sample ballots")
        seen_parties = set()
        for order, (name, printed) in enumerate(clist, start=1):
            if not printed:
                problems.append(f"{rid}: a candidate has no party printed beneath the name (not loaded)")
                continue
            if printed in seen_parties and printed != "Independent":         # a party nominates one; independents qualify each on their own
                raise SystemExit(f"Alabama (state races): {rid}: two {printed} candidates on the sample ballots")
            seen_parties.add(printed)
            display = shown(name)
            if printed != "Independent":
                nominees[rid][printed] = display
            cands.append([rid, "general", GENERAL, display, printed, party_code(printed), order, 0, 0, None, None, None, None, SRC_GENERAL, CAPS])
        if on_every_ballot(r):
            if len(where) != COUNTIES:
                problems.append(f"{rid}: printed on {len(where)} of {COUNTIES} counties' sample ballots")
    set_aside = collections.Counter()
    for (county, date), n in older.items():
        set_aside[date] += n
        if date is None:
            problems.append(f"{county} County's sample ballot draws a contest before any date heading")

    # 4. the primary and the runoff
    paths = {k: local_or_fetch(folder, k, say) for k in ("primary", "runoff", "dem_primary", "rep_primary", "dem_runoff")}
    variants = []
    ptot, pblank, pwhere, pother, pcounties, precon = precinct_results(paths["primary"], geoid_of, problems, variants, "primary")
    rtot, rblank, rwhere, rother, rcounties, rrecon = precinct_results(paths["runoff"], geoid_of, problems, variants, "runoff")
    checks += variants

    def settle(tot, where, what):
        """A district court contest whose results title names no place (DISTRICT COURT JUDGE, PICKENS COUNTY) belongs to
        the county's one district court race on the November ballot, when there is exactly one."""
        out = {}
        for (rid, code), cc in tot.items():
            m = re.fullmatch(rf"2026-{STATE}-DC(\d{{3}})", rid)
            if rid not in races and m:
                one = [r for r in races if r.startswith(f"2026-{STATE}-DC{m.group(1)}-")]
                if len(one) == 1:
                    checks.append(f"{rid} ({PARTY_OF[code]} {what}): the results title names no place; filed under {one[0]}, the "
                                  "county's one district court race on the November ballot")
                    where[one[0]] = where.get(one[0], set()) | where.pop(rid, set())
                    rid = one[0]
            out[(rid, code)] = cc
        return out

    ptot, rtot = settle(ptot, pwhere, "primary"), settle(rtot, rwhere, "runoff")
    fields = collections.Counter()
    orphan = []

    def nominee(rid, code):
        return nominees.get(rid, {}).get(PARTY[code])

    def elsewhere(name, rid):
        """Other races of the same office whose November ballot prints this person."""
        kind = races[rid]["office_kind"]
        return sorted({c[0] for c in cands if c[1] == "general" and c[0] != rid and races[c[0]]["office_kind"] == kind and alike(name, c[3])})

    def not_nominee(rid, code, winner, final, stage):
        other = elsewhere(winner, rid)
        where_now = " and ".join(races[o]["jurisdiction"] for o in other)
        problems.append(f"{rid} {code}: the {stage}'s winner ({winner}) is not the party's candidate on the sample ballots ({final})"
                        + (f"; the ballots print {winner} for {where_now}" if other else ""))
        races[rid]["note"].append(f"The {PARTY_OF[code]} candidate on the November ballot, {final}, is not the winner of the party's "
                                  f"{stage}; the files read here do not say why (the parties' certifications of their nominees are "
                                  "scanned documents).")
        return (f"Won the {PARTY_OF[code]} {stage} for this seat; the November sample ballots print this candidate for {where_now} "
                "instead.") if other else NOT_ON_BALLOT.format(party=PARTY_OF[code])

    for (rid, code), cc in sorted(ptot.items()):
        if rid not in races:
            orphan.append(f"{rid} ({PARTY_OF[code]} primary)")
            continue
        races[rid]["prim_set"] |= pwhere[rid]
        if len(cc) < 2:
            continue
        fields["primary"] += 1
        total = sum(cc.values())
        ranked = sorted(cc, key=lambda n: (-cc[n], n))
        runoff = 2 * cc[ranked[0]] <= total
        won = ranked[:2] if runoff else ranked[:1]
        final = nominee(rid, code)
        held = (rid, code) in rtot
        win_note = None
        if runoff and held:
            pair = rtot[(rid, code)]
            if len(pair) != 2 or not all(any(alike(a, b) for b in pair) for a in won):
                problems.append(f"{rid} {code}: the runoff lists {sorted(pair)}, the primary's top two were {won}")
        elif runoff:
            checks.append(f"{rid} {code}: no majority on May 19 and no runoff in the runoff results "
                          f"(the November ballot prints {final or 'no one'} for the party)")
        elif final and not alike(won[0], final):
            win_note = not_nominee(rid, code, won[0], final, "primary")
        elif not final and rid in general:
            win_note = NOT_ON_BALLOT.format(party=PARTY_OF[code])
        for name in ranked:
            note = None
            if runoff and name in won:
                note = RUNOFF_NOTE if held else NO_RUNOFF
            elif name in won:
                note = win_note
            cands.append([rid, f"primary-{code}", PRIMARY, name, PARTY[code], party_code(PARTY[code]), None, 0, 0, cc[name],
                          round(100 * cc[name] / total, 1) if total else None, "advanced" if name in won else "lost", None, SRC_PRIMARY, note])
    for (rid, code), cc in sorted(rtot.items()):
        if rid not in races:
            orphan.append(f"{rid} ({PARTY_OF[code]} runoff)")
            continue
        races[rid]["prim_set"] |= rwhere[rid]
        if len(cc) < 2:
            continue
        fields["runoff"] += 1
        total = sum(cc.values())
        ranked = sorted(cc, key=lambda n: (-cc[n], n))
        if 2 * cc[ranked[0]] == total:
            problems.append(f"{rid} {code}: the runoff is tied")
        final = nominee(rid, code)
        if (rid, code) not in ptot:
            problems.append(f"{rid} {code}: a runoff with no primary in the May 19 results")
        win_note = None
        if final and not alike(ranked[0], final):
            win_note = not_nominee(rid, code, ranked[0], final, "runoff")
        elif not final and rid in general:
            win_note = NOT_ON_BALLOT.format(party=PARTY_OF[code])
        for name in ranked:
            cands.append([rid, f"runoff-{code}", RUNOFF, name, PARTY[code], party_code(PARTY[code]), None, 0, 0, cc[name],
                          round(100 * cc[name] / total, 1) if total else None, "advanced" if name == ranked[0] else "lost", None, SRC_RUNOFF,
                          win_note if name == ranked[0] else None])
    for o in orphan:
        problems.append(f"{o}: in the results but not printed on any November sample ballot (not stored)")

    # the parties' certified workbooks, as a check on the precinct files
    dem_p, rep_p, dem_r = (dem_workbook(paths["dem_primary"], geoid_of), rep_workbook(paths["rep_primary"], geoid_of),
                           dem_workbook(paths["dem_runoff"], geoid_of))
    chk = {"dem_primary": compare(ptot, dem_p, "DEM", "primary"), "rep_primary": compare(ptot, rep_p, "REP", "primary"),
           "dem_runoff": compare(rtot, dem_r, "DEM", "runoff")}

    # 5. incumbents: a candidate who fits the sitting member, one fit only in each election of the race
    by_election = collections.defaultdict(list)
    for c in cands:
        by_election[(c[0], c[1])].append(c)
    for (rid, _e), rows in by_election.items():
        h = races[rid]["holder"]
        if not h:
            continue
        hit = [c for c in rows if member_fits(c[3], h)]
        if len(hit) == 1:
            hit[0][7], hit[0][12] = 1, h["id"]
        elif len(hit) > 1:
            checks.append(f"{rid} {_e}: more than one name fits the sitting member; none is marked")

    # 6. every seat has its candidates, or says why not
    for rid, r in races.items():
        if rid in split:
            continue
        if not any(c[0] == rid and c[1] == "general" for c in cands):
            r["note"].append(NOT_PRINTED)
            checks.append(f"{rid}: no candidate printed on any November sample ballot")
        if on_every_ballot(r):
            r["county_set"] = set(counties)
            continue
        if r["county_set"] and r["prim_set"] and r["prim_set"] != r["county_set"]:
            name = lambda gs: ", ".join(counties[g].replace(" County", "") for g in sorted(gs))
            checks.append(f"{rid}: the primary was voted in {name(r['prim_set'])}; the November ballots print the seat in "
                          f"{name(r['county_set'])}")
            r["note"].append(f"The party primaries for this seat were voted in {name(r['prim_set'])} "
                             f"{'County' if len(r['prim_set']) == 1 else 'counties'}; the November sample ballots print it in "
                             f"{name(r['county_set'])} {'County' if len(r['county_set']) == 1 else 'counties'}.")
        if r["county_set"]:
            r["note"].append(DERIVED)
        elif r["prim_set"]:
            r["county_set"] = set(r["prim_set"])
            r["note"].append(DERIVED_RESULTS)

    # 7. the last look at every stored text: nothing that looks like a contact detail reaches the database
    for c in cands:
        if CONTACT.search(c[3]) or CONTACT.search(c[4] or ""):
            raise SystemExit(f"Alabama (state races): a stored name for {c[0]} failed the contact-detail check (not shown)")
    keys = [(c[0], c[1], c[3]) for c in cands]
    dup = [k for k, n in collections.Counter(keys).items() if n > 1]
    if dup:
        raise SystemExit(f"Alabama (state races): two candidate rows share race, election and name: {dup[:3]}")

    # 8. places
    place_rows = [("county", g, n, json.dumps([g]), SRC_COUNTY) for g, n in sorted(counties.items())]
    for rid, r in sorted(races.items()):
        kind = {"state_senate": "senate", "state_house": "house", "circuit_court": "judicial", "state_board_of_education": "sboe"}.get(r["office_kind"])
        if kind and r["county_set"]:
            pid = f"{STATE}-{r['district']}" if kind in ("senate", "house", "sboe") else f"{STATE}-{r['jurisdiction_id']}"
            if not any(p[0] == kind and p[1] == pid for p in place_rows):
                place_rows.append((kind, pid, r["jurisdiction"], json.dumps(sorted(r["county_set"])), SRC_GENERAL))

    # 9. write Alabama's rows only, in one transaction
    race_rows = []
    for rid, r in sorted(races.items()):
        h = r["holder"]
        race_rows.append((rid, STATE, r["level"], r["office_kind"], r["office"], r["jurisdiction"], r["jurisdiction_id"],
                          json.dumps(sorted(r["county_set"])) if r["county_set"] else None, r["district"], r["seat"], 0, 1,
                          h["id"] if h else None, h["full"] if h else None, h["party"] if h else None, GENERAL,
                          " ".join(r["note"]) or None))
    gen_rows = [c for c in cands if c[1] == "general"]
    older_note = (", ".join(f"{n} contests dated {d}" for d, n in sorted(set_aside.items(), key=str))
                  + " (in " + ", ".join(sorted({c for c, _d in older})) + ")") if older else "none"
    other_note = "; ".join(f"{why}: {n}" for why, n in sorted(other_general.items())).replace("(local phase)", "(read with the county's own ballot, below)")

    # 8b. the county and local part: the same ballots' county and county school contests
    lfiles, _lpath = read_local_ballots(folder, local_keep, say)
    facts, docs, missing = calendar_findings(local_keep, say)
    local = local_part(lfiles, counties, geoid_of, facts, docs, missing, folder)
    clash = sorted({r[0] for r in race_rows} & {r[0] for r in local["races"]})
    if clash:
        raise SystemExit(f"Alabama: a county race id is also a state race id ({clash[:3]})")
    if {s[0] for s in local["sources"]} & {SRC_GENERAL, SRC_PRIMARY, SRC_RUNOFF, SRC_DEM_P, SRC_REP_P, SRC_DEM_R, SRC_ROSTER, SRC_COUNTY}:
        raise SystemExit("Alabama: a local source id is also a state source id")
    kept_sha = hashlib.sha256(json.dumps(sorted(f["sha256"] for f in files)).encode()).hexdigest()

    def blanks_note(blanks, recon):
        over = sum(v for (r, c, k), v in blanks.items() if k == "Over Votes")
        under = sum(v for (r, c, k), v in blanks.items() if k == "Under Votes")
        ok, n, most, _none = recon
        return (f"{over:,} over votes and {under:,} under votes in these contests are not counted in anyone's share. Each "
                f"statewide contest's votes, with its over and under votes, equal the party's BALLOTS CAST line in {ok} of the {n} "
                f"county returns" + (f" and differ by at most {most:,} in the rest" if ok < n else ""))

    def chk_note(key, stage):
        diffs, carried, of = chk[key]
        return (f"The party's workbook carries {carried} of the {of} contests; " +
                ("where it differs: " + "; ".join(diffs) + "." if diffs else "it agrees with every one."))

    sources = [
        (SRC_GENERAL, STATE, "official sample ballots", AGENCY,
         "2026 General Election Sample Ballots, November 3, 2026 (state offices, every county)", SAMPLE_PAGES["general"], "2026-09-10",
         mdate(kept_path), kept_sha, len(gen_rows),
         f"{len(files)} county sample ballots (PDF; every ballot style of the county drawn one over another), read in drawing order "
         "style by style; only the offices, names, parties, order and Write-in lines of the state contests are kept "
         f"({KEPT}). The sha256 here is of the list of the files' own SHA-256s. Every 2026 style of every county prints each race "
         f"the same way unless named in the loader's report. Styles of an older ballot drawn beneath, set aside: {older_note}. "
         f"Contests on the November styles that are not state offices, left to other pages: {other_note}. The State "
         "Certifications of candidates (August 26, 2026) are scans with no text, so the typeset sample ballots are read "
         "instead. A ballot shows no withdrawn candidates and no declared write-ins, so none are loaded. Ballot order as "
         "printed. Names printed in capitals are shown in ordinary capitals."),
        (SRC_PRIMARY, STATE, "official results", AGENCY, "2026 Primary Election, May 19, 2026: precinct results by county (state offices)",
         SOS + FILES["primary"][0], "2026-07-31", mdate(paths["primary"]), sha256(paths["primary"]),
         sum(len(v) for (rid, _c), v in ptot.items() if rid in races),
         f"{len(pcounties)} county files (.xls; Contest Title, Party, Candidate, a column per precinct, ABSENTEE, PROVISIONAL), "
         f"summed over the counties; {blanks_note(pblank, precon)}. Linked from {DATA_PAGE}. Checked against the parties' certified "
         f"vote totals: Democratic: {chk_note('dem_primary', 'primary')} Republican: {chk_note('rep_primary', 'primary')}"),
        (SRC_RUNOFF, STATE, "official results", AGENCY, "2026 Primary Runoff Election, June 16, 2026: precinct results by county (state offices)",
         SOS + FILES["runoff"][0], "2026-07-31", mdate(paths["runoff"]), sha256(paths["runoff"]),
         sum(len(v) for (rid, _c), v in rtot.items() if rid in races),
         f"{len(rcounties)} county files, read as the primary's; {blanks_note(rblank, rrecon)}. Checked against the Democratic Party's "
         f"certified workbook: {chk_note('dem_runoff', 'runoff')} The Republican Party's certification of the runoff "
         f"({REP_RUNOFF_SCAN}) is a scan with no text, so the Republican runoff is checked only against the November ballot."),
    ]
    for key, sid, title in (("dem_primary", SRC_DEM_P, "Certification of Results, Democratic Party: vote totals, 2026 Primary Election"),
                            ("rep_primary", SRC_REP_P, "Certification of Results, Republican Party: vote totals, 2026 Primary Election"),
                            ("dem_runoff", SRC_DEM_R, "Certification of Results, Democratic Party: vote totals, 2026 Primary Runoff")):
        sources.append((sid, STATE, "official results (a check)", ("Alabama Democratic Party" if key.startswith("dem") else
                        "Alabama Republican Party") + ", posted by the Alabama Secretary of State", title, SOS + FILES[key][0], "",
                        mdate(paths[key]), sha256(paths[key]), 0,
                        "Used only to check the Secretary of State's precinct files; the differences are named in their notes."))
    sources.append((SRC_ROSTER, STATE, "roster", "Open States people project (CC0), as loaded into state_al.sqlite",
                    "Alabama legislators and statewide officials", "https://github.com/openstates/people", "", as_of,
                    "", sum(1 for r in races.values() if r["holder"]),
                    "Who holds each seat today: names, party and ids only."))
    sources.append((SRC_COUNTY, STATE, "county codes", "U.S. Census Bureau", "2024 cartographic boundary file, counties (500k)",
                    COUNTY_URL, "", mdate(county_zip), sha256(county_zip) if os.path.exists(county_zip) else "", len(counties),
                    "County names and FIPS codes only."))

    # Alabama's rows only, state and local together, in one transaction: the state rows are written exactly as before
    con = sqlite3.connect(db_path)
    con.executescript(SCHEMA + EXTRA_SCHEMA)
    with con:
        con.execute("DELETE FROM sl_candidates WHERE race_id IN (SELECT race_id FROM sl_races WHERE state = ?) OR race_id LIKE ?",
                    (STATE, f"2026-{STATE}-%"))
        con.execute("DELETE FROM sl_races WHERE state = ?", (STATE,))
        con.execute("DELETE FROM sl_sources WHERE state = ?", (STATE,))
        con.execute("DELETE FROM sl_places WHERE source_id LIKE ?", (f"{STATE.lower()}-%",))
        con.execute("DELETE FROM sl_gaps WHERE state = ?", (STATE,))
        con.execute("DELETE FROM sl_notes WHERE state = ?", (STATE,))
        con.executemany("INSERT INTO sl_races VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", race_rows + local["races"])
        con.executemany("INSERT INTO sl_candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", [tuple(c) for c in cands + local["cands"]])
        con.executemany("INSERT INTO sl_sources VALUES (?,?,?,?,?,?,?,?,?,?,?)", sources + local["sources"])
        con.executemany("INSERT INTO sl_places VALUES (?,?,?,?,?)", place_rows + local["places"])
        con.executemany("INSERT INTO sl_gaps VALUES (?,?,?,?,?,?,?)", local["gaps"])
        con.executemany("INSERT INTO sl_notes VALUES (?,?,?,?,?)", local["notes"])
    con.close()

    # 10. say what happened
    per = collections.Counter(races[c[0]]["office_kind"] for c in gen_rows)
    seats_per = collections.Counter(r["office_kind"] for r in races.values())
    filled = collections.Counter(races[rid]["office_kind"] for rid in {c[0] for c in gen_rows})
    say(f"    Alabama (state races): {len(races)} races, {len(gen_rows)} candidates on the November ballot; "
        f"{fields['primary']} primary fields on May 19, {fields['runoff']} runoffs on June 16; "
        f"{sum(1 for c in cands if c[7] and c[1] == 'general')} sitting members or officials on the November ballot")
    for kind in sorted(seats_per):
        say(f"      {kind}: {seats_per[kind]} races, {filled[kind]} with candidates, {per[kind]} candidates")
    for key, stage in (("dem_primary", "Democratic primary"), ("rep_primary", "Republican primary"), ("dem_runoff", "Democratic runoff")):
        diffs, carried, of = chk[key]
        say(f"      check against the {stage} workbook: {carried} of {of} contests carried, {len(diffs)} differ")
        for d in diffs:
            say(f"        {d}")
    for stage, (ok, n, most, _none) in (("primary", precon), ("runoff", rrecon)):
        say(f"      check: {stage} statewide contests against the party's ballots cast: {ok} of {n} county returns equal"
            + (f", the rest differ by at most {most:,}" if ok < n else ""))
    for p in problems:
        say(f"      problem: {p}")
    for c in checks:
        say(f"      check: {c}")
    st, lv = local["stats"], local["by_level"]
    say(f"    Alabama (county and school races): {st['races']:,} races ({', '.join(f'{k} {lv[k]:,}' for k in LOCAL_LEVELS if lv.get(k))}), "
        f"{st['cands']:,} candidates on the November ballot, a race in {st['reached']} of {COUNTIES} counties; races with more than one "
        f"candidate: {local['contested']}; with one: {local['single']}; with none printed: {st.get('no_name', 0)}; "
        f"{len(local['places'])} boards of education as places; {len(local['gaps'])} gaps")
    say("      office kinds: " + ", ".join(f"{k} {n:,}" for k, n in sorted(local["by_kind"].items(), key=lambda x: (-x[1], x[0]))))
    say("      offices as printed: " + "; ".join(f"{k} {n:,}" for k, n in sorted(local["offices"].items())))
    if local["new_kinds"]:
        say(f"      office kinds the pages do not know yet: {', '.join(local['new_kinds'])}")
    say(f"      two counts of the {st['files']} ballots: {st['vote_for_lines']:,} vote-for lines against {st['contests_read']:,} contests read "
        f"and {st.get('without_for', 0)} drawn without the word FOR; {st['party_lines']:,} party lines (the straight party boxes' "
        f"{st.get('straight_party_lines', 0):,} apart) against {st['candidate_lines']:,} candidates read and {st.get('without_for_parties', 0)} "
        f"under those; ballots that do not add up: {st.get('unreconciled', 0)}")
    say(f"      county and school contests on the November styles: {st['drawn']:,} drawings carrying {st['lines']:,} candidate lines; "
        f"{st.get('placed_lines', 0):,} lines placed, each in exactly one race ({st['cands']:,} candidates in {st['races']:,} races), "
        f"{st.get('unread_lines', 0):,} under titles not understood, {st.get('problem_lines', 0):,} in races left without candidates; "
        f"{st.get('older_drawn', 0)} drawings on older styles set aside; {st.get('blanked', 0)} name cells blanked, "
        f"{st.get('no_party', 0)} names with no party beneath on any contest of the ballots; races where the versions space a name "
        f"differently (the printing drawn most kept): {st.get('respaced', 0)}")
    if local["absent"]:
        say(f"      no November contest for: {', '.join(local['absent'])}")
    for g in local["gaps"]:
        say(f"      gap ({g[1]} {g[2]}): {g[4]}")
    for c in local["checks"]:
        say(f"      check: {c}")
    return {"races": len(races), "general": len(gen_rows), "fields": dict(fields), "problems": problems, "checks": checks,
            "local": {k: v for k, v in local.items() if k not in ("races", "cands", "places", "sources", "gaps", "notes")},
            "local_gaps": [g[:5] for g in local["gaps"]]}


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    ap.add_argument("db")
    ap.add_argument("--keep", help="folder for the state part's kept-cells JSON (default ballot_cache/al)")
    ap.add_argument("--local-keep", help="folder for the county and local part's kept cells (default ballot_cache/al/local)")
    a = ap.parse_args()
    load(a.db, keep=a.keep, local_keep=a.local_keep)
