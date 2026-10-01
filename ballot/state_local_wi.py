"""
ballot/state_local_wi.py - Wisconsin's state races on the November 3, 2026 ballot: Governor and Lieutenant Governor
(one ticket in November, nominated in separate party primaries in August), Attorney General, Secretary of State, State
Treasurer, the seventeen odd-numbered State Senate seats (four-year terms, the odd half up in 2026) and all 99 seats of
the Assembly, with each party's August 11 primary field and its official votes; and, county by county, the county
offices on the same ballot (see "The local level" below).

Sources, both the Wisconsin Elections Commission's own, both already cached for the federal loader (ballot/lists/wi.py)
because elections.wi.gov refuses scripts and serves browsers; this loader downloads nothing from the Commission:
  - "Candidates on Ballot by Election", 2026 General Election (PDF). Four columns: ballot order, committee ID, candidate,
    party. Under each "Office : ... Incumbent: ..." heading one line per candidate, a long party name wrapping onto the
    next line ("Wisconsin" / "Green"), then "Total Number of ... Candidates :N", which is checked. Columns are found by
    where the page's own column headings start. The committee ID is never kept.
  - The County by County Report of the 2026 Partisan Primary (xlsx), one sheet per office and party, named in its
    Document map. Each sheet's county rows must add up to its "Office Totals:" row, and the candidates plus SCATTERING
    (write-ins nobody registered) to its Total Votes Cast. Registered write-ins are listed as "Name (write-in)".
    The county rows also say which counties each district reaches (their Census codes from the Bureau's county file).
Holders of each seat come from the Open States roster in state_wi.sqlite (legislators, is_current = 1; the officials
table for Governor, Lieutenant Governor and Attorney General). The roster does not carry the Secretary of State or the
Treasurer, so for those two the holder is the incumbent the Commission's list names.

The local level (county offices), county by county
--------------------------------------------------
On November 3, 2026 a Wisconsin county elects its sheriff and its clerk of circuit court, and its coroner where it
still has one (Wis. Stat. 59.20 (2) (b) and (bm)), all on the partisan ballot. Cities, villages, towns, school boards,
county boards, county executives and judges are elected in April (5.02 (21)); county clerks, treasurers, registers of
deeds and district attorneys in November 2028 (59.20 (2) (a), 978.01 (1)). There is no statewide list of county
candidates: each of the 72 county clerks publishes the county's own sample ballots or candidate listing, each in its own
way, and the Commission's list of the clerks' sites is on the site that refuses scripts. So this loader keeps a table of
the clerks' files a script could fetch when the sites were looked at (COUNTY_FILES), reads each through states/net.py
(the honest User-Agent, one request at a time), and says for every other county, in sl_gaps, that its list is not
loaded yet and why. Two kinds of file are read:
  - a sample ballot (one ballot or a whole county's in one PDF). A contest is found by its "Vote for 1" line: the title
    is the line or lines set just above it at the same left edge, the candidates are the indented lines below it (the
    name, then the party in brackets or in smaller type under it) down to the write-in line. The order is the printed
    order. A ballot says which municipality it is "for", and for a municipality that lies in two counties which county;
    a clerk's file can hold the ballot of a neighbouring county's wards (Columbia County's holds the City of Wisconsin
    Dells ballots of Sauk and Adams counties), and those contests are filed under the county the ballot names.
    Every ballot of a county must give the same candidates for a countywide office, or the county is not loaded.
  - a clerk's listing of candidates (Walworth's "Listing of Candidates ...", Winnebago's "Election Information
    Packet"): office, "Vote for One", then name and party. A listing does not state the printed order, so none is
    stored.
Each file is read twice by two routes (the positioned pieces, and the plain rows of the page) and the two must agree
on how many contests a page holds and on every name kept. A scan read by machine (an invisible text layer over page
images) is never trusted for names, and a file with no text at all cannot be read; both leave a gap.
A municipal judge on a November ballot is a special election to fill a vacancy (municipal judges are elected in April;
Wis. Stat. 755.01 (1), 8.50 (4) (fm)); the two that the ballots read here carry are filed under level "court".
Not loaded: ballot questions, the write-in lines, and anyone not printed on the ballot.

A county whose site refuses scripts waits for a sample ballot saved by hand: any PDF in
ballot_cache/wi/local/<county>/ (milwaukee, dane, brown, la-crosse ...) is read as a sample ballot. A clerk's file
that does not answer, or cannot be read, is remembered for a week and not asked for again until then (--refresh asks).

Privacy: the Commission's two files carry no addresses, phones, e-mail or websites. A sample ballot carries none
either; a clerk's packet can (Winnebago's lists polling places with their street addresses on later pages). Whatever
the file, it is read in memory and only the allowed cells are kept, as a small JSON file in ballot_cache/wi/local/:
the office title, the names and parties under it, and the municipality and county a ballot is for. A fetched file is
never saved (a sample ballot saved by hand stays where it was put, with its kept cells beside it), and no line of any
file is ever printed or logged; a file that does not read as expected is named with the page and the check it failed,
never the text. The Census Bureau's list of Wisconsin's cities, villages and towns, which has no contact columns, is
kept whole.

    python -m ballot.state_local_wi <path to a test database> [--refresh]
"""

import collections
import datetime as dt
import glob
import hashlib
import io
import json
import os
import re
import sqlite3
import sys
import time
import unicodedata
import zipfile

import openpyxl

from ballot.check_local import EXTRA_SCHEMA, contact_like
from ballot.common import CACHE, HERE, fold, name_parts, party_code
from ballot.match import fits
from ballot.pdftext import PDF, Ref, join, page_runs, rows
from states import net

STATE, FIPS = "WI", "55"
GENERAL, PRIMARY = "2026-11-03", "2026-08-11"
LIST = "wi_candidates_on_ballot_2026_general.pdf"
BOOK = "wi_county_by_county_2026_partisan_primary.xlsx"
LIST_URL = "https://elections.wi.gov/media/40951/download"
BOOK_URL = ("https://elections.wi.gov/sites/default/files/documents/"
            "County%20by%20County%20Report_Partisan%20Primary%202026_All%20State%20Contests.xlsx")
COUNTY_ZIP = os.path.join(HERE, "states_cache", "census", "cb_2024_us_county_500k.zip")
ROSTER = os.path.join(HERE, "state_wi.sqlite")
SRC_LIST, SRC_BOOK = "wi-wec-2026-state-candidates-on-ballot", "wi-wec-2026-state-primary-county"
SRC_COUNTY, SRC_ROSTER = "wi-census-2024-county-codes", "wi-openstates-roster"

SCHEMA = """
CREATE TABLE IF NOT EXISTS sl_races (race_id TEXT PRIMARY KEY, state TEXT NOT NULL, level TEXT NOT NULL, office_kind TEXT NOT NULL,
  office TEXT NOT NULL, jurisdiction TEXT, jurisdiction_id TEXT, county_ids TEXT, district TEXT, seat TEXT, special INTEGER NOT NULL DEFAULT 0,
  partisan INTEGER NOT NULL, holder_id TEXT, holder_name TEXT, holder_party TEXT, election_date TEXT NOT NULL, note TEXT);
CREATE TABLE IF NOT EXISTS sl_candidates (race_id TEXT NOT NULL, election TEXT NOT NULL, election_date TEXT, name TEXT NOT NULL, party TEXT,
  party_code TEXT, ballot_order INTEGER, incumbent INTEGER NOT NULL DEFAULT 0, write_in INTEGER NOT NULL DEFAULT 0, votes INTEGER, pct REAL,
  outcome TEXT, state_member_id TEXT, source_id TEXT, note TEXT, PRIMARY KEY (race_id, election, name));
CREATE TABLE IF NOT EXISTS sl_sources (source_id TEXT PRIMARY KEY, state TEXT, kind TEXT, agency TEXT, title TEXT, url TEXT, published TEXT,
  fetched TEXT, sha256 TEXT, rows INTEGER, note TEXT);
CREATE TABLE IF NOT EXISTS sl_places (kind TEXT NOT NULL, id TEXT NOT NULL, name TEXT NOT NULL, county_ids TEXT, source_id TEXT,
  PRIMARY KEY (kind, id));
"""

# The Commission's office headings, and what each is here: (race suffix, level, office_kind, office as shown).
STATEWIDE = {
    "GOVERNOR": ("GOV", "governor", "Governor and Lieutenant Governor"),
    "LIEUTENANT GOVERNOR": ("LTG", "lieutenant_governor", "Lieutenant Governor"),
    "ATTORNEY GENERAL": ("AG", "attorney_general", "Attorney General"),
    "SECRETARY OF STATE": ("SOS", "secretary_of_state", "Secretary of State"),
    "STATE TREASURER": ("TREAS", "state_treasurer", "State Treasurer"),
}
ROSTER_OFFICE = {"GOV": "governor", "LTG": "lt_governor", "AG": "attorney general"}     # officials.office in state_wi.sqlite
ON_LIST = ("GOVERNOR", "ATTORNEY GENERAL", "SECRETARY OF STATE", "STATE TREASURER")    # November; LTG rides on the GOV ticket
SENATE = re.compile(r"STATE SENATOR DISTRICT (\d+)")
ASSEMBLY = re.compile(r"REPRESENTATIVE TO THE ASSEMBLY DISTRICT (\d+)")
FEDERAL = re.compile(r"REPRESENTATIVE IN CONGRESS|UNITED STATES SENATOR")
PARTIES = {"Republican", "Democratic", "Libertarian", "Constitution", "Independent", "Wisconsin Green",
           "Serving People Not Parties", "American Solidarity Party"}
PRIMARY_CODE = {"Republican": "REP", "Democratic": "DEM", "Libertarian": "LIB", "Constitution": "CON", "Wisconsin Green": "WGR"}
OFFICE = re.compile(r"^Office : (?P<office>.+?)(?: Incumbent: (?P<inc>.*))?$")
TOTAL = re.compile(r"^Total Number of (?P<office>.+) Candidates :(?P<n>\d+)$")
FURNITURE = {"Candidates on Ballot by Election", "Wisconsin Elections Commission", "Ballot Committee Candidate Party", "Order# ID"}
STAMP = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4}) \d{1,2}:\d{2}:\d{2}")
NONCAND = "(Filed Notification of Noncandidacy)"
WRITE_IN = re.compile(r"\s*\((?:write-in)\)\s*$", re.I)


def race_of(office):
    """(race_id, level, office_kind, office, district) for a Commission office heading; None for a federal office."""
    o = office.strip()
    if FEDERAL.match(o):
        return None
    if o in STATEWIDE:
        kind, okind, shown = STATEWIDE[o]
        return f"2026-{STATE}-{kind}", "statewide", okind, shown, None
    m = SENATE.fullmatch(o)
    if m:
        return f"2026-{STATE}-SS{int(m.group(1))}", "legislature", "state_senate", "State Senator", str(int(m.group(1)))
    m = ASSEMBLY.fullmatch(o)
    if m:
        return f"2026-{STATE}-SH{int(m.group(1))}", "legislature", "state_house", "Representative to the Assembly", str(int(m.group(1)))
    raise SystemExit(f"Wisconsin: an office this loader does not know is on the list: {o!r}")


def key(text):
    return " ".join(fold(text).split())


# ---------------------------------------------------------------------------------------------------- the November list

def general_list(path):
    """[(office heading, incumbent named or '', filed noncandidacy, [(order, name, party)])] in the list's order, and the
    date the list was printed. Each office's candidates are counted against the list's own total."""
    pdf = PDF(open(path, "rb").read())
    offices, printed, cur, cols = [], "", None, None
    for page, res in pdf.pages():
        for _y, rs in rows(pdf, page, res):
            text = join(rs)
            if text.startswith("Ballot Committee"):                      # the column headings give the column edges
                at = {r[3].strip(): r[0] for r in rs}
                cols = (at["Committee"], at["Candidate"], at["Party"])
                continue
            m = STAMP.search(text)
            if m and not text[:1].isdigit() or text in FURNITURE or text.startswith("2026 General Election"):
                if m and not printed:
                    printed = f"{m.group(3)}-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
                continue
            m = OFFICE.match(text)
            if m:
                inc = (m.group("inc") or "").strip()
                cur = [m.group("office").strip(), inc.replace(NONCAND, "").strip(), NONCAND in inc, []]
                offices.append(cur)
                continue
            m = TOTAL.match(text)
            if m:
                if not cur or m.group("office").strip() != cur[0]:
                    raise SystemExit(f"Wisconsin: a total for {m.group('office')!r} came outside its office")
                if len(cur[3]) != int(m.group("n")):
                    raise SystemExit(f"Wisconsin: the list counts {m.group('n')} candidates for {cur[0]}; {len(cur[3])} were read")
                cur = None
                continue
            if cur is None or cols is None:
                continue
            order = join([r for r in rs if r[0] < cols[0] - 5])
            name = join([r for r in rs if cols[1] - 5 <= r[0] < cols[2] - 5])
            party = join([r for r in rs if r[0] >= cols[2] - 5])
            if order:
                if not order.isdigit() or not name:
                    raise SystemExit(f"Wisconsin: a candidate line under {cur[0]} has no ballot order or no name")
                cur[3].append([int(order), name, party])
            elif cur[3]:                                                    # a name or party wrapped onto the next line
                last = cur[3][-1]
                last[1] = f"{last[1]} {name}".strip()
                last[2] = f"{last[2]} {party}".strip()
    for office, _inc, _nc, cands in offices:
        for _o, name, party in cands:
            if party not in PARTIES:
                raise SystemExit(f"Wisconsin: {office} lists a party this loader does not know: {party!r} (add it to PARTIES)")
    return offices, printed


# ----------------------------------------------------------------------------------------------- the August primary

def primary_book(path):
    """{office heading: {"counties": set, "fields": {party: (code, total, [(name, write_in, votes)])}}}, every sheet's
    county rows added to its Office Totals and its candidates to its Total Votes Cast."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    titles = [r[1] for r in wb["Document map"].iter_rows(values_only=True) if len(r) > 1 and r[1]]
    out, problems = {}, []
    for k, title in enumerate(titles):
        title = str(title).strip()
        office, _, party = title.rpartition(" - ")
        if FEDERAL.match(office):
            continue
        race_of(office)                                                     # stops on an office not known here
        ws = wb.worksheets[k + 1]
        sheet = [tuple(r) for r in ws.iter_rows(values_only=True)]
        if not any(str(c or "").strip() == title for r in sheet[:5] for c in r):
            raise SystemExit(f"Wisconsin: sheet {ws.title} does not carry the office its Document map names ({title})")
        head = next(j for j, r in enumerate(sheet) if r and str(r[0] or "").strip() == "County")
        hdr, cands = sheet[head], sheet[head + 1]
        cast = next(j for j, c in enumerate(hdr) if str(c or "").strip() == "Total Votes Cast")
        totals = [r for r in sheet if r and str(r[0] or "").strip().startswith("Office Totals")]
        if len(totals) != 1:
            raise SystemExit(f"Wisconsin: {title} has {len(totals)} Office Totals rows")
        tot = totals[0]
        counties = [r for r in sheet[head + 2:] if r and r[0] and str(r[0]).strip() and not str(r[0]).strip().startswith("Office Totals")]
        cols = [j for j, c in enumerate(cands) if c and str(c).strip()]
        if sum(int(tot[j] or 0) for j in cols) != int(tot[cast] or 0):
            problems.append(f"{title}: candidates and write-ins do not add up to Total Votes Cast")
        for j in cols + [cast]:
            if sum(int(r[j] or 0) for r in counties) != int(tot[j] or 0):
                problems.append(f"{title}: county rows do not add up to Office Totals in column {j + 1}")
        entry = out.setdefault(office, {"counties": set(), "fields": {}})
        entry["counties"] |= {key(str(r[0])) for r in counties}
        abbr = next((str(c).strip() for c in hdr[cast + 1:] if c and str(c).strip()), "") or PRIMARY_CODE.get(party, party[:3].upper())
        field = []
        for j in cols:
            raw = str(cands[j]).strip()
            if raw == "SCATTERING":
                continue
            field.append((WRITE_IN.sub("", raw).strip(), bool(WRITE_IN.search(raw)), int(tot[j] or 0)))
        entry["fields"][party] = (abbr, int(tot[cast] or 0), field)
    if problems:
        raise SystemExit("Wisconsin: the primary report does not reconcile:\n  " + "\n  ".join(problems))
    return out


# -------------------------------------------------------------------------------------------------- the roster

def roster(path):
    con = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    legs = [dict(zip(("id", "first", "last", "full", "other", "party", "chamber", "district"), r)) for r in con.execute(
        "SELECT bioguide_id, first_name, last_name, official_full, other_names, party_name, chamber, district FROM legislators WHERE is_current = 1")]
    offs = {r[1]: dict(zip(("id", "office", "full", "first", "last", "party"), r)) for r in con.execute(
        "SELECT bioguide_id, office, official_full, first_name, last_name, party_name FROM officials")}
    con.close()
    return legs, offs


def person_fits(name, p):
    """A listed name fits a roster person: same family name and a given name that fits (their roster name or any
    full form of it the roster keeps)."""
    cand = name_parts(name)
    forms = [(fold(p["first"] or "").split(), key(p["last"] or ""))]
    forms += [name_parts(o.strip()) for o in (p.get("other") or "").split(";") if o.strip() and "," not in o and "." not in o]
    return any(f[1] and fits(cand, f) for f in forms)


def chamber_words(p):
    return f"{'State Senate' if p['chamber'] == 'Senate' else 'Assembly'}, District {p['district']}"


# ---------------------------------------------------------------------------------------------------- loading

def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def mtime(path):
    return dt.datetime.fromtimestamp(os.path.getmtime(path)).strftime("%Y-%m-%d")


def census_counties(path):
    import shapefile                                                        # pyshp, in the kit's environment
    z = zipfile.ZipFile(path)
    base = next(n for n in z.namelist() if n.endswith(".dbf"))
    rdr = shapefile.Reader(dbf=io.BytesIO(z.read(base)))
    return {key(r["NAME"]): (r["GEOID"], r["NAME"]) for r in (x.as_dict() for x in rdr.iterRecords()) if r["STATEFP"] == FIPS}


# ====================================================================================================================
# The local level: county offices, county by county
# ====================================================================================================================

LOCAL_DIR = os.path.join(CACHE, "wi", "local")
COUSUB_URL = "https://www2.census.gov/geo/docs/reference/codes2020/cousub/st55_wi_cousub2020.txt"
COUSUB_FILE = "census_st55_wi_cousub2020.txt"
COUNTY_URL = "https://www2.census.gov/geo/tiger/GENZ2024/shp/cb_2024_us_county_500k.zip"
STATUTE_URL = "https://docs.legis.wisconsin.gov/statutes/statutes/59/iv/20/2"
SPECIAL_URL = "https://docs.legis.wisconsin.gov/statutes/statutes/8/50"
SRC_COUSUB = "wi-census-2020-cousub"
NONPARTISAN = "Nonpartisan office"
MAX_AGE_DAYS = 7                 # a kept copy of a clerk's file is used again for a week; --refresh asks again
LOOKED = "October 1, 2026"       # the day the clerks' pages were looked at (COUNTY_FILES and NOT_YET below)
WHAT = "candidates for sheriff, clerk of circuit court and, where the county has one, coroner"

# The county clerks' own files a script could fetch on the day the sites were looked at. kind "ballot": sample ballots
# (one, or the whole county's in one PDF); kind "listing": the clerk's listing of candidates. A county with two files is
# read twice and the two must agree.
COUNTY_FILES = (
    {"county": "Winnebago", "kind": "listing", "key": "packet", "title": "Election Information Packet for November 3, 2026 (listing of candidates)",
     "url": "https://www.winnebagocountywi.gov/DocumentCenter/View/4097", "page": "https://www.winnebagocountywi.gov/262"},
    {"county": "Kenosha", "kind": "ballot", "key": "paris", "title": "Colored Sample Ballot, November 3, 2026: Paris",
     "url": "https://www.kenoshacountywi.gov/DocumentCenter/View/31769/COLORED-SAMPLE-BALLOT---PARIS",
     "page": "https://www.kenoshacountywi.gov/139/Election-Information"},
    {"county": "Kenosha", "kind": "ballot", "key": "brighton", "title": "Colored Sample Ballot, November 3, 2026: Brighton",
     "url": "https://www.kenoshacountywi.gov/DocumentCenter/View/31764/COLORED-SAMPLE-BALLOT---BRIGHTON",
     "page": "https://www.kenoshacountywi.gov/139/Election-Information"},
    {"county": "Marathon", "kind": "ballot", "key": "all", "title": "November 2026 General Election Sample Ballots",
     "url": "https://www.marathoncounty.gov/home/showpublisheddocument/19219/639246328067230000",
     "page": "https://www.marathoncounty.gov/services/elections-voting/election-information"},
    {"county": "Walworth", "kind": "listing", "key": "listing",
     "title": "Listing of Candidates for Congressional, Legislative, State and County Offices to be Voted on in Walworth County, November 3, 2026, General Election",
     "url": "https://www.co.walworth.wi.us/DocumentCenter/View/20544/Vote-for-One--Candidates-for-CongressionalLegisState-110326",
     "page": "https://www.co.walworth.wi.us/198/elections"},
    {"county": "Jefferson", "kind": "ballot", "key": "all", "title": "November Sample Ballots (November 3, 2026)",
     "url": "https://www.jeffersoncountywi.gov/County%20Clerk/Elections/November%202026%20Sample%20Ballots.pdf",
     "page": "https://www.jeffersoncountywi.gov/county_clerk/elections/index.php"},
    {"county": "Portage", "kind": "ballot", "key": "all", "title": "Sample Ballots, 2026 General Election",
     "url": "https://www.co.portage.wi.gov/DocumentCenter/View/7365/2026-General-Election-Sample-Ballot---web",
     "page": "https://www.co.portage.wi.gov/693/Election-Information"},
    {"county": "Columbia", "kind": "ballot", "key": "all", "title": "November 3, 2026 Sample Ballots",
     "url": "https://www.co.columbia.wi.us/columbiacounty/LinkClick.aspx?fileticket=z_XyKLKFm44%3d&tabid=102&portalid=2&mid=2148",
     "page": "https://www.co.columbia.wi.us/columbiacounty/countyclerk/ElectionInformation/tabid/102/Default.aspx"},
    {"county": "Douglas", "kind": "ballot", "key": "superior-wards-1-5", "title": "Sample ballot, November 3, 2026: City of Superior Wards 1-5",
     "url": "https://www.douglascountywi.gov/DocumentCenter/View/16414", "page": "https://www.douglascountywi.gov/1131/Sample-Ballots"},
    {"county": "Douglas", "kind": "ballot", "key": "wascott", "title": "Sample ballot, November 3, 2026: Town of Wascott",
     "url": "https://www.douglascountywi.gov/DocumentCenter/View/16439", "page": "https://www.douglascountywi.gov/1131/Sample-Ballots"},
)

# Why a county has no file in the table above, as found on the day the sites were looked at: {county: (why, the
# clerk's page)}. A county in neither table has not been looked at closely; its gap says only that.
REFUSES = "refuses"          # the county's site answered scripts with 403 or a challenge page: never worked around
NOT_YET = "not yet"          # the clerk's election pages linked no November 3 sample ballot or candidate listing
PICTURE = "picture"          # the sample ballots are posted as pictures with no text a script can read
SCAN = "scan"                # the clerk's list is a scanned page; a machine's reading of it is not trusted for names
NO_PARTY = "no party"        # the only file names the candidates without their parties
WHY = {
    "Milwaukee": (REFUSES, "https://county.milwaukee.gov/EN/County-Clerk/Election-Commission"),
    "Dane": (REFUSES, "https://elections.danecounty.gov/"),
    "Brown": (REFUSES, "https://www.browncountywi.gov/departments/county-clerk/elections/sample-ballots/"),
    "La Crosse": (REFUSES, "https://www.lacrossecounty.org/"),
    "Marinette": (REFUSES, "https://www.marinettecountywi.gov/"),
    "Lincoln": (REFUSES, "https://www.co.lincoln.wi.us/"),
    "Langlade": (REFUSES, "https://www.co.langlade.wi.us/"),
    "Taylor": (REFUSES, "https://www.co.taylor.wi.us/"),
    "Shawano": (REFUSES, "https://www.co.shawano.wi.us/"),
    "Waukesha": (NOT_YET, "https://www.waukeshacounty.gov/county-clerks-office/election-information/"),
    "Racine": (NOT_YET, "https://www.racinecounty.gov/departments/county-clerk/election-information"),
    "Outagamie": (NOT_YET, "https://www.outagamie.gov/Our-County/County-Clerk/Elections"),
    "Rock": (NOT_YET, "https://www.co.rock.wi.us/departments/county-clerk/election-information"),
    "Washington": (NOT_YET, "https://www.washcowisco.gov/cms/One.aspx?portalId=16228038&pageId=17298071"),
    "Sheboygan": (NOT_YET, "https://www.sheboygancounty.com/government/upcoming-election-sample-ballots"),
    "Eau Claire": (NOT_YET, "https://www.eauclairecounty.gov/departments/county_clerk/elections/index.php"),
    "Fond du Lac": (NOT_YET, "https://www.fdlco.wi.gov/departments/departments-a-e/county-clerk/election-dates"),
    "St. Croix": (NOT_YET, "https://www.sccwi.gov/464/Voting-Elections"),
    "Manitowoc": (NOT_YET, "https://manitowoccountywi.gov/departments/county-clerk/election-information/"),
    "Wood": (NOT_YET, "https://www.woodcountywi.gov/Departments/Clerk/"),
    "Oneida": (NOT_YET, "https://www.oneidacountywi.gov/government/election-information/"),
    "Grant": (NOT_YET, "https://gcinternetwebsite.co.grant.wi.gov/elections/"),
    "Waupaca": (NOT_YET, "https://www.waupacacounty-wi.gov/departments/government_departments/county_clerk/"),
    "Monroe": (NOT_YET, "https://www.co.monroe.wi.us/"),
    "Green": (NOT_YET, "https://www.greencountywi.org/"),
    "Sawyer": (NOT_YET, "https://www.sawyercountygov.org/"),
    "Pierce": (NOT_YET, "https://www.co.pierce.wi.us/departments/county_clerk/elections.php"),
    "Vernon": (NOT_YET, "https://www.vernoncountywi.gov/departments/county_clerk/elections_and_voting.php"),
    "Waushara": (NOT_YET, "https://www.co.waushara.wi.us/38818/upcoming-elections-and-whats-on-the-ballot"),
    "Juneau": (PICTURE, "https://www.co.juneau.wi.gov/"),
    "Green Lake": (PICTURE, "https://www.greenlakecountywi.gov/"),
    "Marquette": (PICTURE, "https://www.marquettecountywi.gov/government/election-information/"),
    "Polk": (PICTURE, "https://www.polkcountywi.gov/government/elected_officials/county_clerk/elections.php"),
    "Dodge": (SCAN, "https://www.co.dodge.wi.gov/government/county-clerk/elections-voting"),
    "Chippewa": (NO_PARTY, "https://www.chippewacountywi.gov/241"),
}
REASON = {
    REFUSES: "The county clerk publishes the county's sample ballots, but the county's website turns scripts away, so this county's list is not "
             "loaded yet; it waits for a sample ballot saved by hand from a browser.",
    NOT_YET: f"The county clerk publishes the county's sample ballots and candidate list. When the clerk's election pages were looked at on {LOOKED} "
             "they did not yet link either one for November 3, so this county's list is not loaded yet.",
    PICTURE: "The county clerk posts the county's November 3 sample ballots as pictures with no text a script can read, so this county's list is "
             "not loaded yet.",
    SCAN: "The county clerk's list of November 3 races and candidates is posted as a scanned page. The letters a machine reads from a scan are "
          "not reliable enough to give names as filed, so this county's list is not loaded yet.",
    NO_PARTY: "The county clerk's November 3 tally sheet names the candidates without their parties, and no sample ballot was linked when the "
              f"page was looked at on {LOOKED}, so this county's list is not loaded yet.",
    None: "This county's list is not loaded yet. Wisconsin has no statewide list of county candidates: each county clerk publishes the county's "
          "own sample ballots, and this county's have not been read.",
}

VOTE_FOR = re.compile(r"^\(?\s*Vote for (?:not more than |no more than |up to )?(\d+|one|two|three)\s*\)?$", re.I)
VOTE_FOR_ANY = re.compile(r"Vote for (?:not more than |no more than |up to )?(?:\d+|one|two|three)\b", re.I)
WRITE_IN_LINE = re.compile(r"^write-?\s?in:?[\s_]*$", re.I)
IN_BRACKETS = re.compile(r"^\(([^()]+)\)$")
NAME_BRACKETS = re.compile(r"^(.+?)\s*\(([^()]+)\)$")
NAME_DASH = re.compile(r"^(.+?)\s*[\u2013\u2014]\s*(.+)$|^(.+?)\s+-\s+(.+)$")
NUMBER = {"one": 1, "two": 2, "three": 3}
PARTY_WORDS = {"republican": "Republican", "democratic": "Democratic", "libertarian": "Libertarian", "constitution": "Constitution",
               "wisconsin green": "Wisconsin Green", "independent": "Independent"}
GAP = 5.5           # points: a wider space than this between two runs of a printed row starts a new piece (a column gutter)
STATE_OFFICE = re.compile(r"governor|attorney general|secretary of state|state treasurer|congress|united states senator|state senat|assembly", re.I)
# words that say a title names a local office; a listing's title with none of them is a state or federal block whose
# title the clerk set in a way this loader does not read (a district's wards between the title and the "Vote for" line)
LOCAL_WORD = re.compile(r"judge|supervisor|trustee|alder|mayor|board|register|treasurer|surveyor|attorney|clerk|coroner|sheriff|executive|"
                        r"constable|assessor|commissioner", re.I)
COUNTY_OFFICES = ((re.compile(r"(?:county )?sheriff", re.I), "sheriff", "Sheriff"),
                  (re.compile(r"clerk of (?:the )?circuit courts?", re.I), "clerk_of_court", "Clerk of Circuit Court"),
                  (re.compile(r"(?:county )?coroner", re.I), "coroner", "Coroner"))
MUNICIPAL_JUDGE = re.compile(r"(?:(.+?) )?municipal (court )?judge", re.I)
FOR_STOP = re.compile(r"ballot issued|absentee|initial|_{3,}|official ballot|notice to", re.I)
COUNTY_LINE = re.compile(r"^([A-Za-z][A-Za-z .]*?) County(?:,? (?:WI|Wisconsin))?$")
MUNI_A = re.compile(r"^(City|Village|Town) of ([A-Za-z][A-Za-z .'\-]*?)\s*(?:\(.*|W\d.*|Wards?\b.*)?$", re.I)
MUNI_B = re.compile(r"^([A-Z][A-Z .'\-]*?) (C|V|T) WDS? [\d,&\- ]+$")
MUNI_C = re.compile(r"^(C|V|T)\. ([A-Za-z][A-Za-z .'\-]*?) W(?:ards?)? ?[\d,&\- ]+$")
MUNI_KIND = {"c": "city", "v": "village", "t": "town", "city": "city", "village": "village", "town": "town"}
# the page builder's own test for a street address, a little wider than the trial check's (Court, Place)
BUILDER_STREET = re.compile(r"\b\d{1,6}\s+(?:[NSEW]\.?\s+)?[A-Za-z0-9.' -]{1,40}?\s(?:St|Street|Ave|Avenue|Rd|Road|Blvd|Boulevard|Dr|Drive|Ln|Lane|Way|Ct|Court|"
                            r"Cir|Circle|Pkwy|Parkway|Hwy|Highway|Trl|Trail|Pl|Place|Ter|Terrace)\b\.?", re.I)


class Unreadable(Exception):
    """A clerk's file that does not read as this loader expects. Carries the check that failed and the page, never a
    line of the file."""


def slug(text):
    t = "".join(c for c in unicodedata.normalize("NFKD", text or "") if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "-", t.lower()).strip("-")


def looks_like_contact(text):
    return bool(text) and (contact_like(text, True) or bool(BUILDER_STREET.search(str(text))))


def row_groups(runs):
    """Runs grouped into printed rows, top to bottom, as ballot/pdftext.rows does."""
    out = []
    for r in sorted(runs, key=lambda r: (-round(r[1], 1), r[0])):
        if out and abs(out[-1][0] - r[1]) <= max(1.5, 0.35 * r[2]):
            out[-1][1].append(r)
        else:
            out.append([r[1], [r]])
    return out


def pieces(runs):
    """A page's text as pieces, top to bottom: [(y, x0, x1, type size, text)]. Runs on one printed row are joined
    while they touch or stand a word space apart; a wider gap (a column gutter) starts a new piece."""
    out = []
    for _y, rs in row_groups([r for r in runs if r[2] >= 5 and r[3].strip()]):
        rs = sorted(rs, key=lambda r: r[0])
        cur = [rs[0]]
        for r in rs[1:] + [None]:
            if r is not None and r[0] - max(c[4] for c in cur) <= GAP:
                cur.append(r)
                continue
            text = join(cur)
            if text:
                out.append((cur[0][1], min(c[0] for c in cur), max(c[4] for c in cur), max(c[2] for c in cur), text))
            cur = [r]
    return sorted(out, key=lambda p: (-p[0], p[1]))


def party_words(text):
    """A party as this state's rows write it: the five parties with ballot status and "Independent" in their own
    words, anything else (an independent candidate's statement of principle) as printed."""
    t = " ".join((text or "").split())
    return PARTY_WORDS.get(re.sub(r"\s+party$", "", t.lower()), t)


def ballot_contests(ps, page):
    """Every contest on a ballot page, found by its "Vote for" line: [{title, vote_for, write_ins, candidates}]. The
    title is the line or lines set just above the "Vote for" line at the same left edge; the candidates are the
    indented lines below it, each a name with its party in brackets (or in smaller type) under or after it, down to
    the write-in line. A contest that does not read that way carries a "doubt" saying how."""
    out = []
    for v in ps:
        m = VOTE_FOR.match(v[4])
        if not m:
            continue
        title, y = [], v[0]
        while len(title) < 3:
            above = [p for p in ps if abs(p[1] - v[1]) <= 2.5 and 0 < p[0] - y <= 1.9 * max(p[3], v[3])]
            if not above:
                break
            p = min(above, key=lambda p: p[0])
            if VOTE_FOR.match(p[4]) or WRITE_IN_LINE.match(p[4]) or IN_BRACKETS.match(p[4]):
                break
            title.insert(0, p[4])
            y = p[0]
        cands, write_ins, last, doubt = [], 0, None, None          # last: (y, type size) of the last name line
        for p in [p for p in ps if p[0] < v[0] and v[1] - 2.5 <= p[1] <= v[1] + 60]:
            text = p[4]
            if WRITE_IN_LINE.match(text):
                write_ins += 1
                continue
            if write_ins or abs(p[1] - v[1]) <= 2.5:
                break                                              # past the write-in line, or the next title: the contest is over
            bracket, both = IN_BRACKETS.match(text), NAME_BRACKETS.match(text)
            open_name = bool(cands) and cands[-1][1] is None and last is not None
            if bracket and not open_name:
                doubt = "a party line with no name above it"
                break
            if bracket:
                cands[-1][1] = bracket.group(1).strip()
            elif text.startswith("("):
                doubt = "a party line that runs onto a second line"
                break
            elif open_name and last[0] - p[0] <= 1.35 * last[1] and (p[3] <= last[1] - 0.8 or text.lower() in PARTY_WORDS):
                cands[-1][1] = text                                # the party under the name, in smaller type and no brackets
            elif open_name and last[0] - p[0] <= 1.25 * last[1] and not both:
                cands[-1][0] += " " + text                         # a long name set on two lines
                last = (p[0], p[3])
            elif both:
                cands.append([both.group(1).strip(), both.group(2).strip()])
                last = (p[0], p[3])
            else:
                cands.append([text, None])
                last = (p[0], p[3])
        out.append({"title": " ".join(title), "vote_for": int(NUMBER.get(m.group(1).lower(), m.group(1))), "write_ins": write_ins,
                    "candidates": [[c[0], party_words(c[1]) if c[1] else None] for c in cands], **({"doubt": doubt} if doubt else {})})
    return out


def listing_contests(ps, page):
    """Every contest in a clerk's listing of candidates, found by its "Vote for" line: the title is the line or lines
    just above it (in one type size); under it each candidate is "Name - Party" on one line, "Name (Party)", or a name
    with the party on the line below. Lines are aligned with the "Vote for" line by their left edge or by their middle
    (a centred listing). A contest ends at the next title or at the first line that is not a candidate with a party;
    one that does not read that way carries a "doubt" saying how."""
    def middle(p):
        return (p[1] + p[2]) / 2

    def aligned(p, v):
        return abs(p[1] - v[1]) <= 2.5 or abs(middle(p) - middle(v)) <= 8

    votes = [p for p in ps if VOTE_FOR.match(p[4])]
    titles, title_ids = {}, set()
    for v in votes:
        lines = []
        # a listing that sets its titles in larger type can put a line or two between the title and the "Vote for"
        # line (the wards a district covers): the title is then the nearest line above in the larger type
        above = []
        for p in sorted([p for p in ps if p[0] > v[0] and aligned(p, v)], key=lambda p: p[0]):
            if p[0] - v[0] > 140 or VOTE_FOR.match(p[4]):
                break
            if p[3] >= v[3] + 0.8:
                above = [p]
                break
        above = above or [p for p in ps if 0 < p[0] - v[0] <= 2.2 * max(p[3], v[3]) and aligned(p, v)]
        while above and len(lines) < 3:
            top = min(above, key=lambda p: p[0])
            if VOTE_FOR.match(top[4]) or top[4].lower() in PARTY_WORDS:
                break
            lines.insert(0, top)
            above = [p for p in ps if 0 < p[0] - top[0] <= 1.6 * top[3] and abs(p[3] - top[3]) <= 0.3 and aligned(p, v)]
        titles[id(v)] = lines
        title_ids |= {id(t) for t in lines}
    out = []
    for v in votes:
        m = VOTE_FOR.match(v[4])
        cands, pending, last_y, ended, doubt = [], None, v[0], False, None
        under = [p for p in ps if p[0] < v[0] and aligned(p, v)]
        for j, p in enumerate(under):
            text = p[4]
            after = under[j + 1][4].lower() if j + 1 < len(under) else ""
            if id(p) in title_ids or VOTE_FOR.match(text):
                ended = True
                break
            if last_y - p[0] > 3.2 * p[3]:                         # a wide space: the block is over, unless what follows still reads as a candidate
                if NAME_DASH.match(text) or NAME_BRACKETS.match(text) or after in PARTY_WORDS:
                    doubt = "a wider space than expected inside the contest"
                ended = True
                break
            dash, both = NAME_DASH.match(text), NAME_BRACKETS.match(text)
            if pending is not None and text.lower() in PARTY_WORDS:
                cands.append([pending, text])
                pending = None
            elif pending is not None:
                pending, ended = None, True
                break                                              # a name with no party under it: not a candidate; the block is over
            elif dash:
                cands.append([(dash.group(1) or dash.group(3)).strip(), (dash.group(2) or dash.group(4)).strip()])
            elif both:
                cands.append([both.group(1).strip(), both.group(2).strip()])
            elif text.lower() in PARTY_WORDS:
                doubt, ended = "a party line with no name above it", True
                break
            else:
                pending = text
            last_y = p[0]
        out.append({"title": " ".join(t[4] for t in titles[id(v)]), "vote_for": int(NUMBER.get(m.group(1).lower(), m.group(1))), "write_ins": 0,
                    "candidates": [[c[0], party_words(c[1])] for c in cands], "runs_off_page": not ended, **({"doubt": doubt} if doubt else {})})
    return out


def for_block(ps):
    """What a ballot says it is for: {"county": name or None, "places": [[kind, name]]}, read from the lines under the
    word "for" in the ballot's own heading (Official Ballot / ... / November 3, 2026 / for / ...). None when the page
    has no such heading."""
    for f in [p for p in ps if p[4].strip().lower() == "for"]:
        xc = (f[1] + f[2]) / 2
        above = [p for p in ps if 0 < p[0] - f[0] < 40 and abs((p[1] + p[2]) / 2 - xc) < 60]
        if not above or not re.search(r"November \d+, 2026", min(above, key=lambda p: p[0])[4]):
            continue
        county, places = None, []
        for p in ps:
            if not (0 < f[0] - p[0] < 170 and abs((p[1] + p[2]) / 2 - xc) < 60):
                continue
            t = " ".join(p[4].split())
            if FOR_STOP.search(t):
                break
            m = COUNTY_LINE.match(t)
            if m:
                county = m.group(1).strip()
                continue
            a, b, c = MUNI_A.match(t), MUNI_B.match(t), MUNI_C.match(t)
            if a:
                places.append([MUNI_KIND[a.group(1).lower()], a.group(2).strip()])
            elif b:
                places.append([MUNI_KIND[b.group(2).lower()], b.group(1).strip().title()])
            elif c:
                places.append([MUNI_KIND[c.group(1).lower()], c.group(2).strip()])
        return {"county": county, "places": places}
    return None


def office_of(title):
    """(kind, office as shown, the court a judge's title names) for a contest title: kind None for a state or federal
    office, "" for a title this loader does not know."""
    t = re.sub(r"\s*[-\u2013\u2014]\s*Countywide\s*$", "", " ".join((title or "").split()), flags=re.I).strip()
    if STATE_OFFICE.search(t):
        return None, None, None
    bare = re.sub(r"^[A-Za-z][A-Za-z .]*? County (?=(?:Sheriff|Clerk|Coroner)\b)", "", t, flags=re.I)
    for pat, kind, shown in COUNTY_OFFICES:
        if pat.fullmatch(bare):
            return kind, shown, None
    m = MUNICIPAL_JUDGE.fullmatch(t)
    if m:
        court = f"{m.group(1).strip()} Municipal Court" if m.group(1) else None
        return "municipal_court", "Municipal Court Judge" if m.group(2) else "Municipal Judge", court
    return "", t, None


def ocr_layer(pdf, pages):
    """True when the file's text is a machine's reading laid invisibly over scanned pages (text render mode 3)."""
    for page, _res in pages:
        contents = pdf.get(page.get("Contents"))
        parts = contents if isinstance(contents, list) else [page.get("Contents")]
        data = b"\n".join(pdf.stream(p) or b"" for p in parts if isinstance(p, Ref))
        if re.search(rb"(?<![\d.])3\s+Tr\b", data):
            return True
    return False


def read_pdf(data, kind):
    """The allowed cells of a clerk's file, and nothing else of it: for a "ballot" file, one entry per ballot that
    carries a local contest ({page, for, contests}); for a "listing", its local contests ({page, contests}). State and
    federal contests are only counted, by title. Each page is read twice, by its positioned pieces and by its plain
    rows, and the two must agree on the number of contests and on every name kept."""
    if data[:4] != b"%PDF":
        raise Unreadable("the address did not answer with a PDF")
    pdf = PDF(data)
    pages = pdf.pages()
    if ocr_layer(pdf, pages):
        raise Unreadable("the file is a scan read by machine, which is not trusted for names")
    found, others, fors, tops, fronts, blanked, total = [], collections.Counter(), {}, {}, [], 0, 0
    for k, (page, res) in enumerate(pages, 1):
        runs = page_runs(pdf, page, res)
        ps = pieces(runs)
        cs = ballot_contests(ps, k) if kind == "ballot" else listing_contests(ps, k)
        plain = [join(rs) for _y, rs in row_groups([r for r in runs if r[3].strip()])]
        words = {w for t in plain for w in t.split()}
        if sum(len(VOTE_FOR_ANY.findall(t)) for t in plain) != len(cs):
            raise Unreadable(f"the two readings of page {k} do not find the same number of contests")
        total += len(cs)
        fors[k] = for_block(ps) if kind == "ballot" else None
        tops[k] = [p[4].lower() for p in ps[:2]]
        if any(p[4].lower().startswith("notice to voters") for p in ps):
            fronts.append(k)                                       # the notice to voters is printed at the head of a ballot's first page
        heading = next((p[4] for p in ps if re.fullmatch(r"Partisan Office\b[A-Za-z ,]*", p[4])), None)
        local = []
        for c in cs:
            okind, _shown, _court = office_of(c["title"])
            if okind is None:
                others[STATE_OFFICE.search(c["title"]).group(0).lower()] += 1
                continue
            if okind == "" and kind == "listing" and not LOCAL_WORD.search(c["title"]):
                others["a state or federal block whose title was not read"] += 1
                continue
            doubt = c.pop("doubt", None)
            if doubt and okind:
                raise Unreadable(f"{doubt} (page {k})")
            if kind == "ballot" and okind and c["write_ins"] < 1:
                raise Unreadable(f"a contest on page {k} has no write-in line under it, so its candidates may not be complete")
            if kind == "listing" and okind and not c["candidates"]:
                raise Unreadable(f"a county office on page {k} of the listing has no candidate that could be read under it")
            for cand in c["candidates"]:
                if not (any(cand[0] in t for t in plain) or all(w in words for w in cand[0].split())):
                    raise Unreadable(f"the two readings of page {k} do not agree on a name")
                if any(looks_like_contact(x) for x in cand if x):
                    blanked += 1
                    cand[0] = ""
            if looks_like_contact(c["title"]):
                blanked += 1
                continue
            local.append(c)
        if local:
            found.append({"page": k, "contests": local, **({"heading": heading} if heading and not looks_like_contact(heading) else {})})
    if not total:
        raise Unreadable("no contest could be read in the file (its pages may be pictures with no text)")
    for b in found:
        k = b["page"]
        if kind == "ballot":                                       # what the ballot is for is said on one of the ballot's own pages
            first = max([f for f in fronts if f <= k], default=None)
            if first is None:                                      # no notice found: a ballot is taken to be a leaf, front and back
                mine = [k, k + 1 if k % 2 else k - 1]
            else:
                mine = [k] + list(range(first, min([f for f in fronts if f > k], default=len(pages) + 1)))
            b["for"] = next((fors[j] for j in mine if fors.get(j)), None)
        for c in b["contests"]:                                    # a listing's block that runs to the foot of its page must not go on overleaf
            nxt = tops.get(k + 1) or []
            if c.pop("runs_off_page", False) and nxt and (nxt[0] in PARTY_WORDS or (len(nxt) > 1 and nxt[1] in PARTY_WORDS)):
                raise Unreadable(f"a contest on page {k} seems to run onto the next page")
    return {"pages": len(pages), "contests_read": total, "ballots" if kind == "ballot" else "pages_with_local": found,
            "other_contests": dict(sorted(others.items())), "blanked": blanked}


def kept_copy(entry, folder, refresh, say):
    """The kept copy of one clerk's file: fetched (and cut down in memory to the allowed cells) when there is none
    younger than a week, else read from disk. A file that did not answer or could not be read is remembered for a week
    too, so that a site which turns scripts away is not asked again on every run. Returns (kept, why not), exactly one
    of the two set."""
    cslug = slug(entry["county"])
    path = os.path.join(folder, cslug, f"{cslug}_2026_general_{entry['key']}.json")
    old = json.load(open(path, encoding="utf-8")) if os.path.exists(path) else None
    if old and old.get("url") != entry["url"]:
        old = None
    if old and not refresh and (time.time() - os.path.getmtime(path)) < MAX_AGE_DAYS * 86400:
        return (None, old["failed"]) if old.get("failed") else (old, None)
    good = old if old and not old.get("failed") else None
    why, data, read = None, b"", None
    try:
        time.sleep(1.5)
        data = net.get(entry["url"], timeout=180)
    except Exception as err:  # noqa: BLE001  unreachable or refused today
        code = getattr(err, "code", None)
        why = f"the county clerk's file did not answer when it was asked for ({type(err).__name__}{' ' + str(code) if code else ''})"
    else:
        try:
            read = read_pdf(data, entry["kind"])
        except Unreadable as err:
            why = str(err)
    if why and good:                                               # the copy on disk stands; it is asked for again in a week
        say(f"      {entry['county']} County: {why}; using the copy of {good['fetched']}")
        os.utime(path)
        return good, None
    today = dt.date.today().isoformat()
    if why:
        kept = {"county": entry["county"], "kind": entry["kind"], "url": entry["url"], "fetched": today, "failed": why}
    else:
        kept = {"county": entry["county"], "kind": entry["kind"], "title": entry["title"], "url": entry["url"], "page": entry["page"],
                "fetched": today, "sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data),
                "kept": "office titles, the names and parties under them, and the municipality and county a ballot is for; nothing else of the file", **read}
    del data
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path + ".part", "w", encoding="utf-8") as fh:
        json.dump(kept, fh, ensure_ascii=False, indent=1)
    os.replace(path + ".part", path)
    return (None, why) if why else (kept, None)


def saved_ballots(folder, county, page, say):
    """Sample ballots saved by hand for a county whose site refuses scripts: every PDF in the county's folder, read
    in place as a ballot file. The allowed cells are kept beside it (<name>.kept.json, with the file's SHA-256) so that
    a large file is read once. [(entry, kept or None, why not)]."""
    out = []
    for path in sorted(glob.glob(os.path.join(folder, slug(county), "*.pdf"))):
        base = os.path.basename(path)
        title = f"Sample ballot saved by hand ({base})"
        entry = {"county": county, "kind": "ballot", "key": slug(os.path.splitext(base)[0]) or "saved",
                 "title": "Sample ballot saved by hand" if contact_like(title, False) else title, "url": page, "page": page, "saved": True}
        data = open(path, "rb").read()
        digest, side = hashlib.sha256(data).hexdigest(), path + ".kept.json"
        kept = json.load(open(side, encoding="utf-8")) if os.path.exists(side) else None
        if not kept or kept.get("sha256") != digest:
            try:
                read = read_pdf(data, "ballot")
            except Unreadable as err:
                say(f"      {county} County: the saved file {base} could not be read as a sample ballot ({err})")
                out.append((entry, None, f"a file saved by hand could not be read as a sample ballot: {err}"))
                continue
            kept = {"county": county, "kind": "ballot", "title": entry["title"], "url": page, "page": page, "fetched": mtime(path), "sha256": digest,
                    "bytes": len(data), "kept": "office titles, the names and parties under them, and the municipality and county a ballot is for", **read}
            with open(side + ".part", "w", encoding="utf-8") as fh:
                json.dump(kept, fh, ensure_ascii=False, indent=1)
            os.replace(side + ".part", side)
        del data
        out.append((entry, kept, None))
    return out


def county_names(path):
    """{five-digit code: (name, name with its kind word)} for Wisconsin's 72 counties, from the Census county file."""
    import shapefile                                                        # pyshp, in the kit's environment
    z = zipfile.ZipFile(path)
    base = next(n for n in z.namelist() if n.endswith(".dbf"))
    out = {}
    for r in (x.as_dict() for x in shapefile.Reader(dbf=io.BytesIO(z.read(base))).iterRecords()):
        if r["STATEFP"] == FIPS:
            out[str(r["GEOID"])] = (str(r["NAME"]), str(r["NAMELSAD"]))
    if len(out) != 72 or any(not re.fullmatch(FIPS + r"\d{3}", f) for f in out):
        raise SystemExit(f"Wisconsin: the Census county file gives {len(out)} counties for the state, not 72")
    return out


def cousub(folder, say):
    """The Census Bureau's 2020 county subdivisions of Wisconsin (cities, villages and towns): (path, {(kind, folded
    name): {"code", "name", "counties"}}). The file has no contact columns and is kept whole."""
    path = os.path.join(folder, COUSUB_FILE)
    net.download(COUSUB_URL, path, 3650, say=say)
    out = {}
    with open(path, encoding="utf-8", errors="replace") as fh:
        head = fh.readline().rstrip("\r\n").split("|")
        if head[:7] != ["STATE", "STATEFP", "COUNTYFP", "COUNTYNAME", "COUSUBFP", "COUSUBNS", "COUSUBNAME"]:
            raise SystemExit(f"Wisconsin: {COUSUB_FILE}: the header is not the one this loader was checked against; stopping")
        for line in fh:
            f = line.rstrip("\r\n").split("|")
            if len(f) < 7 or f[1] != FIPS:
                continue
            name, _, kind = f[6].rpartition(" ")
            if kind not in ("city", "village", "town"):
                continue
            e = out.setdefault((kind, key(name), f[4]), {"code": f[4], "name": f[6], "counties": set()})
            e["counties"].add(FIPS + f[2])
    return path, out


def local_level(folder, county_zip, say=print, refresh=False):
    """Wisconsin's county contests on the November ballot, as rows ready to write (races, candidates, places, sources,
    gaps, notes) with the counts behind them. Nothing here touches the database."""
    counties = county_names(county_zip)
    by_name = {key(n): f for f, (n, _full) in counties.items()}
    full = {f: fl for f, (_n, fl) in counties.items()}
    os.makedirs(folder, exist_ok=True)
    cpath, places_census = cousub(folder, say)
    problems = []

    # ---- every file: the table's, and any sample ballot saved by hand
    files, failed = [], collections.defaultdict(list)                # files: (entry, kept); failed: county -> [why]
    for entry in COUNTY_FILES:
        kept, why = kept_copy(entry, folder, refresh, say)
        if kept:
            files.append((entry, kept))
        else:
            failed[by_name[key(entry["county"])]].append(why)
            problems.append(f"{entry['county']} County ({entry['key']}): {why}")
    for f, (name, _fl) in sorted(counties.items()):
        for entry, kept, why in saved_ballots(folder, name, WHY.get(name, (None, None))[1], say):
            if kept:
                files.append((entry, kept))
            else:
                failed[f].append(why)
                problems.append(f"{name} County: {why}")

    # ---- which county each ballot is of, and what it says about the county's offices
    seen = collections.defaultdict(lambda: collections.defaultdict(list))   # county -> office kind -> [(candidates, file index, page)]
    judges = collections.defaultdict(list)                                    # (court or None, title) -> [(county, candidates, file index, page, for, heading)]
    unknown = collections.defaultdict(collections.Counter)                    # county -> title -> ballots
    used = collections.defaultdict(lambda: collections.Counter())             # file index -> county -> ballots that gave a county office
    lines_read = collections.Counter()                                        # file index -> candidate lines read from it
    blanked = 0
    for i, (entry, kept) in enumerate(files):
        home = by_name[key(entry["county"])]
        blanked += kept.get("blanked", 0)
        for b in kept.get("ballots") or kept.get("pages_with_local") or []:
            where = b.get("for") or {}
            f = home
            if where.get("county"):
                f = by_name.get(key(where["county"]))
                if f is None:
                    problems.append(f"{entry['county']} County ({entry['key']}), page {b['page']}: the ballot names a county the Census file does not have")
                    continue
            for c in b["contests"]:
                okind, shown, court = office_of(c["title"])
                cands = tuple((n, p) for n, p in c["candidates"])
                lines_read[i] += len(cands)
                if okind == "municipal_court":
                    judges[(court, shown)].append((f, cands, i, b["page"], where, b.get("heading") or ""))
                elif okind:
                    seen[f][okind].append((cands, i, b["page"], c["vote_for"]))
                    used[i][f] += 1
                else:
                    unknown[f][shown] += 1

    races, cands_out, sources, gaps, places = [], [], [], [], []
    loaded, per_file_rows, how_read = {}, collections.Counter(), {}
    office_shown = {k: s for _p, k, s in COUNTY_OFFICES}
    file_sid = {}
    for i, (entry, _kept) in enumerate(files):
        tail = "sample-ballot" if entry["kind"] == "ballot" else "candidate-listing"
        file_sid[i] = f"wi-{slug(entry['county'])}-clerk-2026-general-{tail}-{entry['key']}"

    for f in sorted(seen):
        name, cfull = counties[f][0], full[f]
        bad = []
        for okind, forms in seen[f].items():
            if len({x[0] for x in forms}) != 1:
                bad.append(f"the ballots read do not agree on the candidates for {office_shown[okind].lower()}")
            if any(x[3] != 1 for x in forms):
                bad.append(f"{office_shown[okind].lower()} is not a vote-for-one contest on a ballot read")
            for n, p in forms[0][0]:
                if not n:
                    bad.append("a name on a ballot read like contact details and was set aside")
                elif p is None:
                    bad.append(f"a candidate for {office_shown[okind].lower()} is printed with no party")
        if bad:
            failed[f] += bad
            problems += [f"{cfull}: {b}" for b in bad]
            continue
        own = sorted({x[1] for forms in seen[f].values() for x in forms if key(files[x[1]][0]["county"]) == key(name)})
        lent = sorted({x[1] for forms in seen[f].values() for x in forms} - set(own))
        first = (own or lent)[0]
        n_ballots = max(len(forms) for forms in seen[f].values())
        how_read[f] = (own, lent, n_ballots)
        note = None
        if not own:                                                 # read only from a neighbouring clerk's file
            lender = files[first][0]["county"]
            where = files[first][1]
            muni = next((b.get("for") or {} for b in where.get("ballots", []) if (b.get("for") or {}).get("county") and by_name.get(key(b["for"]["county"])) == f), {})
            town = "; ".join(f"the {k} of {n}" for k, n in muni.get("places", [])) or "a municipality that lies in both counties"
            note = (f"Read from the sample ballots the {lender} County Clerk posts, which include the ballot of {town} for the wards in {cfull}. "
                    "A countywide office has the same candidates on every ballot of the county.")
        for okind in ("sheriff", "clerk_of_court", "coroner"):
            if okind not in seen[f]:
                continue
            listed = seen[f][okind][0][0]
            rid = f"2026-{STATE}-{f}-{okind.replace('_', '-')}"
            ordered = files[first][0]["kind"] == "ballot"
            rnote = [note] if note else []
            if not listed:
                rnote.append("The ballot prints no candidate for this office, only a write-in line.")
            races.append((rid, STATE, "county", okind, office_shown[okind], cfull, f, json.dumps([f]), None, None, 0, 1, None, None, None, GENERAL,
                          " ".join(rnote) or None))
            for k, (n, p) in enumerate(listed, 1):
                cnote = None
                if p.lower() not in PARTY_WORDS:
                    cnote = ("Not the candidate of a party with a ballot line: the words shown as the party are what the county clerk's file prints with "
                             "the name. An independent candidate's line carries the party or principle from the nomination papers, in five words or less "
                             "(Wis. Stat. 5.64 (1) (e)).")
                cands_out.append((rid, "general", GENERAL, n, p, party_code(p), k if ordered else None, 0, 0, None, None, None, None, file_sid[first], cnote))
                per_file_rows[first] += 1
        loaded[f] = sorted(seen[f])
        missing = [office_shown[k] for k in ("sheriff", "clerk_of_court") if k not in seen[f]]
        if missing:
            problems.append(f"{cfull}: no contest for {' or '.join(missing)} was found in the file read, though state law has every county elect one")
        for title, n in sorted(unknown.get(f, {}).items()):
            gaps.append((STATE, "county", f, cfull, f"a contest titled {title}"[:80],
                         f"The county's ballots carry a contest under this title on {n} ballot{'s' if n != 1 else ''} read, which this loader does not "
                         "know how to file, so it is not shown here; the county clerk's sample ballot is the authority.", files[first][0]["page"]))
            problems.append(f"{cfull}: a contest title this loader does not know is on {n} ballots (named in sl_gaps)")

    # ---- municipal judges on these ballots: special elections to fill a vacancy (judges are elected in April)
    judge_races = []
    for (court, shown), hits in sorted(judges.items(), key=lambda kv: (kv[0][0] or "", kv[0][1])):
        groups = collections.defaultdict(list)
        for hit in hits:
            if court:                                              # a court named in the title: one contest wherever it is printed
                groups[court].append(hit)
            else:                                                  # "Municipal Judge" alone: the court of the municipality the ballot is for
                muni = (hit[4] or {}).get("places") or []
                groups[tuple(muni[0]) if len(muni) == 1 else None].append(hit)
        for who, hs in groups.items():
            home = sorted({h[0] for h in hs})
            first = hs[0][2]
            page_url = files[first][0]["page"]
            if who is None or len({h[1] for h in hs}) != 1 or any(not n for n, _p in hs[0][1]):
                for f in home:
                    gaps.append((STATE, "county", f, full[f], "a special election for municipal judge",
                                 "A ballot of this county carries a contest for municipal judge, but the ballots read do not say for which single "
                                 "municipality or do not agree on its candidates, so it is not shown here; the county clerk's sample ballot is the authority.",
                                 page_url))
                problems.append(f"{', '.join(full[f] for f in home)}: a municipal judge contest could not be tied to one court (named in sl_gaps)")
                continue
            reach, named = set(home), []
            for h in hs:
                for kind, nm in (h[4] or {}).get("places") or []:
                    fits_ = [e for (k2, n2, _code), e in places_census.items() if k2 == kind and n2 == key(nm) and h[0] in e["counties"]]
                    if len(fits_) == 1:
                        reach |= fits_[0]["counties"]
                        named.append(fits_[0])
            law = ("municipal judges are elected in April, and a vacancy is filled by special election for the rest of the term "
                   "(Wis. Stat. 755.01 (1), 8.50 (4) (fm))")
            said = all("special" in h[5].lower() for h in hs)      # the ballots' own heading: "Partisan Office, Special Judicial, and Referendum"
            rnote = (f"A special election, as the heading of the ballots that carry it says (\"{hs[0][5]}\"): {law}." if said else
                     f"On the November ballot as a special election: {law}. The ballot itself does not print the word special.")
            if isinstance(who, str):
                jname, jid, pkind = who, f"{STATE}-X-{slug(who)}", "special"
                rid = f"2026-{jid}-S"
                printed = sorted({e["name"] for e in named})
                rnote += (f" Read from the ballots of {and_words(printed) if printed else 'the municipalities the court serves'} in "
                          f"{and_words(full[f] for f in home)}; the court may also serve places whose ballots are not loaded here.")
                psrc = file_sid[first]
            else:
                kind, nm = who
                fits_ = [e for (k2, n2, _code), e in places_census.items() if k2 == kind and n2 == key(nm) and home[0] in e["counties"]]
                if len(home) != 1 or len(fits_) != 1:
                    code, jname = f"{home[0][2:]}-{slug(nm)}-{kind}", f"{nm} {kind}"
                    psrc = file_sid[first]
                else:
                    code, jname, psrc = fits_[0]["code"], fits_[0]["name"], SRC_COUSUB
                    reach = set(fits_[0]["counties"]) | set(home)
                jid, pkind = f"{STATE}-M-{code}", "mcd"
                rid = f"2026-{jid}-municipal-court-S"
            listed = hs[0][1]
            if not listed:
                rnote += " The ballot prints no candidate for this office, only a write-in line."
            judge_races.append(rid)
            races.append((rid, STATE, "court", "municipal_court", shown, jname, jid, json.dumps(sorted(reach)), None, None, 1, 0, None, None, None, GENERAL, rnote))
            places.append((pkind, jid, jname, json.dumps(sorted(reach)), psrc))
            for k, (n, _p) in enumerate(listed, 1):
                cands_out.append((rid, "general", GENERAL, n, NONPARTISAN, "N", k, 0, 0, None, None, None, None, file_sid[first], None))
                per_file_rows[first] += 1

    # ---- sources: one row for every file read
    for i, (entry, kept) in enumerate(files):
        home = by_name[key(entry["county"])]
        n_b = len(kept.get("ballots") or kept.get("pages_with_local") or [])
        state_n = sum(kept.get("other_contests", {}).values())
        lent = sorted(full[f] for f in used[i] if f != home)
        if entry["kind"] == "ballot":
            note = (f"{kept['pages']} pages; {n_b} ballot{'s' if n_b != 1 else ''} with a county contest. Read from each: the county contests (office, names and "
                    "parties in printed order) and the line saying which municipality and county the ballot is for. "
                    + (f"Every ballot of {entry['county']} County in the file gives the same candidates for each county office. " if home in loaded and n_b > 1 else "")
                    + (f"The file also holds the ballot of a municipality that lies partly in {and_words(lent)}; those contests are filed under that county. " if lent else "")
                    + f"The {state_n} state and federal contests on these ballots are only counted (the Elections Commission's list is their source). "
                    "A sample ballot carries no addresses, phones or e-mail; only the cells named here are kept, and the file itself is not saved. "
                    "Each page was read twice, by position and as plain rows, and the two readings agree.")
            kind = "official sample ballot"
        else:
            note = (f"{kept['pages']} pages. Read: the county offices' blocks only (office, \"Vote for\", then each name and party). A listing does not state the "
                    f"printed order, so no ballot order is stored. The {state_n} state and federal contests in it are only counted. "
                    "Nothing else of the file is kept, whatever its other pages hold (a clerk's packet can list polling places with their street addresses), "
                    "and the file itself is not saved. Each page was read twice, by position and as plain rows, and the two readings agree.")
            kind = "official candidate list"
        if entry.get("saved"):
            note = "Saved by hand from a browser, because the county's website turns scripts away. " + note
        if home not in loaded and not lent:
            note = "Not used: " + "; ".join(failed.get(home, ["the county's contests could not be read"])) + ". " + note
        if home in loaded and not per_file_rows[i]:
            note = "Read as a check: its candidates are the same as the other file's for the county. " + note
        sources.append((file_sid[i], STATE, kind, f"{entry['county']} County Clerk", entry["title"], entry["url"], "", kept["fetched"], kept["sha256"],
                        lines_read[i], note))
    sources.append((SRC_COUSUB, STATE, "official place codes", "U.S. Census Bureau", "2020 county subdivision codes, Wisconsin (st55_wi_cousub2020.txt)",
                    COUSUB_URL, "2020", mtime(cpath), sha(cpath), len(places_census),
                    "Names and codes of Wisconsin's cities, villages and towns and the counties each lies in: the code of a municipality whose court a "
                    "ballot names, and the counties a municipality on a ballot reaches. The file has no contact columns."))

    # ---- places and gaps
    places = [("county", f, fl, json.dumps([f]), SRC_COUNTY) for f, fl in sorted(full.items())] + sorted(set(places))
    not_loaded = [f for f in sorted(counties) if f not in loaded]
    for f in not_loaded:
        name = counties[f][0]
        why, page = WHY.get(name, (None, None))
        reason = REASON[why]
        if failed.get(f):
            reason = (f"A file for this county was read but could not be used ({'; '.join(sorted(set(failed[f])))}), so this county's list is not loaded yet; "
                      "the county clerk's sample ballot is the authority.")
            page = page or next((e["page"] for e in COUNTY_FILES if e["county"] == name), None)
        gaps.append((STATE, "county", f, full[f], WHAT, reason, page))
    gaps.append((STATE, "state", STATE, "Wisconsin", "special elections for local offices in counties not loaded",
                 "A city, village, town or school district can call a special election for November 3, and a vacancy for municipal judge is filled that "
                 "way. Such a contest is printed only on the ballots of the places it concerns, so it is shown here only where a county's sample ballots "
                 "were read; in the other counties the county clerk's sample ballot is the authority.", SPECIAL_URL))

    by_kind = collections.Counter(r[3] for r in races if r[2] == "county")
    n_county_cands = sum(1 for c in cands_out if c[0] not in judge_races)
    ballots_from = sorted({files[i][0]["county"] for i in range(len(files)) if files[i][0]["kind"] == "ballot" and by_name[key(files[i][0]["county"])] in loaded})
    listings_from = sorted({files[i][0]["county"] for i in range(len(files)) if files[i][0]["kind"] == "listing" and by_name[key(files[i][0]["county"])] in loaded})
    lent_to = sorted(counties[f][0] for f, (own, _l, _n) in how_read.items() if not own)
    judge_words = ""
    if judge_races:
        n_j = len(judge_races)
        judge_words = (f" {'One special election' if n_j == 1 else NUMBER_WORDS.get(n_j, str(n_j)).capitalize() + ' special elections'} for municipal judge printed on "
                       f"those ballots {'is' if n_j == 1 else 'are'} loaded too, under the courts.")
    notes = [
        (STATE, "local_calendar",
         "On November 3, 2026 each Wisconsin county elects its sheriff and its clerk of circuit court for four years, and a county that still has a coroner "
         "rather than a medical examiner elects the coroner; these are the only county offices on this ballot, and they are on its partisan part. "
         "County clerks, treasurers, registers of deeds, elected surveyors and district attorneys are next elected in November 2028. Cities, villages, towns, "
         "school boards, county boards, county executives and judges are elected at the spring election in April, so one of those offices is on this ballot only "
         "where a special election was called.",
         "Wis. Stat. 59.20 (2) (a), (b) and (bm); 978.01 (1); 5.02 (5) and (21); 8.50 (4) (fm)", STATUTE_URL),
        (STATE, "local_coverage",
         f"Wisconsin has no statewide list of county candidates: each of the 72 county clerks publishes the county's own sample ballots or candidate listing. "
         f"Loaded here: {len(races) - len(judge_races)} county contests ({by_kind.get('sheriff', 0)} for sheriff, {by_kind.get('clerk_of_court', 0)} for clerk of circuit "
         f"court, {by_kind.get('coroner', 0)} for coroner) with {n_county_cands} candidates in {len(loaded)} of the 72 counties, read from the clerks' own files: "
         f"sample ballots for {and_words(ballots_from) or 'none'}"
         + (f" (the {and_words(lent_to)} contests from the ballot of a shared municipality in a neighbouring clerk's file)" if lent_to else "")
         + (f", and candidate listings for {and_words(listings_from)}" if listings_from else "") + "." + judge_words
         + f" Left out: ballot questions, the write-in lines, anyone not printed on the ballot, and the {len(not_loaded)} counties whose list is not loaded yet, "
         "each named among the gaps with the reason. A sample ballot shows who is printed on the ballot, so a candidate who withdrew before printing is simply "
         "absent and cannot be counted; a clerk's listing gives no ballot order.",
         "The county clerks' sample ballots and candidate listings named among the sources", None),
    ]

    # ---- the last look before anything is written: nothing that reads like contact details, by the trial check's test and the page builder's
    for table, strict, items in (("sl_races", True, [(r[0], (r[4], r[5], r[8], r[9], r[16])) for r in races]),
                                 ("sl_candidates", True, [(c[0], (c[3], c[4], c[14])) for c in cands_out]),
                                 ("sl_places", True, [(p[1], (p[2],)) for p in places]),
                                 ("sl_gaps", False, [(g[2], (g[3], g[4], g[5])) for g in gaps]),
                                 ("sl_notes", False, [(x[1], (x[2], x[3])) for x in notes]),
                                 ("sl_sources", False, [(s[0], (s[3], s[4], s[10])) for s in sources])):
        for ident, texts in items:
            if any(t and (contact_like(t, strict) or (strict and BUILDER_STREET.search(str(t)))) for t in texts):
                raise SystemExit(f"Wisconsin: a text for {table} ({ident}) reads like contact details; stopping (the text is not printed)")
    if len({(c[0], c[3]) for c in cands_out}) != len(cands_out):
        raise SystemExit("Wisconsin: a candidate is stored twice in one county contest; stopping")

    return {"races": races, "cands": cands_out, "places": places, "sources": sources, "gaps": gaps, "notes": notes, "problems": problems,
            "loaded": loaded, "how_read": how_read, "counties": counties, "judges": judge_races, "files": len(files), "blanked": blanked,
            "not_loaded": len(not_loaded), "by_kind": dict(by_kind), "county_cands": n_county_cands,
            "state_contests": sum(sum(k.get("other_contests", {}).values()) for _e, k in files), "lines_read": sum(lines_read.values()),
            "ballots_read": sum(len(k.get("ballots") or k.get("pages_with_local") or []) for _e, k in files)}


NUMBER_WORDS = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six"}


def and_words(items):
    items = [str(x) for x in items]
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1] if items else ""


def load(db_path, say=print, cache=CACHE, roster_db=ROSTER, county_zip=COUNTY_ZIP, local_cache=None, refresh=False):
    """Wisconsin's rows into the database at db_path: the state races, then the county ones. local_cache is the folder
    for the kept copies of the clerks' files (ballot_cache/wi/local/ unless told otherwise); refresh=True fetches the
    clerks' files again even when a kept copy is less than a week old."""
    folder = os.path.join(cache, "wi")
    lpath, bpath = os.path.join(folder, LIST), os.path.join(folder, BOOK)
    if not os.path.exists(lpath):
        say(f"    Wisconsin: {LIST} is not in {folder}. elections.wi.gov refuses scripts; the file is carried out of a browser "
            "(\"Candidates on Ballot By Election_November 3 2026 General Election.pdf\" on elections.wi.gov/elections).")
        return 0
    offices, printed = general_list(lpath)
    book = primary_book(bpath) if os.path.exists(bpath) else {}
    legs, offs = roster(roster_db)
    counties = census_counties(county_zip)
    report = []

    state = [o for o in offices if race_of(o[0])]
    heads = {o[0] for o in state}
    senate = sorted(int(SENATE.fullmatch(h).group(1)) for h in heads if SENATE.fullmatch(h))
    assembly = sorted(int(ASSEMBLY.fullmatch(h).group(1)) for h in heads if ASSEMBLY.fullmatch(h))
    if senate != list(range(1, 34, 2)):
        raise SystemExit(f"Wisconsin: the list's Senate seats are not the seventeen odd-numbered ones: {senate}")
    if assembly != list(range(1, 100)):
        raise SystemExit(f"Wisconsin: the list's Assembly seats are not all 99: {len(assembly)} listed")
    missing = [o for o in ON_LIST if o not in heads]
    if missing:
        raise SystemExit(f"Wisconsin: statewide offices missing from the list: {missing}")
    if len(state) != len(heads):
        raise SystemExit("Wisconsin: an office is listed twice")

    races, cands = {}, []
    ticket = {}                                                             # party -> (governor, lieutenant governor)

    def county_ids(office):
        names = book.get(office, {}).get("counties", set())
        bad = sorted(n for n in names if n not in counties)
        if bad:
            raise SystemExit(f"Wisconsin: county names in the primary report not in the Census file: {bad}")
        return ",".join(sorted(counties[n][0] for n in names)) or None

    def holder_of(rid, level, district, okind):
        if level == "legislature":
            chamber = "Senate" if okind == "state_senate" else "House"
            hs = [p for p in legs if p["chamber"] == chamber and str(p["district"]) == district]
            return hs[0] if len(hs) == 1 else None
        rk = ROSTER_OFFICE.get(rid.rsplit("-", 1)[1])
        return offs.get(rk) if rk else None

    def new_race(office, inc, noncand):
        rid, level, okind, shown, district = race_of(office)
        h = holder_of(rid, level, district, okind)
        note = []
        if h is None and level == "legislature":
            report.append(f"{rid}: no sitting member in the roster for this seat")
            note.append("The roster shows no sitting member for this seat.")
        hname, hparty, hid = (h["full"], h["party"], h["id"]) if h else (inc or None, None, None)
        if h is None and inc and level == "statewide":
            note.append("Holder as named on the Elections Commission's list; the Open States roster does not carry this office.")
        if h and inc and not person_fits(inc, h):
            report.append(f"{rid}: the Commission names {inc} as incumbent; the roster's holder is {h['full']}")
        if noncand and inc:
            note.append(f"{inc}, who holds the seat, filed a notification of noncandidacy.")
        if rid.endswith("-GOV"):
            note.append("Wisconsin elects the governor and lieutenant governor together in November, one vote for the pair; "
                        "each party nominates them in separate primaries (the lieutenant governor's under 2026-WI-LTG).")
        jur = "Wisconsin" if level == "statewide" else (f"Senate District {district}" if okind == "state_senate" else f"Assembly District {district}")
        races[rid] = [rid, STATE, level, okind, shown, jur, FIPS if level == "statewide" else district,
                      None if level == "statewide" else county_ids(office), district, None, 0, 1, hid, hname, hparty, GENERAL,
                      " ".join(note) or None]
        return rid, level, okind, h, inc

    def identify(rid, level, okind, h, inc, name, party):
        """(incumbent, state_member_id, note) for one listed name."""
        if h is not None and person_fits(name, h):
            return 1, h["id"], None
        if level == "statewide" and h is None and inc and person_fits(name, {"first": " ".join(name_parts(inc)[0]), "last": name_parts(inc)[1]}):
            return 1, None, "Named as the incumbent on the Elections Commission's list."
        pool = [p for p in legs if person_fits(name, p) and (p["party"] or "") == (party or "")]
        if len(pool) == 1 and not (h is not None and pool[0]["id"] == h["id"]):
            return 0, pool[0]["id"], f"Serves today in the {chamber_words(pool[0])}."
        return 0, None, None

    for office, inc, noncand, listed in state:
        rid, level, okind, h, inc = new_race(office, inc, noncand)
        for order, name, party in listed:
            note, mine = None, name
            if rid.endswith("-GOV"):
                gov, _, ltg = name.partition(" / ")
                if not ltg:
                    raise SystemExit(f"Wisconsin: a governor's line without a running mate: {name!r}")
                ticket[party] = (gov.strip(), ltg.strip())
                mine, note = gov.strip(), "Governor and lieutenant governor on one ticket, as the list prints them."
            incb, mid, n2 = identify(rid, level, okind, h, inc, mine, party)
            cands.append([rid, "general", GENERAL, name, party, party_code(party), order, incb, 0, None, None, None, mid, SRC_LIST,
                          " ".join(x for x in (note, n2) if x) or None])

    # The lieutenant governor: nominated in its own primary, elected on the governor's ticket in November.
    if "LIEUTENANT GOVERNOR" in book:
        new_race("LIEUTENANT GOVERNOR", "", False)
        races["2026-WI-LTG"][16] = ("Nominated in each party's own primary on August 11; in November the lieutenant governor is elected "
                                   "jointly with the governor, on the tickets listed under 2026-WI-GOV.")

    # Primary fields: two or more names on a party's sheet (registered write-ins count), official votes only.
    nominees = {}
    for rid in races:
        for c in cands:
            if c[0] == rid:
                nominees[(rid, c[4])] = c[3]
    for p, (gov, ltg) in ticket.items():
        nominees[("2026-WI-GOV", p)], nominees[("2026-WI-LTG", p)] = gov, ltg
    fields = 0
    for office, entry in book.items():
        rid = race_of(office)[0]
        if rid not in races:
            report.append(f"{rid}: in the primary report but not on the November list")
            continue
        _r, level, okind, _s, district = race_of(office)
        h = holder_of(rid, level, district, okind)
        for party, (code, total, field) in entry["fields"].items():
            if len(field) < 2:
                continue
            fields += 1
            nominee = nominees.get((rid, party))
            won = [n for n, _w, _v in field if nominee and person_fits(n, {"first": " ".join(name_parts(nominee)[0]), "last": name_parts(nominee)[1]})]
            top = max(field, key=lambda f: f[2])[0]
            if len(won) != 1:
                report.append(f"{rid} {party} primary: the November candidate ({nominee or 'none'}) is not one name in the field; outcome left blank")
            elif won[0] != top:
                report.append(f"{rid} {party} primary: {won[0]} is on the November list but {top} had the most votes")
            for name, wi, votes in field:
                incb, mid, n2 = identify(rid, level, okind, h, "", name, party)
                outcome = None if len(won) != 1 else ("advanced" if name == won[0] else "lost")
                cands.append([rid, f"primary-{code}", PRIMARY, name, party, party_code(party), None, incb, int(wi), votes,
                              round(100 * votes / total, 1) if total else None, outcome, mid, SRC_BOOK,
                              " ".join(x for x in ("Registered write-in candidate." if wi else None, n2) if x) or None])

    # Every November candidate of a party that held a primary should be among that party's primary names.
    for c in cands:
        if c[1] != "general" or c[4] not in PRIMARY_CODE:
            continue
        office = next((o for o in book if race_of(o)[0] == c[0]), None)
        name = ticket.get(c[4], (c[3],))[0] if c[0].endswith("-GOV") else c[3]
        names = [n for n, _w, _v in book.get(office, {}).get("fields", {}).get(c[4], ("", 0, []))[2]] if office else []
        if not any(person_fits(n, {"first": " ".join(name_parts(name)[0]), "last": name_parts(name)[1]}) for n in names):
            report.append(f"{c[0]}: {name} ({c[4]}) is on the November list but not among that party's primary names")

    seen = set()
    for c in cands:
        k = (c[0], c[1], c[3])
        if k in seen:
            raise SystemExit(f"Wisconsin: {c[3]} is listed twice in {c[0]} {c[1]}")
        seen.add(k)
    for rid in races:
        if rid != "2026-WI-LTG" and not any(c[0] == rid and c[1] == "general" for c in cands):
            report.append(f"{rid}: no candidates on the November list")

    general = [c for c in cands if c[1] == "general"]
    primary_rows = [c for c in cands if c[1] != "general"]

    # ---- the county offices, county by county (the clerks' own files; nothing here changes a state row)
    local = local_level(local_cache or (LOCAL_DIR if cache == CACHE else os.path.join(cache, "wi", "local")), county_zip, say=say, refresh=refresh)
    clash = sorted({r[0] for r in local["races"]} & set(races))
    if clash:
        raise SystemExit(f"Wisconsin: a county contest shares a race id with a state race ({clash[0]}); stopping")

    con = sqlite3.connect(db_path)
    con.executescript(SCHEMA)
    con.executescript(EXTRA_SCHEMA)
    with con:
        # Wisconsin's rows only: its races and candidates by state and race id, its sources, gaps and notes by state,
        # its places by their wi- source ids (the counties' codes begin 55, so the id cannot pick them out)
        con.execute("DELETE FROM sl_candidates WHERE race_id IN (SELECT race_id FROM sl_races WHERE state = ?)", (STATE,))
        con.execute("DELETE FROM sl_candidates WHERE race_id LIKE ?", (f"2026-{STATE}-%",))
        con.execute("DELETE FROM sl_races WHERE state = ?", (STATE,))
        con.execute("DELETE FROM sl_places WHERE source_id LIKE 'wi-%'")
        con.execute("DELETE FROM sl_sources WHERE state = ?", (STATE,))
        con.execute("DELETE FROM sl_gaps WHERE state = ?", (STATE,))
        con.execute("DELETE FROM sl_notes WHERE state = ?", (STATE,))
        con.executemany("INSERT INTO sl_races VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", list(races.values()) + local["races"])
        con.executemany("INSERT INTO sl_candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", cands + local["cands"])
        con.executemany("INSERT INTO sl_places VALUES (?,?,?,?,?)", local["places"])
        src = [
            (SRC_LIST, STATE, "official candidate list", "Wisconsin Elections Commission",
             "Candidates on Ballot by Election: 2026 General Election, 11/3/2026", LIST_URL, printed, mtime(lpath), sha(lpath), len(general),
             "State offices only (the congressional lines are the federal loader's). Ballot order as printed; each office's count "
             "checked against the list's total. elections.wi.gov refuses scripts; the file was carried out of a browser."),
            (SRC_ROSTER, STATE, "roster", "Open States people project (CC0), as loaded into state_wi.sqlite",
             "Sitting legislators and statewide officials", "https://github.com/openstates/people", "", mtime(roster_db), sha(roster_db),
             len(legs) + len(offs), "Who holds each seat today; the roster does not carry the Secretary of State or the State Treasurer."),
            (SRC_COUNTY, STATE, "county codes", "U.S. Census Bureau", "Cartographic boundary file, counties, 2024 (1:500,000)",
             COUNTY_URL, "2024", mtime(county_zip), sha(county_zip),
             len(counties), "Five-digit county codes (GEOID) and names for Wisconsin's 72 counties: the counties the primary report lists under each "
             "district, and the county of each county office."),
        ]
        if book:
            src.append((SRC_BOOK, STATE, "official results", "Wisconsin Elections Commission",
                        "County by County Report: 2026 Partisan Primary, all state contests (August 11, 2026)", BOOK_URL, "", mtime(bpath),
                        sha(bpath), len(primary_rows),
                        "Votes from each office's Office Totals row; every sheet's county rows and candidates reconciled to its totals. "
                        "Write-ins nobody registered (SCATTERING) count in the total but are not listed; registered write-ins are."))
        con.executemany("INSERT INTO sl_sources VALUES (?,?,?,?,?,?,?,?,?,?,?)", src + local["sources"])
        con.executemany("INSERT INTO sl_gaps VALUES (?,?,?,?,?,?,?)", local["gaps"])
        con.executemany("INSERT INTO sl_notes VALUES (?,?,?,?,?)", local["notes"])
    con.close()

    by = lambda kind: sum(1 for c in general if races[c[0]][3] == kind)
    say(f"    Wisconsin: {len(races)} races ({len(senate)} Senate, {len(assembly)} Assembly, {len(races) - len(senate) - len(assembly)} statewide); "
        f"{len(general)} candidates on the November list (Senate {by('state_senate')}, Assembly {by('state_house')}, statewide "
        f"{len(general) - by('state_senate') - by('state_house')}); {fields} primary fields, {len(primary_rows)} primary rows")
    for c in cands:
        if c[12]:
            say(f"      matched: {c[0]} {c[1]}: {c[3]} -> {c[12]}{' (holds this seat)' if c[7] else ''}")
    for line in report:
        say(f"      check: {line}")
    kinds = local["by_kind"]
    say(f"    Wisconsin: {len(local['races']) - len(local['judges'])} county races in {len(local['loaded'])} of 72 counties (sheriff {kinds.get('sheriff', 0)}, "
        f"clerk of circuit court {kinds.get('clerk_of_court', 0)}, coroner {kinds.get('coroner', 0)}), {local['county_cands']} candidates; "
        f"{len(local['judges'])} special elections for municipal judge; {local['files']} clerks' files read; {local['not_loaded']} counties named as gaps")
    for f, kinds_ in sorted(local["loaded"].items()):
        own, lent, n = local["how_read"][f]
        say(f"      {local['counties'][f][1]}: {', '.join(k.replace('_', ' ') for k in kinds_)}; read {n} time{'s' if n != 1 else ''}"
            + (", all alike" if n > 1 else "") + ("" if own else " (from a neighbouring clerk's file)"))
    say(f"      check: {local['ballots_read']} ballots and listing pages with a local contest; {local['lines_read']} candidate lines read on them and "
        f"{len(local['cands'])} candidates stored, each in one race (the same line on every ballot of a county counts once); "
        f"{local['state_contests']} state and federal contests only counted; {local['blanked']} cells set aside as contact-like")
    for p in local["problems"]:
        say(f"      CHECK {p}")
    return len(general)


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--refresh"]
    if len(args) != 1:
        raise SystemExit("usage: python -m ballot.state_local_wi <database> [--refresh]")
    load(args[0], refresh="--refresh" in sys.argv[1:])
